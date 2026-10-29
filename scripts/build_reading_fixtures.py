"""Builds the synthetic fixture set of the sheet reader (docs/design/read-check-review.md section 10) and the expected map of each file.

    python scripts/build_reading_fixtures.py

Writes backend/tests/fixtures/reading/: five files built here, the expected map of each as JSON, and manifest.json (the sixth fixture, the
test company's files, is sample_data/ and is only listed). Nothing is read from anywhere: every figure and name below is made up, the one
company is "Fixture Co". Running it again changes nothing but the bytes of the workbooks; the expected maps are the contract. A changed
expected map needs the owner's approval, recorded in the PR (section 10).

An expected map holds `as_of_month`, the facts the live run compares (`sheets`, `roles`, `figures`) and `file_map`, the ideal reading that the
offline tests run through the masker, the gateway and the cell checks.
"""
import json
from datetime import datetime
from pathlib import Path

from openpyxl import Workbook

OUT = Path(__file__).resolve().parents[1] / "backend" / "tests" / "fixtures" / "reading"
COMPANY = "Fixture Co"
MONTHS = ["Jan", "Feb", "Mar", "Apr", "May", "Jun", "Jul", "Aug", "Sep", "Oct", "Nov", "Dec"]


def ev(kind, text, **over):
    return {"kind": kind, "text": text, "row": None, "column": None, "cell": None, "cells": None, "other_cells": None, "word": None,
            "multiple": None, **over}


def row_map(row, role, metric, cells, **over):
    base = {"row": row, "label_cell": f"A{row}", "role": role, "metric": metric, "basis": "monthly", "actual_plan": "actual", "boundary_col": None,
            "segment": None, "parent_row": None, "unit": "money", "currency": "EUR", "cells": cells, "variance_of": None, "by_column": False,
            "evidence": f"{role} label and numbers across the period columns"}
    return {**base, **over}


def figure(metric, value, sheet, cell, period, basis="monthly", actual_plan="actual", unit="money", **over):
    base = {"metric": metric, "value": float(value), "unit": unit, "currency": "EUR" if unit == "money" else None, "period": period,
            "basis": basis, "actual_plan": actual_plan, "segment": None,
            "source": {"sheet": sheet, "cell": cell, "page": None, "text_layer": None}, "evidence": f"{metric} row, {period} column"}
    return {**base, **over}


def sheet_map(name, header_row, first, last, granularity, rows, last_actual=None, currency="EUR", currency_cell="A1"):
    return {"sheet": name, "header_row": header_row, "label_col": "A", "first_period_col": first, "last_period_col": last, "granularity": granularity,
            "currency": currency, "currency_cell": currency_cell, "last_actual": last_actual, "rows": rows}


def expected(file, as_of, sheets, roles, figures, file_map, **extra):
    return {"file": file, "as_of_month": as_of, "sheets": sheets, "roles": roles, "figures": [
        {k: f[k] for k in ("metric", "value", "period")} | {"sheet": f["source"]["sheet"], "cell": f["source"]["cell"]} for f in figures],
        "file_map": file_map, **extra}


# --- 1. a monthly report: metrics down the rows, months across, actuals end in June, the plan copied down, a cumulative forecast row ---
def monthly_report():
    name = "Monthly report"
    actual = [412000, 418500, 425250, 431100, 440200, 447223.43]
    plan = [410000, 415000, 420000, 430000, 438000, 445000, 452000, 459000, 466000, 473000, 480000, 487000]
    customers = [120, 123, 127, 131, 134, 138]
    revenue = actual + plan[6:]
    cumulative = [sum(revenue[:i + 1]) for i in range(12)]
    wb = Workbook()
    ws = wb.active
    ws.title = name
    ws["A1"] = f"{COMPANY} monthly report (EUR)"
    ws["A3"] = "Metric"
    for i, m in enumerate(MONTHS):
        ws.cell(3, 2 + i, f"{m}-24")
        ws.cell(4, 2 + i, revenue[i])
        ws.cell(5, 2 + i, plan[i])
        ws.cell(7, 2 + i, cumulative[i])
    for i, c in enumerate(customers):
        ws.cell(6, 2 + i, c)
    for r, label in ((4, "Revenue"), (5, "Revenue plan"), (6, "Customers"), (7, "Cumulative revenue forecast")):
        ws.cell(r, 1, label)
    wb.save(OUT / "monthly_report.xlsx")
    last = {"column": "G", "period": "2024-06", "evidence": [
        ev("row_ends", "The Customers row stops at June", row=6, column="G"),
        ev("equals_range", "From July the Revenue row equals the plan row", cells="H4:M4", other_cells="H5:M5")]}
    rows = [row_map(4, "revenue", "revenue", "B4:M4", actual_plan="mixed", boundary_col="G"),
            row_map(5, "forecast", "revenue", "B5:M5", actual_plan="plan"),
            row_map(6, "customers", "customers", "B6:G6", unit="count"),
            row_map(7, "forecast", "revenue", "B7:M7", basis="cumulative_to_date", actual_plan="mixed", boundary_col="G")]
    figures = [figure("revenue", actual[-1], name, "G4", "2024-06"), figure("customers", customers[-1], name, "G6", "2024-06", unit="count")]
    fmap = {"map_version": 1, "kind": "sheet", "sheets": [sheet_map(name, 3, "B", "M", "monthly", rows, last)], "figures": figures}
    return expected("monthly_report.xlsx", "2024-06",
                    [{"sheet": name, "last_actual_column": "G", "revenue_row": 4, "granularity": "monthly"}],
                    [{"sheet": name, "row": 4, "role": "revenue"}, {"sheet": name, "row": 5, "role": "forecast"},
                     {"sheet": name, "row": 6, "role": "customers"}, {"sheet": name, "row": 7, "role": "forecast"}], figures, fmap)


# --- 2. a yearly P&L, one column per year, the last column a plan with no plan word in its header ---
def yearly_pnl():
    name = "P&L"
    years = ["FY2021", "FY2022", "FY2023", "FY2024", "FY2025"]
    revenue = [2100311, 2650487, 3300125, 4102337, 5200000]
    cost = [610204, 780911, 960430, 1239904, 1400000]
    sm = [402118, 515630, 640771, 880512, 1000000]
    wb = Workbook()
    ws = wb.active
    ws.title = name
    ws["A1"] = f"{COMPANY} profit and loss (EUR)"
    ws["A3"] = "Line"
    for i, y in enumerate(years):
        ws.cell(3, 2 + i, y)
        ws.cell(4, 2 + i, revenue[i])
        ws.cell(5, 2 + i, cost[i])
        ws.cell(6, 2 + i, sm[i])
    for r, label in ((4, "Revenue"), (5, "Cost of revenue"), (6, "Sales and marketing"), (8, "Operating plan"),
                     (9, "Plan: revenue"), (10, "Plan: cost of revenue"), (11, "Plan: sales and marketing")):
        ws.cell(r, 1, label)
    ws.cell(9, 6, revenue[-1])
    ws.cell(10, 6, cost[-1])
    ws.cell(11, 6, sm[-1])
    wb.save(OUT / "yearly_pnl.xlsx")
    last = {"column": "E", "period": "2024", "evidence": [
        ev("equals_range", "The FY2025 column equals the operating plan", cells="F4:F6", other_cells="F9:F11"),
        ev("round_values", "The FY2025 column is round, the actual years are not", cells="F4:F6", multiple=100000)]}
    rows = [row_map(4, "revenue", "revenue", "B4:F4", basis="yearly", actual_plan="mixed", boundary_col="E"),
            row_map(5, "cost_of_revenue", "cost_of_revenue", "B5:F5", basis="yearly", actual_plan="mixed", boundary_col="E"),
            row_map(6, "sales_marketing", "sm_expense", "B6:F6", basis="yearly", actual_plan="mixed", boundary_col="E"),
            row_map(9, "forecast", "revenue", "F9:F9", basis="yearly", actual_plan="plan"),
            row_map(10, "forecast", "cost_of_revenue", "F10:F10", basis="yearly", actual_plan="plan"),
            row_map(11, "forecast", "sm_expense", "F11:F11", basis="yearly", actual_plan="plan")]
    figures = [figure("revenue", revenue[3], name, "E4", "2024", basis="yearly"),
               figure("revenue", revenue[4], name, "F4", "2025", basis="yearly", actual_plan="plan")]
    fmap = {"map_version": 1, "kind": "sheet", "sheets": [sheet_map(name, 3, "B", "F", "yearly", rows, last)], "figures": figures}
    return expected("yearly_pnl.xlsx", "2024-12", [{"sheet": name, "last_actual_column": "E", "revenue_row": 4, "granularity": "yearly"}],
                    [{"sheet": name, "row": 4, "role": "revenue"}, {"sheet": name, "row": 5, "role": "cost_of_revenue"},
                     {"sheet": name, "row": 6, "role": "sales_marketing"}], figures, fmap)


# --- 3. actual | budget | variance columns per month ---
def actual_budget_variance():
    name = "Actual vs budget"
    budget = {5: [400000, 405000, 410000, 415000, 420000, 425000, 430000, 435000, 440000, 445000, 450000, 455000],
              6: [90000, 90000, 92000, 92000, 94000, 94000, 96000, 96000, 98000, 98000, 100000, 100000],
              7: [60000, 62000, 64000, 66000, 68000, 70000, 72000, 74000, 76000, 78000, 80000, 82000]}
    actual = {5: [401500, 409800, 415200, 421500, 418100, 430250], 6: [91200, 90500, 93100, 92800, 94900, 96000],
              7: [58000, 61200, 66100, 65000, 70300, 71200]}
    wb = Workbook()
    ws = wb.active
    ws.title = name
    ws["A1"] = f"{COMPANY} actual against budget (EUR)"
    for m in range(12):
        col = 2 + 3 * m
        ws.cell(3, col, f"{MONTHS[m]}-24")
        ws.merge_cells(start_row=3, start_column=col, end_row=3, end_column=col + 2)
        for k, head in enumerate(("Actual", "Budget", "Variance")):
            ws.cell(4, col + k, head)
    for r, label in ((5, "Revenue"), (6, "Cost of revenue"), (7, "Sales and marketing")):
        ws.cell(r, 1, label)
        for m in range(12):
            col = 2 + 3 * m
            ws.cell(r, col + 1, budget[r][m])
            if m < 6:
                ws.cell(r, col, actual[r][m])
                ws.cell(r, col + 2, actual[r][m] - budget[r][m])
    wb.save(OUT / "actual_budget_variance.xlsx")
    last = {"column": "Q", "period": "2024-06", "evidence": [
        ev("range_empty", "The July actual column is blank", cells="T5:T7"), ev("range_empty", "The August actual column is blank", cells="W5:W7")]}
    rows = [row_map(5, "revenue", "revenue", "B5:AK5", actual_plan="mixed", boundary_col="Q", evidence="Actual, Budget and Variance per month"),
            row_map(6, "cost_of_revenue", "cost_of_revenue", "B6:AK6", actual_plan="mixed", boundary_col="Q"),
            row_map(7, "sales_marketing", "sm_expense", "B7:AK7", actual_plan="mixed", boundary_col="Q")]
    figures = [figure("revenue", actual[5][5], name, "Q5", "2024-06"),
               figure("revenue", budget[5][5], name, "R5", "2024-06", actual_plan="plan"),
               figure("revenue", actual[5][5] - budget[5][5], name, "S5", "2024-06", actual_plan="variance")]
    fmap = {"map_version": 1, "kind": "sheet", "sheets": [sheet_map(name, 3, "B", "AI", "monthly", rows, last)], "figures": figures}
    return expected("actual_budget_variance.xlsx", "2024-06",
                    [{"sheet": name, "last_actual_column": "Q", "revenue_row": 5, "granularity": "monthly"}],
                    [{"sheet": name, "row": 5, "role": "revenue"}], figures, fmap)


# --- 4. the budget on a separate sheet ---
def budget_separate_sheet():
    actual = [398000, 404500, 411200, 417900, 423300, 431800]
    budget = [400000, 405000, 410000, 415000, 420000, 425000, 430000, 435000, 440000, 445000, 450000, 455000]
    customers = [88, 90, 93, 95, 99, 101]
    wb = Workbook()
    ws = wb.active
    ws.title = "Actuals"
    ws["A1"] = f"{COMPANY} actuals (EUR)"
    ws["A2"] = "Metric"
    plan = wb.create_sheet("Budget")
    plan["A1"] = f"{COMPANY} budget (EUR)"
    plan["A2"] = "Metric"
    for i, m in enumerate(MONTHS):
        plan.cell(2, 2 + i, datetime(2024, i + 1, 1)).number_format = "mmm-yy"
        plan.cell(3, 2 + i, budget[i])
        plan.cell(4, 2 + i, 85 + 3 * i)
    for i in range(6):
        ws.cell(2, 2 + i, datetime(2024, i + 1, 1)).number_format = "mmm-yy"
        ws.cell(3, 2 + i, actual[i])
        ws.cell(4, 2 + i, customers[i])
    for sh in (ws, plan):
        sh["A3"], sh["A4"] = "Revenue", "Customers"
    wb.save(OUT / "budget_separate_sheet.xlsx")
    last = {"column": "G", "period": "2024-06", "evidence": [ev("row_ends", "The Actuals sheet stops at June", row=3, column="G")]}
    a_rows = [row_map(3, "revenue", "revenue", "B3:G3"), row_map(4, "customers", "customers", "B4:G4", unit="count")]
    b_rows = [row_map(3, "forecast", "revenue", "B3:M3", actual_plan="plan"), row_map(4, "forecast", "customers", "B4:M4", unit="count", actual_plan="plan")]
    figures = [figure("revenue", actual[-1], "Actuals", "G3", "2024-06"), figure("revenue", budget[5], "Budget", "G3", "2024-06", actual_plan="plan")]
    fmap = {"map_version": 1, "kind": "sheet", "figures": figures, "sheets": [
        sheet_map("Actuals", 2, "B", "G", "monthly", a_rows, last), sheet_map("Budget", 2, "B", "M", "monthly", b_rows)]}
    return expected("budget_separate_sheet.xlsx", "2024-06",
                    [{"sheet": "Actuals", "last_actual_column": "G", "revenue_row": 3, "granularity": "monthly"},
                     {"sheet": "Budget", "last_actual_column": None, "revenue_row": None, "granularity": "monthly"}],
                    [{"sheet": "Actuals", "row": 3, "role": "revenue"}, {"sheet": "Budget", "row": 3, "role": "forecast"}], figures, fmap)


# --- 5. a deck with rotated chart-axis labels whose parts sum to a headline figure (read by the page reader, PR 3) ---
def deck_rotated_axis():
    from pptx import Presentation
    from pptx.chart.data import CategoryChartData
    from pptx.enum.chart import XL_CHART_TYPE
    from pptx.util import Inches, Pt
    parts = [300.0, 280.0, 320.0, 300.0]
    prs = Presentation()
    title = prs.slides.add_slide(prs.slide_layouts[5])
    title.shapes.title.text = f"{COMPANY} update"
    slide = prs.slides.add_slide(prs.slide_layouts[5])
    slide.shapes.title.text = "Revenue by quarter (EUR thousands)"
    box = slide.shapes.add_textbox(Inches(0.5), Inches(1.4), Inches(8), Inches(0.6))
    box.text_frame.text = "FY2024 revenue: 1,200"
    box.text_frame.paragraphs[0].runs[0].font.size = Pt(24)
    data = CategoryChartData()
    data.categories = ["Quarter 1 2024", "Quarter 2 2024", "Quarter 3 2024", "Quarter 4 2024"]
    data.add_series("Revenue", parts)
    chart = slide.shapes.add_chart(XL_CHART_TYPE.COLUMN_CLUSTERED, Inches(0.5), Inches(2.2), Inches(8), Inches(4.5), data).chart
    chart.has_legend = False
    chart.plots[0].has_data_labels = True
    axis = chart.category_axis._element
    axis.get_or_add_txPr().bodyPr.set("rot", "-2700000")              # the axis labels are rotated
    prs.save(OUT / "deck_rotated_axis.pptx")
    figures = [figure("revenue", 1200, None, None, "2024", basis="yearly") for _ in range(1)]
    figures[0]["source"] = {"sheet": None, "cell": None, "page": 2, "text_layer": True}
    figures[0]["evidence"] = "Headline on page 2: FY2024 revenue 1,200"
    for q, v in enumerate(parts, 1):
        f = figure("revenue", v, None, None, f"2024-Q{q}", basis="quarterly")
        f["source"] = {"sheet": None, "cell": None, "page": 2, "text_layer": True}
        f["evidence"] = f"Chart data label of quarter {q} on page 2"
        figures.append(f)
    fmap = {"map_version": 1, "kind": "deck", "sheets": [], "figures": figures}
    out = expected("deck_rotated_axis.pptx", None, [], [], [], fmap, kind="deck", parts_sum={"headline": 1200.0, "parts": parts, "page": 2},
                   read_by="the page reader (PR 3): not run by the sheet-reading job")
    out["figures"] = [{"metric": f["metric"], "value": f["value"], "period": f["period"], "page": 2} for f in figures]
    return out


def write(name, doc):
    (OUT / f"{name}.expected.json").write_text(json.dumps(doc, indent=1, ensure_ascii=False) + "\n", encoding="utf-8")


def main():
    OUT.mkdir(parents=True, exist_ok=True)
    built = {"monthly_report": monthly_report(), "yearly_pnl": yearly_pnl(), "actual_budget_variance": actual_budget_variance(),
             "budget_separate_sheet": budget_separate_sheet(), "deck_rotated_axis": deck_rotated_axis()}
    for name, doc in built.items():
        write(name, doc)
    cols = [row_map(c, role, metric, f"{L}2:{L}15", label_cell=f"{L}1", basis="monthly", by_column=True,
                    evidence=f"{role}: the column headed as in the first row, one figure per month")
            for c, (L, role, metric) in enumerate((("B", "sales_marketing", "sm_expense"), ("C", "revenue", "revenue"),
                                                   ("D", "cost_of_revenue", "cost_of_revenue")), 2)]
    pnl_figures = [figure("revenue", 69000, "CSV", "C15", "2024-02")]
    pnl_map = {"map_version": 1, "kind": "sheet", "figures": pnl_figures,
               "sheets": [sheet_map("CSV", 1, "A", "A", "monthly", cols, currency=None, currency_cell=None)]}
    pnl = expected("sample_data/pnl.csv", "2024-02", [], [], pnl_figures, pnl_map,
                   as_of_source="rules",
                   note="The test company's P&L has one row per month (periods run down the sheet): its as-of month comes from the rules path "
                        "(as_of.pnl_month). The reader may map it; a month it proposes must be 2024-02 or none.")
    pnl["roles"] = [{"sheet": "CSV", "row": 3, "role": "revenue", "by_column": True}]
    pnl["sheets"] = [{"sheet": "CSV", "last_actual_column": None, "revenue_row": 3, "granularity": "monthly"}]
    revenue = {"file": "sample_data/revenue.csv", "as_of_month": None, "read": False,
               "note": "One row per invoice: a records layout is not a recognised sheet, so it is not sent (reading_input.recognised_sheets)."}
    write("testco_pnl", pnl)
    write("testco_revenue", revenue)
    manifest = [{"name": n, "expected": f"{n}.expected.json", "file": d["file"]}
                for n, d in {**built, "testco_pnl": pnl, "testco_revenue": revenue}.items()]
    (OUT / "manifest.json").write_text(json.dumps(manifest, indent=1) + "\n", encoding="utf-8")
    print(f"wrote {len(manifest)} fixtures to {OUT}")


if __name__ == "__main__":
    main()
