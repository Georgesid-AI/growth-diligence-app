"""The two handoffs of docs/specs/interface-contracts.md: engine output -> LLM gateway, engine output -> xlsx export.

One model (schemas/metrics.py) describes the engine output. These tests run the real engine, pass its output through the
gateway's validator and the exporter, open the workbook, and fail when the engine writes a field the model does not
declare. Each failure case is a deliberate violation of the contract, shown to fail. No network, no MongoDB."""
import asyncio
import copy
import logging
import shutil
import subprocess
import sys
import tempfile
from functools import lru_cache
from html.parser import HTMLParser
from pathlib import Path

import pytest

pytest.importorskip("pandas")
pytest.importorskip("openpyxl")
pytest.importorskip("fastapi")
pytest.importorskip("motor")

BACKEND = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(BACKEND))
sys.path.insert(0, str(BACKEND / "tests"))

import openpyxl  # noqa: E402
import pandas as pd  # noqa: E402
from pydantic import ValidationError  # noqa: E402

import contract_fixtures as cf  # noqa: E402
import demo_data  # noqa: E402
import growth_engine as ge  # noqa: E402
import server  # noqa: E402
import test_llm_gateway as t  # noqa: E402
from app import as_of  # noqa: E402
from app import formatting as fmt  # noqa: E402
from app import run_audit  # noqa: E402
from app.llm import gateway  # noqa: E402
from schemas import metrics  # noqa: E402
from schemas.metrics import ContractError, MetricsPayload  # noqa: E402

SAMPLE = BACKEND.parent / "sample_data"
FILES = {"revenue": "revenue.csv", "pnl": "pnl.csv", "crm": "crm.csv"}
FIXED_SENTENCE = "Narrative could not be generated. The computed metrics below are unaffected."


# ---------------------------------------------------------------------------
# Real engine runs: the sample files, and the demo companies in their sparse variants
# ---------------------------------------------------------------------------
def _config(ccy="EUR", target_arr=1_000_000, target_date="2026-12-31", fx=None, uploads=None):
    """The engine config as the server builds it: the as-of month detected from the files as Map detects it (as_of.detect), never
    by the engine."""
    found = as_of.detect(uploads or {})
    return {"reporting_currency": ccy, "target_arr": target_arr, "target_date": target_date,
            "fx": fx or {ccy: 1.0}, "billing_terms": {}, "default_l": 1,
            "as_of_month": as_of.month_end(found["month"]) if found else None}


@lru_cache(maxsize=None)
def _sample_raw() -> dict:
    uploads = {}
    for dtype, name in FILES.items():
        df, sheet = server.parse_file((SAMPLE / name).read_bytes(), name)
        uploads[dtype] = {"file": name, "sheet": sheet, "columns": list(df.columns), "rows": server.df_to_records(df),
                          "mapping": server.suggest_mapping(dtype, list(df.columns))}
    norm = {d: server.normalize(u["rows"], d, u["mapping"]) for d, u in uploads.items()}
    sources = {d: {"file": u["file"], "sheet": u["sheet"]} for d, u in uploads.items()}
    return ge.compute_all_raw(norm["revenue"], norm["crm"], norm["pnl"], _config(uploads=uploads), sources,
                              files=server.candidate_views(uploads))


@lru_cache(maxsize=None)
def _demo_raws() -> dict:
    out = {}
    for i, spec in enumerate(demo_data.DEMO_AUDITS):
        datasets, meta = demo_data.build(spec)
        fx = {k.upper(): v for k, v in meta.get("fx", {}).items()}
        fx[spec["reporting_currency"].upper()] = 1.0
        sources = {d: {"file": meta[d]["file"], "sheet": meta[d]["sheet"]} for d in datasets}

        def run(spec=spec, datasets=datasets, drop=(), target_date=None, months=None, fx=fx, sources=sources):
            if months:
                datasets, _ = demo_data.build({**spec, "months": months})
            norm = {d: server.normalize(server.df_to_records(df), d, m) for d, (df, m) in datasets.items()}
            for d in drop:
                norm[d] = pd.DataFrame()
            uploads = {d: {"rows": server.df_to_records(df), "mapping": m} for d, (df, m) in datasets.items() if d not in drop}
            cfg = _config(spec["reporting_currency"], spec["target_arr"], target_date or spec["target_date"], fx, uploads)
            return ge.compute_all_raw(norm["revenue"], norm["crm"], norm["pnl"], cfg, sources)

        out[f"demo{i}"] = run()
        out[f"demo{i}-no-pnl"] = run(drop=("pnl",))
        out[f"demo{i}-no-crm"] = run(drop=("crm",))
        out[f"demo{i}-bad-target-date"] = run(target_date="1999-01-01")
        out[f"demo{i}-8-months"] = run(months=8)
        out[f"demo{i}-14-months"] = run(months=14)
    return out


@lru_cache(maxsize=None)
def _yearly_raw() -> dict:
    """The sample revenue and CRM files with a yearly P&L read as a statement (chat-upload.md section 26.8): no monthly P&L."""
    from app import column_rules as cr
    from app.structures import sheet as sh
    uploads = {}
    for dtype in ("revenue", "crm"):
        df, sheet = server.parse_file((SAMPLE / FILES[dtype]).read_bytes(), FILES[dtype])
        uploads[dtype] = {"file": FILES[dtype], "sheet": sheet, "columns": list(df.columns), "rows": server.df_to_records(df),
                          "mapping": server.suggest_mapping(dtype, list(df.columns))}
    norm = {d: server.normalize(u["rows"], d, u["mapping"]) for d, u in uploads.items()}
    name = "yearly_pnl_by_product.csv"
    s = cr.read_workbook((BACKEND / "tests" / "fixtures" / "mapping" / name).read_bytes(), name)[0]
    cfg = _config(uploads=uploads)
    fed = sh.feed(sh.build(sh.read(s.grid, s.fmt, 12)), 12, 1000, {"file": name, "sheet": "CSV"}, cfg["as_of_month"])
    yearly = fed["yearly"]
    sources = {d: {"file": u["file"], "sheet": u["sheet"]} for d, u in uploads.items()}
    return ge.compute_all_raw(norm["revenue"], norm["crm"], pd.DataFrame(), cfg, sources,
                              files=server.candidate_views(uploads), yearly=yearly, statement=fed["lines"])


def all_raws() -> dict:
    return {"sample": _sample_raw(), "sample-yearly": _yearly_raw(), **_demo_raws()}


RUNS = ["sample", "sample-yearly"] + [f"demo{i}{v}" for i in range(len(demo_data.DEMO_AUDITS))
                     for v in ("", "-no-pnl", "-no-crm", "-bad-target-date", "-8-months", "-14-months")]


def stored_form(raw: dict) -> dict:
    """What the server stores: the contract's build of the engine output, after sanitize."""
    return server.sanitize(metrics.build(raw))


def undeclared(raw: dict) -> list:
    """Key paths the engine wrote that MetricsPayload does not declare."""
    try:
        MetricsPayload.model_validate({"contract_version": metrics.CONTRACT_VERSION, **raw}, context={"conform": True})
    except ValidationError as exc:
        return sorted(".".join(str(p) for p in e["loc"]) for e in exc.errors() if e["type"] == "extra_forbidden")
    return []


# ---------------------------------------------------------------------------
# 1. The schema covers the engine
# ---------------------------------------------------------------------------
@pytest.mark.parametrize("run", RUNS)
def test_every_engine_field_is_declared_in_the_schema_and_the_output_validates(run):
    raw = all_raws()[run]
    assert undeclared(raw) == [], "the engine writes fields that backend/schemas/metrics.py does not declare"
    stored = stored_form(raw)
    assert metrics.MetricsPayload.model_validate(stored)                      # the stored form validates strictly, no rounding
    assert stored["contract_version"] == metrics.CONTRACT_VERSION


def test_the_coverage_check_fails_when_the_engine_gains_a_field():
    """A deliberate violation: one new field in one block. If this passed, the check above would prove nothing."""
    raw = copy.deepcopy(_sample_raw())
    raw["nrr"]["brand_new_field"] = 1.0
    raw["segment_paths"]["stage_one"]["segments"]["Enterprise"]["another_new_field"] = 3
    assert undeclared(raw) == ["nrr.brand_new_field", "segment_paths.stage_one.segments.Enterprise.another_new_field"]


def test_the_sample_run_exercises_the_blocks_the_contract_exists_for():
    stored = stored_form(_sample_raw())
    for block in ("arr", "nrr", "gross_churn", "cac_payback", "sales_cycle", "win_rate", "acv_path", "segment_paths",
                  "revenue_series", "customers_series", "cohort_retention", "revenue_reconciliation", "won_deal_closes"):
        assert stored[block], f"{block} is empty in the sample run: the contract test would not see it"


CITED_BLOCKS = ("arr", "nrr", "gross_churn", "cac_payback", "sales_cycle", "win_rate", "acv_path", "segment_paths",
                "anomalies", "mrr_series", "revenue_series", "customers_series", "cohort_retention",
                "won_deal_closes",                                  # docs/specs/crm-dedup-clustering.md: the won deals' rows
                "yearly_basis",                                     # chat-upload.md section 26.8: the statement's line rows
                "pnl_lines")                                        # docs/design/pnl-lines.md: the rows of the findings and plan cells
CITATION_BESIDE_BLOCK = {"new_mrr_by_quarter": "new_mrr_by_quarter_source"}      # a dict keyed by quarter has no room for a key
NO_OWN_CITATION = {"revenue_reconciliation": "its months and its window carry both sides' citations",
                   "founder_win_rate": "optional twin of win_rate", "missing_data": "gaps, no figure",
                   "questions_for_management": "text, no figure", "contract_version": "", "reporting_currency": "",
                   "as_of_month": "",
                   # docs/specs/dates-and-periods.md sections 2-3: months as text, read from the uploads, no figure
                   "source_period": "months as text, no figure", "metric_periods": "months as text, no figure",
                   # docs/specs/revenue-file-recognition.md section 9: closed words and counts about how the revenue was read
                   "recognition": "closed words and counts; a sheet with two tables that differ adds, per month, both figures and their cell positions"}


def test_every_block_with_a_figure_carries_a_citation_in_the_model_and_in_the_engine_output():
    """A block added without a citation, or listed nowhere here, fails: the decision is made when the block is added."""
    fields = set(MetricsPayload.model_fields)
    decided = set(CITED_BLOCKS) | set(CITATION_BESIDE_BLOCK) | set(CITATION_BESIDE_BLOCK.values()) | set(NO_OWN_CITATION)
    assert fields == decided, f"decide whether these blocks cite a source: {sorted(fields ^ decided)}"
    assert stored_form(all_raws()["sample-yearly"])["yearly_basis"]["years"], "fixture: the yearly run holds the block"
    assert stored_form(all_raws()["sample-yearly"])["pnl_lines"]["plan"], "fixture: the yearly run holds the statement block"
    for run in ("sample", "sample-yearly", "demo0-no-pnl", "demo1-no-crm", "demo1-14-months"):
        stored = stored_form(all_raws()[run])
        for block in CITED_BLOCKS:
            if stored.get(block) is None:                          # a failed calculation is Missing, not cited; a block not set
                continue
            assert stored[block]["source"]["rule"] and stored[block]["source"]["rows"], f"{run}: {block} has no citation"
        assert stored["new_mrr_by_quarter_source"]["rule"], run
    sample = stored_form(_sample_raw())
    for block in ("segment_paths", "mrr_series", "cohort_retention", "anomalies"):
        assert sample[block]["source"]["row_numbers"] and sample[block]["source"]["file"] == "revenue.csv", block
    assert sample["new_mrr_by_quarter_source"]["row_numbers"]
    assert sample["anomalies"]["deals_source"]["file"] == "crm.csv"
    assert "deals_source" not in stored_form(all_raws()["demo0-no-crm"])["anomalies"], "no CRM, no CRM citation"


def test_the_engine_rounds_counts_and_days_once_and_percents_are_fractions():
    s = stored_form(_sample_raw())
    assert isinstance(s["sales_cycle"]["median_days"], float), "days stay unrounded; the display and the export round up"
    assert isinstance(s["acv_path"]["total_customers_at_target"], int)
    assert isinstance(s["acv_path"]["required_net_new_per_year"], int)
    assert 0 < s["nrr"]["overall_pct"] < 10 and 0 <= s["win_rate"]["win_rate_pct"] <= 1
    assert all(0 <= v <= 10 for r in s["cohort_retention"]["data"] for v in r["values"].values())
    assert s["contract_version"] == metrics.CONTRACT_VERSION == 15 and s["reporting_currency"] == "EUR"


@pytest.mark.parametrize("model, value, expected", [
    (metrics.CountUp, 128.3, 129), (metrics.CountUp, 129.00000000001, 129), (metrics.CountUp, -3.2, -3),
    (metrics.Count, 2.5, 3), (metrics.Count, 2.4, 2)])
def test_the_engine_rounds_up_or_to_nearest_by_the_unit_of_the_field(model, value, expected):
    from pydantic import TypeAdapter
    assert TypeAdapter(model).validate_python(value, context={"conform": True}) == expected


@pytest.mark.parametrize("model", [metrics.Count, metrics.CountUp])
def test_a_stored_count_or_day_count_with_a_fraction_is_refused_not_rounded(model):
    from pydantic import TypeAdapter
    with pytest.raises(ValidationError):
        TypeAdapter(model).validate_python(42.5)


# ---------------------------------------------------------------------------
# 2. The gateway's validator and the exporter accept what the engine stores
# ---------------------------------------------------------------------------
@pytest.mark.parametrize("run", RUNS)
def test_the_stored_output_passes_the_gateway_validator_and_the_exporter(run):
    stored = stored_form(all_raws()[run])
    assert gateway.contract_problems(stored) == []
    metrics.validate_for_export(stored)
    buf = server.build_export_workbook({"company_name": "Acme", "as_of_month": stored["as_of_month"]}, stored)
    assert openpyxl.load_workbook(buf).sheetnames[0] == "Headline"


# ---------------------------------------------------------------------------
# 3. The workbook shows the NRR as a whole percent, not "1%"
# ---------------------------------------------------------------------------
def _cell(wb, sheet, label):
    for row in wb[sheet].iter_rows(min_row=2):
        if row[0].value and str(row[0].value).startswith(label):
            return row[1]
    raise AssertionError(f"no row {label!r} on {sheet}")


def _excel_displays(cell) -> str:
    """What Excel shows for a numeric cell formatted "0%": the value times 100, to the nearest whole, with a %."""
    assert cell.number_format == "0%", cell.number_format
    return f"{int(cell.value * 100 + 0.5)}%"


class _Table(HTMLParser):
    def __init__(self):
        super().__init__()
        self.rows, self._row, self._cell = [], None, None

    def handle_starttag(self, tag, attrs):
        if tag == "tr":
            self._row = []
        elif tag == "td" and self._row is not None:
            self._cell = []

    def handle_data(self, data):
        if self._cell is not None:
            self._cell.append(data.strip())

    def handle_endtag(self, tag):
        if tag == "td" and self._cell is not None:
            self._row.append(" ".join(x for x in self._cell if x))
            self._cell = None
        elif tag == "tr" and self._row is not None:
            self.rows.append(self._row)
            self._row = None


def _libreoffice_shows(xlsx: bytes, label: str):
    """The text a spreadsheet application displays in the cell beside `label` on the first sheet (None: none installed)."""
    soffice = shutil.which("soffice") or shutil.which("libreoffice")
    if not soffice:
        return None
    with tempfile.TemporaryDirectory() as tmp:
        (Path(tmp) / "e.xlsx").write_bytes(xlsx)
        subprocess.run([soffice, "--headless", "--norestore", f"-env:UserInstallation=file://{tmp}/profile",
                        "--convert-to", "html", "--outdir", tmp, str(Path(tmp) / "e.xlsx")],
                       check=True, capture_output=True, timeout=120, cwd=tmp)
        table = _Table()
        table.feed((Path(tmp) / "e.html").read_text(encoding="utf-8", errors="ignore"))
    return next(r[1] for r in table.rows if r and r[0] == label)


def test_the_nrr_cell_of_the_sample_run_displays_the_whole_percent_not_one_percent():
    stored = stored_form(_sample_raw())
    nrr = stored["nrr"]["overall_pct"]
    assert nrr == pytest.approx(1.1268)
    buf = server.build_export_workbook({"company_name": "Acme"}, stored)
    cell = _cell(openpyxl.load_workbook(buf), "Headline", "NRR overall")
    assert cell.value == pytest.approx(nrr), "the cell holds the fraction"
    assert _excel_displays(cell) == "113%" and _excel_displays(cell) != "1%"
    shown = _libreoffice_shows(buf.getvalue(), "NRR overall")
    if shown is not None:
        assert shown == "113%", f"a spreadsheet application shows {shown!r}"


def test_a_stored_nrr_of_1_0641_displays_106_percent():
    stored = cf.stored(nrr=cf.nrr(1.0641, 100), arr=cf.arr())
    buf = server.build_export_workbook({"company_name": "Acme"}, stored)
    cell = _cell(openpyxl.load_workbook(buf), "Headline", "NRR overall")
    assert _excel_displays(cell) == "106%"
    shown = _libreoffice_shows(buf.getvalue(), "NRR overall")
    if shown is not None:
        assert shown == "106%"


# ---------------------------------------------------------------------------
# 4. Violations: each one fails the gateway validator and the exporter (shown to fail)
# ---------------------------------------------------------------------------
def _sample_stored() -> dict:
    return stored_form(_sample_raw())


def _set(path, value):
    def mutate(p):
        node = p
        *head, last = path.split(".")
        for part in head:
            node = node[int(part)] if isinstance(node, list) else node[part]
        node[last] = value
        return p
    return mutate


def _drop(path):
    def mutate(p):
        node = p
        *head, last = path.split(".")
        for part in head:
            node = node[part]
        del node[last]
        return p
    return mutate


VIOLATIONS = {
    "nrr as a whole-number percent": _set("nrr.overall_pct", 112.68),
    "churn as a whole-number percent, above the range": _set("gross_churn.overall_pct", 35.0),
    "a fractional customer count": _set("acv_path.current_customers", 24.5),
    "a fractional required count": _set("acv_path.total_customers_at_target", 24.7),
    "a count as text": _set("nrr.nrr_base_customers", "5"),
    "a missing contract version": _drop("contract_version"),
    "a missing field": _drop("nrr.series"),
    "gross churn without its monthly series": _drop("gross_churn.monthly_series"),
    "a monthly churn rate as a whole-number percent": _set("gross_churn.monthly_series", [{"month": "2026-05", "churn_pct": 35.0}]),
    "an undeclared field": _set("arr.extra", 1),
    "a percent as text": _set("win_rate.win_rate_pct", "40%"),
    "a number that is not finite": _set("arr.value", float("nan")),
    "segment_paths without a citation": _drop("segment_paths.source"),
    "new_mrr_by_quarter without a citation": _drop("new_mrr_by_quarter_source"),
    "cohort_retention without a citation": _drop("cohort_retention.source"),
    "mrr_series without a citation": _drop("mrr_series.source"),
    "anomalies without a citation": _drop("anomalies.source"),
    "a citation without its rule": _drop("mrr_series.source.rule"),
    # docs/specs/revenue-file-recognition.md section 9: closed words, never a string from a file
    "a revenue series basis that is not one of the five": _set("revenue_series.basis", "guess"),
    "a revenue series file type that is a file name": _set("revenue_series.file_type", "ledger.xlsx"),
    "a growth label that is free text": _set("revenue_series.growth_label", "Revenue growth of Acme"),
    "a reconciliation status that is not on the list": _set("revenue_reconciliation.status", "maybe"),
    "a recognition block holding a file name": _set("recognition", {"file_type": "ledger.xlsx", "series_basis": "row", "periods": 12,
                                                                   "excluded_columns": 0, "currency_confirmed": True, "source": "pnl"}),
    "a recognition block with a text field": _set("recognition", {"file_type": "monthly_revenue_report", "series_basis": "row", "periods": 12,
                                                                 "excluded_columns": 0, "currency_confirmed": True, "source": "pnl",
                                                                 "label": "Recognised revenue"}),
    "a recognition gap whose source is text from the file": _set("recognition", {
        "file_type": "monthly_revenue_report", "series_basis": "column", "periods": 6, "excluded_columns": 0, "currency_confirmed": True,
        "source": "revenue_file", "series_source": "dated_list", "cross_check": "mismatch", "cross_check_months": 6,
        "cross_check_gaps": [{"month": "2025-02", "row_value": 1.0, "list_value": 2.0, "row_cell": "Recognised revenue", "list_cells": "J4"}]}),
    "a recognition gap with a month that is not a month": _set("recognition", {
        "file_type": "monthly_revenue_report", "series_basis": "column", "periods": 6, "excluded_columns": 0, "currency_confirmed": True,
        "source": "revenue_file", "series_source": "dated_list", "cross_check": "mismatch", "cross_check_months": 6,
        "cross_check_gaps": [{"month": "February", "row_value": 1.0, "list_value": 2.0, "row_cell": "C3", "list_cells": "J4"}]}),
}


@pytest.mark.parametrize("name", list(VIOLATIONS))
def test_a_violation_fails_the_gateway_validator_and_the_exporter(name):
    good = _sample_stored()
    assert gateway.contract_problems(good) == [] and metrics.validate_for_export(good)
    bad = VIOLATIONS[name](copy.deepcopy(good))
    assert gateway.contract_problems(bad), f"the gateway accepts {name}"
    with pytest.raises(ContractError):
        server.build_export_workbook({"company_name": "Acme"}, bad)


def test_a_unit_slip_inside_the_range_is_what_the_contract_version_is_for():
    """A whole-number percent between 0 and 10 (churn 6.1 for 6.1%) cannot be told from a fraction by its size. The stored
    payload of the old engine has no contract_version, so it is refused whatever its values are."""
    old = copy.deepcopy(_sample_stored())
    del old["contract_version"]
    old["gross_churn"]["overall_pct"] = 6.1
    assert gateway.contract_problems(old)
    with pytest.raises(ContractError):
        metrics.validate_for_export(old)


@pytest.mark.parametrize("path, value, ok", [
    ("gross_churn.overall_pct", 10.0, True), ("gross_churn.overall_pct", 10.01, False), ("gross_churn.overall_pct", -0.01, False),
    ("nrr.overall_pct", -0.5, True), ("nrr.overall_pct", -10.5, False),                       # a credit note can sink NRR
    ("cac_payback.quarters.2023-Q1.gross_margin_pct", -0.3, True),                              # a loss-making quarter
    ("cac_payback.quarters.2023-Q1.gross_margin_pct", 71.2, False),
    ("revenue_reconciliation.gap_pct", 35.0, True)])                                            # a gap against the P&L is unbounded
def test_the_unit_check_bounds_each_percent_by_what_it_measures(path, value, ok):
    stored = _sample_stored()
    node = stored
    *head, last = path.split(".")
    for part in head:
        node = node[part]
    node[last] = value
    assert gateway.contract_problems(stored) == [] if ok else gateway.contract_problems(stored)


def test_the_export_writes_no_file_on_a_violation_and_says_which_field():
    bad = VIOLATIONS["nrr as a whole-number percent"](copy.deepcopy(_sample_stored()))
    with pytest.raises(ContractError) as exc:
        server.build_export_workbook({"company_name": "Acme"}, bad)
    assert "nrr.overall_pct" in str(exc.value) and "112.68" in str(exc.value)
    assert "112.68" not in exc.value.log_text and "overall_pct" in exc.value.log_text, "the log text names fields, no value"
    with pytest.raises(fmt.UnitError):
        fmt.xlsx_value(fmt.PCT, 112.68)           # and no path around the model writes the cell


# ---------------------------------------------------------------------------
# 5. The gateway: a bad payload makes no model call and the metrics still come back
# ---------------------------------------------------------------------------
def _db_with(results):
    db = t.make_db()
    db["audits"].docs[0]["results"] = results
    return db


def _generate(db):
    adapter = t.FakeAdapter()
    result = asyncio.run(gateway.generate_narrative(db, t.RUN_ID, "growth_engine", adapter=adapter, sleep=t._noop_sleep))
    return result, adapter


def test_the_gateway_does_not_call_the_model_on_a_payload_that_breaks_the_contract():
    good_db = _db_with(cf.stored(arr=cf.arr(), nrr=cf.nrr(1.042), gross_churn=cf.gross_churn(0.061), win_rate=cf.win_rate()))
    ok, adapter = _generate(good_db)
    assert ok.narrative_status == "ok" and adapter.calls == 1, "the same payload, unbroken, is read"
    for broken in (cf.broken(good_db["audits"].docs[0]["results"], "nrr.overall_pct", 104.2),
                   cf.broken(good_db["audits"].docs[0]["results"], "contract_version", drop=True),
                   cf.broken(good_db["audits"].docs[0]["results"], "win_rate.won", 26.9)):
        result, adapter = _generate(_db_with(broken))
        assert adapter.calls == 0, "no model call"
        assert (result.narrative_status, result.reason) == ("unavailable", FIXED_SENTENCE)
        assert result.narrative is None


def test_the_metrics_still_come_back_with_their_citation_when_the_narrative_cannot_be_generated():
    stored = cf.broken(cf.stored(arr=cf.arr(), nrr=cf.nrr(1.042)), "nrr.overall_pct", 104.2)
    result, adapter = _generate(_db_with(stored))
    assert result.metrics["nrr"]["overall_pct"] == 104.2 and result.metrics["nrr"]["source"]["rows"]
    assert result.metrics["arr"]["value"] == 3129600.0 and result.metrics["arr"]["source"]["rule"]
    assert adapter.calls == 0


def test_a_read_of_the_cached_narrative_checks_the_contract_too():
    stored = cf.broken(cf.stored(arr=cf.arr(), nrr=cf.nrr(1.042)), "contract_version", drop=True)
    result = asyncio.run(gateway.read_cached_narrative(_db_with(stored), t.RUN_ID, "growth_engine"))
    assert (result.narrative_status, result.reason) == ("unavailable", FIXED_SENTENCE) and result.metrics["arr"]


def test_the_gateway_logs_the_violation_with_paths_and_types_and_no_value(caplog):
    stored = cf.broken(cf.stored(arr=cf.arr(), nrr=cf.nrr(1.042, by_segment={"Zeta Holdings": cf.nrr_group(1.1)})),
                       "nrr.by_segment.Zeta Holdings.nrr_base_customers", "Zeta Holdings 4242 secret")
    with caplog.at_level(logging.INFO):
        _generate(_db_with(stored))
    line = next(r.getMessage() for r in caplog.records if "metrics contract violated" in r.getMessage())
    assert "nrr.by_segment.*.nrr_base_customers" in line, "the segment's name is not in the path"
    assert "Zeta" not in caplog.text and "4242" not in caplog.text and "secret" not in caplog.text


# The data inventory and data integrity steps (A1, A2) read the same stored payload through the same check.
NEW_STEPS = ("data_inventory", "data_integrity")


@pytest.mark.parametrize("step", NEW_STEPS)
def test_the_new_steps_do_not_call_the_model_on_a_payload_that_breaks_the_contract(step):
    good = cf.stored(arr=cf.arr(), nrr=cf.nrr(1.042))
    for broken in (cf.broken(good, "nrr.overall_pct", 104.2), cf.broken(good, "contract_version", drop=True)):
        adapter = t.FakeAdapter()
        result = asyncio.run(gateway.generate_narrative(_db_with(broken), t.RUN_ID, step, adapter=adapter, sleep=t._noop_sleep))
        assert adapter.calls == 0
        assert (result.narrative_status, result.reason) == ("unavailable", FIXED_SENTENCE)
        read = asyncio.run(gateway.read_cached_narrative(_db_with(broken), t.RUN_ID, step))
        assert (read.narrative_status, read.reason) == ("unavailable", FIXED_SENTENCE)


@pytest.mark.parametrize("run", RUNS)
def test_the_new_steps_build_their_payload_from_every_engine_run_with_only_allowlisted_fields(run):
    stored = stored_form(all_raws()[run])
    assert gateway.contract_problems(stored) == []
    for step in NEW_STEPS:
        computed = {"metrics": gateway._slice_for_step(stored, step)}
        out = gateway.build_outbound(computed, {})
        fields = set()

        def walk(node):
            if isinstance(node, dict):
                for k, v in node.items():
                    fields.add(k)
                    walk(v)
            elif isinstance(node, list):
                for i in node:
                    walk(i)
        walk(out["metrics"])
        assert fields <= gateway.OUTBOUND_FIELDS, sorted(fields - gateway.OUTBOUND_FIELDS)


def test_the_integrity_reconciliation_is_the_engines_own_with_no_second_calculation():
    stored = stored_form(all_raws()["sample"])
    rec = stored["revenue_reconciliation"]
    assert rec and rec["available"], "fixture: the sample run reconciles"
    check = next(c for c in gateway._slice_for_step(stored, "data_integrity")["integrity"] if c["check"] == "Revenue file against P&L")
    assert check["gap_pct"] == fmt.fmt_pct(rec["gap_pct"]) and check["window_over_tolerance"] is bool(rec["blocker"])
    over = [m["month"] for m in rec["by_month"] if m["gap_pct"] is not None and abs(m["gap_pct"]) > rec["tolerance_pct"]]
    assert [m["month"] for m in check["months_over_tolerance"]] == over


def test_the_integrity_smoothness_is_the_integrity_modules_own_with_no_second_calculation():
    """The data-integrity payload reads integrity.smoothness (the Dashboard block's function) over the stored revenue series."""
    from app import integrity
    stored = stored_form(all_raws()["sample"])
    read = integrity.smoothness(stored)
    assert read is not None, "fixture: the sample run has 13 consecutive months"
    check = next(c for c in gateway._slice_for_step(stored, "data_integrity")["integrity"] if c["check"] == "Month-on-month revenue smoothness")
    assert check["status"] == "computed" and check["deviation_points"] == f"{read['deviation']:.1f} points"
    assert (check["first_month"], check["last_month"]) == (read["first"], read["last"])
    block = next(r for r in integrity.assess(stored, 0, [])["results"] if r["check"] == "Smoothness")
    assert block["result"] == check["result"], "the screen and the narrative never disagree"


# ---------------------------------------------------------------------------
# Contract 6: what a deck claim carries (chat-upload.md section 28; review of PR 164, H2)
# ---------------------------------------------------------------------------
def test_contract_6_names_the_deck_claim_fields_and_the_register_row_and_the_csv_carry_them():
    assert metrics.CONTRACT_VERSION == 15 and metrics.DECK_CLAIM_FIELDS == ("basis", "scenario", "flags")
    assert {"basis", "scenario"} <= set(cm.FIELDS), "the register row, the baseline CSV and the memo's Appendix A"
    assert "flags" not in cm.FIELDS, "flags ride the deck list, not the register"
    assert not set(metrics.DECK_CLAIM_FIELDS) & gateway.OUTBOUND_FIELDS, "server-only: none of the three reaches the gateway"


# ---------------------------------------------------------------------------
# 6. Results stored under an older contract are marked stale, never recomputed on a read, never shown as current (George, 2026-10-20)
# ---------------------------------------------------------------------------
OLD_RESULTS = {"reporting_currency": "EUR", "nrr": {"overall_pct": 106.41, "n": 100}}
AUDIT_ID = "audit-1"


@pytest.fixture()
def api(monkeypatch):
    from fastapi.testclient import TestClient
    db = t.FakeDB()
    db["audits"].docs.append({"id": AUDIT_ID, "company_name": "Acme", "reporting_currency": "EUR", "status": "computed",
                              "results": copy.deepcopy(OLD_RESULTS)})
    monkeypatch.setattr(server, "db", db)
    calls = []

    async def recompute(audit_id):
        calls.append(audit_id)
        return cf.stored(nrr=cf.nrr(1.0641, 100), arr=cf.arr())

    monkeypatch.setattr(server, "_run_compute", recompute)
    client = TestClient(server.app, raise_server_exceptions=False)
    client.db, client.recomputed = db, calls
    return client


STALE = "Stale – re-run audit"


@pytest.mark.parametrize("stored_version", [None, 1, 2, 3, 4, 8, 9])      # 8: before the simplified flow (George, 2026-10-25); 9: before the annual payback proxies (forecast-and-payback.md)
def test_results_of_an_older_contract_are_stale_not_recomputed_and_not_shown(api, stored_version):
    """Version 0 (no version) holds whole-number percents, 1 to 4 lack the fields added since: whichever it is, a read never calculates
    and never answers with the figures, and `/results` says "Stale – re-run audit" until Run audit calculates again."""
    stored = cf.stored(nrr=cf.nrr(1.0641, 100))
    api.db["audits"].docs[0]["results"] = {**stored, "contract_version": stored_version} if stored_version else copy.deepcopy(OLD_RESULTS)
    r = api.get(f"/api/audits/{AUDIT_ID}/results")
    assert r.status_code == 409 and r.json()["detail"] == {"code": "results_stale", "message": STALE}
    assert "106.41" not in r.text and "1.0641" not in r.text and api.recomputed == []
    assert api.get(f"/api/audits/{AUDIT_ID}/export").status_code == 409 and api.recomputed == []
    claims = api.get(f"/api/audits/{AUDIT_ID}/claims").json()
    assert (claims["results_stale"], claims["register"]) == (True, []) and api.recomputed == []


def test_the_narrative_endpoints_do_not_recompute_stale_results(api):
    api.get(f"/api/runs/{AUDIT_ID}/narrative/growth_engine")
    api.post(f"/api/runs/{AUDIT_ID}/narrative/growth_engine")
    assert api.recomputed == []


def test_current_results_are_not_stale_and_not_recomputed(api):
    api.db["audits"].docs[0]["results"] = cf.stored(nrr=cf.nrr(1.0641, 100))
    assert api.get(f"/api/audits/{AUDIT_ID}/results").status_code == 200 and api.recomputed == []
    assert api.get(f"/api/audits/{AUDIT_ID}/claims").json()["results_stale"] is False


def test_the_contract_version_alone_decides_and_run_audit_counts_the_engine_stage_as_not_done(api):
    from app import run_audit
    audit = api.db["audits"].docs[0]
    assert run_audit.STALE_LABEL == STALE and run_audit.results_stale(audit) is True
    assert run_audit.engine_stale(audit, []) is True
    audit["results"] = cf.stored(nrr=cf.nrr(1.0641, 100))
    assert (run_audit.results_stale(audit), run_audit.engine_stale({**audit, "metrics_stale": False}, [])) == (False, False)
    assert run_audit.results_stale({"results": None}) is False, "no results: nothing is stale, the audit is not computed"


def test_the_banner_is_one_stale_line_and_no_blockers_while_results_are_stale(api):
    r = api.get(f"/api/audits/{AUDIT_ID}/blockers")
    assert r.status_code == 200 and r.json()["blockers"] == [{"kind": "results_stale", "text": STALE}], "one line, no blockers"
    api.db["audits"].docs[0]["results"] = None
    assert [b["kind"] for b in api.get(f"/api/audits/{AUDIT_ID}/blockers").json()["blockers"]] == ["revenue_file_missing"]


def test_the_contract_declares_the_repeat_flag_and_the_won_deal_counts():
    assert set(metrics.Anomalies.model_fields["deals_duplicate_id"].annotation.model_fields) == {
        "rows", "distinct_ids", "collapsed", "rows_list"}
    assert "won_deal_closes" in MetricsPayload.model_fields
    sample = stored_form(_sample_raw())
    assert sample["anomalies"]["deals_duplicate_id"] == {"rows": 0, "distinct_ids": 0, "collapsed": 0, "rows_list": []}
    assert sample["won_deal_closes"]["source"]["file"] == "crm.csv" and sample["won_deal_closes"]["by_month"]
    assert stored_form(all_raws()["demo1-no-crm"])["won_deal_closes"] is None, "no CRM, no won-deal counts"


def test_the_export_endpoint_answers_with_the_reason_and_no_workbook_on_a_violation(api, caplog):
    api.db["audits"].docs[0]["results"] = cf.broken(cf.stored(nrr=cf.nrr(1.042), arr=cf.arr()), "nrr.overall_pct", 104.2)
    with caplog.at_level(logging.INFO):
        r = api.get(f"/api/audits/{AUDIT_ID}/export")
    assert r.status_code == 500 and "Export blocked, no file written" in r.json()["detail"]
    assert "nrr.overall_pct" in r.json()["detail"] and "spreadsheetml" not in r.headers.get("content-type", "")
    assert "export blocked" in caplog.text and "104.2" not in caplog.text


# ---------------------------------------------------------------------------
# 7. One table of units: the formatter reads the schema
# ---------------------------------------------------------------------------
def test_the_formatters_kind_table_is_the_schemas_with_one_stated_exception():
    by_name, by_parent, ambiguous = metrics.kinds_by_name()
    assert set(ambiguous) == {"total"}, f"a name now means two units: {ambiguous}"
    assert fmt.KIND_BY_KEY == {**by_name, "total": fmt.CURRENCY} and fmt.KIND_BY_PARENT == by_parent
    assert by_name["nrr_pct"] == fmt.PCT and by_name["median_days"] == fmt.DAYS and by_name["current_customers"] == fmt.COUNT


def test_every_numeric_leaf_of_a_stored_payload_has_a_display_kind():
    for run in ("sample", "demo0-no-pnl", "demo1-14-months"):
        stored = {k: v for k, v in stored_form(all_raws()[run]).items() if k != "contract_version"}   # server-only, in no slice
        leaves = {(k, parent) for k, parent, _ in _leaves(stored)}
        missing = sorted(k for k, parent in leaves if not (fmt.KIND_BY_KEY.get(k) or fmt.KIND_BY_PARENT.get(parent))
                         and k not in _SEGMENT_NAMED_LEAVES(stored))
        assert missing == [], run


def _leaves(node, key=None, parent=None):
    if isinstance(node, dict):
        for k, v in node.items():
            yield from _leaves(v, k, key)
    elif isinstance(node, list):
        for v in node:
            yield from _leaves(v, key, parent)
    elif isinstance(node, (int, float)) and not isinstance(node, bool):
        yield key, parent, node


def _SEGMENT_NAMED_LEAVES(stored):
    """A monthly row holds one amount per segment under the segment's own name: data, not a field."""
    return set(stored["mrr_series"]["segments"]) | set(stored["customers_series"]["segments"])


def test_a_new_unit_type_without_a_display_kind_is_caught():
    """The kinds in the schema are the formatter's kinds: a unit the formatter cannot show fails here, not in production."""
    kinds = {u.kind for cls in metrics._all_models() for f in cls.model_fields.values() if (u := metrics._field_unit(f))}
    assert kinds <= set(fmt.XLSX_NUMBER_FORMAT), kinds - set(fmt.XLSX_NUMBER_FORMAT)


def test_sales_cycle_days_are_stored_unrounded_and_rounded_up_by_the_export_and_the_display():
    stored = _sample_stored()
    median = stored["sales_cycle"]["median_days"]
    assert median == 58.5, "claim matching reads this value"
    cell = _cell(openpyxl.load_workbook(server.build_export_workbook({"company_name": "Acme"}, stored)), "Headline", "Median sales cycle")
    assert cell.value == 59 and fmt.fmt_days(median) == "59 days"


# ---------------------------------------------------------------------------
# 8. Verdict and IC memo (docs/specs/verdict-and-memo.md section 9)
# ---------------------------------------------------------------------------
import typing  # noqa: E402

from pydantic import BaseModel  # noqa: E402

import test_ic_memo as icm  # noqa: E402
from app import claim_matching as cm  # noqa: E402
from app import ic_memo  # noqa: E402


def _declared(model, parts) -> bool:
    """Whether a source-key path is a path MetricsPayload declares: `*` stands for a data key, a list is transparent."""
    if not parts:
        return True
    head, *rest = parts
    if head in model.model_fields:
        return _declared_in(model.model_fields[head].annotation, rest)
    return model.model_config.get("extra") == "allow" and head == "*" and not rest      # a segment's own amount


def _declared_in(annotation, rest) -> bool:
    if not rest:
        return True
    origin = typing.get_origin(annotation)
    if origin is typing.Union:                                    # Optional[...]
        return any(_declared_in(arg, rest) for arg in typing.get_args(annotation))
    if origin is typing.Annotated:
        return _declared_in(typing.get_args(annotation)[0], rest)
    if origin is dict:                                            # the key is data: `*`
        return rest[0] == "*" and _declared_in(typing.get_args(annotation)[1], rest[1:])
    if origin is list:                                            # a list is transparent
        return _declared_in(typing.get_args(annotation)[0], rest)
    return isinstance(annotation, type) and issubclass(annotation, BaseModel) and _declared(annotation, rest)


def _pattern(key: str) -> str:
    return key.replace("{seg}", "*").replace("{q}", "*").replace("{l}", "1")


def test_every_source_key_of_the_memos_evidence_table_is_a_path_the_model_declares():
    for metric, (_, whole, by_segment) in cm.EVIDENCE.items():
        for key in filter(None, (whole, by_segment)):
            assert _declared(MetricsPayload, _pattern(key).split(".")), f"{metric}: {key} is not declared by MetricsPayload"
    assert set(cm.EVIDENCE) == set(cm.METRICS) - {cm.VOLUME, cm.CONTRACT_VALUE, *cm.UNIT_ECONOMICS.values()}, \
        "every metric of the register has an evidence row, but transaction volume, contract value, ARPA and ARPU (no engine source)"


def test_the_annual_payback_source_keys_are_paths_the_model_declares():
    """forecast-and-payback.md section 4: an annual CAC payback claim cites the yearly-basis block, not a quarter of the monthly P&L."""
    for key in cm.ANNUAL_KEYS.values():
        assert _declared(MetricsPayload, key.split(".")), f"{key} is not declared by MetricsPayload"
    assert not _declared(MetricsPayload, "yearly_basis.years.cac_payback_gross_new_arr_months".split("."))


def test_a_wrong_source_key_is_caught_by_the_declared_path_check():
    """A deliberate violation: if this passed, the check above would prove nothing."""
    for bad in ("nrr.series.nrr", "mrr_series.data.total.extra", "nrr.by_segment.nrr_pct", "arr.value.x", "gross_churn.series.churn_percent"):
        assert not _declared(MetricsPayload, bad.split(".")), bad
    assert _declared(MetricsPayload, "nrr.by_segment.*.nrr_pct".split("."))


def test_a_memo_is_built_from_the_sample_data_engine_run_and_every_figure_in_it_is_in_the_stored_results():
    stored = copy.deepcopy(_sample_stored())
    meta = {"company_name": "TestCo", "as_of_month": stored["as_of_month"], "results": stored}
    tables = ic_memo.workbook_tables(server.build_export_workbook(meta, stored))          # Appendix C: the export's own tables
    assert {t["sheet"] for t in tables} >= {"Headline", "By Segment", "CAC by Quarter", "Path to Plan", "Segment Base", "Anomalies"}
    assert "Missing Data" not in {t["sheet"] for t in tables}
    kw = icm.inputs(stored=stored, tables=tables)
    text = ic_memo.build_memo(**kw)
    assert "## Appendix D" in text and ic_memo.word_count_of(text) <= 1500
    assert "81,431" in text, "a total the export derives from the stored results is in Appendix C"
    assert f"ARR {fmt.fmt_currency(stored['arr']['value'], 'EUR')}" in text


def test_nrr_written_as_106_41_refuses_the_memo_and_names_the_field():
    stored = _set("nrr.overall_pct", 106.41)(copy.deepcopy(_sample_stored()))
    kw = icm.inputs(stored=stored)
    with pytest.raises(ic_memo.MemoRefused) as exc:
        ic_memo.build_memo(**kw)
    assert exc.value.code == "contract" and "nrr.overall_pct" in exc.value.message and "106.41" not in exc.value.message


@pytest.mark.parametrize("run", RUNS)
def test_a_memo_is_built_from_every_engine_run_the_contract_covers_with_no_false_refusal(run):
    """The number check must not refuse a figure the engine stored or the export's own tables derive: the sample run, both demo
    companies and their sparse variants (no P&L, no CRM, 8 and 14 months, a bad target date)."""
    from claim_matching_support import answer_definition
    from test_claim_matching import RUNS as CLAIM_RUNS, candidates_for, settings
    from app import verdict as vd
    stored = copy.deepcopy(stored_form(all_raws()[run]))
    tables = ic_memo.workbook_tables(server.build_export_workbook({"company_name": "TestCo", "results": stored}, stored))
    base = CLAIM_RUNS["A"]
    cands = candidates_for(base)
    sets = {**settings(base), "as_of_month": None}
    first = cm.build_register(cands, stored, sets)
    top = vd.proposal(first)
    # run A's claims against another company's data: the analyst answers every Definition check (claim-matching.md section 4.1)
    for cid in [r["claim_id"] for r in first if r["definition_check"] == cm.DEFINITION_ASK]:
        answer_definition(cands, first, cid, "yes")
    for c in cands:
        for cid in top:
            if cid.split("#")[0] == c["id"]:
                c.setdefault("claim_inputs", {}).setdefault(cid, {}).update({
                    "gate_threshold": 5.0, "gate_budget_decision": "the plan", "gate_date": "2030-01-01", "gate_metric_name": "Pipeline",
                    "gate_direction": "at least"})
    rows = vd.mark_key_gates(cm.build_register(cands, stored, sets), top)
    state = vd.top5_state(rows, {"claim_ids": top})
    text = ic_memo.build_memo(
        audit={"company_name": "TestCo", "target_arr": 1e6, "target_date": "2026-12-31", "as_of_month": stored["as_of_month"]},
        results=stored, rows=rows, ver=vd.verdict(rows, stored, {"claim_ids": top}), key=vd.key_gates(rows, top),
        gaps=vd.data_gaps(stored, rows, state["in_force"]), ic={"thesis": {}, "note": "Checked against the engine run."},      # the thesis is the computed draft
        blockers=[], narratives=[], files=icm.FILES, decks=icm.DECKS, tables=tables, today="2026-10-08")
    assert ic_memo.word_count_of(text) <= 1500 and "## Appendix D" in text


# ---------------------------------------------------------------------------
# 9. Evidence status in the memo (docs/specs/evidence-labels.md section 5): the memo's summary counts Inconclusive under
# Unverified with the Inconclusive count in brackets, and frontend/src/lib/evidence_labels.json (the shared list) says the same.
# ---------------------------------------------------------------------------
import json  # noqa: E402

SHARED_LABELS = Path(__file__).resolve().parents[2] / "frontend" / "src" / "lib" / "evidence_labels.json"


def test_the_memo_summary_counts_inconclusive_under_unverified_with_the_count_in_brackets():
    shared = json.loads(SHARED_LABELS.read_text(encoding="utf-8"))
    assert shared["memo_groups"] == {"Unverified": ["Unverified", "Inconclusive"]}
    kw = icm.inputs()
    rows = [dict(r) for r in kw["rows"]]
    next(r for r in rows if r["evidence_label"] == "Unverified")["evidence_label"] = "Inconclusive"
    counts = cm.label_counts(rows)
    text = ic_memo.build_memo(**{**kw, "rows": rows})
    unverified = sum(counts.get(label, 0) for label in shared["memo_groups"]["Unverified"])
    assert f"{unverified} Unverified ({counts['Inconclusive']} Inconclusive) of {len(rows)} in the register" in text
    assert "Unverified (0 Inconclusive)" in ic_memo.build_memo(**icm.inputs()), "the bracket is printed when none is Inconclusive"


def test_get_audit_never_sends_stale_results_to_the_browser(api):
    """Review of 20 Oct (M1): the stored results of an older contract version are dropped from `GET /audits/{id}`, and the audit says
    `results_stale`; current results go out with `results_stale` false."""
    r = api.get(f"/api/audits/{AUDIT_ID}")
    assert r.status_code == 200 and "106.41" not in r.text
    body = r.json()
    assert body["results"] is None and body["results_stale"] is True
    api.db["audits"].docs[0]["results"] = cf.stored(nrr=cf.nrr(1.0641, 100))
    body = api.get(f"/api/audits/{AUDIT_ID}").json()
    assert body["results"]["nrr"] and body["results_stale"] is False
    assert api.put(f"/api/audits/{AUDIT_ID}", json={"company_name": "Acme"}).json()["results_stale"] is False


def test_run_audit_engine_on_a_stale_audit_clears_the_flag_through_to_results(api, monkeypatch):
    """Review of 20 Oct (L2): the stale flag clears when the engine stage runs, and /results answers 200 afterwards."""
    api.db["audits"].docs[0]["status"] = "mapped"
    api.db["datasets"].docs.append({"audit_id": AUDIT_ID, "dtype": "revenue", "mapped_at": "2026-01-01"})

    async def counts(audit_id, audit, datasets=None):
        return {"open": 0, "missing_required": 0}, True

    async def compute(audit_id):
        fresh = cf.stored(nrr=cf.nrr(1.0641, 100))
        api.db["audits"].docs[0]["results"] = fresh
        api.recomputed.append(audit_id)
        return fresh

    monkeypatch.setattr(server, "_mapping_counts", counts)
    monkeypatch.setattr(server, "_run_compute", compute)
    assert api.get(f"/api/audits/{AUDIT_ID}/results").status_code == 409
    api.db["audits"].docs[0]["as_of_month"] = "2024-02-29"
    r = api.post(f"/api/audits/{AUDIT_ID}/run/engine")
    assert r.status_code == 200 and r.json()["status"] == "done" and api.recomputed == [AUDIT_ID]
    assert api.get(f"/api/audits/{AUDIT_ID}/results").status_code == 200
    assert api.get(f"/api/audits/{AUDIT_ID}").json()["results_stale"] is False


def test_run_audit_engine_runs_with_no_as_of_month_and_says_so(api, monkeypatch):
    """read-check-review.md section 6, PR 2: no confirmed month and none typed is "not determined"; Run audit goes on and calculates
    everything that does not depend on it, with a plain note. A typed month removes the note."""
    api.db["audits"].docs[0].update(status="mapped", as_of_month=None, as_of_by="detected")
    api.db["datasets"].docs.append({"audit_id": AUDIT_ID, "dtype": "revenue", "mapped_at": "2026-01-01"})

    async def counts(audit_id, audit, datasets=None):
        return {"open": 0, "missing_required": 0}, True

    async def compute(audit_id):
        api.recomputed.append(audit_id)
        return cf.stored(nrr=cf.nrr(1.0641, 100))

    monkeypatch.setattr(server, "_mapping_counts", counts)
    monkeypatch.setattr(server, "_run_compute", compute)
    out = api.post(f"/api/audits/{AUDIT_ID}/run/engine").json()
    assert (out["status"], out["message"]) == ("done", run_audit.NO_AS_OF) and api.recomputed == [AUDIT_ID]
    assert "not determined" in run_audit.NO_AS_OF
    api.put(f"/api/audits/{AUDIT_ID}", json={"as_of_month": "2024-02-29"})
    typed = api.post(f"/api/audits/{AUDIT_ID}/run/engine").json()
    assert (typed["status"], typed["message"]) == ("done", None) and api.recomputed == [AUDIT_ID, AUDIT_ID]


def test_stale_results_read_stale_re_run_audit_on_the_review_the_verdict_and_the_memo_not_compute_the_audit_first(api):
    """Decision of 20 Oct (George): where results exist but are stale, the Review summary says so, and the verdict and the memo name
    "Stale – re-run audit" instead of "compute the audit first"; with no results at all the old wording stands."""
    summary = api.get(f"/api/audits/{AUDIT_ID}/review").json()
    assert summary["computed"] is True and summary["results_stale"] is True
    verdict = api.get(f"/api/audits/{AUDIT_ID}/verdict").json()["verdict"]
    assert verdict["status"] == "no_verdict" and STALE in verdict["message"] and "compute the audit first" not in verdict["message"]
    memo = api.get(f"/api/audits/{AUDIT_ID}/memo.md")
    assert memo.status_code == 409 and STALE in memo.json()["detail"] and "compute the audit first" not in memo.json()["detail"]
    api.db["audits"].docs[0]["results"] = None
    assert api.get(f"/api/audits/{AUDIT_ID}/review").json()["results_stale"] is False
    assert "compute the audit first" in api.get(f"/api/audits/{AUDIT_ID}/memo.md").json()["detail"]
    assert "compute the audit first" in api.get(f"/api/audits/{AUDIT_ID}/verdict").json()["verdict"]["message"]


def test_the_demo_audits_seeded_before_contract_5_are_reseeded_and_the_seeded_ones_are_current(monkeypatch):
    """Review of 20 Oct (L2): SEED_VERSION 11 (18 Oct) seeded results of contract 4; the version moves with the contract so a deployed
    database reseeds its demo audits instead of opening them as Stale."""
    from app import run_audit
    assert server.SEED_VERSION > 11
    db = t.FakeDB()
    monkeypatch.setattr(server, "db", db)
    import asyncio
    asyncio.run(server.seed_demo())
    assert db["audits"].docs and not any(run_audit.results_stale(a) for a in db["audits"].docs)


# ---------------------------------------------------------------------------
# The file map (docs/design/read-check-review.md section 4, contract 15): one schema for sheets and decks.
# ---------------------------------------------------------------------------
def _map_row(**over):
    row = {"row": 5, "label_cell": "A5", "role": "revenue", "metric": "revenue", "basis": "monthly", "actual_plan": "mixed",
           "boundary_col": "H", "segment": None, "parent_row": None, "unit": "money", "currency": "EUR", "cells": "C5:N5",
           "variance_of": None, "by_column": False, "evidence": "Revenue label, monthly headers, plan copied after H"}
    return {**row, **over}


def _file_map(**over):
    sheet = {"sheet": "Report", "header_row": 3, "label_col": "A", "first_period_col": "C", "last_period_col": "N",
             "granularity": "monthly", "currency": "EUR", "currency_cell": "A1",
             "last_actual": {"column": "H", "period": "2024-06", "evidence": [
                 {"kind": "row_ends", "text": "Customers row stops at H", "row": 9, "column": "H", "cell": None, "cells": None,
                  "other_cells": None, "word": None, "multiple": None}]},
             "rows": [_map_row()]}
    figure = {"metric": "revenue", "value": 447223.43, "unit": "money", "currency": "EUR", "period": "2024-06", "basis": "monthly",
              "actual_plan": "actual", "segment": None, "source": {"sheet": "Report", "cell": "H5", "page": None, "text_layer": None},
              "evidence": "Revenue row, June column"}
    return {"map_version": 1, "kind": "sheet", "sheets": [sheet], "figures": [figure], **over}


def test_the_file_map_is_one_schema_for_sheets_and_decks():
    from schemas.metrics import FileMap
    assert FileMap.model_validate(_file_map()).sheets[0].rows[0].role == "revenue"
    deck = {"map_version": 1, "kind": "deck", "sheets": [], "figures": [
        {"metric": "arr", "value": 1_200_000.0, "unit": "money", "currency": "USD", "period": "2024", "basis": "yearly",
         "actual_plan": "actual", "segment": None, "source": {"sheet": None, "cell": None, "page": 7, "text_layer": True},
         "evidence": "KPI panel"}]}
    assert FileMap.model_validate(deck).figures[0].source.page == 7


@pytest.mark.parametrize("breakage", [
    lambda m: m["sheets"][0]["rows"][0].update(role="turnover"),                     # a role outside the list
    lambda m: m["sheets"][0]["rows"][0].update(metric="ebitda_margin"),               # a metric the engine does not use
    lambda m: m["sheets"][0]["rows"][0].update(actual_plan="forecast"),               # the word is plan
    lambda m: m["sheets"][0]["rows"][0].update(basis="weekly"),
    lambda m: m["sheets"][0]["rows"][0].update(cells="C5-N5"),                        # not a spreadsheet range
    lambda m: m["sheets"][0]["rows"][0].update(unexpected=1),                         # no extra key
    lambda m: m["sheets"][0].update(label_col="1"),
    lambda m: m["sheets"][0]["last_actual"].update(period="June 2024"),
    lambda m: m["sheets"][0]["last_actual"]["evidence"][0].update(kind="vibes"),
    lambda m: m["figures"][0].update(value=float("nan")),
    lambda m: m["figures"][0]["source"].update(cell="5H"),
    lambda m: m.update(kind="workbook"),
    lambda m: m.update(map_version=2),
], ids=["role", "metric", "actual_plan", "basis", "range", "extra key", "column", "period", "evidence kind", "nan", "cell", "kind", "version"])
def test_the_file_map_refuses_what_is_outside_its_closed_words(breakage):
    from schemas.metrics import FileMap
    m = copy.deepcopy(_file_map())
    breakage(m)
    with pytest.raises(ValidationError):
        FileMap.model_validate(m)


def test_the_file_map_roles_are_the_ones_the_design_note_lists():
    from schemas.metrics import MAP_ROLES
    assert MAP_ROLES == ("revenue", "bookings", "contract_value", "arr", "mrr", "churn_rate", "growth_rate", "margin_rate", "customers",
                         "users", "cost_of_revenue", "sales_marketing", "rnd", "ga", "opex", "ebitda", "cash", "headcount", "forecast",
                         "variance", "other")


def test_the_schema_text_sent_for_a_sheet_reading_is_the_file_map_and_nothing_else():
    """The JSON schema a sheet reading carries as text in its request is generated from FileMap, so the two cannot drift: every closed word
    of one is the other's. No strict output schema is sent for the step (George, 2026-10-30: the provider refused the grammar)."""
    from app.llm import schemas as llm_schemas
    from schemas.metrics import MAP_ROLES, FileMap
    sent = llm_schemas.file_map_schema()
    assert llm_schemas.output_schema(llm_schemas.SHEET_READING) is None
    assert sent["additionalProperties"] is False and set(sent["properties"]) == set(FileMap.model_fields), "FileMap's fields are all sent"
    row = sent["$defs"]["MapRow"]
    assert row["properties"]["role"]["enum"] == list(MAP_ROLES)
    assert set(row["required"]) <= set(row["properties"]) and row["additionalProperties"] is False
    assert "title" not in str(sent), "compact: pydantic's titles are dropped, the names, types, enums, patterns and descriptions stay"
