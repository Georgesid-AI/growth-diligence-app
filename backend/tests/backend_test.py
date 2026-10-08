"""Backend API tests for Growth Diligence engine (Phase 1).

This is an integration suite: it needs a live backend to hit over HTTP. Whether
one is running is decided by actually connecting to it, not by inferring it from
configuration — a `.env` file says where the backend would be, never that it is
up. When the connection fails the whole module skips instead of erroring.
"""
import os
import io
import time
import pytest
import requests

# Every audit names its client (llm-structure-reading.md section 4). Consent is off so an
# integration run never sends anything to a model provider.
ENGAGEMENT = {"client_name": "TEST_Client", "structure_reading_consent": False}

# Where the backend would be, if it is running. Configuration only — reachability
# is established by `require_live_backend` below.
DEFAULT_BACKEND_URL = "http://localhost:8001"

# The probe is deliberately short: a backend that is up answers a local health
# check in milliseconds, so a full second already means "nothing is listening".
PROBE_TIMEOUT_SECONDS = 1.0

# CI sets REQUIRE_LIVE_BACKEND=1. There, a missing backend is a broken pipeline,
# not a reason to pass green — so the module fails instead of skipping, and the
# probe retries first to absorb a server that is still starting up.
REQUIRE_LIVE_BACKEND_ENV = "REQUIRE_LIVE_BACKEND"
STRICT_PROBE_DEADLINE_SECONDS = 15.0
STRICT_PROBE_INTERVAL_SECONDS = 0.5


def _resolve_backend_url():
    url = os.environ.get("REACT_APP_BACKEND_URL")
    if url:
        return url.strip()
    try:
        with open("/app/frontend/.env") as f:
            for line in f:
                if line.startswith("REACT_APP_BACKEND_URL="):
                    value = line.split("=", 1)[1].strip()
                    if value:
                        return value
    except OSError:
        pass
    return DEFAULT_BACKEND_URL


BASE_URL = _resolve_backend_url().rstrip("/")
API = f"{BASE_URL}/api"


def _strict_mode() -> bool:
    """True when the caller insists a backend must be there (CI)."""
    return os.environ.get(REQUIRE_LIVE_BACKEND_ENV, "").strip() == "1"


def _probe_once(timeout: float):
    """Return None when the backend answers, else the connection exception.

    Any HTTP response means the server is up — a 404 or 500 is still a live
    backend, and only a connection-level failure counts as "not running".
    """
    try:
        requests.get(f"{API}/", timeout=timeout)
        return None
    except requests.exceptions.RequestException as exc:
        return exc


@pytest.fixture(scope="session", autouse=True)
def require_live_backend():
    """Gate this module on the backend actually answering.

    Autouse and session-scoped, so the probe runs once for the whole file.

    Default (local): one short probe, skip the module if nothing answers.
    REQUIRE_LIVE_BACKEND=1 (CI): retry for up to 15s, then fail the module — a
    pipeline whose backend never came up must go red, not green-with-skips.
    """
    if not _strict_mode():
        exc = _probe_once(PROBE_TIMEOUT_SECONDS)
        if exc is not None:
            pytest.skip(
                f"no backend answering at {BASE_URL} within {PROBE_TIMEOUT_SECONDS:g}s "
                f"({type(exc).__name__}); this integration suite needs a running server. "
                f"Set {REQUIRE_LIVE_BACKEND_ENV}=1 to make this a failure instead."
            )
        return

    deadline = time.monotonic() + STRICT_PROBE_DEADLINE_SECONDS
    attempts = 0
    while True:
        attempts += 1
        exc = _probe_once(PROBE_TIMEOUT_SECONDS)
        if exc is None:
            return
        if time.monotonic() >= deadline:
            pytest.fail(
                f"{REQUIRE_LIVE_BACKEND_ENV}=1 but no backend answered at {BASE_URL} "
                f"after {STRICT_PROBE_DEADLINE_SECONDS:g}s ({attempts} attempts); "
                f"last error {type(exc).__name__}: {exc}",
                pytrace=False,
            )
        time.sleep(STRICT_PROBE_INTERVAL_SECONDS)

SAMPLES = {
    "revenue": "/app/sample_data/revenue.csv",
    "crm": "/app/sample_data/crm.csv",
    "pnl": "/app/sample_data/pnl.csv",
}


@pytest.fixture(scope="module")
def session():
    return requests.Session()


# --- Demo seeded audits ---
def test_list_audits_has_two_demos(session):
    r = session.get(f"{API}/audits", timeout=30)
    assert r.status_code == 200
    data = r.json()
    computed = [a for a in data if a.get("status") == "computed"]
    assert len(computed) >= 2, f"Expected >=2 computed demo audits, got {len(computed)}"


def test_demo_audit_results_endpoint(session):
    r = session.get(f"{API}/audits", timeout=30)
    demos = [a for a in r.json() if a.get("status") == "computed"]
    assert demos
    aid = demos[0]["id"]
    rr = session.get(f"{API}/audits/{aid}/results", timeout=30)
    assert rr.status_code == 200
    body = rr.json()
    assert "audit" in body and "results" in body
    assert "arr" in body["results"]


def test_demo_audit_get_returns_200_with_datasets(session):
    """Regression: GET /api/audits/{id} must return 200 (not 500) for BOTH seeded demo audits with 'datasets'."""
    r = session.get(f"{API}/audits", timeout=30)
    demos = [a for a in r.json() if a.get("status") == "computed"]
    assert len(demos) >= 2, f"expected 2 demo audits, got {len(demos)}"
    for a in demos[:2]:
        rr = session.get(f"{API}/audits/{a['id']}", timeout=30)
        assert rr.status_code == 200, f"demo {a['id']} GET failed: {rr.status_code} {rr.text[:200]}"
        body = rr.json()
        assert "datasets" in body and isinstance(body["datasets"], dict), \
            f"demo {a['id']} missing datasets: keys={list(body.keys())}"


def delete_audit(session, aid):
    """Delete an audit the way the app does: the company name in the body (CLAUDE.md rule 22)."""
    name = session.get(f"{API}/audits/{aid}", timeout=30).json().get("company_name", "")
    return session.request("DELETE", f"{API}/audits/{aid}", json={"confirm": name}, timeout=30)


# --- CRUD ---
@pytest.fixture(scope="module")
def new_audit(session):
    payload = {**ENGAGEMENT, "company_name": "TEST_Acme", "reporting_currency": "EUR",
               "target_arr": 500000, "target_date": "2025-12-31"}
    r = session.post(f"{API}/audits", json=payload, timeout=30)
    assert r.status_code == 200, r.text
    a = r.json()
    assert a["status"] == "draft" and a["company_name"] == "TEST_Acme"
    yield a
    delete_audit(session, a['id'])


def test_get_audit_has_datasets_field(session, new_audit):
    r = session.get(f"{API}/audits/{new_audit['id']}", timeout=30)
    assert r.status_code == 200
    body = r.json()
    assert "datasets" in body and isinstance(body["datasets"], dict)


def test_results_409_when_not_computed(session, new_audit):
    r = session.get(f"{API}/audits/{new_audit['id']}/results", timeout=30)
    assert r.status_code == 409


# --- Upload flow ---
REQUIRED_FIELDS = {
    "revenue": ["customer_id", "invoice_date", "amount", "currency"],
    "crm": ["deal_id", "created_date", "close_date", "stage", "amount"],
    "pnl": ["month", "sm_expense", "revenue", "cost_of_revenue"],
}


@pytest.fixture(scope="module")
def uploaded(session, new_audit):
    result = {}
    for dtype, path in SAMPLES.items():
        with open(path, "rb") as f:
            files = {"file": (os.path.basename(path), f, "text/csv")}
            r = session.post(f"{API}/audits/{new_audit['id']}/datasets/{dtype}/upload",
                             files=files, timeout=60)
        assert r.status_code == 200, f"{dtype} upload: {r.status_code} {r.text}"
        result[dtype] = r.json()
    return result


def test_revenue_upload_mapping(uploaded):
    body = uploaded["revenue"]
    assert body["row_count"] > 0
    sm = body["suggested_mapping"]
    for f in REQUIRED_FIELDS["revenue"]:
        assert sm.get(f), f"revenue missing required mapping: {f} -> {sm}"


def test_crm_upload_mapping(uploaded):
    body = uploaded["crm"]
    sm = body["suggested_mapping"]
    for f in REQUIRED_FIELDS["crm"]:
        assert sm.get(f), f"crm missing required mapping: {f} -> {sm}"


def test_pnl_upload_mapping_cost_of_revenue_correct(uploaded):
    body = uploaded["pnl"]
    sm = body["suggested_mapping"]
    for f in REQUIRED_FIELDS["pnl"]:
        assert sm.get(f), f"pnl missing required mapping: {f} -> {sm}"
    # Critical: cost_of_revenue must map to 'Cost of Revenue', not 'Revenue'
    assert sm["cost_of_revenue"] == "Cost of Revenue", (
        f"pnl cost_of_revenue mismapped to {sm['cost_of_revenue']}")
    assert sm["revenue"] == "Revenue"


# --- Save mapping + compute ---
def test_save_mapping_and_compute(session, new_audit, uploaded):
    aid = new_audit["id"]
    for dtype, body in uploaded.items():
        payload = {"mapping": body["suggested_mapping"], "fx": {}, "billing_terms": {}}
        r = session.put(f"{API}/audits/{aid}/datasets/{dtype}/mapping",
                        json=payload, timeout=30)
        assert r.status_code == 200, r.text

    r = session.post(f"{API}/audits/{aid}/compute", timeout=120)
    assert r.status_code == 200, r.text
    res = r.json()

    # arr, nrr, sales_cycle, win_rate present with values
    assert res["arr"]["value"] > 0
    assert 100 < res["nrr"]["overall_pct"] < 120, f"NRR expected ~112 got {res['nrr']['overall_pct']}"
    assert 40 < res["sales_cycle"]["median_days"] < 90, f"sales cycle expected ~60 got {res['sales_cycle']['median_days']}"
    assert res["win_rate"]["win_rate_pct"] > 0

    # missing_data empty when all mapped
    assert res.get("missing_data") == [], f"missing_data: {res.get('missing_data')}"

    # every metric has source with file/sheet
    for key in ("arr", "nrr", "sales_cycle", "win_rate"):
        m = res[key]
        assert "source" in m and m["source"].get("file") and m["source"].get("sheet"), \
            f"{key} missing source metadata: {m.get('source')}"


# --- Billing terms (NEW feature) ---
@pytest.fixture(scope="module")
def bt_audit(session):
    r = session.post(f"{API}/audits", json={**ENGAGEMENT, "company_name": "TEST_BT", "reporting_currency": "EUR"}, timeout=30)
    a = r.json()
    # Upload sample revenue (no service_start/end columns)
    with open(SAMPLES["revenue"], "rb") as f:
        up = session.post(f"{API}/audits/{a['id']}/datasets/revenue/upload",
                          files={"file": ("revenue.csv", f, "text/csv")}, timeout=60).json()
    yield a, up
    delete_audit(session, a['id'])


def test_revenue_customers_endpoint_no_service_dates(session, bt_audit):
    a, up = bt_audit
    r = session.get(f"{API}/audits/{a['id']}/datasets/revenue/customers?customer_col=Customer", timeout=30)
    assert r.status_code == 200
    body = r.json()
    assert body["has_service_dates"] is False
    assert body["billing_terms"] == {}
    assert set(body["customers"]) == {f"CUST-{i}" for i in range(1, 6)}


def test_save_and_get_billing_terms_persistence(session, bt_audit):
    a, up = bt_audit
    payload = {"mapping": up["suggested_mapping"], "fx": {}, "billing_terms": {"CUST-1": "annual"}}
    r = session.put(f"{API}/audits/{a['id']}/datasets/revenue/mapping", json=payload, timeout=30)
    assert r.status_code == 200
    got = session.get(f"{API}/audits/{a['id']}", timeout=30).json()
    assert got["datasets"]["revenue"]["billing_terms"] == {"CUST-1": "annual"}


def _single_row_revenue_csv():
    return b"Customer,Invoice Date,Amount,Currency\nACME,2023-01-01,1200,EUR\n"


@pytest.mark.parametrize("term,expected_months,expected_per_month", [
    ("annual", 12, 100.0),
    ("quarterly", 3, 400.0),
    ("monthly", 1, 1200.0),
])
def test_mrr_spread_by_billing_term(session, term, expected_months, expected_per_month):
    """Core rule: 1200 with billing term should spread evenly."""
    r = session.post(f"{API}/audits", json={**ENGAGEMENT, "company_name": f"TEST_SPREAD_{term}", "reporting_currency": "EUR"}, timeout=30)
    aid = r.json()["id"]
    try:
        files = {"file": ("rev.csv", io.BytesIO(_single_row_revenue_csv()), "text/csv")}
        up = session.post(f"{API}/audits/{aid}/datasets/revenue/upload", files=files, timeout=30).json()
        payload = {"mapping": up["suggested_mapping"], "fx": {}, "billing_terms": {"ACME": term}}
        pr = session.put(f"{API}/audits/{aid}/datasets/revenue/mapping", json=payload, timeout=30)
        assert pr.status_code == 200
        cr = session.post(f"{API}/audits/{aid}/compute", timeout=60)
        assert cr.status_code == 200, cr.text
        res = cr.json()
        series = res["mrr_series"]["data"]
        assert len(series) == expected_months, f"expected {expected_months} months, got {len(series)}: {series}"
        for row in series:
            assert abs(row["total"] - expected_per_month) < 0.01, f"month {row['month']} total={row['total']} expected {expected_per_month}"
    finally:
        delete_audit(session, aid)


def test_mrr_spread_defaults_to_monthly_when_absent(session):
    r = session.post(f"{API}/audits", json={**ENGAGEMENT, "company_name": "TEST_SPREAD_default", "reporting_currency": "EUR"}, timeout=30)
    aid = r.json()["id"]
    try:
        files = {"file": ("rev.csv", io.BytesIO(_single_row_revenue_csv()), "text/csv")}
        up = session.post(f"{API}/audits/{aid}/datasets/revenue/upload", files=files, timeout=30).json()
        payload = {"mapping": up["suggested_mapping"], "fx": {}, "billing_terms": {}}
        session.put(f"{API}/audits/{aid}/datasets/revenue/mapping", json=payload, timeout=30)
        cr = session.post(f"{API}/audits/{aid}/compute", timeout=60).json()
        series = cr["mrr_series"]["data"]
        assert len(series) == 1
        assert abs(series[0]["total"] - 1200.0) < 0.01
    finally:
        delete_audit(session, aid)


# --- ITEM 1: as-of month ---
@pytest.fixture(scope="module")
def asof_computed_audit(session):
    """Audit with 3 CSVs uploaded and computed, for as-of tests."""
    r = session.post(f"{API}/audits", json={**ENGAGEMENT, "company_name": "TEST_AsOf", "reporting_currency": "EUR"}, timeout=30)
    a = r.json()
    aid = a["id"]
    ups = {}
    for dtype, path in SAMPLES.items():
        with open(path, "rb") as f:
            up = session.post(f"{API}/audits/{aid}/datasets/{dtype}/upload",
                              files={"file": (os.path.basename(path), f, "text/csv")}, timeout=60).json()
        ups[dtype] = up
        session.put(f"{API}/audits/{aid}/datasets/{dtype}/mapping",
                    json={"mapping": up["suggested_mapping"], "fx": {}, "billing_terms": {}}, timeout=30)
    yield aid
    delete_audit(session, aid)


def test_asof_defaults_to_last_pnl_month(session, asof_computed_audit):
    """When as_of_month is null it defaults to last P&L month (2024-02 from sample)."""
    aid = asof_computed_audit
    r = session.post(f"{API}/audits/{aid}/compute", timeout=60)
    assert r.status_code == 200, r.text
    res = r.json()
    # sample pnl.csv last month is 2024-02
    assert res.get("as_of_month") == "2024-02", f"expected default 2024-02, got {res.get('as_of_month')}"
    assert res["arr"]["month"] == "2024-02"
    assert res["mrr_series"]["months"][-1] == "2024-02"


def test_asof_truncation_via_put_then_compute(session, asof_computed_audit):
    """PUT as_of_month='2024-06' persists, but since data ends 2024-02, min(as_of, last_data_month) applies.
    We test with a value INSIDE the data range: 2023-08."""
    aid = asof_computed_audit
    # PUT the audit with as_of_month=2023-08
    pu = session.put(f"{API}/audits/{aid}", json={"as_of_month": "2023-08"}, timeout=30)
    assert pu.status_code == 200, pu.text
    got = session.get(f"{API}/audits/{aid}", timeout=30).json()
    assert got.get("as_of_month") == "2023-08"
    # Recompute
    cr = session.post(f"{API}/audits/{aid}/compute", timeout=60)
    assert cr.status_code == 200, cr.text
    res = cr.json()
    assert res.get("as_of_month") == "2023-08", f"results.as_of_month={res.get('as_of_month')}"
    assert res["mrr_series"]["months"][-1] == "2023-08", f"last mrr month={res['mrr_series']['months'][-1]}"
    assert res["arr"]["month"] == "2023-08", f"arr.month={res['arr']['month']}"


# --- ITEM 2: win rate excludes invalid deals ---
def test_win_rate_excludes_close_before_created(session):
    """Win rate must exclude deals where close_date < created_date. Craft 3W/2L + 2 invalid = 60%."""
    csv = (
        "Deal ID,Created,Close Date,Stage,Amount\n"
        "W1,2023-01-01,2023-03-01,won,1000\n"
        "W2,2023-01-01,2023-04-01,won,1000\n"
        "W3,2023-01-01,2023-05-01,won,1000\n"
        "L1,2023-01-01,2023-03-01,lost,1000\n"
        "L2,2023-01-01,2023-04-01,lost,1000\n"
        "BAD1,2023-06-01,2023-01-01,won,1000\n"  # close < created - should be excluded
        "BAD2,2023-06-01,2023-02-01,lost,1000\n"  # close < created - should be excluded
    ).encode()
    r = session.post(f"{API}/audits", json={**ENGAGEMENT, "company_name": "TEST_WR", "reporting_currency": "EUR"}, timeout=30)
    aid = r.json()["id"]
    try:
        # Minimal revenue so compute runs
        rev = b"Customer,Invoice Date,Amount,Currency\nA,2023-01-01,100,EUR\n"
        up_rev = session.post(f"{API}/audits/{aid}/datasets/revenue/upload",
                              files={"file": ("r.csv", io.BytesIO(rev), "text/csv")}, timeout=30).json()
        session.put(f"{API}/audits/{aid}/datasets/revenue/mapping",
                    json={"mapping": up_rev["suggested_mapping"], "fx": {}, "billing_terms": {}}, timeout=30)
        up = session.post(f"{API}/audits/{aid}/datasets/crm/upload",
                          files={"file": ("crm.csv", io.BytesIO(csv), "text/csv")}, timeout=30).json()
        session.put(f"{API}/audits/{aid}/datasets/crm/mapping",
                    json={"mapping": up["suggested_mapping"], "fx": {}, "billing_terms": {}}, timeout=30)
        cr = session.post(f"{API}/audits/{aid}/compute", timeout=60)
        assert cr.status_code == 200, cr.text
        res = cr.json()
        wr = res["win_rate"]
        assert wr["won"] == 3, f"won={wr['won']}"
        assert wr["lost"] == 2, f"lost={wr['lost']}"
        assert abs(wr["win_rate_pct"] - 60.0) < 0.01, f"win_rate_pct={wr['win_rate_pct']}"
        assert wr.get("excluded_invalid") == 2, f"excluded_invalid={wr.get('excluded_invalid')}"
    finally:
        delete_audit(session, aid)


# --- ITEM 5: export xlsx ---
def test_export_returns_xlsx_with_sheets(session, asof_computed_audit):
    aid = asof_computed_audit
    # Ensure computed
    session.post(f"{API}/audits/{aid}/compute", timeout=60)
    r = session.get(f"{API}/audits/{aid}/export", timeout=60)
    assert r.status_code == 200, r.text[:200]
    ct = r.headers.get("content-type", "")
    assert "spreadsheetml.sheet" in ct, f"content-type={ct}"
    assert len(r.content) > 1000
    # inspect sheets
    from openpyxl import load_workbook
    wb = load_workbook(io.BytesIO(r.content), read_only=True)
    expected = {"Headline", "By Segment", "NRR by Cohort", "NRR Series",
                "CAC by Quarter", "Path to Plan", "Anomalies", "Missing Data"}
    got = set(wb.sheetnames)
    assert expected.issubset(got), f"missing sheets: {expected - got} (got {got})"
    # Headline includes as-of month
    headline = wb["Headline"]
    text = " ".join(str(c.value) for row in headline.iter_rows() for c in row if c.value is not None)
    assert "as" in text.lower() and ("2024" in text or "2023" in text), f"headline text={text[:400]}"


def test_export_409_when_not_computed(session):
    r = session.post(f"{API}/audits", json={**ENGAGEMENT, "company_name": "TEST_ExpNC", "reporting_currency": "EUR"}, timeout=30)
    aid = r.json()["id"]
    try:
        rr = session.get(f"{API}/audits/{aid}/export", timeout=30)
        assert rr.status_code == 409, f"expected 409 got {rr.status_code}"
    finally:
        delete_audit(session, aid)


# --- Target date validation ---
def test_create_audit_rejects_bad_target_date_year(session):
    r = session.post(f"{API}/audits", json={**ENGAGEMENT, "company_name": "TEST_BadDate", "target_date": "0027-01-01"}, timeout=30)
    assert r.status_code == 422, r.text


def test_create_audit_accepts_valid_target_date(session):
    r = session.post(f"{API}/audits", json={**ENGAGEMENT, "company_name": "TEST_GoodDate", "target_date": "2027-01-01"}, timeout=30)
    assert r.status_code == 200, r.text
    delete_audit(session, r.json()['id'])


# --- Stale metrics: setup changes must recompute automatically ---
def test_fx_change_triggers_recompute(session):
    """Adding an FX rate to an already-computed audit must refresh results —
    the v1/v2 staleness this regression guards against."""
    r = session.post(f"{API}/audits", json={**ENGAGEMENT, "company_name": "TEST_Stale", "reporting_currency": "EUR"}, timeout=30)
    aid = r.json()["id"]
    try:
        csv = b"Customer,Invoice Date,Amount,Currency\nA,2023-01-01,1000,EUR\nB,2023-01-01,500,USD\n"
        up = session.post(f"{API}/audits/{aid}/datasets/revenue/upload",
                          files={"file": ("r.csv", io.BytesIO(csv), "text/csv")}, timeout=30).json()
        # v1: no FX rate for USD -> that row is excluded from MRR
        session.put(f"{API}/audits/{aid}/datasets/revenue/mapping",
                    json={"mapping": up["suggested_mapping"], "fx": {}, "billing_terms": {}}, timeout=30)
        v1 = session.post(f"{API}/audits/{aid}/compute", timeout=30).json()
        assert v1["arr"]["mrr"] == 1000.0, f"v1 mrr={v1['arr']['mrr']}"

        # v2: add a USD rate via the SAME endpoint the FX editor calls — no explicit
        # /compute call here. The stored results must refresh on their own.
        session.put(f"{API}/audits/{aid}/datasets/revenue/mapping",
                    json={"mapping": up["suggested_mapping"], "fx": {"USD": 0.9}, "billing_terms": {}}, timeout=30)
        got = session.get(f"{API}/audits/{aid}/results", timeout=30).json()
        assert got["audit"]["metrics_stale"] is False, "metrics_stale should clear after auto-recompute"
        assert got["results"]["arr"]["mrr"] == 1450.0, f"v2 mrr should include the USD row: {got['results']['arr']['mrr']}"
    finally:
        delete_audit(session, aid)


# --- Overall ACV band summary line ---
def test_acv_path_has_overall_band(session, asof_computed_audit):
    aid = asof_computed_audit
    session.post(f"{API}/audits/{aid}/compute", timeout=60)
    res = session.get(f"{API}/audits/{aid}/results", timeout=30).json()["results"]
    assert res.get("acv_path"), "expected acv_path to be computable for this fixture"
    assert "overall_band" in res["acv_path"]
    ob = res["acv_path"]["overall_band"]
    assert ob is None or {"key", "label", "value_label"} <= ob.keys()


# --- Testco demo regression ---
def test_testco_demo_opens_and_has_asof(session):
    testco_id = "5ae07838-3c1f-452f-8fdf-4e290b3c89bd"
    r = session.get(f"{API}/audits/{testco_id}", timeout=30)
    assert r.status_code == 200, f"Testco GET failed: {r.status_code} {r.text[:200]}"
    rr = session.get(f"{API}/audits/{testco_id}/results", timeout=30)
    assert rr.status_code == 200
    res = rr.json()["results"]
    assert "as_of_month" in res, "results missing as_of_month"


# --- Delete ---
def test_delete_audit_cleans_datasets(session):
    r = session.post(f"{API}/audits", json={**ENGAGEMENT, "company_name": "TEST_Del", "reporting_currency": "EUR"}, timeout=30)
    aid = r.json()["id"]
    with open(SAMPLES["revenue"], "rb") as f:
        session.post(f"{API}/audits/{aid}/datasets/revenue/upload",
                     files={"file": ("revenue.csv", f, "text/csv")}, timeout=30)
    d = delete_audit(session, aid)
    assert d.status_code == 200
    g = session.get(f"{API}/audits/{aid}", timeout=30)
    assert g.status_code == 404
