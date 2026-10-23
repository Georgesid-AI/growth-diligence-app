"""Monthly gross revenue churn (George, 2026-10-23, decision on PR 171; engine contract 8; docs/specs/claim-matching.md, interface-contracts.md).
The engine holds a monthly rate per month beside the trailing-12-month one: MRR lost in the month (churn and contraction, expansion not netted)
over the MRR at the start of the month, the Possibly partial last month skipped as for NRR. A churn claim worded monthly is tested against that
month's rate; a year period against the trailing-12-month rate at year end; no wording and no period is Unverified naming both."""
import copy
import io
import sys
from pathlib import Path

import openpyxl
import pandas as pd
import pytest

BACKEND = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(BACKEND))
sys.path.insert(0, str(BACKEND / "tests"))

import contract_fixtures as cf  # noqa: E402
import growth_engine as ge  # noqa: E402
import server  # noqa: E402
from test_claim_matching import RUNS, results_for, run_claim  # noqa: E402

from app import claim_matching as cm  # noqa: E402
from app import narrative_export as ne  # noqa: E402
from schemas.metrics import MetricsPayload  # noqa: E402

MONTHLY = "Monthly gross revenue churn"


def mrr_table(rows, start="2023-01"):
    """customers x months: rows is {customer: [mrr per month]}."""
    n = len(next(iter(rows.values())))
    return pd.DataFrame(rows, index=pd.period_range(start, periods=n, freq="M")).T


def monthly_of(results):
    return {d["month"]: d["churn_pct"] for d in results["gross_churn"]["monthly_series"]}


def with_monthly(rates, trailing=0.10):
    """Run A's results with a monthly series of fixed rates and a trailing-12-month series, as the engine stores them."""
    results = copy.deepcopy(results_for(RUNS["A"]))
    results["gross_churn"]["monthly_series"] = [{"month": m, "churn_pct": v} for m, v in rates.items()]
    results["gross_churn"]["series"] = [{"month": m, "churn_pct": trailing} for m in rates]
    months = sorted(rates)
    results["metric_periods"]["gross_churn_monthly"] = {"first": months[0], "last": months[-1], "missing": []}
    results["metric_periods"]["gross_churn"] = {"first": months[0], "last": months[-1], "missing": []}
    return results


def churn_row(value, period, snippet="", **extra):
    return run_claim({"claim_type": "churn", "value": value, "unit": "%", "currency": None, "target_date": period, "snippet": snippet}, **extra)


# --- the engine ------------------------------------------------------------------------------------------------------

def test_the_monthly_rate_is_mrr_lost_in_the_month_over_mrr_at_its_start():
    months = 14
    mrr = mrr_table({
        "stays": [100.0] * months,
        "leaves": [100.0] * 5 + [0.0] * (months - 5),            # churns in month 6 (index 5)
        "shrinks": [100.0] * 8 + [60.0] * (months - 8),          # contracts by 40 in month 9
        "grows": [50.0] * 3 + [90.0] * (months - 3),             # expansion is not netted against the losses
        "joins": [0.0] * 10 + [70.0] * (months - 10),            # not in the base of the month it joins
    })
    out = ge.compute_gross_churn(mrr)
    got = {d["month"]: d["churn_pct"] for d in out["monthly_series"]}
    assert list(got)[0] == "2023-02" and list(got)[-1] == "2024-02" and len(got) == months - 1, "one rate per month from the second"
    assert got["2023-06"] == pytest.approx(100 / 390, abs=1e-4)             # leaves: 100 of 390 at the start of the month
    assert got["2023-09"] == pytest.approx(40 / 290, abs=1e-4)              # shrinks: 40 of 290
    assert got["2023-03"] == 0.0 and got["2023-12"] == 0.0
    assert got["2023-11"] == 0.0, "a customer that joins in the month is not in its base"
    assert out["series"] and out["overall_pct"] is not None, "the trailing-12-month figures are unchanged"


def test_a_month_with_no_mrr_at_its_start_has_no_rate():
    mrr = mrr_table({"late": [0.0, 0.0] + [80.0] * 12})
    got = {d["month"]: d["churn_pct"] for d in ge.compute_gross_churn(mrr)["monthly_series"]}
    assert got["2023-02"] is None and got["2023-03"] is None and got["2023-04"] == 0.0


def test_the_engine_stores_the_monthly_series_with_the_trailing_one_and_skips_a_partial_last_month():
    results = results_for(RUNS["A"])
    block = results["gross_churn"]
    assert block["monthly_series"], "every run holds the series"
    assert [d["month"] for d in block["monthly_series"]][-1] == [d["month"] for d in block["series"]][-1], "both end at the last complete month"
    assert results["metric_periods"]["gross_churn_monthly"]["last"] == block["monthly_series"][-1]["month"]
    assert all(0.0 <= d["churn_pct"] <= 1.0 for d in block["monthly_series"] if d["churn_pct"] is not None)


def test_the_series_is_part_of_the_contract_and_a_payload_without_it_fails():
    good = cf.stored(gross_churn=cf.gross_churn(0.061, monthly=[{"month": "2026-05", "churn_pct": 0.01}]))
    MetricsPayload.model_validate(good)
    for bad in (_without(good, "monthly_series"), _with(good, "monthly_series", [{"month": "2026-05", "churn_pct": "1%"}])):
        with pytest.raises(Exception):
            MetricsPayload.model_validate(bad)


def _without(payload, key):
    payload = copy.deepcopy(payload)
    del payload["gross_churn"][key]
    return payload


def _with(payload, key, value):
    payload = copy.deepcopy(payload)
    payload["gross_churn"][key] = value
    return payload


# --- the matching ----------------------------------------------------------------------------------------------------

@pytest.mark.parametrize("text", ["2% monthly churn", "churn of 2% per month", "monthly revenue churn 2%", "2%/mo churn", "churn 2% MoM"])
def test_a_churn_claim_worded_monthly_is_tested_against_that_months_rate(text):
    results = with_monthly({"2024-01": 0.03, "2024-02": 0.02})
    got = churn_row(2, "2024-02", text, results=results)
    assert got["metric"] == MONTHLY and got["direction"] == "lower"
    assert (got["observed_value"], got["observed_at"], got["evidence_label"]) == (2.0, "2024-02", "Verified")
    assert got["evidence_source_key"] == "gross_churn.monthly_series.churn_pct"


def test_a_monthly_claim_beats_or_misses_by_the_rate_rule():
    results = with_monthly({"2024-01": 0.03, "2024-02": 0.02})
    miss = churn_row(1, "2024-01", "1% monthly churn", results=results)
    assert (miss["observed_value"], miss["observed_at"], miss["gap_kind"], miss["evidence_label"]) == (3.0, "2024-01", "miss", "Contradicted")
    beat = churn_row(5, "2024-02", "5% monthly churn", results=results)
    assert (beat["gap_kind"], beat["evidence_label"]) == ("beat", "Verified")


def test_a_monthly_claim_for_a_month_the_engine_has_no_rate_for_names_it():
    results = with_monthly({"2024-01": 0.03, "2024-02": 0.02})
    got = churn_row(2, "2022-05", "2% monthly churn", results=results)
    assert got["evidence_label"] == "Unverified" and "2024-01" in got["reason"] and got["observed_value"] is None


def test_a_churn_claim_for_a_year_is_tested_against_the_trailing_rate_at_year_end():
    rates = {f"2023-{m:02d}": 0.01 for m in range(1, 13)}
    results = with_monthly(rates, trailing=0.09)
    results["gross_churn"]["series"][-1]["churn_pct"] = 0.08
    results["source_period"] = {"month": "2023-12", "status": "Complete", "complete_by": "analyst", "evidence_to": "2023-12"}
    got = churn_row(8, "2023", "annual churn 8%", results=results, as_of_month="2023-12")
    assert got["metric"] == "Gross revenue churn" and (got["observed_value"], got["observed_at"]) == (8.0, "2023-12")
    assert got["evidence_label"] == "Verified"


def test_a_churn_claim_with_no_wording_and_no_period_is_unverified_and_names_both_candidates():
    for claim_type, snippet in (("churn", "churn 2%"), ("churn", ""), ("retention", "gross churn 2%")):
        got = run_claim({"claim_type": claim_type, "value": 2, "unit": "%", "currency": None, "target_date": None, "snippet": snippet})
        assert got["evidence_label"] == "Unverified" and got["reason"] == cm.CHURN_RATE_UNCLEAR_REASON, (claim_type, got["reason"])
        assert "trailing" in got["reason"] and "monthly" in got["reason"]
        assert got["observed_value"] is None and got["gap"] is None


def test_a_monthly_claim_with_no_period_still_asks_for_the_period():
    got = run_claim({"claim_type": "churn", "value": 2, "unit": "%", "currency": None, "target_date": None, "snippet": "2% monthly churn"})
    assert got["evidence_label"] == "Unverified" and got["reason"].startswith(cm.PERIOD_NOT_READ)
    assert got["metric"] == MONTHLY


def test_a_churn_claim_with_a_month_and_no_wording_keeps_the_trailing_reading():
    got = churn_row(2, "2024-02", "churn 2%")
    assert got["metric"] == "Gross revenue churn" and got["reason"] != cm.CHURN_RATE_UNCLEAR_REASON


def test_the_monthly_metric_is_declared_beside_the_others():
    assert cm.METRICS[MONTHLY] == dict(kind="month", unit="%", direction="lower", seg=False, missing=("gross revenue churn",))
    assert cm.EVIDENCE[MONTHLY][1] == "gross_churn.monthly_series.churn_pct"
    assert cm.PERIOD_BLOCK[MONTHLY] == "gross_churn_monthly"
    assert not hasattr(cm, "MONTHLY_CHURN_REASON"), "no claim reads Unverified for want of a monthly figure any more"


# --- the export ------------------------------------------------------------------------------------------------------

def test_the_export_carries_a_churn_series_sheet_with_both_rates():
    results = cf.stored(gross_churn=cf.gross_churn(0.061, series=[{"month": "2026-05", "churn_pct": 0.06}, {"month": "2026-06", "churn_pct": 0.061}],
                                                  monthly=[{"month": "2026-04", "churn_pct": 0.004}, {"month": "2026-05", "churn_pct": 0.005},
                                                           {"month": "2026-06", "churn_pct": 0.006}]))
    wb = openpyxl.load_workbook(io.BytesIO(server.build_export_workbook({"company_name": "X", "reporting_currency": "EUR"}, results, None, None).getvalue()))
    assert "Churn Series" in wb.sheetnames and ne.CHURN_SERIES == "Churn Series" and ne.CHURN_SERIES in ne.DATA_SHEETS
    rows = [[c.value for c in r] for r in wb["Churn Series"].iter_rows()]
    assert rows[0] == ["Month", "Gross revenue churn (trailing 12 months)", "Gross revenue churn (monthly)"]
    assert rows[1:] == [["2026-04", None, 0.004], ["2026-05", 0.06, 0.005], ["2026-06", 0.061, 0.006]]
    for key in ("metrics.gross_churn.monthly_series.churn_pct", "metrics.gross_churn.series.churn_pct"):
        assert ne.data_sheet_for(key, {}) == ne.CHURN_SERIES, key
