"""The page image (docs/design/page-image.md section 1): one PNG of the whole page, 1,568 px on the long edge, drawn from the stored
file. Synthetic files only (CLAUDE.md rule 15). This module runs the real render functions; every other test gets the synthetic
PNG of tests/page_images.py (conftest.py)."""
import asyncio
import io
import shutil

import pytest
from PIL import Image

pytest.importorskip("pypdfium2")

from app.decks import render  # noqa: E402
from pdf_support import pdf_page  # noqa: E402
from test_llm_gateway import make_db  # noqa: E402


def png_size(png: bytes):
    image = Image.open(io.BytesIO(png))
    assert image.format == "PNG"
    return image.size


def two_page_pdf() -> bytes:
    """A landscape page and a portrait page, drawn with Pillow."""
    out = io.BytesIO()
    first, second = Image.new("RGB", (800, 450), "white"), Image.new("RGB", (300, 400), "white")
    first.save(out, format="PDF", save_all=True, append_images=[second])
    return out.getvalue()


def test_a_landscape_page_is_drawn_whole_with_its_long_edge_at_1568_px():
    pdf = pdf_page([(50, 500, 24, "Traction"), (50, 450, 14, "800 Paying Users")])
    width, height = png_size(render.render_page(pdf, 1))
    assert width == render.LONG_EDGE == 1568
    assert height == round(612 * 1568 / 792), "the whole page: the aspect ratio of the file's page, no crop"


def test_a_portrait_page_has_its_height_at_1568_px_and_a_second_page_is_its_own():
    pdf = two_page_pdf()
    assert render.page_count(pdf) == 2
    assert png_size(render.render_page(pdf, 1)) == (1568, round(450 * 1568 / 800))
    assert png_size(render.render_page(pdf, 2)) == (round(300 * 1568 / 400), 1568)


def test_a_page_the_file_does_not_have_is_no_image():
    with pytest.raises(render.PageImageError) as caught:
        render.render_page(two_page_pdf(), 3)
    assert caught.value.reason == render.PAGE_COUNT_DIFFERS


def test_the_stored_file_comes_back_whole_across_its_parts(monkeypatch):
    monkeypatch.setattr(render, "PART_BYTES", 7)
    db, content = make_db(), two_page_pdf()
    asyncio.run(render.save(db, "a1", "d1", content, "pdf"))
    assert len(db[render.COLLECTION].docs) > 1, "the file is kept in parts"
    assert asyncio.run(render._load(db, "a1", "d1", "original")) == content


def test_prepare_gives_the_pdf_of_a_stored_pdf_and_checks_the_page_count():
    db, content = make_db(), two_page_pdf()
    asyncio.run(render.save(db, "a1", "d1", content, "pdf"))
    assert asyncio.run(render.prepare(db, "a1", "d1", "pdf", 2)) == content
    with pytest.raises(render.PageImageError) as caught:
        asyncio.run(render.prepare(db, "a1", "d1", "pdf", 5))
    assert caught.value.reason == render.PAGE_COUNT_DIFFERS, "a rendering with another page count is never drawn"


def test_a_deck_with_no_stored_file_is_no_image():
    with pytest.raises(render.PageImageError) as caught:
        asyncio.run(render.prepare(make_db(), "a1", "d1", "pdf", 1))
    assert caught.value.reason == render.NO_FILE


def test_remove_deletes_one_deck_or_the_whole_audit():
    db = make_db()
    for deck in ("d1", "d2"):
        asyncio.run(render.save(db, "a1", deck, b"x", "pdf"))
    asyncio.run(render.save(db, "a2", "d3", b"x", "pdf"))
    asyncio.run(render.remove(db, "a1", ["d1"]))
    assert {d["deck_id"] for d in db[render.COLLECTION].docs} == {"d2", "d3"}
    asyncio.run(render.remove(db, "a1"))
    assert {d["deck_id"] for d in db[render.COLLECTION].docs} == {"d3"}


def test_a_file_that_cannot_be_converted_is_no_image(monkeypatch):
    monkeypatch.setattr(render.shutil, "which", lambda name: None)
    with pytest.raises(render.PageImageError) as caught:
        render.to_pdf(b"x", "pptx")
    assert caught.value.reason == render.NOT_RENDERED


@pytest.mark.skipif(not (shutil.which("soffice") or shutil.which("libreoffice")), reason="LibreOffice is not installed")
def test_a_pptx_is_drawn_slide_by_slide():
    from pptx import Presentation
    deck = Presentation()
    for title in ("One", "Two"):
        deck.slides.add_slide(deck.slide_layouts[5]).shapes.title.text = title
    out = io.BytesIO()
    deck.save(out)
    pdf = render.to_pdf(out.getvalue(), "pptx")
    assert render.page_count(pdf) == 2, "one page per slide"
    width, height = png_size(render.render_page(pdf, 2))
    assert max(width, height) == 1568
