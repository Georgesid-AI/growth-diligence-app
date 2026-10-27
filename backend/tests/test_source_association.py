"""The source-association rule (docs/specs/source-association.md; section 4 and its acceptance row "Chart/heading checks" of
docs/design/claim-review-simplified-design.md): rotated axes, distant legends, repeated colours and numbers, attached dates, ambiguous
headings, incomplete extraction, the model and log boundaries.

Synthetic pages only (CLAUDE.md rule 15): grids built by hand with the places of the lines and the rectangles of the deck's layout, and
a few small pdf and pptx files drawn here. No client deck, no real company name or figure.

Each test names the behaviour that is new; the ones marked "pin" hold what the rule keeps (they pass without the change).
"""
import asyncio
import io
import json
import logging
import sys
from pathlib import Path

import pytest

BACKEND = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(BACKEND))
sys.path.insert(0, str(BACKEND / "tests"))

pytest.importorskip("pdfplumber")
pytest.importorskip("pptx")

import test_checker as tc  # noqa: E402
import test_llm_gateway as t  # noqa: E402
from pdf_support import pdf_page  # noqa: E402
from test_structure_reading import AUDIT  # noqa: E402

from app import structures  # noqa: E402
from app.decks import CANDIDATES_COLLECTION, TEXT_COLLECTION, pages  # noqa: E402
from app.decks.parser import parse_deck  # noqa: E402
from app.llm import gateway  # noqa: E402
from app.llm.schemas import PAGE_CLAIMS  # noqa: E402
from app.structures import checker, page_check  # noqa: E402
from app.structures import pages as pipeline  # noqa: E402

BLUE, ORANGE = "1F77B4", "FF7F0E"
DISAGREE, AMBIGUOUS, NOT_ASSOCIATED = "sources disagree", "source ambiguous", "source not associated"
INCOMPLETE, NO_SOURCE = "extraction incomplete", "no printed source"


# ---------------------------------------------------------------------------
# Builders
# ---------------------------------------------------------------------------
def cell(row, text, at, box, col=1, **more):
    return {"row": row, "col": col, "text": text, "box": box, "at": list(at), "next_to": [], **more}


def rect(at, fill, scope=0):
    return {"at": list(at), "fill": fill, "scope": scope}


def grid(cells, rects=(), gaps=(), measured=True, tables=None):
    out = {"type": "page", "page": 5, "header_rows": 0, "cells": cells, "tables": tables or []}
    if measured:
        out["layout"] = {"rects": list(rects), "gaps": list(gaps)}
    return out


def claim(value_cell, value, metric="revenue", unit="EUR", period=None, period_cell=None, label_cell=None, basis="actual", scenario="none"):
    return {"metric": metric, "value": value, "value_high": None, "unit": unit, "unit_other": None, "period": period, "basis": basis,
            "scenario": scenario, "value_cell": value_cell, "label_cell": label_cell, "period_cell": period_cell}


def check(grid_, *claims_, fiscal_year_end=12):
    out = checker.check(grid_, list(claims_), fiscal_year_end)
    assert len(out["items"]) == len(claims_), "a claim on a cell of the page is never deleted"
    return out["items"][0] if len(claims_) == 1 else out["items"]


def bar_chart(label_at=(0.31, 0.20, 0.37, 0.27)):
    """Two blue bars, the value of the first printed above it, and a rotated axis title level with them, at the left."""
    return [cell(1, "Growth plan", (0.05, 0.05, 0.40, 0.10), 1, title=True),
            cell(2, "12.5", label_at, 2),
            cell(3, "Revenue (EUR M)", (0.20, 0.30, 0.24, 0.80), 3, rot=270)], [rect((0.30, 0.30, 0.38, 0.80), BLUE), rect((0.40, 0.50, 0.48, 0.80), BLUE)]


def legend_chart(legend_gap=0.0):
    """A blue and an orange bar, a value over the blue one, and a legend far from the bars (bottom right)."""
    cells = [cell(1, "Revenue by segment", (0.05, 0.05, 0.50, 0.10), 1, title=True),
             cell(2, "42", (0.31, 0.20, 0.37, 0.27), 2),
             cell(3, "17", (0.41, 0.45, 0.47, 0.52), 3),
             cell(4, "Enterprise", (0.86 + legend_gap, 0.90, 0.98, 0.93), 4),
             cell(5, "SMB", (0.86 + legend_gap, 0.95, 0.98, 0.98), 5)]
    rects = [rect((0.30, 0.30, 0.38, 0.80), BLUE), rect((0.40, 0.55, 0.48, 0.80), ORANGE),
             rect((0.80, 0.90, 0.84, 0.93), BLUE), rect((0.80, 0.95, 0.84, 0.98), ORANGE)]
    return cells, rects


# ---------------------------------------------------------------------------
# 1. Rotated axes
# ---------------------------------------------------------------------------
def test_a_rotated_axis_title_beside_the_chart_states_the_unit_and_is_the_label():
    cells, rects = bar_chart()
    item = check(grid(cells, rects), claim("r2c1", 12.5, unit="EUR", label_cell="r3c1"))
    assert item["status"] == "matched" and item["flags"] == [], "the currency stands in the axis title: nothing is assumed"
    assert item["label_cell"] == "r3c1" and item["checks"]["unit"] is True
    assert (item["association"]["label"], item["association"]["unit"]) == ("axis", "axis")


def test_the_same_page_without_the_decks_layout_reads_as_before():
    """pin: a grid built without the layout (a deck stored before this field) has no axis, legend or extraction evidence."""
    cells, _ = bar_chart()
    item = check(grid(cells, measured=False), claim("r2c1", 12.5, unit="EUR", label_cell="r3c1"))
    assert item["flags"] == [checker.CURRENCY_ASSUMED, checker.LABEL_MISMATCH] and item["label_cell"] is None and "association" not in item


def test_a_rotated_title_that_is_not_level_with_the_chart_is_no_source():
    cells, rects = bar_chart()
    cells[2] = cell(3, "Revenue (EUR M)", (0.20, 0.88, 0.24, 0.99), 3, rot=270)         # below the bars, not level with them
    item = check(grid(cells, rects), claim("r2c1", 12.5, unit="EUR", label_cell="r3c1"))
    assert item["flags"] == [checker.CURRENCY_ASSUMED, checker.LABEL_MISMATCH] and item["label_cell"] is None and item["association"]["unit"] is None


def test_two_rotated_titles_with_different_words_are_ambiguous_and_none_is_chosen():
    cells, rects = bar_chart()
    cells.append(cell(4, "Customers (thousands)", (0.52, 0.30, 0.56, 0.80), 4, rot=90))        # the right axis
    item = check(grid(cells, rects), claim("r2c1", 12.5, unit="EUR"))
    assert AMBIGUOUS in item["flags"] and item["label_cell"] is None and item["association"]["label"] is None


# ---------------------------------------------------------------------------
# 2. Distant legends
# ---------------------------------------------------------------------------
def test_the_legend_label_of_the_bars_colour_is_the_label_however_far_the_legend_stands():
    cells, rects = legend_chart()
    item = check(grid(cells, rects), claim("r2c1", 42, unit="EUR", label_cell=None))
    assert item["label_cell"] == "r4c1" and item["association"]["label"] == "legend" and item["flags"] == [checker.CURRENCY_ASSUMED]
    other = check(grid(cells, rects), claim("r3c1", 17, unit="EUR"))
    assert other["label_cell"] == "r5c1", "the orange bar takes the orange entry"


def test_a_local_legend_label_takes_precedence_over_the_heading_and_over_another_series():
    cells, rects = legend_chart()
    over_heading = check(grid(cells, rects), claim("r2c1", 42, label_cell="r1c1"))
    assert over_heading["label_cell"] == "r4c1" and DISAGREE not in over_heading["flags"], "a heading and a segment label do not disagree"
    wrong_series = check(grid(cells, rects), claim("r2c1", 42, label_cell="r5c1"))
    assert wrong_series["label_cell"] == "r4c1" and DISAGREE in wrong_series["flags"], "the claim cites the other series' entry: review"


def test_a_label_that_stands_too_far_from_its_swatch_is_no_legend():
    cells, rects = legend_chart(legend_gap=0.12)
    item = check(grid(cells, rects), claim("r2c1", 42))
    assert item["label_cell"] is None and item["association"]["label"] is None


def test_a_lone_card_or_tile_is_no_chart_and_a_sentence_is_no_legend_label():
    """A coloured tile with its value inside is not a bar, and a bullet square of the same colour beside a sentence is not a legend."""
    cells = [cell(1, "Plan", (0.05, 0.02, 0.2, 0.05), 1, title=True), cell(2, "42", (0.12, 0.45, 0.18, 0.50), 2),
             cell(3, "Enterprise", (0.86, 0.90, 0.98, 0.93), 3)]
    rects = [rect((0.10, 0.40, 0.30, 0.60), BLUE), rect((0.80, 0.90, 0.84, 0.93), BLUE)]
    assert check(grid(cells, rects), claim("r2c1", 42))["label_cell"] is None, "one rectangle is a tile, not a chart"
    rects.append(rect((0.32, 0.40, 0.52, 0.60), BLUE))
    assert check(grid(cells, rects), claim("r2c1", 42))["label_cell"] == "r3c1", "two tiles of a colour are a chart with a legend"
    cells[2] = cell(3, "Our mission is to make every team a little better at planning the quarter ahead", (0.86, 0.90, 0.98, 0.93), 3)
    assert check(grid(cells, rects), claim("r2c1", 42))["label_cell"] is None, "a sentence next to a square is not a legend entry"


def test_the_label_beside_the_figure_is_as_local_as_the_legend_and_stays():
    cells, rects = legend_chart()
    cells[1]["next_to"], cells[2]["next_to"] = ["r6c1"], []
    cells.append(cell(6, "Net revenue", (0.31, 0.12, 0.37, 0.16), 6, next_to=["r2c1"]))
    item = check(grid(cells, rects), claim("r2c1", 42, label_cell="r6c1"))
    assert item["label_cell"] == "r6c1" and DISAGREE not in item["flags"] and item["association"]["label"] == "header"


# ---------------------------------------------------------------------------
# 3. Repeated colours and numbers
# ---------------------------------------------------------------------------
def test_one_colour_with_two_legend_labels_is_ambiguous_and_nothing_is_chosen():
    cells, rects = legend_chart()
    cells.append(cell(6, "Mid-market", (0.86, 0.85, 0.98, 0.88), 6))
    rects.append(rect((0.80, 0.85, 0.84, 0.88), BLUE))
    item = check(grid(cells, rects), claim("r2c1", 42, label_cell="r4c1"))
    assert AMBIGUOUS in item["flags"] and item["label_cell"] is None


def test_a_legend_is_read_inside_its_own_chart_only():
    """Two grouped charts (pptx) use the same two colours with a legend each: every bar takes the entry of its own group."""
    cells = [cell(1, "Plan", (0.05, 0.02, 0.2, 0.05), 1, title=True),
             cell(2, "10", (0.11, 0.32, 0.15, 0.37), 2), cell(3, "A-entry", (0.30, 0.20, 0.40, 0.23), 3),
             cell(4, "20", (0.61, 0.32, 0.65, 0.37), 4), cell(5, "B-entry", (0.80, 0.20, 0.90, 0.23), 5)]
    rects = [rect((0.10, 0.40, 0.16, 0.80), BLUE, 1), rect((0.25, 0.20, 0.28, 0.23), BLUE, 1),
             rect((0.60, 0.40, 0.66, 0.80), BLUE, 2), rect((0.75, 0.20, 0.78, 0.23), BLUE, 2)]
    first, second = check(grid(cells, rects), claim("r2c1", 10), claim("r4c1", 20))
    assert (first["label_cell"], second["label_cell"]) == ("r3c1", "r5c1")
    assert AMBIGUOUS not in first["flags"] + second["flags"]


def test_one_legend_colour_used_by_two_charts_with_a_legend_of_neither_is_ambiguous():
    cells = [cell(1, "Plan", (0.05, 0.02, 0.2, 0.05), 1, title=True),
             cell(2, "10", (0.11, 0.32, 0.15, 0.37), 2), cell(3, "20", (0.61, 0.32, 0.65, 0.37), 3), cell(4, "Enterprise", (0.45, 0.93, 0.55, 0.96), 4)]
    rects = [rect((0.10, 0.40, 0.16, 0.80), BLUE), rect((0.18, 0.60, 0.24, 0.80), BLUE), rect((0.60, 0.40, 0.66, 0.80), BLUE),
             rect((0.68, 0.60, 0.74, 0.80), BLUE), rect((0.40, 0.93, 0.43, 0.96), BLUE)]
    assert AMBIGUOUS in check(grid(cells, rects), claim("r2c1", 10))["flags"]
    assert AMBIGUOUS in check(grid(cells, rects), claim("r3c1", 20))["flags"]


def test_a_number_found_only_in_another_cell_does_not_carry_the_period_cited_for_the_cell_it_was_not_in():
    """The model cites a cell that does not hold 150,000 and a period cell that belongs to it; the number stands in another cell.
    The nearest cell is still taken (cell tolerance, 2026-10-23), but the period of the cell it was not in is not that figure's."""
    cells = [cell(1, "Traction", (0.05, 0.02, 0.2, 0.05), 1, title=True),
             cell(2, "$150,000 revenue", (0.10, 0.30, 0.40, 0.34), 2),
             cell(3, "800 users", (0.60, 0.60, 0.80, 0.64), 3, next_to=["r4c1"]),
             cell(4, "Q3 2025", (0.60, 0.66, 0.70, 0.70), 4, next_to=["r3c1"])]
    item = check(grid(cells), claim("r3c1", 150_000, unit="USD", period="2025-Q3", period_cell="r4c1", label_cell="r3c1"))
    assert item["status"] == "matched" and item["value_cell"] == "r2c1" and checker.CELL_CORRECTED in item["flags"]
    assert NOT_ASSOCIATED in item["flags"] and checker.PERIOD_FROM_OTHER_CELL not in item["flags"]
    assert item["period_cells"] == [] and item["label_cell"] is None


def test_a_value_in_the_cited_cell_is_not_touched_by_the_same_number_elsewhere():
    """pin: the cited cell holds it; the same number in another cell changes nothing."""
    cells = [cell(1, "Traction", (0.05, 0.02, 0.2, 0.05), 1, title=True),
             cell(2, "$150,000 revenue", (0.10, 0.30, 0.40, 0.34), 2), cell(3, "$150,000 pipeline", (0.60, 0.60, 0.90, 0.64), 3)]
    item = check(grid(cells), claim("r2c1", 150_000, unit="USD"))
    assert item["flags"] == [] and item["value_cell"] == "r2c1"


# ---------------------------------------------------------------------------
# 4. Attached dates, marker annotations, local over heading
# ---------------------------------------------------------------------------
def test_a_date_attached_to_the_figure_is_its_period():
    """pin: the date line directly under the figure, in the next box, is a source of its period."""
    cells = [cell(1, "Traction", (0.05, 0.02, 0.2, 0.05), 1, title=True),
             cell(2, "$150,000", (0.10, 0.30, 0.25, 0.34), 2, next_to=["r3c1"]), cell(3, "Q3 2025", (0.10, 0.35, 0.25, 0.38), 3, next_to=["r2c1"])]
    item = check(grid(cells), claim("r2c1", 150_000, unit="USD", period="2025-Q3", period_cell="r3c1"))
    assert item["flags"] == [] and item["period"] == "2025-Q3" and item["period_cells"] == ["r3c1"]


def test_a_marker_annotation_date_is_not_the_period_of_another_figure():
    cells, rects = bar_chart()
    cells.append(cell(4, "Mar 2024", (0.40, 0.35, 0.48, 0.38), 4, next_to=["r2c1"]))        # printed in the plot area, above the baseline
    item = check(grid(cells, rects), claim("r2c1", 12.5, period="2024-03", period_cell="r4c1", label_cell="r3c1"))
    assert NOT_ASSOCIATED in item["flags"] and checker.PERIOD_FROM_OTHER_CELL not in item["flags"]
    assert item["period_cells"] == [] and item["association"]["period"] == "annotation"
    assert item["period"] == "2024-03", "the model's period is flagged, never retyped or deleted"


def test_a_date_on_the_axis_below_the_bars_is_not_an_annotation():
    """pin: the category label under the bar is the axis, not a marker."""
    cells, rects = bar_chart()
    cells.append(cell(4, "2024", (0.31, 0.83, 0.37, 0.86), 4, next_to=["r2c1"]))
    item = check(grid(cells, rects), claim("r2c1", 12.5, period="2024", period_cell="r4c1", label_cell="r3c1"))
    assert NOT_ASSOCIATED not in item["flags"] and item["period"] == "2024" and item["period_cells"] == ["r4c1"]


def test_the_figures_own_column_header_fills_a_blank_period_before_the_heading_and_a_model_period_is_kept_over_both():
    g = grid([{"row": 1, "col": 1, "text": "Plan 2025", "box": 1, "title": True, "top_band": True},
              {"row": 2, "col": 2, "text": "FY2024", "table": 1, "trow": 1, "tcol": 2},
              {"row": 3, "col": 1, "text": "Revenue", "table": 1, "trow": 2, "tcol": 1},
              {"row": 3, "col": 2, "text": "1,200", "table": 1, "trow": 2, "tcol": 2}],
             tables=[{"table": 1, "header_rows": 1}], measured=False)
    blank = check(g, claim("r3c2", 1200, unit="EUR", period=None, period_cell="r1c1", label_cell="r3c1"))
    assert blank["period"] == "2024" and blank["period_cells"] == ["r2c2"], "the local header comes first"
    assert DISAGREE in blank["flags"]
    plan, = [p for p in blank["corrections"] if p["flag"] == DISAGREE]
    assert [o["period"] for o in plan["options"]] == ["2024", "2025"], "two buttons, none applied"
    # A period the model set is never replaced (page-image.md section 5, review of PR #193): it stays, the figure's own header is the other button.
    item = check(g, claim("r3c2", 1200, unit="EUR", period="2025", period_cell="r1c1", label_cell="r3c1"))
    assert item["period"] == "2025" and item["period_cells"] == [] and DISAGREE in item["flags"]
    plan, = [p for p in item["corrections"] if p["flag"] == DISAGREE]
    assert [o["period"] for o in plan["options"]] == ["2025", "2024"], "the model's period and the local source, not the heading"


def test_a_heading_that_agrees_with_the_figures_own_header_raises_nothing():
    """pin"""
    g = grid([{"row": 1, "col": 1, "text": "Plan 2024", "box": 1, "title": True, "top_band": True},
              {"row": 2, "col": 2, "text": "FY2024", "table": 1, "trow": 1, "tcol": 2},
              {"row": 3, "col": 1, "text": "Revenue", "table": 1, "trow": 2, "tcol": 1},
              {"row": 3, "col": 2, "text": "1,200", "table": 1, "trow": 2, "tcol": 2}],
             tables=[{"table": 1, "header_rows": 1}], measured=False)
    item = check(g, claim("r3c2", 1200, unit="EUR", period="2024", period_cell="r1c1", label_cell="r3c1"))
    assert DISAGREE not in item["flags"] and item["period"] == "2024"


def test_no_year_is_inferred_for_a_heading_that_states_none():
    """pin (spec deck-parser.md, period rules): "Q3" alone is a part, never a period."""
    g = grid([cell(1, "Q3 outlook", (0.05, 0.02, 0.3, 0.06), 1, title=True, top_band=True), cell(2, "$150,000 revenue", (0.1, 0.3, 0.4, 0.34), 2)],
             measured=False)
    item = check(g, claim("r2c1", 150_000, unit="USD", period="2025-Q3"))
    assert checker.PERIOD_NOT_ON_PAGE in item["flags"] and item["corrections"] == []


# ---------------------------------------------------------------------------
# 5. Ambiguous headings
# ---------------------------------------------------------------------------
def test_two_headings_that_state_different_periods_select_nothing():
    g = grid([cell(1, "Plan 2025E", (0.05, 0.02, 0.3, 0.06), 1, title=True, top_band=True),
              cell(2, "Actuals 2024", (0.05, 0.08, 0.3, 0.12), 2, top_band=True),
              cell(3, "$150,000 revenue", (0.1, 0.4, 0.4, 0.44), 3)], measured=False)
    item = check(g, claim("r3c1", 150_000, unit="USD", period="2025", period_cell="r1c1"))
    assert AMBIGUOUS in item["flags"] and item["period"] == "2025", "the model's period stays, nothing is selected"
    plan, = [p for p in item["corrections"] if p["flag"] == AMBIGUOUS]
    assert [(o["period"], o["period_cells"], o["marker"]) for o in plan["options"]] == [("2025", ["r1c1"], "forecast"), ("2024", ["r2c1"], "actual")]


def test_one_heading_stays_informational():
    """pin: a single page heading is an applicable heading (George, 2026-10-23)."""
    g = grid([cell(1, "Plan 2025", (0.05, 0.02, 0.3, 0.06), 1, title=True, top_band=True), cell(2, "$150,000 revenue", (0.1, 0.4, 0.4, 0.44), 2)],
             measured=False)
    item = check(g, claim("r2c1", 150_000, unit="USD", period="2025", period_cell="r1c1"))
    assert item["flags"] == [checker.PERIOD_FROM_HEADING]


# ---------------------------------------------------------------------------
# 6. Incomplete extraction
# ---------------------------------------------------------------------------
PAGE = [cell(1, "Traction", (0.05, 0.02, 0.2, 0.05), 1, title=True), cell(2, "$150,000 revenue", (0.1, 0.3, 0.4, 0.34), 2)]


def test_a_value_on_no_cell_of_a_completely_read_page_is_dropped_with_its_reason():
    item = check(grid(PAGE, gaps=()), claim("r2c1", 999_000, unit="USD"))
    assert item["drop"] == NO_SOURCE and checker.VALUE_NOT_ON_PAGE in item["flags"] and INCOMPLETE not in item["flags"]


@pytest.mark.parametrize("gaps", [["picture"], ["chart"], ["graphic"], ["glyphs"], ["format"], ["shapes"], ["picture", "chart"]])
def test_the_same_value_on_a_page_the_reader_could_not_read_completely_is_kept_for_extraction_review(gaps):
    item = check(grid(PAGE, gaps=gaps), claim("r2c1", 999_000, unit="USD"))
    assert item["drop"] is None and item["flags"] == [checker.VALUE_NOT_ON_PAGE, INCOMPLETE] and item["status"] == "unmatched"
    assert INCOMPLETE in checker.REVIEW_FLAGS


def test_a_page_whose_extraction_was_not_measured_is_neither_dropped_nor_flagged():
    """pin: a deck stored before the layout field, and a hand-built grid."""
    item = check(grid(PAGE, measured=False), claim("r2c1", 999_000, unit="USD"))
    assert item["drop"] is None and item["flags"] == [checker.VALUE_NOT_ON_PAGE]


def test_a_value_the_page_holds_is_never_dropped_and_a_date_claim_is_not_a_value():
    held = check(grid(PAGE, gaps=()), claim("r2c1", 150_000, unit="USD"))
    assert held["drop"] is None and held["flags"] == []
    qualitative = check(grid(PAGE, gaps=()), {**claim("r2c1", None, metric="product", unit=None), "value": None})
    assert qualitative["drop"] is None, "a claim with no value is not dropped for having no printed value"


def test_a_number_elsewhere_on_the_page_does_not_carry_the_metric_cited_for_another_cell():
    """The value is in another cell (cell tolerance), the claim cites a label cell that is not next to it: the figure stands, the metric
    does not follow it (source not associated)."""
    cells = PAGE + [cell(3, "Footnote words", (0.6, 0.8, 0.9, 0.84), 3)]
    item = check(grid(cells, gaps=()), claim("r3c1", 150_000, unit="USD", label_cell="r3c1"))
    assert item["value_cell"] == "r2c1" and NOT_ASSOCIATED in item["flags"] and item["drop"] is None


# ---------------------------------------------------------------------------
# 7. Correction records: the token, the cells, the marker; never the heading's prose
# ---------------------------------------------------------------------------
def _db_with_layout(parsed):
    db = tc._page_db(parsed)
    db[TEXT_COLLECTION].docs[0]["layout"] = parsed["layout"]
    return db


def read_with_layout(parsed, *claims_):
    db = _db_with_layout(parsed)
    adapter = t.FakeAdapter(replies=[json.dumps({"type": PAGE_CLAIMS, "claims": list(claims_)})])
    assert asyncio.run(structures.process_deck(db, AUDIT, "d1", adapter=adapter, sleep=t._noop_sleep)) == "read"
    deck_doc, = db[TEXT_COLLECTION].docs
    return db, sorted(db[CANDIDATES_COLLECTION].docs, key=lambda r: r["order"]), deck_doc


def test_a_blank_period_filled_from_another_cell_keeps_the_token_the_cells_and_the_marker_and_no_heading_prose():
    """A heading or another cell fills a blank period only (page-image.md section 5): the correction record keeps the token, the cell and the marker."""
    parsed = tc.deck([tc.TITLE, (50, 480, 12, "$150,000 annual revenue run rate"), (400, 300, 12, "Plan year 2025E strategy for Northwind")])
    cell_ = f"p1{tc.cell_of(parsed, '$150,000 annual revenue run rate')}"
    other = f"p1{tc.cell_of(parsed, 'Plan year 2025E strategy for Northwind')}"
    db, rows, deck_doc = read_with_layout(parsed, claim(cell_, 150_000, unit="USD", period=None, period_cell=other))
    row, = rows
    entry, = row["corrections"]
    assert (entry["state"], entry["reason"], entry["original"], entry["corrected"]) == ("applied", checker.PERIOD_FROM_OTHER_CELL, None, "2025")
    assert (entry["cells"], entry["marker"]) == ([other[2:]], "forecast")
    line, = [d for d in deck_doc["diagnostics"] if d["kind"] == "corrected"]
    assert (line["corrected"], line["cells"], line["marker"]) == ("2025", [other[2:]], "forecast")
    stored = json.dumps(db[gateway.STRUCTURES_COLLECTION].docs[0]["items"]) + json.dumps(deck_doc["diagnostics"])
    for prose in ("strategy", "Northwind", "Plan year 2025E"):
        assert prose not in stored, f"{prose!r} in a diagnostics line or a stored item"
    kept, = db[gateway.STRUCTURES_COLLECTION].docs[0]["items"]["claims"][0]["corrections"]
    assert kept["options"] == ["2025"] and kept["sources"] == [{"cells": [other[2:]], "marker": "forecast"}]


def test_the_marker_is_kept_apart_from_the_period_token():
    from app.structures import association
    assert [association.period_marker(x) for x in ("2025E", "2025 Forecast", "FY2026B", "2024A", "2024 Actuals", "Q3 2025", None, "")] == \
        ["forecast", "forecast", "forecast", "actual", "actual", None, None, None]


# ---------------------------------------------------------------------------
# 8. The pipeline: a dropped claim leaves a Diagnostics line, an incomplete page keeps its row
# ---------------------------------------------------------------------------
def test_a_value_with_no_printed_source_leaves_one_diagnostics_line_and_no_row(caplog):
    parsed = tc.deck([tc.TITLE, *tc.LINES])
    assert parsed["layout"] == [{"page": 1, "rects": [], "gaps": []}]
    with caplog.at_level(logging.INFO):
        db, rows, deck_doc = read_with_layout(parsed, claim("p1r3c1", 999_000, unit="USD", period="2025"))
    assert rows == []
    line, = deck_doc["diagnostics"]
    assert line == {"page": 1, "kind": NO_SOURCE, "metric": "revenue", "value": 999_000, "period": "2025"}
    assert "Paying Users" not in caplog.text and "annual revenue run rate" not in caplog.text
    assert db[gateway.STRUCTURES_COLLECTION].docs[0]["items"]["claims"][0]["drop"] == NO_SOURCE


def test_the_same_claim_on_a_page_with_a_picture_keeps_its_row_flagged_for_extraction_review():
    parsed = tc.deck([tc.TITLE, *tc.LINES])
    parsed["layout"][0]["gaps"] = ["picture"]
    db, rows, deck_doc = read_with_layout(parsed, claim("p1r3c1", 999_000, unit="USD"))
    row, = rows
    assert row["flags"] == [checker.VALUE_NOT_ON_PAGE, INCOMPLETE] and row["ai_status"] == "unmatched"
    assert deck_doc["diagnostics"] == []


# ---------------------------------------------------------------------------
# 9. The reader: rotated text, rectangles, gaps (pdf and pptx)
# ---------------------------------------------------------------------------
def test_the_pdf_reader_reads_text_turned_a_quarter_as_one_line_with_its_turn():
    data = pdf_page([(120, 300, 12, "42")], rotated=[(60, 200, 12, "Revenue EUR M", "up"), (700, 400, 12, "Customers", "down")])
    blocks = parse_deck(data, "deck.pdf")["blocks"]
    turned = {b["text"]: b["rot"] for b in blocks if b.get("rot")}
    assert turned == {"Revenue EUR M": 270, "Customers": 90}, "one line each, never single characters"
    assert not [b for b in blocks if len(b["text"]) == 1], "no character of a turned line is left over as a line of its own"
    assert next(b for b in blocks if b["text"] == "42").get("rot") is None


def test_the_pdf_reader_keeps_filled_rectangles_and_the_pictures_it_cannot_read():
    data = pdf_page([(120, 300, 12, "42")], rects=[(300, 300, 100, 30)], fills=[(100, 100, 40, 150, (0.1, 0.4, 0.8)), (500, 100, 12, 12, (1, 1, 1)),
                                                                              (10, 10, 770, 590, (0.9, 0.9, 0.9))],
                    pictures=[(300, 100, 400, 200)])
    layout, = parse_deck(data, "deck.pdf")["layout"]
    assert [(r["fill"], r["scope"]) for r in layout["rects"]] == [("1A66CC", 0)], "a ruled box, a white fill and a page-sized background are no bars"
    assert layout["gaps"] == ["picture"]
    small = parse_deck(pdf_page([(120, 300, 12, "42")], pictures=[(10, 10, 20, 20)]), "deck.pdf")["layout"][0]
    assert small["gaps"] == [], "a logo is not a gap"


def _pptx(build):
    from pptx import Presentation
    from pptx.util import Emu
    prs = Presentation()
    prs.slide_width, prs.slide_height = Emu(9144000), Emu(5143500)
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    build(slide)
    out = io.BytesIO()
    prs.save(out)
    return out.getvalue()


def test_the_pptx_reader_keeps_turned_text_groups_colours_pictures_and_charts():
    from PIL import Image
    from pptx.chart.data import CategoryChartData
    from pptx.dml.color import RGBColor
    from pptx.enum.chart import XL_CHART_TYPE
    from pptx.enum.shapes import MSO_SHAPE
    from pptx.util import Inches

    def build(slide):
        title = slide.shapes.add_textbox(Inches(0.2), Inches(1), Inches(0.5), Inches(3))
        title.text_frame.text = "Revenue EUR M"
        title.rotation = 270
        vertical = slide.shapes.add_textbox(Inches(9), Inches(1), Inches(0.5), Inches(3))
        vertical.text_frame.text = "Customers"
        vertical.text_frame._txBody.bodyPr.set("vert", "vert270")
        bar = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(2), Inches(2), Inches(1), Inches(2))
        bar.fill.solid()
        bar.fill.fore_color.rgb = RGBColor(0x1F, 0x77, 0xB4)
        group = slide.shapes.add_group_shape()
        member = group.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(5), Inches(2), Inches(1), Inches(1))
        member.fill.solid()
        member.fill.fore_color.rgb = RGBColor(0xFF, 0x7F, 0x0E)
        group.shapes.add_textbox(Inches(5), Inches(1), Inches(1), Inches(0.5)).text_frame.text = "42"
        image = io.BytesIO()
        Image.new("RGB", (4, 4)).save(image, "PNG")
        image.seek(0)
        slide.shapes.add_picture(image, Inches(1), Inches(1), Inches(6), Inches(3))
        data = CategoryChartData()
        data.categories = ["A", "B"]
        data.add_series("S", (1, 2))
        slide.shapes.add_chart(XL_CHART_TYPE.COLUMN_CLUSTERED, Inches(1), Inches(1), Inches(2), Inches(2), data)

    parsed = parse_deck(_pptx(build), "deck.pptx")
    blocks = {b["text"]: b for b in parsed["blocks"] if b["kind"] == "text"}
    assert blocks["Revenue EUR M"]["rot"] == 270 and blocks["Customers"]["rot"] == 270
    assert "bbox" not in blocks["42"] and len(blocks["42"]["lbbox"]) == 4, "a line in a group keeps its slide place apart from the unchanged bbox"
    layout, = parsed["layout"]
    assert sorted((r["fill"], r["scope"]) for r in layout["rects"]) == [("1F77B4", 0), ("FF7F0E", 1)]
    assert layout["gaps"] == ["picture", "chart"]


def test_a_pdf_with_bars_swatches_a_legend_and_a_rotated_title_is_associated_end_to_end():
    texts = [(50, 560, 20, "Revenue by segment"), (120, 460, 12, "42"), (600, 70, 12, "Enterprise"), (600, 40, 12, "SMB")]
    fills = [(100, 100, 60, 340, (0.12, 0.47, 0.71)), (170, 100, 60, 150, (1.0, 0.5, 0.05)),
             (560, 68, 22, 18, (0.12, 0.47, 0.71)), (560, 38, 22, 18, (1.0, 0.5, 0.05))]
    parsed = parse_deck(pdf_page(texts, fills=fills, rotated=[(60, 150, 12, "Revenue (EUR M)", "up")]), "deck.pdf")
    g, = pages.page_grids(parsed["blocks"], parsed["layout"])
    ids = {c["text"]: page_check._id(c) for c in g["cells"]}
    item, = checker.check(g, [claim(ids["42"], 42, unit="EUR")], 12)["items"]
    assert item["label_cell"] == ids["Enterprise"] and item["flags"] == [] and item["association"]["unit"] == "axis"


# ---------------------------------------------------------------------------
# 10. The model and log boundaries
# ---------------------------------------------------------------------------
def test_no_layout_field_reaches_the_page_text_and_the_turned_line_goes_as_any_other_line():
    import re

    from app import decks as decks_module
    from app.structures import redact
    parsed = parse_deck(pdf_page([(50, 560, 20, "Revenue by segment"), (120, 460, 12, "42")], fills=[(100, 100, 60, 340, (0.12, 0.47, 0.71))],
                                 rotated=[(60, 150, 12, "Revenue (EUR M)", "up")]), "deck.pdf")
    g, = pages.page_grids(parsed["blocks"], parsed["layout"])
    assert g["layout"]["rects"] and any(c.get("rot") for c in g["cells"]) and all(c.get("at") for c in g["cells"])
    text = pipeline.batch_text(pipeline.batches([g], "Target Co", {})[0])
    assert all(re.fullmatch(r"p1r\d+c\d+(?: title)?: .*", line) for line in text.split("\n")), "one line per cell: an id and a text"
    assert "1F78B5" not in text and "0.0" not in text and "scope" not in text
    cells = redact.parse_page_text(text)[1]
    assert sorted(c["text"] for c in cells) == sorted(c["text"] for c in g["cells"]), "the lines, as before, and nothing else"
    assert all(set(c) <= {"row", "col", "text", "title", "row_span", "col_span"} for c in cells)
    assert not {"at", "rot", "layout", "rects", "fill", "scope", "lbbox", "association", "gaps"} & set(decks_module.GATEWAY_READABLE_FIELDS)


def test_the_new_words_are_closed_and_an_item_carries_words_cell_ids_and_period_tokens_only():
    from app.structures import association
    assert (checker.SOURCES_DISAGREE, checker.SOURCE_AMBIGUOUS, checker.SOURCE_NOT_ASSOCIATED, checker.EXTRACTION_INCOMPLETE, checker.NO_PRINTED_SOURCE) == \
        (DISAGREE, AMBIGUOUS, NOT_ASSOCIATED, INCOMPLETE, NO_SOURCE)
    assert set(association.FLAGS) <= set(checker.FLAGS) and INCOMPLETE in checker.FLAGS
    assert association.SOURCES == ("own", "header", "axis", "legend", "heading", "other", "annotation")
    from app.decks import parser
    assert parser.GAPS == ("picture", "chart", "graphic", "glyphs", "format", "shapes")
    cells, rects = legend_chart()
    cells[0]["text"] = "Revenue by segment at Northwind Holdings"
    item = check(grid(cells, rects), claim("r2c1", 42, label_cell="r5c1", period="2025", period_cell="r1c1"))
    stored = json.dumps(pipeline._without_text(item))
    for prose in ("Northwind", "Enterprise", "SMB", "Revenue by segment"):
        assert prose not in stored, f"{prose!r} in a stored item"
    assert set(item["association"]) == {"label", "period", "unit", "flags", "options"}
    assert {item["association"][k] for k in ("label", "period", "unit")} <= set(association.SOURCES) | {None}
