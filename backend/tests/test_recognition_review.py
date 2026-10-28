"""The review of PR 135 (Opus review, owner decisions of 19 Oct 2026): one test per finding, each built from a synthetic sheet.
Each High was written first as a failing test (CLAUDE.md rule 13) and shown failing again with its fix switched off (rule 11).

Recognition-level tests use the `rec` fixture; upload-level tests use the `api` fixture of conftest.py."""
import importlib
import io
import json

import pytest

pytest.importorskip("fastapi")
pytest.importorskip("motor")

from openpyxl import Workbook  # noqa: E402

from app import column_rules as cr  # noqa: E402
from test_chat_upload import AUDIT, adapter_of, upload  # noqa: E402
from test_revenue_recognition import (FIXTURES, MONTHLY, SIDE, csv_sheets, decide, monthly_bytes, recognise,  # noqa: E402,F401
                                      rules_only, run_stage, set_as_of)

YEAR = "2025-12-31"
MONTHS = "Jan-25,Feb-25,Mar-25,Apr-25,May-25,Jun-25,Jul-25,Aug-25,Sep-25,Oct-25,Nov-25,Dec-25"


@pytest.fixture()
def rec():
    return importlib.import_module("app.recognition")


def sheet_text(title, header, rows):
    return "\n".join([title, header] + rows) + "\n"


def run(rec, text, as_of=YEAR, **kw):
    return rec.recognise(csv_sheets(text), as_of_month=as_of, **kw)


# ---------------------------------------------------------------------------
# H1: a word in brackets marks the row; "Revenue (forecast)" is never a revenue row
# ---------------------------------------------------------------------------
H1 = sheet_text("Management pack (EUR),,,,,,,,,,,,", "Metric," + MONTHS, [
    "Revenue (actual)," + ",".join(str(100 + i) for i in range(6)) + ",,,,,,",
    "Revenue (forecast)," + ",".join(str(200 + i) for i in range(12)),
    "Customers," + ",".join(str(30 + i) for i in range(12)),
    "Cost of sales," + ",".join(str(40 + i) for i in range(12)),
    "Payroll," + ",".join(str(50 + i) for i in range(12))])


@pytest.mark.parametrize("word", ["forecast", "budget", "target", "plan", "pipeline", "bookings", "run rate"])
def test_h1_a_kind_word_in_brackets_marks_the_label(rec, word):
    kind = rec.non_revenue_kind(f"Revenue ({word})")
    assert (kind is not None) == (word != "plan"), "plan is not a kind word (a pricing plan is an actual row)"
    assert rec.is_candidate_label(f"Revenue ({word})") is (word == "plan")


def test_h1_a_forecast_row_written_in_brackets_is_never_the_revenue_row(rec):
    got = run(rec, H1, as_of="2025-06-30")
    assert got["revenue_row"]["label"] == "Revenue (actual)"
    assert [m["label"] for m in got["non_revenue_rows"]] == ["Revenue (forecast)"]
    assert [s["value"] for s in got["series"]] == [100 + i for i in range(6)]


# ---------------------------------------------------------------------------
# Decision 2: contracted and committed recurring revenue are contract values
# ---------------------------------------------------------------------------
@pytest.mark.parametrize("label", ["Contracted ARR", "Committed ARR", "Committed MRR", "CARR", "carr (EUR)"])
def test_d2_contracted_and_committed_recurring_revenue_are_contract_values_never_revenue(rec, label):
    assert rec.non_revenue_kind(label) == "contract_value" and rec.is_candidate_label(label) is False


def test_d2_an_arr_row_beside_a_contracted_arr_row_stays_the_revenue_row(rec):
    text = sheet_text("Management pack (EUR),,,", "Metric,Jan-25,Feb-25,Mar-25", [
        "Contracted ARR,900,910,920", "ARR,100,110,120", "Customers,3,4,5", "Cost of sales,4,5,6"])
    got = run(rec, text, as_of="2025-03-31")
    assert got["revenue_row"]["label"] == "ARR" and [m["kind"] for m in got["non_revenue_rows"]] == ["contract_value"]


# ---------------------------------------------------------------------------
# H2: a summed series never includes forecast, budget, bookings, pipeline, contract value or run-rate rows
# ---------------------------------------------------------------------------
def product_rows(extra):
    return [f"Product {i},{10 * i},{10 * i + 1},{10 * i + 2}" for i in range(1, 13)] + extra


def test_h2_a_budget_row_among_twelve_products_is_not_added(rec):
    text = sheet_text("Revenue by product line (EUR),,,", "Product line,Jan-25,Feb-25,Mar-25",
                      product_rows(["Budget,5000,5000,5000"]))
    got = run(rec, text, as_of="2025-03-31")
    assert got["series_basis"] == "sum_of_rows"
    assert [s["value"] for s in got["series"]] == [780, 792, 804]
    assert 15 not in got["series_rows"]


@pytest.mark.parametrize("label", ["Forecast", "Bookings", "Pipeline", "Contract value", "Run rate", "Revenue (forecast)"])
def test_h2_every_marked_kind_is_left_out_of_a_sum(rec, label):
    text = sheet_text("Revenue by product line (EUR),,,", "Product line,Jan-25,Feb-25,Mar-25",
                      ["Hardware,10,11,12", "Software,20,21,22", f"{label},900,900,900"])
    assert [s["value"] for s in run(rec, text, as_of="2025-03-31")["series"]] == [30, 32, 34], label


def test_h2_the_monthly_report_changed_to_a_revenue_summary_does_not_sum_forecast_bookings_or_pipeline(rec):
    got = rec.recognise(rec_sheets(MONTHLY), as_of_month="2025-06-30", decisions={"file_type": "revenue_summary"})
    assert got["series_basis"] == "row", "a metrics sheet reads its revenue row and never sums (re-check, owner decision)"
    assert not {9, 10, 11} & set(got["series_rows"]), "Forecast revenue, Bookings and Pipeline are rows 9, 10 and 11"
    assert got["series_rows"] == [4], "the revenue row only: no MRR, ARR, customers, % rows and no forecast, bookings or pipeline rows"
    assert got["series"][0]["value"] == 10570


def rec_sheets(name):
    return cr.read_workbook((FIXTURES / name).read_bytes(), name)


# ---------------------------------------------------------------------------
# H3: subtotals and grand totals are never added; a total is detected by value only from two or more rows
# ---------------------------------------------------------------------------
REGIONS = sheet_text("Revenue by region (EUR),,,", "Region,Jan-25,Feb-25,Mar-25", [
    "EMEA North,10,10,10", "EMEA South,20,20,20", "EMEA subtotal,30,30,30", "US East,40,40,40", "US West,50,50,50",
    "US subtotal,90,90,90", "Grand total,120,120,120"])


def test_h3_subtotals_and_a_grand_total_are_found_by_their_label_and_never_added(rec):
    got = run(rec, REGIONS, as_of="2025-03-31")
    assert [s["value"] for s in got["series"]] == [120, 120, 120]
    assert [(t["label"], t["matches"]) for t in got["total_rows"]] == [("EMEA subtotal", True), ("US subtotal", True),
                                                                      ("Grand total", True)]


def test_h3_subtotals_without_the_word_are_found_by_value_against_the_rows_since_the_previous_total(rec):
    text = REGIONS.replace("EMEA subtotal", "EMEA").replace("US subtotal", "US").replace("Grand total", "All regions")
    got = run(rec, text, as_of="2025-03-31")
    assert [s["value"] for s in got["series"]] == [120, 120, 120]
    assert [t["label"] for t in got["total_rows"]] == ["EMEA", "US", "All regions"]


def test_h3_a_row_equal_to_the_single_row_above_is_not_a_total(rec):
    text = sheet_text("Revenue by plan (EUR),,,", "Plan,Jan-25,Feb-25,Mar-25", ["Starter,600,650,700", "Team,600,650,700",
                                                                            "Business,100,100,100"])
    got = run(rec, text, as_of="2025-03-31")
    assert [s["value"] for s in got["series"]] == [1300, 1400, 1500] and got["total_rows"] == []


# ---------------------------------------------------------------------------
# H4: the analyst's last actual month applies in the dates-down layout
# ---------------------------------------------------------------------------
DOWN = sheet_text("", "Month,Revenue,Cost of sales,Payroll", [
    "Jan-25,100,10,5", "Feb-25,110,10,5", "Mar-25,120,10,5", "Apr-25,999,10,5", "May-25,999,10,5", "Jun-25,999,10,5"]).lstrip("\n")


def test_h4_the_as_of_month_ends_the_series_and_the_feed(rec):
    free = run(rec, DOWN, as_of="2025-06-30")
    assert [s["month"] for s in free["series"]][-1] == "2025-06"
    got = run(rec, DOWN, as_of="2025-03-31")
    assert got["last_actual"]["month"] == "2025-03" and got["last_actual"]["state"] == "auto"
    assert [s["value"] for s in got["series"]] == [100, 110, 120]
    feed = rec.engine_feed(got, "f.csv", got["sheet"])
    assert [x["value"] for x in feed["series"]] == [100, 110, 120] and feed["rows"] == [2, 3, 4]


# ---------------------------------------------------------------------------
# H5: columns marked budget, forecast or plan are excluded wherever they sit; one value per month
# ---------------------------------------------------------------------------
ALT_HEAD = "Pack (EUR),,,,,,\n,Actual,Budget,Actual,Budget,Actual,Budget\nMetric,Jan-25,Jan-25,Feb-25,Feb-25,Mar-25,Mar-25"


def test_h5_budget_columns_before_the_boundary_do_not_enter_the_series(rec):
    text = ALT_HEAD + "\nRevenue,100,150,110,160,120,170\nCustomers,3,3,4,4,5,5\nCost of sales,9,9,9,9,9,9\n"
    got = run(rec, text, as_of="2025-03-31")
    assert [(s["month"], s["value"]) for s in got["series"]] == [("2025-01", 100), ("2025-02", 110), ("2025-03", 120)]
    assert got["last_actual"]["column"] == 6 and got["last_actual"]["candidates"] == [], "read from the as-of month, never chosen"
    feed = rec.engine_feed(got, "f.csv", got["sheet"])
    assert [x["month"] for x in feed["series"]] == ["2025-01", "2025-02", "2025-03"], "one value per month"


def test_h5_a_sum_of_rows_leaves_the_budget_columns_out_too(rec):
    text = ALT_HEAD.replace("Pack (EUR)", "Revenue by product (EUR)").replace("Metric", "Product line") + "\nHardware,100,150,110,160,120,170\nSoftware,10,15,11,16,12,17\n"
    got = run(rec, text, as_of="2025-03-31")
    assert got["series_basis"] == "sum_of_rows"
    assert [(s["month"], s["value"]) for s in got["series"]] == [("2025-01", 110), ("2025-02", 121), ("2025-03", 132)]


# ---------------------------------------------------------------------------
# Decision 1 and H6: a file that gives no monthly revenue series is not the revenue file
# ---------------------------------------------------------------------------
def blocker_kinds(api):
    return [b["kind"] for b in api.get(f"/api/audits/{AUDIT}/blockers").json()["blockers"]]


def crm_workbook():
    from mapping_shapes import crm_anonymous_to_2028
    return crm_anonymous_to_2028()


def only_recognition_reply(word):
    return json.dumps({"type": "file_recognition", "sheet": 1, "file_type": word, "revenue_row": None, "last_actual_column": None,
                       "currency": None, "non_revenue_rows": []})


@pytest.mark.parametrize("scenario", ["no_consent", "failed_call", "model_says_not_revenue"])
def test_h6_a_crm_workbook_is_never_stored_as_the_revenue_file_and_the_banner_stays(api, scenario):
    import test_llm_gateway as t
    set_as_of(api, "2028-12-31")
    if scenario == "no_consent":
        api.db["audits"].docs[0]["structure_reading_consent"] = False
    elif scenario == "failed_call":
        api.state["adapter"] = t.FakeAdapter(replies=["not json"])
    else:
        api.state["adapter"] = t.FakeAdapter(replies=[only_recognition_reply("not_revenue")])
    body = upload(api, "CRM_Anonymous_to_2028.xlsx", crm_workbook()).json()
    assert body["status"] == "unknown_type", "the type picker, as before the recognition PR"
    assert api.db["datasets"].docs == []
    assert "revenue_file_missing" in blocker_kinds(api)
    assert run_stage(api, "mapping")["status"] == "failed"


def test_d1_a_snapshot_and_a_yearly_only_report_are_not_the_revenue_file(api):
    rules_only(api, "2025-12-31")
    for name in ("arr_snapshot.csv", "revenue_summary_by_product.csv"):
        body = upload(api, name, (FIXTURES / name).read_bytes()).json()
        assert body["status"] == "unknown_type", name
        assert api.db["datasets"].docs == [], name
    assert "revenue_file_missing" in blocker_kinds(api)


def test_d1_a_monthly_report_is_still_the_revenue_file(api):
    rules_only(api, "2025-06-30")
    body = upload(api, MONTHLY, monthly_bytes()).json()
    assert body["dtype"] == "revenue" and body["recognition"]["recognised"]


PNL_MONTHLY = ("Profit and loss (EUR),,,\nLine,Jan-25,Feb-25,Mar-25\nRevenue,100,110,120\nCost of sales,40,41,42\nGross profit,60,69,78\n").encode()
PNL_YEARLY = (FIXTURES / "pnl_multi_period.csv").read_bytes()


def test_d1_a_confirmed_yearly_pnl_does_not_clear_the_banner_and_a_confirmed_monthly_pnl_does(api):
    rules_only(api, "2025-12-31")
    assert upload(api, "pnl_multi_period.csv", PNL_YEARLY).json()["dtype"] == "pnl"
    assert decide(api, dtype="pnl", action="confirm", part="revenue_row").status_code == 200
    assert "revenue_file_missing" in blocker_kinds(api), "three years are no monthly series"
    assert run_stage(api, "mapping")["status"] == "failed"
    api.db["datasets"].docs.clear()
    set_as_of(api, "2025-03-31")
    assert upload(api, "pnl_monthly.csv", PNL_MONTHLY).json()["dtype"] == "pnl"
    decide(api, dtype="pnl", action="confirm", part="revenue_row")
    assert "revenue_file_missing" not in blocker_kinds(api)


# ---------------------------------------------------------------------------
# M4: the Worth flagging line gives both figures and both sources for each month that differs
# ---------------------------------------------------------------------------
def test_m4_the_mismatch_names_each_month_that_differs_with_both_figures_and_both_cell_ranges(rec):
    got = rec.recognise(csv_sheets((FIXTURES / SIDE).read_text(), "pack.csv"), as_of_month=YEAR)
    assert got["series_source"]["gaps"] == [
        {"month": "2025-02", "row_value": 12130, "list_value": 12615, "row_cell": "C3", "list_cells": "J4"},
        {"month": "2025-05", "row_value": 12490, "list_value": 11865, "row_cell": "F3", "list_cells": "J7"}]
    feed = rec.engine_feed(got, "pack.csv", "pack.csv", unit_scale=1000)
    assert feed["cross_check_gaps"][0]["row_value"] == 12_130_000 and feed["cross_check_gaps"][0]["list_value"] == 12_615_000


def test_m4_the_memo_line_counts_the_months_that_differ_not_the_months_compared():
    from claim_matching_support import engine_results
    from test_ic_memo import inputs
    from app import ic_memo
    rec = importlib.import_module("app.recognition")
    got = rec.recognise(csv_sheets((FIXTURES / SIDE).read_text(), "pack.csv"), as_of_month=YEAR, decisions={"series_source": "dated_list"})
    feed = rec.engine_feed(got, "pack.csv", "pack.csv")
    base = engine_results()
    stored = {**base, "recognition": {"file_type": "monthly_revenue_report", "series_basis": "column", "periods": 6, "excluded_columns": 0,
                                       "currency_confirmed": True, "source": "revenue_file", "series_source": "dated_list",
                                       "cross_check": "mismatch", "cross_check_months": feed["cross_check_months"],
                                       "cross_check_gaps": feed["cross_check_gaps"]}}
    text = ic_memo.build_memo(**inputs(stored=stored))
    line = next(x for x in text.split("## Worth flagging")[1].split("## Data gaps")[0].splitlines() if "months differ" in x)
    assert line.startswith("- 2 of 6 months differ")
    assert "February 2025: revenue row" in line and "12,130" in line and "12,615" in line and ", C3)" in line and ", J4)" in line
    assert "May 2025: revenue row" in line and ", F3)" in line and ", J7)" in line
    assert line.count("revenue row") >= 3 and "read from the dated list" in line


# ---------------------------------------------------------------------------
# H7, H8, M1, M2, L4: what is stored, and the size limit measured as MongoDB measures it
# ---------------------------------------------------------------------------
def stored_ok(api):
    """Every stored dataset can be encoded as MongoDB stores it (the fake database applies none of its rules)."""
    import bson
    for doc in api.db["datasets"].docs:
        bson.decode(bson.encode(doc))


def bulk_report(n=1500):
    rows = [f"Customers,{','.join(str((i * 7 + k) % 90) for k in range(12))}" for i in range(n)]
    return ("Management pack (EUR)\nMetric," + MONTHS + "\nRecognised revenue," + ",".join(str(1000 + k) for k in range(12)) + "\n"
            + "\n".join(rows) + "\n").encode()


def test_h7_the_limit_is_measured_in_bson_bytes_so_a_file_that_passes_the_json_check_is_refused(api, monkeypatch):
    import bson
    import server
    rules_only(api, YEAR)
    assert upload(api, "bulk.csv", bulk_report()).status_code == 200
    doc = api.db["datasets"].docs[0]
    json_size, bson_size = len(json.dumps(doc, default=str)), len(bson.encode(doc))
    assert json_size * 1.5 < bson_size, "a grid of small integers is much larger in BSON than in JSON"
    api.db["datasets"].docs.clear()
    monkeypatch.setattr(server, "DATASET_MAX_BYTES", (json_size + bson_size) // 2)
    res = upload(api, "bulk.csv", bulk_report())
    assert res.status_code == 413 and res.json()["detail"] == server.DATASET_TOO_LARGE, "not 'Unexpected server error (DocumentTooLarge)'"
    assert api.db["datasets"].docs == []


def report_with_time_cells():
    import datetime
    from test_revenue_recognition import MONTHLY
    wb = Workbook()
    ws = wb.active
    ws.title = "Report"
    for r, line in enumerate((FIXTURES / MONTHLY).read_text().splitlines(), 1):
        for c, text in enumerate(line.split(","), 1):
            if text:
                ws.cell(r, c, float(text) if text.replace(".", "").isdigit() else text)
    ws.cell(1, 16, datetime.time(9, 30))
    ws.cell(2, 16, datetime.timedelta(hours=30))
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def test_h8_a_time_and_a_duration_in_a_recognised_sheet_are_stored_as_plain_values(api):
    import datetime
    rules_only(api, "2025-06-30")
    content = report_with_time_cells()
    grid = cr.read_workbook(content, "pack.xlsx")[0].grid
    kinds = {type(v) for r in grid for v in r}
    assert datetime.time in kinds and datetime.timedelta in kinds, "the fixture holds a time and a duration cell"
    res = upload(api, "pack.xlsx", content)
    assert res.status_code == 200, res.text
    stored_ok(api)
    kept = api.db["datasets"].docs[0]["recognition_sheets"][0]["grid"]
    assert "09:30:00" in [v for r in kept for v in r] and 108000.0 in [v for r in kept for v in r]
    assert res.json()["recognition"]["recognised"]


def three_sheet_workbook():
    from test_revenue_recognition import MONTHLY
    wb = Workbook()
    ws = wb.active
    ws.title = "Report"
    for r, line in enumerate((FIXTURES / MONTHLY).read_text().splitlines(), 1):
        for c, text in enumerate(line.split(","), 1):
            if text:
                ws.cell(r, c, float(text) if text.replace(".", "").isdigit() else text)
    for name, rows in (("Call log", [["Caller", "Minutes"], ["Alder Works", 12], ["Birch Ltd", 7]]),
                       ("Staff list", [["Person", "Team"], ["Jane Doe", "Sales"], ["Max Roe", "Ops"]])):
        other = wb.create_sheet(name)
        for r, row in enumerate(rows, 1):
            for c, v in enumerate(row, 1):
                other.cell(r, c, v)
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def test_m1_only_the_recognised_sheets_are_stored_not_the_other_tabs_of_the_workbook(api):
    rules_only(api, "2025-06-30")
    assert upload(api, "pack.xlsx", three_sheet_workbook()).status_code == 200
    doc = api.db["datasets"].docs[0]
    assert [s["name"] for s in doc["recognition_sheets"]] == ["Report"]
    blob = json.dumps(doc, default=str)
    assert "Jane Doe" not in blob and "Alder Works" not in blob and "Call log" not in blob


INVOICES = ("Customer,Invoice Date,Amount,Currency\n" + "".join(
    f"Customer {c},2025-{m:02d}-05,{1000 + 10 * c + m},EUR\n" for m in range(1, 7) for c in range(1, 5))).encode()


def test_d3_an_invoice_file_keeps_its_rows_and_stores_no_sheet_and_still_reads_its_recognition_back(api):
    rules_only(api, "2025-06-30")
    body = upload(api, "invoices.csv", INVOICES).json()
    doc = api.db["datasets"].docs[0]
    assert doc["row_count"] == 24 and doc["recognition_sheets"] == [] and doc["recognition_facts"]["month_count"] == 6
    assert "Customer 1" not in json.dumps(doc["recognition_facts"])
    again = api.get(f"/api/audits/{AUDIT}/datasets").json()["datasets"][0]["recognition"]
    assert again == body["recognition"] and again["file_type"] == "invoice_transactions" and again["currency"] == {"code": "EUR", "state": "seen"}


def test_d3_an_invoice_file_without_a_currency_column_still_confirms_its_currency_from_the_facts(api):
    rules_only(api, "2025-06-30")
    upload(api, "no_currency.csv", INVOICES.replace(b",Currency", b"").replace(b",EUR", b""))
    assert api.db["datasets"].docs[0]["recognition_sheets"] == []
    done = decide(api, action="change", part="currency", currency="EUR")
    assert done.status_code == 200 and done.json()["recognition"]["currency"] == {"code": "EUR", "state": "you"}


def test_m2_every_upload_over_the_limit_returns_the_message_a_crm_export_included(api, monkeypatch):
    import server
    rules_only(api, YEAR)
    monkeypatch.setattr(server, "DATASET_MAX_BYTES", 500)
    from test_mapping_miss import FIXTURES as MAPPING_FIXTURES
    res = upload(api, "hubspot_deals.csv", (MAPPING_FIXTURES / "hubspot_deals.csv").read_bytes())
    assert res.status_code == 413 and res.json()["detail"] == server.DATASET_TOO_LARGE
    assert api.db["datasets"].docs == []
    again = upload(api, "unknown.csv", (MAPPING_FIXTURES / "hubspot_deals.csv").read_bytes(), dtype="crm")
    assert again.status_code == 413


def test_l4_a_file_over_the_limit_is_refused_before_the_model_is_called_or_anything_is_kept(api, monkeypatch):
    import server
    set_as_of(api, "2025-06-30")                  # consent given: the model would be asked for a file the rules cannot type
    monkeypatch.setattr(server, "DATASET_MAX_BYTES", 500)
    res = upload(api, MONTHLY, monthly_bytes())
    assert res.status_code == 413 and res.json()["detail"] == server.DATASET_TOO_LARGE
    assert adapter_of(api).calls == 0 and api.db["column_mappings"].docs == [] and api.db["datasets"].docs == []
    assert "recognition" not in api.db["audits"].docs[0]["usage"]["steps"] or not api.db["audits"].docs[0]["usage"]["steps"]["recognition"]


# ---------------------------------------------------------------------------
# M3: after "Use the dated list" the last actual month can still be confirmed or changed
# ---------------------------------------------------------------------------
def test_m3_the_dated_list_is_read_up_to_the_as_of_month_and_the_last_actual_month_is_no_decision(api):
    rules_only(api, "2025-06")
    got = upload(api, SIDE, (FIXTURES / SIDE).read_bytes()).json()["recognition"]
    assert (got["last_actual"]["state"], got["last_actual"]["reason"]) == ("auto", None)
    chosen = decide(api, action="change", part="series_source", series_source="dated_list").json()["recognition"]
    assert chosen["series_basis"] == "column" and chosen["last_actual"]["state"] == "auto"
    assert chosen["last_actual"]["as_of"] == "2025-06" and chosen["series"][-1]["month"] <= "2025-06"
    assert decide(api, action="change", part="last_actual", column=7).status_code == 422


# ---------------------------------------------------------------------------
# M5: the filter that refuses tester notes quoting file text also reads the recognised sheets
# ---------------------------------------------------------------------------
SUMMARY = b"Revenue by product line (EUR),,,\nProduct line,Jan-25,Feb-25,Mar-25\nPlatform licences,10,11,12\nSupport plans,20,21,22\nTraining,30,31,32\n"


def note_on_crm(api, note):
    from test_mapping_miss import FIXTURES as MAPPING_FIXTURES
    rules_only(api, "2025-03-31")
    assert upload(api, "summary.csv", SUMMARY).json()["recognition_only"] is True
    upload(api, "hubspot_deals.csv", (MAPPING_FIXTURES / "hubspot_deals.csv").read_bytes())
    return api.post(f"/api/audits/{AUDIT}/datasets/crm/decisions", json=[
        {"column": "Deal owner", "action": "correct", "field": None, "reason": "other", "note": note}])


def test_m5_a_note_quoting_a_label_of_a_file_read_by_recognition_only_is_refused(api):
    import server
    res = note_on_crm(api, "Platform licences is wrong")
    assert res.status_code == 400 and res.json()["detail"] == server.NOTE_REFUSED, "refused for the quoted label, not for another reason"
    assert api.db["audits"].docs[0]["usage"]["other_notes"] == []


def test_m5_a_note_that_quotes_nothing_from_the_file_is_still_kept(api):
    assert note_on_crm(api, "That column is not needed").status_code == 200


# ---------------------------------------------------------------------------
# Low findings
# ---------------------------------------------------------------------------
def test_l1_a_column_headed_name_is_entity_labelled_whatever_its_values_and_its_labels_are_masked(rec):
    names = ["Revenue Partners", "Cash Converters", "Sales Hub", "Cost Cutters", "Margin Makers", "Churn Chasers"]
    for head in ("Name", "Customer", "Client", "Account", "Company"):
        text = sheet_text("Revenue by customer (EUR),,,", f"{head},Jan-25,Feb-25,Mar-25", [f"{n},{10 + i},{11 + i},{12 + i}" for i, n in enumerate(names)])
        got = run(rec, text, as_of="2025-03-31")
        assert got["label_mode"] != "metric", head
        sent = rec.model_text(csv_sheets(text), None, {}, (), "0" * 64)
        assert not any(n in sent for n in names), f"{head}: a customer name reached the text"
    text = sheet_text("Revenue by metric (EUR),,,", "Metric name,Jan-25,Feb-25,Mar-25", [f"{n},{10 + i},{11 + i},{12 + i}" for i, n in enumerate(names)])
    assert run(rec, text, as_of="2025-03-31")["label_mode"] == "metric", "only a column headed Name itself is masked"


def test_l3_a_file_whose_amounts_show_two_currencies_asks_for_no_single_currency_and_compute_stops(api):
    rules_only(api, "2024-12-31")
    from test_chat_upload import MESSY
    body = upload(api, "mixed_currency.csv", (MESSY / "mixed_currency.csv").read_bytes()).json()
    assert body["recognition"]["currency"] == {"code": None, "state": "confirm", "mixed": True, "codes": ["EUR", "USD"]}
    assert body["missing_required"] == ["currency"] and not body.get("currency_pending")
    assert decide(api, action="change", part="currency", currency="EUR").status_code == 400
    assert api.post(f"/api/audits/{AUDIT}/compute").status_code == 409


def test_l3_amounts_that_all_show_the_same_sign_name_the_currency(rec):
    text = "Customer,Invoice Date,Amount\n" + "".join(f"Customer {c},2025-0{m}-05,\"€1,2{c}{m}\"\n" for m in range(1, 4) for c in range(1, 4))
    got = rec.recognise(csv_sheets(text), as_of_month="2025-03-31")
    assert got["currency"] == {"code": "EUR", "state": "seen"}


def test_l5_the_dated_list_cites_only_the_rows_up_to_the_boundary(rec):
    got = rec.recognise(csv_sheets((FIXTURES / SIDE).read_text(), "pack.csv"), as_of_month="2025-03-31",
                        decisions={"series_source": "dated_list", "last_actual": 4})
    assert [s["month"] for s in got["series"]] == ["2025-01", "2025-02", "2025-03"]
    assert got["series_rows"] == [3, 4, 5], "rows 6 to 8 are after the boundary"


def test_l6_the_recognition_of_unchanged_inputs_is_not_computed_again(api, monkeypatch):
    import server
    rules_only(api, "2025-03-31")
    upload(api, "summary.csv", SUMMARY)
    calls = []
    original = server.rec.recognise
    monkeypatch.setattr(server.rec, "recognise", lambda *a, **k: calls.append(1) or original(*a, **k))
    server._RECOGNITION_CACHE.clear()
    for _ in range(3):
        assert api.get(f"/api/audits/{AUDIT}/datasets").status_code == 200
        assert api.get(f"/api/audits/{AUDIT}/blockers").status_code == 200
    assert len(calls) == 1, "read once for six requests"
    decide(api, action="change", part="file_type", file_type="customer_revenue")
    api.get(f"/api/audits/{AUDIT}/datasets")
    assert len(calls) == 2, "a decision changes the inputs and is read again"


def test_l6_the_audit_view_does_not_load_the_stored_sheets(api, monkeypatch):
    rules_only(api, "2025-03-31")
    upload(api, "summary.csv", SUMMARY)
    seen = []
    original = api.db["datasets"].find

    def spy(query=None, projection=None, *a, **k):
        seen.append(projection)
        return original(query, projection, *a, **k) if projection is not None else original(query)
    monkeypatch.setattr(api.db["datasets"], "find", spy)
    assert api.get(f"/api/audits/{AUDIT}").status_code == 200
    assert any(p and p.get("recognition_sheets") == 0 for p in seen)


def test_l2_a_version_4_audit_is_stale_until_run_audit_calculates_it_again_with_the_real_engine(api):
    """The contract test stubs the engine; this one runs it: results stored before the recognition block are marked stale on a read
    (George, 2026-10-20) and carry the version 5 fields once the Run audit engine stage has calculated again."""
    import copy
    rules_only(api, "2025-06-30")
    assert upload(api, MONTHLY, monthly_bytes()).status_code == 200
    assert decide(api, action="confirm", part="pending").status_code == 200
    assert run_stage(api, "engine")["status"] == "done"
    audit = api.db["audits"].docs[0]
    old = copy.deepcopy(audit["results"])
    for key in ("recognition",):
        old.pop(key, None)
    for key in ("basis", "file_type", "growth_label"):
        old["revenue_series"].pop(key, None)
    (old.get("revenue_reconciliation") or {}).pop("status", None)
    audit["results"] = {**old, "contract_version": 4}
    res = api.get(f"/api/audits/{AUDIT}/results")
    assert res.status_code == 409 and res.json()["detail"]["message"] == "Stale – re-run audit"
    assert api.get(f"/api/audits/{AUDIT}/review").json()["engine_stale"] is True
    assert run_stage(api, "engine")["status"] == "done"
    res = api.get(f"/api/audits/{AUDIT}/results")
    assert res.status_code == 200
    again = res.json()["results"]
    from schemas.metrics import CONTRACT_VERSION
    assert again["contract_version"] == CONTRACT_VERSION == 12 and again["recognition"]["series_basis"] == "row"
    assert again["revenue_series"]["basis"] == "row" and again["revenue_series"]["data"]
