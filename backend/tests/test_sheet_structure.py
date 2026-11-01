"""Sheet structure reading and yearly-basis metrics (docs/specs/chat-upload.md section 26). A P&L the rules cannot map by columns is
read as a statement: Python reads the skeleton, the model labels lines through the gateway, Python checks every total and builds
what the map feeds. Every test runs on the in-memory Mongo stub and the fake adapter: nothing reaches a provider. The fixtures in
tests/fixtures/mapping/ are synthetic (rule 15)."""
import io
import json
from pathlib import Path

import pytest

pytest.importorskip("fastapi")

import growth_engine as ge  # noqa: E402
import server  # noqa: E402
import test_llm_gateway as t  # noqa: E402
from app import column_rules as cr  # noqa: E402
from app.llm import gateway  # noqa: E402
from app.structures import redact  # noqa: E402
from app.structures import pnl_lines as pl  # noqa: E402
from app.structures import sheet as sh  # noqa: E402
from test_chat_upload import AUDIT, COMPANY, upload, view  # noqa: E402

MAPPING = Path(__file__).resolve().parent / "fixtures" / "mapping"
SAGE, KPI, YEARLY, CUSTOMERS = "sage_statement.csv", "monthly_kpi.csv", "yearly_pnl_by_product.csv", "customer_list.csv"
KPI_NO_REVENUE = "kpi_no_revenue.csv"      # customers, churn and subscriptions: no revenue, MRR or ARR row (revenue-file-recognition.md decision 3)
NAMES = ("Alder Works", "Birch Labs", "Jasmine Tea", "Heather Spa")


def sheet_of(name: str, content: bytes = None) -> cr.Sheet:
    return cr.read_workbook(content if content is not None else (MAPPING / name).read_bytes(), name)[0]


def reading(name: str, content: bytes = None, fye: int = 12) -> dict:
    s = sheet_of(name, content)
    return sh.read(s.grid, s.fmt, fye)


def lines_by_label(view: dict) -> dict:
    return {line["label"]: line for line in view["lines"]}


def stored_map(name: str, reply=None) -> dict:
    return sh.build(reading(name), reply)


def put(api, name: str, **params):
    return upload(api, name, (MAPPING / name).read_bytes(), **params)


def sent_texts(api) -> list:
    return [json.loads(p)["text"] for p in api.state["adapter"].payloads if json.loads(p).get("type") == "sheet_structure"]


# A recorded reply for the yearly P&L: it reads "Sales" as revenue and the margin % line as money; Python's rules win (26.2).
YEARLY_REPLY = json.dumps({"type": "sheet_structure", "header_row": 2, "label_col": 1,
                           "periods": [{"col": c, "grain": "year", "start": f"{2019 + c}-01", "tag": "forecast" if c > 4 else "actual"}
                                       for c in range(2, 8)],
                           "rows": [{"row": 3, "role": "total", "parent_row": None, "meaning": "revenue"},
                                    {"row": 4, "role": "subline", "parent_row": 3, "meaning": "revenue"},
                                    {"row": 5, "role": "subline", "parent_row": 3, "meaning": "revenue"},
                                    {"row": 6, "role": "other", "parent_row": None, "meaning": "cost_of_revenue"},
                                    {"row": 7, "role": "other", "parent_row": None, "meaning": "gross_margin"},
                                    {"row": 8, "role": "other", "parent_row": None, "meaning": "gross_margin"},
                                    {"row": 9, "role": "total", "parent_row": None, "meaning": "opex"},
                                    {"row": 10, "role": "total", "parent_row": 9, "meaning": "revenue"},
                                    {"row": 11, "role": "subline", "parent_row": 10, "meaning": "sm_expense"},
                                    {"row": 12, "role": "subline", "parent_row": 10, "meaning": "sm_expense"},
                                    {"row": 13, "role": "subline", "parent_row": 9, "meaning": "sm_expense"},
                                    {"row": 14, "role": "subline", "parent_row": 9, "meaning": "opex"},
                                    {"row": 15, "role": "other", "parent_row": None, "meaning": "ebitda"}]})


# ---------------------------------------------------------------------------
# 26.2: Python's reading of the three statements of the task
# ---------------------------------------------------------------------------
def test_a_sage_statement_reads_one_period_only_and_asks_for_a_monthly_export():
    read = reading(SAGE)
    assert (read["header_row"], read["label_col"]) == (4, 2), "the account codes are numbers; the descriptions are the labels"
    assert [(p["header"], p["grain"]) for p in read["periods"]] == [("Period To Date", "period_to_date"), ("Year To Date", "ytd")]
    view = sh.view(sh.build(read), 12, 1)
    assert view["one_period"] and view["grain"] == "yearly" and not view["monthly_pnl"]
    assert sh.feed(sh.build(read), 12, 1, {}) == {"pnl": None, "rows": [], "revenue_row": None, "yearly": None,
                                                  "lines": {"basis": "yearly", "findings": [], "plan": []}}
    lines = lines_by_label(view)
    assert lines["Total Sales"]["role"] == "total" and lines["Sales Type A"]["parent_row"] == 7, "a Total adds up the lines above it"
    assert lines["Total Overheads"]["checks"] == {} and lines["Salaries"]["parent_row"] == 16


def test_a_monthly_kpi_sheet_stores_its_series_and_checks_the_sub_lines_against_the_total():
    m = stored_map(KPI)
    view = sh.view(m, 12, 1)
    lines = lines_by_label(view)
    assert lines["Total MRR"]["checks"] == {"total_mismatch": ["Jun-25"]}, "Jun-25: 1,500 + 6,000 + 10,000 is not 17,600"
    assert [lines[f"MRR - {k}"]["parent_row"] for k in ("Starter", "Pro", "Enterprise")] == [6, 6, 6]
    series = sh.kpi_series(m["lines"], sh.with_periods(m["periods"], 12), 1)
    assert [(s["meaning"], s["row"], s["label"]) for s in series] == [
        ("mrr", 6, "Unverified"), ("new_arr", 7, "Unverified"), ("new_arr", 8, "Unverified"), ("churn_rate", 9, "Unverified")]
    assert series[0]["values"][0] == {"month": "2025-01", "value": 15000, "cell": "B6"}
    assert lines["MRR growth % MoM"]["role"] == "rate" and not lines["MRR growth % MoM"]["series"]
    assert view["grain"] == "monthly" and not view["monthly_pnl"]
    assert view["missing_fields"] == ["revenue", "cost_of_revenue", "sm_expense"]
    assert sh.feed(m, 12, 1, {})["pnl"] is None, "KPI series never feed the engine"


def test_a_yearly_pnl_by_product_reads_its_header_tags_sections_and_checks():
    read = reading(YEARLY)
    assert read["header_row"] == 2, "row 1 holds the title and the forecast tags; row 2 the six FY periods"
    assert [(p["header"], p["tag"]) for p in read["periods"]] == [
        ("FY2021", "actual"), ("FY2022", "actual"), ("FY2023", "actual"),
        ("FY2024", "forecast"), ("FY2025", "forecast"), ("FY2026", "forecast")]
    view = sh.view(sh.build(read), 12, 1000)
    lines = lines_by_label(view)
    assert lines["Sales"]["meaning"] == "sm_expense", "Sales inside the operating expenses is S&M, never revenue"
    assert lines["Sales team"]["meaning"] == lines["Sales tools"]["meaning"] == "sm_expense"
    assert lines["Gross margin as % of revenue"]["role"] == "rate"
    assert lines["Operating expenses"]["checks"] == {"no_breakdown": ["FY2021", "FY2022"]}
    assert lines["Revenue"]["checks"] == {} and lines["Sales"]["checks"] == {}
    assert view["grain"] == "yearly" and view["yearly"] and not view["monthly_pnl"]


def test_the_yearly_pnl_is_never_split_into_months_and_its_actual_years_are_the_yearly_input():
    m = stored_map(YEARLY)
    fed = sh.feed(m, 12, 1000, {"file": YEARLY, "sheet": "CSV"}, "2023-12")     # no as-of month: no actual year (is_actual)
    assert fed["pnl"] is None, "no month column: nothing reaches the engine as a monthly P&L"
    years = {y["period"]: y["fields"] for y in fed["yearly"]["actual"]}
    assert list(years) == ["FY2021", "FY2022", "FY2023"], "only the actual-tagged year columns"
    assert years["FY2023"]["sm_expense"] == {"value": 550000.0, "cells": ["D10", "D13"], "scope": "sales and marketing"}, "Sales + Marketing, ×1,000"
    assert "sm_expense" not in years["FY2021"], "a blank breakdown is no figure, never 0"
    assert [p["period"] for p in fed["yearly"]["planned"]] == ["FY2024", "FY2025", "FY2026"]


def test_the_forecast_columns_alone_become_claims_each_with_its_tag():
    found = sh.claims(stored_map(YEARLY), {"sheet": "CSV", "unit_scale": 1000, "title_currency": "EUR"}, 12)
    revenue = next(line for line in found["lines"] if line["label"] == "Revenue")
    assert [(v["target_date"], v["value"], v["cell"], v["tag"]) for v in revenue["values"]] == [
        ("2024", 1700000, "E3", "forecast"), ("2025", 2200000, "F3", "forecast"), ("2026", 2800000, "G3", "forecast")]
    assert all(v["target_date"] >= "2024" for line in found["lines"] for v in line["values"]), "no actual year is a claim"
    assert not any(line["label"] in ("Sales", "Marketing", "Operating expenses") for line in found["lines"]), "cost lines"


FOUR_FORECASTS = (MAPPING / YEARLY).read_bytes().replace(b",,,,forecast,forecast,forecast", b",,,forecast,forecast,forecast,forecast")


def test_a_yearly_pnl_gives_revenue_gross_profit_and_ebitda_claims_plus_the_sm_sum_and_no_rate_claim():
    """The yearly P&L with four forecast years (FY2023 to FY2026): 3 lines x 4 years = 12, and 4 S&M sums."""
    m = sh.build(reading("yearly_four_forecasts.csv", FOUR_FORECASTS))
    found = sh.claims(m, {"sheet": "CSV", "unit_scale": 1000, "title_currency": "EUR"}, 12)
    by_label = {line["label"]: line for line in found["lines"]}
    sm = by_label[sh.SM_CLAIM_LABEL]
    assert sorted(by_label) == sorted(["Revenue", "Gross profit", "EBITDA", sh.SM_CLAIM_LABEL]), "no rate line is a claim"
    assert sum(len(line["values"]) for label, line in by_label.items() if label != sh.SM_CLAIM_LABEL) == 12
    assert len(sm["values"]) == 4 and found["count"] == 16
    assert [(v["target_date"], v["value"], v["cell"], v["tag"]) for v in sm["values"]] == [
        ("2023", 550000, "D10+D13", "forecast"), ("2024", 660000, "E10+E13", "forecast"),
        ("2025", 830000, "F10+F13", "forecast"), ("2026", 1000000, "G10+G13", "forecast")]
    assert (sm["claim_type"], sm["currency"]) == ("costs", "EUR")


def test_the_sm_claim_is_cited_to_its_cells_and_saved_as_a_sheet_candidate():
    from app import forecast_claims as fc
    m = sh.build(reading("yearly_four_forecasts.csv", FOUR_FORECASTS))
    ds = {"sheet": "CSV", "unit_scale": 1000, "title_currency": "EUR", "file": "p.csv", "file_hash": "h", "dtype": "pnl"}
    found = sh.claims(m, ds, 12)
    (made,) = fc.candidates(found, [next(x["row"] for x in found["lines"] if x["label"] == sh.SM_CLAIM_LABEL)], ds, 12)
    assert made["snippet"] == sh.SM_CLAIM_LABEL and made["claim_type"] == "costs" and made["origin"] == fc.ORIGIN
    assert [s["cell"] for s in made["sources"]] == ["D10+D13", "E10+E13", "F10+F13", "G10+G13"]


def test_a_line_with_the_role_rate_is_never_a_claim_whatever_its_label_says():
    m = stored_map(YEARLY)
    for line in m["lines"]:
        if line["label"] == "Revenue":
            line["role"] = "rate"       # the analyst (or the reply) calls the revenue line a rate
    found = sh.claims(m, {"sheet": "CSV", "unit_scale": 1000, "title_currency": "EUR"}, 12)
    assert "Revenue" not in [line["label"] for line in found["lines"]]
    assert "Gross margin as % of revenue" not in [line["label"] for line in found["lines"]], "gross margin % included"


def test_a_customer_list_column_is_withheld_and_counted():
    text, counts = sh.skeleton(reading(CUSTOMERS), COMPANY, {}, (COMPANY,))
    assert not any(name in text for name in NAMES)
    assert "r2c1: Customer" in text and "c1 withheld: text, 36 distinct values" in text.splitlines()
    assert counts["columns"] == 1 and redact.sheet_text_problem(text) is None


@pytest.mark.parametrize("name", [SAGE, KPI, YEARLY, CUSTOMERS])
def test_the_skeleton_holds_every_label_and_no_figure(name):
    read = reading(name)
    text, _ = sh.skeleton(read, COMPANY, {}, (COMPANY,))
    assert redact.sheet_text_problem(text) is None
    cells = redact.parse_sheet_text(text)["cells"]
    assert not any(redact.sheet_figure(c["text"]) for c in cells)
    for line in read["lines"]:
        values = [v for v in line["values"] if v is not None]
        assert not any(f"{v:g}" in [c["text"] for c in cells] for v in values), "a figure is num in its row line, never a cell"
    if name != CUSTOMERS:
        assert all(f"r{line['row']}c{read['label_col']}: " in text for line in read["lines"]), "every row label is sent"


def test_a_figure_in_a_title_row_or_a_body_row_is_never_a_cell_line():
    content = (b"Plan,1250000,,\nLine,FY2023,FY2024,FY2025\nRevenue,1,2,3\nCustomers,1.250,40,\"1,300\"\n"
               b"Spend,2015,2016,2017\n")
    text, _ = sh.skeleton(reading("f.csv", content), COMPANY, {}, ())
    cells = [c["text"] for c in redact.parse_sheet_text(text)["cells"]]
    assert cells == ["Plan", "Line", "FY2023", "FY2024", "FY2025", "Revenue", "Customers", "Spend"]
    assert "r1 indent 0 num c2 blank c3-c4" in text.splitlines()
    assert "r5 indent 0 num c2-c4" in text.splitlines(), "below the header a year-like figure is a number, never a period"


def test_the_company_name_and_an_email_are_masked_in_the_skeleton():
    content = (MAPPING / SAGE).read_bytes().replace(b"Profit and Loss,,,", b"Profit and Loss for ops@fixture-target.example,,,")
    text, counts = sh.skeleton(reading(SAGE, content), COMPANY, {}, (COMPANY,))
    assert COMPANY not in text and "ops@fixture-target.example" not in text
    assert "r1c1: [redacted]" in text and "[email]" in text and (counts["company"], counts["email"]) == (1, 1)


def test_past_400_rows_only_the_first_400_are_read():
    body = "\n".join(f"Line {i},{i},{i + 1},{i + 2}" for i in range(450))
    read = reading("long.csv", f"Line,FY2021,FY2022,FY2023\n{body}\n".encode())
    assert read["too_long"] and len({c["row"] for c in read["cells"]}) == 400
    assert sh.view(sh.build(read), 12, 1)["too_long"]


def test_a_merged_forecast_cell_and_excel_indents_are_read_from_the_workbook():
    from openpyxl import Workbook
    from openpyxl.styles import Alignment
    wb = Workbook()
    ws = wb.active
    for row in [["P&L"], ["Line", "FY2022", "FY2023", "FY2024", "FY2025", "FY2026"], ["Revenue", 10, 12, 15, 18, 22],
                ["Product A", 4, 5, 6, 7, 9], ["Product B", 6, 7, 9, 11, 13]]:
        ws.append(row)
    ws["D1"] = "Forecast"
    ws.merge_cells("D1:E1")
    for cell in ("A4", "A5"):
        ws[cell].alignment = Alignment(indent=1)
    buf = io.BytesIO()
    wb.save(buf)
    read = reading("merged.xlsx", buf.getvalue())
    assert [p["tag"] for p in read["periods"]] == ["actual", "actual", "forecast", "forecast", "unknown"], \
        "the merged cell tags its span and no column after it"
    assert [line["indent"] for line in read["lines"]] == [0, 1, 1]
    assert lines_by_label(sh.view(sh.build(read), 12, 1))["Product A"]["parent_row"] == 3


def test_a_tag_in_the_header_and_a_tag_carried_to_the_right():
    content = b"Line,2023A,2024A,2025 Budget,2026\nRevenue,1,2,3,4\n"
    assert [p["tag"] for p in reading("t.csv", content)["periods"]] == ["actual", "actual", "budget", "budget"]
    content = b",Actual,,Forecast,\nLine,2023,2024,2025,2026\nRevenue,1,2,3,4\n"
    assert [p["tag"] for p in reading("t.csv", content)["periods"]] == ["actual", "actual", "forecast", "forecast"]
    content = b"Line,2023,2024,2025\nRevenue,1,2,3\n"
    assert [p["tag"] for p in reading("t.csv", content)["periods"]] == ["unknown"] * 3, "no tag in the sheet: unknown"


def test_three_dates_of_one_data_row_are_no_period_header():
    sheet = sheet_of("pipedrive_deals.csv")
    assert sheet.header_row == 1, "the header stays the first row: a data row with numbers holds no period header"
    content = b"Deal,Amount,Created,Expected close,Closed\nD1,500,2024-01-04,2024-02-28,2024-03-25\nD2,700,2024-02-01,2024-03-30,\n"
    assert sheet_of("deals.csv", content).header_row == 1, "three months in one data row are no period header either"


def test_rules_first_the_column_rules_cover_a_column_pnl_and_not_a_statement():
    assert sh.rules_cover(cr.analyse("pnl", sheet_of("accountant_pnl.csv"))[0])
    assert not sh.rules_cover(cr.analyse("pnl", sheet_of(YEARLY))[0])


# ---------------------------------------------------------------------------
# 26.4: the reply, and the rules that win over it
# ---------------------------------------------------------------------------
def test_the_rules_win_over_the_reply_for_sales_in_a_cost_section_and_for_a_rate_line():
    read = reading(YEARLY)
    lines = {line["label"]: line for line in sh.label_lines(read, json.loads(YEARLY_REPLY))}
    assert lines["Sales"]["meaning"] == "sm_expense" and lines["Sales"]["source"] == "rules"
    assert lines["Gross margin as % of revenue"]["role"] == "rate", "the reply's money margin line is a rate"
    assert lines["Cost of revenue"]["source"] == "rules", "the reply agrees with Python: Python's confidence stands"


def test_python_keeps_its_periods_and_the_header_with_more_periods():
    read = reading(YEARLY)
    wrong = {**json.loads(YEARLY_REPLY), "header_row": 1,
             "periods": [{"col": 2, "grain": "month", "start": "2021-01", "tag": "forecast"}]}
    again = sh.merged(sheet_of(YEARLY).grid, {}, 12, read, wrong)
    assert again["header_row"] == 2, "row 1 holds no period: Python's header stays"
    assert [(p["grain"], p["tag"]) for p in again["periods"]][0] == ("year", "actual"), "a period is rebuilt from its header cell"


def test_a_monthly_statement_feeds_the_engine_with_sm_as_the_sum_of_the_sales_and_marketing_totals():
    content = (b"Monthly P&L,,,forecast\n"
               b"Line,Jan-25,Feb-25,Mar-25\n"
               b"Revenue,100,110,120\nCost of sales,-30,-33,-36\n"
               b"Operating expenses,,,\n  Sales,20,21,22\n    Sales salaries,15,16,17\n    Sales tools,5,5,5\n"
               b"  Marketing,10,11,12\n")
    m = sh.build(reading("m.csv", content))
    fed = sh.feed(m, 12, 1, {})
    assert fed["pnl"] == [{"month": "2025-01", "revenue": 100.0, "cost_of_revenue": 30.0, "sm_expense": 30.0},
                          {"month": "2025-02", "revenue": 110.0, "cost_of_revenue": 33.0, "sm_expense": 32.0}], \
        "the forecast month is no engine input; costs written as negatives are read as positive; S&M = Sales + Marketing"
    assert fed["rows"] == [3, 4, 6, 9] and fed["revenue_row"] == 3
    view = sh.view(m, 12, 1)
    assert view["monthly_pnl"] and view["grain"] == "monthly"
    assert lines_by_label(view)["Cost of sales"]["checks"] == {"sign_convention": []}
    frame = server.statement_pnl(fed)
    assert list(frame.columns) == ["_row", "month", "revenue", "cost_of_revenue", "sm_expense"]


# ---------------------------------------------------------------------------
# 26.8: yearly-basis metrics (engine, Python only)
# ---------------------------------------------------------------------------
def _yearly_block(name=YEARLY, content=None, results=None, series=None, as_of="2023-12"):
    read = reading(name, content)
    fed = sh.feed(sh.build(read), 12, cr.read_units(read["titles"])[0], {"file": name, "sheet": "CSV"}, as_of)
    return ge.compute_yearly_basis(fed["yearly"], results or {}, series)


def test_the_yearly_basis_metrics_of_the_actual_years_each_labelled():
    block = _yearly_block()
    years = {y["period"]: y for y in block["years"]}
    assert list(years) == ["FY2021", "FY2022", "FY2023"]
    y23 = years["FY2023"]
    assert y23["sm_pct_revenue"] == round(550 / 1300, 4)
    assert y23["gross_margin_pct"] == round((1300 - 360) / 1300, 4)
    assert y23["gross_margin_change"] == round((1300 - 360) / 1300 - (1000 - 290) / 1000, 4)
    assert y23["new_revenue"] == 300000
    assert y23["cac_payback_months"] == round(12 * 550000 / (300000 * (1300 - 360) / 1300), 2)
    assert years["FY2022"]["cac_payback_months"] is None and years["FY2022"]["cac_payback_reason"] == "needs a sales_marketing row"
    assert all(y["basis"] == "yearly basis, not monthly" for y in block["years"]) and block["basis"] == "yearly basis, not monthly"
    assert "yearly basis, not monthly" in y23["source"]["rule"] and "D10" in y23["source"]["rule"]
    assert [m["text"] for m in block["missing"]] == ["Missing: monthly P&L / revenue by customer, last 24 months"] * 5


def test_cac_payback_is_shown_only_when_new_revenue_is_positive():
    content = b"Line,2022A,2023A\nRevenue,1000,900\nCost of revenue,300,300\nSales and marketing,200,250\n"
    (y22, y23) = _yearly_block("d.csv", content)["years"]
    assert (y23["new_revenue"], y23["cac_payback_months"], y23["cac_payback_reason"]) == (-100, None, "new revenue not positive")
    assert y22["cac_payback_reason"] == "no revenue for the year before"


def test_the_track_record_compares_a_years_budget_with_its_actual():
    content = b"Line,2023 Budget,2023 Actual,2024 Forecast\nRevenue,1000,900,1200\n"
    (row,) = _yearly_block("b.csv", content)["track_record"]
    assert (row["period"], row["line"], row["tag"], row["forecast"], row["actual"], row["gap"], row["gap_pct"]) == \
        ("2023 Actual", "Revenue", "budget", 1000, 900, -100, -0.1)
    assert row["basis"] == "yearly basis, not monthly"


def test_a_reported_actual_year_is_compared_with_the_revenue_file_when_it_covers_the_year():
    content = b"Line,2023A,2024A\nRevenue,120,130\n"
    months = [f"2024-{m:02d}" for m in range(1, 13)]
    series = {"data": [{"month": m, "total": 10.5} for m in months]}
    (row,) = _yearly_block("r.csv", content, series=series, as_of="2024-12")["reported_vs_file"]
    assert (row["period"], row["reported"], row["revenue_file"], row["gap"]) == ("2024A", 130, 126, -4)


def test_the_metrics_computed_from_the_revenue_file_are_not_listed_missing():
    results = {"nrr": {"overall_pct": 1.0}, "gross_churn": {"overall_pct": 0.1}}
    missing = [m["metric"] for m in _yearly_block(results=results)["missing"]]
    assert missing == ["Integrity tests (revenue reconciliation)", "Median sales cycle", "CAC payback (monthly timing)"]


# ---------------------------------------------------------------------------
# The upload: rules first, the call, the table, decisions, the saved map, compute (26.1, 26.4, 26.6, 26.7)
# ---------------------------------------------------------------------------
def test_a_pnl_the_rules_map_by_columns_makes_no_sheet_call(api):
    out = put(api, "accountant_pnl.csv").json()
    assert out["sheet_map"] is None and out["pnl_grain"] == "monthly"
    assert sent_texts(api) == []


def test_a_yearly_statement_is_read_with_one_call_and_the_reply_labels_the_lines(api):
    api.state["adapter"] = t.FakeAdapter(replies=[YEARLY_REPLY])
    out = put(api, YEARLY, dtype="pnl").json()
    (text,) = sent_texts(api)
    assert redact.sheet_text_problem(text) is None and "r10c1: Sales" in text
    assert out["ai_reading"]["status"] == "read" and out["pnl_grain"] == "yearly" and out["header_row"] == 2
    lines = lines_by_label(out["sheet_map"])
    assert lines["Sales"]["meaning"] == "sm_expense" and lines["Gross margin as % of revenue"]["role"] == "rate"
    assert out["sheet_map"]["counts"]["lines"] > 0 and out["columns"] == [] and out["miss"] is None
    assert out["pending"] == 0 and out["saved"], "every used line agrees with Python at 100: the map saves"
    assert [line["label"] for line in out["sheet_map"]["lines"]][:1] == ["Operating expenses"], "a failed check comes first"
    assert out["forecast_claims"]["count"] == 12, "Revenue, Gross profit and EBITDA (9) and S&M (3) over three forecast years, no rate"


def test_the_company_name_in_a_title_is_masked_in_the_sent_skeleton(api):
    content = (MAPPING / YEARLY).read_bytes().replace(b"Product P&L (EUR thousands)", COMPANY.encode() + b" P&L (EUR thousands)")
    api.state["adapter"] = t.FakeAdapter(replies=[YEARLY_REPLY])
    out = upload(api, YEARLY, content, dtype="pnl").json()
    (text,) = sent_texts(api)
    assert COMPANY not in text and "r1c1: [redacted] P&L (EUR thousands)" in text
    assert out["ai_reading"]["status"] == "read" and out["sheet_map"]["counts"]["company"] == 1


def test_without_consent_the_statement_is_read_by_python_alone(api):
    api.db["audits"].docs[0]["structure_reading_consent"] = False
    out = put(api, YEARLY, dtype="pnl").json()
    assert api.state["adapter"].calls == 0 and out["ai_reading"]["status"] == "no_consent"
    assert lines_by_label(out["sheet_map"])["Sales"]["meaning"] == "sm_expense"


@pytest.mark.parametrize("adapter", [t.FakeAdapter(raise_with=RuntimeError("provider down")),
                                     t.FakeAdapter(replies=[json.dumps({**json.loads(YEARLY_REPLY), "rows": [
                                         {"row": 99, "role": "total", "parent_row": None, "meaning": "revenue"}]})])])
def test_when_the_model_fails_or_names_a_row_not_sent_pythons_reading_stands(api, adapter):
    api.state["adapter"] = adapter
    out = put(api, YEARLY, dtype="pnl").json()
    assert out["ai_reading"]["status"] == "not_read"
    assert lines_by_label(out["sheet_map"])["Sales"]["source"] == "rules"


def test_lines_wait_for_a_click_save_the_map_by_hash_and_reuse_it(api):
    out = put(api, KPI, dtype="pnl").json()
    waiting = [line["row"] for line in out["sheet_map"]["lines"] if line["pending"]]
    assert waiting == [7, 8, 9] and out["pending"] == 3 and not out["saved"], "series lines Python reads at 80 wait"
    assert api.post(f"/api/audits/{AUDIT}/compute").status_code == 409
    r = api.post(f"/api/audits/{AUDIT}/datasets/pnl/lines", json=[{"row": 7, "action": "confirm"}, {"row": 8, "action": "confirm"},
                                                               {"row": 9, "action": "change", "meaning": "churn_rate", "role": "rate"}])
    out = r.json()
    assert r.status_code == 200 and out["pending"] == 0 and out["saved"] and out["version"] == 1
    saved = next(v for v in api.db["column_mappings"].docs if v.get("kind") == "sheet_map")
    assert not any("MRR" in json.dumps(v) or "New ARR" in json.dumps(v) for v in saved["lines"]), "no label text in the saved map"
    calls = api.state["adapter"].calls
    api.db["datasets"].docs.clear()
    again = put(api, KPI, dtype="pnl").json()
    assert api.state["adapter"].calls == calls, "the same bytes: the saved map, no model call"
    assert again["pending"] == 0 and {line["source"] for line in again["sheet_map"]["lines"]} == {"saved"}


def test_a_line_decision_outside_the_lists_or_the_sheet_is_refused(api):
    put(api, KPI, dtype="pnl")
    assert api.post(f"/api/audits/{AUDIT}/datasets/pnl/lines", json=[{"row": 99, "action": "confirm"}]).status_code == 400
    assert api.post(f"/api/audits/{AUDIT}/datasets/pnl/lines",
                    json=[{"row": 7, "action": "change", "meaning": "profit"}]).status_code == 400


def test_compute_reads_the_yearly_statement_and_the_results_carry_the_yearly_basis(api):
    upload(api, "blank_vs_zero.csv")
    put(api, YEARLY, dtype="pnl")
    results = api.post(f"/api/audits/{AUDIT}/compute").json()
    block = results["yearly_basis"]
    assert [y["period"] for y in block["years"]] == ["FY2021", "FY2022", "FY2023"]
    assert all(y["basis"] == "yearly basis, not monthly" for y in block["years"])
    assert results["cac_payback"] is None and "CAC payback (monthly timing)" in [m["metric"] for m in block["missing"]]
    assert results["revenue_reconciliation"] is None, "a yearly statement is no monthly P&L"


def test_compute_reads_a_monthly_statement_as_the_pnl(api):
    upload(api, "blank_vs_zero.csv")
    content = (b"Line,Jan-24,Feb-24,Mar-24,Apr-24\nRevenue,1000,1100,1200,1300\nCost of sales,300,330,360,390\n"
               b"Sales and marketing,200,210,220,230\n")
    out = upload(api, "monthly_statement.csv", content, dtype="pnl").json()
    assert out["sheet_map"]["monthly_pnl"] and out["pnl_grain"] == "monthly" and out["months"]["count"] == 4
    results = api.post(f"/api/audits/{AUDIT}/compute").json()
    assert results.get("yearly_basis") is None
    rec = results["revenue_reconciliation"]
    assert rec is not None and rec["source"]["pnl"]["row_numbers"] == [2], "the revenue line's row"


def test_saved_claims_keep_a_statement_file_an_engine_input(api):
    upload(api, "blank_vs_zero.csv")
    out = put(api, YEARLY, dtype="pnl").json()
    rows = [line["row"] for line in out["forecast_claims"]["lines"]]
    assert api.post(f"/api/audits/{AUDIT}/datasets/pnl/forecast-claims", json={"rows": rows}).json()["saved"] == 12
    assert view(api, "pnl")["read_as_claims"]["claims"] == 12
    assert api.post(f"/api/audits/{AUDIT}/compute").json()["yearly_basis"]["years"]


def test_delete_audit_removes_the_saved_line_map(api):
    put(api, YEARLY, dtype="pnl")
    assert any(v.get("kind") == "sheet_map" for v in api.db["column_mappings"].docs)
    assert api.request("DELETE", f"/api/audits/{AUDIT}", json={"confirm": COMPANY}).status_code == 200
    assert not [v for v in api.db["column_mappings"].docs if v.get("audit_id") == AUDIT]


def test_the_sheet_call_has_its_own_caps_and_stays_under_the_daily_spend_cap(api, monkeypatch):
    async def over(_db):
        from app.llm import guards
        raise guards.GuardRefusal("daily_spend_cap_exceeded", "over")
    monkeypatch.setattr(gateway.guards, "check_spend_cap", over)
    api.state["adapter"] = t.FakeAdapter(replies=[YEARLY_REPLY])
    out = put(api, YEARLY, dtype="pnl").json()
    assert out["ai_reading"]["status"] == "not_read" and api.state["adapter"].calls == 0
    assert gateway.SHEET_INPUT_CAP == gateway.SHEET_MAX_TOKENS == 12_000


# ---------------------------------------------------------------------------
# The line table (26.7): the Used column's data, Confirm all, the labels; the KPI sheet (26.14)


def test_every_line_says_whether_the_calculations_use_it(api):
    out = put(api, KPI, dtype="pnl").json()
    lines = lines_by_label(out["sheet_map"])
    assert lines["Total MRR"]["used"] is True and lines["New ARR - inbound"]["used"] is True
    assert lines["Customers"]["used"] is False, "customers is no stored series"
    assert lines["MRR growth % MoM"]["used"] is False
    assert all(isinstance(line["used"], bool) for line in out["sheet_map"]["lines"])


def test_confirm_all_clears_every_pending_line_in_one_request(api):
    out = put(api, KPI, dtype="pnl").json()
    waiting = [line["row"] for line in out["sheet_map"]["lines"] if line["pending"]]
    assert len(waiting) > 1 and not out["saved"]
    r = api.post(f"/api/audits/{AUDIT}/datasets/pnl/lines", json=[{"row": n, "action": "confirm"} for n in waiting])
    assert r.status_code == 200
    out = r.json()
    assert out["pending"] == 0 and out["sheet_map"]["pending"] == 0 and out["saved"]
    assert not [line for line in out["sheet_map"]["lines"] if line["pending"]]
    assert {line["decision"] for line in out["sheet_map"]["lines"] if line["row"] in waiting} == {"confirm"}


def test_the_map_remembers_whether_the_sheet_call_ran(api):
    api.state["adapter"] = t.FakeAdapter(replies=[YEARLY_REPLY])
    assert put(api, YEARLY, dtype="pnl").json()["sheet_map"]["ai_read"] is True
    api.db["audits"].docs[0]["structure_reading_consent"] = False
    assert put(api, KPI, dtype="pnl", replace="true").json()["sheet_map"]["ai_read"] is False


def test_a_sheet_of_kpi_lines_is_a_kpi_sheet_and_a_pnl_is_not():
    assert sh.is_kpi_sheet(sh.build(reading(KPI))["lines"])
    assert not sh.is_kpi_sheet(sh.build(reading(YEARLY))["lines"])
    mixed = sh.build(reading(KPI))["lines"]
    for line in mixed:
        if line["label"] == "Customers":
            line["meaning"] = "sm_expense"
    assert not sh.is_kpi_sheet(mixed), "an S&M expense line makes it a P&L"


def test_one_audit_keeps_a_pnl_and_a_kpi_sheet(api):
    """A KPI sheet with no revenue, invoiced or MRR/ARR row keeps chat-upload.md 26.14 (one with such a row is a monthly revenue report)."""
    file_type = json.dumps({"type": "file_recognition", "sheet": 1, "file_type": "pnl", "revenue_row": None, "currency": None, "non_revenue_rows": []})
    api.state["adapter"] = t.FakeAdapter(replies=[YEARLY_REPLY, file_type, "{}"])
    assert put(api, YEARLY, dtype="pnl").json()["dtype"] == "pnl"
    out = put(api, KPI_NO_REVENUE).json()
    assert out["dtype"] == "kpi" and out["sheet_map"] and out["type_source"] == "ai"
    kinds = {d["dtype"]: d for d in api.get(f"/api/audits/{AUDIT}/datasets").json()["datasets"]}
    assert sorted(kinds) == ["kpi", "pnl"] and kinds["pnl"]["file"] == YEARLY and kinds["kpi"]["file"] == KPI_NO_REVENUE
    assert all(line["series"] for line in kinds["kpi"]["sheet_map"]["lines"] if line["label"] == "Logo churn rate %")
    same = put(api, KPI_NO_REVENUE).json()
    assert same["dtype"] == "kpi" and same["status"] == "same_file", "the same bytes again are the KPI sheet, no replace prompt"


def test_a_kpi_sheet_with_an_mrr_row_is_a_monthly_revenue_report_and_the_revenue_file(api):
    file_type = json.dumps({"type": "file_recognition", "sheet": 1, "file_type": "pnl", "revenue_row": None, "currency": None, "non_revenue_rows": []})
    api.state["adapter"] = t.FakeAdapter(replies=[file_type])
    out = put(api, KPI).json()
    assert out["dtype"] == "revenue" and out["recognition"]["file_type"] == "monthly_revenue_report"
    assert out["recognition"]["revenue_row"]["kind"] == "recurring", "an MRR row is a recurring row, never recognised revenue"


def test_a_kpi_sheet_does_not_feed_the_engine_and_waits_for_its_pending_lines(api):
    upload(api, "blank_vs_zero.csv")
    put(api, YEARLY, dtype="pnl")
    out = put(api, KPI, dtype="kpi").json()
    assert out["dtype"] == "kpi" and out["pending"] > 0
    assert api.post(f"/api/audits/{AUDIT}/compute").status_code == 409, "a line the calculations use waits for the analyst"
    waiting = [line["row"] for line in out["sheet_map"]["lines"] if line["pending"]]
    api.post(f"/api/audits/{AUDIT}/datasets/kpi/lines", json=[{"row": n, "action": "confirm"} for n in waiting])
    results = api.post(f"/api/audits/{AUDIT}/compute").json()
    assert [y["period"] for y in results["yearly_basis"]["years"]] == ["FY2021", "FY2022", "FY2023"], "the P&L alone is the input"


def test_the_analyst_changes_a_statement_between_pnl_and_kpi_sheet(api):
    put(api, KPI, dtype="pnl")
    r = api.post(f"/api/audits/{AUDIT}/datasets/pnl/type", json={"to": "kpi"})
    assert r.status_code == 200 and r.json()["dtype"] == "kpi" and r.json()["type_source"] == "you"
    assert [d["dtype"] for d in api.get(f"/api/audits/{AUDIT}/datasets").json()["datasets"]] == ["kpi"]
    back = api.post(f"/api/audits/{AUDIT}/datasets/kpi/type", json={"to": "pnl"})
    assert back.status_code == 200 and back.json()["dtype"] == "pnl" and back.json()["sheet_map"]


def test_a_sheet_with_no_statement_cannot_be_a_kpi_sheet(api):
    r = upload(api, "blank_vs_zero.csv", dtype="kpi")
    assert r.status_code == 400 and "KPI sheet" in r.json()["detail"]


# ---------------------------------------------------------------------------
# 26.15: how a line is measured (basis), row tags and the new meanings. kpi_bases.csv is synthetic (rule 15): one monthly revenue
# line, its ytd, ltm and quarterly versions, a forecast row, NEW ... (ARR), ARPA, ACV and the other new meanings, Jan-24 to Dec-24.
# ---------------------------------------------------------------------------
BASES = "kpi_bases.csv"
COST_ROWS = "\nCost of revenue," + ",".join(["40"] * 12) + "\nSales and marketing," + ",".join(["30"] * 12) + "\n"


def bases_map(content: bytes = None, reply=None) -> dict:
    return sh.build(reading(BASES, content), reply)


def bases_view(m: dict = None) -> dict:
    return lines_by_label(sh.view(m or bases_map(), 12, 1))


def test_every_line_gets_its_basis_from_its_label_its_meaning_or_its_values():
    lines = bases_view()
    assert {k: lines[k]["basis"] for k in lines} == {
        "Revenue (monthly)": "month", "Revenue current year (YTD)": "ytd", "Revenue last 12 months (LTM)": "ltm",
        "Revenue - quarterly": "quarter", "Forecast revenue": "month", "ARR": "balance", "NEW Enterprise (ARR)": "month",
        "NEW SMB (ARR)": "month", "ARPA": "month", "ARPA - Enterprise": "month", "New sales (ACV)": "month",
        "Annual Contract Value": "month", "Subscriptions": "month", "Contracted revenue": "month",
        "Enterprise % share of ARR": "month", "Net promoter score": "quarter", "Total customers achieved": "lifetime"}
    assert sh.basis_cue("Revenue TTM") == "ltm" and sh.basis_cue("Bookings, last 12 months") == "ltm"
    assert sh.basis_cue("Revenue YTD") == "ytd" and sh.basis_cue("Revenue") is None


def test_new_arr_acv_arpa_and_the_other_new_meanings_read_right():
    lines = bases_view()
    assert {k: lines[k]["meaning"] for k in ("NEW Enterprise (ARR)", "NEW SMB (ARR)", "New sales (ACV)", "Annual Contract Value", "ARPA",
                                             "ARPA - Enterprise", "Subscriptions", "Contracted revenue", "ARR")} == {
        "NEW Enterprise (ARR)": "new_arr", "NEW SMB (ARR)": "new_arr", "New sales (ACV)": "acv", "Annual Contract Value": "acv",
        "ARPA": "arpa", "ARPA - Enterprise": "arpa", "Subscriptions": "subscriptions", "Contracted revenue": "contracted_revenue",
        "ARR": "arr"}
    assert sh.meaning_of("Average revenue per user")[0] == "arpa" and sh.meaning_of("NEW logos (ARR)")[0] == "new_arr"
    share = lines["Enterprise % share of ARR"]
    assert (share["role"], share["meaning"]) == ("rate", "other"), "a % share line is a rate, never ARR"
    assert not share["used"] and not share["series"]
    assert "arpa" in sh.MONEY and "contracted_revenue" in sh.MONEY and "subscriptions" not in sh.MONEY, "ARPA is money; subscriptions a count"


def test_a_label_with_forecast_budget_plan_or_target_tags_its_row_over_its_columns():
    lines = bases_view()
    assert {k for k, v in lines.items() if v["tag"] != "actual"} == {"Forecast revenue"}
    assert [sh.row_tag(x) for x in ("Revenue forecast", "Budget revenue", "Plan ARR", "Target ARR", "Revenue", "Planning")] == [
        "forecast", "budget", "budget", "forecast", "actual", "actual"]


def test_only_the_monthly_line_feeds_the_engine_the_others_show_their_use_and_are_never_summed():
    m = bases_map(content=(MAPPING / BASES).read_bytes() + COST_ROWS.encode())
    periods = sh.with_periods(m["periods"], 12)
    assert [x["row"] for x in sh.pick(m["lines"], "revenue")] == [3], "not the ytd, ltm, quarterly or forecast line"
    pnl = sh.feed(m, 12, 1, {"file": BASES, "sheet": "CSV"})["pnl"]
    assert [r["revenue"] for r in pnl] == [100 + 10 * i for i in range(12)], "never 100 + 100 + ... summed with the ytd line"
    assert pnl[0]["cost_of_revenue"] == 40 and pnl[0]["sm_expense"] == 30
    view = lines_by_label(sh.view(m, 12, 1))
    assert view["Revenue (monthly)"]["used"] and view["Revenue (monthly)"]["use"] == "series"
    assert {k: view[k]["use"] for k in ("Revenue current year (YTD)", "Revenue last 12 months (LTM)", "Revenue - quarterly",
                                        "Forecast revenue")} == {
        "Revenue current year (YTD)": "recompute check", "Revenue last 12 months (LTM)": "recompute check",
        "Revenue - quarterly": "excluded", "Forecast revenue": "series"}, "checked against the months, left out, or a plan series: never Not used"
    assert not any(view[k]["used"] for k in ("Revenue current year (YTD)", "Revenue last 12 months (LTM)", "Forecast revenue"))
    assert periods[0]["grain"] == "month"


def test_a_balance_feeds_the_series_for_arr_but_a_ytd_arr_line_does_not():
    m = bases_map()
    assert [(s["meaning"], s["row"]) for s in sh.kpi_series(m["lines"], sh.with_periods(m["periods"], 12), 1)] == [
        ("arr", 9), ("new_arr", 10), ("new_arr", 11)]
    ytd = bases_map(content=(MAPPING / BASES).read_bytes().replace(b"ARR,1200", b"ARR YTD,1200"))
    assert "arr" not in [s["meaning"] for s in sh.kpi_series(ytd["lines"], sh.with_periods(ytd["periods"], 12), 1)]


def test_the_december_of_a_ytd_line_and_the_last_12_months_of_an_ltm_line_match_the_monthly_line():
    m = bases_map()
    periods = sh.with_periods(m["periods"], 12)
    found = sh.basis_checks(m["lines"], periods)
    assert [(c["row"], c["basis"], c["period"], c["line"], c["reference"], c["ok"]) for c in found] == [
        (4, "ytd", "Dec-24", 1860, 1860, True), (5, "ltm", "Dec-24", 1860, 1860, True)], "Jan-Nov LTM has no 12 months: not checked"
    lines = bases_view(m)
    assert lines["Revenue current year (YTD)"]["checks"] == {} and lines["Revenue last 12 months (LTM)"]["checks"] == {}


def test_a_ytd_or_ltm_line_off_by_more_than_half_a_percent_is_flagged():
    lines = (MAPPING / BASES).read_text().split("\n")
    lines[3] = lines[3].rsplit(",", 1)[0] + ",1900"        # YTD December 1,900 against 1,860 = +2.2%
    lines[4] = lines[4].rsplit(",", 1)[0] + ",1865"        # LTM December 1,865 against 1,860 = +0.27%, inside the tolerance
    bad = bases_view(bases_map("\n".join(lines).encode()))
    assert bad["Revenue current year (YTD)"]["checks"] == {"ytd_mismatch": ["Dec-24"]}
    assert bad["Revenue last 12 months (LTM)"]["checks"] == {}
    lines[4] = lines[4].rsplit(",", 1)[0] + ",1700"
    worse = bases_view(bases_map("\n".join(lines).encode()))
    assert worse["Revenue last 12 months (LTM)"]["checks"] == {"ltm_mismatch": ["Dec-24"]}
    order = [x["label"] for x in sh.view(bases_map("\n".join(lines).encode()), 12, 1)["lines"]][:2]
    assert order == ["Revenue current year (YTD)", "Revenue last 12 months (LTM)"], "lines with a failed check come first"


def test_a_forecast_row_never_feeds_actuals_or_series_and_becomes_a_claim_for_the_months_it_covers():
    m = bases_map()
    periods = sh.with_periods(m["periods"], 12)
    assert 7 not in [x["row"] for x in sh.pick(m["lines"], "revenue")] and sh.pick(m["lines"], "revenue", "forecast")[0]["row"] == 7
    assert 7 not in sh.used_rows(m["lines"], periods)
    found = {c["label"]: c for c in sh.claims(m, {}, 12)["lines"]}
    plan = found["Forecast revenue"]
    assert (plan["metric"], plan["claim_type"]) == ("revenue", "revenue")
    assert [(v["period"], v["value"], v["cell"], v["tag"]) for v in plan["values"]] == [("Jan-24–Dec-24", 1920, "B7:M7", "forecast")], \
        "the twelve months of the row summed, cited to its cells, tagged forecast"
    assert not {"Revenue current year (YTD)", "Revenue last 12 months (LTM)", "Revenue - quarterly", "Contracted revenue",
                "ARPA"} & set(found), "other bases and money-per-customer lines are no plan line"


def test_when_a_monthly_pnl_feeds_the_engine_only_the_forecast_rows_are_claims():
    m = bases_map(content=(MAPPING / BASES).read_bytes() + COST_ROWS.encode())
    found = sh.claims(m, {}, 12)
    assert [c["label"] for c in found["lines"]] == ["Forecast revenue"], "the monthly P&L is engine input, not a claim"
    assert found["lines"][0]["values"][0]["tag"] == "forecast"


def test_a_stored_map_from_before_basis_and_tag_is_still_read():
    m = bases_map()
    for line in m["lines"]:
        line.pop("basis"), line.pop("tag")
    view = sh.view(m, 12, 1)
    assert all(line["use"] in (None, *pl.USES) for line in view["lines"]) and view["series"] >= 1


def test_a_reply_basis_is_used_but_a_cue_in_the_label_wins_and_a_reply_without_one_keeps_pythons():
    reply = {"type": "sheet_structure", "header_row": 2, "label_col": 1, "periods": [],
             "rows": [{"row": 3, "role": "other", "parent_row": None, "meaning": "revenue", "basis": "ltm"},
                      {"row": 4, "role": "other", "parent_row": None, "meaning": "revenue", "basis": "month"},
                      {"row": 6, "role": "other", "parent_row": None, "meaning": "revenue"}]}
    lines = bases_view(bases_map(reply=reply))
    assert lines["Revenue (monthly)"]["basis"] == "ltm", "the model's reading"
    assert lines["Revenue current year (YTD)"]["basis"] == "ytd", "the label says YTD"
    assert lines["Revenue - quarterly"]["basis"] == "quarter", "no basis in the reply: Python's"


def test_the_analyst_changes_a_basis_in_one_click_and_the_saved_map_keeps_it(api):
    out = put(api, BASES, dtype="pnl").json()
    assert lines_by_label(out["sheet_map"])["Revenue (monthly)"]["basis"] == "month"
    r = api.post(f"/api/audits/{AUDIT}/datasets/pnl/lines", json=[{"row": 3, "action": "change", "basis": "ytd"}])
    assert r.status_code == 200
    changed = lines_by_label(r.json()["sheet_map"])["Revenue (monthly)"]
    assert (changed["basis"], changed["use"], changed["decision"]) == ("ytd", "recompute check", "change")
    assert api.post(f"/api/audits/{AUDIT}/datasets/pnl/lines", json=[{"row": 3, "action": "change", "basis": "yearly"}]).status_code == 400


def test_the_saved_line_map_holds_the_basis_and_no_label(api):
    out = put(api, BASES, dtype="pnl").json()
    api.post(f"/api/audits/{AUDIT}/datasets/pnl/lines", json=[{"row": line["row"], "action": "confirm"}
                                                             for line in out["sheet_map"]["lines"] if line["pending"]])
    versions = lambda: sorted((v for v in api.db["column_mappings"].docs if v.get("kind") == "sheet_map"), key=lambda v: v["version"])  # noqa: E731
    assert len(versions()) == 1
    api.post(f"/api/audits/{AUDIT}/datasets/pnl/lines", json=[{"row": 3, "action": "change", "basis": "ltm"}])
    assert len(versions()) == 2, "a change of basis alone is a new version"
    saved = versions()[-1]
    assert {line["basis"] for line in saved["lines"] if line["row"] == 3} == {"ltm"}
    assert "Revenue" not in json.dumps(saved["lines"]) and "Forecast" not in json.dumps(saved["lines"])


def test_plan_tags_a_row_only_as_the_plan_figure_never_a_pricing_or_subscription_plan():
    assert [sh.row_tag(x) for x in ("NEW ANNUAL PLAN SUBSCRIPTIONS (ARR)", "MRR - Pro plan", "ARR plan")] == ["actual"] * 3
    assert [sh.row_tag(x) for x in ("Plan: revenue", "Revenue (plan)", "Planned ARR", "Business plan revenue", "Revenue plan: 2026")] == ["budget"] * 5
    assert sh.row_tag("Revenue budget") == "budget" and sh.row_tag("Revenue target") == "forecast"
