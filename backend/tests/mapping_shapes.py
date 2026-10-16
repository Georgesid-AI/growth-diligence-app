"""Synthetic workbooks in the shape of the three files that missed their mapping (2026-10-16), and the fixture folder of
the export layouts (tests/fixtures/mapping/). Invented values only (CLAUDE.md rule 15): no client name, figure or deck text."""
import io
from functools import lru_cache
from pathlib import Path

from openpyxl import Workbook

FIXTURES = Path(__file__).resolve().parent / "fixtures" / "mapping"
PERIODS = ["2023", "Q1:24", "Q2:24", "Q3:24", "2024 YTD", "2024", "2025", "2026", "2027", "2028"]


def _save(wb) -> bytes:
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def _period_table(ws, top: int, title: str, lines, first: int = 3):
    """A wide table: columns A and B empty, the title in C, the header on row top + 3, one line per metric."""
    ws.cell(top, first, title)
    ws.cell(top + 1, first, "EUR thousands, unaudited")
    ws.cell(top + 3, first, "Line")
    for j, p in enumerate(PERIODS):
        ws.cell(top + 3, first + 1 + j, p)
    for i, line in enumerate(lines):
        ws.cell(top + 4 + i, first, line)
        for j in range(len(PERIODS)):
            ws.cell(top + 4 + i, first + 1 + j, 100 + 10 * i + j)
    return top + 4 + len(lines)


def _stacked_wide(title: str, lines) -> bytes:
    """Title rows in C2:C3, the header on row 5, then a forecast table, a budget comparison and notes below it."""
    wb = Workbook()
    ws = wb.active
    ws.title = "Revenue" if title.startswith("Revenue") else "P&L"
    end = _period_table(ws, 2, title, lines)
    ws.cell(end + 1, 3, "Forecast")
    for j, p in enumerate(PERIODS):
        ws.cell(end + 2, 4 + j, p)
    ws.cell(end + 3, 3, "Forecast line")
    for j in range(len(PERIODS)):
        ws.cell(end + 3, 4 + j, 9000 + j)
    ws.cell(end + 6, 3, "Budget comparison")
    ws.cell(end + 7, 3, "Budget line")
    ws.cell(end + 7, 4, 77777)
    ws.cell(end + 10, 3, "Notes")
    ws.cell(end + 11, 3, "Figures are management estimates and are not audited.")
    other = wb.create_sheet("Source inputs")
    other.append(["Input", "Value"])
    other.append(["Growth assumption", 0.2])
    other.append(["Churn assumption", 0.05])
    return _save(wb)


@lru_cache(maxsize=None)
def revenue_to_2028() -> bytes:
    return _stacked_wide("Revenue through 31 December 2028", ["Revenue", "Gross profit", "New customers", "Customers at year-end", "ARR"])


@lru_cache(maxsize=None)
def pnl_to_2028() -> bytes:
    return _stacked_wide("Profit and loss through 31 December 2028", ["Revenue", "Cost of sales", "Payroll", "Sales and marketing", "EBITDA"])


@lru_cache(maxsize=None)
def crm_anonymous_to_2028() -> bytes:
    wb = Workbook()
    ws = wb.active
    ws.title = "Accounts"
    ws["C2"] = "Accounts through 31 December 2028"
    ws["C3"] = "EUR; accounts anonymised"
    header = ["Account ID", "Segment", "Fee EUR / month", "Add-on EUR", "Existing units", "Existing MRR EUR", "2024 target units",
              "2024 target MRR EUR", "Probability claim", "Incremental MRR EUR", "Weighted increment EUR", "Source / date",
              "Evidence label", "Payer evidence", "Notes / contradictions", "Action owner", "Due"]
    for j, h in enumerate(header):
        ws.cell(6, 3 + j, h)
    for i in range(6):
        row = [f"ACC-{i + 1:04d}", "SMB" if i % 2 else "Mid-market", 100 + i, 10, 5 + i, 500 + 50 * i, 8 + i, 800 + 50 * i, "60%",
               300, 180, f"email 2024-0{i + 1}-02", "Claimed", "none yet", "no contradiction found", "Owner A", f"2024-0{i + 1}-28"]
        for j, v in enumerate(row):
            ws.cell(7 + i, 3 + j, v)
    tracker = wb.create_sheet("Annual tracker")
    for j, h in enumerate(["Account ID", "Period end", "Basis", "Units at year-end", "MRR EUR at year-end", "ARR EUR",
                           "Evidence / update", "Owner", "Next review"]):
        tracker.cell(6, 3 + j, h)
    for i in range(3):
        for j, v in enumerate([f"ACC-{i + 1:04d}", f"2024-12-31", "Plan", 6 + i, 600 + 20 * i, 7200 + 240 * i, "plan", "Owner A", "2025-03-31"]):
            tracker.cell(7 + i, 3 + j, v)
    funnel = wb.create_sheet("Funnel and evidence")
    funnel["C2"] = "Funnel"
    for j, h in enumerate(["Step", "Count", "Evidence"]):
        funnel.cell(5, 3 + j, h)
    funnel.append([None, None, "Lead", 10, "listed"])
    return _save(wb)
