"""Model reading of deck structures and spreadsheet headers (docs/specs/llm-structure-reading.md).

Recorded replies (tests/fixtures/structure_replies/, public test decks only) are replayed through the
fake adapter: no test reaches a live API. They are written in the model's output format against the
structures the parser finds on the public decks; no live call was made to produce them.
"""
import asyncio
import copy
import json
import os
import sys
from pathlib import Path

import pytest

BACKEND = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(BACKEND))
sys.path.insert(0, str(BACKEND / "tests"))

from app.llm import gateway, guards  # noqa: E402
from app.llm.schemas import StructureReply  # noqa: E402
from app.structures import redact, verify  # noqa: E402
import test_llm_gateway as t  # noqa: E402

REPLIES = BACKEND / "tests" / "fixtures" / "structure_replies"
DECKS = BACKEND.parent / "tests" / "fixtures" / "decks" / "decks"
AUDIT = "audit-s"
AUDIT_DOC = {"id": AUDIT, "company_name": "Zero2Hero", "client_name": "Northbridge Capital",
             "engagement_reference": "ENG-2026-041", "structure_reading_consent": True, "fiscal_year_end": 12,
             "results": None}
# A structure as the model reads it: its cells, then the items Python listed (docs/specs/structure-labelling.md).
TEXT = ('r1c2: FY2025\nr1c3: FY2026\nr2c1: Revenue\nr2c2: £1,200,000\nr2c3: £1,500,000\nitems:\n'
        'i1 r2c2 "1,200,000" 1200000 h r2c1 r1c2\ni2 r2c3 "1,500,000" 1500000 h r2c1 r1c3')


def _label(item, metric="revenue", period=None, unit=None, unit_other=None, actual_or_forecast="forecast"):
    """One label of the labelling reply: exactly these fields, no value, cell or flag."""
    return {"item": item, "metric": metric, "period": period, "unit": unit, "unit_other": unit_other,
            "actual_or_forecast": actual_or_forecast}


REPLY = {"type": "table", "labels": [_label("i1", period="FY2025", unit="GBP"), _label("i2", period="FY2026", unit="GBP")],
         "pairs": []}
NOTHING = {"type": "table", "labels": [], "pairs": []}          # the reply to a structure with no item
ROADMAP_TEXT = ('r1c1: Launch the API\nr2c1: Q3 2025\nr3c1: Hire 5 engineers\nr4c1: Q4 2025\nitems:\n'
                'i1 r3c1 "5" 5 h r1c1\nd1 r2c1\nd2 r4c1\nt1 r1c1\nt2 r3c1')
ROADMAP_REPLY = {"type": "roadmap", "labels": [_label("i1", "people", unit="count")],
                 "pairs": [{"line": "t1", "date": "d1", "category": "launch"},
                           {"line": "t2", "date": "d2", "category": "hiring"}]}


def _db(**audit):
    db = t.FakeDB()
    db["audits"].docs.append({**copy.deepcopy(AUDIT_DOC), **audit})
    return db


def _read(db, text=TEXT, kind="table", replies=None, adapter=None, **kwargs):
    adapter = adapter or t.FakeAdapter(replies=[json.dumps(r) for r in (replies or [REPLY])])
    result = asyncio.run(gateway.read_structure(db, AUDIT, text, kind, adapter=adapter, sleep=t._noop_sleep, **kwargs))
    return result, adapter


def _fixtures():
    return [json.loads(p.read_text(encoding="utf-8")) for p in sorted(REPLIES.glob("*.json"))]


def _structure(fixture):
    from app.decks import parser
    deck = parser.parse_deck((DECKS / fixture["file"]).read_bytes(), fixture["file"])
    return next(s for s in deck["structures"]
                if (s.get("slide") or s.get("page")) == fixture["page"] and s["type"] == fixture["type"])


# ---------------------------------------------------------------------------
# Recorded replies on the public test decks, replayed and verified
# ---------------------------------------------------------------------------
@pytest.mark.parametrize("fixture", _fixtures(), ids=lambda f: f"{f['file']}-p{f['page']}-{f['type']}")
def test_recorded_replies_replay_through_the_gateway_and_the_verifier(fixture):
    pytest.importorskip("pdfplumber")
    pytest.importorskip("pptx")
    structure = _structure(fixture)
    cells, _ = redact.redact_structure(structure["cells"], "Zero2Hero", {})
    listed = structure_items.list_items({**structure, "cells": cells})
    text = structure_items.text({**structure, "cells": cells}, listed)
    assert fixture["match"] in text, "--fake and RecordedAdapter find a recorded reply by this text: it must be today's"
    db = _db()
    result, adapter = _read(db, text, fixture["type"], [fixture["reply"]])
    assert result.status == "read" and adapter.calls == 1, result.reason
    sent = json.loads(adapter.payloads[0])
    assert sent == {"type": fixture["type"], "text": text}, "the type, the structure text and its item list, nothing else"
    checked = verify.verify(structure, listed, result.labels, result.pairs)
    statuses = [i["status"] for i in checked["items"]]
    assert {"verified": statuses.count("verified"), "suggestion": statuses.count("suggestion")} == fixture["expected"]


def _moz_text(page):
    from app.decks import parser
    deck = parser.parse_deck((DECKS / "02-moz.pdf").read_bytes(), "02-moz.pdf")
    structure, = [s for s in deck["structures"] if s.get("page") == page and s["type"] == "kpi_panel"]
    cells, _ = redact.redact_structure(structure["cells"], "Moz", {})
    return structure_items.text({**structure, "cells": cells}, structure_items.list_items({**structure, "cells": cells}))


def test_moz_p20_and_p21_reach_the_model_with_every_value_beside_its_label_and_the_title_marked():
    """Issue #48: the nine value boxes that hold only a figure reach the model, each in the row of its label and citing
    it as its header; "2011 Estimated Revenue" goes as the title, the label of "$12 -$13 million"."""
    pytest.importorskip("pdfplumber")
    sent = {}
    for page in (20, 21):
        text = _moz_text(page)
        result, adapter = _read(_db(), text, "kpi_panel", [unlabelled({"type": "kpi_panel", "text": text})])
        assert result.status == "read", result.reason
        sent[page] = json.loads(adapter.payloads[0])["text"]
    assert sent[20].splitlines()[:2] == ["r1c1 title: 2011 Estimated Revenue", "r1c2: $12 -$13 million"]
    assert sent[21].splitlines()[:2] == ["r1c1 title: % of Free Trials Converting to Paid", "r1c2: ~57%"]
    assert 'i1 r1c2#1 "12" 12000000 h r1c1' in sent[20].splitlines()
    pairs = {(20, "Number of PRO Subscribers", "~13,500"), (20, "# of New Free Trials / Day", "~100"),
             (20, "Avg. Customer Lifetime Value", "~$900"), (20, "Avg. Cost of Paid Acquisition", "~$100"),
             (20, "Avg. Monthly Revenue / Subscriber", "~$93"), (21, "Email Subscribers", "~300K"),
             (21, "Gross Margins", "~82%"), (21, "% of Free Trials Converting to Paid", "~57%"),
             (21, "Churn Rate in 1st2 Paid Months", "~25%")}
    for page, label, value in pairs:
        lines = sent[page].splitlines()
        cell = next(line.split(": ", 1)[0] for line in lines if line.endswith(f": {value}"))
        row = cell.split("c")[0]
        assert any(line.startswith(f"{row}c1") and line.endswith(f": {label}") for line in lines), (page, value)
        items = [line.split() for line in lines[lines.index("items:") + 1:]]
        assert any(item[1] == cell and item[-2:] == ["h", f"{row}c1"] for item in items), \
            f"{value}'s item cites {label} as its header"


def _deck_structure(file, page, kind):
    from app.decks import parser
    deck = parser.parse_deck((DECKS / file).read_bytes(), file)
    structure, = [s for s in deck["structures"] if (s.get("slide") or s.get("page")) == page and s["type"] == kind]
    return structure


def _front_b_items(page):
    """(cell text, raw text, header texts) per item of front-b's KPI panel on a slide, as Python lists them."""
    structure = _deck_structure("01-front-b.pptx", page, "kpi_panel")
    text = {redact.cell_id(c): c["text"] for c in structure["cells"]}
    return [(text[i["cell"]], i["raw"], [text[h] for h in i["headers"]])
            for i in structure_items.list_items(structure)["items"]]


def test_front_b_a_kpi_value_cites_a_label_from_another_box_only_when_it_sits_directly_next_to_it():
    """Issue #50: a tall box merges two visual rows into one grid row, so a value's row neighbour need not sit next to
    it on the page. A wrong label is worse than a missing one: the verifier checks the value and the period, never the
    metric. The top line of the value's own box always counts."""
    p15 = _front_b_items(15)
    assert ("18 months", "18", []) in p15, "\"Cash on hand\" sits above \"$7m left\", not next to \"18 months\""
    assert ("$7m left", "7m", ["Cash on hand"]) in p15, "the label directly above its value still counts"
    p12 = _front_b_items(12)
    assert [i for i in p12 if i[0].endswith("ago")] == [
        ("2 weeks ago", "2", []), ("10 months ago", "10", []),
        ("joined 6 months ago", "6", ["The team isn’t one year old"]), ("2 months ago", "2", []),
        ("10 months ago", "10", [])], "no tenure value cites \"LTV / CAC\" or \"Spend as\""
    assert ("2.5 2.6 4.4", "2.5", ["On a sustainable trajectory"]) in p12, "the top line of its own box"
    p16 = _front_b_items(16)
    assert len(p16) == 12 and not [i for i in p16 if "Cash" in i[2]], "a chart legend entry never labels an axis date"


def test_roadmaps_keep_their_header_cells():
    """The guard is a KPI panel's: moz p2's "1.1M" still cites the cells left of it in its row (its paragraph is one
    cell since issue #49, so it is its box's top line), and tea p11's "layer-2" the cell left of it and its box's top
    line."""
    pytest.importorskip("pdfplumber")
    roadmap = _deck_structure("02-moz.pdf", 2, "roadmap")
    item, = [i for i in structure_items.list_items(roadmap)["items"] if i["raw"] == "1.1M"]
    assert item["headers"] == ["r1c2", "r1c1"]
    tea = _deck_structure("10-tea.pdf", 11, "roadmap")
    item, = [i for i in structure_items.list_items(tea)["items"] if i["raw"] == "2"]
    assert (v._text(tea, item["cell"]), item["headers"]) == ("layer-1 to layer-2", ["r10c1", "r9c2"])


def test_a_kpi_value_takes_a_neighbours_scale_only_when_it_sits_directly_next_to_it():
    """Issue #50 (decision of 2026-10-06): a tall box puts the label row ("Headcount", "Revenue £m") and the value row
    ("12", "4.5") in one grid row, so "Revenue £m" is a grid neighbour of "12" as well. Its scale is for "4.5", the value
    right under it; the headcount stays 12."""
    pptx = pytest.importorskip("pptx")
    import io
    from pptx.util import Inches
    from app.decks import parser
    prs = pptx.Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    for text, left, top, height in (("Headcount", 1, 1, 0.5), ("Revenue £m", 3.2, 1, 0.5), ("12", 1, 1.6, 0.5),
                                    ("4.5", 3.2, 1.6, 0.5), ("Runway\n18 months", 5.4, 1, 1.1)):
        slide.shapes.add_textbox(Inches(left), Inches(top), Inches(2), Inches(height)).text_frame.text = text
    buf = io.BytesIO()
    prs.save(buf)
    panel, = parser.parse_deck(buf.getvalue(), "d.pptx")["structures"]
    text = {redact.cell_id(c): c["text"] for c in panel["cells"]}
    assert [text[f"r1c{n}"] for n in range(1, 6)] == ["Headcount", "12", "Revenue £m", "4.5", "Runway"], "one grid row"
    assert [(text[i["cell"]], i["values"][0]["value"]) for i in structure_items.list_items(panel)["items"]] == \
        [("12", 12), ("4.5", 4500000), ("18 months", 18)]


def test_the_call_is_pinned_to_one_model_with_no_temperature_no_tools_and_the_structure_schema(monkeypatch):
    monkeypatch.setenv("NARRATIVE_MODEL", "claude-opus-5-5")      # the narrative setting does not move it
    result, adapter = _read(_db())
    assert result.status == "read" and result.model == "claude-sonnet-5-5" == gateway.STRUCTURE_MODEL
    request, = adapter.requests
    assert request["model"] == "claude-sonnet-5-5" and request["temperature"] is None
    assert request["max_tokens"] == gateway.STRUCTURE_MAX_TOKENS
    assert set(request["json_schema"]["properties"]) == {"type", "labels", "pairs"}, "a deck structure is labelled"
    assert "claude-sonnet-5-5" in gateway.MODEL_PRICING_USD, "the pinned model has a price entry"


# ---------------------------------------------------------------------------
# The item list (docs/specs/structure-labelling.md sections 1 and 3): Python lists every figure in every
# redacted cell, with its cell, position, raw text, value(s) and header cells; the model only labels them
# ---------------------------------------------------------------------------
import test_structure_verifier as v  # noqa: E402
from app.structures import items as structure_items  # noqa: E402


def _listed(rows, header_rows=1, kind="table", spans=None):
    return structure_items.list_items(v._struct(rows, header_rows, kind, spans))["items"]


def _brief(listed):
    """(id, cell, position, raw text, values with the default first, header cells) per item."""
    return [(i["id"], i["cell"], i["position"], i["raw"], [x["value"] for x in i["values"]], i["headers"])
            for i in listed]


SOCIAL = [["", "Members"], ["Web", "12"], ["Social", "Discord(150) Telegram(30K)"]]


def test_a_cell_with_several_figures_gives_one_item_each_with_its_position_in_reading_order():
    assert _brief(_listed(SOCIAL)) == [
        ("i1", "r2c2", 1, "12", [12], ["r2c1", "r1c2"]),
        ("i2", "r3c2", 1, "(150)", [150, -150], ["r3c1", "r1c2"]),
        ("i3", "r3c2", 2, "(30K)", [30000, -30000], ["r3c1", "r1c2"])]


def test_the_item_list_comes_below_the_structure_text_one_line_per_item():
    structure = v._struct(SOCIAL)
    text = structure_items.text(structure, structure_items.list_items(structure))
    assert text == ("r1c2: Members\nr2c1: Web\nr2c2: 12\nr3c1: Social\nr3c2: Discord(150) Telegram(30K)\n"
                    "items:\n"
                    'i1 r2c2 "12" 12 h r2c1 r1c2\n'
                    'i2 r3c2#1 "(150)" 150 or -150 h r3c1 r1c2\n'
                    'i3 r3c2#2 "(30K)" 30000 or -30000 h r3c1 r1c2'), "the spec's line format"
    assert structure_items.text(v._struct([["Plan"]], header_rows=0), {"items": [], "dates": [], "lines": []}) == \
        "r1c1: Plan\nitems:", "a structure with no figure still sends the items line"


def test_dates_are_periods_and_are_left_out():
    listed = _listed([["", "FY2025", "Q3", "M1", "Year 2"], ["Revenue in 2024", "£1.2m in 2025", "", "", ""],
                      ["Users", "Mar 2025: 5K", "", "", ""]])
    assert [(i["cell"], i["raw"], i["values"][0]["value"]) for i in listed] == \
        [("r2c2", "1.2m", 1200000), ("r3c2", "5K", 5000)]


def test_a_digit_inside_a_word_is_no_figure_as_the_parser_reads_it():
    """Issue #50: "zero2hero" (zero2hero p11) and "Churn Rate in 1st2 Paid Months" (moz p21) give no item, as the
    parser's reader (claims.figures) already reads them: no letter right before a number, no ordinal ending. A scale,
    a % or brackets stay part of their figure."""
    from app.decks import claims
    words = ["zero2hero", "Churn Rate in 1st2 Paid Months", "Go2Market", "Web3 Foundation Open Grant"]
    for text in words:
        assert (verify.figures(text), claims.figures(text)) == ([], []), text
    assert _listed([["", "Plan"]] + [["Name", w] for w in words]) == []
    assert [i["raw"] for i in _listed([["", "Plan"], ["Social", "Telegram(30K) ~$93 62% £1.2m 2.5x"]])] == \
        ["(30K)", "93", "62%", "1.2m", "2.5"]
    pytest.importorskip("pdfplumber")
    for file, page, text in (("05-zero2hero.pdf", 11, "zero2hero"), ("02-moz.pdf", 21, "Churn Rate in 1st2 Paid Months")):
        structure = _deck_structure(file, page, "kpi_panel")
        cell, = [redact.cell_id(c) for c in structure["cells"] if c["text"] == text]
        assert not [i for i in structure_items.list_items(structure)["items"] if i["cell"] == cell], (file, page)


def test_a_title_cell_is_marked_in_its_line_read_back_and_never_an_item():
    """deck-parser.md section 7: a slide title a KPI panel takes as a label goes as `r<row>c<col> title: <text>`; it is
    a label, so the item list skips it, a digit in it too, and redaction keeps the mark for the gateway's check."""
    cells = [{"row": 1, "col": 1, "text": "2011 Estimated Revenue", "title": True},
             {"row": 1, "col": 2, "text": "$12 -$13 million"}]
    text = redact.structure_text(cells)
    assert text == "r1c1 title: 2011 Estimated Revenue\nr1c2: $12 -$13 million"
    assert redact.parse_structure_text(text) == cells
    assert redact.redact_structure(cells, "Target Co", {})[0] == cells
    structure = {"type": "kpi_panel", "header_rows": 0, "cells": [{**c, "box": n} for n, c in enumerate(cells, 1)]}
    structure["cells"][1]["next_to"] = ["r1c1"]           # the title sits beside it (deck-parser.md section 7)
    assert [(i["cell"], i["headers"]) for i in structure_items.list_items(structure)["items"]] == \
        [("r1c2", ["r1c1"]), ("r1c2", ["r1c1"])], "the title is the range's header"
    structure["cells"][0]["text"] = "Top 3 Metrics"
    assert [i["cell"] for i in structure_items.list_items(structure)["items"]] == ["r1c2", "r1c2"]


def test_values_are_in_full_units_under_todays_normalisation():
    assert [i["values"][0]["value"] for i in _listed([["£m", "2025", "2026 (€'000)"],
                                                      ["Revenue", "1.2", "4.5"], ["EBITDA", "(0.3)", "1,234"]])] \
        == [1200000, 4500, -300000, 1234000], "a scale in the corner or a header, never an item of its own; brackets"
    assert [i["values"][0]["value"] for i in _listed([["", "Plan"], ["Revenue", "€1.234,5"]])] == [1234.5], \
        "the structure writes a decimal comma"
    assert [i["values"][0]["value"] for i in _listed([["", "Plan"], ["Revenue", "$3.6m"], ["ARR", "2bn"],
                                                      ["Margin", "62%"], ["Loss", "-$1,200"]])] == \
        [3600000, 2000000000, 62, -1200]


def test_a_range_is_two_figures_and_its_low_end_takes_the_high_ends_scale():
    """moz p20: "$12 -$13 million" is twelve to thirteen million, not 12 and minus thirteen million."""
    assert [i["values"] for i in _listed([["", "Plan"], ["Revenue", "$12 -$13 million"], ["Margin", "5 – 10%"]])] \
        == [[{"value": 12000000, "dot_reading": None, "bracket_reading": None}],
            [{"value": 13000000, "dot_reading": None, "bracket_reading": None}],
            [{"value": 5, "dot_reading": None, "bracket_reading": None}],
            [{"value": 10, "dot_reading": None, "bracket_reading": None}]]


def test_a_ranges_two_items_are_marked_as_one_range_low_and_high():
    """Decision of 2026-10-06: the two figures of a range stay two items for the model (spec section 1), and the item
    list stores them as one range (low, high)."""
    listed = _listed([["", "Plan"], ["Revenue", "$12 -$13 million"], ["Margin", "5 – 10%"], ["Users", "12 and 13"]])
    assert [(i["id"], i.get("range")) for i in listed] == [
        ("i1", {"items": ["i1", "i2"], "low": 12000000, "high": 13000000}),
        ("i2", {"items": ["i1", "i2"], "low": 12000000, "high": 13000000}),
        ("i3", {"items": ["i3", "i4"], "low": 5, "high": 10}),
        ("i4", {"items": ["i3", "i4"], "low": 5, "high": 10}),
        ("i5", None), ("i6", None)], "two figures with no dash between them are no range"
    stored = structure_items.stored({"items": listed})["items"]
    assert stored[0]["range"] == {"items": ["i1", "i2"], "low": 12000000, "high": 13000000} and "raw" not in stored[0]
    structure = v._struct([["", "Plan"], ["Revenue", "$12 -$13 million"]])
    assert structure_items.text(structure, structure_items.list_items(structure)).endswith(
        'i1 r2c2#1 "12" 12000000 h r2c1 r1c2\ni2 r2c2#2 "13 million" 13000000 h r2c1 r1c2'), "the item lines are unchanged"


def _readings(text):
    item, = _listed([["", "Plan"], ["Hours", text]])
    return [(x["value"], x["dot_reading"], x["bracket_reading"]) for x in item["values"]]


@pytest.mark.parametrize("text, readings", [
    ("Approx. 2.500 hours", [(2500, "thousands", None), (2.5, "decimal", None)]),
    ("12.500", [(12500, "thousands", None), (12.5, "decimal", None)]),
    ("1.250M", [(1250000, "decimal", None), (1250000000, "thousands", None)]),       # a suffix: decimal
    ("$2.500bn", [(2500000000, "decimal", None), (2500000000000, "thousands", None)]),
    ("0.500", [(0.5, None, None)]),                                                     # not ambiguous
    ("2.50", [(2.5, None, None)]),
])
def test_a_dot_before_three_digits_carries_both_values_thousands_first_unless_a_suffix(text, readings):
    assert _readings(text) == readings


def test_with_a_decimal_comma_a_dot_before_three_digits_is_a_thousands_separator_only():
    listed = _listed([["", "Plan", ""], ["Revenue", "2.500", "1.234,5"]])
    assert [x["value"] for x in listed[0]["values"]] == [2500]


@pytest.mark.parametrize("text, readings", [
    ("Telegram(30K)", [(30000, None, "positive"), (-30000, None, "negative")]),
    ("Discord (150)", [(150, None, "positive"), (-150, None, "negative")]),
    ("Net loss (1,200)", [(-1200, None, "negative"), (1200, None, "positive")]),
    ("LOSSES (7)", [(-7, None, "negative"), (7, None, "positive")]),
    ("Deficit(5)", [(-5, None, "negative"), (5, None, "positive")]),
    ("Negative cash flow (2)", [(-2, None, "negative"), (2, None, "positive")]),
    ("Revenue decline (3%)", [(-3, None, "negative"), (3, None, "positive")]),
    ("(1,200)", [(-1200, None, None)]),                              # wholly bracketed: one reading
    ("£(1,200)", [(-1200, None, None)]),
    ("Net loss -1,200", [(-1200, None, None)]),
    ("Net loss (2.500)", [(-2500, "thousands", "negative"), (2500, "thousands", "positive"),
                          (-2.5, "decimal", "negative"), (2.5, "decimal", "positive")]),
])
def test_a_bracketed_number_after_text_carries_both_signs_negative_first_after_a_loss_word(text, readings):
    assert _readings(text) == readings


def test_a_loss_word_counts_only_in_the_text_before_the_figure():
    first, second = _listed([["", "Plan"], ["P&L", "Gross (5) then net loss (3)"]])
    assert [x["bracket_reading"] for x in first["values"]] == ["positive", "negative"]
    assert [x["bracket_reading"] for x in second["values"]] == ["negative", "positive"]


def test_the_same_structure_always_gives_the_same_list():
    structure = v._struct([["", "2024", "2025"], ["Revenue", "£1M", "£2M"], ["Users", "5K / 7K", "9K"]])
    shuffled = {**structure, "cells": list(reversed(structure["cells"]))}
    first = structure_items.list_items(structure)
    assert first == structure_items.list_items(structure) == structure_items.list_items(shuffled)
    assert [(i["id"], i["cell"], i["position"]) for i in first["items"]] == [
        ("i1", "r2c2", 1), ("i2", "r2c3", 1), ("i3", "r3c2", 1), ("i4", "r3c2", 2), ("i5", "r3c3", 1)]


def test_a_figure_spanning_a_line_break_or_double_space_is_cut_as_the_cell_line_writes_it():
    """The structure text writes each cell on one line with single spaces; an item's raw text is cut from that same
    text, so it is always inside its cell line and the structure is sent, not refused (bad_item_line)."""
    structure = v._struct([["", "Plan"], ["Revenue", "$12 -\n$13\nmillion"], ["EBITDA", "Net  loss  (  1,200 )"]])
    listed = structure_items.list_items(structure)
    assert [(i["raw"], i["values"][0]["value"]) for i in listed["items"]] == \
        [("12", 12000000), ("13 million", 13000000), ("( 1,200 )", -1200)]
    text = structure_items.text(structure, listed)
    assert "\n".join(text.split("\n")[:5]) == "r1c2: Plan\nr2c1: Revenue\nr2c2: $12 - $13 million\nr3c1: EBITDA\n" \
        "r3c2: Net loss ( 1,200 )"
    result, adapter = _read(_db(), text, replies=[{"type": "table", "pairs": [], "labels": [
        _label(i["id"], "not_a_metric") for i in listed["items"]]}])
    assert (result.status, adapter.calls) == ("read", 1), result.reason


def test_the_item_list_is_built_from_the_redacted_cells():
    cells = [{"row": 1, "col": 1, "text": "Call +44 20 7946 0958"}, {"row": 1, "col": 2, "text": "£1M"}]
    redacted, _ = redact.redact_structure(cells, "Zero2Hero", {})
    listed = structure_items.list_items({"type": "table", "header_rows": 0, "cells": redacted})
    assert [(i["cell"], i["raw"]) for i in listed["items"]] == [("r1c2", "1M")], "no figure of the phone number"


def _roadmap(rows, boxes):
    """A roadmap from a grid of texts; boxes: the box of each column, or {(row, col): box}."""
    structure = v._struct(rows, header_rows=0, kind="roadmap")
    for c in structure["cells"]:
        c["box"] = boxes[(c["row"], c["col"])] if isinstance(boxes, dict) else boxes[c["col"] - 1]
    return structure


def test_a_roadmap_lists_its_date_cells_and_text_lines_below_its_items():
    """Spec section 2: date cells are date labels (deck-parser.md section 7: a date with at most two other words);
    every other non-empty cell is a text line. A figure in a line is an item too."""
    roadmap = _roadmap([["Launch the API", "2026"], ["Launch Q3 2025", "Break even"],
                        ["Hire 5 engineers", "Launched the web app to our first users in January 2011"]], [1, 2])
    listed = structure_items.list_items(roadmap)
    assert listed["dates"] == [{"id": "d1", "cell": "r1c2"}, {"id": "d2", "cell": "r2c1"}]
    assert listed["lines"] == [{"id": "t1", "cell": "r1c1"}, {"id": "t2", "cell": "r2c2"}, {"id": "t3", "cell": "r3c1"},
                               {"id": "t4", "cell": "r3c2"}], "a sentence holding a date is a line"
    assert structure_items.text(roadmap, listed).split("items:\n")[1] == \
        'i1 r3c1 "5" 5 h r1c1\nd1 r1c2\nd2 r2c1\nt1 r1c1\nt2 r2c2\nt3 r3c1\nt4 r3c2'
    table = {**roadmap, "type": "table"}
    assert (structure_items.list_items(table)["dates"], structure_items.list_items(table)["lines"]) == ([], []), \
        "only a roadmap lists its dates and lines"


def test_the_buffer_timeline_lists_six_dates_and_six_lines():
    pytest.importorskip("pptx")
    roadmap = _structure({"file": "03-buffer.pptx", "page": 6, "type": "roadmap"})
    listed = structure_items.list_items(roadmap)
    assert [d["cell"] for d in listed["dates"]] == [f"r{n}c1" for n in (2, 4, 6, 8, 10, 12)]
    assert [t["cell"] for t in listed["lines"]] == [f"r{n}c1" for n in (1, 3, 5, 7, 9, 11)]


# ---------------------------------------------------------------------------
# Units: 20 listed currencies; any other ISO currency is "other", with its code in unit_other
# ---------------------------------------------------------------------------
LISTED_CURRENCIES = ("EUR", "USD", "GBP", "CHF", "BGN", "RON", "PLN", "CZK", "HUF", "SEK", "NOK", "DKK", "TRY", "UAH",
                     "RSD", "JPY", "CNY", "INR", "AUD", "CAD")


def test_deck_structures_use_the_labelling_schema_and_a_column_mapping_keeps_its_own():
    """Spec section 3: one labelling schema for every deck type, {"type", "labels", "pairs"}; a label holds exactly
    item, metric, period, unit, unit_other and actual_or_forecast: no value, cell id or flag. Column mapping keeps
    its own schema, and the gateway picks the schema by type."""
    from app.llm import schemas
    schema = schemas.labelling_output_schema()
    assert list(schema["properties"]) == ["type", "labels", "pairs"] == schema["required"]
    assert schema["additionalProperties"] is False and schema["properties"]["type"]["enum"] == list(schemas.DECK_TYPES)
    label = schema["properties"]["labels"]["items"]
    assert list(label["properties"]) == ["item", "metric", "period", "unit", "unit_other", "actual_or_forecast"]
    assert label["required"] == list(label["properties"]) and label["additionalProperties"] is False
    assert label["properties"]["metric"]["enum"] == [*schemas.CLAIM_METRICS, "use_of_funds", "other", "not_a_metric"]
    pair = schema["properties"]["pairs"]["items"]
    assert pair["properties"] == {"line": {"type": "string"}, "date": {"type": "string"}, "category": {
        "type": "string", "enum": ["launch", "feature", "expansion", "partnership", "hiring", "break_even", "funding",
                                   "certification", "other"]}}
    assert pair["required"] == ["line", "date", "category"] and pair["additionalProperties"] is False
    assert all(schemas.output_schema(kind) == schema for kind in schemas.DECK_TYPES)
    assert schemas.output_schema("column_mapping") == schemas.structure_output_schema()
    assert set(schemas.output_schema("column_mapping")["properties"]) == {"type", "items"}
    result, adapter = _read(_db(), MAPPING_TEXT, "column_mapping", replies=[ONE_MAPPING])
    assert (result.status, result.items, result.labels) == ("read", ONE_MAPPING["items"], [])
    assert set(adapter.requests[0]["json_schema"]["properties"]) == {"type", "items"}
    assert set(adapter.count_requests[-1]["json_schema"]["properties"]) == {"type", "items"}


def test_each_schemas_hash_enters_the_cache_key_of_its_own_type(monkeypatch):
    """Spec section 3: each schema's hash enters the cache key, so a reading stored under another schema is never
    served: the next read misses the cache and calls the model. A change to one schema moves only its own keys."""
    from app.llm import schemas
    db = _db()
    first, _ = _read(db)
    mapping, _ = _read(db, MAPPING_TEXT, "column_mapping", replies=[ONE_MAPPING])
    again, adapter = _read(db)
    assert (again.cache_hit, adapter.calls, again.key) == (True, 0, first.key)
    labelling = schemas.labelling_output_schema()
    monkeypatch.setattr(schemas, "labelling_output_schema", lambda: {**labelling, "required": ["labels", "type", "pairs"]})
    changed, adapter = _read(db)
    assert (changed.status, changed.cache_hit, adapter.calls) == ("read", False, 1) and changed.key != first.key
    kept, adapter = _read(db, MAPPING_TEXT, "column_mapping", replies=[ONE_MAPPING])
    assert (kept.cache_hit, kept.key, adapter.calls) == (True, mapping.key, 0), "the column-mapping key stays"
    old = schemas.structure_output_schema()
    monkeypatch.setattr(schemas, "structure_output_schema", lambda: {**old, "required": ["items", "type"]})
    moved, adapter = _read(db, MAPPING_TEXT, "column_mapping", replies=[ONE_MAPPING])
    assert (moved.cache_hit, adapter.calls) == (False, 1) and moved.key != mapping.key


def test_the_prompt_is_v4_and_names_every_label_field_metric_category_and_tie_break():
    """The prompt lists exactly the label fields, every metric (other and not_a_metric too), the roadmap categories
    and the tie-breaks; a unit is one of the 20 listed currency codes, or "other" with its ISO code in unit_other.
    The column-mapping section keeps its own item fields. v4 (issue #48): its input section reads a title cell."""
    import re
    from app.llm import prompt_store, schemas
    prompt = prompt_store.load(gateway.STRUCTURE_PROMPT)
    assert prompt.version == "v4"
    text = prompt.text
    given = text.split("# The input")[1].split("# Deck structures")[0]
    assert "`r1c1 title: 2011 Estimated Revenue`" in given and "label of the value" in given
    fields = text.split("Each label has exactly these fields:")[1].split("\n\n")[1]
    named = re.findall(r"^- `(\w+)`:", fields, re.M)
    assert named == list(schemas.labelling_output_schema()["properties"]["labels"]["items"]["properties"])
    unit, unit_other = (re.search(rf"^- `{f}`:(.*?)(?=^- `)", fields, re.M | re.S).group(1)
                        for f in ("unit", "unit_other"))
    assert "20 currency codes" in unit and "`other`" in unit
    assert "`other`" in unit_other and "ISO code" in unit_other and "null" in unit_other
    for word in [*schemas.LABEL_METRICS, *schemas.ROADMAP_CATEGORIES]:
        assert f"`{word}`" in text, word
    ties = text.split("Tie-breaks:")[1].split("\n\n")[0]
    assert "time figures are `product`, unless a user count is named" in ties
    assert '"% of marketplace" and market share are `market`' in ties
    assert "commission and take rate are `sales`" in ties
    mapping = text.split("# Column mapping")[1].split("Each item has exactly these fields:")[1]
    heads = [line.split(":")[0] for line in mapping.splitlines() if line.startswith("- `")]
    named = [name for head in heads for name in re.findall(r"`(\w+)`", head)]
    assert sorted(named) == sorted(schemas.structure_output_schema()["properties"]["items"]["items"]["properties"])


def test_the_labelling_schema_lists_20_currencies_and_other_with_the_code_in_unit_other():
    from app.llm import schemas
    schema = schemas.labelling_output_schema()
    label = schema["properties"]["labels"]["items"]
    assert label["properties"]["unit"] == {"anyOf": [
        {"type": "string", "enum": [*LISTED_CURRENCIES, "other", "%", "x", "count", "days", "months", "years"]},
        {"type": "null"}]}
    assert label["properties"]["unit_other"] == {"anyOf": [{"type": "string"}, {"type": "null"}]}
    assert set(LISTED_CURRENCIES) < set(schemas.ISO_CURRENCIES), "listed codes are ISO 4217 codes"
    assert not [code for code in schemas.ISO_CURRENCIES if code not in LISTED_CURRENCIES and code in json.dumps(schema)]


ABSENT = object()


@pytest.mark.parametrize("unit, unit_other, valid", [
    ("other", "ZAR", True),
    ("GBP", None, True),
    ("%", None, True),
    (None, None, True),
    ("ZAR", None, False),                   # an unlisted code is not a unit of its own any more
    ("other", None, False),                 # "other" names its code
    ("other", "USD", False),                # a listed currency has its own value
    ("other", "XYZ", False),                # not an ISO 4217 code
    ("other", "zar", False),
    ("other", "Jane Doe (CEO)", False),     # free text never passes
    ("GBP", "GBP", False),                  # a code stands only beside "other"
    ("%", "ZAR", False),
    ("GBP", ABSENT, False),                 # exactly the schema's fields: unit_other is always written
])
def test_unit_other_holds_an_unlisted_iso_code_and_stands_only_beside_other(unit, unit_other, valid):
    label = {**REPLY["labels"][0], "unit": unit, "unit_other": unit_other}
    if unit_other is ABSENT:
        del label["unit_other"]
    reply = json.dumps({**REPLY, "labels": [label, REPLY["labels"][1]]})
    if valid:
        parsed = gateway.parse_structure_reply(reply, "table", TEXT)
        assert (parsed.labels[0].unit, parsed.labels[0].unit_other) == (unit, unit_other)
    else:
        with pytest.raises(gateway.GatewayError, match="did not match the schema"):
            gateway.parse_structure_reply(reply, "table", TEXT)


def _rows(structure, labels, model_type=None, pairs=()):
    """The approval rows process_deck would add for a structure read with these labels."""
    listed = structure_items.list_items(structure)
    checked = verify.verify(structure, listed, labels, pairs)
    return [structures.candidate_from_item(item, structure, {"file": "plan.pdf"}, model_type, 12)
            for item in structures.approval_items(checked["items"])]


def test_a_currency_outside_the_list_is_read_as_other_and_reaches_the_approval_row_as_its_code():
    label = _label("i1", period="FY2025", unit="other", unit_other="ZAR")
    result, _ = _read(_db(), replies=[{**REPLY, "labels": [label, REPLY["labels"][1]]}])
    assert result.status == "read", result.reason
    assert (result.labels[0]["unit"], result.labels[0]["unit_other"]) == ("other", "ZAR")

    def row(**unit):
        found, _ = _rows(PLAN, [{**label, **unit}, REPLY["labels"][1]])
        return found["currency"], found["unit"], found["ai_label"]
    assert row() == ("ZAR", None, "Verified")
    assert row(unit="GBP", unit_other=None) == ("GBP", None, "Verified")
    assert row(unit="%", unit_other=None) == (None, "%", "Verified")


PLAN = v._struct([["", "FY2025", "FY2026"], ["Revenue", "£1,200,000", "£1,500,000"]])


def test_an_ambiguous_figure_reaches_its_row_with_both_readings_and_the_default_first():
    """Spec section 1: the approval row shows both readings; the default is the row's value, pre-selected."""
    structure = v._struct([["", "Plan"], ["Hours", "Approx. 2.500 hours"], ["P&L", "Net loss (1,200)"], ["Users", "5K"]])
    rows = _rows(structure, [_label("i1", "product"), _label("i2", "net_profit"), _label("i3", "users", unit="count")])
    assert [(r["value"], r["readings"]) for r in rows] == [
        (2500, [{"value": 2500, "dot_reading": "thousands", "bracket_reading": None},
                {"value": 2.5, "dot_reading": "decimal", "bracket_reading": None}]),
        (-1200, [{"value": -1200, "dot_reading": None, "bracket_reading": "negative"},
                 {"value": 1200, "dot_reading": None, "bracket_reading": "positive"}]),
        (5000, [])], "one reading: nothing to choose"
    assert [r["ai_label"] for r in rows] == ["Verified"] * 3


def test_a_ranges_two_items_make_one_approval_row_low_and_high():
    """A range is one row, its value the low end and value_high the high end, from the low end's label. It is Verified
    only when both ends are labelled with the same metric and Verified."""
    structure = v._struct([["", "Plan"], ["Revenue", "$12 -$13 million"], ["Users", "5K"]])
    rows = _rows(structure, [_label("i1", unit="USD"), _label("i2", unit="USD"), _label("i3", "users", unit="count")])
    assert [(r["item"], r["value"], r["value_high"], r["ai_label"], r["readings"]) for r in rows] == [
        ("i1", 12000000, 13000000, "Verified", []), ("i3", 5000, None, "Verified", [])]
    for labels in ([_label("i1", unit="USD"), _label("i2", "not_a_metric"), _label("i3", "users")],
                   [_label("i1", unit="USD"), _label("i2", "costs", unit="USD"), _label("i3", "users")],
                   [_label("i1", unit="USD"), _label("i2", "other", unit="USD"), _label("i3", "users")]):
        row = _rows(structure, labels)[0]
        assert (row["value"], row["value_high"], row["ai_label"]) == (12000000, 13000000, "AI suggestion, not verified"), \
            labels[1]["metric"]
    assert [r["item"] for r in _rows(structure, [_label("i1", "not_a_metric"), _label("i2", unit="USD"),
                                                 _label("i3", "users")])] == ["i2", "i3"], "from the end that is kept"


def test_the_rows_cell_citation_is_unchanged_and_names_the_type_python_sent():
    structure = v._struct([["", "Members"], ["Social", "Discord(150) Telegram(30K)"]])
    rows = _rows(structure, [_label("i1", "users", unit="count"), _label("i2", "users", unit="count")], "kpi_panel")
    assert [(r["sources"][0]["cell"], r["cell"], r["position"], r["item"]) for r in rows] == \
        [("r2c2", "r2c2", 1, "i1"), ("r2c2", "r2c2", 2, "i2")], "the cell id as before; the position beside it"
    assert {r["sources"][0]["structure"] for r in rows} == {"table"}, "a corrected type is only logged"
    assert [r["snippet"] for r in rows] == ["Discord(150) Telegram(30K)"] * 2


def test_an_other_item_is_listed_as_type_other_and_approved_only_once_its_type_is_edited(monkeypatch):
    """Spec section 4: the server refuses to approve an "other" item until its type is edited to a claim type."""
    client, db, adapter = _deck_api(monkeypatch)
    _map_revenue(client)
    _upload_deck(client, "02-moz.pdf")
    p20 = [c for c in _deck(client)["candidates"] if c.get("origin") == "ai" and c["sources"][0].get("page") == 20]
    ranged, = [c for c in p20 if c["sources"][0]["cell"] == "r1c2"]
    assert (ranged["value"], ranged["value_high"], ranged["ai_label"]) == (12000000, 13000000, "Verified"), \
        "moz p20 \"$12 -$13 million\": one row, twelve to thirteen million"
    other, = [c for c in p20 if c.get("claim_type") == "other"]
    assert (other["value"], other["ai_label"], other["unit"]) == (9, "AI suggestion, not verified", "months")
    url = f"/api/audits/{AUDIT}/decks/candidates/{other['id']}"
    refused = client.put(url, json={"status": "approved"})
    assert refused.status_code == 400 and "type" in refused.json()["detail"]
    assert client.put(url, json={"value": 10}).status_code == 400, "an edit approves, so it needs a type too"
    assert client.put(url, json={"claim_type": "other"}).status_code == 422, "other is no claim type to choose"
    assert next(c for c in db["deck_candidates"].docs if c["id"] == other["id"])["status"] == "pending"
    edited = client.put(url, json={"claim_type": "product"})
    assert edited.status_code == 200 and edited.json()["status"] == "edited"
    assert edited.json()["parsed"]["claim_type"] == "other", "what the model read stays next to the edit"
    assert client.put(url, json={"status": "approved"}).json()["status"] == "edited"
    rejected = next(c for c in _deck(client)["candidates"] if c.get("ai_label") == "Verified")
    assert client.put(f"/api/audits/{AUDIT}/decks/candidates/{rejected['id']}",
                      json={"status": "rejected"}).status_code == 200


def test_the_adapter_sends_no_tools_and_no_temperature_to_the_provider(monkeypatch):
    sent = {}

    class Messages:
        def create(self, **kwargs):
            sent.update(kwargs)

            class Response:
                content, stop_reason = [type("B", (), {"type": "text", "text": json.dumps(REPLY)})()], "end_turn"
                usage = type("U", (), {"input_tokens": 10, "output_tokens": 5})()
            return Response()

    adapter = gateway.AnthropicAdapter(api_key="test")
    adapter._client = type("C", (), {"messages": Messages()})()
    adapter.complete(model=gateway.STRUCTURE_MODEL, system="s", user_payload="u", max_tokens=10, temperature=None,
                     json_schema={"type": "object"})
    assert "tools" not in sent and "temperature" not in sent and sent["output_config"]["format"]["type"] == "json_schema"


def test_the_adapter_counts_the_text_alone_with_no_system_prompt_and_no_schema():
    sent = []

    class Messages:
        def count_tokens(self, **kwargs):
            sent.append(kwargs)
            return type("R", (), {"input_tokens": 7})()

    adapter = gateway.AnthropicAdapter(api_key="test")
    adapter._client = type("C", (), {"messages": Messages()})()
    assert adapter.count_tokens(model=gateway.STRUCTURE_MODEL, system=None, user_payload="r1c1: Revenue",
                                json_schema=None) == 7
    adapter.count_tokens(model=gateway.STRUCTURE_MODEL, system="s", user_payload="u", json_schema={"type": "object"})
    assert sent[0] == {"model": gateway.STRUCTURE_MODEL, "messages": [{"role": "user", "content": "r1c1: Revenue"}]}
    assert sent[1]["system"] == "s" and sent[1]["output_config"]["format"]["schema"] == {"type": "object"}


def test_the_adapter_raises_on_a_refusal_stop_reason():
    class Messages:
        def create(self, **kwargs):
            class Response:
                content, stop_reason = [], "refusal"
                usage = type("U", (), {"input_tokens": 10, "output_tokens": 0})()
            return Response()
    adapter = gateway.AnthropicAdapter(api_key="test")
    adapter._client = type("C", (), {"messages": Messages()})()
    with pytest.raises(gateway.GatewayError) as exc:
        adapter.complete(model=gateway.STRUCTURE_MODEL, system="s", user_payload="u", max_tokens=10, temperature=None,
                         json_schema={"type": "object"})
    assert exc.value.reason == "model_refused"


def test_a_refusal_is_not_retried_and_the_structure_is_not_read():
    class Refusing(t.FakeAdapter):
        def complete(self, **kwargs):
            self.calls += 1
            raise gateway.GatewayError("model_refused", "declined")
    result, adapter = _read(_db(), adapter=Refusing())
    assert (result.status, result.reason, adapter.calls) == ("not_read", "Not read by AI", 1)


# ---------------------------------------------------------------------------
# Schema violations (spec section 3): one reask, then "Not read by AI"; the type
# ---------------------------------------------------------------------------
@pytest.mark.parametrize("bad", [
    {**REPLY, "comment": "extra field"},
    {**REPLY, "labels": [{**REPLY["labels"][0], "value": 1200000}, REPLY["labels"][1]]},          # no value
    {**REPLY, "labels": [{**REPLY["labels"][0], "value_cell": "r2c2"}, REPLY["labels"][1]]},      # no cell id
    {**REPLY, "labels": [{**REPLY["labels"][0], "proposed_flags": []}, REPLY["labels"][1]]},      # no flag
    {**REPLY, "labels": [{**REPLY["labels"][0], "note": "Revenue grew strongly"}, REPLY["labels"][1]]},
    {**REPLY, "labels": [{**REPLY["labels"][0], "period": "next year"}, REPLY["labels"][1]]},
    {**REPLY, "labels": REPLY["labels"] + [{**REPLY["labels"][1], "item": "i9"}]},                # an unknown id
    {**REPLY, "labels": REPLY["labels"] + [REPLY["labels"][0]]},                                   # a duplicate id
    {**REPLY, "labels": [REPLY["labels"][0]]},                                                     # a missing id
    {**REPLY, "labels": REPLY["labels"] + [{**REPLY["labels"][1], "item": "r2c3"}]},              # a cell, not an id
    {**REPLY, "labels": [REPLY["labels"][0], {**REPLY["labels"][1], "metric": "invoice_date"}]},  # a sheet field
    {**REPLY, "pairs": [{"line": "t1", "date": "d1", "category": "launch"}]},     # pairs, not sent as a roadmap
    {**REPLY, "type": "roadmap", "pairs": [{"line": "t1", "date": "d1", "category": "launch"}]},
    {"type": "column_mapping", "items": []},                                       # a deck structure is never one
    {"type": "table", "items": []},                                                # the old reply
    "not json",
])
def test_a_reply_that_fails_validation_is_asked_again_once_then_not_read(bad):
    reply = bad if isinstance(bad, str) else json.dumps(bad)
    adapter = t.FakeAdapter(replies=[reply, reply])
    result, _ = _read(_db(), adapter=adapter)
    assert (result.status, result.reason, adapter.calls) == ("not_read", "Not read by AI", 2)
    adapter = t.FakeAdapter(replies=[reply, json.dumps(REPLY)])
    result, _ = _read(_db(), adapter=adapter)
    assert (result.status, adapter.calls) == ("read", 2), "the one reask recovers"


def _pair(line, date, category="launch"):
    return {"line": line, "date": date, "category": category}


@pytest.mark.parametrize("pairs", [
    [_pair("t9", "d1")],                                  # a line that was not listed
    [_pair("t1", "d9")],                                  # a date that was not listed
    [_pair("d1", "d2")],                                  # a date is not a line
    [_pair("t1", "t2")],                                  # a line is not a date
    [_pair("i1", "d1")],                                  # an item is not a line
    [_pair("t1", "d1"), _pair("t1", "d2", "hiring")],     # a line has at most one pair
    [_pair("t1", "d1", "ipo")],                           # not a category
    [{**_pair("t1", "d1"), "period": "2025-Q3"}],         # exactly line, date and category
])
def test_a_roadmap_pair_must_be_one_listed_line_and_one_listed_date_with_a_category(pairs):
    bad = json.dumps({**ROADMAP_REPLY, "pairs": pairs})
    result, adapter = _read(_db(), ROADMAP_TEXT, "roadmap", adapter=t.FakeAdapter(replies=[bad, bad]))
    assert (result.status, result.reason, adapter.calls) == ("not_read", "Not read by AI", 2)


def test_pairs_stand_only_on_a_structure_python_sent_as_a_roadmap_whatever_type_the_reply_gives():
    """The reply follows the type Python sent: pairs of listed lines and dates are a violation on any other type,
    even when the reply corrects the type to roadmap."""
    for reply_type in ("table", "roadmap"):
        reply = json.dumps({**ROADMAP_REPLY, "type": reply_type})
        with pytest.raises(gateway.GatewayError, match="pairs on a structure not sent as a roadmap"):
            gateway.parse_labelling_reply(reply, "table", ROADMAP_TEXT)
    assert gateway.parse_labelling_reply(json.dumps({**ROADMAP_REPLY, "type": "table"}), "roadmap", ROADMAP_TEXT).pairs


def test_a_roadmap_reply_pairs_lines_with_dates_and_a_date_may_serve_several_lines():
    result, _ = _read(_db(), ROADMAP_TEXT, "roadmap", replies=[ROADMAP_REPLY])
    assert result.status == "read", result.reason
    assert (result.labels, result.pairs) == (ROADMAP_REPLY["labels"], ROADMAP_REPLY["pairs"])
    shared = {**ROADMAP_REPLY, "pairs": [_pair("t1", "d1"), _pair("t2", "d1", "hiring")]}
    result, _ = _read(_db(), ROADMAP_TEXT, "roadmap", replies=[shared])
    assert result.status == "read" and len(result.pairs) == 2
    result, _ = _read(_db(), ROADMAP_TEXT, "roadmap", replies=[{**ROADMAP_REPLY, "pairs": []}])
    assert result.status == "read", "a line may stay unpaired"


def test_each_roadmap_category_maps_to_a_claim_type():
    """Spec section 2: hiring -> people, break_even -> ebitda, funding -> other (type Other), the rest -> product."""
    from app.llm import schemas
    assert verify.MILESTONE_TYPES == {"launch": "product", "feature": "product", "expansion": "product",
                                      "partnership": "product", "hiring": "people", "break_even": "ebitda",
                                      "funding": "other", "certification": "product", "other": "product"}
    assert tuple(verify.MILESTONE_TYPES) == schemas.ROADMAP_CATEGORIES


def test_the_reply_follows_the_type_python_sent_and_a_corrected_type_is_only_logged(caplog):
    import logging
    db = _db()
    with caplog.at_level(logging.INFO):
        result, _ = _read(db, replies=[{**REPLY, "type": "unit_economics"}])
    assert (result.status, result.type, result.model_type) == ("read", "table", "unit_economics")
    assert result.labels == REPLY["labels"]
    stored, = db[gateway.STRUCTURES_COLLECTION].docs
    assert (stored["type"], stored["model_type"]) == ("table", "unit_economics")
    assert "type_change=table->unit_economics" in caplog.text
    result, adapter = _read(_db(), MAPPING_TEXT, "column_mapping", replies=[REPLY, REPLY])
    assert result.status == "not_read" and adapter.calls == 2, "a column mapping stays a column mapping"


def test_the_schema_lists_match_the_claim_types_and_the_mapping_fields():
    os.environ.setdefault("MONGO_URL", "mongodb://localhost:27017")
    os.environ.setdefault("DB_NAME", "structure_reading_test")
    from app.decks import claims
    from app.llm import schemas
    import server
    fields = {f for d in server.FIELD_DEFS.values() for group in ("required", "optional") for f in d[group]}
    assert schemas.CLAIM_METRICS == claims.CLAIM_TYPES
    assert set(schemas.MAPPING_FIELDS) == fields
    assert set(schemas.STRUCTURE_METRICS) == set(claims.CLAIM_TYPES) | {"use_of_funds"} | fields
    assert schemas.LABEL_METRICS == claims.CLAIM_TYPES + ("use_of_funds", "other", "not_a_metric")
    assert StructureReply.model_json_schema()["additionalProperties"] is False


# ---------------------------------------------------------------------------
# Consent, cache, caps
# ---------------------------------------------------------------------------
@pytest.mark.parametrize("consent", [False, None])
def test_without_consent_no_call_is_made(consent):
    audit = {"structure_reading_consent": consent} if consent is not None else {}
    db = t.FakeDB()
    db["audits"].docs.append({k: v for k, v in {**AUDIT_DOC, **audit}.items()
                              if consent is not None or k != "structure_reading_consent"})
    result, adapter = _read(db)
    assert result.status == "no_consent" and adapter.calls == 0 and getattr(adapter, "counted", 0) == 0


def test_a_cache_hit_makes_no_call_and_no_audit_is_served_another_audits_result():
    db = _db()
    first, adapter = _read(db)
    again, cached = _read(db)
    assert first.status == again.status == "read" and again.cache_hit and cached.calls == 0
    assert (again.labels, again.pairs) == (first.labels, first.pairs) and again.key == first.key
    db["audits"].docs.append({**AUDIT_DOC, "id": "audit-other"})
    other = t.FakeAdapter(replies=[json.dumps(REPLY)])
    asyncio.run(gateway.read_structure(db, "audit-other", TEXT, "table", adapter=other, sleep=t._noop_sleep))
    assert other.calls == 1, "the same text in another audit is read again, not served from this audit"
    bypass, adapter = _read(db, use_cache=False)
    assert adapter.calls == 1 and not bypass.cache_hit


def test_the_cache_key_covers_text_type_prompt_and_model():
    key = gateway.structure_key(TEXT, "table", "r4:v1", "claude-sonnet-5-5")
    assert key != gateway.structure_key(TEXT + "x", "table", "r4:v1", "claude-sonnet-5-5")
    assert key != gateway.structure_key(TEXT, "chart", "r4:v1", "claude-sonnet-5-5")
    assert key != gateway.structure_key(TEXT, "table", "r5:v1", "claude-sonnet-5-5")
    assert key != gateway.structure_key(TEXT, "table", "r4:v1", "claude-opus-5-5")


LISTED = structure_items.text(v._struct([["", "FY2025", "FY2026"], ["Revenue", "£1,200,000", "£1,500,000"]]),
                              structure_items.list_items(v._struct([["", "FY2025", "FY2026"],
                                                                    ["Revenue", "£1,200,000", "£1,500,000"]])))


def test_a_structure_over_4000_tokens_of_text_and_item_list_is_not_sent():
    """Spec section 5: the cap is 4,000 tokens per structure, counted on the structure text plus the item list."""
    assert gateway.STRUCTURE_INPUT_CAP == 4000 and "\nitems:\ni1 r2c2 " in LISTED
    adapter = t.FakeAdapter(replies=[json.dumps(REPLY)])
    adapter.text_tokens, adapter.input_tokens = 4001, 4500
    result, _ = _read(_db(), LISTED, adapter=adapter)
    assert (result.status, result.reason, adapter.calls) == ("too_large", "Too large for AI reading", 0)
    assert adapter.count_requests == [{"system": None, "user_payload": LISTED, "json_schema": None}], \
        "the structure text with its item list is counted, and nothing more once it is over"
    adapter = t.FakeAdapter(replies=[json.dumps(REPLY)])
    adapter.text_tokens, adapter.input_tokens = 4000, 4500
    result, _ = _read(_db(), LISTED, adapter=adapter)
    assert result.status != "too_large" and adapter.calls == 1, "exactly at the cap: sent"


def test_the_4000_cap_is_on_the_text_and_item_list_not_the_prompt_and_schema():
    db = _db()
    adapter = t.FakeAdapter(replies=[json.dumps(REPLY)])
    adapter.text_tokens, adapter.input_tokens = 3900, 5000
    result, _ = _read(db, adapter=adapter)
    assert (result.status, adapter.calls) == ("read", 1)
    whole = adapter.count_requests[1]
    assert whole["system"] and whole["json_schema"] and json.loads(whole["user_payload"])["text"] == TEXT
    # The 400,000 cap counts the whole call: 391,300 used + 5,000 input + 4,000 max_tokens is over
    # (with the text's 3,900 it would fit).
    db["llm_calls"].docs.append({"run_id": AUDIT, "step": "structures", "cache_hit": False, "input_tokens": 385000,
                                 "output_tokens": 1000, "estimated_cost_usd": 0.0, "timestamp": "2026-10-05T00:00:00"})
    adapter = t.FakeAdapter(replies=[json.dumps(REPLY)])
    adapter.text_tokens, adapter.input_tokens = 3900, 5000
    result, _ = _read(db, adapter=adapter, use_cache=False)
    assert result.status == "stopped" and adapter.calls == 0


def test_the_400000_token_cap_counts_billed_input_and_output_and_stops_with_the_spec_message():
    db = _db()
    db["llm_calls"].docs.append({"run_id": AUDIT, "step": "structures", "cache_hit": False, "input_tokens": 390000,
                                 "output_tokens": 4000, "estimated_cost_usd": 0.0, "timestamp": "2026-10-05T00:00:00"})
    adapter = t.FakeAdapter(replies=[json.dumps(REPLY)])
    adapter.input_tokens = 2001            # 394,000 used + 2,001 input + 4,000 max_tokens > 400,000
    result, _ = _read(db, adapter=adapter)
    assert result.status == "stopped" and adapter.calls == 0
    assert result.reason == ("AI reading stopped: this audit reached its 400,000-token limit. The remaining "
                             "structures were read by Python only.")
    adapter.input_tokens = 2000            # exactly at the cap: the call goes out
    result, _ = _read(db, adapter=adapter)
    assert result.status == "read" and adapter.calls == 1


def test_an_audit_of_8_decks_of_5_structures_fits_the_token_cap():
    """At the live run's 4,156 billed input tokens a read and about 300 output tokens, the 40 reads of an audit of 8
    decks of 5 structures each use about 178,000 tokens. The cap stops reading at the 89th read: each call must leave
    room for its input and its 4,000 max_tokens."""
    class Read(t.FakeAdapter):
        input_tokens, text_tokens = 4156, 100

        def complete(self, **kwargs):
            reply, tokens_in, _ = super().complete(**kwargs)
            return reply, tokens_in, 300
    db, adapter = _db(), Read(replies=[json.dumps({"type": "table", "labels": [], "pairs": []})])
    statuses = [_read(db, f"r1c1: Revenue\nr1c2: {n}", adapter=adapter)[0].status for n in range(89)]
    assert statuses == ["read"] * 88 + ["stopped"]
    assert asyncio.run(guards.structure_tokens_used(db, AUDIT)) == 88 * 4456


def test_a_cap_refusal_logs_a_not_read_line(monkeypatch, caplog):
    """A structure the daily spend cap or the 400,000-token cap refuses was not read and printed nothing, so a run
    that hit a cap showed only "0 of N structures read". It now logs reason=spend_cap or reason=token_cap."""
    import logging
    from datetime import datetime, timezone
    db = _db()
    db["llm_calls"].docs.append({"run_id": AUDIT, "step": "structures", "cache_hit": False, "input_tokens": 394000,
                                 "output_tokens": 4000, "estimated_cost_usd": 0.0, "timestamp": "2026-10-05T00:00:00"})
    with caplog.at_level(logging.INFO):
        result, adapter = _read(db)
    assert (result.status, adapter.calls) == ("stopped", 0)
    assert caplog.text.count("structure not read:") == 1
    assert "structure not read: run_id=audit-s step=structures hash=" in caplog.text
    assert "reason=token_cap status=- type=-" in caplog.text
    caplog.clear()
    monkeypatch.setenv("LLM_DAILY_SPEND_CAP_USD", "0.01")
    db = _db()
    db["llm_calls"].docs.append({"run_id": "elsewhere", "cache_hit": False, "estimated_cost_usd": 1.0,
                                 "timestamp": datetime.now(timezone.utc).isoformat()})
    with caplog.at_level(logging.INFO):
        result, adapter = _read(db)
    assert (result.status, adapter.calls) == ("not_read", 0)
    assert caplog.text.count("structure not read:") == 1 and "reason=spend_cap status=- type=-" in caplog.text
    assert db["llm_calls"].docs[-1]["status"] == "daily_spend_cap_exceeded", "llm_calls keeps the guard's code"


def test_the_15_call_cap_counts_narrative_calls_only():
    db = t.make_db()
    for _ in range(guards.MAX_CALLS_PER_RUN):
        db["llm_calls"].docs.append({"run_id": t.RUN_ID, "step": "structures", "cache_hit": False,
                                     "input_tokens": 10, "output_tokens": 10, "estimated_cost_usd": 0.0,
                                     "timestamp": "2026-10-05T00:00:00+00:00"})
    adapter = t.FakeAdapter()
    result = asyncio.run(gateway.generate_narrative(db, t.RUN_ID, "growth_engine", adapter=adapter, sleep=t._noop_sleep))
    assert adapter.calls == 1 and result.narrative_status in ("ok", "flagged"), "structure calls never use the call cap"
    assert asyncio.run(guards.calls_made(db, t.RUN_ID)) == 1


def test_the_daily_spend_cap_and_the_lock_apply_to_structure_calls(monkeypatch):
    from datetime import datetime, timezone
    monkeypatch.setenv("LLM_DAILY_SPEND_CAP_USD", "0.01")
    db = _db()
    db["llm_calls"].docs.append({"run_id": "elsewhere", "cache_hit": False, "estimated_cost_usd": 1.0,
                                 "timestamp": datetime.now(timezone.utc).isoformat()})
    result, adapter = _read(db)
    assert result.status == "not_read" and adapter.calls == 0
    assert db["llm_locks"].docs and all(not lock["held"] for lock in db["llm_locks"].docs), "lock taken, then released"
    assert {lock["step"] for lock in db["llm_locks"].docs} == {"structures"}


# ---------------------------------------------------------------------------
# Logging, usage, Delete audit
# ---------------------------------------------------------------------------
def test_storage_and_logs_hold_the_model_output_and_metadata_never_the_text_sent():
    db = _db()
    result, _ = _read(db, deck_id="deck-1", page=19)
    stored, = db[gateway.STRUCTURES_COLLECTION].docs
    assert {"audit_id", "key", "content_hash", "type", "model_type", "output", "prompt_version", "model",
            "input_tokens", "output_tokens", "estimated_cost_usd", "deck_id", "page", "created_at"} <= set(stored)
    assert (stored["deck_id"], stored["page"], stored["model"]) == ("deck-1", 19, "claude-sonnet-5-5")
    call, = db["llm_calls"].docs
    assert (call["step"], call["deck_id"], call["content_hash"]) == ("structures", "deck-1", stored["content_hash"])
    blob = json.dumps([stored, call])
    for needle in ("Revenue", "£1,200,000", TEXT):
        assert needle not in blob, f"{needle!r} was stored"


def test_usage_gains_per_deck_cost_and_keeps_narrative_calls_apart():
    db = _db()
    _read(db, deck_id="deck-1")
    _read(db, deck_id="deck-1")             # a cache hit
    usage = asyncio.run(gateway.usage_for_run(db, AUDIT))
    assert usage.calls == 0 and usage.structure_calls == 1 and usage.structure_token_cap == 400000
    deck = usage.by_deck["deck-1"]
    assert (deck.calls, deck.cache_hits, deck.input_tokens, deck.output_tokens) == (1, 1, 1200, 300)
    assert deck.estimated_cost_usd == pytest.approx(gateway.estimate_cost_usd("claude-sonnet-5-5", 1200, 300))


def test_purge_run_removes_the_stored_structure_readings():
    db = _db()
    _read(db, deck_id="deck-1")
    purged = asyncio.run(gateway.purge_run(db, AUDIT))
    assert purged["llm_structures"] == 1 and db[gateway.STRUCTURES_COLLECTION].docs == []
    assert db["llm_calls"].docs == []


# ---------------------------------------------------------------------------
# Column mapping: header stack of at most 3 rows, samples for numeric and date columns, a profile
# only for text columns; one-click confirmation; confirmed mappings reused per header set
# ---------------------------------------------------------------------------
import app.structures as structures  # noqa: E402

SECRET = "Jane Doe Holdings"           # a text cell value: never on the column-mapping path
MAPPING_TEXT = "r1c1: Customer\nr1c2: Amount\nc2 sample: 1200"


def _sheet():
    columns = ["Kunde", "Rechnungsdatum", "Betrag", "Notiz", "Kundennummer"]
    rows = [{"Kunde": f"{SECRET} {i}", "Rechnungsdatum": f"2025-0{i}-28T00:00:00", "Betrag": 100.5 * i,
             "Notiz": f"call {SECRET}", "Kundennummer": 10000 + i} for i in range(1, 6)]
    return columns, rows


def test_numeric_and_date_columns_send_up_to_3_samples_and_text_columns_a_profile_only():
    columns, rows = _sheet()
    text = structures.column_mapping_text(columns, rows, "Target", {}, ("Kundennummer",))
    parsed = redact.parse_column_text(text)
    assert parsed["samples"] == {2: ["2025-01-28", "2025-02-28", "2025-03-28"], 3: ["100.5", "201", "301.5"]}
    assert set(parsed["profiles"]) == {1, 4, 5}, "text columns and the customer column send a profile"
    assert parsed["profiles"][1] == {"distinct": 5, "length": 19, "shape": "Aa Aa Aa 0"}
    assert SECRET.lower() not in text.lower() and "10001" not in text, "no text cell value, no customer number"
    assert redact.column_text_problem(text) is None


def test_the_header_stack_is_capped_at_the_3_rows_nearest_the_data():
    columns = ["Revenue report 2025", "Unnamed: 1", "Unnamed: 2"]
    rows = [{"Revenue report 2025": "Prepared by finance", "Unnamed: 1": None, "Unnamed: 2": None},
            {"Revenue report 2025": None, "Unnamed: 1": 2025, "Unnamed: 2": None},
            {"Revenue report 2025": "Month", "Unnamed: 1": "Jan", "Unnamed: 2": "Feb"},
            {"Revenue report 2025": "Revenue", "Unnamed: 1": 100, "Unnamed: 2": 200}]
    headers, samples, _ = structures.column_mapping_input(columns, rows)
    assert [(c["row"], c["col"], c["text"]) for c in headers] == [
        (1, 1, "Prepared by finance"), (2, 2, "2025"), (3, 1, "Month"), (3, 2, "Jan"), (3, 3, "Feb")]
    assert "Revenue report 2025" not in {c["text"] for c in headers}, "the 4th row from the data is left out"
    assert samples == {2: ["100"], 3: ["200"]}


def _api(monkeypatch, replies, consent=True):
    pytest.importorskip("fastapi")
    os.environ.setdefault("MONGO_URL", "mongodb://localhost:27017")
    os.environ.setdefault("DB_NAME", "structure_reading_test")
    from fastapi.testclient import TestClient
    import server
    db = t.FakeDB()
    db["audits"].docs.append({**AUDIT_DOC, "structure_reading_consent": consent})
    monkeypatch.setattr(server, "db", db)
    adapter = t.FakeAdapter(replies=[json.dumps(r) for r in replies])
    monkeypatch.setattr(gateway, "AnthropicAdapter", lambda *a, **k: adapter)
    return TestClient(server.app, raise_server_exceptions=False), db, adapter


GERMAN = "Kunde,Rechnungsdatum,Betrag,Waehrung\nAcme GmbH,2025-01-31,100,EUR\nBeta AG,2025-02-28,200,EUR\n"
MAPPING_REPLY = {"type": "column_mapping", "items": [
    {"metric": f, "period": None, "value": None, "unit": None, "actual_or_forecast": "unknown", "value_cell": cell,
     "unit_other": None, "period_cells": [], "proposed_flags": []}
    for f, cell in (("customer_id", "r1c1"), ("invoice_date", "r1c2"), ("amount", "r1c3"), ("currency", "r1c4"),
                    ("deal_id", "r1c1"))]}


ONE_MAPPING = {"type": "column_mapping", "items": [MAPPING_REPLY["items"][0]]}     # fits MAPPING_TEXT


def test_the_model_proposes_what_the_aliases_miss_and_the_screen_marks_it_as_a_suggestion(monkeypatch):
    client, db, adapter = _api(monkeypatch, [MAPPING_REPLY])
    r = client.post(f"/api/audits/{AUDIT}/datasets/revenue/upload", files={"file": ("rev.csv", GERMAN)})
    assert r.status_code == 200, r.text
    body = r.json()
    assert adapter.calls == 1 and json.loads(adapter.payloads[0])["type"] == "column_mapping"
    assert body["suggested_mapping"]["customer_id"] == "Kunde" and body["mapping_source"]["customer_id"] == "ai"
    assert body["suggested_mapping"]["invoice_date"] == "Rechnungsdatum"
    assert "deal_id" not in body["suggested_mapping"], "a field of another file type is left out"
    sent = json.loads(adapter.payloads[0])["text"]
    assert "Acme" not in sent and "Beta" not in sent, "customer cells travel as a profile"
    stored, = db[gateway.STRUCTURES_COLLECTION].docs
    assert stored["statuses"] == ["suggestion"] * 5, "a column mapping is never Verified"


def test_one_click_confirms_and_the_confirmed_mapping_is_reused_for_the_same_headers(monkeypatch):
    client, db, adapter = _api(monkeypatch, [MAPPING_REPLY])
    client.post(f"/api/audits/{AUDIT}/datasets/revenue/upload", files={"file": ("rev.csv", GERMAN)})
    mapping = {"customer_id": "Kunde", "invoice_date": "Rechnungsdatum", "amount": "Betrag", "currency": "Waehrung"}
    assert client.put(f"/api/audits/{AUDIT}/datasets/revenue/mapping", json={"mapping": mapping}).status_code == 200
    again = client.post(f"/api/audits/{AUDIT}/datasets/revenue/upload", files={"file": ("rev2.csv", GERMAN)}).json()
    assert adapter.calls == 1, "the stored correction is used: no second model call"
    assert {f: again["suggested_mapping"][f] for f in mapping} == mapping
    assert set(again["mapping_source"].values()) == {"stored"}
    client.delete(f"/api/audits/{AUDIT}")
    assert db[structures.COLUMN_MAPPINGS_COLLECTION].docs == [], "Delete audit removes the stored mapping"


def test_without_consent_the_mapping_comes_from_the_aliases_only(monkeypatch):
    client, _, adapter = _api(monkeypatch, [MAPPING_REPLY], consent=False)
    body = client.post(f"/api/audits/{AUDIT}/datasets/revenue/upload", files={"file": ("rev.csv", GERMAN)}).json()
    assert adapter.calls == 0 and body["ai_reading"]["status"] == "no_consent"
    assert "ai" not in body["mapping_source"].values()


def test_a_proposal_names_a_field_of_this_file_type_and_cites_an_existing_header():
    items = [dict(MAPPING_REPLY["items"][0]), {**MAPPING_REPLY["items"][0], "metric": "deal_id", "value_cell": "r1c3"},
             {**MAPPING_REPLY["items"][0], "metric": "amount", "value_cell": "r1c9"},
             {**MAPPING_REPLY["items"][0], "metric": "currency", "value_cell": "r1c1"}]
    fields = ["customer_id", "invoice_date", "amount", "currency"]
    assert structures.proposed_mapping(items, ["Kunde", "Datum", "Betrag"], fields) == {"customer_id": "Kunde"}, \
        "deal_id is a CRM field, r1c9 does not exist, Kunde is already proposed"


def test_the_aliases_keep_their_fields_and_the_model_fills_the_rest(monkeypatch):
    reply = {"type": "column_mapping", "items": [
        {**MAPPING_REPLY["items"][0], "metric": "amount", "value_cell": "r1c4"},
        {**MAPPING_REPLY["items"][0], "metric": "customer_id", "value_cell": "r1c1"}]}
    client, _, _ = _api(monkeypatch, [reply])
    sheet = "Kunde,Rechnungsdatum,Amount,Betrag,Waehrung\nAcme,2025-01-31,1,2,EUR\n"
    body = client.post(f"/api/audits/{AUDIT}/datasets/revenue/upload", files={"file": ("r.csv", sheet)}).json()
    assert (body["suggested_mapping"]["amount"], body["mapping_source"]["amount"]) == ("Amount", "rules")
    assert (body["suggested_mapping"]["customer_id"], body["mapping_source"]["customer_id"]) == ("Kunde", "ai")


# ---------------------------------------------------------------------------
# Consent (spec section 4)
# ---------------------------------------------------------------------------
NEW_AUDIT = {"company_name": "Zero2Hero", "client_name": "Northbridge Capital", "engagement_reference": "ENG-2026-041"}


def test_consent_is_ticked_at_creation_by_default_and_every_change_is_logged_with_its_time(monkeypatch):
    client, db, _ = _api(monkeypatch, [])
    created = client.post("/api/audits", json=NEW_AUDIT).json()
    assert created["structure_reading_consent"] is True
    assert [e["value"] for e in created["consent_log"]] == [True] and created["consent_log"][0]["at"]
    assert (created["client_name"], created["engagement_reference"]) == ("Northbridge Capital", "ENG-2026-041")
    client.put(f"/api/audits/{created['id']}", json={"structure_reading_consent": False})
    client.put(f"/api/audits/{created['id']}", json={"structure_reading_consent": False})      # no change: no entry
    stored = next(a for a in db["audits"].docs if a["id"] == created["id"])
    assert [e["value"] for e in stored["consent_log"]] == [True, False] and all(e["at"] for e in stored["consent_log"])
    unticked = client.post("/api/audits", json={**NEW_AUDIT, "structure_reading_consent": False}).json()
    assert unticked["structure_reading_consent"] is False and [e["value"] for e in unticked["consent_log"]] == [False]


@pytest.mark.parametrize("missing", ["client_name", "engagement_reference"])
def test_audit_creation_refuses_a_missing_client_name_or_engagement_reference(monkeypatch, missing):
    client, _, _ = _api(monkeypatch, [])
    assert client.post("/api/audits", json={k: v for k, v in NEW_AUDIT.items() if k != missing}).status_code == 422
    assert client.post("/api/audits", json={**NEW_AUDIT, missing: " "}).status_code == 422


def test_an_audit_created_before_consent_stays_unticked_until_it_has_an_engagement_reference(monkeypatch):
    client, db, adapter = _api(monkeypatch, [REPLY])
    db["audits"].docs.append({"id": "old", "company_name": "Old Co", "results": None})
    result = asyncio.run(gateway.read_structure(db, "old", TEXT, "table", adapter=adapter, sleep=t._noop_sleep))
    assert result.status == "no_consent" and adapter.calls == 0, "no field means unticked"
    assert client.put("/api/audits/old", json={"structure_reading_consent": True}).status_code == 400
    assert client.put("/api/audits/old", json={"structure_reading_consent": True,
                                              "engagement_reference": "ENG-9"}).status_code == 200


# ---------------------------------------------------------------------------
# Orchestration: decks wait for the mapped revenue file, then are read, verified and listed
# ---------------------------------------------------------------------------
def unlabelled(sent):
    """The reply to a structure with no recorded reply: every listed item labelled not_a_metric, no pair."""
    listed = redact.parse_item_lines(redact.split_items(sent["text"])[1] or []) or {"items": []}
    return {"type": sent["type"], "pairs": [], "labels": [
        {"item": i["id"], "metric": "not_a_metric", "period": None, "unit": None, "unit_other": None,
         "actual_or_forecast": "unknown"} for i in listed["items"]]}


class RecordedAdapter(t.FakeAdapter):
    """Replays the recorded reply whose structure the request carries; every item not_a_metric otherwise."""

    def __init__(self):
        super().__init__()
        self.fixtures = _fixtures()

    def complete(self, *, model, system, user_payload, max_tokens, temperature, json_schema):
        sent = json.loads(user_payload)
        reply = next((f["reply"] for f in self.fixtures if f["type"] == sent["type"] and f["match"] in sent["text"]),
                     None) or (unlabelled(sent) if sent["type"] != "column_mapping" else {"type": sent["type"], "items": []})
        self._replies = [json.dumps(reply)]
        self.calls_before = self.calls
        return super().complete(model=model, system=system, user_payload=user_payload, max_tokens=max_tokens,
                                temperature=temperature, json_schema=json_schema)


REVENUE = "Customer,Invoice Date,Amount,Currency\nAcme Retail,2025-01-31,100,EUR\nBeta Stores,2025-02-28,200,EUR\n"
REVENUE_MAPPING = {"customer_id": "Customer", "invoice_date": "Invoice Date", "amount": "Amount", "currency": "Currency"}


def _deck_api(monkeypatch, consent=True):
    pytest.importorskip("pdfplumber")
    client, db, _ = _api(monkeypatch, [], consent=consent)
    adapter = RecordedAdapter()
    monkeypatch.setattr(gateway, "AnthropicAdapter", lambda *a, **k: adapter)
    return client, db, adapter


def _upload_deck(client, name="05-zero2hero.pdf"):
    r = client.post(f"/api/audits/{AUDIT}/decks/upload", files={"file": (name, (DECKS / name).read_bytes())})
    assert r.status_code == 200, r.text
    return r.json()


def _map_revenue(client):
    client.post(f"/api/audits/{AUDIT}/datasets/revenue/upload", files={"file": ("rev.csv", REVENUE)})
    assert client.put(f"/api/audits/{AUDIT}/datasets/revenue/mapping", json={"mapping": REVENUE_MAPPING}).status_code == 200


def _deck(client):
    return client.get(f"/api/audits/{AUDIT}/decks").json()


def test_a_client_name_or_engagement_reference_in_a_deck_cell_goes_out_as_redacted_and_the_structure_is_read():
    from app import structures
    from app.decks import TEXT_COLLECTION
    db = _db()
    db["datasets"].docs.append({"audit_id": AUDIT, "dtype": "revenue", "mapping": REVENUE_MAPPING,
                                "mapped_at": "2026-10-05T00:00:00"})
    cells = [{"row": 1, "col": 1, "text": "Prepared for Northbridge Capital"}, {"row": 1, "col": 2, "text": "FY2025"},
             {"row": 2, "col": 1, "text": "Revenue (eng-2026-041)"}, {"row": 2, "col": 2, "text": "£1,200,000"}]
    db[TEXT_COLLECTION].docs.append({"audit_id": AUDIT, "deck_id": "d1", "file": "plan.pptx", "page_unit": "slide",
                                     "structures": [{"type": "table", "slide": 1, "header_rows": 1, "cells": cells}]})
    adapter = RecordedAdapter()
    status = asyncio.run(structures.process_deck(db, AUDIT, "d1", adapter=adapter, sleep=t._noop_sleep))
    assert (status, adapter.calls) == (structures.READ, 1)
    assert json.loads(adapter.payloads[0])["text"] == \
        "r1c1: Prepared for [redacted]\nr1c2: FY2025\nr2c1: Revenue ([redacted])\nr2c2: £1,200,000\nitems:\n" \
        'i1 r2c2 "1,200,000" 1200000 h r2c1 r1c2', "the item list is built from the redacted cells"


def test_a_client_name_in_a_spreadsheet_header_goes_out_as_redacted(monkeypatch):
    client, _, adapter = _api(monkeypatch, [MAPPING_REPLY])
    sheet = "Kunde,Rechnungsdatum,Betrag Northbridge Capital,Waehrung\nAcme,2025-01-31,1,EUR\n"
    client.post(f"/api/audits/{AUDIT}/datasets/revenue/upload", files={"file": ("r.csv", sheet)})
    sent = json.loads(adapter.payloads[0])["text"]
    assert "r1c3: Betrag [redacted]" in sent and "Northbridge" not in sent


def test_a_corrected_period_is_verified_counted_on_the_deck_and_the_models_period_stays_stored(monkeypatch):
    from app import structures
    from app.decks import TEXT_COLLECTION, CANDIDATES_COLLECTION
    literal = {"type": "table", "pairs": [], "labels": [_label("i1", period="2025-04", unit="USD")]}
    client, db, adapter = _api(monkeypatch, [literal])
    db["audits"].docs[0]["fiscal_year_end"] = 3
    db["datasets"].docs.append({"audit_id": AUDIT, "dtype": "revenue", "mapping": REVENUE_MAPPING,
                                "mapped_at": "2026-10-05T00:00:00"})
    cells = [{"row": 1, "col": 2, "text": "FY2025"}, {"row": 2, "col": 2, "text": "Apr"},
             {"row": 3, "col": 1, "text": "Revenue"}, {"row": 3, "col": 2, "text": "$5M"}]
    db[TEXT_COLLECTION].docs.append({"audit_id": AUDIT, "deck_id": "d1", "file": "plan.pptx", "page_unit": "slide",
                                     "uploaded_at": "2026-10-05T00:00:00",
                                     "structures": [{"type": "table", "slide": 2, "header_rows": 2, "cells": cells}]})
    asyncio.run(structures.process_deck(db, AUDIT, "d1", adapter=adapter, sleep=t._noop_sleep))
    row, = [c for c in db[CANDIDATES_COLLECTION].docs if c.get("origin") == "ai"]
    assert (row["ai_label"], row["target_date"], row["period_start"], row["period_end"]) == \
        ("Verified", "FY2025-04", "2024-04-01", "2024-04-30")
    stored, = db[gateway.STRUCTURES_COLLECTION].docs
    assert stored["output"]["labels"][0]["period"] == "2025-04", "the stored reading keeps the model's period"
    assert stored["periods_corrected"] == 1
    deck, = client.get(f"/api/audits/{AUDIT}/decks").json()["decks"]
    assert deck["periods_corrected"] == 1
    # December: the model's reading holds, so nothing is corrected; the row keeps Python's label.
    client.put(f"/api/audits/{AUDIT}", json={"fiscal_year_end": 12})
    deck, = client.get(f"/api/audits/{AUDIT}/decks").json()["decks"]
    row = next(c for c in db[CANDIDATES_COLLECTION].docs if c.get("origin") == "ai")
    assert deck["periods_corrected"] == 0 and db[gateway.STRUCTURES_COLLECTION].docs[0]["periods_corrected"] == 0
    assert (row["ai_label"], row["target_date"], row["period_start"]) == ("Verified", "FY2025-04", "2025-04-01")


def test_a_deck_uploaded_while_consent_was_off_is_flagged_once_consent_is_on_and_is_not_read(monkeypatch):
    from app.decks import TEXT_COLLECTION
    client, db, adapter = _deck_api(monkeypatch, consent=False)
    _upload_deck(client)
    db[TEXT_COLLECTION].docs.append({"audit_id": AUDIT, "deck_id": "old", "file": "old.pptx", "structures": [],
                                     "uploaded_at": "2020-01-01T00:00:00"})      # uploaded before AI reading existed
    assert [d["uploaded_before_consent"] for d in _deck(client)["decks"]] == [False, False], "off: nothing to say"
    client.put(f"/api/audits/{AUDIT}", json={"structure_reading_consent": True})
    _map_revenue(client)
    assert [d["uploaded_before_consent"] for d in _deck(client)["decks"]] == [True, True]
    deck_reads = lambda: [p for p in adapter.payloads if json.loads(p)["type"] != "column_mapping"]  # noqa: E731
    assert deck_reads() == [], "not read until it is uploaded again"
    _upload_deck(client)
    newest = _deck(client)["decks"][0]
    assert (newest["ai_status"], newest["uploaded_before_consent"]) == ("read", False) and deck_reads()


def test_a_deck_waits_for_the_mapped_revenue_file_and_is_read_once_it_is(monkeypatch):
    client, db, adapter = _deck_api(monkeypatch)
    upload = _upload_deck(client)
    listed = _deck(client)
    assert listed["decks"][0]["ai_status"] == "waiting" and listed["decks"][0]["ai_message"] == "waiting for revenue file"
    assert adapter.calls == 0, "nothing is sent before the customers are in the mapping"
    assert not [c for c in listed["candidates"] if c.get("origin") == "ai"]
    _map_revenue(client)
    calls_for_mapping = 1                    # the revenue upload's column-mapping call is not queued
    assert adapter.calls == calls_for_mapping + upload["structures"], "every structure read once"
    deck = _deck(client)["decks"][0]
    assert deck["ai_status"] == "read" and deck["sent_pages"] == [11, 17, 19, 22]
    assert deck["ai_cost_usd"] == pytest.approx(upload["structures"] * gateway.estimate_cost_usd("claude-sonnet-5-5", 1200, 300))
    usage = client.get(f"/api/runs/{AUDIT}/llm-usage").json()
    assert usage["by_deck"][upload["deck_id"]]["calls"] == upload["structures"] and usage["calls"] == 0
    assert usage["structure_calls"] == upload["structures"] + calls_for_mapping


def test_results_show_in_the_approval_list_labelled_and_citing_their_cell(monkeypatch):
    client, db, adapter = _deck_api(monkeypatch)
    _map_revenue(client)
    _upload_deck(client)
    ai = [c for c in _deck(client)["candidates"] if c.get("origin") == "ai"]
    panel = [c for c in ai if c["sources"][0]["structure"] == "kpi_panel" and c["sources"][0]["page"] == 19]
    assert sorted((c["claim_type"], c["value"], c["ai_label"], c["sources"][0]["cell"]) for c in panel) == [
        ("gross_profit", 150000, "Verified", "r2c1"), ("users", 5000, "Verified", "r3c1")]
    assert panel[0]["period_text"] == "23 Y/E" and panel[0]["target_date"] == "2023"
    table = [c for c in ai if c["sources"][0]["structure"] == "table"]
    assert table == [], "the 24 table values Python already lists in the same cells are not listed twice"
    stored = next(s["ai"] for s in db["deck_text"].docs[0]["structures"] if s["type"] == "table")
    assert (stored["status"], stored["not_a_metric"]) == ("read", 0)
    assert all(c["status"] == "pending" for c in ai)


def test_a_roadmaps_pairs_become_milestone_rows_dated_by_their_date_cells(monkeypatch):
    """Buffer p6, recorded: each line paired with the date below it. A milestone row cites its line's cell, takes
    its category's claim type and its date cell's period, has no value and is never Verified."""
    pytest.importorskip("pptx")
    client, db, adapter = _deck_api(monkeypatch)
    _map_revenue(client)
    _upload_deck(client, "03-buffer.pptx")
    rows = [c for c in _deck(client)["candidates"] if c.get("origin") == "ai" and c["value"] is None
            and c["sources"][0]["structure"] == "roadmap"]
    assert sorted((c["sources"][0]["cell"], c["claim_type"], c["target_date"], c["period_text"], c["ai_label"])
                  for c in rows) == [
        ("r11c1", "product", "2013-01", "January 2013", "AI suggestion, not verified"),
        ("r1c1", "product", "2011-01", "January 2011", "AI suggestion, not verified"),
        ("r3c1", "product", "2011-10", "October 2011", "AI suggestion, not verified"),
        ("r5c1", "product", "2011-10", "October 2011", "AI suggestion, not verified"),
        ("r7c1", "product", "2011-12", "December 2011", "AI suggestion, not verified"),
        ("r9c1", "product", "2012-01", "January 2012", "AI suggestion, not verified")]


# moz p2's 9 dated paragraphs, each with its date, as the page prints them (issue #49).
MOZ_P2_MILESTONES = [
    ("Rand starts working w/ Gillian building websites for small, local businesses", "1997"),
    ("Deeply in debt, and failing to get traffic to clients’ sites, Rand starts the SEOmoz Blog as part of learning the "
     "SEO process.", "2004"),
    ("SEOmoz takes an investment of $1.1M from Ignition Partners & Curious Office", "Nov. 2007"),
    ("Moz’scollection of tools becomes a singular, campaign-based web app. Prices rise to $99 / $499 / $1999 per month.",
     "Sept. 2010"),
    ("Gillian (Rand’s Mom) founds the company that will become SEOmoz", "1981"),
    ("Rand drops out of UW, 2 classes from graduation to work full time w/ Gillian", "2001"),
    ("SEOmoz launches its first subscription software product, “PRO” for $39/month", "Feb. 2007"),
    ("Linkscape, SEOmoz’sweb index and link graph, launches. By December, mozis profitable.", "Oct. 2008"),
    ("SEOmozis moving from just “SEO” to social media, content marketing, analytics, local and video. To this end, we’ve "
     "acquired “Moz.com.”", "July 2011"),
]


def column_dates(sent):
    """Issue #49: the reply the live reading gave on moz p2, as a rule: every text line paired with the nearest date
    cell in its grid column; every item not_a_metric."""
    listed = redact.parse_item_lines(redact.split_items(sent["text"])[1] or [])

    def at(cell):
        row, col = cell[1:].split("c")
        return int(row), int(col)
    pairs = []
    for line in listed["lines"]:
        row, col = at(line["cell"])
        dates = [d for d in listed["dates"] if at(d["cell"])[1] == col]
        if dates:
            date = min(dates, key=lambda d: abs(at(d["cell"])[0] - row))
            pairs.append({"line": line["id"], "date": date["id"], "category": "launch"})
    return {**unlabelled(sent), "pairs": pairs}


class ColumnDatesAdapter(RecordedAdapter):
    """A roadmap read with column_dates, whatever its text; every other structure as RecordedAdapter reads it."""

    def complete(self, *, model, system, user_payload, max_tokens, temperature, json_schema):
        sent = json.loads(user_payload)
        if sent["type"] != "roadmap":
            return super().complete(model=model, system=system, user_payload=user_payload, max_tokens=max_tokens,
                                    temperature=temperature, json_schema=json_schema)
        self._replies = [json.dumps(column_dates(sent))]
        self.calls_before = self.calls
        return t.FakeAdapter.complete(self, model=model, system=system, user_payload=user_payload,
                                      max_tokens=max_tokens, temperature=temperature, json_schema=json_schema)


def _moz_p2_roadmap_rows(monkeypatch):
    """The approval rows of moz p2's roadmap, read with column_dates."""
    client, _, _ = _deck_api(monkeypatch)
    adapter = ColumnDatesAdapter()
    monkeypatch.setattr(gateway, "AnthropicAdapter", lambda *a, **k: adapter)
    _map_revenue(client)
    _upload_deck(client, "02-moz.pdf")
    return [c for c in _deck(client)["candidates"] if c.get("origin") == "ai"
            and c["sources"][0]["structure"] == "roadmap" and c["sources"][0].get("page") == 2]


def test_moz_p2_gives_one_milestone_row_per_dated_paragraph(monkeypatch):
    """Issue #49: read as the live reading read it, every line paired with the date in its column, a row per wrapped
    line gave 40 milestone rows for the page's 9 dated paragraphs. Each paragraph is one row, with its date."""
    rows = [c for c in _moz_p2_roadmap_rows(monkeypatch) if c["value"] is None]
    assert sorted((c["snippet"], c["period_text"]) for c in rows) == sorted(MOZ_P2_MILESTONES)


def test_a_moz_p2_roadmap_row_takes_no_label_from_another_box(monkeypatch):
    """Issue #49: "Label from" was the cell left of the row in its grid row, a line of the neighbouring paragraph
    ("just “SEO” to social media," showed "Label from: index and link graph,")."""
    roadmap = _deck_structure("02-moz.pdf", 2, "roadmap")
    box_of = {f"r{c['row']}c{c['col']}": c["box"] for c in roadmap["cells"]}
    own = {}
    for c in roadmap["cells"]:
        own.setdefault(c["box"], set()).add(c["text"])
    rows = _moz_p2_roadmap_rows(monkeypatch)
    assert rows and [(c["snippet"], c["label_from"]) for c in rows
                     if c["label_from"] and c["label_from"] not in own[box_of[c["cell"]]]] == []


def test_a_roadmap_row_takes_no_label_from_another_box_while_a_kpi_value_keeps_the_label_beside_it():
    """Issue #49: in a roadmap the cell left of a row in its grid row is another text box. In a KPI panel the label
    beside the value is its label, and it stays."""
    roadmap = {"type": "roadmap", "header_rows": 0, "cells": [
        {"row": 1, "col": 1, "text": "Launch the API", "box": 1}, {"row": 1, "col": 2, "text": "Hire 5 engineers", "box": 2},
        {"row": 2, "col": 1, "text": "Q3 2025", "box": 1}, {"row": 2, "col": 2, "text": "Q4 2025", "box": 2}]}
    rows = _rows(roadmap, [_label("i1", "people", unit="count")], pairs=[{"line": "t2", "date": "d2", "category": "hiring"}])
    assert [(r["cell"], r["value"], r["label_from"]) for r in rows] == [("r1c2", 5, None), ("r1c2", None, None)], \
        "the figure and the milestone of the line"
    kpi = {"type": "kpi_panel", "header_rows": 0, "cells": [
        {"row": 1, "col": 1, "text": "Gross churn", "box": 1, "next_to": ["r1c2"]},
        {"row": 1, "col": 2, "text": "4.5%", "box": 2, "next_to": ["r1c1"]}]}
    row, = _rows(kpi, [_label("i1", "retention", unit="%")])
    assert row["label_from"] == "Gross churn"


@pytest.mark.parametrize("value, high, known, found", [
    (12000000, None, [(12000000, None)], "exact"),
    (12000000, 13000000, [(12000000, 13000000)], "exact"),            # the same range
    (12500000, None, [(12000000, 13000000)], "in range"),             # a value inside a range the cell lists
    (12000000, None, [(12000000, 13000000)], "in range"),             # an end is inside its range
    (12000000, 13000000, [(12500000, None)], "in range"),             # a range holding a value the cell lists
    (12000000, 13000000, [(12000000, 14000000)], "in range"),         # a range inside the range the cell lists
    (12000000, 15000000, [(12000000, 14000000)], None),               # a range reaching past it
    (14000000, None, [(12000000, 13000000)], None),
    (12000000, None, [(13000000, None)], None),
    (None, None, [(None, None)], "exact"),                            # a milestone kept from an earlier upload
    (None, None, [(12000000, 13000000)], None),
])
def test_the_dedupe_matcher_returns_exact_in_range_or_no_match(value, high, known, found):
    """Decision of 2026-10-06: an AI row is not added when its cell already lists the same value ("exact"), or when one
    falls inside the other's range ("in range")."""
    assert structures.value_match(value, high, known) == found


def test_an_ai_row_inside_a_range_python_lists_in_the_same_cell_is_not_added_twice():
    from app import structures
    from app.decks import CANDIDATES_COLLECTION, TEXT_COLLECTION
    cells = [{"row": 1, "col": 2, "text": "Plan"}, {"row": 2, "col": 1, "text": "Revenue"},
             {"row": 2, "col": 2, "text": "$12 -$13 million"}, {"row": 3, "col": 1, "text": "Costs"},
             {"row": 3, "col": 2, "text": "$12.5 million"}, {"row": 4, "col": 1, "text": "EBITDA"},
             {"row": 4, "col": 2, "text": "$20 million"}]

    def python_row(row, value, high):
        return {"id": f"p{row}", "audit_id": AUDIT, "deck_id": "d1", "file": "plan.pdf", "status": "pending",
                "claim_type": "revenue", "value": value, "value_high": high, "target_date": None,
                "sources": [{"file": "plan.pdf", "page": 3, "kind": "table", "table": 1, "row": row, "col": 2}]}
    db = _db()
    db["datasets"].docs.append({"audit_id": AUDIT, "dtype": "revenue", "mapping": REVENUE_MAPPING,
                                "mapped_at": "2026-10-05T00:00:00"})
    db[TEXT_COLLECTION].docs.append({"audit_id": AUDIT, "deck_id": "d1", "file": "plan.pdf", "page_unit": "page",
                                     "structures": [{"type": "table", "page": 3, "table": 1, "header_rows": 1,
                                                     "cells": cells}]})
    db[CANDIDATES_COLLECTION].docs += [python_row(2, 12000000, 13000000), python_row(3, 12000000, 13000000),
                                       python_row(4, 19000000, None)]
    labels = [_label("i1", unit="USD"), _label("i2", unit="USD"), _label("i3", "costs", unit="USD"),
              _label("i4", "ebitda", unit="USD")]
    adapter = t.FakeAdapter(replies=[json.dumps({"type": "table", "labels": labels, "pairs": []})])
    asyncio.run(structures.process_deck(db, AUDIT, "d1", adapter=adapter, sleep=t._noop_sleep))
    ai = [(c["sources"][0]["row"], c["value"], c["value_high"]) for c in db[CANDIDATES_COLLECTION].docs
          if c.get("origin") == "ai"]
    assert ai == [(4, 20000000, None)], "row 2: the same range (exact); row 3: inside Python's range; row 4: added"


def test_unticked_consent_is_the_python_only_path(monkeypatch):
    client, db, adapter = _deck_api(monkeypatch, consent=False)
    _map_revenue(client)
    _upload_deck(client)
    listed = _deck(client)
    assert adapter.calls == 0 and listed["decks"][0]["ai_status"] == "python_only"
    assert listed["candidates"] and not [c for c in listed["candidates"] if c.get("origin") == "ai"]


def test_a_deck_already_read_is_not_sent_again_when_the_crm_file_is_mapped_later(monkeypatch):
    client, db, adapter = _deck_api(monkeypatch)
    _map_revenue(client)
    _upload_deck(client)
    before = adapter.calls
    crm = "Deal ID,Account,Created,Close Date,Stage,Amount\nD1,Zeta,2025-01-01,2025-02-01,won,10\n"
    client.post(f"/api/audits/{AUDIT}/datasets/crm/upload", files={"file": ("crm.csv", crm)})
    client.put(f"/api/audits/{AUDIT}/datasets/crm/mapping", json={"mapping": {"deal_id": "Deal ID"}})
    assert adapter.calls == before + 1, "only the CRM file's own column-mapping call"


def test_uploading_the_same_deck_again_is_served_from_the_cache_and_lists_no_row_twice(monkeypatch):
    client, db, adapter = _deck_api(monkeypatch)
    _map_revenue(client)
    _upload_deck(client)
    first = [c for c in _deck(client)["candidates"] if c.get("origin") == "ai"]
    approved = first[0]
    client.put(f"/api/audits/{AUDIT}/decks/candidates/{approved['id']}", json={"status": "approved"})
    before = adapter.calls
    _upload_deck(client)
    again = [c for c in _deck(client)["candidates"] if c.get("origin") == "ai"]
    assert adapter.calls == before, "every structure is a cache hit"
    assert len(again) == len(first), "the approved row is kept and not added again"
    assert any(c["id"] == approved["id"] and c["status"] == "approved" for c in again)


def test_the_token_cap_stops_the_remaining_structures_and_the_deck_says_so(monkeypatch):
    client, db, adapter = _deck_api(monkeypatch)
    _map_revenue(client)
    db["llm_calls"].docs.append({"run_id": AUDIT, "step": "structures", "cache_hit": False, "input_tokens": 394000,
                                 "output_tokens": 0, "estimated_cost_usd": 0.0, "timestamp": "2026-10-05T00:00:00"})
    before = adapter.calls
    _upload_deck(client)
    deck = _deck(client)["decks"][0]
    assert adapter.calls == before and deck["ai_status"] == "stopped"
    assert deck["ai_message"] == ("AI reading stopped: this audit reached its 400,000-token limit. The remaining "
                                  "structures were read by Python only.")
    stored = db["deck_text"].docs[0]["structures"]
    assert {s["ai"]["status"] for s in stored} == {"stopped"}


def test_a_new_fiscal_year_end_re_verifies_the_model_readings(monkeypatch):
    client, db, adapter = _deck_api(monkeypatch)
    _map_revenue(client)
    _upload_deck(client)
    panel = [c for c in db["deck_candidates"].docs if c.get("origin") == "ai" and c["claim_type"] == "gross_profit"
             and c["sources"][0]["structure"] == "kpi_panel"]
    assert panel[0]["ai_label"] == "Verified"
    client.put(f"/api/audits/{AUDIT}/decks/candidates/{panel[0]['id']}", json={"status": "pending"})
    # The recorded reading cites "23 Y/E" as FY2023: it stays verified under any year-end, its range moves.
    client.put(f"/api/audits/{AUDIT}", json={"fiscal_year_end": 3})
    after = next(c for c in db["deck_candidates"].docs if c["id"] == panel[0]["id"])
    assert after["ai_label"] == "Verified" and (after["period_start"], after["period_end"]) == ("2022-04-01", "2023-03-31")


def _processed(cells, labels, kind="table", header_rows=1, year_end=12):
    """A deck holding one structure, read through process_deck with these labels (the revenue file mapped)."""
    from app import structures
    from app.decks import TEXT_COLLECTION
    db = _db(fiscal_year_end=year_end)
    db["datasets"].docs.append({"audit_id": AUDIT, "dtype": "revenue", "mapping": REVENUE_MAPPING,
                                "mapped_at": "2026-10-05T00:00:00"})
    db[TEXT_COLLECTION].docs.append({"audit_id": AUDIT, "deck_id": "d1", "file": "plan.pptx", "page_unit": "slide",
                                     "structures": [{"type": kind, "slide": 3, "header_rows": header_rows,
                                                     "cells": cells}]})
    adapter = t.FakeAdapter(replies=[json.dumps({"type": kind, "labels": labels, "pairs": []})])
    status = asyncio.run(structures.process_deck(db, AUDIT, "d1", adapter=adapter, sleep=t._noop_sleep))
    assert status == structures.READ
    return db


def test_llm_structures_stores_the_reply_and_the_item_list_without_raw_text():
    """Spec section 5, CLAUDE.md rule 17: the reply and the item list (id, cell, position, values, header ids), so every
    label resolves to a value with its cell reference; no deck text, the raw text of an item included."""
    cells = [{"row": 1, "col": 1, "text": "Channel"}, {"row": 1, "col": 2, "text": "Members"},
             {"row": 2, "col": 1, "text": "Social"}, {"row": 2, "col": 2, "text": "Telegram(30K) Discord(150)"},
             {"row": 3, "col": 1, "text": "Hours"}, {"row": 3, "col": 2, "text": "Approx. 2.500 hours"}]
    db = _processed(cells, [_label("i1", "users", unit="count"), _label("i2", "users", unit="count"),
                            _label("i3", "product")])
    stored, = db[gateway.STRUCTURES_COLLECTION].docs
    assert stored["output"] == {"type": "table", "labels": [_label("i1", "users", unit="count"),
                                                            _label("i2", "users", unit="count"), _label("i3", "product")],
                                "pairs": []}, "the reply as it came"
    assert stored["items"] == {"dates": [], "lines": [], "items": [
        {"id": "i1", "cell": "r2c2", "position": 1, "headers": ["r2c1", "r1c2"],
         "values": [{"value": 30000, "dot_reading": None, "bracket_reading": "positive"},
                    {"value": -30000, "dot_reading": None, "bracket_reading": "negative"}]},
        {"id": "i2", "cell": "r2c2", "position": 2, "headers": ["r2c1", "r1c2"],
         "values": [{"value": 150, "dot_reading": None, "bracket_reading": "positive"},
                    {"value": -150, "dot_reading": None, "bracket_reading": "negative"}]},
        {"id": "i3", "cell": "r3c2", "position": 1, "headers": ["r3c1", "r1c2"],
         "values": [{"value": 2500, "dot_reading": "thousands", "bracket_reading": None},
                    {"value": 2.5, "dot_reading": "decimal", "bracket_reading": None}]}]}
    items = {i["id"]: i for i in stored["items"]["items"]}
    assert all(items[label["item"]]["cell"] for label in stored["output"]["labels"]), "every label resolves to a cell"
    flat = json.dumps(db[gateway.STRUCTURES_COLLECTION].docs + db["llm_calls"].docs, ensure_ascii=False, default=str)
    for needle in ("Telegram", "(30K)", "30K", "Discord", "(150)", "Approx", "2.500", "Members", "Social", "raw"):
        assert needle not in flat, f"{needle!r} was stored"


def test_a_new_fiscal_year_end_re_verifies_a_reading_from_its_stored_labels_and_item_list():
    """"Year 1" counted from "Start: Jan 2025" is January to December 2025: a reading of it as the year 2025 holds
    with a December year-end only. Re-verification reads the stored reply and item list, never the deck text."""
    from app import structures
    from app.decks import CANDIDATES_COLLECTION
    cells = [{"row": 1, "col": 1, "text": "Start: Jan 2025"}, {"row": 1, "col": 2, "text": "Year 1"},
             {"row": 2, "col": 1, "text": "Revenue"}, {"row": 2, "col": 2, "text": "£2M"}]
    db = _processed(cells, [_label("i1", period="2025", unit="GBP")])
    row, = db[CANDIDATES_COLLECTION].docs
    assert (row["ai_label"], row["item"]) == ("Verified", "i1")
    labels = []
    for year_end in (3, 12):
        asyncio.run(structures.reverify_audit(db, AUDIT, year_end))
        labels.append(db[CANDIDATES_COLLECTION].docs[0]["ai_label"])
    assert labels == ["AI suggestion, not verified", "Verified"]


def test_re_verification_labels_a_range_row_from_both_its_ends():
    from app import structures
    from app.decks import CANDIDATES_COLLECTION
    cells = [{"row": 1, "col": 2, "text": "Plan"}, {"row": 2, "col": 1, "text": "Revenue"},
             {"row": 2, "col": 2, "text": "$12 -$13 million"}]
    db = _processed(cells, [_label("i1", unit="USD"), _label("i2", "other", unit="USD")])
    row, = db[CANDIDATES_COLLECTION].docs
    assert (row["value"], row["value_high"], row["ai_label"]) == (12000000, 13000000, "AI suggestion, not verified")
    asyncio.run(structures.reverify_audit(db, AUDIT, 12))
    assert db[CANDIDATES_COLLECTION].docs[0]["ai_label"] == "AI suggestion, not verified", "the high end is other"


def test_re_verification_skips_a_reading_stored_under_v2_and_its_row_keeps_its_label():
    """Spec section 5: reverify_audit skips readings stored under v2 (they hold no labels or item list)."""
    from app import structures
    from app.decks import CANDIDATES_COLLECTION, TEXT_COLLECTION
    structure = {"type": "table", "header_rows": 1, "ai": {"key": "k1", "periods_corrected": 1}, "cells": [
        {"row": 1, "col": 2, "text": "FY2025"}, {"row": 2, "col": 1, "text": "Revenue"}, {"row": 2, "col": 2, "text": "£2M"}]}
    v2 = {"metric": "revenue", "period": "2024", "value": 2000000, "unit": "GBP", "actual_or_forecast": "forecast",
          "unit_other": None, "value_cell": "r2c2", "period_cells": ["r1c2"], "proposed_flags": []}
    db = _db()
    db[TEXT_COLLECTION].docs.append({"audit_id": AUDIT, "deck_id": "d1", "structures": [structure],
                                     "periods_corrected": 1})
    db[gateway.STRUCTURES_COLLECTION].docs.append({"audit_id": AUDIT, "key": "k1", "type": "table", "prompt_version": "v2",
                                                   "output": {"type": "table", "items": [v2]}})
    db[CANDIDATES_COLLECTION].docs.append({"id": "c1", "audit_id": AUDIT, "deck_id": "d1", "structure_key": "k1",
                                           "status": "pending", "cell": "r2c2", "claim_type": "revenue",
                                           "value": 2000000, "ai_status": "verified", "ai_label": "Verified"})
    assert asyncio.run(structures.reverify_audit(db, AUDIT, 3)) == 0
    assert db[CANDIDATES_COLLECTION].docs[0]["ai_label"] == "Verified"
    assert db[TEXT_COLLECTION].docs[0]["periods_corrected"] == 1


def test_delete_audit_removes_model_outputs_cache_and_mapping(monkeypatch):
    client, db, adapter = _deck_api(monkeypatch)
    _map_revenue(client)
    _upload_deck(client)
    assert db[gateway.STRUCTURES_COLLECTION].docs and db["pseudonym_map"].docs and db["column_mappings"].docs
    assert all(doc.get("items") for doc in db[gateway.STRUCTURES_COLLECTION].docs if "labels" in doc["output"])
    purged = client.delete(f"/api/audits/{AUDIT}").json()
    assert purged["llm_purged"]["llm_structures"] > 0, "the replies and their item lists go with the audit"
    for name in (gateway.STRUCTURES_COLLECTION, "pseudonym_map", "column_mappings", "llm_calls", "deck_candidates",
                 "deck_text"):
        assert db[name].docs == [], name


def _consistency_script():
    import importlib.util
    spec = importlib.util.spec_from_file_location("consistency_run", BACKEND.parent / "scripts" / "consistency_run.py")
    script = importlib.util.module_from_spec(spec)
    spec.loader.exec_module(script)
    return script


def test_live_run_reports_and_diagnostics_are_kept_by_git_in_docs_test_runs():
    """docs/test-runs/README.md (decision of 2026-10-06): a live run's report and diagnostic are committed, so their
    numbers are not lost; nothing may ignore them."""
    import shutil
    import subprocess
    script = _consistency_script()
    assert script.REPORTS == BACKEND.parent / "docs" / "test-runs" and (script.REPORTS / "README.md").is_file()
    if not shutil.which("git") or not (BACKEND.parent / ".git").exists():
        pytest.skip("not a git checkout")
    for name in ("consistency_2026-10-06.md", "consistency_2026-10-06-2_diagnostic.md"):
        found = subprocess.run(["git", "check-ignore", "-q", f"docs/test-runs/{name}"], cwd=BACKEND.parent)
        assert found.returncode == 1, f"{name} is ignored by git"


def test_the_consistency_script_reports_agreement_match_rate_cost_and_cache_hits_without_a_live_call(
        monkeypatch, tmp_path):
    """scripts/consistency_run.py is manual (live API, costs money); --fake checks its arithmetic. Its FakeAdapter
    replays the recorded labelling replies and labels every other listed item not_a_metric."""
    pytest.importorskip("pdfplumber")
    import tempfile
    script = _consistency_script()
    monkeypatch.setattr(script, "REPORTS", tmp_path / "docs" / "test-runs")
    monkeypatch.setattr(tempfile, "tempdir", str(tmp_path / "temp"))
    (tmp_path / "temp").mkdir()
    report = script.main(["--fake", "--deck", "05-zero2hero.pdf"])
    assert [p.name.split("_")[0] for p in (tmp_path / "temp").iterdir()] == ["consistency"]
    assert not script.REPORTS.exists(), "a fake run's report never lands in docs/test-runs"
    assert report["agreement_pct"] == {"hiring_table": 100.0, "kpi_panel": 100.0, "table": 100.0}
    assert (report["agreement_pct_all"], report["agreement_pct_all_old"]) == (100.0, 100.0)
    assert report["verifier_match_rate_pct"] == 100.0, "the table's 24 and the panel's 2 per pass"
    assert report["counts"]["table"] == {"not_a_metric": 0, "other": 0, "ambiguous": 0, "flags": 0}
    assert report["counts"]["hiring_table"]["not_a_metric"] == 9, "3 items, 3 passes"
    assert report["counts"]["kpi_panel"] == {"not_a_metric": 27, "other": 0, "ambiguous": 18, "flags": 0}, \
        "pages 11 and 17 (9 items: the 2 in page 11's label box \"zero2hero\" is no figure), 3 passes; page 17's " \
        "six bracketed figures after text read two ways"
    assert report["cache_hit_rate_pct"] == {"pass_2": 100.0, "pass_3": 100.0}
    assert report["per_deck"]["05-zero2hero.pdf"]["structures"] == 5
    assert report["model_reads"] == 5 * 3, "passes 2 and 3 bypass the cache after counting its hits"
    assert report["period_corrected"] == 0

    class WrongYear(script.FakeAdapter):
        """Replays the recorded replies with one revenue period a year off: Python rebuilds it from its header."""
        def complete(self, **kwargs):
            reply, tokens_in, tokens_out = super().complete(**kwargs)
            body = json.loads(reply)
            for label in body.get("labels", []):
                if (label["metric"], label["period"]) == ("revenue", "FY2023"):
                    label["period"] = "FY2021"
            return json.dumps(body), tokens_in, tokens_out
    report = asyncio.run(script.run(["05-zero2hero.pdf"], 3, script.MemoryDB(), WrongYear()))
    assert report["period_corrected"] == 3, "one per pass"
    assert report["verifier_match_rate_pct"] == 100.0, "a corrected item is verified"
    assert report["agreement_pct_all"] == 100.0, "the period is compared as the verifier keeps it"
    os.environ.pop("ANTHROPIC_API_KEY", None)
    with pytest.raises(SystemExit, match="ANTHROPIC_API_KEY"):
        script.main(["--yes", "--deck", "05-zero2hero.pdf"])


def _drifting(script):
    """The recorded replies, read a little differently from pass to pass (05-zero2hero.pdf): pass 2 writes the FY2023
    revenue's unit as USD, pass 3 writes its period as "2023" and labels the page 19 panel's users as customers. In
    every pass the FY2022 gross profit is labelled other and page 17's Telegram members are read as users in 2023,
    a period no cell of theirs holds: both stay unverified."""
    class Drifting(script.FakeAdapter):
        def __init__(self):
            super().__init__()
            self.calls = {}

        def complete(self, **kwargs):
            reply, tokens_in, tokens_out = super().complete(**kwargs)
            n = self.calls[kwargs["user_payload"]] = self.calls.get(kwargs["user_payload"], 0) + 1   # the pass
            text, body = json.loads(kwargs["user_payload"])["text"], json.loads(reply)
            for label in body.get("labels", []):
                if "r4c1: Revenue" in text and label["item"] == "i11":
                    label.update({2: {"unit": "USD"}, 3: {"period": "2023"}}.get(n, {}))
                if "r4c1: Revenue" in text and label["item"] == "i20":
                    label["metric"] = "other"
                if "r2c1: Gross Profit £150K" in text and label["item"] == "i2" and n == 3:
                    label["metric"] = "customers"
                if "Telegram(30K)" in text and label["item"] == "i1":
                    label.update(metric="users", period="2023", unit="count", actual_or_forecast="actual")
            return json.dumps(body), tokens_in, tokens_out
    return Drifting


def test_the_consistency_report_shows_where_passes_disagree_why_items_are_unverified_and_the_fixed_prompt(
        monkeypatch, tmp_path):
    """Spec sections 6 and 7: agreement on metric and the period the verifier keeps over the items Python listed,
    with the old method (period as written, unit, actual or forecast) beside it; the fields that differ per
    disagreeing structure; a reason per unverified item; the counts per type; the fixed prompt's tokens and the
    average text plus item list against 4,000."""
    pytest.importorskip("pdfplumber")
    from app.decks import parser
    from app.llm import cache, prompt_store, schemas
    script = _consistency_script()
    report = asyncio.run(script.run(["05-zero2hero.pdf"], 3, script.MemoryDB(), _drifting(script)()))
    # Table: 24 items, the FY2023 revenue written three ways but kept as one period: 23 of 24 as written, 24 of 24
    # normalised. KPI panels: 3 + 6 + 2 items (the 2 in page 11's label box "zero2hero" is no figure), the page 19
    # users read as customers in pass 3: 10 of 11 either way.
    assert report["agreement_pct_old"] == {"hiring_table": 100.0, "kpi_panel": 90.9, "table": 95.8}
    assert report["agreement_pct"] == {"hiring_table": 100.0, "kpi_panel": 90.9, "table": 100.0}
    assert (report["agreement_pct_all_old"], report["agreement_pct_all"]) == (94.7, 97.4), "36 of 38; 37 of 38"
    assert report["disagreements"] == [
        {"deck": "05-zero2hero.pdf", "page": 19, "type": "table", "fields": ["period", "unit"],
         "same_after_normalisation": True},
        {"deck": "05-zero2hero.pdf", "page": 19, "type": "kpi_panel", "fields": ["metric"],
         "same_after_normalisation": False}]
    zero = dict.fromkeys(script.FIELDS, 0)
    assert script.FIELDS == ("metric", "period", "unit", "actual_or_forecast")
    assert report["disagreement_fields"] == {
        "hiring_table": {"structures": 1, "disagreeing": 0, **zero},
        "kpi_panel": {"structures": 3, "disagreeing": 1, **zero, "metric": 1},
        "table": {"structures": 1, "disagreeing": 1, **zero, "period": 1, "unit": 1}}
    assert script.REASONS == ("period not rebuilt", "metric invalid", "other")
    assert report["unverified_reasons"] == {"kpi_panel": {**dict.fromkeys(script.REASONS, 0), "period not rebuilt": 3},
                                            "table": {**dict.fromkeys(script.REASONS, 0), "other": 3}}
    assert report["unverified_items"] == [
        {"deck": "05-zero2hero.pdf", "page": 17, "type": "kpi_panel", "item": "i1", "cell": "r1c1#1",
         "reason": "period not rebuilt", "detail": None, "passes": 3},
        {"deck": "05-zero2hero.pdf", "page": 19, "type": "table", "item": "i20", "cell": "r6c2#1",
         "reason": "other", "detail": "type Other", "passes": 3}]
    assert report["verifier_match_rate_pct"] == 92.6, "25 of 27 per pass: 24 + 2 + Telegram, less two"
    assert report["counts"]["table"] == {"not_a_metric": 0, "other": 3, "ambiguous": 0, "flags": 0}

    # Tokens: the fake counter is one token per 2 characters of what would be sent, as measured on the live runs of
    # 2026-10-06, and a read bills what it sends at that rate.
    def count(system=None, payload="", schema=None):
        return (len(system or "") + len(payload) + (len(json.dumps(schema)) if schema else 0)) // 2
    assert script.FakeAdapter.CHARS_PER_TOKEN == 2
    system, schema = prompt_store.load(gateway.STRUCTURE_PROMPT).text, schemas.labelling_output_schema()
    empty = cache.canonical_json({"type": "table", "text": ""})
    sent = []
    for structure in parser.parse_deck((DECKS / "05-zero2hero.pdf").read_bytes(), "05-zero2hero.pdf")["structures"]:
        cells = redact.redact_structure(structure["cells"], "05-zero2hero", {}, ("Consistency run", "CONSISTENCY"))[0]
        redacted = {**structure, "cells": cells}
        sent.append((structure["type"], structure_items.text(redacted, structure_items.list_items(redacted))))
    billed = [count(system, cache.canonical_json({"type": kind, "text": text}), schema) for kind, text in sent]
    assert report["tokens"] == {
        "fixed_prompt": count(system, empty, schema), "system_prompt": count(system, empty) - count(payload=empty),
        "output_schema": count(payload=empty, schema=schema) - count(payload=empty),
        "empty_message": count(payload=empty),
        "text_and_items_avg": round(sum(count(payload=text) for _, text in sent) / 5, 1), "structures_counted": 5,
        "billed_input_per_model_read": round(sum(billed) / 5, 1)}, "each structure read once per pass"

    path = script.write_report(report, tmp_path, 3, fake=True)
    text = path.read_text(encoding="utf-8")
    tokens = report["tokens"]
    for line in (
            "| table | 100.0% | 95.8% |", "| kpi_panel | 90.9% | 90.9% |", "| all | 97.4% | 94.7% |",
            "| Type | Structures | Disagreeing | metric | period | unit | actual_or_forecast |",
            "| kpi_panel | 3 | 1 | 1 | 0 | 0 | 0 |", "| table | 1 | 1 | 0 | 1 | 1 | 0 |",
            "| 05-zero2hero.pdf | 19 | table | period, unit | yes |",
            "| 05-zero2hero.pdf | 19 | kpi_panel | metric | no |",
            "| Type | period not rebuilt | metric invalid | other | Unverified |",
            "| kpi_panel | 3 | 0 | 0 | 3 |", "| table | 0 | 0 | 3 | 3 |", "| all | 3 | 0 | 3 | 6 |",
            "| 05-zero2hero.pdf | 17 | kpi_panel | i1 | r1c1#1 | period not rebuilt | 3 |",
            "| 05-zero2hero.pdf | 19 | table | i20 | r6c2#1 | other (type Other) | 3 |",
            "| Type | not_a_metric | other | ambiguous readings | flags |", "| table | 0 | 3 | 0 | 0 |",
            f"- Fixed prompt: {tokens['fixed_prompt']:,} (system prompt {tokens['system_prompt']:,}, output schema "
            f"{tokens['output_schema']:,}, empty message {tokens['empty_message']:,})",
            f"- Structure text and item list, average of 5 structures: {tokens['text_and_items_avg']:,} against the "
            "gateway's 4,000-token cap",
            f"- Billed input per model read: {tokens['billed_input_per_model_read']:,}"):
        assert line in text, line
    assert script.summary(report).startswith("Agreement 97.4% (target 95.0%: met; old method 94.7%); verified 92.6%")

    class Uncountable:
        def count_tokens(self, **kwargs):
            raise RuntimeError("count_tokens unavailable")
    failed = asyncio.run(script.fixed_tokens(Uncountable()))
    assert failed == dict.fromkeys(("fixed_prompt", "system_prompt", "output_schema", "empty_message")), \
        "a failed count is n/a and never stops a paid run"
    report["tokens"].update(failed)
    assert "- Fixed prompt: n/a (system prompt n/a, output schema n/a, empty message n/a)" in \
        script.write_report(report, tmp_path, 3, fake=True).read_text(encoding="utf-8")


def test_roadmap_lines_are_reported_apart_with_their_pairs_and_dates_rebuilt_from_cells(tmp_path):
    """Spec section 6: "roadmap lines: N, same pair and category in every pass: M, date rebuilt from cell: K". K counts
    the lines whose period Python rebuilds from their own period cells, whatever the pair says."""
    pytest.importorskip("pdfplumber")
    pytest.importorskip("pptx")
    script = _consistency_script()

    class Tea(script.FakeAdapter):
        """The TEA roadmap (page 11): two milestones paired, one of them with another date in pass 3."""
        def __init__(self):
            super().__init__()
            self.calls = {}

        def complete(self, **kwargs):
            reply, tokens_in, tokens_out = super().complete(**kwargs)
            sent = json.loads(kwargs["user_payload"])
            if sent["type"] != "roadmap" or "r8c2: Rich dApps running on network" not in sent["text"]:
                return reply, tokens_in, tokens_out
            n = self.calls[sent["text"]] = self.calls.get(sent["text"], 0) + 1
            body = json.loads(reply)
            body["pairs"] = [{"line": "t4", "date": "d1", "category": "launch"},
                             {"line": "t16", "date": "d4" if n == 3 else "d3", "category": "launch"}]
            return json.dumps(body), tokens_in, tokens_out
    report = asyncio.run(script.run(["03-buffer.pptx", "05-zero2hero.pdf", "10-tea.pdf"], 3, script.MemoryDB(), Tea()))
    assert (report["roadmap_lines"], report["roadmap_same_pair"], report["roadmap_dates_rebuilt"]) == (36, 35, 27), \
        "Buffer: 6 lines, all paired alike, all dated from the date line below them; TEA: 30 lines, one paired " \
        "with another date in pass 3, 21 with a period in their own cells or headers"
    # The match rate keeps roadmap figures (milestones left out) and is also given apart: financial (outside
    # roadmaps: zero2hero's 26 a pass) and roadmap (Buffer's 7 pair-dated figures a pass, each paired with the date
    # line its own timeline gives it).
    assert (report["verifier_match_rate_pct"], report["verifier_match_rate_financial_pct"],
            report["verifier_match_rate_roadmap_pct"]) == (100.0, 100.0, 100.0), "99 of 99; 78 of 78; 21 of 21"
    assert "verified 100.0% (financial 100.0%, roadmap 100.0%)" in script.summary(report)
    assert report["agreement_pct"]["roadmap"] == 100.0, "agreement counts a roadmap's items like any other"
    assert "roadmap lines: 36, same pair and category in every pass: 35, date rebuilt from cell: 27;" in \
        script.summary(report)
    text = script.write_report(report, tmp_path, 3, fake=True).read_text(encoding="utf-8")
    assert "- Roadmap lines: 36, same pair and category in every pass: 35, date rebuilt from cell: 27" in text
    for line in ("- Match rate: 100.0%", "- Match rate, financial (outside roadmaps): 100.0%",
                 "- Match rate, roadmap (figures in roadmaps, milestones left out): 100.0%"):
        assert line in text.splitlines(), line


@pytest.mark.parametrize("label, reason", [
    (dict(period="2025"), ("period not rebuilt", None)),         # a period no cell of the item holds
    (dict(metric="other"), ("other", "type Other")),
    (dict(period=None), None),                                   # verified
])
def test_each_unverified_item_gets_its_reason(label, reason):
    import test_structure_verifier as tv
    script = _consistency_script()
    structure = v._struct([["", "Plan"], ["Revenue", "£2M"]])
    checked, = tv._verify(structure, {"r2c2": tv._lab(**label)})["items"]
    assert (None if checked["status"] == verify.VERIFIED else script.unverified_reason(structure, checked)) == reason
    assert script.unverified_reason(structure, {**checked, "metric": "amount"}) == ("metric invalid", None)


@pytest.mark.parametrize("change, fields", [
    ({}, []),
    ({"metric": "sales"}, ["metric"]),
    ({"period": "2023"}, ["period"]),                            # as written: FY2023 and 2023 differ here
    ({"unit": "USD"}, ["unit"]),
    ({"actual_or_forecast": "actual"}, ["actual_or_forecast"]),
    ({"metric": "sales", "unit": None}, ["metric", "unit"]),
])
def test_the_fields_that_differ_are_found_by_lining_labels_up_by_item_id(change, fields):
    script = _consistency_script()
    kept, item = _label("i1", period="FY2022", unit="GBP"), _label("i2", period="FY2023", unit="GBP")
    assert script.differing_fields([[kept, item], [kept, {**item, **change}], [item, kept]]) == fields


def test_agreement_compares_the_metric_and_the_period_the_verifier_keeps_by_its_dates():
    """Spec section 6: the period as the verifier keeps it (start and end dates), not as the model wrote it; a
    not_a_metric label counts as a metric, with no period kept."""
    script = _consistency_script()
    written = _label("i1", period="FY2021")
    assert script._normalised_key(written, {"period": "2023"}) == ("i1", "revenue", ("2023-01-01", "2023-12-31")), \
        "the verifier rebuilt 2023 from the header: the model's FY2021 is not compared"
    assert script._normalised_key(written, {"period": "2023"}) == script._normalised_key(written, {"period": "FY2023"})
    assert script._normalised_key(written, {"period": "2023"}) != script._normalised_key(written, {"period": "2024"})
    assert script._normalised_key(written, {"period": None}) == ("i1", "revenue", None)
    dropped = _label("i2", metric="not_a_metric", period="2023")
    assert script._normalised_key(dropped, None) == ("i2", "not_a_metric", None)
    assert script._normalised_key(dropped, None) != script._normalised_key({**dropped, "metric": "revenue"}, None)


def test_two_passes_writing_other_with_different_codes_differ_in_unit():
    script = _consistency_script()
    zar = _label("i1", period="FY2023", unit="other", unit_other="ZAR")
    mxn = {**zar, "unit_other": "MXN"}
    assert script.differing_fields([[zar], [mxn], [zar]]) == ["unit"]
    assert script.differing_fields([[zar], [zar], [zar]]) == []
    assert script._item_key(zar) != script._item_key(mxn), "the old agreement key tells them apart too"
    assert script._item_key(zar) == ("i1", "revenue", "FY2023", "ZAR", "forecast")


def _motor_without_a_server(monkeypatch, names, drop_fails=False):
    """Real motor: it binds a client to the event loop of the client's first call and runs every later call
    on that loop. Only the blocking pymongo call is answered here, so no MongoDB server is needed."""
    import motor.frameworks.asyncio as framework
    from pymongo.errors import AutoReconnect
    real, dropped = framework.run_on_executor, []

    def answer(sync_method, delegate, *args, **kwargs):
        if sync_method.__name__ == "list_database_names":
            return list(names)
        if sync_method.__name__ == "drop_database":
            if drop_fails:
                raise AutoReconnect("connection lost")
            dropped.append(args[0])
            return None
        raise AssertionError(f"unexpected server call {sync_method.__name__}")
    monkeypatch.setattr(framework, "run_on_executor", lambda loop, fn, *a, **k: real(loop, answer, fn, *a, **k))
    return dropped


def test_the_live_run_drops_its_scratch_database_in_the_event_loop_it_ran_in(monkeypatch, tmp_path, capsys):
    """`consistency_run.py --yes` crashed with "RuntimeError: Event loop is closed" at drop_database and wrote
    no report: motor had bound the client to the run's loop, which asyncio.run had closed."""
    pytest.importorskip("pdfplumber")
    import datetime
    script = _consistency_script()
    dropped = _motor_without_a_server(monkeypatch, ["consistency_run", "consistency_run_old", "growth_diligence"])
    run = script.run

    async def live_run(decks, passes, db, adapter=None, **kwargs):
        await db.client.list_database_names()            # a motor call binds the client to this loop, as an insert does
        return await run(decks, passes, script.MemoryDB(), Billed())     # no MongoDB, no live API

    class Billed(script.FakeAdapter):
        def complete(self, **kwargs):
            return super().complete(**kwargs)[0], 1000, 20          # a round bill per read, so the sums show
    monkeypatch.setattr(script, "run", live_run)
    monkeypatch.setattr(script, "REPORTS", tmp_path / "docs" / "test-runs")
    monkeypatch.setenv("ANTHROPIC_API_KEY", "not-used")
    report = script.main(["--yes", "--deck", "05-zero2hero.pdf"])
    assert report["per_deck"]["05-zero2hero.pdf"]["structures"] == 5
    assert dropped == ["consistency_run", "consistency_run_old", "consistency_run"], \
        "leftovers of earlier runs at start, by name pattern; this run's own at the end; never another database"
    day = datetime.date.today().isoformat()
    path = tmp_path / "docs" / "test-runs" / f"consistency_{day}.md"
    out = capsys.readouterr().out.splitlines()
    assert out[:2] == ["Dropped scratch database consistency_run left by an earlier run",
                       "Dropped scratch database consistency_run_old left by an earlier run"]
    assert [line.split(":")[0] for line in out[2:5]] == [f"[1/1] 05-zero2hero.pdf pass {n}/3" for n in (1, 2, 3)]
    assert all("5 of 5 structures read" in line for line in out[2:5]), "one progress line per deck and pass"
    assert out[5:7] == [f"Report: {path}", script.summary(report)], "the report path and a summary line"
    assert out[6].startswith("Agreement 100.0% (target 95.0%: met; old method 100.0%); verified 100.0%")
    text = path.read_text(encoding="utf-8")
    assert "\n".join(out[7:]) + "\n" == text, \
        "then the whole report, as the committed file holds it"
    assert text.startswith(f"# Consistency run {day}\n\nLive API. Decks: 1. Passes: 3. Model: {gateway.STRUCTURE_MODEL}.")
    for line in ("| table | 100.0% |", "| kpi_panel | 100.0% |", "- Match rate: 100.0%", "| 2 | 100.0% |", "| 3 | 100.0% |",
                 "| 05-zero2hero.pdf | 5 | 15,000 | 300 |"):
        assert line in text, line
    script.main(["--yes", "--deck", "05-zero2hero.pdf"])
    assert path.read_text(encoding="utf-8") == text, "a later run the same day never overwrites a report"
    assert (path.parent / f"consistency_{day}-2.md").exists()
    del dropped[:]
    with pytest.raises(SystemExit, match="--db must be consistency_run"):
        script.main(["--yes", "--db", "growth_diligence"])
    assert dropped == [], "the run drops only a database it may clean up at the next start"


def test_a_failed_drop_keeps_the_paid_report(monkeypatch, tmp_path):
    """The report is written before the scratch database is dropped; the next run drops what is left."""
    pytest.importorskip("pdfplumber")
    from pymongo.errors import AutoReconnect
    script = _consistency_script()
    _motor_without_a_server(monkeypatch, [], drop_fails=True)
    run = script.run

    async def live_run(decks, passes, db, adapter=None, **kwargs):
        return await run(decks, passes, script.MemoryDB(), script.FakeAdapter())
    monkeypatch.setattr(script, "run", live_run)
    monkeypatch.setattr(script, "REPORTS", tmp_path)
    monkeypatch.setenv("ANTHROPIC_API_KEY", "not-used")
    with pytest.raises(AutoReconnect):
        script.main(["--yes", "--deck", "05-zero2hero.pdf"])
    assert [p.name.split("_")[0] for p in tmp_path.iterdir()] == ["consistency"]


def test_a_live_run_writes_its_diagnostic_beside_the_report_before_the_drop(monkeypatch, tmp_path):
    pytest.importorskip("pdfplumber")
    import datetime
    from pymongo.errors import AutoReconnect
    script = _consistency_script()
    _motor_without_a_server(monkeypatch, [], drop_fails=True)
    run = script.run

    async def live_run(decks, passes, db, adapter=None, **kwargs):
        return await run(decks, passes, script.MemoryDB(), _drifting(script)(), diagnostic=kwargs["diagnostic"])
    monkeypatch.setattr(script, "run", live_run)
    monkeypatch.setattr(script, "REPORTS", tmp_path)
    monkeypatch.setenv("ANTHROPIC_API_KEY", "not-used")
    with pytest.raises(AutoReconnect):
        script.main(["--yes", "--diagnostic", "--deck", "05-zero2hero.pdf"])
    day = datetime.date.today().isoformat()
    assert sorted(p.name for p in tmp_path.iterdir()) == [f"consistency_{day}.md", f"consistency_{day}_diagnostic.md"]
    assert "£ 150,000" in (tmp_path / f"consistency_{day}_diagnostic.md").read_text(encoding="utf-8")


# ---------------------------------------------------------------------------
# Provider errors: the HTTP status and error type are recorded, 429 and 529 are retried
# ---------------------------------------------------------------------------
CELL_MESSAGE = "cannot read cell 'Jane Doe (CEO)'"


def _sdk_error(status, error_type, retry_after=None):
    """The error the Anthropic SDK raises for this HTTP status. Its message and body quote a cell, as a provider
    message may: neither reaches a log or llm_calls."""
    anthropic = pytest.importorskip("anthropic")
    httpx2 = pytest.importorskip("httpx2")
    headers = {} if retry_after is None else {"retry-after": str(retry_after)}
    response = httpx2.Response(status, headers=headers,
                               request=httpx2.Request("POST", "https://api.anthropic.com/v1/messages"))
    body = {"type": "error", "error": {"type": error_type, "message": CELL_MESSAGE}}
    return anthropic.Anthropic(api_key="test")._make_status_error(CELL_MESSAGE, body=body, response=response)


class _Failing(t.FakeAdapter):
    """Raises `errors` from complete() and `count_errors` from count_tokens() in turn, then replays REPLY."""

    def __init__(self, errors=(), count_errors=()):
        super().__init__(replies=[json.dumps(REPLY)])
        self.errors, self.count_errors = list(errors), list(count_errors)

    def complete(self, **kwargs):
        if self.errors:
            self.calls += 1
            raise self.errors.pop(0)
        return super().complete(**kwargs)

    def count_tokens(self, **kwargs):
        if self.count_errors:
            raise self.count_errors.pop(0)
        return super().count_tokens(**kwargs)


def _read_slept(db, adapter):
    """read_structure with a sleep that records the backoff instead of waiting."""
    slept = []

    async def sleep(seconds):
        slept.append(seconds)
    return asyncio.run(gateway.read_structure(db, AUDIT, TEXT, "table", adapter=adapter, sleep=sleep)), slept


def test_a_provider_error_is_logged_and_recorded_with_its_http_status_and_type(caplog):
    """The live consistency run logged "reason=provider_error" and 0 tokens for every call after the first deck,
    and nothing said why. The not-read line and the llm_calls record carry the HTTP status and the provider's
    error type, for the call and for the token count before it; the message is never read."""
    import logging
    db = _db()
    with caplog.at_level(logging.INFO):
        result, slept = _read_slept(db, adapter := _Failing([_sdk_error(400, "invalid_request_error")]))
    assert (result.status, adapter.calls, slept) == ("not_read", 1, []), "a 400 is not retried"
    call, = db["llm_calls"].docs
    assert (call["status"], call["http_status"], call["error_type"]) == ("provider_error", 400, "invalid_request_error")
    assert "tokens=0/0 cost=0.000000 reason=provider_error status=400 type=invalid_request_error" in caplog.text

    db = _db()
    with caplog.at_level(logging.INFO):
        result, _ = _read_slept(db, adapter := _Failing(count_errors=[_sdk_error(403, "permission_error")]))
    assert (result.status, adapter.calls) == ("not_read", 0)
    call, = db["llm_calls"].docs
    assert (call["status"], call["http_status"], call["error_type"], call["input_tokens"], call["cache_hit"]) == \
        ("provider_error", 403, "permission_error", 0, False), "a failed token count is recorded too"
    assert "reason=provider_error status=403 type=permission_error" in caplog.text
    assert "Jane Doe" not in caplog.text + json.dumps(db["llm_calls"].docs)


def test_429_and_529_are_retried_three_times_honouring_retry_after(caplog):
    """A burst of rate limits or overloads does not fail the structure: 3 attempts, each wait the longer of the
    backoff (1 s, 2 s) and the provider's retry-after, at most RETRY_AFTER_MAX_SECONDS so the lock (300 s) cannot
    expire mid-call. The token count is retried the same way."""
    import logging
    burst = [_sdk_error(429, "rate_limit_error", retry_after=7), _sdk_error(529, "overloaded_error")]
    result, slept = _read_slept(_db(), adapter := _Failing(burst))
    assert (result.status, adapter.calls, slept) == ("read", 3, [7.0, 2.0]), "retry-after 7 s, then the 2 s backoff"

    db = _db()
    with caplog.at_level(logging.INFO):
        result, slept = _read_slept(db, adapter := _Failing([_sdk_error(429, "rate_limit_error", retry_after=600)] * 3))
    assert (result.status, adapter.calls) == ("not_read", 3)
    assert slept == [gateway.RETRY_AFTER_MAX_SECONDS] * 2 == [20.0, 20.0]
    assert 4 * gateway.REQUEST_TIMEOUT_SECONDS + 2 * gateway.RETRY_AFTER_MAX_SECONDS < guards.LOCK_TTL_SECONDS, \
        "the worst call under the lock (4 attempts at the timeout, 2 capped waits) fits in it"
    call, = db["llm_calls"].docs
    assert (call["status"], call["http_status"], call["error_type"]) == ("provider_unreachable", 429, "rate_limit_error")
    assert "reason=provider_unreachable status=429 type=rate_limit_error" in caplog.text

    result, slept = _read_slept(_db(), _Failing(count_errors=[_sdk_error(429, "rate_limit_error", retry_after=5)]))
    assert (result.status, slept) == ("read", [5.0])


def test_the_consistency_run_logs_one_not_read_line_per_structure_and_pass(caplog):
    """On passes 2 and 3 each "structure not read" line printed twice: the lookup that counts cache hits called
    the model when the cache missed, then the call with the cache bypassed did again. The lookup reads the cache
    only, so a failed structure is sent once per pass, and its line names the HTTP status and error type."""
    pytest.importorskip("pdfplumber")
    import logging
    script = _consistency_script()

    class Failing(script.FakeAdapter):
        calls = 0

        def complete(self, **kwargs):
            self.calls += 1
            raise _sdk_error(400, "invalid_request_error")
    adapter = Failing()
    with caplog.at_level(logging.WARNING):
        report = asyncio.run(script.run(["05-zero2hero.pdf"], 3, script.MemoryDB(), adapter))
    lines = [r.getMessage() for r in caplog.records if r.getMessage().startswith("structure not read:")]
    assert (adapter.calls, len(lines), report["not_read"]) == (15, 15, 15), "5 structures, 3 passes, once each"
    assert all(line.endswith("reason=provider_error status=400 type=invalid_request_error") for line in lines)
    assert report["cache_hit_rate_pct"] == {"pass_2": 0.0, "pass_3": 0.0}


def test_one_provider_error_does_not_stop_the_calls_after_it_in_the_deck_or_the_next():
    """The circuit breaker (spec section 7: the per-audit lock, the daily spend cap, the retry policy) does not open
    on an error: after a failed call the next structure and the next deck are still sent. Each deck is read in its
    own audit, so it has its own lock, and the lock is released after every call."""
    pytest.importorskip("pdfplumber")
    script = _consistency_script()

    class OnceFailing(script.FakeAdapter):
        failed = False

        def complete(self, **kwargs):
            if not self.failed:
                self.failed = True
                raise _sdk_error(400, "invalid_request_error")
            return super().complete(**kwargs)
    db = script.MemoryDB()
    report = asyncio.run(script.run(["05-zero2hero.pdf", "10-tea.pdf"], 1, db, OnceFailing()))
    assert (report["not_read"], report["model_reads"]) == (1, 8), "the first of 5 + 3 structures only"
    locks = db[guards.LOCKS_COLLECTION].docs
    assert len(locks) == 2 and not any(lock["held"] for lock in locks), "one lock per deck, none left held"


def test_the_consistency_run_pauses_between_structure_calls(monkeypatch, tmp_path):
    """--pause (default 2 s) waits between structure calls, not after the last; --fake does not wait."""
    pytest.importorskip("pdfplumber")
    script = _consistency_script()
    slept = []

    async def sleep(seconds):
        slept.append(seconds)
    asyncio.run(script.run(["05-zero2hero.pdf"], 3, script.MemoryDB(), script.FakeAdapter(), pause=2.0, sleep=sleep))
    assert slept == [2.0] * 14, "between the 15 structure calls"

    _motor_without_a_server(monkeypatch, [])
    run, pauses = script.run, []

    async def live_run(decks, passes, db, adapter=None, **kwargs):
        pauses.append(kwargs["pause"])
        return await run(decks, passes, script.MemoryDB(), script.FakeAdapter())
    monkeypatch.setattr(script, "run", live_run)
    monkeypatch.setattr(script, "REPORTS", tmp_path)
    monkeypatch.setenv("ANTHROPIC_API_KEY", "not-used")
    script.main(["--yes", "--deck", "05-zero2hero.pdf"])
    script.main(["--yes", "--deck", "05-zero2hero.pdf", "--pause", "0.5"])
    assert pauses == [2.0, 0.5]


# ---------------------------------------------------------------------------
# --diagnostic: the cell text behind each unverified item and disagreeing structure, public test decks only
# ---------------------------------------------------------------------------
def _seen(metric, period, unit, actual_or_forecast, verifier):
    return {"metric": metric, "period": period, "unit": unit, "actual_or_forecast": actual_or_forecast,
            "verifier": verifier}


def test_the_diagnostic_gives_the_cell_text_and_every_pass_of_each_unverified_and_disagreeing_item(tmp_path):
    """Spec section 7: for every unverified item and every item whose labels differ between passes: deck, page, type,
    item id, cell#position, the cell's text as sent, Python's values, the reason, then per pass the model's metric,
    period, unit and actual_or_forecast and the verifier's result, in a file of its own beside the report. The report
    and its JSON keys are unchanged by it."""
    pytest.importorskip("pdfplumber")
    script = _consistency_script()
    rows = []
    report = asyncio.run(script.run(["05-zero2hero.pdf"], 3, script.MemoryDB(), _drifting(script)(), diagnostic=rows))
    assert report == asyncio.run(script.run(["05-zero2hero.pdf"], 3, script.MemoryDB(), _drifting(script)())), \
        "the report is the same with or without the diagnostic"
    panel = {"deck": "05-zero2hero.pdf", "type": "kpi_panel"}
    table = {"deck": "05-zero2hero.pdf", "page": 19, "type": "table"}
    assert [r for r in rows if r["section"] == "unverified"] == [
        {**panel, "page": 17, "section": "unverified", "item": "i1", "cell": "r1c1#1", "cell_text": "Telegram(30K)",
         "values": [30000, -30000], "reason": "period not rebuilt",
         "passes": [_seen("users", "2023", "count", "actual", "period not rebuilt")] * 3},
        {**table, "section": "unverified", "item": "i20", "cell": "r6c2#1", "cell_text": "£ 42,638",
         "values": [42638], "reason": "other (type Other)",
         "passes": [_seen("other", "FY2022", "GBP", "actual", "other (type Other)")] * 3}]
    assert [r for r in rows if r["section"] == "disagreeing"] == [
        {**table, "section": "disagreeing", "item": "i11", "cell": "r4c3#1", "cell_text": "£ 150,000",
         "values": [150000], "reason": None,
         "passes": [_seen("revenue", "FY2023", "GBP", "forecast", "verified"),
                    _seen("revenue", "FY2023", "USD", "forecast", "verified"),
                    _seen("revenue", "2023", "GBP", "forecast", "verified")]},
        {**panel, "page": 19, "section": "disagreeing", "item": "i2", "cell": "r3c1#1", "cell_text": "5K Users",
         "values": [5000], "reason": None,
         "passes": [_seen("users", "FY2023", "count", "forecast", "verified")] * 2
         + [_seen("customers", "FY2023", "count", "forecast", "verified")]}], "the items labelled differently"

    path = script.write_report(report, tmp_path, 3, fake=True)
    rows.append({**table, "section": "unverified", "item": "i9", "cell": "r9c9#1", "cell_text": "Plan | B\nnext",
                 "values": [2500, 2.5], "reason": "other", "passes": [None, _seen("product", None, None, "unknown",
                                                                               "other (type Other)"), None]})
    diagnostic = script.write_diagnostic(rows, path, 3)
    assert diagnostic == tmp_path / f"{path.stem}_diagnostic.md", "beside its report, named after it"
    text = diagnostic.read_text(encoding="utf-8")
    head = "| Deck | Page | Type | Item | Cell | Cell text | Python's values | Reason | Pass 1 | Pass 2 | Pass 3 |"
    for line in (
            f"# Consistency run {path.stem.split('_', 1)[1]}: diagnostic", head,
            "| 05-zero2hero.pdf | 17 | kpi_panel | i1 | r1c1#1 | Telegram(30K) | 30000 or -30000 | period not rebuilt | "
            + " | ".join(["users 2023 count actual: period not rebuilt"] * 3) + " |",
            "| 05-zero2hero.pdf | 19 | table | i9 | r9c9#1 | Plan \\| B next | 2500 or 2.5 | other | not read | "
            "product null null unknown: other (type Other) | not read |",
            "| 05-zero2hero.pdf | 19 | table | i11 | r4c3#1 | £ 150,000 | 150000 |  | revenue FY2023 GBP forecast: verified | "
            "revenue FY2023 USD forecast: verified | revenue 2023 GBP forecast: verified |"):
        assert line in text.splitlines(), line
    assert text.count(head) == 2, "both sections take the same columns"
    report_text = path.read_text(encoding="utf-8")
    assert "5K Users" not in report_text and "£ 150,000" not in report_text, "the report itself holds no cell text"


def test_the_diagnostic_runs_only_on_the_10_public_test_decks(monkeypatch, tmp_path, capsys):
    """--diagnostic writes cell text, so it refuses any deck that is not one of the 10 public test decks, by file
    name and SHA-256: another file in the folder, a public name holding other bytes, a path to a public deck. It
    refuses before anything is read, sent or written; without --diagnostic every deck runs as before."""
    pytest.importorskip("pdfplumber")
    import hashlib
    import shutil
    import tempfile
    script = _consistency_script()
    on_disk = {p.name: hashlib.sha256(p.read_bytes()).hexdigest()
               for p in DECKS.iterdir() if not p.name.startswith(".")}
    assert script.PUBLIC_DECKS == on_disk and len(on_disk) == 10, "the 10 decks of the folder, byte for byte"
    folder, temp = tmp_path / "decks", tmp_path / "temp"
    folder.mkdir()
    temp.mkdir()
    shutil.copy(DECKS / "05-zero2hero.pdf", folder / "05-zero2hero.pdf")
    shutil.copy(DECKS / "05-zero2hero.pdf", folder / "client-deck.pdf")
    (folder / "10-tea.pdf").write_bytes((DECKS / "10-tea.pdf").read_bytes() + b"\n")
    monkeypatch.setattr(script, "DECKS", folder)
    monkeypatch.setattr(tempfile, "tempdir", str(temp))
    for deck in ("client-deck.pdf", "10-tea.pdf", "../decks/05-zero2hero.pdf", str(folder / "05-zero2hero.pdf")):
        with pytest.raises(SystemExit, match=f"--diagnostic runs on the 10 public test decks only: {deck} is not one"):
            script.main(["--fake", "--diagnostic", "--deck", "05-zero2hero.pdf", "--deck", deck])
    monkeypatch.delenv("ANTHROPIC_API_KEY", raising=False)
    with pytest.raises(SystemExit, match="public test decks only: client-deck.pdf"):
        script.main(["--yes", "--diagnostic", "--deck", "client-deck.pdf"])
    assert not any(temp.iterdir()), "refused before anything is written"
    script.main(["--fake", "--deck", "client-deck.pdf"])
    assert [p.name.endswith("_diagnostic.md") for p in temp.iterdir()] == [False], "without --diagnostic, as before"
    capsys.readouterr()

    monkeypatch.setattr(script, "DECKS", DECKS)
    script.main(["--fake", "--diagnostic", "--deck", "02-moz.pdf"])
    out = capsys.readouterr().out
    diagnostic, = temp.glob("consistency_*_diagnostic.md")
    assert diagnostic.with_name(diagnostic.name.replace("_diagnostic", "")).exists()
    assert f"Diagnostic: {diagnostic}" in out.splitlines(), "its path is printed"
    assert "~9 Months" in diagnostic.read_text(encoding="utf-8") and "~9 Months" not in out, \
        "its cell text is not: stdout carries the report only (moz p20: the months figure, labelled other)"


# ---------------------------------------------------------------------------
# Leftover risks closed in this change
# ---------------------------------------------------------------------------
def test_a_background_failure_marks_the_deck_not_read_and_logs_the_error_type_only(monkeypatch, caplog):
    import logging
    client, db, adapter = _deck_api(monkeypatch)

    async def boom(*args, **kwargs):
        raise RuntimeError("cannot read cell 'Jane Doe (CEO)'")
    monkeypatch.setattr(structures, "process_deck", boom)
    with caplog.at_level(logging.INFO):
        _upload_deck(client)
    assert _deck(client)["decks"][0]["ai_status"] == "not_read", "never left 'reading'"
    assert "structure reading failed: run_id=audit-s" in caplog.text and "error=RuntimeError" in caplog.text
    assert "Jane Doe" not in caplog.text


def test_a_short_engagement_reference_is_withheld_as_a_word_and_never_inside_one():
    db = _db(engagement_reference="E7")
    result, adapter = _read(db, "r1c1: Plan E7\nr1c2: £1M")
    assert (result.reason, adapter.calls) == ("refused: engagement_reference", 0), "sent as written, it is refused"
    cells, _ = redact.redact_structure([{"row": 1, "col": 1, "text": "Plan E7"}, {"row": 1, "col": 2, "text": "£1M"}],
                                       "Zero2Hero", {}, redact.withheld_values(db["audits"].docs[0]))
    result, adapter = _read(db, redact.structure_text(cells), replies=[NOTHING])
    assert result.status == "read" and json.loads(adapter.payloads[0])["text"] == "r1c1: Plan [redacted]\nr1c2: £1M"
    result, adapter = _read(db, "r1c1: Plan E70\nr1c2: £1M", replies=[NOTHING])
    assert result.status == "read", "E70 is not the reference"


def test_ordinary_text_with_a_dot_is_not_taken_for_a_file_name():
    result, _ = _read(_db(), "r1c1: 2.key metrics\nr1c2: £1M", replies=[NOTHING])
    assert result.status == "read"
    result, adapter = _read(_db(), "r1c1: See plan_v2.xlsx\nr1c2: £1M")
    assert result.status == "refused" and adapter.calls == 0


def test_a_metric_must_belong_to_the_kind_of_structure_read():
    deck_reply = {**REPLY, "labels": [{**REPLY["labels"][0], "metric": "invoice_date"}, REPLY["labels"][1]]}
    result, adapter = _read(_db(), replies=[deck_reply, deck_reply])
    assert result.status == "not_read" and adapter.calls == 2, "a spreadsheet field is not a deck claim"
    mapping_reply = {"type": "column_mapping", "items": [{**MAPPING_REPLY["items"][0], "metric": "revenue_growth"}]}
    result, adapter = _read(_db(), "r1c1: Customer\nr1c2: Amount\nc2 sample: 1200", "column_mapping",
                            replies=[mapping_reply, mapping_reply])
    assert result.status == "not_read" and adapter.calls == 2, "a claim type is not a spreadsheet field"
    both = {"type": "column_mapping", "items": [{**MAPPING_REPLY["items"][0], "metric": "revenue", "value_cell": "r1c2"}]}
    result, _ = _read(_db(), "r1c1: Customer\nr1c2: Revenue\nc2 sample: 1200", "column_mapping", replies=[both])
    assert result.status == "read", "revenue is a P&L field and a claim type"


def test_an_edited_ai_row_keeps_the_label_under_parsed_never_next_to_the_analysts_value(monkeypatch):
    client, db, adapter = _deck_api(monkeypatch)
    _map_revenue(client)
    _upload_deck(client)
    row = next(c for c in _deck(client)["candidates"] if c.get("ai_label") == "Verified")
    edited = client.put(f"/api/audits/{AUDIT}/decks/candidates/{row['id']}", json={"value": 1}).json()
    assert edited["status"] == "edited" and edited["ai_label"] is None
    assert edited["parsed"]["ai_label"] == "Verified" and edited["parsed"]["value"] == row["value"]


def test_every_structure_refusal_reason_in_the_code_is_declared_in_the_boundary_test():
    import ast as _ast
    import test_gateway_data_boundary as boundary
    found = set()
    for path, function in ((BACKEND / "app" / "llm" / "gateway.py", "structure_text_problem"),
                           (BACKEND / "app" / "structures" / "redact.py", "column_text_problem")):
        tree = _ast.parse(path.read_text(encoding="utf-8"))
        body = next(n for n in _ast.walk(tree) if isinstance(n, _ast.FunctionDef) and n.name == function)
        found |= {n.value.value for n in _ast.walk(body) if isinstance(n, _ast.Return)
                  and isinstance(n.value, _ast.Constant) and isinstance(n.value.value, str)}
        found |= {"client_name", "engagement_reference"} if function == "structure_text_problem" else set()
    assert found == boundary.STRUCTURE_REFUSALS, "a new refusal reason is declared in test_gateway_data_boundary.py"
