"""Revenue file recognition (docs/specs/revenue-file-recognition.md): a file is a revenue file when it holds a revenue measure, a
reporting reference and a number (section 2), and it is then classified into one of six types (section 3), per sheet, with the
labels of its rows read as metric names or as entities (section 4). Python checks every proposal of the model before anything is
used (section 6). Synthetic files only (CLAUDE.md rule 15): tests/fixtures/recognition/ and the workbooks built here.

The module is imported through the `rec` fixture. The payload the model reads is tested in test_gateway_data_boundary.py (rule 14)."""
import importlib
import io
import json
import logging
from pathlib import Path

import pytest

pytest.importorskip("fastapi")
pytest.importorskip("motor")

import test_llm_gateway as t  # noqa: E402
from test_chat_upload import AUDIT, adapter_of, upload  # noqa: E402

from app import column_rules as cr  # noqa: E402

FIXTURES = Path(__file__).resolve().parent / "fixtures" / "recognition"
MAPPING = Path(__file__).resolve().parent / "fixtures" / "mapping"
SAMPLE = Path(__file__).resolve().parents[2] / "sample_data"
MONTHLY = "monthly_report_actual_forecast.csv"
AS_OF = "2025-06-30"                        # the last actual month of MONTHLY is June 2025
AS_OF_YEAR_END = "2025-12-31"


@pytest.fixture()
def rec():
    return importlib.import_module("app.recognition")


def sheets_of(name):
    return cr.read_workbook((FIXTURES / name).read_bytes(), name)


def csv_sheets(text, name="x.csv"):
    return cr.read_workbook(text.encode(), name)


def recognise(rec, name, as_of=AS_OF_YEAR_END, proposal=None):
    return rec.recognise(sheets_of(name), as_of_month=as_of, proposal=proposal)


def proposal(**over):
    """What a validated model reply hands Python (section 5): the sheet it read, the type, the revenue row with its label, the last
    actual column, the currency and the rows that are not revenue. Rows and columns are sheet positions (A = 1, row 1 = 1)."""
    body = {"sheet": 1, "file_type": "monthly_revenue_report", "revenue_row": {"row": 4, "label": "Recognised revenue (EUR)"},
            "last_actual_column": 7, "currency": "EUR",
            "non_revenue_rows": [{"row": 9, "kind": "forecast"}, {"row": 10, "kind": "bookings"}, {"row": 11, "kind": "pipeline"}]}
    body.update(over)
    return body


def recognition_reply(**over):
    return json.dumps({"type": "file_recognition", **proposal(**over)})


# ---------------------------------------------------------------------------
# Section 3: six types, classified per sheet
# ---------------------------------------------------------------------------
def test_a_metrics_by_months_sheet_with_actual_and_forecast_columns_is_a_monthly_revenue_report_with_its_revenue_row(rec):
    got = recognise(rec, MONTHLY, as_of=AS_OF)
    assert got["recognised"] is True and got["missing"] == []
    assert got["file_type"] == "monthly_revenue_report" and got["type_state"] == "auto" and got["type_source"] == "rules"
    assert got["label_mode"] == "metric"
    row = got["revenue_row"]
    assert (row["row"], row["label"], row["kind"], row["state"], row["confidence"]) == (4, "Recognised revenue (EUR)", "recognised", "auto", 100)
    assert got["series_basis"] == "row"


def test_the_boundary_is_the_last_actual_column_and_the_forecast_columns_are_excluded(rec):
    got = recognise(rec, MONTHLY, as_of=AS_OF)
    last = got["last_actual"]
    assert (last["column"], last["month"], last["state"], last["reason"]) == (7, "2025-06", "auto", None)
    assert got["excluded_columns"] == [8, 9, 10, 11, 12, 13], "July to December are tagged Forecast above the header"
    assert got["needs_click"] is False


def test_forecast_bookings_and_pipeline_rows_are_recorded_and_never_a_revenue_candidate(rec):
    got = recognise(rec, MONTHLY, as_of=AS_OF)
    assert {(r["row"], r["kind"]) for r in got["non_revenue_rows"]} == {(9, "forecast"), (10, "bookings"), (11, "pipeline")}
    assert {c["row"] for c in got["revenue_row"]["candidates"]} == {4, 5, 6}, "revenue, MRR and ARR: not 'Forecast revenue'"
    assert [c["kind"] for c in got["revenue_row"]["candidates"]] == ["recognised", "recurring", "recurring"], "recognised revenue ranks first"


def test_customers_by_periods_is_customer_revenue_and_the_customer_level_analyses_are_enabled(rec):
    got = recognise(rec, "customer_revenue_months.csv")
    assert got["recognised"] and got["file_type"] == "customer_revenue" and got["label_mode"] == "entity"
    assert got["label_share"] == pytest.approx(1 / 12, abs=1e-4) and got["confidence"] == 92, "one name of twelve holds a metric word"
    assert got["supports"] == ["customer_level", "trend"] and got["blocked"] == []
    assert got["series_basis"] == "sum_of_rows" and got["revenue_row"] is None, "the series is the sum of the customer rows"


def test_one_row_per_invoice_is_invoice_transactions(rec):
    got = recognise(rec, "invoice_transactions.csv")
    assert got["recognised"] and got["file_type"] == "invoice_transactions" and got["series_basis"] == "records"
    assert got["supports"] == ["customer_level", "trend"] and got["blocked"] == []


def test_invoices_without_a_customer_column_are_recognised_and_only_the_customer_level_analyses_are_blocked(rec):
    text = (FIXTURES / "invoice_transactions.csv").read_text().splitlines()
    kept = [",".join(c for i, c in enumerate(line.split(",")) if i != 2) for line in text]       # drop the Customer column
    got = rec.recognise(csv_sheets("\n".join(kept) + "\n"), as_of_month=AS_OF_YEAR_END)
    assert got["recognised"] and got["file_type"] == "invoice_transactions", "missing information narrows the analyses"
    assert got["supports"] == ["trend"]
    assert got["blocked"] == [{"analysis": "customer_level", "needs": rec.NEEDS_CUSTOMER_COLUMN}]


def test_categories_by_periods_is_a_revenue_summary_with_trend_and_no_customer_level(rec):
    got = recognise(rec, "revenue_summary_by_product.csv")
    assert got["recognised"] and got["file_type"] == "revenue_summary" and got["label_mode"] == "entity"
    assert got["supports"] == ["trend"] and [b["analysis"] for b in got["blocked"]] == ["customer_level", "monthly_series"]
    assert got["blocked"][0]["needs"] == rec.NEEDS_BY_CUSTOMER and got["blocked"][1]["needs"] == rec.NEEDS_MONTHLY


def test_a_multi_period_pnl_is_financial_statements_and_trend_is_enabled(rec):
    got = recognise(rec, "pnl_multi_period.csv")
    assert got["recognised"] and got["file_type"] == "financial_statements" and got["label_mode"] == "metric"
    assert got["periods"] == 3 and got["supports"] == ["trend"]
    assert (got["revenue_row"]["row"], got["revenue_row"]["state"]) == (3, "auto")


def test_yearly_periods_support_trend_and_the_blocked_list_says_the_monthly_series_is_missing(rec):
    got = recognise(rec, "pnl_multi_period.csv")
    assert got["supports"] == ["trend"], "three yearly columns: trend is supported"
    assert {"analysis": "monthly_series", "needs": rec.NEEDS_MONTHLY} in got["blocked"]
    assert rec.NEEDS_MONTHLY == "one column per month"
    assert any(line.startswith("Monthly series missing") and line.endswith("one column per month") for line in rec.missing_lines(got))
    for name in (MONTHLY, "customer_revenue_months.csv", "monthly_report_no_currency.csv"):
        assert "monthly_series" not in [b["analysis"] for b in recognise(rec, name)["blocked"]], name
    for name in ("arr_snapshot.csv", "invoice_transactions.csv"):
        assert "monthly_series" not in [b["analysis"] for b in recognise(rec, name)["blocked"]], name


def test_a_single_period_pnl_blocks_trend_and_names_the_data_that_would_unblock_it(rec):
    got = recognise(rec, "pnl_single_period.csv")
    assert got["recognised"] and got["file_type"] == "financial_statements" and got["periods"] == 1
    assert "trend" not in got["supports"]
    trend = next(b for b in got["blocked"] if b["analysis"] == "trend")
    assert trend["needs"] == rec.NEEDS_PERIODS and "at least one more period" in trend["needs"]
    assert (got["revenue_row"]["row"], got["revenue_row"]["state"]) == (3, "auto"), "one period is numeric in all its columns"


def test_a_single_period_arr_table_is_a_snapshot_for_claim_matching_only(rec):
    got = recognise(rec, "arr_snapshot.csv")
    assert got["recognised"] and got["file_type"] == "arr_mrr_snapshot" and got["series_basis"] == "snapshot"
    assert got["supports"] == ["claim_matching"]
    assert {b["analysis"]: b["needs"] for b in got["blocked"]} == {"customer_level": rec.NEEDS_BY_CUSTOMER, "trend": rec.NEEDS_SNAPSHOT_SERIES}


def test_every_blocked_analysis_names_the_data_that_would_unblock_it_in_the_data_inventory_lines(rec):
    """V6: never write Missing without naming the data. The lines go to missing_data (section 3)."""
    for name in ("pnl_single_period.csv", "arr_snapshot.csv", "revenue_summary_by_product.csv", MONTHLY):
        got = recognise(rec, name)
        lines = rec.missing_lines(got)
        assert len(lines) == len(got["blocked"]) > 0, name
        for line, blocked in zip(lines, got["blocked"]):
            assert line.endswith(blocked["needs"]) and "Missing: " in line and blocked["needs"].strip(), line
    assert rec.missing_lines(recognise(rec, "customer_revenue_months.csv")) == []
    assert any("at least one more period" in line for line in rec.missing_lines(recognise(rec, "pnl_single_period.csv")))


def test_names_the_rules_cannot_class_leave_the_type_to_the_analyst_with_both_candidates(rec):
    got = recognise(rec, "unclassed_names_by_quarter.csv")
    assert got["recognised"], "it is a revenue file; only its type is open"
    assert got["file_type"] is None and got["type_state"] == "click" and got["needs_click"] is True
    assert set(got["type_candidates"]) == {"customer_revenue", "revenue_summary"}


# ---------------------------------------------------------------------------
# Section 2: the three essentials, all required; currency; actuals against the rest
# ---------------------------------------------------------------------------
def test_a_sheet_with_bookings_and_pipeline_rows_but_no_revenue_row_is_not_a_revenue_file_and_the_message_names_the_missing_measure(rec):
    got = recognise(rec, "bookings_pipeline_only.csv")
    assert got["recognised"] is False and got["missing"] == ["revenue_measure"]
    assert "revenue measure" in got["message"] and "MRR" in got["message"], "the message names what is missing"
    assert got["file_type"] is None and got["revenue_row"] is None
    assert {r["kind"] for r in got["non_revenue_rows"]} == {"bookings", "pipeline", "forecast", "contract_value", "run_rate", "budget"}, \
        "the rows are recognised, none of them as revenue: 'Forecast revenue' is a forecast and 'Sales pipeline' a pipeline"


def test_a_file_with_a_revenue_measure_and_numbers_but_no_date_period_or_id_is_not_a_revenue_file(rec):
    got = rec.recognise(csv_sheets("Item,Value\nRevenue,1000\nCustomers,40\nChurn,3\n"), as_of_month=AS_OF_YEAR_END)
    assert got["recognised"] is False and "reference" in got["missing"]
    assert "date, period or transaction ID" in got["message"]


def test_a_revenue_row_with_no_number_is_not_a_revenue_file(rec):
    got = rec.recognise(csv_sheets("Metric,Jan-25,Feb-25,Mar-25\nRevenue,n/a,n/a,n/a\nCustomers,4,5,6\nChurn,1,2,3\nCash,7,8,9\n"),
                        as_of_month=AS_OF_YEAR_END)
    assert got["recognised"] is False and got["missing"] == ["numbers"]


def test_no_currency_anywhere_recognises_the_file_and_asks_the_analyst_to_confirm_it(rec):
    got = recognise(rec, "monthly_report_no_currency.csv")
    assert got["recognised"] is True
    assert got["currency"] == {"code": None, "state": "confirm"} and got["needs_click"] is True
    assert recognise(rec, MONTHLY, as_of=AS_OF)["currency"] == {"code": "EUR", "state": "seen"}


def test_a_currency_the_model_names_but_the_file_never_shows_is_dropped(rec):
    got = recognise(rec, "monthly_report_no_currency.csv", proposal=proposal(currency="USD", revenue_row={"row": 3, "label": "Revenue"},
                                                                          last_actual_column=13, non_revenue_rows=[]))
    assert got["currency"] == {"code": None, "state": "confirm"}, "Python finds the code in the cells or the analyst confirms it"


def test_a_turnover_row_waits_for_the_revenue_or_volume_answer(rec):
    text = "Metric,Jan-25,Feb-25,Mar-25,Apr-25\nTurnover (EUR),1100,1200,1300,1400\nCustomers,4,5,6,7\nChurn,1,2,3,4\nCash,5,6,7,8\n"
    got = rec.recognise(csv_sheets(text), as_of_month="2025-04-30")
    assert got["recognised"] and got["money_kind"] == {"asked": True}
    assert (got["revenue_row"]["state"], got["revenue_row"]["reason"]) == ("click", "turnover_question")


# ---------------------------------------------------------------------------
# Section 4: row labels or entities, a confidence per sheet
# ---------------------------------------------------------------------------
def test_the_label_share_threshold_separates_the_metric_sheets_from_the_entity_sheets_and_a_confidence_is_recorded_per_sheet(rec):
    assert (rec.METRIC_LABEL_SHARE, rec.ENTITY_LABEL_SHARE) == (0.5, 0.2)
    metric = ("monthly_report_actual_forecast.csv", "pnl_multi_period.csv", "pnl_single_period.csv", "bookings_pipeline_only.csv",
              "monthly_report_no_currency.csv")
    entity = ("customer_revenue_months.csv", "revenue_summary_by_product.csv", "unclassed_names_by_quarter.csv")
    for name in metric:
        got = recognise(rec, name)
        assert got["label_mode"] == "metric" and got["label_share"] > rec.METRIC_LABEL_SHARE, (name, got["label_share"])
        assert isinstance(got["confidence"], int) and 0 <= got["confidence"] <= 100
    for name in entity:
        got = recognise(rec, name)
        assert got["label_mode"] == "entity" and got["label_share"] <= rec.ENTITY_LABEL_SHARE, (name, got["label_share"])
        assert isinstance(got["confidence"], int) and 0 <= got["confidence"] <= 100


def test_a_label_column_headed_customer_is_never_metric_labelled_even_when_every_name_holds_a_metric_word(rec):
    text = ("Revenue by customer,,,\nCustomer,Jan-25,Feb-25,Mar-25\nRevenue Partners,1,2,3\nSales Hub,4,5,6\nCash Converters,7,8,9\n"
            "Margin Foods,1,2,3\n")
    got = rec.recognise(csv_sheets(text), as_of_month="2025-03-31")
    assert got["label_mode"] != "metric", "four of four names match the vocabulary, but the header says customer"


def test_the_share_is_a_share_of_label_rows_with_a_number_so_notes_and_blank_lines_do_not_move_it(rec):
    text = ("Metric,Jan-25,Feb-25,Mar-25\nRevenue,1,2,3\nCustomers,4,5,6\nChurn,7,8,9\nNotes: figures are unaudited,,,\n,,,\nSee the appendix,,,\n")
    got = rec.recognise(csv_sheets(text), as_of_month="2025-03-31")
    assert got["label_mode"] == "metric" and got["label_share"] == 1.0


# ---------------------------------------------------------------------------
# Workbooks: the best sheet, the others recorded
# ---------------------------------------------------------------------------
def kpi_and_pnl_bytes(order=("KPIs", "P&L")):
    from openpyxl import Workbook
    months = ["Jan-25", "Feb-25", "Mar-25", "Apr-25", "May-25", "Jun-25", "Jul-25", "Aug-25", "Sep-25", "Oct-25", "Nov-25", "Dec-25"]
    sheets = {"KPIs": ["MRR", "ARR", "Customers", "Logo churn %"],
              "P&L": ["Revenue", "Cost of sales", "Gross profit", "Payroll", "EBITDA"]}
    wb = Workbook()
    wb.remove(wb.active)
    for name in order:
        ws = wb.create_sheet(name)
        ws.append([f"{name} 2025"] + [None] * 12)
        ws.append(["Metric"] + months)
        for r, label in enumerate(sheets[name]):
            ws.append([label] + [20000 + 100 * r + c for c in range(12)])
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def kpi_and_pnl_workbook(order=("KPIs", "P&L")):
    return cr.read_workbook(kpi_and_pnl_bytes(order), "pack.xlsx")


def test_a_workbook_with_a_kpi_sheet_and_a_pnl_sheet_takes_the_best_sheet_and_records_the_other(rec):
    got = rec.recognise(kpi_and_pnl_workbook(), as_of_month=AS_OF_YEAR_END)
    assert got["sheet"] == "P&L" and got["file_type"] == "financial_statements", "a recognised revenue row beats an MRR row"
    assert got["other_sheets"] == [{"sheet": "KPIs", "file_type": "monthly_revenue_report", "recognised": True}]


def test_the_best_sheet_does_not_depend_on_the_sheet_order(rec):
    first = rec.recognise(kpi_and_pnl_workbook(("KPIs", "P&L")), as_of_month=AS_OF_YEAR_END)
    second = rec.recognise(kpi_and_pnl_workbook(("P&L", "KPIs")), as_of_month=AS_OF_YEAR_END)
    assert first["sheet"] == second["sheet"] == "P&L"
    assert second["other_sheets"] == first["other_sheets"]


def test_a_workbook_with_no_recognised_sheet_reports_why_for_its_best_candidate(rec):
    from openpyxl import Workbook
    wb = Workbook()
    wb.active.title = "Cover"
    wb.active.append(["Notes only"])
    ws = wb.create_sheet("Pipeline")
    ws.append(["Metric", "Jan-25", "Feb-25", "Mar-25"])
    ws.append(["Bookings", 1, 2, 3])
    buf = io.BytesIO()
    wb.save(buf)
    got = rec.recognise(cr.read_workbook(buf.getvalue(), "x.xlsx"), as_of_month=AS_OF_YEAR_END)
    assert got["recognised"] is False and "revenue_measure" in got["missing"]


# ---------------------------------------------------------------------------
# Section 6: Python checks every proposal before anything is used
# ---------------------------------------------------------------------------
@pytest.mark.parametrize("revenue_row, reason", [
    ({"row": 7, "label": "Customers"}, "no_revenue_word"),                         # numeric, no revenue-vocabulary word
    ({"row": 999, "label": "Revenue"}, "row_missing"),
    ({"row": 9, "label": "Forecast revenue"}, "marked_not_revenue"),               # a forecast row is never revenue
    ({"row": 4, "label": "Total turnover"}, "label_mismatch"),                     # the label is not the label of that row
    ({"row": 5, "label": "MRR"}, "rules_and_model_disagree"),                      # a revenue word and numeric, but not the rules' row
])
def test_a_proposed_revenue_row_that_fails_a_check_is_unsure_and_needs_a_click(rec, revenue_row, reason):
    got = recognise(rec, MONTHLY, as_of=AS_OF, proposal=proposal(revenue_row=revenue_row))
    row = got["revenue_row"]
    assert (row["state"], row["reason"]) == ("click", reason)
    assert got["needs_click"] is True
    assert {c["row"] for c in row["candidates"]} >= {4}, "the analyst chooses among the candidate rows"
    assert 9 not in {c["row"] for c in row["candidates"]}


def test_a_proposed_revenue_row_numeric_in_fewer_than_three_columns_is_unsure(rec):
    text = ("Metric,Jan-25,Feb-25,Mar-25,Apr-25,May-25\nRevenue,100,110,120,130,140\nRevenue (restated),,,,,150\nCustomers,4,5,6,7,8\n")
    got = rec.recognise(csv_sheets(text), as_of_month="2025-05-31", proposal=proposal(
        revenue_row={"row": 3, "label": "Revenue (restated)"}, last_actual_column=6, currency=None, non_revenue_rows=[]))
    assert (got["revenue_row"]["state"], got["revenue_row"]["reason"]) == ("click", "not_numeric_enough")


def test_a_sheet_with_fewer_than_three_periods_needs_its_row_numeric_in_each_of_them(rec):
    """The three-column check rejects a stray row; with two periods the row needs both, with one period it needs that one."""
    two = rec.recognise(csv_sheets("Line,FY2024,FY2025\nRevenue,100,110\nCost of sales,50,55\nPayroll,20,22\nEBITDA,30,33\n"),
                        as_of_month=AS_OF_YEAR_END)
    assert two["recognised"] and two["revenue_row"]["state"] == "auto" and two["supports"] == ["trend"]


def test_the_model_and_the_rules_agree_the_revenue_row_and_the_boundary_are_accepted(rec):
    got = recognise(rec, MONTHLY, as_of=AS_OF, proposal=proposal())
    assert (got["revenue_row"]["row"], got["revenue_row"]["state"]) == (4, "auto")
    assert (got["last_actual"]["column"], got["last_actual"]["state"]) == (7, "auto")
    assert got["needs_click"] is False and got["type_source"] == "rules"


def test_a_rule_confidence_of_exactly_80_needs_a_click_and_81_or_more_does_not(rec):
    """The thresholds of the column mapping (chat-upload.md 4.1, 21): a label with a revenue word among other words scores 80."""
    text = "Metric,Jan-25,Feb-25,Mar-25\nRevenue - EMEA,100,110,120\nCustomers,4,5,6\nChurn,1,2,3\n"
    got = rec.recognise(csv_sheets(text), as_of_month="2025-03-31")
    assert (got["revenue_row"]["confidence"], got["revenue_row"]["state"], got["revenue_row"]["reason"]) == (80, "click", "confidence_80")
    exact = rec.recognise(csv_sheets(text.replace("Revenue - EMEA", "Revenue")), as_of_month="2025-03-31")
    assert (exact["revenue_row"]["confidence"], exact["revenue_row"]["state"]) == (100, "auto")


def test_two_revenue_rows_of_the_same_kind_and_score_are_a_tie_and_wait_for_a_click(rec):
    text = "Metric,Jan-25,Feb-25,Mar-25\nRevenue,100,110,120\nNet sales,101,111,121\nCustomers,4,5,6\n"
    got = rec.recognise(csv_sheets(text), as_of_month="2025-03-31")
    assert (got["revenue_row"]["state"], got["revenue_row"]["reason"]) == ("click", "tie")


def test_the_model_cannot_move_the_boundary_the_as_of_month_sets_it(rec):
    got = recognise(rec, MONTHLY, as_of=AS_OF, proposal=proposal(last_actual_column=13))        # December 2025: ignored
    last = got["last_actual"]
    assert (last["column"], last["state"], last["reason"], last["candidates"]) == (7, "auto", None, [])
    assert last["as_of"] == AS_OF[:7] and got["needs_click"] is False


def test_the_boundary_is_the_last_column_up_to_the_as_of_month_and_the_columns_after_it_are_plan(rec):
    got = recognise(rec, MONTHLY, as_of="2025-03-31")
    last = got["last_actual"]
    assert (last["state"], last["reason"], last["month"]) == ("auto", None, "2025-03")
    assert last["column"] < 7 and got["excluded_columns"], "the months after the as-of month are not read"


def test_with_no_as_of_month_the_boundary_is_the_rules_last_actual_column_and_waits_for_nothing(rec):
    got = recognise(rec, MONTHLY, as_of=None, proposal=proposal())
    last = got["last_actual"]
    assert (last["column"], last["state"], last["reason"]) == (7, "auto", None)
    assert got["revenue_row"]["state"] == "auto" and got["needs_click"] is False


@pytest.mark.parametrize("column", [1, 99, 5])
def test_a_proposed_boundary_is_ignored_whatever_it_names(rec, column):
    got = recognise(rec, MONTHLY, as_of=AS_OF, proposal=proposal(last_actual_column=column))
    assert (got["last_actual"]["column"], got["last_actual"]["state"], got["last_actual"]["reason"]) == (7, "auto", None)


def test_a_file_that_starts_after_the_as_of_month_waits_with_after_as_of(rec):
    got = recognise(rec, MONTHLY, as_of="2024-01-31")
    assert (got["last_actual"]["state"], got["last_actual"]["reason"]) == ("click", "after_as_of") and got["needs_click"] is True


def test_the_model_marking_the_revenue_row_as_a_forecast_removes_it_from_the_candidates(rec):
    got = recognise(rec, MONTHLY, as_of=AS_OF, proposal=proposal(non_revenue_rows=[{"row": 4, "kind": "forecast"}]))
    assert 4 not in {c["row"] for c in got["revenue_row"]["candidates"]}
    assert (4, "forecast") in {(r["row"], r["kind"]) for r in got["non_revenue_rows"]}
    assert got["revenue_row"]["state"] == "click"


def test_a_row_the_model_marks_as_non_revenue_is_kept_as_a_claim_row_with_its_closed_kind(rec):
    got = recognise(rec, MONTHLY, as_of=AS_OF, proposal=proposal(non_revenue_rows=[{"row": 12, "kind": "run_rate"}]))
    assert (12, "run_rate") in {(r["row"], r["kind"]) for r in got["non_revenue_rows"]}


@pytest.mark.parametrize("header", ["Month", "Period", "Date"])
def test_a_header_reading_month_period_or_date_is_a_reporting_reference(rec, header):
    text = f"{header},Revenue,Cost of sales,Payroll\n2025-01-01,100,40,20\n2025-02-01,110,44,20\n2025-03-01,120,48,20\n"
    got = rec.recognise(csv_sheets(text), as_of_month="2025-03-31")
    assert got["recognised"] is True and "reference" not in got["missing"], got["message"]


def test_a_pnl_with_dates_down_the_rows_and_measures_across_is_recognised_as_financial_statements(rec):
    sheets = cr.read_workbook((MAPPING / "accountant_pnl.csv").read_bytes(), "accountant_pnl.csv")
    got = rec.recognise(sheets, as_of_month="2024-12-31")
    assert got["recognised"] is True and got["missing"] == [] and got["message"] is None
    assert got["file_type"] == "financial_statements" and got["series_basis"] == "column"
    assert got["revenue_column"]["label"] == "Revenue" and got["revenue_row"] is None


def last_date(path, column):
    rows = [line.split(",") for line in path.read_text().splitlines()[1:] if line.strip()]
    return max(r[column][:10] for r in rows)


@pytest.mark.parametrize("name, column", [("pnl.csv", 0), ("revenue.csv", 1)])
def test_the_clean_run_files_get_no_click_and_no_not_a_revenue_file_line(rec, name, column):
    """sample_data: the P&L (dates down the rows, no currency in the file) and the revenue file (customer, date, amount, currency)."""
    path = SAMPLE / name
    got = rec.recognise(cr.read_workbook(path.read_bytes(), name), as_of_month=last_date(path, column), reporting_currency="EUR")
    assert got["recognised"] is True and got["message"] is None
    assert got["needs_click"] is False, "recognition adds no click to a clean run"
    assert got["file_type"] == {"pnl.csv": "financial_statements", "revenue.csv": "invoice_transactions"}[name]
    if name == "pnl.csv":
        assert got["currency"] == {"code": "EUR", "state": "audit"}, "a P&L takes the audit's reporting currency, as it does today"


def test_the_extra_unsure_sheets_are_unsure_with_confidence_zero_and_their_labels_are_masked(rec):
    for name, share in (("unsure_share_labels.csv", 0.3), ("unsure_few_hits.csv", 2 / 3)):
        got = recognise(rec, name, as_of="2025-04-30")
        assert got["label_mode"] == "unsure" and got["confidence"] == 0, name
        assert got["label_share"] == pytest.approx(share, abs=1e-4), name
    assert recognise(rec, "unsure_few_hits.csv", as_of="2025-03-31")["label_share"] > rec.METRIC_LABEL_SHARE, "above 0.5 but only 2 hits"


def category_sums(rec, name):
    sheet = sheets_of(name)[0]
    body = [r for r in sheet.grid[sheet.header_row:] if r and r[0] and not str(r[0]).startswith("Total")]
    return [sum(float(r[c]) for r in body) for c in range(1, 4)]


def test_a_total_row_is_never_added_to_the_series_and_stays_available_for_the_cross_check(rec):
    got = recognise(rec, "revenue_summary_by_product.csv")
    assert [s["value"] for s in got["series"]] == category_sums(rec, "revenue_summary_by_product.csv"), "the categories once, not twice"
    assert got["total_rows"] == [{"row": 9, "label": "Total revenue", "matches": True}]
    snap = recognise(rec, "arr_snapshot.csv", as_of="2025-06-30")
    assert [s["value"] for s in snap["series"]] == [1070000], "120000 + 340000 + 610000, not 2,140,000"
    assert snap["total_rows"] == [{"row": 6, "label": "Total ARR", "matches": True}]


def test_a_row_whose_values_equal_the_sum_of_the_rows_above_is_a_total_row_even_without_the_word(rec):
    text = ("Revenue by product line (EUR),,,\nProduct line,FY2023,FY2024,FY2025\nHardware,100,110,120\nSoftware,200,210,220\n"
            "Services,300,310,320\nAll products,600,630,660\n")
    got = rec.recognise(csv_sheets(text), as_of_month=AS_OF_YEAR_END)
    assert [s["value"] for s in got["series"]] == [600, 630, 660]
    assert got["total_rows"] == [{"row": 6, "label": "All products", "matches": True}]


def test_a_total_that_does_not_match_its_rows_is_excluded_and_flagged(rec):
    text = ("Revenue by product line (EUR),,,\nProduct line,FY2023,FY2024,FY2025\nHardware,100,110,120\nSoftware,200,210,220\n"
            "Total,999,630,660\n")
    got = rec.recognise(csv_sheets(text), as_of_month=AS_OF_YEAR_END)
    assert [s["value"] for s in got["series"]] == [300, 320, 340]
    assert got["total_rows"] == [{"row": 5, "label": "Total", "matches": False}]


def test_the_model_naming_another_type_than_the_rules_makes_the_type_a_click_with_both_candidates(rec):
    got = recognise(rec, MONTHLY, as_of=AS_OF, proposal=proposal(file_type="revenue_summary"))
    assert (got["file_type"], got["type_state"], got["type_reason"], got["type_source"]) == \
        ("monthly_revenue_report", "click", "rules_and_model_disagree", "rules")
    assert set(got["type_candidates"]) == {"monthly_revenue_report", "revenue_summary"} and got["needs_click"] is True


def test_with_no_rules_type_the_models_choice_among_the_candidates_is_an_ai_suggestion_that_waits(rec):
    base = dict(revenue_row=None, last_actual_column=5, currency=None, non_revenue_rows=[])
    got = recognise(rec, "unclassed_names_by_quarter.csv", proposal=proposal(file_type="customer_revenue", **base))
    assert (got["file_type"], got["type_state"], got["type_reason"], got["type_source"]) == ("customer_revenue", "click", "type_unsure", "ai")
    outside = recognise(rec, "unclassed_names_by_quarter.csv", proposal=proposal(file_type="monthly_revenue_report", **base))
    assert outside["file_type"] is None and outside["type_state"] == "click", "a type outside the candidates is ignored"


def test_with_two_periods_a_row_numeric_in_only_one_is_unsure(rec):
    text = "Line,FY2024,FY2025\nRevenue,100,\nCost of sales,50,55\nPayroll,20,22\nEBITDA,30,33\n"
    got = rec.recognise(csv_sheets(text), as_of_month=AS_OF_YEAR_END,
                        proposal=proposal(revenue_row={"row": 2, "label": "Revenue"}, last_actual_column=3, currency=None, non_revenue_rows=[]))
    assert (got["revenue_row"]["state"], got["revenue_row"]["reason"]) == ("click", "not_numeric_enough")


def test_the_choices_waiting_count_a_turnover_rows_two_clicks_as_one(rec):
    text = "Metric,Jan-25,Feb-25,Mar-25,Apr-25\nTurnover (EUR),1100,1200,1300,1400\nCustomers,4,5,6,7\nChurn,1,2,3,4\nCash,5,6,7,8\n"
    got = rec.recognise(csv_sheets(text), as_of_month="2025-04-30")
    assert got["revenue_row"]["state"] == "click" and got["money_kind"]["asked"] is True
    assert rec.choices_waiting(got) == 1
    assert rec.choices_waiting(recognise(rec, MONTHLY, as_of=AS_OF)) == 0
    assert rec.choices_waiting(recognise(rec, MONTHLY, as_of=None)) == 0, "the boundary waits for nothing"


def test_growth_labels_name_the_row_and_never_call_a_recurring_row_revenue(rec):
    assert rec.growth_label(recognise(rec, MONTHLY, as_of=AS_OF)) == "Revenue growth"
    mrr = rec.recognise(csv_sheets("Metric,Jan-25,Feb-25,Mar-25\nMRR,100,110,120\nCustomers,4,5,6\nChurn,1,2,3\nCash,5,6,7\n"), as_of_month="2025-03-31")
    arr = rec.recognise(csv_sheets("Metric,Jan-25,Feb-25,Mar-25\nARR,100,110,120\nCustomers,4,5,6\nChurn,1,2,3\nCash,5,6,7\n"), as_of_month="2025-03-31")
    assert (rec.growth_label(mrr), rec.growth_label(arr)) == ("MRR growth", "ARR growth")
    assert mrr["revenue_row"]["kind"] == "recurring" and arr["revenue_row"]["kind"] == "recurring"


# ---------------------------------------------------------------------------
# Determinism
# ---------------------------------------------------------------------------
def test_the_same_input_gives_the_same_output_twice_and_after_reading_the_bytes_again(rec):
    first = recognise(rec, MONTHLY, as_of=AS_OF, proposal=proposal())
    second = recognise(rec, MONTHLY, as_of=AS_OF, proposal=proposal())
    assert first == second
    assert json.dumps(first, sort_keys=True) == json.dumps(second, sort_keys=True)
    for name in sorted(p.name for p in FIXTURES.glob("*.csv")):
        assert recognise(rec, name) == recognise(rec, name), name


# ---------------------------------------------------------------------------
# The upload: one model call per new file, cached by the file hash (section 5)
# ---------------------------------------------------------------------------
def monthly_bytes():
    return (FIXTURES / MONTHLY).read_bytes()


def set_as_of(api, value=AS_OF):
    api.db["audits"].docs[0].update(as_of_month=value, as_of_by="analyst")        # typed by the analyst: Map never overwrites it


def recognition_calls(api):
    return [p for p in adapter_of(api).payloads if json.loads(p)["type"] == "file_recognition"]


def test_an_upload_is_recognised_stored_as_the_revenue_file_and_asks_the_model_once(api):
    set_as_of(api)
    api.state["adapter"] = t.FakeAdapter(replies=[recognition_reply()])
    res = upload(api, MONTHLY, monthly_bytes())
    assert res.status_code == 200, res.text
    body = res.json()
    assert body["dtype"] == "revenue" and body["recognition"]["recognised"] is True
    got = body["recognition"]
    assert (got["file_type"], got["type_source"], got["type_state"]) == ("monthly_revenue_report", "rules", "auto")
    assert (got["revenue_row"]["row"], got["revenue_row"]["state"]) == (4, "auto")
    assert (got["last_actual"]["column"], got["last_actual"]["state"]) == (7, "auto")
    assert len(recognition_calls(api)) == 1 and adapter_of(api).calls == 1


def test_the_same_file_twice_makes_one_model_call_and_a_new_file_makes_another(api):
    set_as_of(api)
    api.state["adapter"] = t.FakeAdapter(replies=[recognition_reply()])
    upload(api, MONTHLY, monthly_bytes())
    api.db["datasets"].docs.clear()
    again = upload(api, MONTHLY, monthly_bytes())
    assert again.status_code == 200 and again.json()["recognition"]["recognised"] is True
    assert len(recognition_calls(api)) == 1, "the answer is kept by the file hash"
    kept = [d for d in api.db["column_mappings"].docs if d.get("kind") == "file_recognition"]
    assert len(kept) == 1 and len(kept[0]["file_hash"]) == 64
    assert set(kept[0]) <= {"audit_id", "kind", "file_hash", "prompt_tag", "reply", "saved_at"} and kept[0]["prompt_tag"]
    other = (FIXTURES / "monthly_report_no_currency.csv").read_bytes()
    api.db["datasets"].docs.clear()
    upload(api, "monthly_report_no_currency.csv", other)
    assert len(recognition_calls(api)) == 2, "a new file is a new call"


def test_the_payload_and_the_row_labels_are_never_stored_or_logged(api, caplog):
    caplog.set_level(logging.DEBUG)
    set_as_of(api)
    api.state["adapter"] = t.FakeAdapter(replies=[recognition_reply()])
    upload(api, MONTHLY, monthly_bytes())
    stored = json.dumps([api.db[n].docs for n in ("llm_structures", "llm_calls", "column_mappings", "audits")], default=str)
    sent = [json.loads(p)["text"] for p in adapter_of(api).payloads]
    log = "\n".join(r.getMessage() for r in caplog.records)
    secrets = ("Northwind Test Co", "Recognised revenue", "Forecast revenue", "Logo churn", "Gross margin", MONTHLY)
    for secret in secrets[:-1]:
        assert secret in sent[0], f"{secret!r}: the title and the metric labels are sent (section 5)"
    for secret in secrets:
        assert secret not in stored, f"{secret!r} reached a collection"
        assert secret not in log, f"{secret!r} reached a log line"
    row = next(c for c in api.db["llm_calls"].docs if c.get("structure_type") == "file_recognition")
    assert row["status"] == "ok"


def test_no_consent_means_no_model_call_and_the_rules_recognise_the_file_alone(api):
    set_as_of(api)
    api.db["audits"].docs[0]["structure_reading_consent"] = False
    api.state["adapter"] = t.FakeAdapter(replies=[recognition_reply()])
    body = upload(api, MONTHLY, monthly_bytes()).json()
    assert adapter_of(api).calls == 0
    assert body["recognition"]["recognised"] and body["recognition"]["type_source"] == "rules"
    assert body["recognition"]["revenue_row"]["state"] == "auto", "rules alone above 80 need no click"


def test_a_failed_model_call_leaves_the_rules_outcome_and_the_file_is_still_recognised(api):
    set_as_of(api)
    api.state["adapter"] = t.FakeAdapter(replies=["not json"])
    body = upload(api, MONTHLY, monthly_bytes()).json()
    assert body["recognition"]["recognised"] and body["recognition"]["revenue_row"]["row"] == 4
    assert body["recognition"]["type_source"] == "rules"


def test_a_file_the_rules_do_not_recognise_is_never_made_a_revenue_file_by_the_model_and_the_message_names_the_missing_measure(api):
    """The rules give it no type either, so the one call of section 5.6 runs; whatever it answers, recognition is Python's."""
    set_as_of(api)
    api.state["adapter"] = t.FakeAdapter(replies=[recognition_reply(non_revenue_rows=[])])         # the model claims a monthly revenue report
    body = upload(api, "bookings_pipeline_only.csv", (FIXTURES / "bookings_pipeline_only.csv").read_bytes()).json()
    assert body["status"] == "unknown_type" and body["recognition"]["recognised"] is False
    assert body["recognition"]["missing"] == ["revenue_measure"] and "revenue measure" in body["recognition"]["message"]
    assert len(recognition_calls(api)) == 1 and api.db["datasets"].docs == [], "nothing is stored"
    assert set(body["types"]) >= {"revenue", "crm", "pnl"}, "the S6 type picker stays as the fallback"


def test_without_consent_a_file_the_rules_do_not_recognise_makes_no_call_and_shows_the_same_message(api):
    api.db["audits"].docs[0]["structure_reading_consent"] = False
    body = upload(api, "bookings_pipeline_only.csv", (FIXTURES / "bookings_pipeline_only.csv").read_bytes()).json()
    assert body["status"] == "unknown_type" and "revenue measure" in body["recognition"]["message"]
    assert adapter_of(api).calls == 0


def test_a_recognition_call_counts_toward_the_tester_allowance_and_stops_at_100_percent(api):
    from app import access
    from test_usage_allowance import ledger, own, spent
    own(api.db)
    set_as_of(api)
    api.state["adapter"] = t.FakeAdapter(replies=[recognition_reply()])
    upload(api, MONTHLY, monthly_bytes())
    total = ledger(api.db)["all"]
    assert total["files"] == 1 and total["calls"] == 1 and total["cost_usd"] > 0
    api.db[access.COLLECTION].docs[0].update(spent(10))
    calls = adapter_of(api).calls
    api.db["datasets"].docs.clear()          # as the tests above: a second revenue file with other bytes needs `replace` otherwise
    body = upload(api, "monthly_report_no_currency.csv", (FIXTURES / "monthly_report_no_currency.csv").read_bytes()).json()
    assert adapter_of(api).calls == calls, "no provider call at 100%"
    assert body["recognition"]["recognised"] and api.db["llm_calls"].docs[-1]["status"] == "monthly_allowance_reached"


# ---------------------------------------------------------------------------
# The analyst: Correct / Change (section 7)
# ---------------------------------------------------------------------------
def decide(api, dtype="revenue", **body):
    return api.post(f"/api/audits/{AUDIT}/datasets/{dtype}/recognition", json=body)


def test_correct_confirms_a_part_and_change_chooses_among_the_candidates(api):
    set_as_of(api)
    api.state["adapter"] = t.FakeAdapter(replies=[recognition_reply(revenue_row={"row": 5, "label": "MRR"})])
    got = upload(api, MONTHLY, monthly_bytes()).json()["recognition"]
    assert (got["revenue_row"]["state"], got["revenue_row"]["reason"]) == ("click", "rules_and_model_disagree")
    assert {c["row"] for c in got["revenue_row"]["candidates"]} == {4, 5, 6}
    changed = decide(api, action="change", part="revenue_row", row=4)
    assert changed.status_code == 200, changed.text
    assert (changed.json()["recognition"]["revenue_row"]["row"], changed.json()["recognition"]["revenue_row"]["state"]) == (4, "you")
    assert decide(api, action="confirm", part="last_actual").status_code == 422, "the last actual column is no decision any more"
    stored = api.get(f"/api/audits/{AUDIT}/datasets").json()["datasets"]
    assert next(d for d in stored if d["dtype"] == "revenue")["recognition"]["revenue_row"]["state"] == "you", "the decision is saved"


def test_a_change_to_a_row_that_is_not_a_candidate_or_to_the_last_actual_column_is_refused(api):
    set_as_of(api)
    api.state["adapter"] = t.FakeAdapter(replies=[recognition_reply()])
    upload(api, MONTHLY, monthly_bytes())
    assert decide(api, action="change", part="revenue_row", row=7).status_code == 400, "Customers is no candidate"
    assert decide(api, action="change", part="revenue_row", row=9).status_code == 400, "Forecast revenue is never revenue"
    assert decide(api, action="change", part="last_actual", column=13).status_code == 422, "the as-of month sets the last actual column"
    assert decide(api, action="change", part="file_type", file_type="crm").status_code == 400
    assert decide(api, action="change", part="currency", currency="XXX").status_code == 400


def test_the_analyst_confirms_the_currency_when_the_file_names_none_and_the_as_of_month_sets_the_boundary(api):
    api.state["adapter"] = t.FakeAdapter(replies=[recognition_reply(
        revenue_row={"row": 3, "label": "Revenue"}, last_actual_column=13, currency=None, non_revenue_rows=[])])
    got = upload(api, "monthly_report_no_currency.csv", (FIXTURES / "monthly_report_no_currency.csv").read_bytes()).json()["recognition"]
    assert got["needs_click"] is True and got["currency"]["state"] == "confirm"
    assert (got["last_actual"]["state"], got["last_actual"]["reason"]) == ("auto", None), "the boundary waits for nothing"
    set_as_of(api, AS_OF_YEAR_END)
    final = decide(api, action="change", part="currency", currency="EUR").json()["recognition"]
    assert final["currency"] == {"code": "EUR", "state": "you"}
    assert final["last_actual"]["state"] == "auto" and final["excluded_columns"] == [], "read to December from the as-of month"


def rules_only(api, as_of=AS_OF_YEAR_END):
    """No consent: the rules recognise the file alone, so the test reads Python's outcome and makes no model call."""
    api.db["audits"].docs[0]["structure_reading_consent"] = False
    set_as_of(api, as_of)


def test_a_workbook_stored_as_the_pnl_is_recognised_through_the_pnl_path_and_another_sheet_can_be_used(api):
    rules_only(api)
    body = upload(api, "pack.xlsx", kpi_and_pnl_bytes()).json()
    assert body["dtype"] == "pnl", "a P&L stays the P&L (section 8)"
    got = body["recognition"]
    assert got["sheet"] == "P&L" and got["file_type"] == "financial_statements"
    assert got["other_sheets"] == [{"sheet": "KPIs", "file_type": "monthly_revenue_report", "recognised": True}]
    other = decide(api, dtype="pnl", action="change", part="sheet", sheet="KPIs")
    assert other.status_code == 200, other.text
    assert other.json()["recognition"]["sheet"] == "KPIs" and other.json()["recognition"]["file_type"] == "monthly_revenue_report"
    assert decide(api, dtype="pnl", action="change", part="sheet", sheet="Missing").status_code == 400
    assert decide(api, action="confirm", part="revenue_row").status_code == 404, "there is no revenue dataset: the P&L has its own recognition"


def test_a_proposal_for_a_sheet_python_did_not_pick_is_dropped(api):
    set_as_of(api, AS_OF_YEAR_END)
    api.state["adapter"] = t.FakeAdapter(replies=[recognition_reply(sheet=1, revenue_row={"row": 3, "label": "MRR"}, last_actual_column=13,
                                                                    non_revenue_rows=[])])
    got = upload(api, "pack.xlsx", kpi_and_pnl_bytes()).json()["recognition"]
    assert got["sheet"] == "P&L" and got["type_source"] == "rules"
    assert got["revenue_row"]["row"] == 3 and got["revenue_row"]["state"] == "auto", "the proposal named the KPIs sheet: it is not applied to P&L"


def blocker_kinds(api):
    return [b["kind"] for b in api.get(f"/api/audits/{AUDIT}/blockers").json()["blockers"]]


def test_a_pnl_with_two_or_more_months_satisfies_the_revenue_file_once_the_analyst_confirms_its_recognition(api):
    rules_only(api, "2025-06-30")
    body = upload(api, "pnl_monthly.csv", (FIXTURES / "pnl_monthly.csv").read_bytes()).json()
    assert body["dtype"] == "pnl" and body["recognition"]["recognised"] and body["recognition"]["periods"] == 6
    assert "revenue_file_missing" in blocker_kinds(api), "recognised is not confirmed"
    done = decide(api, dtype="pnl", action="confirm", part="revenue_row")
    assert done.status_code == 200 and done.json()["recognition"]["revenue_row"]["state"] == "you"
    assert "revenue_file_missing" not in blocker_kinds(api), "a confirmed P&L with 2+ months is the revenue file"


def test_a_pnl_with_one_period_never_satisfies_the_revenue_file(api):
    rules_only(api)
    body = upload(api, "pnl_single_period.csv", (FIXTURES / "pnl_single_period.csv").read_bytes()).json()
    assert body["dtype"] == "pnl" and body["recognition"]["periods"] == 1
    decide(api, dtype="pnl", action="confirm", part="revenue_row")
    assert "revenue_file_missing" in blocker_kinds(api)


def test_a_confirm_only_reason_is_settled_by_one_correct(api):
    api.db["audits"].docs[0]["structure_reading_consent"] = False        # no as-of month stored, no P&L: the boundary waits for nothing
    text = "Metric,Jan-25,Feb-25,Mar-25,Apr-25\nRevenue - EMEA (EUR),100,110,120,130\nCustomers,4,5,6,7\nChurn,1,2,3,4\nCash,5,6,7,8\n"
    got = upload(api, "emea.csv", text.encode()).json()["recognition"]
    assert (got["revenue_row"]["state"], got["revenue_row"]["reason"]) == ("click", "confidence_80")
    assert (got["last_actual"]["state"], got["last_actual"]["reason"]) == ("auto", None)
    done = decide(api, action="confirm", part="pending")
    assert done.status_code == 200, done.text
    after = done.json()["recognition"]
    assert (after["revenue_row"]["state"], after["last_actual"]["state"]) == ("you", "auto") and after["needs_click"] is False


def test_one_correct_does_not_settle_a_part_that_needs_a_choice(api):
    set_as_of(api)
    api.state["adapter"] = t.FakeAdapter(replies=[recognition_reply(revenue_row={"row": 5, "label": "MRR"})])
    upload(api, MONTHLY, monthly_bytes())
    after = decide(api, action="confirm", part="pending").json()["recognition"]
    assert (after["revenue_row"]["state"], after["revenue_row"]["reason"]) == ("click", "rules_and_model_disagree")


def test_a_new_prompt_version_asks_again_for_the_same_bytes(api):
    set_as_of(api)
    api.state["adapter"] = t.FakeAdapter(replies=[recognition_reply()])
    upload(api, MONTHLY, monthly_bytes())
    api.db["datasets"].docs.clear()
    api.db["prompts"].docs.append({"name": "file_type", "version": "v2", "text": "<!-- TEST ONLY --> v2 stub", "uploaded_at": "2026-02-01T00:00:00+00:00",
                                   "sha256": "0" * 64})
    upload(api, MONTHLY, monthly_bytes())
    assert len(recognition_calls(api)) == 2, "the cached answer belongs to the prompt version that gave it"
    tags = {d["prompt_tag"] for d in api.db["column_mappings"].docs if d.get("kind") == "file_recognition"}
    assert len(tags) == 2, "one kept answer per prompt version"


def test_an_entity_labelled_upload_stores_and_logs_no_customer_name(api, caplog):
    caplog.set_level(logging.DEBUG)
    set_as_of(api, AS_OF_YEAR_END)
    api.state["adapter"] = t.FakeAdapter(replies=[recognition_reply(file_type="customer_revenue", revenue_row=None, last_actual_column=13,
                                                                    non_revenue_rows=[])])
    upload(api, "customer_revenue_months.csv", (FIXTURES / "customer_revenue_months.csv").read_bytes())
    decide(api, action="confirm", part="last_actual")
    sent = " ".join(json.loads(p)["text"] for p in recognition_calls(api))
    stored = json.dumps([api.db[n].docs for n in ("llm_structures", "llm_calls", "column_mappings", "audits")], default=str)
    log = "\n".join(r.getMessage() for r in caplog.records)
    names = ("Alder Works", "Birch Labs", "Cedar Foods", "Dune Print", "Elm Studio", "Fern Tools", "Gorse Yard", "Heath Pumps", "Iris Optics",
             "Juniper Mills", "Margin Foods", "Larch Freight")
    for name in names:
        assert name not in sent, f"{name!r} reached the model"
        assert name not in stored, f"{name!r} reached a collection"
        assert name not in log, f"{name!r} reached a log line"
    assert "Row 1" in sent


def test_the_decision_endpoint_logs_codes_only(api, caplog):
    caplog.set_level(logging.DEBUG)
    set_as_of(api)
    api.state["adapter"] = t.FakeAdapter(replies=[recognition_reply()])
    upload(api, MONTHLY, monthly_bytes())
    caplog.clear()
    assert decide(api, action="change", part="revenue_row", row=5).status_code == 200
    log = "\n".join(r.getMessage() for r in caplog.records)
    for secret in ("Recognised revenue", "MRR", "Forecast revenue", "Northwind", MONTHLY):
        assert secret not in log, f"{secret!r} reached a log line"


def test_the_analyst_answers_the_turnover_question_and_the_row_is_then_decided(api):
    set_as_of(api, "2025-12-31")
    text = ("Metric,Jan-25,Feb-25,Mar-25,Apr-25,May-25,Jun-25,Jul-25,Aug-25,Sep-25,Oct-25,Nov-25,Dec-25\n"
            "Turnover (EUR),11,12,13,14,15,16,17,18,19,20,21,22\nCustomers,4,5,6,7,8,9,10,11,12,13,14,15\n"
            "Churn,1,2,3,4,5,6,7,8,9,10,11,12\nCash,5,6,7,8,9,10,11,12,13,14,15,16\n")
    api.state["adapter"] = t.FakeAdapter(replies=[recognition_reply(revenue_row={"row": 2, "label": "Turnover (EUR)"}, last_actual_column=13,
                                                                    non_revenue_rows=[])])
    got = upload(api, "turnover.csv", text.encode()).json()["recognition"]
    assert (got["revenue_row"]["state"], got["revenue_row"]["reason"], got["money_kind"]) == ("click", "turnover_question", {"asked": True})
    done = decide(api, action="change", part="money_kind", money_kind="revenue").json()["recognition"]
    assert done["revenue_row"]["state"] == "you" and done["money_kind"] == {"asked": True, "answer": "revenue"}
    assert decide(api, action="change", part="money_kind", money_kind="maybe").status_code == 400


def test_delete_audit_removes_the_kept_recognition_and_the_decisions(api):
    from test_chat_upload import COMPANY
    set_as_of(api)
    api.state["adapter"] = t.FakeAdapter(replies=[recognition_reply()])
    upload(api, MONTHLY, monthly_bytes())
    decide(api, action="confirm", part="revenue_row")
    assert any(d.get("kind") == "file_recognition" for d in api.db["column_mappings"].docs)
    assert api.request("DELETE", f"/api/audits/{AUDIT}", json={"confirm": COMPANY}).status_code == 200
    for name, col in api.db._cols.items():
        assert not [d for d in col.docs if AUDIT in json.dumps(d, default=str)], f"{name} still holds a document of the audit"


# ---------------------------------------------------------------------------
# Run audit without a revenue file (follow-up of the specification PR)
# ---------------------------------------------------------------------------
def run_stage(api, stage):
    return api.post(f"/api/audits/{AUDIT}/run/{stage}").json()


def test_run_audit_without_a_revenue_file_stops_at_the_mapping_stage_with_the_missing_file_message(api):
    rules_only(api)
    assert run_stage(api, "mapping")["message"] == "No file has been uploaded yet. Add the revenue file and run the audit again."
    body = upload(api, "pnl_multi_period.csv", (FIXTURES / "pnl_multi_period.csv").read_bytes()).json()
    assert body["dtype"] == "pnl" and body["recognition"]["recognised"]
    first = run_stage(api, "mapping")
    assert (first["status"], first["message"]) == ("failed", "The revenue file is missing. Add it and run the audit again."), \
        "a recognised P&L that the analyst has not confirmed is not a revenue file"
    assert run_stage(api, "engine")["status"] == "waiting", "nothing is calculated without a revenue file"


def test_a_confirmed_pnl_lets_run_audit_past_the_mapping_stage_and_the_engine_reads_its_revenue(api):
    rules_only(api, "2025-06-30")
    upload(api, "pnl_monthly.csv", (FIXTURES / "pnl_monthly.csv").read_bytes())
    assert decide(api, dtype="pnl", action="confirm", part="revenue_row").status_code == 200
    assert run_stage(api, "mapping")["status"] == "done"
    assert run_stage(api, "engine")["status"] == "done"
    results = api.db["audits"].docs[0]["results"]
    assert results["revenue_reconciliation"]["status"] == "not_applicable", "the P&L is the only revenue source"
    assert results["recognition"]["source"] == "pnl" and results["recognition"]["file_type"] == "financial_statements"


# ---------------------------------------------------------------------------
# A missing currency asks the analyst to confirm it (owner request of 19 Oct 2026)
# ---------------------------------------------------------------------------
NO_CURRENCY_COLUMN = ("Customer,Invoice Date,Amount\n" + "".join(
    f"Customer {c},2025-{m:02d}-05,{1000 + 10 * c + m}\n" for m in range(1, 7) for c in range(1, 5))).encode()


def test_a_file_with_no_currency_column_asks_to_confirm_the_currency_and_never_says_a_required_field_is_not_mapped(api):
    rules_only(api, "2025-06-30")
    body = upload(api, "no_currency.csv", NO_CURRENCY_COLUMN).json()
    assert body["dtype"] == "revenue" and body["missing_required"] == [] and body["currency_pending"] is True
    assert body["recognition"]["currency"] == {"code": None, "state": "confirm"}
    assert run_stage(api, "mapping")["status"] == "needs_review", "the audit does not run on a currency nobody confirmed"
    done = decide(api, action="change", part="currency", currency="EUR").json()
    assert done["mapping"]["currency"] == "Currency (confirmed)" and not done.get("currency_pending")
    assert done["recognition"]["currency"] == {"code": "EUR", "state": "you"}
    assert run_stage(api, "mapping")["status"] == "done" and run_stage(api, "engine")["status"] == "done"
    results = api.db["audits"].docs[0]["results"]
    assert not any("unmapped currency" in m["metric"] for m in results["missing_data"]), "every row took the confirmed currency"
    assert results["revenue_series"]["data"], "the confirmed currency reached the engine"


def test_a_monthly_revenue_report_needs_no_customer_id_column_and_holds_no_required_field(api):
    rules_only(api, AS_OF)
    body = upload(api, MONTHLY, monthly_bytes()).json()
    assert body["dtype"] == "revenue" and body["recognition_only"] is True
    assert body["missing_required"] == [] and body["miss"] is None and body["columns"] == [] and body["fields"]["required"] == []
    assert "customer_id" not in body["mapping"] and body["pending"] == 0 and body["saved"] is True
    monthly_summary = b"Revenue by product line (EUR),,,\nProduct line,Jan-25,Feb-25,Mar-25\nPlatform,10,11,12\nSupport,20,21,22\nTraining,30,31,32\n"
    summary = upload(api, "revenue_summary_monthly.csv", monthly_summary, replace="true").json()
    assert summary["recognition"]["file_type"] == "revenue_summary" and summary["missing_required"] == [] and summary["fields"]["required"] == []


# ---------------------------------------------------------------------------
# Section 6.7: a sheet with a metrics block and a dated list (owner request of 19 Oct 2026)
# ---------------------------------------------------------------------------
REVENUE_ROW = [12010, 12130, 12250, 12370, 12490, 12610]
STACKED, SIDE = "two_tables_stacked_match.csv", "two_tables_side_mismatch.csv"


def test_a_sheet_with_a_metrics_block_and_a_dated_list_detects_both_and_classifies_each(rec):
    for name in (STACKED, SIDE):
        got = recognise(rec, name)
        assert got["recognised"] and got["label_mode"] == "metric", f"{name}: the rows of the list are no labels of the block"
        assert [t["kind"] for t in got["tables"]] == ["metrics_block", "dated_list"], name
        assert [t["file_type"] for t in got["tables"]] == ["monthly_revenue_report", "monthly_revenue_report"], name
        assert got["tables"][1]["months"] == 6
        assert (got["revenue_row"]["row"], got["revenue_row"]["label"]) == (3, "Recognised revenue"), name


def test_a_dated_list_that_sums_to_the_recognised_revenue_row_month_by_month_is_a_cross_check_and_not_a_second_source(rec):
    got = recognise(rec, STACKED)
    assert got["series_source"]["supplier"] == "metrics_block" and got["series_source"]["state"] == "auto"
    assert got["series_source"]["cross_check"] == {"status": "match", "months": 6, "months_off": 0, "max_gap_pct": 0}
    assert got["series_basis"] == "row" and [s["value"] for s in got["series"]] == REVENUE_ROW
    assert got["needs_click"] is False and rec.choices_waiting(got) == 0, "a match asks the analyst nothing"


def test_a_dated_list_within_one_percent_of_the_row_still_matches_and_one_over_does_not(rec):
    near = STACKED_TEXT.replace("2025-01-10,9010", "2025-01-10,9100")           # Jan: 12,100 against 12,010 is 0.7% over
    far = STACKED_TEXT.replace("2025-01-10,9010", "2025-01-10,9300")            # Jan: 12,300 against 12,010 is 2.4% over
    assert rec.recognise(csv_sheets(near), as_of_month=AS_OF_YEAR_END)["series_source"]["cross_check"]["status"] == "match"
    off = rec.recognise(csv_sheets(far), as_of_month=AS_OF_YEAR_END)["series_source"]
    assert (off["cross_check"]["status"], off["cross_check"]["months_off"], off["state"]) == ("mismatch", 1, "click")


STACKED_TEXT = (FIXTURES / STACKED).read_text()


def test_a_dated_list_that_does_not_match_leaves_the_choice_to_the_analyst_until_it_is_made(rec):
    got = recognise(rec, SIDE)
    source = got["series_source"]
    assert (source["supplier"], source["state"]) == ("metrics_block", "click")
    assert source["cross_check"]["status"] == "mismatch" and source["cross_check"]["months_off"] == 2
    assert source["cross_check"]["max_gap_pct"] == pytest.approx(0.05, abs=1e-4)
    assert got["needs_click"] is True and rec.choices_waiting(got) == 1
    assert [s["value"] for s in got["series"]] == REVENUE_ROW, "until the choice, the revenue row is the series"


@pytest.mark.parametrize("choice, basis, first", [("metrics_block", "row", 12010), ("dated_list", "column", 12010)])
def test_the_analysts_choice_decides_which_table_supplies_the_series(rec, choice, basis, first):
    got = rec.recognise(sheets_of(SIDE), as_of_month=AS_OF_YEAR_END, decisions={"series_source": choice})
    assert (got["series_source"]["supplier"], got["series_source"]["state"]) == (choice, "you")
    assert got["series_basis"] == basis and got["needs_click"] is False
    assert got["series"][0]["value"] == first and got["series"][0]["month"] == "2025-01"
    values = [s["value"] for s in got["series"]]
    assert values == (REVENUE_ROW if choice == "metrics_block" else [12010, 12615, 12250, 12370, 11865, 12610])
    if choice == "dated_list":
        assert got["series_rows"] == [3, 4, 5, 6, 7, 8], "the citation names the rows of the list"


def test_a_row_that_is_not_recognised_revenue_never_wins_by_itself_even_when_the_list_matches(rec):
    text = STACKED_TEXT.replace("Recognised revenue", "MRR")
    got = rec.recognise(csv_sheets(text), as_of_month=AS_OF_YEAR_END)
    assert got["revenue_row"]["kind"] == "recurring" and got["series_source"]["cross_check"]["status"] == "match"
    assert got["series_source"]["state"] == "click", "an MRR row does not win: the analyst chooses"


def test_a_sheet_with_one_table_has_no_series_source_and_no_tables(rec):
    got = recognise(rec, MONTHLY, as_of=AS_OF)
    assert got["tables"] == [] and got["series_source"] is None


def test_the_list_beside_the_block_is_sent_without_its_dates_as_labels(rec):
    """The block's labels go to the model as written; the dated list is not a label column (rule 14, section 5.3)."""
    text = rec.model_text(sheets_of(SIDE), company_name="Fixture Target Ltd", mapping={}, withheld=(), file_hash="0" * 64)
    labels = [line.split(": ", 1)[1] for line in text.split("\n") if " label r" in line]
    assert labels == ["Recognised revenue", "Customers", "Logo churn %"]


def test_the_upload_stores_both_tables_and_the_choice_travels_to_the_engine_and_the_memo(api):
    rules_only(api)
    body = upload(api, SIDE, (FIXTURES / SIDE).read_bytes()).json()
    got = body["recognition"]
    assert [t["kind"] for t in got["tables"]] == ["metrics_block", "dated_list"] and got["series_source"]["state"] == "click"
    assert body["pending"] == 1 and run_stage(api, "mapping")["status"] == "needs_review", "the audit waits for the choice"
    assert decide(api, part="series_source", action="change", series_source="maybe").status_code == 400
    done = decide(api, action="change", part="series_source", series_source="dated_list").json()["recognition"]
    assert done["series_source"]["supplier"] == "dated_list" and done["needs_click"] is False
    assert run_stage(api, "mapping")["status"] == "done" and run_stage(api, "engine")["status"] == "done"
    results = api.db["audits"].docs[0]["results"]
    assert results["recognition"]["series_source"] == "dated_list" and results["recognition"]["cross_check"] == "mismatch"
    assert [d["total"] for d in results["revenue_series"]["data"]] == [12010, 12615, 12250, 12370, 11865, 12610]
    assert results["revenue_series"]["basis"] == "column"


# ---------------------------------------------------------------------------
# The stored sheets: the whole sheet within the dataset size limit, never cut (owner decision on PR 135)
# ---------------------------------------------------------------------------
def test_a_sheet_over_the_size_limit_is_refused_with_a_message_and_nothing_is_stored_or_cut(api, monkeypatch):
    import server
    rules_only(api)
    monkeypatch.setattr(server, "DATASET_MAX_BYTES", 2000)
    for name, content in ((MONTHLY, monthly_bytes()), ("pnl_multi_period.csv", (FIXTURES / "pnl_multi_period.csv").read_bytes())):
        res = upload(api, name, content)
        assert res.status_code == 413 and res.json()["detail"] == server.DATASET_TOO_LARGE, name
    assert api.db["datasets"].docs == [], "nothing is stored"


def test_a_long_sheet_under_the_limit_is_kept_whole(api):
    rules_only(api, "2025-12-31")
    rows = "".join(f"Customer {i},{10000 + i},{20000 + i},{30000 + i}\n" for i in range(2600))
    text = "Revenue by customer (EUR),,,\nCustomer,Jan-25,Feb-25,Mar-25\n" + rows
    assert upload(api, "long.csv", text.encode()).status_code == 200
    kept = api.db["datasets"].docs[0]["recognition_sheets"][0]["grid"]
    assert len(kept) == 2602, "every row of the sheet is kept"


METRICS_SUMMARY = """Test Co - metrics,,,
Metric,Jan-25,Feb-25,Mar-25
Product A revenue (EUR),100,110,120
Product B revenue (EUR),50,60,70
MRR,900,910,920
ARR,10800,10920,11040
Customers,40,41,42
Logo churn %,0.05,0.04,0.03
Gross margin %,0.6,0.61,0.62
Cost of sales,30,31,32
"""


def test_a_revenue_summary_from_a_metrics_sheet_reads_the_revenue_row_and_never_sums_rows(rec):
    sheets = csv_sheets(METRICS_SUMMARY, "metrics.csv")
    got = rec.recognise(sheets, as_of_month="2025-03-31", decisions={"file_type": "revenue_summary"})
    assert got["file_type"] == "revenue_summary" and got["label_mode"] == "metric" and got["series_basis"] == "row"
    assert got["revenue_row"]["state"] == "click" and got["revenue_row"]["reason"] == "tie", "two revenue rows: the analyst chooses"
    assert [s["value"] for s in got["series"]] == [100, 110, 120] and got["series_rows"] == [3], "never MRR, ARR, customers, % or cost rows"
    chosen = rec.recognise(sheets, as_of_month="2025-03-31", decisions={"file_type": "revenue_summary", "revenue_row": 4})
    assert [s["value"] for s in chosen["series"]] == [50, 60, 70] and chosen["series_rows"] == [4]