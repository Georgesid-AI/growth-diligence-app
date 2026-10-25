"""Cost is hidden from testers (docs/specs/usage-and-allowance.md section 9): no response a tester token receives holds a cost, token
or USD field, on any audit route, and the memo and the export hold no usage table. The owner routes keep the figures. The app still
stores them. Every GET route of the audit is read from the app's own router, so a route added later is covered. No network."""
import hashlib
import io
import re

import pytest

pytest.importorskip("fastapi")
pytest.importorskip("motor")

from app import access, tester_view  # noqa: E402
from test_access_links import audit_routes, fill  # noqa: E402
from test_run_audit import AUDIT, THESIS, ready_for_review, write_all  # noqa: E402

OWNER = "owner-token-for-the-test-suite-only"
TESTER_TOKEN = "tester-token-for-the-test-suite-only"
# Keys that contain the words but are not an AI figure: a P&L line, and the uploaded file's own column headers (a preview is keyed by them).
NOT_AI = {"costs", "cost_of_revenue"}
DATA_KEYED = ("/preview",)
WORDS = re.compile(r"(?i)cost|usd|token")


def keys_of(value, path=""):
    if isinstance(value, dict):
        for k, v in value.items():
            yield f"{path}/{k}", str(k)
            yield from keys_of(v, f"{path}/{k}")
    elif isinstance(value, list):
        for v in value:
            yield from keys_of(v, path)


def offending(value):
    return sorted({p for p, k in keys_of(value) if WORDS.search(k) and k not in NOT_AI and not p.startswith(DATA_KEYED) and not any(d in p for d in DATA_KEYED)})


@pytest.fixture()
def as_tester(run_api, monkeypatch):
    """A computed audit with narratives written (so llm_calls hold cost and tokens), then the real access guard on, as a tester."""
    ready_for_review(run_api)
    write_all(run_api)
    assert run_api.db["llm_calls"].docs and all("estimated_cost_usd" in c for c in run_api.db["llm_calls"].docs), "the cost is stored"
    monkeypatch.delenv("ACCESS_DISABLED", raising=False)
    monkeypatch.setenv("OWNER_TOKEN_SHA256", hashlib.sha256(OWNER.encode()).hexdigest())
    run_api.db[access.COLLECTION].docs.append({
        "id": "tester-1", "label": "Tester One", "hash": access.digest(TESTER_TOKEN), "created_at": "2026-10-01T00:00:00+00:00",
        "expires_at": "2099-01-01T00:00:00+00:00", "revoked_at": None})
    run_api.db["audits"].docs[0]["owner"] = "tester-1"
    run_api.headers = {"Authorization": f"Bearer {TESTER_TOKEN}"}
    return run_api


def test_no_tester_response_on_any_audit_route_holds_a_cost_token_or_usd_field(as_tester):
    seen = 0
    for method, path in audit_routes():
        if method != "GET":
            continue
        r = as_tester.get(fill(path, AUDIT), headers=as_tester.headers)
        if r.headers.get("content-type", "").startswith("application/json"):
            seen += 1
            assert offending(r.json()) == [], (path, offending(r.json()))
    assert seen > 10, "the router walk reached the JSON routes"


def test_the_generate_narrative_the_estimate_the_usage_and_the_allowance_routes_hold_no_cost_for_a_tester(as_tester):
    h = as_tester.headers
    for step in ("growth_engine", "data_inventory", "data_integrity"):
        assert offending(as_tester.get(f"/api/runs/{AUDIT}/narrative/{step}", headers=h).json()) == []
        r = as_tester.post(f"/api/runs/{AUDIT}/narrative/{step}", headers=h)
        assert r.status_code == 200 and offending(r.json()) == []
    usage = as_tester.get(f"/api/runs/{AUDIT}/llm-usage", headers=h).json()
    assert offending(usage) == [] and usage["calls"] > 0 and "by_step" in usage, "calls and cache hits stay"
    estimate = as_tester.get(f"/api/runs/{AUDIT}/narratives/estimate", headers=h).json()
    assert offending(estimate) == [] and all("cached" in s for s in estimate["steps"])
    as_tester.db[access.COLLECTION].docs[0]["ledger"] = {"all": {"cost_usd": 8.5}, "2026-10": {"cost_usd": 8.5, "calls": 3, "tokens": 90}}
    assert as_tester.get("/api/allowance", headers=h).json().keys() == {"percent", "state"}
    decks = as_tester.get(f"/api/audits/{AUDIT}/decks", headers=h).json()
    assert decks["decks"] and offending(decks) == [] and all("ai_cost_usd" not in d for d in decks["decks"])


def test_the_memo_and_the_xlsx_export_carry_no_usage_table(as_tester):
    h = as_tester.headers
    ids = as_tester.get(f"/api/audits/{AUDIT}/verdict", headers=h).json()["verdict"]["top5"]["proposed"]
    for body in ({"top5": ids}, {"thesis": THESIS, "note": "A note."}):
        as_tester.put(f"/api/audits/{AUDIT}/ic-inputs", json=body, headers=h)
    memo = as_tester.get(f"/api/audits/{AUDIT}/memo.md", headers=h)
    assert memo.status_code == 200, "the memo needs no rating and no judgment"
    if memo.status_code == 200:
        assert not re.search(r"(?i)AI usage|Cost \(USD\)|Input tokens|Output tokens", memo.text)
    export = as_tester.get(f"/api/audits/{AUDIT}/export", headers=h)
    assert export.status_code == 200
    from openpyxl import load_workbook
    book = load_workbook(io.BytesIO(export.content))
    text = " ".join(str(c.value) for sheet in book for row in sheet.iter_rows() for c in row if c.value is not None) + " ".join(book.sheetnames)
    assert not re.search(r"(?i)AI usage|Cost \(USD\)|Input tokens|Output tokens|estimated_cost", text)


def test_the_owner_routes_keep_the_figures_and_the_filter_removes_a_key_at_any_depth(as_tester):
    owner = {"Authorization": f"Bearer {OWNER}"}
    costs = as_tester.get("/api/admin/audit-costs", headers=owner).json()
    assert costs["audits"] and {"cost_usd", "tokens", "calls"} <= set(costs["audits"][0])
    assert "cost_usd" in as_tester.get("/api/admin/usage", headers=owner).json()["totals"]["all_time"]
    assert tester_view.strip({"a": [{"input_tokens": 1, "x": {"estimated_cost_usd": 2, "keep": 3}}], "total_usd": 4}) == {"a": [{"x": {"keep": 3}}]}
    assert tester_view.strip_body(b"not json") == b"not json"
