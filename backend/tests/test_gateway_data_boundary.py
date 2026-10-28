"""The gateway's data boundary: what may leave the server, and what never does.

The outbound payload is built from an allowlist. File names, column
headers and raw cell values stay in Mongo for the dashboard; segment names go
out as stable labels and come back as real names. Every test runs against the
in-memory Mongo stub and the fake adapter - no test can reach a real provider.
"""
import asyncio
import copy
import inspect
import json
import os
import sys
from pathlib import Path

import pytest

BACKEND = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(BACKEND))
sys.path.insert(0, str(BACKEND / "tests"))

from app.llm import gateway, prompt_store, redaction  # noqa: E402
import contract_fixtures as cf  # noqa: E402
import page_images  # noqa: E402
import test_llm_gateway as t  # noqa: E402

RUN_ID = t.RUN_ID
FILE = "Revenue_Lines.xlsx"
SHEET = "Billing"
PERSON = "Jane Doe (CEO)"                     # a raw cell value from the founder-involved column
HEADERS = ("Deal Opened On", "Signed Date")   # column headers in the upload

SOURCE = {"file": FILE, "sheet": SHEET, "rows": "rows 2–40 (39 rows)", "row_numbers": [2, 3, 4],
          "rule": "ARR = current-month recurring MRR × 12"}

RESULTS = cf.stored(
    as_of_month="2026-12",
    arr=cf.arr(3129600.0, 260800.0, month="2026-12", source=SOURCE),
    nrr=cf.nrr(1.04, 110, month="2026-12",
               by_segment={"Enterprise": cf.nrr_group(1.12, 40), "Mid-Market": cf.nrr_group(0.96, 70)},
               source={**SOURCE, "rule": "NRR = base-cohort MRR now ÷ MRR 12 months ago"}),
    win_rate=cf.win_rate(0.30, 30, 70,
                         founder_involved_excluded={"count": 1, "rows": [7], "values": [PERSON]},
                         source={"file": "CRM_Deals.csv", "sheet": "Deals", "rows": "no rows", "row_numbers": [],
                                 "rule": "Win rate = won ÷ (won + lost)"}),
    sales_cycle=cf.sales_cycle(45, 30, status="Computed - management to explain",
                               source={"file": FILE, "sheet": SHEET, "rows": "no rows", "row_numbers": [],
                                       "rule": "Median days from created to close", "dataset": "revenue",
                                       "columns": {"created_date": HEADERS[0], "close_date": HEADERS[1]}}),
    segment_paths=cf.segment_paths(
        available=True, source={**SOURCE, "rule": "Segment paths = start ARR at the current NRR"},
        stage_one={"segments": {"Enterprise": cf.seg_base(2000000.0, 40), "Mid-Market": cf.seg_base(1129600.0, 70)},
                   "start_arr_total": 3129600.0},
        reverse_solve={"12": {"window_months": 12, "computable": True, "reachable": None, "best_segment": "Enterprise",
                              "reason": "no landed ACV for Mid-Market"}}),
    cohort_retention={"cohorts": [], "max_offset": 0, "data": [],
                      "source": {**SOURCE, "rule": "Retention = cohort MRR ÷ starting MRR"}},
    new_mrr_by_quarter_source=SOURCE,
    missing_data=[{"metric": "CRM rows with unrecognized founder-involved value", "status": "Missing",
                   "reason": f"1 row(s) have a founder-involved value that isn't yes/no-like ({PERSON})",
                   "unlocked_by": "Use a yes/no style value", "file": "CRM_Deals.csv"}],
    questions_for_management=[{
        "metric": "Sales cycle", "status": "Computed - management to explain", "result_key": "sales_cycle",
        "dataset": "revenue", "file": FILE,
        "columns": {"created_date": HEADERS[0], "close_date": HEADERS[1]},
        "question": (f"Sales cycle was computed from the revenue upload (created_date = '{HEADERS[0]}', "
                     f"close_date = '{HEADERS[1]}') because it was not available from the crm upload."),
    }],
)

NARRATIVE = {
    "headline": "Segment A retains 112% of its revenue over twelve months.",
    "what_this_means": "Segment A expands while Segment B contracts.",
    "table_rows": [{"label": "Segment A NRR", "value": "112%",
                    "source_key": "metrics.nrr.by_segment.Segment A.nrr_pct"}],
    "worth_flagging": ["Segment B is at 96%."],
    "next_actions": ["Ask why Segment B churns."],
    "source_keys": ["metrics.nrr.by_segment"],
}


def _db():
    db = t.FakeDB()
    doc = copy.deepcopy(t.RESULTS_DOC)
    doc["results"] = copy.deepcopy(RESULTS)
    db["audits"].docs.append(doc)
    return db


def _generate(db, reply=NARRATIVE):
    adapter = t.FakeAdapter(replies=[json.dumps(reply)])
    result = asyncio.run(gateway.generate_narrative(db, RUN_ID, "growth_engine", adapter=adapter, sleep=t._noop_sleep))
    assert adapter.calls == 1
    return result, json.loads(adapter.payloads[0])


def _keys(node):
    if isinstance(node, dict):
        for k, v in node.items():
            yield k
            yield from _keys(v)
    elif isinstance(node, list):
        for i in node:
            yield from _keys(i)


def test_outbound_payload_has_no_file_sheet_header_or_raw_cell_value():
    db = _db()
    _, sent = _generate(db)
    blob = json.dumps(sent, ensure_ascii=False)
    for secret in (FILE, "CRM_Deals.csv", SHEET, "Deals", *HEADERS, PERSON, "Jane Doe"):
        assert secret not in blob, f"{secret!r} reached the provider"
    keys = set(_keys(sent))
    for dropped in ("file", "sheet", "columns", "rows", "row_numbers", "unlocked_by", "run_id", "computed_at"):
        assert dropped not in keys, dropped
    assert "values" not in sent["metrics"]["win_rate"]["founder_involved_excluded"]
    # What the model needs is still there.
    assert sent["metrics"]["arr"]["source"] == {"rule": "ARR = current-month recurring MRR × 12"}
    # The citations added to segment_paths and cohort_retention go out as the rule only, like every other block's.
    assert sent["metrics"]["segment_paths"]["source"] == {"rule": "Segment paths = start ARR at the current NRR"}
    assert not {"new_mrr_by_quarter_source", "mrr_series", "anomalies"} & set(sent["metrics"])
    assert sent["metrics"]["win_rate"]["founder_involved_excluded"]["count"] == "1"
    question = sent["metrics"]["questions_for_management"][0]["question"]
    assert "close_date, created_date" in question and "revenue upload" in question
    assert sent["metrics"]["missing_data"][0]["reason"] == (
        "1 row(s) have a founder_involved value that isn't yes/no-like — excluded from the founder split, not guessed")


def test_the_cohort_retention_citation_goes_out_as_the_rule_only():
    out = gateway.build_outbound({"metrics": gateway._slice_for_step(RESULTS, "cohort_retention")}, {})
    assert out["metrics"]["cohort_retention"]["source"] == {"rule": "Retention = cohort MRR ÷ starting MRR"}
    assert FILE not in json.dumps(out, ensure_ascii=False)


def test_full_versions_stay_in_mongo_for_the_dashboard():
    db = _db()
    _generate(db)
    stored = db["audits"].docs[0]["results"]
    assert stored == RESULTS
    assert HEADERS[0] in stored["questions_for_management"][0]["question"]


def test_segment_names_are_labelled_outbound_and_restored_in_the_reply():
    db = _db()
    result, sent = _generate(db)
    blob = json.dumps(sent)
    assert "Enterprise" not in blob and "Mid-Market" not in blob
    assert set(sent["metrics"]["nrr"]["by_segment"]) == {"Segment A", "Segment B"}
    assert set(sent["metrics"]["segment_paths"]["stage_one"]["segments"]) == {"Segment A", "Segment B"}
    assert sent["metrics"]["segment_paths"]["reverse_solve"]["12"]["best_segment"] == "Segment A"
    assert sent["metrics"]["segment_paths"]["reverse_solve"]["12"]["reason"] == "no landed ACV for Segment B"

    assert result.narrative_status == "ok", result.reason
    n = result.narrative
    assert n.headline == "Enterprise retains 112% of its revenue over twelve months."
    assert n.what_this_means == "Enterprise expands while Mid-Market contracts."
    assert n.table_rows[0].label == "Enterprise NRR"
    assert n.table_rows[0].source_key == "metrics.nrr.by_segment.Enterprise.nrr_pct"
    assert n.worth_flagging == ["Mid-Market is at 96%."]


def test_segment_labels_are_stable_for_a_run():
    db = _db()
    computed = asyncio.run(gateway.load_computed_results(db, RUN_ID, "growth_engine"))
    first = asyncio.run(redaction.get_or_create_map(db, RUN_ID, computed))
    second = asyncio.run(redaction.get_or_create_map(db, RUN_ID, computed))
    assert first == second
    assert first["Enterprise"] == "Segment A" and first["Mid-Market"] == "Segment B"


def test_identifier_12_never_changes_a_date_or_a_number():
    mapping = {"12": "Customer_01", "Acme": "Customer_02"}
    payload = {"month": "2026-12", "target_date": "2027-12-31", "trailing_window_months": "12",
               "value": "1,212 EUR", "nrr_pct": "112%", "ratio": "1.12x", "window": "12-month NRR",
               "reason": "Acme churned in 2026-12", "series": ["12", "2026-12"], "n": 12}
    out = redaction.redact(payload, mapping)
    assert out == {**payload, "reason": "Customer_02 churned in 2026-12"}
    assert redaction.substitute("2026-12", mapping) == "2026-12"
    assert redaction.find_leaks(out, mapping) == []


def test_field_name_keys_are_never_rewritten():
    mapping = {"arr": "Customer_01", "value": "Customer_02"}
    payload = {"arr": {"value": "1 EUR"}, "by_segment": {"arr": {"value": "2 EUR"}}}
    out = redaction.redact(payload, mapping)
    assert set(out) == {"arr", "by_segment"} and set(out["arr"]) == {"value"}
    assert set(out["by_segment"]) == {"Customer_01"}, "only segment-container keys are data"


def test_word_boundary_matching_does_not_touch_longer_words():
    mapping = {"Acme": "Customer_01"}
    assert redaction.substitute("Acme, Acmes and AcmeCorp", mapping) == "Customer_01, Acmes and AcmeCorp"
    assert redaction.restore("Customer_01 and Customer_010", {"Acme": "Customer_01"}) == "Acme and Customer_010"


def test_numeric_guard_allowed_set_is_unchanged_by_redaction():
    db = _db()
    computed = asyncio.run(gateway.load_computed_results(db, RUN_ID, "growth_engine"))
    mapping = {**asyncio.run(redaction.get_or_create_map(db, RUN_ID, computed)), "12": "Customer_01"}
    windows = gateway.STEP_CONFIG["growth_engine"]["windows"]
    plain = gateway.build_outbound(computed, {})
    redacted = gateway.build_outbound(computed, mapping)
    assert redacted != plain, "segments were relabelled"
    assert gateway.allowed_numerals(redacted, windows) == gateway.allowed_numerals(plain, windows)
    assert gateway.allowed_numerals(redaction.redact(plain, mapping), windows) == gateway.allowed_numerals(plain, windows)


def test_demo_upload_leaks_no_file_sheet_header_or_cell_value():
    """A real engine run over the demo upload, renamed to carry identifying strings."""
    pytest.importorskip("pandas")
    pytest.importorskip("fastapi")
    pytest.importorskip("motor")
    os.environ.setdefault("MONGO_URL", "mongodb://localhost:27017")
    os.environ.setdefault("DB_NAME", "data_boundary_test")
    import demo_data
    import growth_engine as ge
    import server

    file, sheet, header, cell = "Acme_Revenue.xlsx", "Acme Data", "Acme Segment", "Jane Doe (CEO)"
    spec = demo_data.DEMO_AUDITS[0]
    datasets, meta = demo_data.build(spec)
    uploads = {}
    for dtype, (df, mapping) in datasets.items():
        rows, mapping = server.df_to_records(df), dict(mapping)
        if dtype == "revenue":
            for r in rows:
                r[header] = r.pop(mapping["segment"])
            mapping["segment"] = header
            rows[0][mapping["currency"]] = cell            # a currency cell the engine quotes in a reason
        if dtype == "crm":
            rows[0][mapping["founder_involved"]] = cell    # a founder cell the engine quotes in a reason
        name, tab = (file, sheet) if dtype == "revenue" else (meta[dtype]["file"], meta[dtype]["sheet"])
        uploads[dtype] = {"file": name, "sheet": tab, "columns": list(rows[0]), "rows": rows, "mapping": mapping}

    norm = {d: server.normalize(u["rows"], d, u["mapping"]) for d, u in uploads.items()}
    fx = {k.upper(): v for k, v in meta.get("fx", {}).items()}
    fx[spec["reporting_currency"].upper()] = 1.0
    cfg = {"reporting_currency": spec["reporting_currency"], "target_arr": spec["target_arr"],
           "target_date": spec["target_date"], "fx": fx, "billing_terms": {}, "default_l": 1, "as_of_month": None}
    sources = {d: {"file": u["file"], "sheet": u["sheet"]} for d, u in uploads.items()}
    results = server.sanitize(ge.compute_all(norm["revenue"], norm["crm"], norm["pnl"], cfg, sources,
                                             files=server.candidate_views(uploads)))

    stored = json.dumps(results, ensure_ascii=False).lower()
    # The header is stored only when a management question cites it; that path is covered above.
    for needle in (file, sheet, cell):
        assert needle.lower() in stored, f"fixture did not exercise {needle!r}"

    db = t.FakeDB()
    db["audits"].docs.append({"id": RUN_ID, "results": results, "reporting_currency": spec["reporting_currency"],
                              "target_arr": spec["target_arr"], "target_date": spec["target_date"]})
    adapter = t.FakeAdapter()
    asyncio.run(gateway.generate_narrative(db, RUN_ID, "growth_engine", adapter=adapter, sleep=t._noop_sleep))
    assert adapter.calls == 1
    sent = adapter.payloads[0].lower()
    for needle in (file, sheet, header, cell, "jane doe", "acme"):
        assert needle.lower() not in sent, f"{needle!r} reached the provider"
    assert json.dumps(db["audits"].docs[0]["results"], ensure_ascii=False).lower() == stored, "Mongo keeps the full text"


# ---------------------------------------------------------------------------
# Boundary invariants (CLAUDE.md rule 14). What may reach the model or the logs
# is enforced here, not in review. A PR that adds a results key, a reason string,
# a log line or a guard word extends the lists in this section.
# ---------------------------------------------------------------------------
import ast  # noqa: E402
import re  # noqa: E402

import test_narrative_model_guard as guard_t  # noqa: E402

SENTINEL = "JANEDOE"            # written into one cell of every mapped column
HEADER_SENTINEL = "Jane Doe"    # every mapped column header is renamed to carry it

# Results keys the engine writes that never leave the server. A key that is in
# neither this set nor gateway.OUTBOUND_FIELDS fails test_every_results_key_is_
# allowlisted_or_declared_server_only: decide where it belongs when you add it.
SERVER_ONLY_TOP_LEVEL = frozenset({"contract_version",                   # interface-contracts.md: the version the stored payload was built under
                                   "anomalies", "mrr_series", "new_mrr_by_quarter", "new_mrr_by_quarter_source",
                                   "revenue_series", "customers_series",    # claim-matching.md table 2a
                                   "revenue_reconciliation",                # chat-upload.md section 6.2; data_integrity reads a summary (step_payloads.py)
                                   "source_period", "metric_periods",       # dates-and-periods.md sections 2-3: months as text, no figure
                                   "won_deal_closes",                       # crm-dedup-clustering.md: data_integrity reads a closed sentence
                                   "yearly_basis",                          # chat-upload.md section 26.8: no narrative reads it
                                   "pnl_lines",                             # docs/design/pnl-lines.md: line labels and plan cells of the file; no narrative reads it
                                   "recognition",                           # revenue-file-recognition.md section 9: closed words and counts, in no slice
                                   "as_of_month"})                          # the analyst's setting: the model is told evidence_to and last_month_status instead
SERVER_ONLY_FIELDS = frozenset({
    "file", "sheet", "columns", "rows", "row_numbers", "unlocked_by",
    "values",                       # founder_involved_excluded.values are raw cells (cohort `data.values` is allowed by parent)
    "value_label",                  # acv_path.overall_band: display string, dashboard only
    "required_vs_observed_12m_reason", "required_vs_observed_24m_reason",   # acv_path: not yet allowlisted
    "monthly_series",               # gross_churn: the register's monthly rate (engine contract 8); test_the_monthly_churn_series_never_reaches_a_narrative_payload
})

# (file, function) of log calls that may carry a traceback. Empty: every log call carries error
# type, run id and engine step only, since an exception message or traceback can quote uploaded cells.
TRACEBACK_ALLOWED = frozenset()


def _engine_results(spec, inject: bool) -> dict:
    """A real engine run over a demo upload. With `inject`, one cell per mapped column and
    every mapped header carry a sentinel, so any reason or metric name that quotes upload
    content shows it."""
    import demo_data
    import growth_engine as ge
    import server

    datasets, meta = demo_data.build(spec)
    uploads = {}
    for dtype, (df, mapping) in datasets.items():
        rows, mapping = server.df_to_records(df), dict(mapping)
        if inject:
            for i, col in enumerate(c for c in mapping.values() if c):
                rows[i][col] = f"{SENTINEL} {col}"                # row i, column i: the other cells stay valid
            for field, col in list(mapping.items()):
                if col:
                    mapping[field] = f"{HEADER_SENTINEL} {field}"
                    for r in rows:
                        r[mapping[field]] = r.pop(col)
        uploads[dtype] = {"file": meta[dtype]["file"], "sheet": meta[dtype]["sheet"],
                          "columns": list(rows[0]), "rows": rows, "mapping": mapping}
    norm = {d: server.normalize(u["rows"], d, u["mapping"]) for d, u in uploads.items()}
    fx = {k.upper(): v for k, v in meta.get("fx", {}).items()}
    fx[spec["reporting_currency"].upper()] = 1.0
    cfg = {"reporting_currency": spec["reporting_currency"], "target_arr": spec["target_arr"],
           "target_date": spec["target_date"], "fx": fx, "billing_terms": {}, "default_l": 1, "as_of_month": None}
    sources = {d: {"file": u["file"], "sheet": u["sheet"]} for d, u in uploads.items()}
    return server.sanitize(ge.compute_all(norm["revenue"], norm["crm"], norm["pnl"], cfg, sources,
                                          files=server.candidate_views(uploads)))


@pytest.fixture(scope="module")
def engine_runs():
    pytest.importorskip("pandas")
    pytest.importorskip("fastapi")
    pytest.importorskip("motor")
    os.environ.setdefault("MONGO_URL", "mongodb://localhost:27017")
    os.environ.setdefault("DB_NAME", "data_boundary_test")
    import demo_data
    return [(spec, inject, _engine_results(spec, inject)) for spec in demo_data.DEMO_AUDITS for inject in (False, True)]


def _field_keys(node, parent, path, out):
    """Every (parent, key, path) in `node` whose key is a field name, not data."""
    if isinstance(node, list):
        for item in node:
            _field_keys(item, parent, path, out)
    elif isinstance(node, dict):
        for k, v in node.items():
            if parent in redaction.SEGMENT_CONTAINERS or parent in gateway.PERIOD_KEYED:
                _field_keys(v, "_data", f"{path}/{k}", out)       # segment names, periods, window lengths
            elif parent in gateway.DATASET_KEYED:
                assert k in gateway.DATASET_KEYED[parent], f"{path}/{k}: unknown dataset"
                _field_keys(v, k, f"{path}/{k}", out)
            else:
                out.append((parent, k, f"{path}/{k}"))
                _field_keys(v, k, f"{path}/{k}", out)


def test_every_results_key_is_allowlisted_or_declared_server_only(engine_runs):
    """A results key the engine writes is in OUTBOUND_FIELDS (the model sees it) or declared
    server-only here. An unknown key fails: a field added without a decision is otherwise
    dropped silently and the model never sees it (log 2026-10-02 gateway-data-boundary)."""
    assert not (SERVER_ONLY_FIELDS & gateway.OUTBOUND_FIELDS), "a key cannot be both"
    model_facing_top = set()
    for step in gateway.STEP_CONFIG:
        model_facing_top |= set(gateway._slice_for_step({k: [] for k in gateway.OUTBOUND_FIELDS}, step))
    model_facing_top |= set(gateway._GAP_FIELDS)
    allowed_by_parent = set().union(*gateway.OUTBOUND_FIELDS_BY_PARENT.values())
    unknown = []
    for spec, inject, results in engine_runs:
        for top, value in results.items():
            if top in SERVER_ONLY_TOP_LEVEL:
                continue
            if top not in model_facing_top:
                unknown.append(f"/{top} (top-level: not in any step slice and not in SERVER_ONLY_TOP_LEVEL)")
                continue
            found = []
            _field_keys(value, top, f"/{top}", found)
            for parent, key, path in found:
                if key not in gateway.OUTBOUND_FIELDS and key not in SERVER_ONLY_FIELDS \
                        and not (key in allowed_by_parent and key in gateway.OUTBOUND_FIELDS_BY_PARENT.get(parent, ())):
                    unknown.append(path)
    assert not unknown, "results keys that are neither allowlisted nor declared server-only:\n  " + "\n  ".join(sorted(set(unknown)))


def test_no_reason_or_metric_name_reaches_the_model_with_upload_content(engine_runs):
    """Every Missing Data reason and metric name the engine can write, with a sentinel in one cell
    of every mapped column and in every header: the model's version carries none of it. A reason
    that quotes a cell value needs a _NEUTRAL_REASONS entry (log 2026-10-02 gateway-data-boundary-2)."""
    injected = [(spec, results) for spec, inject, results in engine_runs if inject]
    stored_all = json.dumps([r for _, r in injected], ensure_ascii=False).lower()
    assert SENTINEL.lower() in stored_all, "fixture did not exercise the cell sentinel"
    # Headers reach results only through a management question (covered by the fixture test above).
    quoting = [item["metric"] for _, r in injected for item in r["missing_data"] if SENTINEL.lower() in item["reason"].lower()]
    assert quoting, "fixture: at least one stored reason quotes a cell value, so the neutral path is exercised"

    for spec, results in injected:
        db = t.FakeDB()
        db["audits"].docs.append({"id": RUN_ID, "results": results, "reporting_currency": spec["reporting_currency"],
                                  "target_arr": spec["target_arr"], "target_date": spec["target_date"]})
        computed = asyncio.run(gateway.load_computed_results(db, RUN_ID, "growth_engine"))
        mapping = asyncio.run(redaction.get_or_create_map(db, RUN_ID, computed))
        for item in results["missing_data"]:
            assert SENTINEL.lower() not in item["metric"].lower(), f"metric name quotes a cell: {item['metric']!r}"
            assert HEADER_SENTINEL.lower() not in item["metric"].lower(), f"metric name quotes a header: {item['metric']!r}"
            reason = redaction.substitute(gateway._outbound_gap_value("reason", item), mapping)
            assert SENTINEL.lower() not in reason.lower(), (
                f"{item['metric']!r} quotes a cell value in its reason; add a _NEUTRAL_REASONS entry: {reason!r}")
            assert HEADER_SENTINEL.lower() not in reason.lower(), f"{item['metric']!r} quotes a header: {reason!r}"
        adapter = t.FakeAdapter()
        asyncio.run(gateway.generate_narrative(db, RUN_ID, "growth_engine", adapter=adapter, sleep=t._noop_sleep))
        assert adapter.calls == 1
        sent = adapter.payloads[0].lower()
        for needle in (SENTINEL, HEADER_SENTINEL):
            assert needle.lower() not in sent, f"{needle!r} reached the provider"


_LOG_LEVELS = frozenset({"debug", "info", "warning", "warn", "error", "critical", "exception", "log"})
# The only calls that may take the exception object: they return a class name or a code name, never its text; except the
# outermost 500 handler's helper (George, 2026-10-23): the traceback's last frame as file:line in function, never the message.
# tests/test_endpoint_errors.py pins what it returns and that the request never reaches the log.
SAFE_EXCEPTION_READERS = frozenset({"type", "_engine_step", "_last_frame"})
# A parameter with one of these names, or typed as an exception, is treated as the exception object.
_EXCEPTION_PARAM_NAMES = frozenset({"exc", "e", "err", "error", "exception"})


def _log_calls():
    """(file, function, lineno, call text, problems) for every log call in the backend."""
    out = []
    files = [BACKEND / "server.py", BACKEND / "growth_engine.py", BACKEND / "demo_data.py",
             *sorted((BACKEND / "app").rglob("*.py"))]
    for path in files:
        tree = ast.parse(path.read_text(encoding="utf-8"), filename=str(path))
        parents = {child: node for node in ast.walk(tree) for child in ast.iter_child_nodes(node)}
        for node in ast.walk(tree):
            if not (isinstance(node, ast.Call) and isinstance(node.func, ast.Attribute)
                    and node.func.attr in _LOG_LEVELS and isinstance(node.func.value, ast.Name)
                    and "log" in node.func.value.id.lower()):
                continue
            # Names bound to an exception where this call sits: `except X as name` around it, and
            # parameters of the enclosing functions that are typed or named as an exception.
            bound, function, up = set(), "<module>", node
            while up in parents:
                up = parents[up]
                if isinstance(up, ast.ExceptHandler) and up.name:
                    bound.add(up.name)
                if isinstance(up, (ast.FunctionDef, ast.AsyncFunctionDef)):
                    function = up.name if function == "<module>" else function
                    for a in [*up.args.posonlyargs, *up.args.args, *up.args.kwonlyargs]:
                        annotation = ast.unparse(a.annotation) if a.annotation else ""
                        if a.arg in _EXCEPTION_PARAM_NAMES or annotation.endswith(("Exception", "Error")):
                            bound.add(a.arg)
            problems = []
            if node.func.attr == "exception":
                problems.append("logger.exception appends the traceback")
            for kw in node.keywords:
                if kw.arg in ("exc_info", "stack_info"):
                    problems.append(f"{kw.arg}= appends the traceback")
            for arg in [*node.args, *(kw.value for kw in node.keywords)]:
                for sub in ast.walk(arg):
                    if isinstance(sub, ast.Call) and ast.unparse(sub.func).endswith(("format_exc", "format_exception")):
                        problems.append("traceback text is formatted into the message")
                    if isinstance(sub, ast.Name) and sub.id in bound:
                        owner = parents.get(sub)
                        if not (isinstance(owner, ast.Call) and ast.unparse(owner.func) in SAFE_EXCEPTION_READERS):
                            problems.append(f"the exception {sub.id!r} is formatted into the message")
            out.append((str(path.relative_to(BACKEND)), function, node.lineno, ast.unparse(node), problems))
    return out


def test_no_log_call_formats_an_exception_message_or_traceback():
    """A log call may use an exception only through SAFE_EXCEPTION_READERS, whether it sits in an
    `except` block or in a helper that takes the exception as a parameter. The message and the
    traceback can quote uploaded cell values (log 2026-10-02 fix-calc-correctness-2). The two sites
    TRACEBACK_ALLOWED is empty; a new entry is a decision, not a default."""
    calls = _log_calls()
    assert len(calls) >= 8, "the scan found fewer log calls than the codebase is known to have"
    violations = [f"{file}:{line} in {func}(): {text}  <- {'; '.join(problems)}"
                  for file, func, line, text, problems in calls
                  if problems and (file, func) not in TRACEBACK_ALLOWED]
    assert not violations, "log calls that can carry upload content:\n  " + "\n  ".join(violations)
    allowed_in_use = {(file, func) for file, func, _, _, problems in calls if problems}
    assert allowed_in_use == set(TRACEBACK_ALLOWED), "TRACEBACK_ALLOWED lists a site that no longer exists; remove it"


def _alternatives(pattern: re.Pattern) -> list:
    """The top-level words of a `\\b(?:a|b(?:s|es)|c)\\b` direction regex."""
    body = pattern.pattern
    assert body.startswith(r"\b(?:") and body.endswith(r")\b"), "direction regex shape changed; update _alternatives"
    body, out, depth, cur = body[5:-3], [], 0, ""
    for ch in body:
        depth += (ch == "(") - (ch == ")")
        if ch == "|" and depth == 0:
            out.append(cur)
            cur = ""
        else:
            cur += ch
    return out + [cur]


# One false friend per direction word: the word in another sense, or the confusable form that
# sits next to it, with the verdict the strict reading gives. "passes" means the dropped minus
# sign on change_arr -77,949 is accepted; "flagged" means 77949 stays unverified.
# A realistic false friend that passes is a hole in the list: "contracts", "down" and "lower" were
# removed for that (test_removed_direction_words_no_longer_accept_a_dropped_sign).
FALSE_FRIENDS = {
    # decline words
    "fall": ("The fall in ARR, 77,949 EUR, is in Segment B.", "passes"),
    "fell": ("Costs rose, ARR fell 77,949 EUR.", "flagged"),
    "decline": ("The decline of 77,949 EUR is in Segment B.", "passes"),
    "drop": ("The drop-off is 77,949 EUR.", "passes"),
    "decrease": ("The decrease is 77,949 EUR.", "passes"),
    "shrink": ("The shrink in Segment B is 77,949 EUR.", "passes"),
    "loss of": ("A loss of 77,949 EUR of ARR.", "passes"),
    # growth words: a growth word in a decline sentence keeps the figure unverified (a reask, never a wrong figure)
    "grows": ("Segment B grows while ARR falls by 77,949 EUR.", "flagged"),
    "grew": ("Costs grew; ARR fell by 77,949 EUR.", "flagged"),
    "rise": ("The rise in churn cut ARR by 77,949 EUR.", "flagged"),
    "rose": ("Churn rose, so ARR fell by 77,949 EUR.", "flagged"),
    "increase": ("The increase in churn cut ARR by 77,949 EUR.", "flagged"),
    "up": ("ARR falls by up to 77,949 EUR.", "flagged"),
    "higher": ("Higher churn cut ARR by 77,949 EUR.", "flagged"),
    "gains": ("Gains of 77,949 EUR offset the decline.", "flagged"),
    "adds": ("Segment B adds 77,949 in new contracts.", "flagged"),
    "expands": ("ARR declines by 77,949 EUR as the base expands.", "flagged"),
    "improves": ("Margin improves as ARR declines by 77,949 EUR.", "flagged"),
    "accelerates": ("Churn accelerates; ARR falls by 77,949 EUR.", "flagged"),
    "wins": ("ARR declines by 77,949 EUR after two wins.", "flagged"),
    "won": ("Won deals fell; ARR dropped 77,949 EUR.", "flagged"),
}


def test_every_direction_word_has_a_false_friend_case():
    """A word added to _DECLINE or _GROWTH without its own entry in FALSE_FRIENDS fails here.
    Keys are the form of the word its case uses; each key must match exactly one direction word."""
    uncovered, matched = [], {}
    for name, rx in (("_DECLINE", gateway._DECLINE), ("_GROWTH", gateway._GROWTH)):
        for alt in _alternatives(rx):
            keys = [k for k in FALSE_FRIENDS if re.fullmatch(alt, k, re.IGNORECASE)]
            if not keys:
                uncovered.append(f"{name}: {alt}")
            for k in keys:
                matched[k] = matched.get(k, 0) + 1
    assert not uncovered, "direction words without their own case in FALSE_FRIENDS:\n  " + "\n  ".join(uncovered)
    for key, (text, _) in FALSE_FRIENDS.items():
        assert matched.get(key) == 1, f"{key!r} is not exactly one direction word"
        assert re.search(rf"\b{re.escape(key)}\b", text, re.IGNORECASE), f"{key!r}: its case does not contain the word"


@pytest.mark.parametrize("word", sorted(FALSE_FRIENDS))
def test_false_friend_verdict(word):
    text, expected = FALSE_FRIENDS[word]
    payload = gateway.build_outbound(guard_t.COMPUTED, {})
    assert "-77,949" in str(payload), "fixture: the engine figure is signed"
    flagged = gateway.numeric_guard(guard_t._narrative(text), payload).all
    assert flagged == ([] if expected == "passes" else ["77949"]), f"{word!r}: {text!r} -> {flagged}"


@pytest.mark.parametrize("text", ["A 77,949 EUR down payment was booked.", "The lower band is 77,949 EUR.",
                                  "ARR is down 77,949 EUR.", "ARR is lower by 77,949 EUR."])
def test_removed_direction_words_no_longer_accept_a_dropped_sign(text):
    """"down" and "lower" read as a decline in a down payment or the lower band, so the words were
    dropped; the figure stays unverified until the sentence uses a decline verb."""
    payload = gateway.build_outbound(guard_t.COMPUTED, {})
    assert gateway.numeric_guard(guard_t._narrative(text), payload).all == ["77949"]
    for word in ("down", "lower"):
        assert not gateway._DECLINE.search(word)


def test_the_outermost_500_handler_and_the_gateway_failure_log_type_run_id_and_step_only(monkeypatch, caplog):
    """An error whose message quotes a cell reaches both catch-all handlers: the log carries the error type and the request or run,
    never the message or a traceback. The outermost 500 handler adds the last frame (file:line in function, George 2026-10-23)."""
    pytest.importorskip("fastapi")
    import logging
    from fastapi.testclient import TestClient
    import server

    def boom(*args, **kwargs):
        raise RuntimeError("cannot parse cell 'Jane Doe (CEO)'")

    async def aboom(*args, **kwargs):
        boom()

    generate = gateway.generate_narrative            # server.llm_gateway is this module; keep the real one
    monkeypatch.setattr(server, "db", t.make_db())
    monkeypatch.setattr(server.llm_gateway, "generate_narrative", aboom)
    client = TestClient(server.app, raise_server_exceptions=False)
    with caplog.at_level(logging.DEBUG):
        assert client.post("/api/runs/run-abc/narrative/growth_engine").status_code == 500
    assert "unexpected error handling POST /api/runs/run-abc/narrative/growth_engine: error=RuntimeError at=" in caplog.text
    assert "Jane Doe" not in caplog.text
    caplog.clear()

    db = _db()
    monkeypatch.setattr(gateway.guards, "check_call_cap", aboom)
    with caplog.at_level(logging.DEBUG):
        result = asyncio.run(generate(db, RUN_ID, "growth_engine", adapter=t.FakeAdapter(), sleep=t._noop_sleep))
    assert result.narrative_status != "ok"
    assert f"unexpected gateway failure for run {RUN_ID} step growth_engine: error=RuntimeError" in caplog.text
    for text in (caplog.text, *[r.getMessage() for r in caplog.records]):
        assert "Jane Doe" not in text and "Traceback" not in text
    assert all(r.exc_info is None for r in caplog.records), "no traceback in the log"


# ---------------------------------------------------------------------------
# Deck text never reaches the model (docs/specs/deck-parser.md section 4). The deck parser has
# no path to the gateway, and the gateway reads no parsed text, no snippet and no source
# reference. If it ever reads claim candidates, it projects decks.GATEWAY_READABLE_FIELDS only.
# ---------------------------------------------------------------------------
import subprocess  # noqa: E402

from app import decks  # noqa: E402

DECK_SENTINEL = "DECKTEXT Jane Doe"            # in the parsed text and the snippet
DECK_FILE = "Acme_Board_Q3.pptx"               # in the source references


def _imports(tree):
    """(module, name) for every import at any depth, plus importlib / __import__ calls."""
    out = []
    for node in ast.walk(tree):
        if isinstance(node, ast.Import):
            out += [(a.name, a.name) for a in node.names]
        elif isinstance(node, ast.ImportFrom):
            out += [("." * node.level + (node.module or ""), a.name) for a in node.names]
        elif isinstance(node, ast.Call) and ast.unparse(node.func).endswith(("import_module", "__import__")):
            out.append((ast.unparse(node), "<dynamic>"))
    return out


def test_the_deck_parser_never_imports_or_calls_the_gateway():
    offenders = []
    for path in sorted((BACKEND / "app" / "decks").rglob("*.py")):
        tree = ast.parse(path.read_text(encoding="utf-8"))
        for module, name in _imports(tree):
            parts = set(re.split(r"[.\s\"'()]+", module)) | {name}
            if parts & {"llm", "gateway", "anthropic", "<dynamic>"}:
                offenders.append(f"{path.name}: {module} -> {name}")
        for node in ast.walk(tree):
            if isinstance(node, ast.Name) and node.id in ("gateway", "llm_gateway") or \
                    isinstance(node, ast.Attribute) and node.attr in ("gateway", "llm_gateway", "generate_narrative"):
                offenders.append(f"{path.name}:{node.lineno}: {ast.unparse(node)}")
    assert not offenders, "the deck parser reaches for the gateway:\n  " + "\n  ".join(offenders)
    # Transitively too: importing the parser loads neither the gateway nor a provider SDK.
    code = ("import sys, app.decks.parser, app.decks.claims; "
            "print(sorted(m for m in sys.modules if m.startswith('app.llm') or m.split('.')[0] == 'anthropic'))")
    out = subprocess.run([sys.executable, "-c", code], cwd=BACKEND, capture_output=True, text=True, check=True)
    assert out.stdout.strip() == "[]", out.stdout


def test_the_gateway_names_no_deck_text_snippet_or_source_field():
    """Static: no gateway module imports the deck package, or names the parsed-text collection
    or the parsed-text, snippet, borrowed-label or source-reference fields, or a table row's values
    by period (each keeps its column header's text and its cell), or a claim's confidence, where its type came from
    or where it sits on the page, its direction with no figure, its type group and its FX rate (layout, parser checks and
    the saved rate, computed in Python and shown to the analyst only)."""
    forbidden = {"blocks", "snippet", "label_from", "date_from", "sources", "by_period",
                 "confidence", "type_from", "reading", "claim_direction", "group", "fx", "fx_rate", "fx_date",
                 "mixed_slide", "inconsistencies", "page_segment", "page_share_of"}
    offenders = []
    for path in sorted((BACKEND / "app" / "llm").glob("*.py")):
        tree = ast.parse(path.read_text(encoding="utf-8"))
        offenders += [f"{path.name}: {module} -> {name}" for module, name in _imports(tree)
                      if "decks" in re.split(r"[.\s\"'()]+", module) or name == "decks"]
        offenders += [f"{path.name}:{node.lineno}: {node.value!r}" for node in ast.walk(tree)
                      if isinstance(node, ast.Constant) and node.value in forbidden]
    assert not offenders, "the gateway names deck text:\n  " + "\n  ".join(offenders)
    # The parsed-text collection is named once, to count the audit's decks for the data inventory (George, 2026-10-11).
    named = [f"{p.name}:{n.lineno}" for p in sorted((BACKEND / "app" / "llm").glob("*.py"))
             for n in ast.walk(ast.parse(p.read_text(encoding="utf-8")))
             if isinstance(n, ast.Constant) and n.value == decks.TEXT_COLLECTION]
    assert len(named) == 1 and named[0].startswith("gateway.py"), named
    assert gateway.DECK_COUNT_COLLECTION == decks.TEXT_COLLECTION


def _recording_db():
    """The usual stub plus a stored deck, recording every read the gateway makes."""
    reads = []

    class Collection(t.FakeCollection):
        def __init__(self, name):
            super().__init__()
            self.name = name

        async def find_one(self, flt, projection=None):
            doc = await super().find_one(flt, projection)
            reads.append((self.name, projection, doc))
            return doc

        def find(self, flt, projection=None):
            cursor = super().find(flt, projection)
            reads.append((self.name, projection, cursor._docs))
            return cursor

        def aggregate(self, pipeline):
            cursor = super().aggregate(pipeline)
            reads.append((self.name, None, cursor._docs))
            return cursor

    class DB(t.FakeDB):
        def __getitem__(self, name):
            return self._cols.setdefault(name, Collection(name))

    db = DB()
    doc = copy.deepcopy(t.RESULTS_DOC)
    doc["results"] = copy.deepcopy(RESULTS)
    db["audits"].docs.append(doc)
    db[decks.TEXT_COLLECTION].docs.append({"audit_id": RUN_ID, "file": DECK_FILE, "format": "pptx",
                                           "blocks": [{"slide": 3, "kind": "text", "text": DECK_SENTINEL}]})
    db[decks.CANDIDATES_COLLECTION].docs.append({
        "audit_id": RUN_ID, "id": "c1", "claim_type": "revenue", "value": 3600000, "unit": None, "currency": "USD",
        "target_date": "2024", "status": "approved", "file": DECK_FILE, "snippet": DECK_SENTINEL,
        "label_from": DECK_SENTINEL, "date_from": DECK_SENTINEL, "parsed": {"value": 3600000},
        "sources": [{"file": DECK_FILE, "slide": 3, "kind": "text"}], "inconsistent_dates": ["2024"],
        "by_period": [{"value": 3600000, "value_high": None, "target_date": "2024", "period": DECK_SENTINEL,
                       "source": {"file": DECK_FILE, "slide": 3, "kind": "table", "table": 1, "row": 2, "col": 2}}]})
    return db, reads


def test_the_gateway_reads_no_deck_text_snippet_or_source():
    """Dynamic: every gateway path runs over a database holding a parsed deck. Nothing the gateway
    reads, and nothing it sends, carries the deck's text, snippet or file name."""
    # Spec section 4: type, value, high value (of a range), unit, date, status - nothing else.
    assert decks.GATEWAY_READABLE_FIELDS == {"claim_type", "value", "value_high", "unit", "target_date", "status"}
    db, reads = _recording_db()
    adapter = t.FakeAdapter(replies=[json.dumps(NARRATIVE)])
    asyncio.run(gateway.generate_narrative(db, RUN_ID, "growth_engine", adapter=adapter, sleep=t._noop_sleep))
    asyncio.run(gateway.read_cached_narrative(db, RUN_ID, "growth_engine"))
    asyncio.run(gateway.narratives_for_run(db, RUN_ID))
    asyncio.run(gateway.disclosure_for_run(db, RUN_ID))
    asyncio.run(gateway.usage_for_run(db, RUN_ID))
    assert adapter.calls == 1 and reads, "fixture: the gateway ran and read Mongo"

    assert not [r for r in reads if r[0] == decks.TEXT_COLLECTION], "the gateway read the parsed-text collection"
    for name, projection, _ in reads:
        if name == decks.CANDIDATES_COLLECTION:
            fields = {k for k, v in (projection or {}).items() if v and k != "_id"}
            assert fields and fields <= decks.GATEWAY_READABLE_FIELDS, f"candidate projection {projection!r}"
    returned = json.dumps([doc for _, _, doc in reads], ensure_ascii=False, default=str)
    sent = adapter.payloads[0]
    for needle in (DECK_SENTINEL, DECK_FILE):
        assert needle not in returned, f"{needle!r} was read by the gateway"
        assert needle not in sent, f"{needle!r} reached the provider"


# ---------------------------------------------------------------------------
# The structure path (CLAUDE.md rules 16-18, docs/specs/llm-structure-reading.md section 10 and
# docs/specs/structure-labelling.md section 9). Redacted deck structure cells with Python's item list (each
# item's raw text a figure inside its cell) and column-mapping texts (at most 3 header rows, at most 3 samples
# per numeric or date column, a profile per text column) may reach the provider, with the audit's consent, as
# extracted text with cell positions. Nothing else does, and no sent text, raw item text included, is stored or
# logged.
# ---------------------------------------------------------------------------
from app.structures import redact as structure_redact  # noqa: E402
from app import structures  # noqa: E402

STRUCTURE_AUDIT = {"id": "audit-boundary", "company_name": "Target Co",
                   "structure_reading_consent": True, "results": None}
GOOD_STRUCTURE = "r1c2: FY2025\nr1c3: FY2026\nr2c1: Revenue\nr2c2: £1,200,000\nr2c3: £1,500,000"
GOOD_MAPPING = ("r1c1: Customer\nr1c2: Invoice Date\nr1c3: Amount\nc2 sample: 2025-01-31\nc2 sample: 2025-02-28\n"
                "c2 sample: 2025-03-31\nc3 sample: 1200.50\nc1 profile: distinct 42, typical length 12, shape Aa Aa")
# The labelling reply to GOOD_ITEMS: one label per listed item, no value, cell or flag.
STRUCTURE_REPLY = json.dumps({"type": "table", "pairs": [], "labels": [
    {"item": f"i{n}", "metric": "revenue", "period": period, "unit": "GBP", "unit_other": None,
     "actual_or_forecast": "forecast"} for n, period in ((1, "FY2025"), (2, "FY2026"))]})
NO_LABELS = json.dumps({"type": "table", "labels": [], "pairs": []})       # the reply to a text with no item
MAPPING_REPLY = json.dumps({"type": "column_mapping", "items": [
    {"metric": "customer_id", "period": None, "value": None, "unit": None, "actual_or_forecast": "unknown",
     "unit_other": None, "value_cell": "r1c1", "period_cells": [], "proposed_flags": []}]})
# Every reason a structure text is refused: codes, never text from the structure.
STRUCTURE_REFUSALS = frozenset({"not_text", "empty", "raw_bytes", "file_name",
                                "not_cells", "cell_too_long", "redaction_changed", "too_many_header_rows",
                                "too_many_samples", "text_value", "bad_item_line", "bad_hint",
                                "no_figure", "too_many_pages", "pipeline_rules"})   # the page_claims step (2026-10-22)
# A deck structure's item list (docs/specs/structure-labelling.md section 3): one line per figure Python listed, its
# raw text a figure cut from its cell; a roadmap adds its date cells and text lines. Nothing else passes.
GOOD_ITEMS = GOOD_STRUCTURE + ('\nitems:\ni1 r2c2 "1,200,000" 1200000 h r2c1 r1c2\n'
                               'i2 r2c3 "1,500,000" 1500000 h r2c1 r1c3')
# A KPI panel whose value has no label of its own takes the slide title as its label, marked (deck-parser.md section 7):
# a title cell is a label, never a value.
TITLED = ('r1c1 title: 2011 Estimated Revenue\nr1c2: $12 -$13 million\nitems:\n'
          'i1 r1c2#1 "12" 12000000 h r1c1\ni2 r1c2#2 "13 million" 13000000 h r1c1')


def _structure_db(mapping=None, **audit):
    db = t.FakeDB()
    db["audits"].docs.append({**STRUCTURE_AUDIT, **audit})
    if mapping:
        db["pseudonym_map"].docs.append({"run_id": STRUCTURE_AUDIT["id"], "mapping": mapping})
    return db


def _send(db, text, kind="table", reply=STRUCTURE_REPLY):
    adapter = t.FakeAdapter(replies=[reply])
    images = [page_images.png(1)] if kind == "page_claims" else None          # a page goes with its image (page-image.md)
    result = asyncio.run(gateway.read_structure(db, STRUCTURE_AUDIT["id"], text, kind, adapter=adapter,
                                                sleep=t._noop_sleep, images=images))
    return result, adapter


def test_redacted_structure_cells_reach_the_provider_as_cell_lines_only():
    cells = [{"row": 1, "col": 1, "text": "Contact"}, {"row": 2, "col": 1, "text": "jane.doe@northwind.com"},
             {"row": 3, "col": 1, "text": "Northwind Trading renewed"}, {"row": 4, "col": 1, "text": "Michael Smith"},
             {"row": 5, "col": 1, "text": "+44 20 7946 0958"}, {"row": 6, "col": 1, "text": "£1,200,000"}]
    mapping = {"Northwind Trading": "Customer_01"}
    redacted, _ = structure_redact.redact_structure(cells, "Target Co", mapping)
    text = structure_redact.structure_text(redacted)
    result, adapter = _send(_structure_db(mapping), text, reply=NO_LABELS)
    assert result.status == "read" and adapter.calls == 1, result.reason
    sent = json.loads(adapter.payloads[0])
    assert set(sent) == {"type", "text"} and sent["text"] == text
    for needle in ("jane.doe@northwind.com", "Northwind Trading", "Michael Smith", "+44 20 7946 0958"):
        assert needle not in sent["text"], f"{needle!r} reached the provider"
    assert all(re.match(r"^r\d+c\d+: ", line) for line in sent["text"].splitlines())


def test_no_client_name_is_read_or_refused_and_a_stored_one_on_an_old_audit_changes_nothing():
    """Chat-upload.md section 23: the client name is gone from the context, the checks and the refusal reasons."""
    text = "r1c1: Prepared for Northbridge Capital\nr1c2: FY2025\nr2c1: Revenue\nr2c2: £1,200,000"
    for extra in ({}, {"client_name": "Northbridge Capital"}):
        result, adapter = _send(_structure_db(**extra), text, reply=NO_LABELS)
        assert result.status == "read" and adapter.calls == 1, result.reason
    assert "client_name" not in STRUCTURE_REFUSALS
    assert "client_name" not in inspect.getsource(gateway.load_structure_context)
    assert "client_name" not in inspect.getsource(gateway.structure_text_problem)


def test_a_column_mapping_text_within_the_caps_reaches_the_provider():
    result, adapter = _send(_structure_db(), GOOD_MAPPING, "column_mapping", MAPPING_REPLY)
    assert result.status == "read" and adapter.calls == 1, result.reason
    sent = json.loads(adapter.payloads[0])["text"]
    parsed = structure_redact.parse_column_text(sent)
    assert len({c["row"] for c in parsed["headers"]}) <= 3 and all(len(v) <= 3 for v in parsed["samples"].values())
    assert set(parsed["profiles"]) == {1}, "the text column sends a profile"


def test_a_sheet_built_by_the_app_sends_no_text_cell_value():
    columns = ["Customer", "Invoice Date", "Amount", "Note"]
    rows = [{"Customer": f"JANEDOE Retail {i}", "Invoice Date": "2025-01-31T00:00:00", "Amount": 100 + i,
             "Note": "JANEDOE called"} for i in range(6)]
    text = structures.column_mapping_text(columns, rows, "Target Co", {}, ("Customer",))
    assert SENTINEL not in text and structure_redact.column_text_problem(text) is None
    result, adapter = _send(_structure_db(), text, "column_mapping", MAPPING_REPLY)
    assert adapter.calls == 1 and SENTINEL not in adapter.payloads[0]


@pytest.mark.parametrize("name, text, kind, reason", [
    ("raw bytes", b"PK\x03\x04 r1c1: x", "table", "not_text"),
    ("control bytes", "r1c1: x\x00\x01", "table", "raw_bytes"),
    ("a full page", "Our mission\nWe grew revenue strongly in 2025 and plan to triple it.\nTeam\nOur founders met at...",
     "table", "not_cells"),
    ("a prose snippet", "Revenue grew from £1.2m to £1.5m between FY2025 and FY2026.", "table", "not_cells"),
    ("a cell over 200 characters", "r1c1: " + "x" * 201, "table", "cell_too_long"),
    ("a file name", "r1c1: See Acme_Board_Q3.pptx", "table", "file_name"),
    ("an unredacted email", "r1c1: jane.doe@northwind.com", "table", "redaction_changed"),
    ("an unredacted phone number", "r1c1: +44 20 7946 0958", "table", "redaction_changed"),
    ("an unredacted name", "r1c1: Michael Smith", "table", "redaction_changed"),
    ("an unredacted customer name", "r1c1: Northwind Trading renewed", "table", "redaction_changed"),
    ("more than 3 samples", GOOD_MAPPING + "\nc3 sample: 1\nc3 sample: 2\nc3 sample: 3", "column_mapping",
     "too_many_samples"),
    ("more than 3 header rows", GOOD_MAPPING + "\nr4c1: Customer name", "column_mapping", "too_many_header_rows"),
    ("a text cell value on the column-mapping path", GOOD_MAPPING + "\nr1c4: Note\nc4 sample: Northwind Trading",
     "column_mapping", "text_value"),
    ("a text value as a profile and a sample", GOOD_MAPPING + "\nc1 sample: 12", "column_mapping", "text_value"),
    # A rules' hint is one field of the closed list and a confidence from 0 to 100, once per column (chat-upload.md section
    # 4.2, 2026-10-17): nothing else rides on a hint line.
    ("a hint naming no field", GOOD_MAPPING + "\nc1 rule: Northwind 60", "column_mapping", "bad_hint"),
    ("a hint holding text", GOOD_MAPPING + "\nc1 rule: customer_id Northwind Trading 60", "column_mapping", "bad_hint"),
    ("a hint over 100", GOOD_MAPPING + "\nc3 rule: amount 160", "column_mapping", "bad_hint"),
    ("two hints for a column", GOOD_MAPPING + "\nc3 rule: amount 60\nc3 rule: currency 30", "column_mapping", "bad_hint"),
    # An item line passes only in format, its raw text a figure inside its cell.
    ("an item line out of format", GOOD_STRUCTURE + "\nitems:\ni1 r2c2 Revenue grew strongly 1200000", "table",
     "bad_item_line"),
    ("prose after an item line", GOOD_ITEMS + "\nRevenue grew strongly in 2025", "table", "bad_item_line"),
    ("raw text that is not in its cell", GOOD_STRUCTURE + '\nitems:\ni1 r2c2 "9,999" 9999', "table", "bad_item_line"),
    ("raw text that is not a figure", GOOD_STRUCTURE + '\nitems:\ni1 r2c1 "Revenue" 1', "table", "bad_item_line"),
    ("raw text holding more than a figure", "r1c1: Revenue grew 12% in 2025\nitems:\n"
     'i1 r1c1 "Revenue grew 12%" 12', "table", "bad_item_line"),
    ("an item citing a cell not sent", GOOD_STRUCTURE + '\nitems:\ni1 r9c9 "5" 5', "table", "bad_item_line"),
    ("a header cell not sent", GOOD_STRUCTURE + '\nitems:\ni1 r2c2 "1,200,000" 1200000 h r9c9', "table",
     "bad_item_line"),
    ("item ids out of order", GOOD_STRUCTURE + '\nitems:\ni2 r2c2 "1,200,000" 1200000', "table", "bad_item_line"),
    ("a date line on a structure not sent as a roadmap", GOOD_ITEMS + "\nd1 r1c2", "table", "bad_item_line"),
    ("a date line citing a cell not sent", "r1c1: Launch the API\nr2c1: Q3 2025\nitems:\nd1 r9c1", "roadmap",
     "bad_item_line"),
    # A title line passes only in format, redacted like any cell, and no item line may cite it.
    ("an item citing a title cell", 'r1c1 title: Top 3 Metrics\nr1c2: $7m\nitems:\ni1 r1c1 "3" 3', "kpi_panel",
     "bad_item_line"),
    ("a title mark out of format", "title r1c1: 2011 Estimated Revenue\nr1c2: $7m", "kpi_panel", "not_cells"),
    ("an unredacted name in a title", "r1c1 title: Michael Smith\nr1c2: $7m", "kpi_panel", "redaction_changed"),
    ("a title over 200 characters", "r1c1 title: " + "x" * 201 + "\nr1c2: $7m", "kpi_panel", "cell_too_long"),
])
def test_the_structure_path_refuses_what_rule_16_does_not_allow(name, text, kind, reason):
    assert reason in STRUCTURE_REFUSALS
    db = _structure_db({"Northwind Trading": "Customer_01"})
    result, adapter = _send(db, text, kind)
    assert adapter.calls == 0 and getattr(adapter, "counted", 0) == 0, f"{name} reached the provider"
    assert (result.status, result.reason) == ("refused", f"refused: {reason}"), name


# ---------------------------------------------------------------------------
# The sheet skeleton (chat-upload.md section 26; CLAUDE.md rule 16 as amended 2026-10-18, George): for the sheet_structure step
# only, every text cell of one sheet (titles, headers, row labels), indents and num/blank markers, after redaction and the
# withhold rule; never a figure; at most 400 rows with text.
# ---------------------------------------------------------------------------
GOOD_SHEET = ("r1c1: Product P&L\nr1 indent 0 blank c2-c3\nr2c1: Line\nr2c2: FY2021\nr2c3: FY2022\nr2 indent 0\n"
              "r3c1: Revenue\nr3 indent 0 num c2-c3\nr4c1: Product A\nr4 indent 1 num c2-c3\nc5 withheld: text, 40 distinct values")
SHEET_REPLY = json.dumps({"type": "sheet_structure", "header_row": 2, "label_col": 1,
                          "periods": [{"col": 2, "grain": "year", "start": "2021-01", "tag": "actual"}],
                          "rows": [{"row": 3, "role": "total", "parent_row": None, "meaning": "revenue"},
                                   {"row": 4, "role": "subline", "parent_row": 3, "meaning": "revenue"}]})
SHEET_REFUSALS = frozenset({"not_cells", "too_many_rows", "figure", "withhold_missed", "cell_too_long", "redaction_changed",
                            "file_name", "raw_bytes"})
_CUSTOMER_COLUMN = "\n".join(f"r{r}c1: Shop {chr(65 + r % 26)}{chr(65 + r // 26)} Ltd" for r in range(3, 40))


def test_a_sheet_skeleton_within_the_limits_reaches_the_provider_as_sent():
    result, adapter = _send(_structure_db(), GOOD_SHEET, "sheet_structure", SHEET_REPLY)
    assert result.status == "read" and adapter.calls == 1, result.reason
    sent = json.loads(adapter.payloads[0])
    assert set(sent) == {"type", "text"} and sent == {"type": "sheet_structure", "text": GOOD_SHEET}
    assert adapter.requests[0]["max_tokens"] == gateway.SHEET_MAX_TOKENS
    assert result.sheet["rows"][0] == {"row": 3, "role": "total", "parent_row": None, "meaning": "revenue", "basis": None}


def test_a_sheet_reply_may_carry_a_basis_from_the_closed_list_and_nothing_else():
    with_basis = json.loads(SHEET_REPLY)
    with_basis["rows"][0]["basis"] = "ytd"
    result, adapter = _send(_structure_db(), GOOD_SHEET, "sheet_structure", json.dumps(with_basis))
    assert result.status == "read" and result.sheet["rows"][0]["basis"] == "ytd", result.reason
    from app.llm import schemas
    assert set(schemas.SHEET_BASES) == {"month", "quarter", "ytd", "ltm", "lifetime", "balance"}


@pytest.mark.parametrize("name,text,reason", [
    ("a figure in a cell line", GOOD_SHEET + "\nr3c2: 1,200", "figure"),
    ("an amount with a currency", GOOD_SHEET + "\nr3c2: £1,200,000", "figure"),
    ("a 401st row with text", "\n".join(f"r{r}c1: Line {r}" for r in range(1, 402)), "too_many_rows"),
    ("a customer list not withheld", "r1c1: Customer\n" + _CUSTOMER_COLUMN, "withhold_missed"),
    ("an unredacted email", GOOD_SHEET + "\nr5c1: jane.doe@northwind.com", "redaction_changed"),
    ("an unredacted phone number", GOOD_SHEET + "\nr5c1: +44 20 7946 0958", "redaction_changed"),
    ("an unredacted name", GOOD_SHEET + "\nr5c1: Michael Smith", "redaction_changed"),
    ("a customer name, not its pseudonym", GOOD_SHEET + "\nr5c1: Northwind Trading revenue", "redaction_changed"),
    ("the company name", "r1c1: Target Co P&L\n" + GOOD_SHEET, "redaction_changed"),
    ("a file name", GOOD_SHEET + "\nr5c1: see pnl_2025.xlsx", "file_name"),
    ("a cell over 200 characters", GOOD_SHEET + "\nr5c1: " + "x" * 201, "cell_too_long"),
    ("prose that is no line", GOOD_SHEET + "\nRevenue grew strongly in 2025", "not_cells"),
    ("a title mark", "r1c1 title: Product P&L\nr2c1: Revenue", "not_cells"),
])
def test_the_sheet_path_refuses_what_the_amended_rule_16_does_not_allow(name, text, reason):
    assert reason in SHEET_REFUSALS
    result, adapter = _send(_structure_db({"Northwind Trading": "Customer_01"}), text, "sheet_structure", SHEET_REPLY)
    assert adapter.calls == 0 and getattr(adapter, "counted", 0) == 0, f"{name} reached the provider"
    assert (result.status, result.reason) == ("refused", f"refused: {reason}"), name


def test_a_year_header_is_a_period_and_the_skeleton_of_a_statement_passes_with_a_customer_pseudonym():
    text = "r1c1: Line\nr1c2: 2024\nr1c3: 2025B\nr2c1: Customer_01 revenue\nr2 indent 0 num c2-c3"
    result, adapter = _send(_structure_db({"Northwind Trading": "Customer_01"}), text, "sheet_structure", SHEET_REPLY.replace(
        '"row": 4, "role": "subline", "parent_row": 3', '"row": 2, "role": "subline", "parent_row": null').replace(
        '"header_row": 2', '"header_row": 1').replace('"row": 3, "role": "total"', '"row": 1, "role": "heading"'))
    assert result.status == "read" and adapter.calls == 1, result.reason


@pytest.mark.parametrize("reply,check", [
    ({"rows": [{"row": 9, "role": "total", "parent_row": None, "meaning": "revenue"}]}, "other"),
    ({"rows": [{"row": 3, "role": "total", "parent_row": 3, "meaning": "revenue"}]}, "other"),
    ({"rows": [{"row": 3, "role": "total", "parent_row": None, "meaning": "revenue"}] * 2}, "item_count"),
    ({"rows": [{"row": 3, "role": "total", "parent_row": None, "meaning": "profit"}]}, "metric_list"),
    ({"rows": [{"row": 3, "role": "total", "parent_row": None, "meaning": "revenue", "basis": "last year"}]}, "metric_list"),
    ({"label_col": 9}, "other"),
    ({"rows": [{"row": 3, "role": "total", "parent_row": None, "meaning": "revenue", "label": "Revenue"}]}, "schema"),
])
def test_a_sheet_reply_naming_anything_not_sent_or_carrying_text_is_rejected_with_a_closed_check(reply, check, caplog):
    import logging
    bad = json.dumps({**json.loads(SHEET_REPLY), **reply})
    with caplog.at_level(logging.WARNING):
        result, adapter = _send(_structure_db(), GOOD_SHEET, "sheet_structure", bad)
    assert result.status == "not_read" and adapter.calls == 2, "one reask, then not read"
    assert f"check={check}" in caplog.text and "Revenue" not in caplog.text


def test_no_sheet_skeleton_or_label_in_a_log_llm_calls_llm_structures_or_a_dataset(api, caplog):
    import logging
    from app.structures import sheet as sheet_mod
    name = "yearly_pnl_by_product.csv"
    content = (BACKEND / "tests" / "fixtures" / "mapping" / name).read_bytes()
    with caplog.at_level(logging.DEBUG):
        chat.upload(api, name, content, dtype="pnl")
    (payload,) = [json.loads(p) for p in chat.adapter_of(api).payloads if json.loads(p)["type"] == "sheet_structure"][:1]
    sent = payload["text"]
    assert "r10c1: Sales" in sent and not any(structure_redact.sheet_figure(c["text"])
                                               for c in structure_redact.parse_sheet_text(sent)["cells"])
    logs = caplog.text + "".join(r.getMessage() for r in caplog.records)
    for needle in ("Product A", "Sales team", "General and administrative", name):
        assert needle not in logs, f"{needle!r} was logged"
    for coll in ("llm_calls", gateway.STRUCTURES_COLLECTION, "column_mappings", "audits"):
        assert "Product A" not in json.dumps(api.db[coll].docs, default=str), f"a label reached {coll}"
    assert sent not in json.dumps(api.db["datasets"].docs, default=str), "the skeleton is never stored"
    assert "sheet=statement" in caplog.text
    assert sheet_mod.MAX_ROWS == structure_redact.SHEET_MAX_ROWS == 400


def test_the_yearly_basis_block_never_reaches_a_narrative_payload():
    from schemas.metrics import MetricsPayload
    assert "yearly_basis" in MetricsPayload.model_fields and "yearly_basis" in SERVER_ONLY_TOP_LEVEL
    assert "yearly_basis" not in gateway.OUTBOUND_TOP_LEVEL and "yearly_basis" not in gateway.OUTBOUND_FIELDS


# docs/design/pnl-lines.md: the statement's findings carry the file's own line labels and its plan cells. The block is server-only, pure Python, and
# never reaches a narrative payload, a log line or the provider (rules 14 and 17).
PNL_SENTINEL = "Falcon Orchid ledger line"


def _results_with_pnl_lines():
    results = copy.deepcopy(RESULTS)
    cite = {**SOURCE, "file": "falcon.csv", "rule": "Sum check"}
    results["pnl_lines"] = {
        "basis": "yearly", "source": cite,
        "findings": [{"kind": "sum", "line": PNL_SENTINEL, "period": "FY 2021", "gap": 7_654_321.0, "gap_points": None, "source": cite,
                      "text": f"{PNL_SENTINEL} differs from the sum of its lines by 7,654,321 in FY 2021"}],
        "plan": [{"metric": "Revenue", "period": "FY 2022", "period_start": "2022-01", "period_end": "2022-12", "value": 8_765_432.0, "cell": "E3",
                  "source_text": "File Falcon!E3 (plan)", "measure": "growth", "implied": 0.6, "observed": 0.875, "gap_points": -27.5,
                  "compared": "FY 2021 → FY 2022 against FY 2020 → FY 2021", "reason": None, "source": cite}]}
    return results


def test_the_pnl_lines_block_is_declared_server_only_and_in_no_step_slice():
    from schemas.metrics import MetricsPayload
    assert "pnl_lines" in MetricsPayload.model_fields and "pnl_lines" in SERVER_ONLY_TOP_LEVEL
    assert "pnl_lines" not in gateway.OUTBOUND_TOP_LEVEL and "pnl_lines" not in gateway.OUTBOUND_FIELDS
    for step in gateway.STEP_CONFIG:
        assert "pnl_lines" not in gateway._slice_for_step(_results_with_pnl_lines(), step), f"{step} slices the statement block"


def test_the_pnl_lines_block_never_reaches_a_narrative_payload_or_the_provider():
    db = t.FakeDB()
    doc = copy.deepcopy(t.RESULTS_DOC)
    doc["results"] = _results_with_pnl_lines()
    db["audits"].docs.append(doc)
    needles = ("pnl_lines", PNL_SENTINEL, "7,654,321", "7654321", "8765432", "falcon.csv", "File Falcon!E3")
    for step in gateway.STEP_CONFIG:
        computed = asyncio.run(gateway.load_computed_results(db, RUN_ID, step))
        sent = json.dumps(gateway.build_outbound(computed, {}), ensure_ascii=False)
        for needle in needles:
            assert needle not in sent, f"{step}: {needle!r} reached the outbound payload"
    adapter = t.FakeAdapter(replies=[json.dumps(NARRATIVE)])
    asyncio.run(gateway.generate_narrative(db, RUN_ID, "growth_engine", adapter=adapter, sleep=t._noop_sleep))
    for needle in needles:
        assert needle not in adapter.payloads[0], f"{needle!r} reached the provider"


def test_the_pnl_lines_module_imports_nothing_from_the_gateway_and_writes_no_log_line():
    tree = ast.parse((BACKEND / "app" / "structures" / "pnl_lines.py").read_text(encoding="utf-8"))
    offenders = [f"{m} -> {n}" for m, n in _imports(tree) if (set(re.split(r"[.\s\"'()]+", m)) | {n}) & {"llm", "gateway", "anthropic", "server", "logging"}]
    assert not offenders, "pnl_lines reaches for the gateway or logs:\n  " + "\n  ".join(offenders)
    calls = [ast.unparse(n.func) for n in ast.walk(tree) if isinstance(n, ast.Call) and isinstance(n.func, ast.Attribute)
             and n.func.attr in ("debug", "info", "warning", "error", "exception", "critical", "log")]
    assert not calls, f"pnl_lines writes a log line: {calls}"


def test_the_sheet_lines_log_line_names_the_run_the_type_and_two_counts():
    tree = ast.parse((BACKEND / "server.py").read_text(encoding="utf-8"))
    fn = next(n for n in ast.walk(tree) if isinstance(n, ast.AsyncFunctionDef) and n.name == "decide_lines")
    calls = [n for n in ast.walk(fn) if isinstance(n, ast.Call) and isinstance(n.func, ast.Attribute) and n.func.attr in _LOG_LEVELS
             and isinstance(n.func.value, ast.Name) and "log" in n.func.value.id.lower()]
    (call,) = calls
    assert {ast.unparse(a) for a in call.args[1:]} == {"audit_id", "dtype", "len(decisions)",
                                                       "sheet_mod.pending_count(sheet_map)"}, ast.unparse(call)


def test_an_item_list_in_format_reaches_the_provider():
    result, adapter = _send(_structure_db(), GOOD_ITEMS)
    assert result.status != "refused" and adapter.calls == 1, result.reason
    assert json.loads(adapter.payloads[0])["text"] == GOOD_ITEMS
    roadmap = 'r1c1: Launch the API\nr2c1: Q3 2025\nr3c1: 5K users\nitems:\ni1 r3c1 "5K" 5000\nd1 r2c1\nt1 r1c1\nt2 r3c1'
    result, adapter = _send(_structure_db(), roadmap, "roadmap")
    assert result.status != "refused" and adapter.calls > 0, result.reason


def test_a_title_line_on_a_kpi_panel_reaches_the_provider_marked_with_no_item_in_it():
    reply = json.dumps({"type": "kpi_panel", "pairs": [], "labels": [
        {"item": f"i{n}", "metric": "revenue", "period": "2011", "unit": "USD", "unit_other": None,
         "actual_or_forecast": "forecast"} for n in (1, 2)]})
    result, adapter = _send(_structure_db(), TITLED, "kpi_panel", reply)
    assert result.status == "read" and adapter.calls == 1, result.reason
    sent = json.loads(adapter.payloads[0])
    assert sent == {"type": "kpi_panel", "text": TITLED}
    assert [line for line in sent["text"].splitlines() if " title: " in line] == ["r1c1 title: 2011 Estimated Revenue"]


@pytest.mark.parametrize("deck, page, kind", [("01-front-b.pptx", 15, "kpi_panel"), ("10-tea.pdf", 11, "roadmap")])
def test_a_kpi_or_roadmap_cells_next_to_is_cell_ids_kept_with_the_structure_and_never_sent(deck, page, kind):
    """deck-parser.md section 7 (issues #50 and #56): a KPI panel or roadmap cell keeps next_to, the ids of the cells
    of other boxes beside its line, stored with the structure. What reaches the provider is the cell lines and the item
    lines alone (a roadmap's date and text lines among them)."""
    pytest.importorskip("pdfplumber")
    from app.decks import parser
    from app.structures import items as structure_items
    deck_file = BACKEND.parent / "tests" / "fixtures" / "decks" / "decks" / deck
    panel, = [s for s in parser.parse_deck(deck_file.read_bytes(), deck_file.name)["structures"]
              if (s.get("slide") or s.get("page")) == page and s["type"] == kind]
    near = [n for c in panel["cells"] for n in c.get("next_to", ())]
    assert near and all(re.fullmatch(r"r\d+c\d+", n) for n in near)
    boxes = [c["date_box"] for c in panel["cells"] if "date_box" in c]
    assert bool(boxes) == (kind == "roadmap") and all(isinstance(b, int) for b in boxes), \
        "a roadmap line keeps the number of the date box beside it (issue #47), stored and never sent"
    cells, _ = structure_redact.redact_structure(panel["cells"], "Target Co", {})
    redacted = {**panel, "cells": cells}
    listed = structure_items.list_items(redacted)
    text = structure_items.text(redacted, listed)
    reply = json.dumps({"type": kind, "pairs": [], "labels": [
        {"item": i["id"], "metric": "not_a_metric", "period": None, "unit": None, "unit_other": None,
         "actual_or_forecast": "unknown"} for i in listed["items"]]})
    result, adapter = _send(_structure_db(), text, kind, reply)
    assert result.status == "read" and adapter.calls == 1, result.reason
    sent = json.loads(adapter.payloads[0])
    assert sent == {"type": kind, "text": text} and "next_to" not in sent["text"]
    lines = sent["text"].splitlines()
    assert lines[:lines.index("items:")] == [f"r{c['row']}c{c['col']}{' title' if c.get('title') else ''}: {c['text']}"
                                             for c in sorted(cells, key=lambda c: (c["row"], c["col"]))]
    assert lines[lines.index("items:") + 1:] == structure_redact.item_lines(listed)


def test_a_kpi_panel_cells_bar_year_is_kept_with_the_structure_and_never_sent():
    """deck-parser.md section 7 (issue #148): a bar's year label by place is Python's, kept as "bar_year" on the cell and stored with the
    structure; what reaches the provider is the cell lines and the item lines alone, the year label of the axis not among them."""
    pytest.importorskip("pdfplumber")
    from pdf_support import pdf_page
    from app.decks import parser
    from app.structures import items as structure_items
    texts = [(50, 560, 20, "Plan")] + [(x, 380 + 20 * i, 12, bar) for i, (bar, x) in enumerate(zip(("€1.48MM", "€3.39MM", "€5.3MM"), (100, 255, 410)))]
    texts += [(x, 340, 12, year) for year, x in zip(("2022", "2023", "2024"), (150, 300, 450))]
    panel, = [s for s in parser.parse_deck(pdf_page(texts, []), "deck.pdf")["structures"] if s["type"] == "kpi_panel"]
    assert sorted(c["bar_year"] for c in panel["cells"]) == ["2022", "2023", "2024"]
    cells, _ = structure_redact.redact_structure(panel["cells"], "Target Co", {})
    redacted = {**panel, "cells": cells}
    listed = structure_items.list_items(redacted)
    text = structure_items.text(redacted, listed)
    reply = json.dumps({"type": "kpi_panel", "pairs": [], "labels": [
        {"item": i["id"], "metric": "not_a_metric", "period": None, "unit": None, "unit_other": None,
         "actual_or_forecast": "unknown"} for i in listed["items"]]})
    result, adapter = _send(_structure_db(), text, "kpi_panel", reply)
    assert result.status == "read" and adapter.calls == 1, result.reason
    sent = json.loads(adapter.payloads[0])
    assert sent == {"type": "kpi_panel", "text": text}
    assert "bar_year" not in sent["text"] and not any(year in sent["text"] for year in ("2022", "2023", "2024"))


def test_no_call_is_made_without_consent():
    for audit in ({"structure_reading_consent": False}, {"structure_reading_consent": None}):
        result, adapter = _send(_structure_db(**audit), GOOD_STRUCTURE)
        assert result.status == "no_consent" and adapter.calls == 0
        result, adapter = _send(_structure_db(**audit), GOOD_MAPPING, "column_mapping", MAPPING_REPLY)
        assert result.status == "no_consent" and adapter.calls == 0


def test_no_sent_text_in_a_log_llm_calls_or_llm_structures(caplog):
    import logging
    db = _structure_db()
    with caplog.at_level(logging.DEBUG):
        _send(db, GOOD_ITEMS)
        _send(db, GOOD_ITEMS.replace("FY2025", "FY2027"),
              reply=json.dumps({"type": "table", "items": [{"bad": 1}]}))           # rejected twice: not read
        _send(db, "r1c1: jane.doe@northwind.com")                                   # refused
    stored = json.dumps(db[gateway.STRUCTURES_COLLECTION].docs + db["llm_calls"].docs, ensure_ascii=False, default=str)
    logs = caplog.text + "".join(r.getMessage() for r in caplog.records)
    for needle in ("Revenue", "£1,200,000", "1,200,000", "1,500,000", "jane.doe", GOOD_STRUCTURE, GOOD_ITEMS):
        assert needle not in stored, f"{needle!r} was stored"
        assert needle not in logs, f"{needle!r} was logged"
    assert "structure read: run_id=audit-boundary step=structures hash=" in caplog.text
    assert "structure refused: run_id=audit-boundary step=structures reason=redaction_changed" in caplog.text
    assert "[redacted]" not in logs, "the placeholder count is not logged either"
    assert all(r.exc_info is None for r in caplog.records)


def test_unit_other_carries_an_iso_code_only_so_no_cell_text_is_stored_or_logged(caplog):
    """unit_other is the one string a reply item fills itself (spec section 1): the ISO code of a currency outside the
    listed 20, beside unit "other". A reply that writes cell text there is rejected (one reask, then not read) and
    the text reaches neither llm_structures, llm_calls nor a log."""
    import logging
    first, second = json.loads(STRUCTURE_REPLY)["labels"]
    item = {**first, "unit": "other", "unit_other": "Revenue £1,200,000"}
    db = _structure_db()
    with caplog.at_level(logging.DEBUG):
        result, adapter = _send(db, GOOD_ITEMS, reply=json.dumps({"type": "table", "labels": [item, second],
                                                                  "pairs": []}))
    assert (result.status, adapter.calls) == ("not_read", 2)
    stored = json.dumps(db[gateway.STRUCTURES_COLLECTION].docs + db["llm_calls"].docs, ensure_ascii=False, default=str)
    logs = caplog.text + "".join(r.getMessage() for r in caplog.records)
    for needle in ("Revenue", "£1,200,000"):
        assert needle not in stored and needle not in logs, needle
    db = _structure_db()
    result, _ = _send(db, GOOD_ITEMS, reply=json.dumps({"type": "table", "labels": [{**item, "unit_other": "ZAR"}, second],
                                                         "pairs": []}))
    stored, = db[gateway.STRUCTURES_COLLECTION].docs
    assert result.status == "read" and stored["output"]["labels"][0]["unit_other"] == "ZAR"


# Every key an llm_calls record may carry (rule 17): metadata, never text sent to or received from the provider.
# http_status and error_type describe a provider error: the HTTP status and the provider's error code (or, when it
# sends none that reads as a code, the exception's class name), never its message.
LLM_CALLS_KEYS = frozenset({"run_id", "step", "prompt_version", "model", "input_tokens", "output_tokens",
                            "estimated_cost_usd", "image_tokens", "image_cost_usd", "cache_hit", "status", "unmatched_numbers",
                            "timestamp", "content_hash", "deck_id", "http_status", "error_type", "structure_type"})


def test_a_provider_error_reaches_llm_calls_and_the_log_as_a_status_and_a_type_code_only(caplog):
    import logging
    anthropic = pytest.importorskip("anthropic")
    httpx2 = pytest.importorskip("httpx2")
    message = "cannot read cell 'Jane Doe (CEO)' in r2c2: £1,200,000"

    def error(status, error_type):
        response = httpx2.Response(status, request=httpx2.Request("POST", "https://api.anthropic.com/v1/messages"))
        body = {"type": "error", "error": {"type": error_type, "message": message}}
        return anthropic.Anthropic(api_key="test")._make_status_error(message, body=body, response=response)

    class Failing(t.FakeAdapter):
        def __init__(self, raised, on_count=False, replies=None):
            super().__init__(replies=replies, raise_with=None if on_count else raised)
            self.on_count, self.raised = on_count, raised

        def count_tokens(self, **kwargs):
            if self.on_count:
                raise self.raised
            return super().count_tokens(**kwargs)

    db, narrative_db = _structure_db(), t.make_db()
    with caplog.at_level(logging.DEBUG):
        _send(db, GOOD_ITEMS)                                                        # read: the ok record's keys
        for raised, on_count in ((error(400, "invalid_request_error"), False), (error(403, "permission_error"), True),
                                 (error(400, "Jane Doe (CEO) may not"), False), (error(529, "overloaded_error"), False)):
            asyncio.run(gateway.read_structure(db, STRUCTURE_AUDIT["id"], GOOD_ITEMS.replace("FY2025", "FY2027"),
                                               "table", adapter=Failing(raised, on_count), sleep=t._noop_sleep))
        asyncio.run(gateway.generate_narrative(narrative_db, RUN_ID, "growth_engine",
                                               adapter=Failing(error(400, "invalid_request_error")),
                                               sleep=t._noop_sleep))
    records = db["llm_calls"].docs + narrative_db["llm_calls"].docs
    assert [(r.get("http_status"), r.get("error_type")) for r in records] == [
        (None, None), (400, "invalid_request_error"), (403, "permission_error"), (400, "BadRequestError"),
        (529, "overloaded_error"), (400, "invalid_request_error")], "an error type that is not a code: the class name"
    for record in records:
        assert set(record) <= LLM_CALLS_KEYS, set(record) - LLM_CALLS_KEYS
    logs = caplog.text + "".join(r.getMessage() for r in caplog.records)
    for needle in ("Jane Doe", "£1,200,000", "cannot read cell"):
        assert needle not in json.dumps(records, ensure_ascii=False), f"{needle!r} was stored"
        assert needle not in logs, f"{needle!r} was logged"
    assert "reason=provider_unreachable status=529 type=overloaded_error" in logs


# The reason a cap refusal's "structure not read" line gives: a closed word per guard, never the guard's message.
CAP_REASONS = frozenset({"spend_cap", "token_cap", "allowance"})


def test_a_cap_refusal_logs_a_closed_reason_and_no_text(monkeypatch, caplog):
    import logging
    from datetime import datetime, timezone
    assert set(gateway.CAP_REASONS.values()) == CAP_REASONS
    assert set(gateway.CAP_REASONS) == {"daily_spend_cap_exceeded", "structure_token_cap_reached", "monthly_allowance_reached"}
    db = _structure_db()
    db["llm_calls"].docs.append({"run_id": STRUCTURE_AUDIT["id"], "step": "structures", "cache_hit": False,
                                 "input_tokens": 398000, "output_tokens": 0, "estimated_cost_usd": 0.0,
                                 "timestamp": "2026-10-05T00:00:00"})
    with caplog.at_level(logging.DEBUG):
        _send(db, GOOD_ITEMS)
        monkeypatch.setenv("LLM_DAILY_SPEND_CAP_USD", "0.01")
        spent = _structure_db()
        spent["llm_calls"].docs.append({"run_id": "elsewhere", "cache_hit": False, "estimated_cost_usd": 1.0,
                                        "timestamp": datetime.now(timezone.utc).isoformat()})
        _send(spent, GOOD_ITEMS)
    lines = [r.getMessage() for r in caplog.records if r.getMessage().startswith("structure not read:")]
    assert [line.split("reason=")[1] for line in lines] == ["token_cap status=- type=-", "spend_cap status=- type=-"]
    for needle in ("Revenue", "£1,200,000", "400,000-token", "spent today", "$"):
        assert needle not in "".join(lines), f"{needle!r} was logged"


def test_an_allowance_refusal_logs_the_closed_word_and_no_amount_label_or_id(caplog):
    import logging
    db = _structure_db()
    db["access_tokens"].docs.append({"id": "tester-9", "label": "Label Nine", "ledger": {
        "all": {"cost_usd": 10.0}, __import__("app.allowance", fromlist=["x"]).month_key(): {"cost_usd": 10.0}}})
    db["audits"].docs[0]["owner"] = "tester-9"
    with caplog.at_level(logging.DEBUG):
        result, adapter = _send(db, GOOD_ITEMS)
    lines = "".join(r.getMessage() for r in caplog.records if r.getMessage().startswith("structure not read:"))
    assert result.status == "not_read" and adapter.calls == 0
    assert "reason=allowance status=- type=-" in lines
    for needle in ("tester-9", "Label Nine", "10.0", "USD", "$"):
        assert needle not in "".join(r.getMessage() for r in caplog.records), f"{needle!r} was logged"


def test_the_owner_audit_cost_rows_hold_a_closed_key_set_and_the_tester_filter_covers_every_cost_key(caplog):
    """usage-and-allowance.md section 9: the owner's row has the tester label, a time and three counts, never an audit, company or
    file name; and every cost, token or USD key the app writes to the call log is in the set a tester never receives."""
    from app import allowance, tester_view
    db = t.FakeDB()
    db["audits"].docs.append({"id": "a1", "owner": "t1", "created_at": "2026-10-18T12:32:00+00:00", "company_name": "Acme Zed"})
    db["access_tokens"].docs.append({"id": "t1", "label": "Label One"})
    db["llm_calls"].docs.append({"run_id": "a1", "cache_hit": False, "input_tokens": 5, "output_tokens": 5, "estimated_cost_usd": 0.1,
                                 "file": "deck_zed.pptx"})
    rows = asyncio.run(allowance.audit_costs(db))["audits"]
    assert rows == [{"tester": "Label One", "created_at": "2026-10-18T12:32:00+00:00", "calls": 1, "tokens": 10, "cost_usd": 0.1,
                     "image_tokens": 0, "image_cost_usd": 0.0}]
    assert "Acme Zed" not in json.dumps(rows) and "deck_zed" not in json.dumps(rows)
    assert {"input_tokens", "output_tokens", "estimated_cost_usd", "image_tokens", "image_cost_usd"} <= tester_view.HIDDEN_KEYS


# The check a "structure not read" line names for a reply that fails validation (issue #64): a closed word, never the
# reply, the check's message or cell text.
PARSE_CHECKS = frozenset({"schema", "metric_list", "item_count", "other"})


def test_a_failed_reply_logs_a_closed_check_and_no_text(caplog):
    import logging
    assert set(gateway.PARSE_CHECKS) == PARSE_CHECKS
    labels = json.loads(STRUCTURE_REPLY)["labels"]

    def reply(**first):
        return json.dumps({"type": "table", "pairs": [], "labels": [{**labels[0], **first}, labels[1]]})
    with caplog.at_level(logging.DEBUG):
        for sent in (reply(period="Northbridge Capital £1,200,000"), reply(metric="Revenue at Target Co"),
                     reply(item="Target Co"), "Revenue £1,200,000 at Target Co"):
            result, adapter = _send(_structure_db(), GOOD_ITEMS, reply=sent)
            assert (result.status, adapter.calls) == ("not_read", 2)
    lines = [r.getMessage() for r in caplog.records if r.getMessage().startswith("structure not read:")]
    assert [line.rsplit(" check=", 1)[1] for line in lines] == ["schema", "metric_list", "item_count", "schema"]
    logs = caplog.text + "".join(r.getMessage() for r in caplog.records)
    for needle in ("Northbridge", "£1,200,000", "Target Co", "Revenue", "period must be", "Input should", "labelled"):
        assert needle not in logs, f"{needle!r} was logged"


def test_a_corrected_period_stores_a_count_and_never_the_cell_text_the_raw_text_or_the_rebuilt_period():
    """Rule 17: llm_structures keeps the model's reply as it came (its own period), the item list without raw text
    and the number of periods Python corrected; the header text the correction came from is not stored there."""
    from app.decks import TEXT_COLLECTION
    db = _structure_db(fiscal_year_end=3)
    db["datasets"].docs.append({"audit_id": STRUCTURE_AUDIT["id"], "dtype": "revenue", "mapped_at": "2026-10-05",
                                "mapping": {"customer_id": "Customer"}})
    cells = [{"row": 1, "col": 2, "text": "FY2025"}, {"row": 2, "col": 2, "text": "Apr"},
             {"row": 3, "col": 1, "text": "Revenue"}, {"row": 3, "col": 2, "text": "$5M"}]
    db[TEXT_COLLECTION].docs.append({"audit_id": STRUCTURE_AUDIT["id"], "deck_id": "d1", "file": "plan.pptx",
                                     "structures": [{"type": "table", "slide": 2, "header_rows": 2, "cells": cells}]})
    reply = json.dumps({"type": "table", "pairs": [], "labels": [
        {"item": "i1", "metric": "revenue", "period": "2025-04", "unit": "USD", "unit_other": None,
         "actual_or_forecast": "forecast"}]})
    adapter = t.FakeAdapter(replies=[reply])
    with mock.patch.dict(os.environ, {"DECK_PIPELINE": "rules"}):          # the structure read: the former path
        asyncio.run(structures.process_deck(db, STRUCTURE_AUDIT["id"], "d1", adapter=adapter, sleep=t._noop_sleep))
    assert '"$5M"' not in adapter.payloads[0] and '\\"5M\\"' in adapter.payloads[0], "the raw text went out once"
    stored, = db[gateway.STRUCTURES_COLLECTION].docs
    assert stored["output"]["labels"][0]["period"] == "2025-04" and stored["periods_corrected"] == 1
    assert stored["items"]["items"] == [{"id": "i1", "cell": "r3c2", "position": 1, "headers": ["r3c1", "r2c2", "r1c2"],
                                         "values": [{"value": 5000000, "dot_reading": None, "bracket_reading": None}]}]
    flat = json.dumps(stored, default=str)
    for needle in ("FY2025-04", "Apr", "Revenue", "$5M", '"5M"', "raw"):
        assert needle not in flat, f"{needle!r} was stored"


# Every check the verifier puts on an item, stored on its approval row as ai_checks: a boolean, or the closed
# word naming the default reading of a "2.500"-style number (dot_reading) or a bracketed number after text
# (bracket_reading). Never cell text. Values and cells are Python's, so there is no value check any more.
VERIFIER_CHECKS = {"period": (True, False), "period_corrected": (True, False),
                   "dot_reading": ("decimal", "thousands", None), "bracket_reading": ("negative", "positive", None)}


def test_the_verifier_checks_are_booleans_and_closed_words_only():
    from app.structures import items
    structure = {"type": "kpi_panel", "header_rows": 0, "cells": [
        {"row": 1, "col": 1, "text": "Approx. 2.500 hours", "box": 1},
        {"row": 2, "col": 1, "text": "Jane Doe 5", "box": 2},
        {"row": 3, "col": 1, "text": "Net loss (1,200)", "box": 3},
        {"row": 4, "col": 1, "text": "Telegram (30K)", "box": 4}]}
    listed = items.list_items(structure)
    labels = [{"item": i["id"], "metric": "sales", "period": None, "unit": None, "unit_other": None,
               "actual_or_forecast": "actual"} for i in listed["items"]]
    checked = structures.verify.verify(structure, listed, labels)
    readings = [("thousands", None), (None, None), (None, "negative"), (None, "positive")]
    assert len(checked["items"]) == len(readings)
    for got, (dot, bracket) in zip(checked["items"], readings):
        assert set(got["checks"]) == set(VERIFIER_CHECKS)
        assert (got["checks"]["dot_reading"], got["checks"]["bracket_reading"]) == (dot, bracket)
        assert all(value in VERIFIER_CHECKS[key] for key, value in got["checks"].items()), got["checks"]


def test_the_structure_path_never_reads_parsed_deck_text():
    """read_structure is handed the text; it reads consent and names from the audit, the pseudonym map,
    its own cache and the call log, never the parsed-text or candidate collections."""
    db, reads = _recording_db()
    db["audits"].docs[0].update({"company_name": "Target Co",
                                 "structure_reading_consent": True})
    adapter = t.FakeAdapter(replies=[STRUCTURE_REPLY])
    asyncio.run(gateway.read_structure(db, RUN_ID, GOOD_ITEMS, "table", adapter=adapter, sleep=t._noop_sleep))
    assert adapter.calls == 1
    touched = {name for name, _, _ in reads}
    assert not touched & {decks.TEXT_COLLECTION, decks.CANDIDATES_COLLECTION}, touched
    audit_reads = [projection for name, projection, _ in reads if name == "audits"]
    assert audit_reads and all(set(p) - {"_id"} <= {"id", "company_name",
                                                       "structure_reading_consent", "owner"} for p in audit_reads)
    assert DECK_SENTINEL not in adapter.payloads[0] and DECK_FILE not in adapter.payloads[0]


def test_the_deck_parser_still_has_no_link_to_the_structure_path_or_the_gateway():
    """app.structures may read deck structures and call the gateway; the deck package imports neither."""
    for path in sorted((BACKEND / "app" / "decks").rglob("*.py")):
        for module, name in _imports(ast.parse(path.read_text(encoding="utf-8"))):
            assert "structures" not in re.split(r"[.\s\"'()]+", module), f"{path.name} imports {module}"


# The --diagnostic file of scripts/consistency_run.py, beside its report: it holds the cell text of the 10 public
# test decks (the script refuses any other deck), so the check below leaves it out by this name only.
DIAGNOSTIC_SUFFIX = "_diagnostic.md"


def test_the_consistency_run_prints_and_reports_no_text_sent_to_the_model(monkeypatch, tmp_path, capsys):
    """Rule 17 for scripts/consistency_run.py: its progress lines and its report (kept in docs/test-runs) carry
    deck file names, counts, rates, tokens and cost, never the text of a structure. Where passes disagree and why
    items are unverified are told in closed words only: deck, page, type, cell id, field names, reasons.
    The one file excluded, by name, is the --diagnostic file beside the report, which holds cell text and is
    written for the 10 public test decks only (spec section 11)."""
    pytest.importorskip("pdfplumber")
    import importlib.util
    import tempfile
    from app.decks import parser
    from app.llm.schemas import DECK_TYPES
    from app.structures import items, redact
    spec = importlib.util.spec_from_file_location("consistency_run", BACKEND.parent / "scripts" / "consistency_run.py")
    script = importlib.util.module_from_spec(spec)
    spec.loader.exec_module(script)

    class Drifting(script.FakeAdapter):
        """Each pass reads a little differently, so both diagnostic sections have rows."""
        def __init__(self):
            super().__init__()
            self.calls = {}

        def complete(self, **kwargs):
            reply, tokens_in, tokens_out = super().complete(**kwargs)
            n = self.calls[kwargs["user_payload"]] = self.calls.get(kwargs["user_payload"], 0) + 1
            body = json.loads(reply)
            labels = [{**label, "unit": None if n == 2 else label["unit"]} for label in body["labels"]]
            if n == 3 and labels:
                labels[-1] = {**labels[-1], "metric": "other"}
            return json.dumps({**body, "labels": labels}), tokens_in, tokens_out
    monkeypatch.setattr(script, "FakeAdapter", Drifting)
    monkeypatch.setattr(tempfile, "tempdir", str(tmp_path))
    names = ["05-zero2hero.pdf", "02-moz.pdf"]
    report = script.main(["--fake", "--diagnostic", "--deck", names[0], "--deck", names[1],
                          "--out", str(tmp_path / "consistency_report.json")])
    assert report["disagreements"] and report["unverified_items"], "the diagnostic sections have rows to check"
    for row in report["disagreements"]:
        assert set(row) == {"deck", "page", "type", "fields", "same_after_normalisation"}
        assert row["deck"] in names and row["type"] in DECK_TYPES and set(row["fields"]) <= set(script.FIELDS)
    for row in report["unverified_items"]:
        assert set(row) == {"deck", "page", "type", "item", "cell", "reason", "detail", "passes"}
        assert row["deck"] in names and row["type"] in DECK_TYPES and re.fullmatch(r"i\d+", row["item"])
        assert re.fullmatch(r"r\d+c\d+#\d+", row["cell"])
        assert row["reason"] in script.REASONS and row["detail"] in (None, "type Other")
    for counts in [*report["disagreement_fields"].values(), *report["unverified_reasons"].values(),
                   *report["counts"].values()]:
        assert set(counts) <= {"structures", "disagreeing", *script.FIELDS, *script.REASONS, *script.COUNTS}
        assert all(isinstance(n, int) for n in counts.values())
    assert set(report["tokens"]) == {"fixed_prompt", "system_prompt", "output_schema", "empty_message",
                                     "text_and_items_avg", "structures_counted", "billed_input_per_model_read"}
    assert all(isinstance(n, (int, float)) for n in report["tokens"].values())
    assert all(isinstance(report[k], int) for k in ("roadmap_lines", "roadmap_same_pair", "roadmap_dates_rebuilt"))
    assert all(report[k] is None or isinstance(report[k], float) for k in
               ("verifier_match_rate_pct", "verifier_match_rate_financial_pct", "verifier_match_rate_roadmap_pct"))
    files = sorted(tmp_path.glob("consistency_*"))
    excluded = [p for p in files if p.name.endswith(DIAGNOSTIC_SUFFIX)]
    report_md, = [p for p in files if p.suffix == ".md" and p not in excluded]
    assert [p.name for p in excluded] == [report_md.stem + DIAGNOSTIC_SUFFIX], "only the run's own diagnostic file"
    written = capsys.readouterr().out + "".join(p.read_text(encoding="utf-8") for p in files if p not in excluded)
    sent = set()
    for name in names:
        for structure in parser.parse_deck((script.DECKS / name).read_bytes(), name)["structures"]:
            cells, _ = redact.redact_structure(structure["cells"], Path(name).stem, {}, set())
            sent |= {line.split(": ", 1)[1] for line in redact.structure_text(cells).splitlines()}
            listed = items.list_items({**structure, "cells": cells})
            sent |= {item["raw"] for item in listed["items"] if len(item["raw"]) >= 4 and not item["raw"].isdigit()}
    # A cell that only names the company ("zero2hero", zero2hero p11) is part of the deck's file name, which the
    # report carries.
    words = {text for text in sent if len(text) >= 4 and any(c.isalpha() for c in text)
             and not any(text in name for name in names)}
    assert len(words) > 50 and "Gross Profit" in words, "the check sees the decks' cell text"
    assert not sorted(text for text in words if text in written)
    assert any(text in excluded[0].read_text(encoding="utf-8") for text in words), "the excluded file holds cell text"
    with pytest.raises(SystemExit, match="public test decks only"):
        script.main(["--fake", "--diagnostic", "--deck", "client-deck.pdf"])


# ---------------------------------------------------------------------------
# Claim matching (docs/specs/claim-matching.md section 9, CLAUDE.md rules 14 and 17). Python only: the matching
# module has no link to the gateway, the two monthly engine series (revenue, customers) and every register field
# stay out of every narrative payload and out of everything the gateway reads, and a log line about the register
# holds counts per label, never a value or a gate text.
# ---------------------------------------------------------------------------
from app import claim_matching  # noqa: E402

SERIES_SENTINEL = 7123456.78            # a monthly revenue figure only the matcher may read
SERIES_FILE = "Secret_Billing_Export.xlsx"
GATE_SENTINEL = "the Falcon hiring plan"
MATCHING_SERVER_ONLY = ("revenue_series", "customers_series")
# Names of the register that the gateway must never mention: the analyst's inputs, the row's fields, the series.
REGISTER_NAMES = frozenset({"claim_inputs", "gate_sentence", "gate_budget_decision", "gate_threshold", "gate_date",
                            "observed_source", "evidence_label", "value_at_stake_arr", "claim_matching",
                            "turnover_choices", "turnover_as", "turnover_reason", "turnover_state", "deck_take_rate",
                            "implied_take_rate", "implied_take_rate_source", "turnover_suggested", "deck_revenue_note", "period_basis", "inconsistencies",
                            "turnover_views", "file_note", "NO_FILE_PERIOD",
                            "inconsistent_dates", "mixed_slide", "USE_OF_FUNDS_NOTE", "MIXED_SLIDE", "ACQUISITION_NOTE", "ACQUISITION_CHANNELS",
                            "UNTESTED_TYPES", "channel", "auto_accepted", "auto_note", "AUTO_REVOKED", "auto_approval", "auto_citation", "AUTO_LABEL", "ANALYST_LABEL", "auto_views", "entered_by_python", "needs_a_word", "MATCHED_STATUSES", "_enter_matched_claims", "_entry_plan", "_with_entries", "BEAT_PREFIX", "is_beat", "flag_text", "flagged_rows",
                            "qualifier", "QUALIFIERS", "QUALIFIED_REASON", "qualified_label", "definition_check", "definition_question",
                            "definition_views", "definition_rows", "definition_basis", "DEFINITION_QUESTION", "DEFINITION_OPEN_REASON",
                            "DEFINITION_OPEN_REASON_RATE", "DEFINITION_MISMATCH", "BIG_GAP_RATE_POINTS",
                            # docs/specs/evidence-labels.md: the confidence column, the statuses, their explainers and grouping
                            "extraction_confidence", "with_confidence", "confidence_level", "EVIDENCE_STATUSES", "MEMO_GROUPS", "QUALIFIED_NO_FIGURE",
                            "EXTRACTION_EXPLAINER", "EVIDENCE_EXPLAINER", "POPULATION_PREFIX", "NOT_VERIFIED",
                            # docs/specs/dates-and-periods.md: coverage, the growth basis and their reasons
                            "period_coverage", "months_missing", "growth_basis", "analysis_date", "as_of_by", "as_of_file", "as_of_cue",
                            "NOT_COVERED_REASON", "PARTLY_REASON", "PERIOD_NOT_COVERED", "AFTER_SOURCE_DETAIL", "GROWTH_UNKNOWN_REASON",
                            "GROWTH_WINDOW_REASON", "GROWTH_BASE_REASON", "GROWTH_RULES", "GROWTH_CUES",
                            # docs/specs/claim-types-and-correction.md: the claim fields, the edit log, Add from deck and their reasons
                            "population", "figure_type", "required_unknown", "edit_log", "deck_analyst", "claim_fields",
                            "candidate_fields", "REQUIRED_REASON", "TARGET_REASON", "POPULATION_REASON", "POPULATION_WORDS",
                            "TARGET_CUE", "EDIT_REASONS", "ADDED_FROM_DECK", "add_from_deck", "deck_page",
                            # docs/specs/integrity-coverage.md: the coverage, its conclusion and the cross-deck reading
                            "integrity_mod", "assess", "data_requests", "completed_text", "cross_deck_inconsistencies",
                            "INSUFFICIENT", "INSUFFICIENT_NOTE", "COMPLETED_TEXT",
                            # docs/specs/smoothness.md: the deviation limit and the window
                            "SMOOTHNESS_LIMIT_POINTS", "SMOOTHNESS_MONTHS",
                            # docs/specs/claim-groups.md: the relationships between claims, the saved choices and their texts
                            "group_id", "group_role", "relationship", "group_choices", "claim_groups", "choose_group", "undo_group_choice",
                            "VAS_IN_TOTAL", "VAS_PENDING", "VAS_REPEAT", "VAS_RUN_RATE", "NOTE_TOTAL", "NOTE_PART", "NOTE_ALSO", "NOTE_REPEAT", "NOTE_PENDING",
                            "NOTE_DIFFERS", "NOTE_PARTIAL", "NOTE_AMBIGUOUS", "NOTE_TAKEN", "NOTE_CONFLICT", "NOTE_ALTERNATIVE", "NOTE_CHOOSE",
                            "REASON_TOTAL", "NO_EXPOSURE", "SUBORDINATE", "context_hash", "annotate", "relate", "GroupChoice",
                            "has_open_question", "scope_period", "is_repeat_of_listed",
                            # docs/design/segment-partitions.md: the segment a label names, the segment parts and their notes
                            "label_segment", "SEGMENT_PART", "SEGMENT_LIMIT", "NOTE_SEGMENTS", "NOTE_SEGMENT_PART", "NOTE_SEGMENTS_DIFFER",
                            "segment_parts", "segment_set_by",
                            # George, 2026-10-27: the segment slices, the oversize card, the page reader's segment, the undated figures of a page
                            "SLICE", "KIND_SIZE", "NOTE_SLICE", "NOTE_OVERSIZE", "REASON_SIZE", "page_segment", "undated_kept", "undated_by_page",
                            "keep_undated", "SOURCE_AMBIGUOUS_NOTED", "label_is_title",
                            # docs/specs/crm-dedup-clustering.md: the won-deal counts, the repeated rows and the clustering bands
                            "won_deal_closes", "rows_list", "CLUSTERING_DAYS", "CLUSTERING_MIN_DEALS", "CLUSTERING_FEW",
                            "CLUSTERING_AMBER_PCT", "CLUSTERING_RED_PCT", "CLUSTERING_FINDING",
                            # docs/specs/claim-review-conclusion.md: the analyst judgment and the contradictions outside the five
                            "judgment", "reasons", "unresolved_questions", "other_contradictions", "other_contradictions_text",
                            "SCOPE_SENTENCE", "NO_OTHER_CONTRADICTIONS",
                            # docs/design/simplified-flow.md: the gate written from the claim, its status, the note and the computed ratings
                            "gate_status", "GATE_PROPOSED", "GATE_SET", "AT_MOST_TYPES", "NO_CLAIMED_VALUE", "claim_name", "ic_inputs.note",
                            "selection_key", "rate_reliability", "rate_growth", "default_thesis", "unresolved_questions", "mark_key_gates", "gate_text",
                            "SELECTION_SENTENCE", "SELECTION_CHANGED", "NOT_RATED", "BENCHMARK", "OPEN_QUESTION", "OPEN_QUESTIONS", "first_quarterly_review",
                            # issue #136: the citation beside an observed figure, the period that was not read, the contract value
                            "observed_citation", "period_words", "citation", "PERIOD_NOT_READ", "CONTRACT_VALUE", "CONTRACT_REASON",
                            "CONTRACT_TERM", "CONTRACT_LABEL", "_type_choice_open", "metric_labels",
                            # 2026-10-20: the period that cannot be paired, the unit-economics reason, the stale marker
                            "PERIOD_UNCLEAR", "period_unclear", "UNIT_ECONOMICS_REASON", "UNIT_ECONOMICS", "results_stale",
                            # 2026-10-23 (engine contract 8): the monthly-churn reason and the metric each recurring type is tested against
                            "CHURN_RATE_UNCLEAR_REASON", "RECURRING_METRIC",
                            # docs/specs/forecast-and-payback.md: the calculation lines, the forecast reason, the annual payback and the S&M scope
                            "calc_lines", "calc_details", "sm_scope", "net_new_arr", "FORECAST_REASON", "FORECAST_NO_CALC", "FORECAST_METRICS",
                            "FORECAST_BASES", "FORECAST_NOTE", "FOUR_COLUMNS", "forecast_calc", "payback_definition", "diagnose", "claimed_definition",
                            "ANNUAL_ANALYSIS", "ANNUAL_KEYS", "ANNUAL_NOT_CALCULATED", "ANNUAL_NO_YEAR", "POPULATION_REASON_FC", "RULE_CONFLICTS",
                            "is_forecast_row", "has_forecast", "evidence_cell", "result_cell", "source_cell",
                            *MATCHING_SERVER_ONLY})
# The functions that log about the register, and the only things their log calls may name.
REGISTER_LOGGERS = frozenset({"_claim_rows", "claim_register", "claim_register_csv", "update_claim_inputs", "answer_turnover",
                              "answer_definition_check", "choose_group", "undo_group_choice"})
REGISTER_LOG_ARGS = frozenset({"audit_id", "counts", "len(rows)"})      # exactly these expressions, as format arguments
# The turnover answer logs two closed words (Literal fields of the request), never a claim or a figure.
TURNOVER_LOG_ARGS = frozenset({"audit_id", "payload.answer", "payload.reason"})
# The Definition check answer (claim-matching.md section 4.1) logs the run and yes or no (a Literal of the request), nothing else.
DEFINITION_LOG_ARGS = frozenset({"audit_id", "payload.answer"})
# A saved or undone choice among the candidates of a group question (claim-groups.md section 5) logs the run and the kind of question, a closed word.
GROUP_LOG_ARGS = frozenset({"audit_id", "kind"})


def _results_with_the_series():
    results = copy.deepcopy(RESULTS)
    results["revenue_series"] = {"months": ["2026-11", "2026-12"], "segments": ["Enterprise"],
                                 "data": [{"month": "2026-11", "total": 1.5, "Enterprise": 1.5},
                                          {"month": "2026-12", "total": SERIES_SENTINEL, "Enterprise": SERIES_SENTINEL}],
                                 "source": {**SOURCE, "file": SERIES_FILE}}
    results["customers_series"] = {"months": ["2026-12"], "segments": ["Enterprise"], "data": [{"month": "2026-12", "total": 4242,
                                                                                               "Enterprise": 4242}],
                                   "source": {**SOURCE, "file": SERIES_FILE}}
    return results


def _db_with_the_series():
    db = t.FakeDB()
    doc = copy.deepcopy(t.RESULTS_DOC)
    doc["results"] = _results_with_the_series()
    db["audits"].docs.append(doc)
    return db


def test_the_two_monthly_series_are_declared_server_only_and_in_no_step_slice():
    assert set(MATCHING_SERVER_ONLY) <= SERVER_ONLY_TOP_LEVEL
    assert not set(MATCHING_SERVER_ONLY) & gateway.OUTBOUND_FIELDS
    for step in gateway.STEP_CONFIG:
        sliced = gateway._slice_for_step(_results_with_the_series(), step)
        assert not set(MATCHING_SERVER_ONLY) & set(sliced), f"{step} slices a matching series"


def test_the_two_monthly_series_never_reach_a_narrative_payload():
    db = _db_with_the_series()
    for step in gateway.STEP_CONFIG:
        computed = asyncio.run(gateway.load_computed_results(db, RUN_ID, step))
        assert set(MATCHING_SERVER_ONLY) <= set(db["audits"].docs[0]["results"]), "fixture: the series are in the stored results"
        sent = json.dumps(gateway.build_outbound(computed, {}), ensure_ascii=False)
        for needle in (*MATCHING_SERVER_ONLY, str(SERIES_SENTINEL), "7,123,456", "7123456", SERIES_FILE, "4242"):
            assert needle not in sent, f"{step}: {needle!r} reached the outbound payload"
    adapter = t.FakeAdapter(replies=[json.dumps(NARRATIVE)])
    asyncio.run(gateway.generate_narrative(db, RUN_ID, "growth_engine", adapter=adapter, sleep=t._noop_sleep))
    assert adapter.calls == 1
    for needle in (*MATCHING_SERVER_ONLY, "7123456", "7,123,456", SERIES_FILE):
        assert needle not in adapter.payloads[0], f"{needle!r} reached the provider"
    assert db["audits"].docs[0]["results"]["revenue_series"] == _results_with_the_series()["revenue_series"], \
        "Mongo keeps the series for the matcher"


@pytest.mark.parametrize("source", ["claim_matching", "claim_groups"])
def test_the_matching_module_imports_nothing_from_the_gateway_or_a_provider_and_no_gateway_module_imports_it(source):
    path = BACKEND / "app" / f"{source}.py"
    offenders = []
    for module, name in _imports(ast.parse(path.read_text(encoding="utf-8"))):
        parts = set(re.split(r"[.\s\"'()]+", module)) | {name}
        if parts & {"llm", "gateway", "anthropic", "structures", "<dynamic>", "server"}:
            offenders.append(f"{module} -> {name}")
    assert not offenders, "the matching module reaches for the gateway:\n  " + "\n  ".join(offenders)
    code = (f"import sys, app.{source}; "
            "print(sorted(m for m in sys.modules if m.startswith('app.llm') or m.split('.')[0] == 'anthropic'))")
    out = subprocess.run([sys.executable, "-c", code], cwd=BACKEND, capture_output=True, text=True, check=True)
    assert out.stdout.strip() == "[]", out.stdout
    for gateway_file in sorted((BACKEND / "app" / "llm").glob("*.py")):
        for module, name in _imports(ast.parse(gateway_file.read_text(encoding="utf-8"))):
            assert "claim_matching" not in re.split(r"[.\s\"'()]+", module) and name != "claim_matching", \
                f"{gateway_file.name} imports the matching module"


def test_the_gateway_names_no_register_field_input_or_series():
    offenders = []
    for path in sorted((BACKEND / "app" / "llm").glob("*.py")):
        tree = ast.parse(path.read_text(encoding="utf-8"))
        offenders += [f"{path.name}:{node.lineno}: {node.value!r}" for node in ast.walk(tree)
                      if isinstance(node, ast.Constant) and node.value in REGISTER_NAMES]
        offenders += [f"{path.name}:{node.lineno}: {ast.unparse(node)}" for node in ast.walk(tree)
                      if isinstance(node, (ast.Name, ast.Attribute)) and getattr(node, "id", getattr(node, "attr", None)) in REGISTER_NAMES]
    assert not offenders, "the gateway names the claim register:\n  " + "\n  ".join(offenders)
    assert not REGISTER_NAMES & decks.GATEWAY_READABLE_FIELDS
    assert not set(claim_matching.FIELDS) & decks.GATEWAY_READABLE_FIELDS - {"claim_type", "status", "unit"}, \
        "the readable candidate fields stay the six of the narrative path"


def test_the_gateway_reads_no_analyst_input_register_row_or_series():
    """Dynamic: a stored candidate carrying the analyst's gate and the stored series, every gateway path run: what the gateway
    reads is limited to the readable fields, and nothing it reads or sends holds the input or the series."""
    db, reads = _recording_db()
    db[decks.CANDIDATES_COLLECTION].docs[0]["claim_inputs"] = {"c1": {"gate_budget_decision": GATE_SENTINEL, "segment": "Enterprise",
                                                                      "gate_threshold": 4242.4242}}
    db["audits"].docs[0]["results"] = _results_with_the_series()
    adapter = t.FakeAdapter(replies=[json.dumps(NARRATIVE)])
    asyncio.run(gateway.generate_narrative(db, RUN_ID, "growth_engine", adapter=adapter, sleep=t._noop_sleep))
    asyncio.run(gateway.read_cached_narrative(db, RUN_ID, "growth_engine"))
    asyncio.run(gateway.narratives_for_run(db, RUN_ID))
    asyncio.run(gateway.disclosure_for_run(db, RUN_ID))
    asyncio.run(gateway.usage_for_run(db, RUN_ID))
    for name, projection, _ in reads:
        if name == decks.CANDIDATES_COLLECTION:
            assert {k for k, v in (projection or {}).items() if v and k != "_id"} <= decks.GATEWAY_READABLE_FIELDS
    sent = adapter.payloads[0]
    for needle in (GATE_SENTINEL, "4242.4242", "claim_inputs", "7123456", SERIES_FILE):
        assert needle not in sent, f"{needle!r} reached the provider"
        assert needle not in json.dumps([doc for name, _, doc in reads if name == decks.CANDIDATES_COLLECTION], default=str)


def test_the_gateway_reads_no_saved_group_choice_and_the_provider_payload_holds_none():
    """A saved choice among the candidates of a group question (claim-groups.md section 5) sits on the audit under hashes: no gateway path reads
    it, and nothing it sends holds its key, its context or the relationship text of a claim."""
    db, reads = _recording_db()
    db["audits"].docs[0]["group_choices"] = {"feedfacecafe0001": {"kind": "total", "chosen": "deadbeefdeadbeef", "context": "0123456789abcdef",
                                                                  "saved_at": "2026-10-26T10:00:00+00:00"}}
    db[decks.CANDIDATES_COLLECTION].docs[0]["relationship"] = GATE_SENTINEL
    adapter = t.FakeAdapter(replies=[json.dumps(NARRATIVE)])
    asyncio.run(gateway.generate_narrative(db, RUN_ID, "growth_engine", adapter=adapter, sleep=t._noop_sleep))
    asyncio.run(gateway.narratives_for_run(db, RUN_ID))
    asyncio.run(gateway.disclosure_for_run(db, RUN_ID))
    for name, projection, _ in reads:
        if name == "audits":
            assert "group_choices" not in {k for k, v in (projection or {}).items() if v}, "no gateway read asks for the saved choices"
    sent = adapter.payloads[0]
    for needle in ("group_choices", "feedfacecafe0001", "deadbeefdeadbeef", "0123456789abcdef", GATE_SENTINEL, "relationship"):
        assert needle not in sent, f"{needle!r} reached the provider"
    assert not {"group_id", "group_role", "relationship"} & decks.GATEWAY_READABLE_FIELDS


def test_a_relationship_note_and_an_exposure_reason_name_a_rank_and_a_page_never_a_file_or_deck_text():
    """The notes of a repeat, a part under another set, a part waiting for a conflict and a total (claim-groups.md sections 3 and 4) are built
    from the rank and the page label of a row and closed words; the file name and the snippet of a deck never enter them."""
    from test_claim_groups import annual, halves, quarters, related
    secret = "client_secret_board.pptx"
    rows = [annual(1_000_000, cid="a1"), annual(1_004_000, cid="a2", rank=2), annual(1_300_000, cid="a3", rank=3), *quarters([240_000, 250_000, 250_000, 260_000]),
            *halves([480_000, 520_000])]
    for r in rows:
        r.update(deck_file=secret, snippet=f"{secret} Revenue FY2024 EUR secret-snippet")
    related(rows)
    texts = [r["relationship"] or "" for r in rows] + [r["vas_reason"] or "" for r in rows]
    assert any("Same figure as" in t for t in texts) and any("waiting for the choice" in t for t in texts)
    for text in texts:
        assert secret not in text and "secret-snippet" not in text, text
    done = [annual(1_000_000), *quarters([240_000, 250_000, 250_000, 260_000]), *halves([480_000, 520_000])]
    for r in done:
        r.update(deck_file=secret)
    related(done)
    assert any("also reconcile" in (r["relationship"] or "") for r in done) and not any(secret in (r["relationship"] or "") for r in done)


SEGMENT_FIELDS = ("segment", "relationship", "gate_sentence")       # the segment, and the sentences built from it (the gate names the segment: "Revenue (Product A) must ...")


def test_the_segment_a_label_names_is_the_segment_field_and_the_notes_built_from_it_and_no_other_field():
    """docs/design/segment-partitions.md section 1: the cited label of a money figure may become the row's segment, the one deck text a register
    row holds (computed on read, never stored, logged or sent to a model). A sentinel in the label reaches that field and the sentences built from
    it (the gate sentence, the relationship note); the snippet, the file name, the reasons, the gloss and every other field never hold it."""
    from app.decks import claims as deck_claims
    sentinel, secret = "Zephyr Holdings", "Secret_Board_Deck.pptx"

    def cand(cid, value, label):
        c = {"id": cid, "status": "approved", "claim_type": "revenue", "value": value, "currency": "EUR", "target_date": "FY2026",
             "period_text": "FY2026", "snippet": f"{secret} revenue EUR {value}", "label_from": label, "file": secret, "order": int(cid[1:]),
             "sources": [{"file": secret, "slide": int(cid[1:])}]}
        return deck_claims.resolve_period(c, 12)
    rows = claim_matching.build_register([cand("c1", 100.0, "Total revenue"), cand("c2", 60.0, sentinel), cand("c3", 40.0, "Other Holdings")], RESULTS,
                                         {"fiscal_year_end": 12, "reporting_currency": "EUR", "fx": {"EUR": 1.0}})
    got = {r["claim_id"]: r for r in rows}
    assert got["c2"]["segment"] == sentinel and got["c2"]["group_role"] == "segment_part", "fixture: the label became the segment of a part"
    for r in rows:
        assert sentinel not in json.dumps({k: v for k, v in r.items() if k not in SEGMENT_FIELDS}, default=str), r["claim_id"]
        assert secret not in json.dumps({k: v for k, v in r.items() if k not in ("deck_file",)}, default=str), r["claim_id"]
        assert sentinel not in (r["relationship"] or "") or r["claim_id"] == "c2", "only the segment's own note names it"
    assert not {"segment", "segment_set_by", "segment_parts"} & decks.GATEWAY_READABLE_FIELDS


def test_the_sales_question_the_made_of_line_and_the_legend_items_hold_no_deck_text_outside_their_own_fields(caplog):
    """The claims list changes of 2026-10-26: the Sales reasons are closed words whatever the deck says; the answer is kept under the deck id; the
    "part of <metric> <period> total" line holds the metric and the period only; a legend item reaches the candidate's own correction options and,
    once clicked, the segment, never the stored page items, the Diagnostics lines, a log line or a model."""
    import logging

    from app import claim_groups
    from app.decks import claims as deck_claims
    from app.structures import pages as structure_pages
    sentinel = "Zephyr Holdings"
    c = {"id": "c1", "status": "approved", "claim_type": "revenue", "value": 100.0, "currency": "EUR", "target_date": "FY2020", "period_text": "FY2020",
         "snippet": f"{sentinel} EUR 100", "label_from": "Sales", "file": "Secret_Board_Deck.pptx", "deck_id": "d1", "order": 1, "basis": "actual",
         "sources": [{"file": "Secret_Board_Deck.pptx", "slide": 1}]}
    deck_claims.resolve_period(c, 12)
    base = {"fiscal_year_end": 12, "reporting_currency": "EUR", "fx": {"EUR": 1.0}}
    for choices in ({}, {"d1": {"as": "revenue"}}, {"d1": {"as": "bookings"}}):
        row, = claim_matching.build_register([c], RESULTS, {**base, "sales_choices": choices})
        assert sentinel not in json.dumps(row, default=str), choices
        assert row["reason"] in (claim_matching.SALES_ASK_REASON, claim_matching.BOOKINGS_REASON) or sentinel not in row["reason"]
    questions = claim_matching.sales_questions([c], RESULTS, base)
    assert set(questions[0]) == {"deck_id", "file", "claims"} and sentinel not in json.dumps(questions)
    assert not {"made_of", "sales_questions", "sales_choices", "adds_up_to"} & decks.GATEWAY_READABLE_FIELDS
    # the made-of line
    parts = [{**c, "id": "c2", "value": 60.0, "label_from": sentinel, "order": 2, "snippet": "Revenue EUR 60"},
             {**c, "id": "c3", "value": 40.0, "label_from": "Other Holdings", "order": 3, "snippet": "Revenue EUR 40"},
             {**c, "id": "c4", "value": 100.0, "label_from": "Total revenue", "order": 4, "snippet": "Revenue EUR 100"}]
    rows = claim_matching.build_register(parts, RESULTS, base)
    for r in rows:
        made = claim_groups.made_of(r, rows)
        assert sentinel not in (made["part_of"] or "") and "Other Holdings" not in (made["part_of"] or "")
    # a legend item reaches no stored page item, no Diagnostics line, no log line
    kept = structure_pages._without_text({"corrections": [
        {"flag": "label mismatch", "field": "segment", "original": "r9c1", "original_text": None,
         "options": [{"value": sentinel, "ref": "r7c1", "segment": sentinel}]}]})
    assert sentinel not in json.dumps(kept)
    with caplog.at_level(logging.DEBUG):
        claim_matching.sales_questions([c], RESULTS, base)
    assert sentinel not in caplog.text


def test_the_heading_metric_is_one_closed_word_and_never_the_title_text():
    """A title that names a metric gives the candidate `heading_metric`: revenue, sales, ARR or MRR, never the title (which may hold a company name)."""
    from app.structures import association
    sentinel = "Zephyr Holdings"
    assert association.heading_word([f"{sentinel} Revenue by Channel"]) == "revenue"
    assert association.heading_word([f"{sentinel} growth"]) is None and association.heading_word(["Revenue and ARR"]) is None
    assert "heading_metric" not in decks.GATEWAY_READABLE_FIELDS
    assert set(claim_matching.HEADING_METRIC) == {"revenue", "ARR", "MRR"}


def test_the_register_log_lines_name_counts_per_label_and_the_run_only():
    """Static: every log call in a register function names the run, the label counts or a row count, as format arguments
    and nothing else (no row, no claim, no gate, no value)."""
    tree = ast.parse((BACKEND / "server.py").read_text(encoding="utf-8"))
    found = []
    for fn in ast.walk(tree):
        if not (isinstance(fn, (ast.FunctionDef, ast.AsyncFunctionDef)) and fn.name in REGISTER_LOGGERS):
            continue
        for node in ast.walk(fn):
            if isinstance(node, ast.Call) and isinstance(node.func, ast.Attribute) and node.func.attr in _LOG_LEVELS \
                    and isinstance(node.func.value, ast.Name) and "log" in node.func.value.id.lower():
                found.append((fn.name, ast.unparse(node)))
                assert not node.keywords and isinstance(node.args[0], ast.Constant) and "{" not in node.args[0].value, ast.unparse(node)
                for arg in node.args[1:]:
                    allowed = {"answer_turnover": TURNOVER_LOG_ARGS, "answer_definition_check": DEFINITION_LOG_ARGS,
                               "choose_group": GROUP_LOG_ARGS, "undo_group_choice": GROUP_LOG_ARGS}.get(fn.name, REGISTER_LOG_ARGS)
                    assert ast.unparse(arg) in allowed, f"{fn.name}: {ast.unparse(node)} logs {ast.unparse(arg)}"
    assert found, "the register logs its label counts"
    module = ast.parse((BACKEND / "app" / "claim_matching.py").read_text(encoding="utf-8"))
    assert not [n for n in ast.walk(module) if isinstance(n, ast.Name) and n.id in ("logger", "logging", "log", "print")], \
        "the matching module logs and prints nothing"


def test_the_fx_reason_names_two_currency_codes_and_the_new_candidate_fields_are_not_readable():
    """The reason of a claim with no rate holds the claim's currency code and the reporting currency, nothing from the deck;
    period_basis and the inconsistency pairs (values and places) stay out of what the gateway may read."""
    import re
    row = {"id": "x1", "status": "approved", "claim_type": "revenue", "value": 5e6, "currency": "GBP", "target_date": "2021",
           "period_text": "FY2021", "snippet": "Zephyr ARR GBP 5M", "file": "deck.pptx", "order": 0, "sources": [{"file": "deck.pptx", "slide": 3}]}
    from app.decks import claims as deck_claims
    deck_claims.resolve_period(row, 12)
    out = claim_matching.build_register([row], RESULTS, {"fiscal_year_end": 12, "reporting_currency": "EUR", "fx": {"EUR": 1.0}})[0]
    assert re.fullmatch(r"FX rate needed: [A-Z]{3}→[A-Z]{3}", out["reason"]), out["reason"]
    assert not {"period_basis", "inconsistencies", "inconsistent_dates"} & decks.GATEWAY_READABLE_FIELDS


def test_the_beat_reason_is_the_closed_prefix_and_the_gloss_and_holds_no_deck_text():
    """A beat outside tolerance (claim-matching.md section 4, 2026-10-13) is Verified with reason "beat: observed beats the claim
    by {gloss}": the gloss is a figure and closed words (section 3), so a sentinel in the snippet or the label reaches no field."""
    sentinel = "Zephyr Holdings"
    row = {"id": "x1", "status": "approved", "claim_type": "revenue", "value": 4800.0, "currency": "EUR", "target_date": "2026-12",
           "period_text": "Dec 2026", "snippet": f"{sentinel} ARR EUR 4,800", "label_from": "ARR", "file": "deck.pptx",
           "order": 0, "sources": [{"file": "deck.pptx", "slide": 3}]}
    from app.decks import claims as deck_claims
    deck_claims.resolve_period(row, 12)
    results = {**_results_with_the_series(), "mrr_series": {"months": ["2026-12"], "segments": [], "source": SOURCE,
                                                            "data": [{"month": "2026-12", "total": 500.0}]}}
    out = claim_matching.build_register([row], results, {"fiscal_year_end": 12, "reporting_currency": "EUR", "fx": {"EUR": 1.0}})[0]
    assert (out["evidence_label"], out["gap_kind"]) == ("Verified", "beat"), out["reason"]
    assert out["reason"] == claim_matching.BEAT_PREFIX + out["gloss"]
    assert re.fullmatch(r"€[\d,]+ higher \(beat\)", out["gloss"]), out["gloss"]
    assert sentinel not in json.dumps(out, default=str)


def test_the_definition_check_question_and_reasons_hold_figures_and_closed_words_never_deck_text():
    """claim-matching.md section 4.1: the question is two figures in the metric's unit and closed words; the reasons are closed
    sentences. A sentinel in the snippet or the label reaches no field of the row, asked or answered, nor the deck list's view."""
    sentinel = "Zephyr Holdings"
    row = {"id": "x1", "status": "approved", "claim_type": "revenue", "value": 100.0, "currency": "EUR", "target_date": "2026-12",
           "period_text": "Dec 2026", "snippet": f"{sentinel} ARR EUR 100", "label_from": "ARR", "file": "deck.pptx",
           "order": 0, "sources": [{"file": "deck.pptx", "slide": 3}]}
    from app.decks import claims as deck_claims
    deck_claims.resolve_period(row, 12)
    results = {**_results_with_the_series(), "mrr_series": {"months": ["2026-12"], "segments": [], "source": SOURCE,
                                                            "data": [{"month": "2026-12", "total": 500.0}]}}
    settings = {"fiscal_year_end": 12, "reporting_currency": "EUR", "fx": {"EUR": 1.0}}
    asked = claim_matching.build_register([row], results, settings)[0]
    assert re.fullmatch(r"Observed €[\d,]+, claimed €[\d,]+\. Same metric and period\?", asked["definition_question"])
    assert asked["reason"] == claim_matching.DEFINITION_OPEN_REASON.format(share="50%")
    texts = [asked]
    for answer in claim_matching.DEFINITION_ANSWERS:
        answered = claim_matching.build_register([{**row, "claim_inputs": {"x1": {"definition_check": answer, "definition_basis": claim_matching.definition_basis(asked)}}}], results, settings)[0]
        assert answered["reason"] in (claim_matching.DEFINITION_MISMATCH, claim_matching.BEAT_PREFIX + answered["gloss"])
        texts.append(answered)
    texts.append(claim_matching.definition_views([row], results, settings))
    rate = {**row, "id": "x2", "claim_type": "sales", "value": 60.0, "unit": "%", "currency": None,
            "snippet": f"{sentinel} win rate 60%", "label_from": f"{sentinel} win rate"}
    rated = claim_matching.build_register([rate], {**results, "win_rate": {"win_rate_pct": 0.40, "source": SOURCE}}, settings)[0]
    assert rated["reason"] == claim_matching.DEFINITION_OPEN_REASON_RATE.format(points="10"), "a rate: 20 points, asked in points"
    texts += [rated, claim_matching.definition_basis(rated)]
    assert sentinel not in json.dumps(texts, default=str)
    assert not {"definition_check", "definition_question"} & decks.GATEWAY_READABLE_FIELDS


def test_the_qualified_users_reason_holds_two_closed_words_and_the_qualifier_never_reaches_the_gateway_or_a_log(caplog):
    """A customers noun with a qualifier is Users (deck-parser.md section 2, George, 2026-10-13). The register reason names the
    qualifier and the noun, both from closed lists (QUALIFIERS, the customers keywords); the rest of the row label (a sentinel
    here) reaches no register field. `qualifier` is not gateway-readable: run over a database holding the candidate, nothing
    the gateway reads, sends or logs carries the sentinel or the reason."""
    import logging
    from app.decks import claims as deck_claims
    sentinel = "Zephyr"
    assert "qualifier" not in decks.GATEWAY_READABLE_FIELDS
    found, = deck_claims.line_candidates(f"{sentinel} registered institutions 40", [(0, 40, {"file": "deck.pptx", "slide": 3,
                                                                                              "kind": "text"})])
    assert (found["claim_type"], found["qualifier"]) == ("users", "registered") and sentinel in found["snippet"], "fixture"
    row = {**found, "id": "x1", "status": "approved", "file": "deck.pptx", "order": 0, "target_date": "2026-12"}
    deck_claims.resolve_period(row, 12)
    out = claim_matching.build_register([row], RESULTS, {"fiscal_year_end": 12, "reporting_currency": "EUR", "fx": {"EUR": 1.0}})[0]
    nouns = r"customers?|clients?|paying users?|accounts?|compan(?:y|ies)|agenc(?:y|ies)|subscribers?|institutions?"
    words = "|".join(re.escape(q) for q in deck_claims.QUALIFIERS)
    # no paying-customer figure in these results: Unverified (a stated figure the app cannot test), with the reason that names the population
    closed = re.escape(claim_matching.QUALIFIED_NO_FIGURE).replace(re.escape("{label}"), f"(?:{words}) (?:{nouns})")
    assert re.fullmatch(closed, out["reason"]), out["reason"]
    assert out["evidence_label"] == "Unverified" and sentinel not in json.dumps(out, default=str)
    db, reads = _recording_db()
    db[decks.CANDIDATES_COLLECTION].docs.append({**row, "audit_id": RUN_ID})
    adapter = t.FakeAdapter(replies=[json.dumps(NARRATIVE)])
    with caplog.at_level(logging.DEBUG):
        asyncio.run(gateway.generate_narrative(db, RUN_ID, "growth_engine", adapter=adapter, sleep=t._noop_sleep))
    returned = json.dumps([doc for _, _, doc in reads], ensure_ascii=False, default=str)
    logs = caplog.text + "".join(r.getMessage() for r in caplog.records)
    for where, text in (("read by the gateway", returned), ("sent to the provider", adapter.payloads[0]), ("logged", logs)):
        assert sentinel not in text and '"qualifier"' not in text, f"the row label or its qualifier was {where}"
        assert "is not paying customers" not in text, f"the register reason was {where}"


def test_the_period_unclear_unit_economics_and_stale_texts_are_closed_sentences():
    """2026-10-20: three fixed texts, none built from a deck line, a figure or a file name, so none carries client text."""
    from app import run_audit
    assert claim_matching.PERIOD_UNCLEAR == "Period unclear – confirm"
    assert claim_matching.UNIT_ECONOMICS_REASON == "Unit economics – not computed by the app"
    assert run_audit.STALE_LABEL == "Stale – re-run audit"


def test_the_required_field_target_and_population_reasons_are_closed_sentences_with_no_deck_text():
    """claim-types-and-correction.md sections 2 and 6: the three reasons name field names, closed population words and nothing from
    the deck; a sentinel in the snippet or the label reaches no field of the row."""
    sentinel = "Zephyr Holdings"
    from app.decks import claims as deck_claims
    base = {"id": "x1", "status": "approved", "currency": None, "file": "deck.pptx", "order": 0, "label_from": f"{sentinel} label",
            "sources": [{"file": "deck.pptx", "slide": 3}]}
    rows = []
    for extra in ({"claim_type": "sales", "value": 45.0, "unit": "days", "target_date": None, "snippet": f"{sentinel} sales cycle 45 days"},
                  {"claim_type": "revenue", "value": 1.5, "currency": "EUR", "target_date": "2026-11", "snippet": f"{sentinel} target revenue"},
                  {"claim_type": "users", "value": 4000.0, "unit": "users", "target_date": "2026-12", "population": "registered",
                   "snippet": f"{sentinel} users"}):
        c = {**base, **extra}
        deck_claims.resolve_period(c, 12)
        rows.append(claim_matching.build_register([c], _results_with_the_series(),
                                                  {"fiscal_year_end": 12, "reporting_currency": "EUR", "fx": {"EUR": 1.0}})[0])
    # issue #136: a period that was not read has its own closed sentence, the other unknown fields follow it
    assert re.fullmatch(r"Period not read – set it in the claim row(; required field unknown: [a-z ,]+; set it with Edit in the claims list)?",
                        rows[0]["reason"]), rows[0]["reason"]
    assert rows[1]["reason"] == claim_matching.TARGET_REASON
    assert rows[2]["reason"] == claim_matching.POPULATION_REASON.format(population="registered users", measured="paying customers")
    assert sentinel not in json.dumps(rows, default=str)
    assert not {"population", "figure_type", "required_unknown", "edit_log"} & decks.GATEWAY_READABLE_FIELDS


# The two Add from deck endpoints log the run and one count, nothing of the page (claim-types-and-correction.md section 6).
DECK_PAGE_LOG = {"deck_page": ("deck page: run_id=%s lines=%d", ["audit_id", "len(lines)"]),
                 "add_from_deck": ("add from deck: run_id=%s figures=%d", ["audit_id", "len(found)"])}


def test_the_page_view_and_add_from_deck_log_the_run_and_a_count_only():
    tree = ast.parse((BACKEND / "server.py").read_text(encoding="utf-8"))
    seen = {}
    for fn in ast.walk(tree):
        if isinstance(fn, (ast.FunctionDef, ast.AsyncFunctionDef)) and fn.name in DECK_PAGE_LOG:
            seen[fn.name] = [(node.args[0].value, [ast.unparse(a) for a in node.args[1:]]) for node in ast.walk(fn)
                             if isinstance(node, ast.Call) and isinstance(node.func, ast.Attribute) and node.func.attr in _LOG_LEVELS
                             and isinstance(node.func.value, ast.Name) and "log" in node.func.value.id.lower() and not node.keywords]
    assert seen == {name: [(fmt, list(args))] for name, (fmt, args) in DECK_PAGE_LOG.items()}


def test_the_use_of_funds_reason_is_a_closed_sentence_and_mixed_slide_and_the_note_never_reach_the_gateway_or_a_log(caplog):
    """Use of funds (deck-parser.md section 2): the register reason is one fixed sentence, `mixed_slide` is a layout flag of the
    parser, and the inconsistency note holds a category name from the deck. None of them is gateway-readable; run over a
    database holding a use-of-funds slide with a sentinel category, nothing the gateway reads, sends or logs carries it."""
    import logging
    import test_deck_parser as dp
    sentinel = "Zephyr Holdings Marketing"
    assert not {"mixed_slide", "inconsistencies", "note"} & decks.GATEWAY_READABLE_FIELDS
    slide = dp._funds_slide([("Cost control", 50), ("Team hiring", 30), (sentinel, 20)], [63, 38])
    found = [c for c in dp._funds_candidates([slide]) if c["unit"] == "%"]
    assert any(sentinel in json.dumps(c["inconsistencies"]) for c in found), "fixture: the note holds the deck's category"
    mixed = [c for c in dp._funds_candidates([["Use of funds", "Revenue by region", sentinel + " (60%)", "Cost control (40%)"]])
             if c["unit"] == "%"]
    assert mixed and all(c["mixed_slide"] for c in mixed), "fixture: a mixed slide"
    # the register row: the fixed sentence, no deck text
    row = {"id": "x1", "status": "approved", "claim_type": "use_of_funds", "value": 20, "unit": "%", "currency": None,
           "target_date": None, "snippet": f"{sentinel} (20%)", "label_from": sentinel, "file": "deck.pptx", "order": 0,
           "sources": [{"file": "deck.pptx", "slide": 22}]}
    out = claim_matching.build_register([row], RESULTS, {"fiscal_year_end": 12, "reporting_currency": "EUR", "fx": {"EUR": 1.0}})[0]
    assert out["reason"] == claim_matching.USE_OF_FUNDS_NOTE and sentinel not in json.dumps(out, default=str)
    # the gateway over a database that holds them
    db, reads = _recording_db()
    for i, c in enumerate([*found, *mixed]):
        db[decks.CANDIDATES_COLLECTION].docs.append({**c, "audit_id": RUN_ID, "id": f"u{i}", "status": "approved"})
    adapter = t.FakeAdapter(replies=[json.dumps(NARRATIVE)])
    with caplog.at_level(logging.DEBUG):
        asyncio.run(gateway.generate_narrative(db, RUN_ID, "growth_engine", adapter=adapter, sleep=t._noop_sleep))
    returned = json.dumps([doc for _, _, doc in reads], ensure_ascii=False, default=str)
    logs = caplog.text + "".join(r.getMessage() for r in caplog.records)
    for where, text in (("read by the gateway", returned), ("sent to the provider", adapter.payloads[0]), ("logged", logs)):
        assert sentinel not in text, f"the deck's category was {where}"
        assert claim_matching.USE_OF_FUNDS_NOTE not in text and "mixed slide" not in text, f"the register reason or tag was {where}"


def test_the_acquisition_channels_reason_is_a_closed_sentence_and_the_platform_name_never_reaches_the_gateway_or_a_log(caplog):
    """Acquisition channels (chat-upload.md section 17): the register reason is one fixed sentence; the platform name is deck text
    kept on the candidate as `channel`, which the gateway may not read. Run over a database holding a funnel slide with a
    sentinel platform, nothing the gateway reads, sends or logs carries it."""
    import logging
    import test_deck_parser as dp
    sentinel = "ZephyrChat"
    assert "channel" not in decks.GATEWAY_READABLE_FIELDS
    found = [c for c in dp._funds_candidates([["User acquisition funnel", f"{sentinel}(30K)", "Telegram(12K)"]])
             if c["claim_type"] == "acquisition_channels"]
    assert {c["channel"] for c in found} == {sentinel, "Telegram"}, "fixture: the candidate holds the deck's platform name"
    row = {**next(c for c in found if c["channel"] == sentinel), "id": "x1", "status": "approved", "file": "deck.pptx", "order": 0}
    out = claim_matching.build_register([row], RESULTS, {"fiscal_year_end": 12, "reporting_currency": "EUR", "fx": {"EUR": 1.0}})[0]
    assert out["reason"] == claim_matching.ACQUISITION_NOTE and sentinel not in json.dumps(out, default=str)
    db, reads = _recording_db()
    for i, c in enumerate(found):
        db[decks.CANDIDATES_COLLECTION].docs.append({**c, "audit_id": RUN_ID, "id": f"ch{i}", "status": "approved"})
    adapter = t.FakeAdapter(replies=[json.dumps(NARRATIVE)])
    with caplog.at_level(logging.DEBUG):
        asyncio.run(gateway.generate_narrative(db, RUN_ID, "growth_engine", adapter=adapter, sleep=t._noop_sleep))
    returned = json.dumps([doc for _, _, doc in reads], ensure_ascii=False, default=str)
    logs = caplog.text + "".join(r.getMessage() for r in caplog.records)
    for where, text in (("read by the gateway", returned), ("sent to the provider", adapter.payloads[0]), ("logged", logs)):
        assert sentinel not in text, f"the deck's platform name was {where}"
        assert claim_matching.ACQUISITION_NOTE not in text, f"the register reason was {where}"


def test_the_turnover_fields_hold_closed_words_figures_and_sources_never_deck_text():
    """The row of a turnover claim whose snippet holds a sentinel: the sentinel is in no register field, the audit's stored
    answer is keyed by a hash, and the take-rate scan yields a yes or no."""
    sentinel = "Zephyr Holdings take rate is 2.5%"
    row = {"id": "x1", "status": "approved", "claim_type": "revenue", "value": 5e6, "currency": "EUR", "target_date": "2021",
           "period_text": "FY2021", "snippet": f"GMV EUR 5M. {sentinel}", "label_from": None, "file": "deck.pptx", "order": 0,
           "sources": [{"file": "deck.pptx", "slide": 3}]}
    out = claim_matching.build_register([row], RESULTS, {"fiscal_year_end": 12, "reporting_currency": "EUR", "fx": {"EUR": 1.0},
                                                         "deck_take_rate": claim_matching.mentions_take_rate(sentinel)})[0]
    assert "Zephyr" not in json.dumps(out, default=str) and "take rate" not in json.dumps(out, default=str)
    # the take-rate scan pre-selects Volume; the row asks until the analyst answers (claim-matching.md section 11 point 4)
    assert (out["turnover_state"], out["turnover_suggested"]) == ("ask", "volume")
    assert out["turnover_note"] == "Revenue or volume? Confirm below" and out["deck_revenue_note"] is None
    key = claim_matching.turnover_key(row)
    assert len(key) == 16 and re.fullmatch(r"[0-9a-f]{16}", key)         # the stored answer's whole key; the endpoint test in
    # test_turnover_api.py checks the stored audit and candidate documents for deck text
    assert claim_matching.mentions_take_rate(sentinel) is True


# ---------------------------------------------------------------------------
# Chat upload and column mapping (docs/specs/chat-upload.md section 9). The assertions run through the real endpoints
# on the in-memory Mongo stub and the fake adapter; the fixtures are synthetic (rule 15).
# ---------------------------------------------------------------------------
import logging  # noqa: E402

import test_chat_upload as chat  # noqa: E402

FIXTURE_CELLS = ("Alder Works", "Birch Labs", "Cedar Foods", "Dune Print", "Elm Studio", "Fern Tools", "Birch Labs Ltd",
                 "C-100", "C-101", "C-102", "C-103", "C-104", "Quarterly revenue export", "Large", "Small", "Gold", "Silver")
FIXTURE_FILES = ("blank_vs_zero", "mixed_currency", "duplicate_ids_header_row3")


def _all_documents(db):
    return {name: json.dumps(col.docs, default=str) for name, col in db._cols.items()}


def test_chat_the_mapping_text_holds_every_column_a_masked_sample_and_the_rules_hint_from_at_most_20_rows(api):
    rows = "\n".join(f"Customer {i},2024-01-{i % 27 + 1:02d},{100 + i},EUR,note {i} {chr(97 + i % 26)}" for i in range(60))
    chat.upload(api, "wide.csv", ("Customer ID,Invoice Date,Amount,Currency,Remarks\n" + rows + "\n").encode())
    text = chat.model_text(api)
    parsed = structure_redact.parse_column_text(text)
    assert [c["text"] for c in parsed["headers"]] == ["Customer ID", "Invoice Date", "Amount", "Currency", "Remarks"]
    assert parsed["profiles"][5]["distinct"] == 20 and parsed["profiles"][1]["distinct"] == 20, "built from the first 20 data rows"
    assert parsed["samples"] == {2: ["2024-01-01", "2024-01-02", "2024-01-03"], 3: ["100", "101", "102"]}
    assert parsed["hints"] == [(1, "customer_id", 100), (2, "invoice_date", 100), (3, "amount", 100), (4, "currency", 100)]
    assert not re.search(r"Customer \d|note \d|EUR", text), "no text cell value"


@pytest.mark.parametrize("name", ["duplicate_ids_header_row3.csv", "blank_vs_zero.csv", "mixed_currency.csv"])
def test_chat_no_fixture_customer_name_and_no_text_cell_reaches_the_model(api, name):
    chat.upload(api, name)
    text = chat.model_text(api)
    assert structure_redact.column_text_problem(text) is None
    for cell in FIXTURE_CELLS:
        assert cell.lower() not in text.lower(), cell


def test_chat_a_customer_alias_column_sends_a_profile_only_even_when_its_values_are_numbers(api):
    csv = "Date,Amount,Currency,Customer ID,Customer ID\n" + "\n".join(
        f"2024-01-{i + 1:02d},{50 + i},EUR,{70000 + i},{81000 + i}" for i in range(8)) + "\n"
    chat.upload(api, "numeric_ids.csv", csv.encode())
    text = chat.model_text(api)
    parsed = structure_redact.parse_column_text(text)
    assert {4, 5} <= set(parsed["profiles"]) and not {4, 5} & set(parsed["samples"])
    assert "70000" not in text and "81000" not in text and structure_redact.column_text_problem(text) is None


@pytest.mark.parametrize("name", [n + ".csv" for n in FIXTURE_FILES])
def test_chat_no_collection_holds_the_sent_text_after_an_upload_with_consent(api, name):
    chat.upload(api, name)
    texts = [json.loads(p)["text"] for p in api.state["adapter"].payloads]
    assert texts, "the model was asked"
    stored = _all_documents(api.db)
    for text in texts:
        for line in text.split("\n"):
            assert line not in stored.get("llm_structures", "") + stored.get("llm_calls", "") + stored.get("column_mappings", ""), line
        assert all(text not in blob for blob in stored.values())
    calls = api.db["llm_calls"].docs
    assert all(set(c) <= {"run_id", "step", "prompt_version", "model", "input_tokens", "output_tokens", "estimated_cost_usd",
                          "image_tokens", "image_cost_usd", "cache_hit", "status", "unmatched_numbers", "timestamp", "content_hash", "deck_id", "http_status",
                          "error_type", "structure_type"} for c in calls)


def test_chat_log_lines_hold_no_file_name_header_cell_value_company_name_or_note(api, caplog):
    caplog.set_level(logging.INFO)
    note = "Adj is not the type"
    for name in FIXTURE_FILES:
        chat.upload(api, name + ".csv", replace="true")
    chat.upload(api, "blank_vs_zero.csv", replace="true")
    api.post(f"/api/audits/{chat.AUDIT}/datasets/revenue/decisions", json=[
        {"column": "Adj", "action": "correct", "field": "revenue_type", "reason": "other", "note": note}])
    api.post(f"/api/audits/{chat.AUDIT}/usage", json={"screen": "mapping", "rejected_extension": "pptx"})
    api.post(f"/api/audits/{chat.AUDIT}/compute")
    api.get(f"/api/audits/{chat.AUDIT}/blockers")
    api.get("/api/usage/totals")
    api.request("DELETE", f"/api/audits/{chat.AUDIT}", json={"confirm": chat.COMPANY})
    log = "\n".join(r.getMessage() for r in caplog.records)
    assert "upload: run_id=" in log and "decisions: run_id=" in log, "the new lines were written"
    secrets = [*FIXTURE_FILES, *FIXTURE_CELLS, "Customer ID", "Invoice Date", "Adj", chat.COMPANY, "ENG-2026-041",
               note, "pptx.", "wide.csv"]
    for secret in secrets:
        assert secret not in log, f"{secret!r} reached a log line"


def test_chat_the_counters_and_the_totals_hold_no_file_name_cell_value_or_company_name_and_header_text_only_in_a_note(api):
    note = "Adj is not the type"
    chat.upload(api, "blank_vs_zero.csv")
    api.post(f"/api/audits/{chat.AUDIT}/datasets/revenue/decisions", json=[
        {"column": "Adj", "action": "correct", "field": "revenue_type", "reason": "other", "note": note}])
    api.post(f"/api/audits/{chat.AUDIT}/usage", json={"rejected_extension": "pptx"})
    usage = api.db["audits"].docs[0]["usage"]
    totals = api.get("/api/usage/totals").json()
    for name, doc in (("audits.usage", usage), ("totals", totals)):
        blob = json.dumps(doc)
        for secret in (*FIXTURE_FILES, *FIXTURE_CELLS, chat.COMPANY, "ENG-2026-041", "Customer ID", "Invoice Date"):
            assert secret not in blob, f"{secret!r} in {name}"
        assert blob.count("Adj") == blob.count(note) == 1, f"header text only inside the kept note ({name})"
    assert usage["columns"]["reasons"] == {"other": 1}
    # A reason code outside the fixed list is refused and counted nowhere.
    refused = api.post(f"/api/audits/{chat.AUDIT}/datasets/revenue/decisions", json=[
        {"column": "Adj", "action": "correct", "field": "revenue_type", "reason": "Customer ID is wrong"}])
    assert refused.status_code == 400 and "Customer ID is wrong" not in json.dumps(api.db["audits"].docs[0]["usage"])


@pytest.mark.parametrize("note,refused", [
    ("Revenue 2024 is wrong", True), ("a" * 61, True), ("Adj is not the type", False),
    ("see blank_vs_zero", True), ("see blank_vs_zero.csv", True), ("Alder is wrong", True),
    ("Fixture Target Ltd", True), ("ENG-2026-041", True),
])
def test_chat_an_other_note_is_refused_for_a_digit_a_file_name_a_cell_text_a_name_or_length(api, note, refused):
    chat.upload(api, "blank_vs_zero.csv")
    r = api.post(f"/api/audits/{chat.AUDIT}/datasets/revenue/decisions", json=[
        {"column": "Adj", "action": "correct", "field": "revenue_type", "reason": "other", "note": note}])
    assert (r.status_code == 400) is refused, r.text


def test_chat_a_note_quoting_a_label_of_a_file_read_by_recognition_only_is_refused(api):
    """Review M5 on PR 135: such a file keeps no rows, only its recognised sheets, and the note filter reads their text too."""
    import server
    from mapping_shapes import FIXTURES as MAPPING_FILES
    api.db["audits"].docs[0]["structure_reading_consent"] = False
    api.db["audits"].docs[0]["as_of_month"] = "2025-03-31"
    summary = b"Revenue by product line (EUR),,,\nProduct line,Jan-25,Feb-25,Mar-25\nPlatform licences,10,11,12\nSupport plans,20,21,22\n"
    assert chat.upload(api, "summary.csv", summary).json()["recognition_only"] is True
    assert api.db["datasets"].docs[0]["rows"] == [], "the label is in no row: only the stored sheet holds it"
    chat.upload(api, "hubspot_deals.csv", (MAPPING_FILES / "hubspot_deals.csv").read_bytes())
    url = f"/api/audits/{chat.AUDIT}/datasets/crm/decisions"
    refused = api.post(url, json=[{"column": "Deal owner", "action": "correct", "field": None, "reason": "other", "note": "Platform licences is wrong"}])
    assert refused.status_code == 400 and refused.json()["detail"] == server.NOTE_REFUSED
    kept = api.post(url, json=[{"column": "Deal owner", "action": "correct", "field": None, "reason": "other", "note": "That column is not needed"}])
    assert kept.status_code == 200


def test_chat_a_note_that_quotes_a_header_is_kept_including_a_header_holding_a_cell_word_and_one_with_a_digit_is_refused(api):
    head = "Customer ID,Invoice Date,Amount,Currency,Segment,Large Account Flag,Plan 2024\nAlder Works,2024-01-01,5,EUR,Large,x,y\n"
    chat.upload(api, "h.csv", head.encode())
    ok = lambda note, col: api.post(f"/api/audits/{chat.AUDIT}/datasets/revenue/decisions", json=[  # noqa: E731
        {"column": col, "action": "correct", "field": None, "reason": "other", "note": note}])
    assert ok("Large Account Flag is not needed", "Large Account Flag").status_code == 200
    assert ok("Plan 2024 is not needed", "Plan 2024").status_code == 400
    assert ok("Segment is not needed", "Large Account Flag").status_code == 200, "a quoted header with no digit"


def test_chat_a_kept_note_is_found_only_in_the_usage_counters_and_the_totals(api, caplog):
    caplog.set_level(logging.INFO)
    note = "Adj is not the type"
    chat.upload(api, "blank_vs_zero.csv")
    api.post(f"/api/audits/{chat.AUDIT}/datasets/revenue/decisions", json=[
        {"column": "Adj", "action": "correct", "field": "revenue_type", "reason": "other", "note": note}])
    api.get("/api/usage/totals")
    holders = {name for name, blob in _all_documents(api.db).items() if note in blob}
    assert holders == {"audits"}, "not in llm_calls, column_mappings, llm_structures or any other collection"
    assert note in json.dumps(api.db["audits"].docs[0]["usage"]["other_notes"]) and note in json.dumps(api.get("/api/usage/totals").json())
    assert not any(note in p for p in api.state["adapter"].payloads), "not in a model call"
    assert note not in "\n".join(r.getMessage() for r in caplog.records)
    assert note not in json.dumps({k: v for k, v in api.db["audits"].docs[0].items() if k != "usage"}, default=str)


def test_chat_revenue_reconciliation_is_in_no_narrative_slice_and_never_reaches_the_provider():
    results = copy.deepcopy(RESULTS)
    cite = {"revenue_file": SOURCE, "pnl": SOURCE}
    results["revenue_reconciliation"] = {
        "available": True, "first": "2024-01", "last": "2024-12", "tolerance_pct": 0.02, "file_total": 7654321.0,
        "pnl_total": 7000000.0, "gap": 654321.0, "gap_pct": 0.0935, "blocker": True, "source": cite,
        "by_month": [{"month": "2024-01", "revenue_file": 1.0, "pnl": 1.0, "gap": 31337.0, "gap_pct": 0.0935, "source": cite}]}
    for step in gateway.STEP_CONFIG:
        assert "revenue_reconciliation" not in gateway._slice_for_step(results, step), step
    db = _db()
    db["audits"].docs[0]["results"] = results
    adapter = t.FakeAdapter(replies=[json.dumps(NARRATIVE)])
    asyncio.run(gateway.generate_narrative(db, RUN_ID, "growth_engine", adapter=adapter, sleep=t._noop_sleep))
    for needle in ("revenue_reconciliation", "7654321", "654321", "31337", "0.0935", "9.35", "by_month", "gap_pct"):
        assert needle not in adapter.payloads[0], f"{needle!r} reached the provider"


# ---------------------------------------------------------------------------
# Interface contract (docs/specs/interface-contracts.md section 4, CLAUDE.md rules 14 and 23): the stored metrics are checked
# against backend/schemas/metrics.py before any prompt is built. The new results key `contract_version` is declared server-only
# above; the one new reason string is a constant sentence; the one new log line carries paths and error types.
# ---------------------------------------------------------------------------
def test_a_payload_that_breaks_the_contract_reaches_no_provider_and_no_log_line_holds_its_content(caplog):
    import logging
    secrets = ("Zeta Holdings", "4242.4242", "Jane Roe", "Hidden_Billing.xlsx")
    results = cf.stored(arr=cf.arr(source={**SOURCE, "file": secrets[3]}),
                        nrr=cf.nrr(1.04, 110, by_segment={secrets[0]: cf.nrr_group(1.1, 4242.4242)}),
                        missing_data=[{"metric": "m", "status": "Missing", "reason": f"{secrets[2]} left",
                                       "unlocked_by": "x", "file": secrets[3]}])
    db = t.FakeDB()
    doc = copy.deepcopy(t.RESULTS_DOC)
    doc["results"] = results
    db["audits"].docs.append(doc)
    adapter = t.FakeAdapter()
    with caplog.at_level(logging.DEBUG):
        result = asyncio.run(gateway.generate_narrative(db, RUN_ID, "growth_engine", adapter=adapter, sleep=t._noop_sleep))
        asyncio.run(gateway.read_cached_narrative(db, RUN_ID, "growth_engine"))
    assert adapter.calls == 0 and result.reason == gateway.NARRATIVE_NOT_GENERATED
    assert "metrics contract violated" in caplog.text, "fixture: the contract line was logged"
    for text in (caplog.text, result.reason, *[r.getMessage() for r in caplog.records]):
        for secret in secrets:
            assert secret not in text, f"{secret!r} reached a log line or the reason"
    assert db["llm_calls"].docs == [] and db["llm_narratives"].docs == [], "nothing is stored about a call that was not made"


# ---------------------------------------------------------------------------
# Verdict and IC memo (docs/specs/verdict-and-memo.md section 9; CLAUDE.md rules 14 and 17). The gate's metric name, the budget
# decision, the thesis, the ratings and the top 5 are the analyst's inputs: they are stored on the audit and the candidate rows,
# read by two pure modules, and reach no model, no log line and no collection of the gateway.
# ---------------------------------------------------------------------------
import test_claim_register_api as register_api_tests  # noqa: E402

IC_SENTINELS = ("the Falcon hiring plan", "Zephyr pipeline cover", "Quartz thesis plan", "Basalt thesis evidence", "Opal thesis condition",
                "Garnet analyst note", "Jasper stored judgment")      # the last two: the optional note (simplified-flow.md) and a stored judgment no longer read
VERDICT_LOGGERS = frozenset({"get_verdict", "export_memo"})
VERDICT_LOG_FORMATS = {
    "get_verdict": "verdict: run_id=%s outcome=%s top5=%s changed=%s open=%d",
    "export_memo": "memo export: run_id=%s status=%s reason=%s words=%d unmatched=%d"}
VERDICT_LOG_ARGS = frozenset({"audit_id", "ver['outcome_code']", "top5", "ver['top5']['changed']", "ctx['open_questions']", "'refused' if refusal else 'ok'",
                              "refusal.code if refusal else 'none'", "words", "len(refusal.unmatched) if refusal else 0"})   # ast.unparse


def test_the_new_register_fields_and_the_analysts_inputs_are_names_the_gateway_never_mentions():
    names = {"ic_inputs", "gate_metric_name", "gate_direction", "key_gate", "overlaps_with", "evidence_analysis",
             "evidence_source_key", "shortfall", "thesis", "first_quarterly_review", "gap_target_dates"}
    assert names <= set(claim_matching.FIELDS) | {"ic_inputs", "thesis", "first_quarterly_review", "gap_target_dates"}
    offenders = []
    for path in sorted((BACKEND / "app" / "llm").glob("*.py")):
        tree = ast.parse(path.read_text(encoding="utf-8"))
        offenders += [f"{path.name}:{n.lineno}: {n.value!r}" for n in ast.walk(tree) if isinstance(n, ast.Constant) and n.value in names]
        offenders += [f"{path.name}:{n.lineno}: {ast.unparse(n)}" for n in ast.walk(tree)
                      if isinstance(n, (ast.Name, ast.Attribute)) and getattr(n, "id", getattr(n, "attr", None)) in names]
    assert not offenders, "the gateway names the verdict inputs:\n  " + "\n  ".join(offenders)
    assert not names & gateway.OUTBOUND_FIELDS and not names & decks.GATEWAY_READABLE_FIELDS


def test_the_integrity_module_imports_nothing_from_the_gateway_and_logs_nothing_and_the_endpoint_logs_codes_and_counts():
    """integrity-coverage.md section 5: a pure module; the endpoint's one log line holds the run, the conclusion code and two counts."""
    path = BACKEND / "app" / "integrity.py"
    tree = ast.parse(path.read_text(encoding="utf-8"))
    offenders = [f"{m} -> {n}" for m, n in _imports(tree)
                 if (set(re.split(r"[.\s\"'()]+", m)) | {n}) & {"llm", "gateway", "anthropic", "structures", "<dynamic>", "server"}]
    assert not offenders, "integrity.py reaches for the gateway:\n  " + "\n  ".join(offenders)
    assert not [n for n in ast.walk(tree) if isinstance(n, ast.Name) and n.id in ("logger", "logging", "log", "print")]
    code = "import sys, app.integrity; print(sorted(m for m in sys.modules if m.startswith('app.llm') or m.split('.')[0] == 'anthropic'))"
    assert subprocess.run([sys.executable, "-c", code], cwd=BACKEND, capture_output=True, text=True, check=True).stdout.strip() == "[]"
    server_tree = ast.parse((BACKEND / "server.py").read_text(encoding="utf-8"))
    fn = next(f for f in ast.walk(server_tree) if isinstance(f, ast.AsyncFunctionDef) and f.name == "get_integrity")
    calls = [n for n in ast.walk(fn) if isinstance(n, ast.Call) and isinstance(n.func, ast.Attribute) and n.func.attr in _LOG_LEVELS
             and isinstance(n.func.value, ast.Name) and "log" in n.func.value.id.lower()]
    assert [(c.args[0].value, [ast.unparse(a) for a in c.args[1:]]) for c in calls] == [
        ("integrity: run_id=%s conclusion=%s completed=%d applicable=%d", ["audit_id", "out['code']", "out['completed']", "out['applicable']"])]


def test_the_verdict_and_memo_modules_import_nothing_from_the_gateway_a_provider_or_the_server_and_log_nothing():
    for path in (BACKEND / "app" / "verdict.py", BACKEND / "app" / "ic_memo.py"):
        tree = ast.parse(path.read_text(encoding="utf-8"))
        offenders = [f"{m} -> {n}" for m, n in _imports(tree)
                     if (set(re.split(r"[.\s\"'()]+", m)) | {n}) & {"llm", "gateway", "anthropic", "structures", "<dynamic>", "server"}]
        assert not offenders, f"{path.name} reaches for the gateway:\n  " + "\n  ".join(offenders)
        assert not [n for n in ast.walk(tree) if isinstance(n, ast.Name) and n.id in ("logger", "logging", "log", "print")], \
            f"{path.name} logs or prints"
    code = ("import sys, app.verdict, app.ic_memo; "
            "print(sorted(m for m in sys.modules if m.startswith('app.llm') or m.split('.')[0] == 'anthropic'))")
    out = subprocess.run([sys.executable, "-c", code], cwd=BACKEND, capture_output=True, text=True, check=True)
    assert out.stdout.strip() == "[]", out.stdout
    for gateway_file in sorted((BACKEND / "app" / "llm").glob("*.py")):
        for module, name in _imports(ast.parse(gateway_file.read_text(encoding="utf-8"))):
            parts = re.split(r"[.\s\"'()]+", module)
            assert not ({"verdict", "ic_memo"} & set(parts)) and name not in ("verdict", "ic_memo"), \
                f"{gateway_file.name} imports the verdict or memo module"


def test_the_two_new_log_lines_are_fixed_formats_with_codes_and_counts_as_their_only_arguments():
    tree = ast.parse((BACKEND / "server.py").read_text(encoding="utf-8"))
    seen = {}
    for fn in ast.walk(tree):
        if not (isinstance(fn, (ast.FunctionDef, ast.AsyncFunctionDef)) and fn.name in VERDICT_LOGGERS):
            continue
        for node in ast.walk(fn):
            if isinstance(node, ast.Call) and isinstance(node.func, ast.Attribute) and node.func.attr in _LOG_LEVELS \
                    and isinstance(node.func.value, ast.Name) and "log" in node.func.value.id.lower():
                assert not node.keywords and isinstance(node.args[0], ast.Constant), ast.unparse(node)
                seen.setdefault(fn.name, []).append(node.args[0].value)
                for arg in node.args[1:]:
                    assert ast.unparse(arg) in VERDICT_LOG_ARGS, f"{fn.name}: {ast.unparse(node)} logs {ast.unparse(arg)}"
    assert {k: v for k, v in seen.items()} == {k: [v] for k, v in VERDICT_LOG_FORMATS.items()}


def _with_the_analysts_inputs(client, db):
    ids = client.get(f"/api/audits/{register_api_tests.AUDIT}/verdict").json()["verdict"]["top5"]["proposed"]
    base = f"/api/audits/{register_api_tests.AUDIT}"
    assert client.put(f"{base}/ic-inputs", json={"top5": ids}).status_code == 200
    for n, cid in enumerate(ids):
        body = {"gate_threshold": 10.0 + n, "gate_budget_decision": IC_SENTINELS[0], "gate_date": "2024-06-30",
                "gate_metric_name": IC_SENTINELS[1], "gate_direction": "at least"}
        assert register_api_tests._put(client, cid, body).status_code == 200
    r = client.put(f"{base}/ic-inputs", json={"thesis": {"plan": IC_SENTINELS[2], "evidence": IC_SENTINELS[3], "condition": IC_SENTINELS[4]},
                                              "note": IC_SENTINELS[5]})
    assert r.status_code == 200, r.text


def test_the_analysts_inputs_reach_no_log_line_no_gateway_collection_and_no_provider(register_api, caplog):
    client, db = register_api
    base = f"/api/audits/{register_api_tests.AUDIT}"
    with caplog.at_level(logging.DEBUG):
        _with_the_analysts_inputs(client, db)
        for url in ("/verdict", "/memo.md", "/claims.csv", "/claims", "/blockers"):
            client.get(base + url)
    memo = client.get(f"{base}/memo.md")
    assert memo.status_code == 200 and IC_SENTINELS[3] in memo.text and IC_SENTINELS[5] in memo.text, \
        "fixture: the memo is built from the inputs and returned, not logged"
    for needle in IC_SENTINELS:
        assert needle not in caplog.text, f"{needle!r} reached a log line"
    stored = {name: json.dumps(col.docs, default=str) for name, col in db._cols.items()}
    for name, blob in stored.items():
        if name not in ("audits", decks.CANDIDATES_COLLECTION):
            for needle in IC_SENTINELS:
                assert needle not in blob, f"{needle!r} was stored in {name}"
    assert IC_SENTINELS[2] in stored["audits"] and IC_SENTINELS[1] in stored[decks.CANDIDATES_COLLECTION], "fixture: they are stored where they belong"
    assert not db["llm_calls"].docs


def test_the_gateway_reads_no_ic_inputs_and_the_provider_payload_holds_none(register_api):
    db, reads = _recording_db()
    db["audits"].docs[0]["ic_inputs"] = {"thesis": {"plan": IC_SENTINELS[2]}, "ratings": {"growth_engine": "Weak"},
                                         "top5": {"claim_ids": ["c1"], "set_at": "x"}, "first_quarterly_review": "2031-04-05",
                                         "note": IC_SENTINELS[5], "judgment": {"reasons": IC_SENTINELS[6], "unresolved_questions": IC_SENTINELS[6]}}
    db[decks.CANDIDATES_COLLECTION].docs[0]["claim_inputs"] = {"c1": {"gate_metric_name": IC_SENTINELS[1], "key_gate": True,
                                                                      "gate_direction": "at most"}}
    adapter = t.FakeAdapter(replies=[json.dumps(NARRATIVE)])
    asyncio.run(gateway.generate_narrative(db, RUN_ID, "growth_engine", adapter=adapter, sleep=t._noop_sleep))
    asyncio.run(gateway.narratives_for_run(db, RUN_ID))
    asyncio.run(gateway.usage_for_run(db, RUN_ID))
    returned = json.dumps([doc for name, _, doc in reads if name in ("audits", decks.CANDIDATES_COLLECTION)], default=str)
    for needle in (*IC_SENTINELS[1:3], *IC_SENTINELS[5:], "ic_inputs", "2031-04-05", "gate_metric_name", "key_gate", "judgment", "note"):
        assert needle not in returned, f"{needle!r} was read by the gateway"
        assert needle not in adapter.payloads[0], f"{needle!r} reached the provider"


# ---------------------------------------------------------------------------
# Data inventory and data integrity (A1, A2; CLAUDE.md rules 14 and 17). Their payload is built from the stored engine output
# by app/llm/step_payloads.py: the audit's source keys with the engine's generic description, statuses, closed sentences and
# display strings. No file or sheet name, header, cell, customer, segment or row number, and none of the server-only blocks
# (anomalies, the monthly series, the reconciliation's by_month) as such. The new payload fields are in OUTBOUND_FIELDS; the
# narratives add no results key, no log line and no guard word beyond the closed reasons of step_problem.
# ---------------------------------------------------------------------------
import test_step_payloads as steps  # noqa: E402

NEW_STEPS = ("data_inventory", "data_integrity")
STEP_FIELDS = frozenset({"source_list", "inventory", "integrity", "source_key", "description", "category", "in_request_list",
                         "quality_issues", "analyses_blocked", "months_supplied", "months_required", "check", "first_month",
                         "last_month", "revenue_file_total", "pnl_total", "gap", "gap_pct", "tolerance", "window_over_tolerance",
                         "months_over_tolerance", "deviation_points", "limit_points", "result", "threshold_note", "finding"})


def _strings(node):
    if isinstance(node, dict):
        for v in node.values():
            yield from _strings(v)
    elif isinstance(node, list):
        for v in node:
            yield from _strings(v)
    elif isinstance(node, str):
        yield node


def test_the_step_payload_fields_are_allowlisted_and_not_server_only():
    assert STEP_FIELDS <= gateway.OUTBOUND_FIELDS
    assert not STEP_FIELDS & SERVER_ONLY_FIELDS, "a key cannot be both"
    assert not {"anomalies", "revenue_reconciliation", "by_month", "mrr_series", "revenue_series"} & gateway.OUTBOUND_FIELDS


@pytest.mark.parametrize("step", NEW_STEPS)
def test_the_step_payload_holds_no_file_sheet_header_cell_customer_segment_or_row_number(step):
    results = steps.full_results()
    results["mrr_series"]["segments"] = ["Enterprise"]
    results["anomalies"]["revenue_gap_then_resume"] = [steps.CUSTOMER]
    results["missing_data"] = [{"metric": "m", "status": "Missing", "reason": f"{PERSON} left", "unlocked_by": "x",
                                "file": steps.FILE_NAMES[0]}]
    out = steps.payload_of(step, results)
    blob = json.dumps(out, ensure_ascii=False)
    for secret in (*steps.FILE_NAMES, *steps.SHEETS, steps.CUSTOMER, PERSON, "Enterprise", "row_numbers", "0.1", "0.0935"):
        assert secret not in blob, f"{secret!r} reached the {step} payload"
    assert "missing_data" not in out["metrics"] and "questions_for_management" not in out["metrics"], "no reason string leaves"
    # Every string is one the builder wrote: a closed sentence, a status, a label, a key or a figure.
    for text in _strings(out):
        assert not re.search(r"\.(xlsx|csv|pptx)\b", text, re.I), text


@pytest.mark.parametrize("step", NEW_STEPS)
def test_the_step_payload_of_every_engine_run_holds_no_file_sheet_sentinel_or_stored_customer_id(step, engine_runs):
    for spec, inject, results in engine_runs:
        stored = json.dumps(results, ensure_ascii=False)
        files = set(re.findall(r'"file": "([^"]+)"', stored))
        sheets = set(re.findall(r'"sheet": "([^"]+)"', stored))
        customers = set(results["anomalies"]["revenue_gap_then_resume"]) if results.get("anomalies") else set()
        blob = json.dumps(gateway.build_outbound({"metrics": gateway._slice_for_step(results, step)}, {}), ensure_ascii=False)
        for secret in (*files, *sheets, *customers, SENTINEL, HEADER_SENTINEL):
            assert secret not in blob, f"{secret!r} reached the {step} payload of {spec['reporting_currency']}"


def test_the_integrity_payload_holds_the_window_figures_as_display_strings_and_never_the_monthly_block():
    results = steps.full_results()
    for m in results["revenue_reconciliation"]["by_month"]:
        m["revenue_file"], m["pnl"], m["gap"] = 123456.0, 111111.0, 12345.0      # figures only the monthly block holds
    blob = json.dumps(steps.payload_of("data_integrity", results), ensure_ascii=False)
    assert "1,100,000 EUR" in blob and "10%" in blob, "the stored window total and gap, as display strings"
    for needle in ("revenue_reconciliation", "by_month", "123456", "123,456", "111,111", "12,345", "0.1"):
        assert needle not in blob, f"{needle!r} (a month of the monthly block) reached the integrity payload"
    inventory = json.dumps(steps.payload_of("data_inventory", results), ensure_ascii=False)
    for needle in ("1,100,000", "1,000,000", "gap_pct", "window_over_tolerance"):
        assert needle not in inventory, f"{needle!r} reached the inventory payload"


def test_the_clustering_sentence_reaches_the_integrity_payload_and_the_year_end_and_the_deal_rows_never_do():
    """crm-dedup-clustering.md section 2: the gateway reads the audit's fiscal year-end beside the results to set the quarters; the
    model is told the finding and the result, never the year-end, the monthly counts or a row."""
    results = steps.full_results()
    results["won_deal_closes"] = {"by_month": [{"month": "2026-03", "won": 10, "last_10_days": 10},
                                               {"month": "2026-04", "won": 15, "last_10_days": 15}],
                                  "source": steps.cite(1, "rule")}
    db = steps.db_for(results)
    db["audits"].docs[0]["fiscal_year_end"] = 1             # quarters end in Jan, Apr, Jul, Oct: April's closes count
    computed = asyncio.run(gateway.load_computed_results(db, steps.RUN_ID, "data_integrity"))
    blob = json.dumps(gateway.build_outbound(computed, {}), ensure_ascii=False)
    assert "60% of 25 won deals closed in the last 10 days of a quarter (chance: 11%)" in blob and '"Red"' in blob
    for needle in ("fiscal_year_end", "won_deal_closes", "by_month", "last_10_days", "2026-04", steps.FILE_NAMES[1], "row_numbers"):
        assert needle not in blob, f"{needle!r} reached the integrity payload"
    assert "fiscal_year_end" not in gateway.OUTBOUND_TOP_LEVEL | gateway.OUTBOUND_FIELDS


def test_the_repeat_flag_reaches_the_inventory_as_counts_and_never_its_rows():
    results = steps.full_results()
    results["anomalies"]["deals_duplicate_id"] = {"rows": 9, "distinct_ids": 3, "collapsed": 6, "rows_list": [717171, 727272]}
    blob = json.dumps(steps.payload_of("data_inventory", results), ensure_ascii=False)
    assert "9 CRM row(s) share a deal ID with another row (3 deal ID(s)); 6 exact repeat(s) counted once" in blob
    for needle in ("717171", "717,171", "rows_list", "deals_duplicate_id"):
        assert needle not in blob, f"{needle!r} reached the inventory payload"


@pytest.mark.parametrize("step", NEW_STEPS)
def test_a_rejected_or_failed_step_logs_closed_reasons_and_stores_no_payload_text(step, caplog):
    import logging
    db = steps.db_for(steps.full_results())
    reply = steps.GOOD[step](what_this_means="Tier 3 only")
    with caplog.at_level(logging.DEBUG):
        result, _ = steps.run(db, step, reply)
    assert result.narrative_status == "unavailable"
    stored = json.dumps([db["llm_calls"].docs, db["llm_narratives"].docs], default=str)
    for text in (caplog.text, result.reason, stored):
        for secret in (*steps.FILE_NAMES, *steps.SHEETS, steps.CUSTOMER):
            assert secret not in text, f"{secret!r} reached a log line, a reason or a collection"
    assert db["llm_narratives"].docs == [], "a rejected narrative is not stored"
    assert all(set(c) <= {"run_id", "step", "prompt_version", "model", "input_tokens", "output_tokens", "estimated_cost_usd",
                          "image_tokens", "image_cost_usd", "cache_hit", "status", "unmatched_numbers", "timestamp", "http_status",
                          "error_type"}
               for c in db["llm_calls"].docs)


@pytest.mark.parametrize("step", NEW_STEPS)
def test_the_stored_narrative_holds_model_output_only(step):
    db = steps.db_for(steps.full_results())
    steps.run(db, step, steps.GOOD[step]())
    doc = db["llm_narratives"].docs[0]
    assert doc["step"] == step and set(doc["narrative"]) >= {"headline", "table_rows", "next_actions"}
    assert not {"payload", "outbound", "prompt", "system", "metrics"} & set(doc)
    for secret in (*steps.FILE_NAMES, steps.CUSTOMER):
        assert secret not in json.dumps(doc, default=str)




# ---------------------------------------------------------------------------
# Board decks in the data inventory (George, 2026-10-11): the payload carries the number of decks and nothing of them. The
# gateway counts the audit's parsed decks with count_documents on the audit id; it reads no document, text, file name or field.
# ---------------------------------------------------------------------------
def _db_with_decks(n):
    db, reads = _recording_db()
    counts = []

    async def count_documents(flt, _orig=db[decks.TEXT_COLLECTION].count_documents):
        counts.append(flt)
        return await _orig(flt)

    db[decks.TEXT_COLLECTION].count_documents = count_documents
    db[decks.TEXT_COLLECTION].docs.extend({"audit_id": RUN_ID, "file": f"{DECK_FILE}{i}", "blocks": [{"text": DECK_SENTINEL}]}
                                          for i in range(n - 1))
    return db, reads, counts


def test_the_inventory_counts_decks_and_reads_none():
    db, reads, counts = _db_with_decks(3)
    adapter = t.FakeAdapter(replies=[json.dumps(steps.inventory_reply())])
    result = asyncio.run(gateway.generate_narrative(db, RUN_ID, "data_inventory", adapter=adapter, sleep=t._noop_sleep))
    asyncio.run(gateway.read_cached_narrative(db, RUN_ID, "data_inventory"))
    asyncio.run(gateway.narratives_for_run(db, RUN_ID))
    assert counts and all(c == {"audit_id": RUN_ID} for c in counts), counts
    assert not [r for r in reads if r[0] == decks.TEXT_COLLECTION], "a document of the parsed-text collection was read"
    sent = json.loads(adapter.payloads[0])
    deck = next(r for r in sent["metrics"]["inventory"] if r["category"] == "Board decks")
    assert deck == {"category": "Board decks", "status": "Provided", "in_request_list": False, "quality_issues": [],
                    "analyses_blocked": [], "count": "3"}
    assert len(sent["metrics"]["inventory"]) == 7, "all seven categories"
    for needle in (DECK_SENTINEL, DECK_FILE):
        assert needle not in adapter.payloads[0] and needle not in json.dumps(db["llm_narratives"].docs, default=str)
    assert result.narrative_status == "ok", result.reason


def test_no_deck_is_missing_and_on_the_request_list_and_an_unreadable_count_leaves_the_row_out():
    from app.llm import step_payloads
    none = step_payloads.build_inventory(steps.full_results(), 0)["inventory"]
    deck = next(r for r in none if r["category"] == "Board decks")
    assert deck["status"] == "Missing" and deck["in_request_list"] is True and "count" not in deck
    assert len(none) == 7 and len(step_payloads.build_inventory(steps.full_results(), None)["inventory"]) == 6


# ---------------------------------------------------------------------------
# Run audit (CLAUDE.md rule 14): the stage endpoint adds two log lines, the claims it accepts carry one new stored flag, and the
# narrative estimate reads the payload a generation builds. None of them sends or logs anything the gateway did not already.
# ---------------------------------------------------------------------------
import test_run_audit as run_tests  # noqa: E402


def test_the_run_audit_log_lines_hold_the_run_the_stage_the_status_and_an_error_type_only(run_api, monkeypatch, caplog):
    import logging
    run_tests.upload_files(run_api)
    run_tests.confirm_open_columns(run_api)
    run_tests.add_deck(run_api)

    async def broken(audit_id):
        raise ValueError(f"{PERSON} 4242.4242 {FILE}")
    monkeypatch.setattr(run_tests.server, "_run_compute", broken)
    with caplog.at_level(logging.DEBUG):
        for stage in ("mapping", "engine", "claims"):
            run_tests.stage(run_api, stage)
    lines = [r.getMessage() for r in caplog.records if r.getMessage().startswith("run audit")]
    assert len(lines) == 4, lines
    for line in lines:
        assert re.fullmatch(r"run audit(?: stage failed)?: run_id=\S+ stage=(mapping|engine|claims) (status|error)=\w+", line), line
    for secret in (PERSON, "4242.4242", FILE, "revenue.csv", "testco_board"):
        assert secret not in caplog.text


def test_the_auto_accepted_flag_is_a_register_name_the_gateway_never_mentions():
    assert "auto_accepted" in REGISTER_NAMES and "auto_note" in REGISTER_NAMES and "AUTO_REVOKED" in REGISTER_NAMES
    for path in sorted((BACKEND / "app" / "llm").glob("*.py")):
        assert "auto_accepted" not in path.read_text(encoding="utf-8"), path.name


def test_the_narrative_estimate_sends_nothing_and_names_no_file(run_api):
    run_tests.ready_for_review(run_api)
    sent_before = len(run_api.adapter.payloads)
    body = run_api.get(f"/api/runs/{run_tests.AUDIT}/narratives/estimate").text
    assert len(run_api.adapter.payloads) == sent_before, "an estimate is not a call"
    for secret in ("revenue.csv", "crm.csv", "pnl.csv", "testco_board"):
        assert secret not in body


def test_the_period_the_model_is_told_is_two_closed_values_and_never_the_as_of_setting():
    """Decision 7 (15 Oct 2026): evidence_to is a month and last_month_status is one closed word (always Complete since 2026-10-28,
    when the as-of month always counts as complete). The as-of setting, who set it and the requested analysis date are not outbound
    fields."""
    assert {"evidence_to", "last_month_status"} <= gateway.OUTBOUND_TOP_LEVEL
    assert not ({"as_of_month", "as_of_by", "as_of_file", "as_of_cue", "requested_analysis_date"} & (gateway.OUTBOUND_TOP_LEVEL | gateway.OUTBOUND_FIELDS))
    out = gateway.period_fields({"source_period": {"month": "2019-06", "status": "Complete", "complete_by": "as_of",
                                                   "evidence_to": "2019-06"}, "as_of_month": "2026-06"})
    assert out == {"evidence_to": "2019-06", "last_month_status": "Complete"}
    assert gateway.period_fields({}) == {"evidence_to": None, "last_month_status": None}


def test_the_access_module_logs_nothing_and_imports_nothing_from_the_gateway():
    """access-links.md section 4: a token never reaches a log line. The module has no logging call at all, and the server's
    guard and admin routes log nothing of the request's headers or links."""
    tree = ast.parse((BACKEND / "app" / "access.py").read_text())
    imported = {n.module or "" for n in ast.walk(tree) if isinstance(n, ast.ImportFrom)} | \
        {a.name for n in ast.walk(tree) if isinstance(n, ast.Import) for a in n.names}
    assert not any(m.startswith(("app.llm", "logging", "anthropic")) or m == "logging" for m in imported), imported
    assert not [n for n in ast.walk(tree) if isinstance(n, ast.Call) and getattr(n.func, "id", "") == "print"]
    server_tree = ast.parse((BACKEND / "server.py").read_text())
    fns = {"access_guard", "issue_tester", "list_testers", "revoke_tester", "reissue_tester"}
    for fn in (f for f in ast.walk(server_tree) if isinstance(f, (ast.AsyncFunctionDef, ast.FunctionDef)) and f.name in fns):
        calls = [ast.unparse(c) for c in ast.walk(fn) if isinstance(c, ast.Call) and "logger" in ast.unparse(c.func)]
        assert not calls, f"{fn.name} logs: {calls}"


# ---------------------------------------------------------------------------
# The prompt store (docs/specs/prompt-store.md, section 5; CLAUDE.md rules 14 and 26). Prompt text is the methodology: it goes from
# the owner page into `prompts`, from `prompts` into the model request, and nowhere else.
# ---------------------------------------------------------------------------
PROMPT_SENTINEL = "PROMPT-TEXT-SENTINEL-8812"


def _run_with_sentinel_prompt():
    class Seen(t.FakeAdapter):
        def complete(self, *, system, **kwargs):
            self.systems = getattr(self, "systems", []) + [system]
            return super().complete(system=system, **kwargs)

    async def go():
        db = t.make_db()
        await prompt_store.upload(db, "growth_engine", f"Test only. {PROMPT_SENTINEL}")
        adapter = Seen()
        result = await gateway.generate_narrative(db, RUN_ID, "growth_engine", adapter=adapter, sleep=t._noop_sleep)
        return db, adapter, result

    return asyncio.run(go())


def test_prompt_text_reaches_the_provider_as_the_system_prompt_only_and_is_stored_in_prompts_only(caplog):
    import logging
    with caplog.at_level(logging.DEBUG):
        db, adapter, result = _run_with_sentinel_prompt()
    assert adapter.systems == [f"Test only. {PROMPT_SENTINEL}"] and not any(PROMPT_SENTINEL in p for p in adapter.payloads)
    assert PROMPT_SENTINEL not in json.dumps(result.model_dump(), default=str), "the response holds the version, never the text"
    assert PROMPT_SENTINEL not in caplog.text
    holders = {name for name, col in db._cols.items() if PROMPT_SENTINEL in json.dumps(col.docs, default=str)}
    assert holders == {"prompts"}, "no other collection holds prompt text"


def test_the_prompts_document_has_fixed_fields():
    db = t.FakeDB()
    db["prompts"].docs.clear()
    asyncio.run(prompt_store.upload(db, "growth_engine", "Test only."))
    assert set(db["prompts"].docs[0]) == {"name", "version", "text", "uploaded_at", "sha256"}


def test_the_prompt_list_adds_a_character_count_and_never_the_text():
    db = t.FakeDB()
    db["prompts"].docs.clear()
    asyncio.run(prompt_store.upload(db, "growth_engine", f"{PROMPT_SENTINEL} test only"))
    rows = asyncio.run(prompt_store.list_versions(db))
    assert [set(r) for r in rows] == [{"name", "version", "uploaded_at", "sha256", "chars"}]
    assert isinstance(rows[0]["chars"], int) and PROMPT_SENTINEL not in json.dumps(rows)


def test_the_missing_prompt_reason_is_the_closed_sentence_and_holds_no_name_text_or_path():
    async def go():
        db = t.make_db()
        db["prompts"].docs.clear()
        return await gateway.generate_narrative(db, RUN_ID, "growth_engine", adapter=t.FakeAdapter(), sleep=t._noop_sleep)

    assert asyncio.run(go()).reason == prompt_store.NO_PROMPT == "No prompt loaded — upload prompts on the owner page"


def test_no_module_of_the_gateway_reads_a_token_and_the_prompt_store_logs_nothing():
    for path in (BACKEND / "app" / "llm").glob("*.py"):
        text = path.read_text(encoding="utf-8")
        assert "app.access" not in text and "OWNER_TOKEN" not in text and "authorization" not in text.lower(), path.name
    text = (BACKEND / "app" / "llm" / "prompt_store.py").read_text(encoding="utf-8")
    assert "logging" not in text and "logger" not in text and "print(" not in text


def test_the_prompt_upload_route_logs_the_name_and_version_only():
    source = (BACKEND / "server.py").read_text(encoding="utf-8")
    fn = source[source.index("async def admin_upload_prompt"):source.index("@api.post(\"/admin/testers\")")]
    assert re.findall(r'logger\.\w+\((.*)\)', fn) == ['"prompt uploaded: name=%s version=%s", row["name"], row["version"]']


# ---------------------------------------------------------------------------
# The file type (docs/specs/chat-upload.md section 19, 2026-10-16): the model may read spreadsheet sheet names, title rows and header
# rows and the masked sample, with the company name masked. Nothing else, and none of it is kept.
# ---------------------------------------------------------------------------
import mapping_shapes as shapes  # noqa: E402

FILE_TYPE_TITLES = ("Revenue through 31 December 2028", "Accounts through 31 December 2028", "EUR; accounts anonymised")
FILE_TYPE_CONTEXT = {"company_name": "Fixture Target Ltd"}


def _file_type_text(**sheet):
    base = {"name": "Accounts", "titles": [{"row": 2, "col": 1, "text": "Accounts through 2028"}],
            "headers": [{"row": 5, "col": 1, "text": "Account ID"}], "samples": {2: ["100", "200"]},
            "profiles": {1: {"distinct": 6, "length": 8, "shape": "A-0"}}}
    return structure_redact.file_type_text([{**base, **sheet}])


def test_chat_a_file_type_text_passes_with_title_rows_headers_a_masked_sample_and_nothing_else():
    assert gateway.structure_text_problem(_file_type_text(), "file_type", FILE_TYPE_CONTEXT, {}) is None


@pytest.mark.parametrize("sheet,problem", [
    ({"samples": {2: ["Alder Works"]}}, "text_value"),                                   # a text cell value
    ({"samples": {2: ["1", "2", "3", "4"]}}, "too_many_samples"),                        # more than 3 samples
    ({"samples": {1: ["1"]}}, "text_value"),                                             # a sample and a profile for one column
    ({"titles": [{"row": r, "col": 1, "text": "t"} for r in range(1, 8)]}, "too_many_header_rows"),
    ({"headers": [{"row": 5, "col": 1, "text": "a"}, {"row": 6, "col": 1, "text": "b"}]}, "too_many_header_rows"),
    ({"titles": [{"row": 2, "col": 1, "text": "x" * 201}]}, "cell_too_long"),
    ({"titles": [{"row": 2, "col": 1, "text": "Contact jane.doe@example.com"}]}, "redaction_changed"),
    ({"name": "export.xlsx"}, "file_name"),
])
def test_chat_a_file_type_text_with_a_text_value_too_many_rows_a_long_title_a_name_an_address_or_a_file_name_is_refused(sheet, problem):
    assert gateway.structure_text_problem(_file_type_text(**sheet), "file_type", FILE_TYPE_CONTEXT, {}) == problem


def test_chat_a_file_type_text_with_more_than_six_sheets_or_a_line_of_no_kind_is_refused():
    base = {"name": "S", "titles": [], "headers": [{"row": 1, "col": 1, "text": "h"}], "samples": {}, "profiles": {}}
    seven = structure_redact.file_type_text([{**base, "name": f"S{i}"} for i in range(7)])
    assert gateway.structure_text_problem(seven, "file_type", FILE_TYPE_CONTEXT, {}) == "too_many_header_rows"
    assert gateway.structure_text_problem("s1: Accounts\nfree prose about the company", "file_type", FILE_TYPE_CONTEXT, {}) == "not_cells"
    assert gateway.structure_text_problem("r1c1: Account ID", "file_type", FILE_TYPE_CONTEXT, {}) == "not_cells"


def test_chat_the_file_type_reply_is_one_closed_word():
    for reply in ('{"type":"file_type","file_type":"deck"}', '{"type":"file_type","file_type":"crm","note":"x"}', '{"file_type":"crm"}'):
        with pytest.raises(gateway.GatewayError):
            gateway.parse_structure_reply(reply, "file_type", "")
    assert gateway.parse_structure_reply('{"type":"file_type","file_type":"crm"}', "file_type", "").file_type == "crm"


def _recognition_word(word):
    """The reply of the file_type prompt v2 (revenue-file-recognition.md 5.7) that names only the file's type."""
    return json.dumps({"type": "file_recognition", "sheet": 1, "file_type": word, "revenue_row": None, "currency": None, "non_revenue_rows": []})


def _file_type_upload(api):
    api.state["adapter"] = t.FakeAdapter(replies=[_recognition_word("crm"), chat.mapping_reply()])
    return chat.upload(api, "CRM_Anonymous_to_2028.xlsx", shapes.crm_anonymous_to_2028())


def test_chat_the_file_type_text_holds_titles_and_headers_and_no_data_text_or_file_name(api):
    _file_type_upload(api)
    text = chat.model_text(api)
    parsed = structure_redact.parse_recognition_text(text)
    assert any("Accounts through" in c["text"] for c in parsed[0]["titles"]), "the title rows are sent"
    for secret in ("ACC-0001", "Mid-market", "no contradiction found", "Owner A", "Claimed", "CRM_Anonymous", ".xlsx"):
        assert secret not in text, secret
    assert gateway.structure_text_problem(text, "file_recognition", FILE_TYPE_CONTEXT, {}) is None


def test_chat_no_collection_holds_the_file_type_text_and_the_kept_answer_is_a_hash_and_a_word(api, caplog):
    caplog.set_level(logging.DEBUG)
    _file_type_upload(api)
    text = chat.model_text(api)
    stored = _all_documents(api.db)
    for line in text.split("\n"):
        assert line not in stored.get("llm_structures", "") + stored.get("llm_calls", "") + stored.get("column_mappings", ""), line
    for title in FILE_TYPE_TITLES + ("Source inputs", "Annual tracker", "Funnel and evidence"):
        for name in ("llm_structures", "llm_calls", "column_mappings", "audits"):
            assert title not in stored.get(name, ""), (title, name)
    kept = [d for d in api.db["column_mappings"].docs if d.get("kind") == "file_recognition"]
    assert len(kept) == 1 and set(kept[0]) <= {"audit_id", "kind", "file_hash", "prompt_tag", "reply", "saved_at"}
    assert kept[0]["reply"]["file_type"] == "crm" and set(kept[0]["reply"]) <= {
        "sheet", "file_type", "revenue_row", "currency", "non_revenue_rows", "last_actual_column", "boundary_cue"}
    assert kept[0]["reply"].get("last_actual_column") is None and kept[0]["reply"].get("boundary_cue") is None, "a CRM has no boundary"
    log = "\n".join(r.getMessage() for r in caplog.records)
    for title in FILE_TYPE_TITLES + ("Source inputs", "Annual tracker", "Funnel and evidence", "Account ID", "CRM_Anonymous", chat.COMPANY):
        assert title not in log, f"{title!r} reached a log line"
    row = next(c for c in api.db["llm_calls"].docs if c.get("structure_type") == "file_recognition")
    assert set(row) <= LLM_CALLS_KEYS and row["status"] == "ok"


def test_chat_a_title_row_holding_an_address_reaches_the_model_masked(api):
    from openpyxl import Workbook
    import io
    wb = Workbook()
    ws = wb.active
    ws["C2"] = "Monthly pack, contact jane.doe@example.com"
    ws["C3"] = "Fixture Target Ltd management accounts"
    for j, h in enumerate(["Line", "2023", "2024", "2025"]):
        ws.cell(5, 3 + j, h)
    ws.cell(6, 3, "Metric")
    for j in range(3):
        ws.cell(6, 4 + j, 10 + j)
    buf = io.BytesIO()
    wb.save(buf)
    api.state["adapter"] = t.FakeAdapter(replies=[json.dumps({"type": "file_type", "file_type": "unknown"})])
    chat.upload(api, "x.xlsx", buf.getvalue())
    text = chat.model_text(api)
    assert "jane.doe" not in text and "[redacted]" in text and "[email]" in text
    assert "Fixture Target Ltd" not in text, "the company's own name is masked in the file-type text (chat-upload.md section 19)"


def _name_workbook(sheet_name, title):
    from openpyxl import Workbook
    import io
    wb = Workbook()
    ws = wb.active
    ws.title = sheet_name
    ws["A1"] = title
    for j, h in enumerate(["Line", "2023", "2024", "2025"]):
        ws.cell(3, 1 + j, h)
    ws.cell(4, 1, "Metric")
    for j in range(3):
        ws.cell(4, 2 + j, 10 + j)
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def test_chat_the_company_name_in_a_sheet_name_or_a_title_row_never_reaches_the_model_and_the_unit_marker_stays(api):
    api.state["adapter"] = t.FakeAdapter(replies=[_recognition_word("not_revenue")])
    chat.upload(api, "x.xlsx", _name_workbook("Fixture Target Ltd P&L", "Fixture Target Ltd \u2013 April 2024 Report, EUR thousands"))
    text = chat.model_text(api)
    assert "fixture target" not in text.lower(), "the company name is absent from the model payload"
    titles = [c["text"] for s in structure_redact.parse_recognition_text(text) for c in s["titles"]]
    assert any("[redacted]" in x and "EUR thousands" in x for x in titles), titles
    assert gateway.structure_text_problem(text, "file_recognition", FILE_TYPE_CONTEXT, {}) is None


def test_chat_a_file_type_text_still_holding_the_company_name_is_refused():
    text = _file_type_text(name="Fixture Target Ltd", titles=[{"row": 1, "col": 1, "text": "Fixture Target Ltd report"}])
    assert gateway.structure_text_problem(text, "file_type", FILE_TYPE_CONTEXT, {}) == "redaction_changed"


# ---------------------------------------------------------------------------
# Forecast tables read as claims (chat-upload.md section 20, issue #114): rules only, no model call; the line labels, the sheet and
# the cells are deck-candidate evidence the gateway may not read; the register reason is one fixed sentence; the log line holds the
# run and two counts.
# ---------------------------------------------------------------------------
import server  # noqa: E402
from app import column_rules as cr_mod, forecast_calc, forecast_claims  # noqa: E402

FORECAST_SENTINEL = "Zephyr Holdings Revenue"           # a line label: client text, only on the candidate
FORECAST_LOG_ARGS = frozenset({"audit_id", "len(made)", "saved"})


def _forecast_file():
    head = ("Zephyr Holdings plan,,,\nEUR thousands,,,\n,,,\nLine,2026F,2027F,2028F\n"
            f"{FORECAST_SENTINEL},4321,5432,6543\nCost of sales,-1,-2,-3\nCustomers,11,22,33\n")
    return head.encode()


def test_forecast_claims_are_read_and_saved_with_no_model_call_and_no_log_line_holds_a_label_value_or_file_name(api, caplog):
    caplog.set_level(logging.DEBUG)
    out = chat.upload(api, "Zephyr_plan.csv", _forecast_file(), dtype="pnl").json()     # the type as the analyst picks it (S6)
    assert out["forecast_claims"]["count"] == 6, "fixture: two plan lines over three years"
    # A P&L statement makes its one sheet_structure call at upload (section 26.4); the claims are read by rules (section 20.4).
    made = chat.adapter_of(api).calls
    assert all(json.loads(p)["type"] == "sheet_structure" for p in chat.adapter_of(api).payloads)
    saved = api.post(f"/api/audits/{chat.AUDIT}/datasets/{out['dtype']}/forecast-claims", json={"rows": [5, 7]})
    assert saved.json() == {"saved": 6, "lines": 2}
    assert chat.adapter_of(api).calls == made, "reading and saving the claims calls no model (section 20.4)"
    log = "\n".join(r.getMessage() for r in caplog.records)
    assert "forecast claims: run_id=" in log
    for secret in (FORECAST_SENTINEL, "Zephyr", "4321000", "4321", "Zephyr_plan.csv", "E5", "Customers"):
        assert secret not in log, f"{secret!r} reached a log line"
    stored = [c for c in api.db[decks.CANDIDATES_COLLECTION].docs if c.get("origin") == forecast_claims.ORIGIN]
    assert any(c["snippet"] == FORECAST_SENTINEL for c in stored), "fixture: the label is kept on the candidate as evidence"
    for name in ("llm_calls", "llm_structures", "column_mappings", "audits"):
        assert FORECAST_SENTINEL not in json.dumps(api.db[name].docs, default=str), f"the label reached {name}"


def test_the_forecast_claims_log_line_is_a_fixed_format_with_the_run_and_two_counts():
    tree = ast.parse((BACKEND / "server.py").read_text(encoding="utf-8"))
    fn = next(n for n in ast.walk(tree) if isinstance(n, ast.AsyncFunctionDef) and n.name == "confirm_forecast_claims")
    calls = [n for n in ast.walk(fn) if isinstance(n, ast.Call) and isinstance(n.func, ast.Attribute) and n.func.attr in _LOG_LEVELS
             and isinstance(n.func.value, ast.Name) and "log" in n.func.value.id.lower()]
    assert len(calls) == 1, "one log line"
    (call,) = calls
    assert not call.keywords and isinstance(call.args[0], ast.Constant) and "{" not in call.args[0].value
    assert {ast.unparse(a) for a in call.args[1:]} <= FORECAST_LOG_ARGS, ast.unparse(call)
    module = ast.parse((BACKEND / "app" / "forecast_claims.py").read_text(encoding="utf-8"))
    names = {n.id for n in ast.walk(module) if isinstance(n, ast.Name)} | \
        {a.name for n in ast.walk(module) if isinstance(n, (ast.Import, ast.ImportFrom)) for a in n.names}
    assert not names & {"logger", "logging", "log", "print", "gateway", "llm"}, "the reader logs nothing and imports no gateway"
    assert not any("llm" in (n.module or "") for n in ast.walk(module) if isinstance(n, ast.ImportFrom))


def test_the_forecast_reason_is_a_closed_sentence_and_the_line_label_never_reaches_the_gateway_or_a_log(caplog):
    assert not {"snippet", "sources", "by_period", "sheet_row", "file_hash", "dtype", "origin"} & decks.GATEWAY_READABLE_FIELDS
    sheet = cr_mod.read_workbook(_forecast_file(), "plan.csv")[0]
    ds = {"columns": sheet.columns, "rows": server.df_to_records(sheet.frame), "row_numbers": sheet.row_numbers, "wide": sheet.wide,
          "raw_headers": sheet.raw_headers, "first_col": sheet.first_col, "title_currency": sheet.currency,
          "unit_scale": sheet.unit_scale, "sheet": sheet.sheet, "file": "plan.csv", "file_hash": "h", "dtype": "pnl"}
    found = forecast_claims.candidates(forecast_claims.read(ds), {5, 7}, ds)
    rows = claim_matching.build_register([{**c, "id": f"fc{i}", "status": "approved", "file": "plan.csv", "order": i}
                                          for i, c in enumerate(found)], RESULTS,
                                         {"fiscal_year_end": 12, "reporting_currency": "EUR", "fx": {"EUR": 1.0}})
    assert rows and all(r["reason"] == claim_matching.FORECAST_NOTE for r in rows)
    assert FORECAST_SENTINEL not in json.dumps(rows, default=str), "the register holds no line label"
    db, reads = _recording_db()
    for i, c in enumerate(found):
        db[decks.CANDIDATES_COLLECTION].docs.append({**c, "audit_id": RUN_ID, "id": f"fc{i}", "status": "approved"})
    adapter = t.FakeAdapter(replies=[json.dumps(NARRATIVE)])
    with caplog.at_level(logging.DEBUG):
        asyncio.run(gateway.generate_narrative(db, RUN_ID, "growth_engine", adapter=adapter, sleep=t._noop_sleep))
    returned = json.dumps([doc for _, _, doc in reads], ensure_ascii=False, default=str)
    logs = caplog.text + "".join(r.getMessage() for r in caplog.records)
    for where, text in (("read by the gateway", returned), ("sent to the provider", adapter.payloads[0]), ("logged", logs)):
        assert FORECAST_SENTINEL not in text, f"the line label was {where}"
        assert claim_matching.FORECAST_NOTE not in text and "plan.csv" not in text, f"the reason or the file was {where}"


# ---------------------------------------------------------------------------
# Revenue file recognition (docs/specs/revenue-file-recognition.md section 5, CLAUDE.md rule 14). The text the model reads is built in
# memory by app/recognition.py `model_text`; the gateway checks it again (`structure_text_problem`, structure type "file_recognition").

# ---------------------------------------------------------------------------
import importlib  # noqa: E402
import hashlib  # noqa: E402

RECOGNITION_FIXTURES = BACKEND / "tests" / "fixtures" / "recognition"
RECOGNITION_COMPANY = "Northwind Test Co"
RECOGNITION_CONTEXT = {"company_name": RECOGNITION_COMPANY}
RECOGNITION_WITHHELD = (RECOGNITION_COMPANY,)       # what structure_redact.file_type_withheld(RECOGNITION_CONTEXT) returns
MASKED_FIXTURES = ("customer_revenue_months.csv", "unclassed_names_by_quarter.csv", "revenue_summary_by_product.csv",
                   "unsure_share_labels.csv", "unsure_few_hits.csv")      # the last two: label mode "unsure" (review H4)
FILE_HASH = hashlib.sha256(b"synthetic file one").hexdigest()
OTHER_HASH = hashlib.sha256(b"synthetic file two").hexdigest()
METRIC_FIXTURES = ("monthly_report_actual_forecast.csv", "pnl_multi_period.csv", "bookings_pipeline_only.csv")


def _recognition_sheets(name):
    return cr_mod.read_workbook((RECOGNITION_FIXTURES / name).read_bytes(), name)


def _recognition_text(name, mapping=None, file_hash=FILE_HASH):
    rec = importlib.import_module("app.recognition")
    return rec.model_text(_recognition_sheets(name), company_name=RECOGNITION_COMPANY, mapping=mapping or {}, withheld=RECOGNITION_WITHHELD,
                          file_hash=file_hash)


def _labels(text):
    return [m.group(2) for m in (re.match(r"^s1 label r(\d+): (.*)$", line) for line in text.split("\n")) if m]


def _label_cells(sheet):
    """The row labels of a sheet as written: the text of the first column below the header."""
    return [str(row[0]) for r, row in enumerate(sheet.grid, 1) if r > sheet.header_row and row and row[0] not in (None, "")]


def _number(v):
    try:
        return float(str(v).replace(",", ""))
    except ValueError:
        return None


def _numbers_by_column(sheet):
    out = {}
    for r, row in enumerate(sheet.grid, 1):
        if r <= sheet.header_row:
            continue
        for c, v in enumerate(row, 1):
            if c > 1 and _number(v) is not None:
                out.setdefault(c, []).append(_number(v))
    return out


@pytest.mark.parametrize("name", MASKED_FIXTURES)
def test_recognition_an_entity_labelled_sheet_sends_no_row_label_text_only_row_1_row_2(name):
    sheet = _recognition_sheets(name)[0]
    text = _recognition_text(name)
    assert "s1 labels: masked" in text.split("\n")
    labels = _labels(text)
    assert labels == [f"Row {i}" for i in range(1, len(_label_cells(sheet)) + 1)]
    for original in _label_cells(sheet):
        assert original not in text, f"the row label {original!r} reached the model text"


def test_recognition_a_label_column_headed_customer_is_masked_even_when_every_name_holds_a_metric_word():
    rec = importlib.import_module("app.recognition")
    sheets = cr_mod.read_workbook(("Revenue by customer,,,\nCustomer,Jan-25,Feb-25,Mar-25\nRevenue Partners,1,2,3\nSales Hub,4,5,6\n"
                                   "Cash Converters,7,8,9\nMargin Foods,1,2,3\n").encode(), "x.csv")
    text = rec.model_text(sheets, company_name=RECOGNITION_COMPANY, mapping={}, withheld=RECOGNITION_WITHHELD,
                          file_hash=FILE_HASH)
    assert _labels(text) == ["Row 1", "Row 2", "Row 3", "Row 4"]
    for name in ("Revenue Partners", "Sales Hub", "Cash Converters", "Margin Foods"):
        assert name not in text


@pytest.mark.parametrize("name", METRIC_FIXTURES)
def test_recognition_a_metric_labelled_sheet_sends_its_row_labels(name):
    sheet = _recognition_sheets(name)[0]
    text = _recognition_text(name)
    assert "s1 labels: metric" in text.split("\n")
    assert _labels(text) == _label_cells(sheet)


def _samples_by_column(text):
    samples = {}
    for line in text.split("\n"):
        m = re.match(r"^s1 c(\d+) sample: (.+)$", line)
        if m:
            samples.setdefault(int(m.group(1)), []).append(_number(m.group(2)))
    return samples


def _source_rows(sheet):
    """{column: {value: sheet row}}: every figure of these fixtures is unique, so a sample names its source row."""
    out = {}
    for r, row in enumerate(sheet.grid, 1):
        if r <= sheet.header_row:
            continue
        for c, v in enumerate(row, 1):
            if c > 1 and _number(v) is not None:
                out.setdefault(c, {})[_number(v)] = r
    return out


ALL_FIXTURES = MASKED_FIXTURES + METRIC_FIXTURES + ("arr_snapshot.csv", "pnl_single_period.csv")


@pytest.mark.parametrize("name", ALL_FIXTURES)
def test_recognition_samples_are_never_tied_to_a_row_no_two_neighbouring_columns_share_a_source_row(name):
    """Review H3, decision of 2026-10-19: on every sheet the example values of a column come from different rows chosen from the file hash,
    never the same row in two neighbouring columns, so no row's series can be rebuilt from the text."""
    sheet = _recognition_sheets(name)[0]
    text = _recognition_text(name)
    samples, source = _samples_by_column(text), _source_rows(sheet)
    assert set(samples) <= set(source)
    taken = {}
    for column in sorted(source):
        values = samples.get(column, [])
        assert len(values) <= min(3, len(source[column]) // 2), f"column {column}: at most min(3, rows ÷ 2) samples"
        assert all(v in source[column] for v in values), f"column {column}: a sample that is not a figure of that column"
        taken[column] = {source[column][v] for v in values}
    columns = sorted(source)
    for left, right in zip(columns, columns[1:]):
        assert not taken[left] & taken[right], f"columns {left} and {right} share a source row: a row's figures can be linked"
    leaked = set()
    for column, values in source.items():
        leaked |= set(values) - set(samples.get(column, []))
    unsampled = "\n".join(line for line in text.split("\n") if "sample:" not in line)
    found = {float(x) for x in re.findall(r"\d+(?:\.\d+)?", unsampled)} & leaked
    assert not found, f"figures beyond the sample rule: {sorted(found)[:5]}"


def test_recognition_the_sample_rows_come_from_the_file_hash_not_from_the_sheet_order():
    name = "customer_revenue_months.csv"
    sheet = _recognition_sheets(name)[0]
    source = _source_rows(sheet)
    rows_of = lambda text: {c: sorted(source[c][v] for v in vals) for c, vals in _samples_by_column(text).items()}  # noqa: E731
    one, again, other = rows_of(_recognition_text(name)), rows_of(_recognition_text(name)), rows_of(_recognition_text(name, file_hash=OTHER_HASH))
    assert one == again, "the same bytes give the same text"
    assert one != other, "other bytes give another choice"
    first_rows = {c: [min(source[c].values()) + i for i in range(3)] for c in source}
    assert any(one[c] != first_rows[c] for c in one), "not simply the first three rows of every column"


def test_recognition_the_text_holds_sheet_names_title_rows_and_headers_with_the_company_name_redacted():
    text = _recognition_text("monthly_report_actual_forecast.csv")
    assert RECOGNITION_COMPANY not in text and "Northwind" not in text and "[redacted]" in text
    assert text.split("\n")[0] == "s1: CSV"
    assert any(line.startswith("s1 title r1c1: ") for line in text.split("\n"))
    assert [m.group(1) for m in (re.match(r"^s1 header r3c(\d+): ", line) for line in text.split("\n")) if m] == [str(c) for c in range(1, 14)]


def test_recognition_the_columns_are_printed_as_the_sheet_counts_them_so_the_suggested_column_and_cue_address_are_sheet_positions():
    rec = importlib.import_module("app.recognition")
    shifted = cr_mod.read_workbook(b",Plan,\n,Metric,Jan-25,Feb-25\n,Revenue,1000,2000\n,Customers,30,40\n,Churn,5,6\n", "x.csv")
    assert shifted[0].first_col == 2
    text = rec.model_text(shifted, company_name=RECOGNITION_COMPANY, mapping={}, withheld=RECOGNITION_WITHHELD, file_hash=FILE_HASH)
    lines = text.split("\n")
    assert "s1 title r1c2: Plan" in lines and "s1 header r2c3: Jan-25" in lines and "s1 header r2c4: Feb-25" in lines, "B = 2, as in the sheet"
    assert not any(re.match(r"^s1 (title|header) r\d+c1: ", line) for line in lines), "column A is empty and is not printed as column 1"
    assert {int(m.group(1)) for m in (re.match(r"^s1 c(\d+) (sample|profile)", line) for line in lines) if m} <= {2, 3, 4}


def test_recognition_a_boundary_cue_is_kept_with_the_model_reply_and_the_audit_and_reaches_no_log_line(api, caplog):
    caplog.set_level(logging.DEBUG)
    cue = "Forecast scenario"
    csv = f"Test report,,,\n,Jan 2021,Feb 2021,{cue}\nRevenue,10,20,30\nCustomers,3,4,5\nChurn,1,2,3\n".encode()
    reply = json.dumps({"type": "file_recognition", "sheet": 1, "file_type": "monthly_revenue_report", "revenue_row": {"row": 3, "label": "Revenue"},
                        "currency": None, "non_revenue_rows": [], "last_actual_column": 3, "boundary_cue": {"cell": "D2", "text": cue}})
    api.state["adapter"] = t.FakeAdapter(replies=[reply])
    assert chat.upload(api, "report.csv", csv).status_code == 200
    log = "\n".join(r.getMessage() for r in caplog.records)
    assert cue not in log, "the cue text reached a log line"
    stored = _all_documents(api.db)
    assert cue not in stored.get("llm_calls", ""), "the call row holds counts and hashes, never text"
    kept = [d for d in api.db["column_mappings"].docs if d.get("kind") == "file_recognition"]
    assert kept and kept[0]["reply"]["boundary_cue"] == {"cell": "D2", "text": cue} and kept[0]["reply"]["last_actual_column"] == 3
    assert api.db["audits"].docs[0].get("as_of_cue") == cue, "the audit keeps the one cue its as-of month was detected by"


def test_recognition_the_customer_names_the_audit_already_pseudonymises_take_their_pseudonyms_in_the_title_and_headers():
    rec = importlib.import_module("app.recognition")
    sheets = cr_mod.read_workbook(b"Revenue of Alder Works by month,,\nMetric,Jan-25,Feb-25\nRevenue,1,2\nCustomers,3,4\nChurn,5,6\n", "x.csv")
    text = rec.model_text(sheets, company_name=RECOGNITION_COMPANY, mapping={"Alder Works": "Customer_01"}, withheld=RECOGNITION_WITHHELD,
                          file_hash=FILE_HASH)
    assert "Alder Works" not in text and "Customer_01" in text


def test_recognition_at_most_400_row_labels_are_sent_for_a_sheet():
    rec = importlib.import_module("app.recognition")
    rows = "".join(f"Customer {i},{10000 + i},{20000 + i}\n" for i in range(450))
    text = rec.model_text(cr_mod.read_workbook(("Revenue by customer,,\nCustomer,Jan-25,Feb-25\n" + rows).encode(), "x.csv"),
                          company_name=RECOGNITION_COMPANY, mapping={}, withheld=RECOGNITION_WITHHELD, file_hash=FILE_HASH)
    assert len(_labels(text)) == 398 and _labels(text)[-1] == "Row 398", "400 rows with text, the title row and the header row included"


def test_recognition_the_text_is_built_the_same_twice():
    for name in MASKED_FIXTURES + METRIC_FIXTURES:
        assert _recognition_text(name) == _recognition_text(name)


@pytest.mark.parametrize("name", MASKED_FIXTURES + METRIC_FIXTURES)
def test_recognition_the_gateway_accepts_the_text_the_builder_writes(name):
    assert gateway.structure_text_problem(_recognition_text(name), "file_recognition", RECOGNITION_CONTEXT, {}) is None


def test_recognition_the_gateway_refuses_a_row_label_on_a_masked_sheet_a_figure_a_file_name_an_unredacted_name_and_a_401st_row():
    masked = _recognition_text("customer_revenue_months.csv")
    metric = _recognition_text("monthly_report_actual_forecast.csv")
    check = lambda text: gateway.structure_text_problem(text, "file_recognition", RECOGNITION_CONTEXT, {})  # noqa: E731
    assert check(masked.replace("s1 label r3: Row 1", "s1 label r3: Alder Works")) == "label_on_masked_sheet"
    assert check(metric.replace("s1 label r4: Recognised revenue (EUR)", "s1 label r4: 10570")) == "figure_in_label"
    assert check(metric.replace("Recognised revenue (EUR)", "ledger_2025.xlsx")) == "file_name"
    assert check(metric.replace("[redacted]", RECOGNITION_COMPANY)) == "redaction_changed"
    extra = "\n".join(f"s1 label r{500 + i}: Row {13 + i}" for i in range(401))      # the numbering goes on from Row 13
    assert check(masked + "\n" + extra) == "too_many_rows"
    assert check("s1 labels: metric\ns1 label r4: Revenue") == "not_cells", "a line before the sheet line"
    assert check(metric + "\nfree prose about the company") == "not_cells"


def test_recognition_the_reply_is_closed_json_with_no_confidence_and_no_free_text():
    good = {"type": "file_recognition", "sheet": 1, "file_type": "monthly_revenue_report",
            "revenue_row": {"row": 4, "label": "Recognised revenue (EUR)"}, "currency": "EUR",
            "non_revenue_rows": [{"row": 9, "kind": "forecast"}]}
    parse = lambda body: gateway.parse_structure_reply(json.dumps(body), "file_recognition", "")  # noqa: E731
    assert parse(good).file_type == "monthly_revenue_report"
    assert parse({**good, "file_type": "not_revenue", "revenue_row": None, "currency": None,
                  "non_revenue_rows": []}).file_type == "not_revenue"
    for other in ("crm", "pnl"):                           # the same call answers the file type of chat-upload.md section 19 (spec 5.6)
        assert parse({**good, "file_type": other, "revenue_row": None, "currency": None, "non_revenue_rows": []})
    for bad in ({**good, "confidence": 90}, {**good, "note": "x"}, {**good, "file_type": "deck"}, {**good, "currency": "EURO"},
                {**good, "non_revenue_rows": [{"row": 9, "kind": "maybe"}]}, {**good, "non_revenue_rows": [{"row": 9, "kind": "forecast", "why": "x"}]},
                {**good, "last_actual_column": 0}, {**good, "last_actual_column": 7, "boundary_cue": {"cell": "Sheet1!H2", "text": "Forecast"}},
                {**good, "boundary_cue": {"cell": "H2"}}, {**good, "boundary_cue": {"cell": "H2", "text": "Forecast", "why": "x"}},
                {**good, "boundary_cue": {"cell": "H2", "text": "x" * 201}}, {**good, "boundary_cue": "H2"},
                {**good, "revenue_row": {"row": 4}}, {**good, "sheet": 0}):
        with pytest.raises(gateway.GatewayError):
            parse(bad)
    cued = parse({**good, "last_actual_column": 7, "boundary_cue": {"cell": "H2", "text": "Forecast"}})
    assert (cued.last_actual_column, cued.boundary_cue.cell, cued.boundary_cue.text) == (7, "H2", "Forecast")
    assert parse({**good, "last_actual_column": None, "boundary_cue": None}).boundary_cue is None
    for cell in ("D3", "AA12", "r3c4", "r12"):
        assert parse({**good, "last_actual_column": 4, "boundary_cue": {"cell": cell, "text": "Plan"}})
    kinds = ("forecast", "budget", "bookings", "pipeline", "contract_value", "run_rate")
    for kind in kinds:
        assert parse({**good, "non_revenue_rows": [{"row": 9, "kind": kind}]})


def test_recognition_the_words_a_result_can_carry_are_closed_lists_and_hold_no_text_from_a_file():
    """A result key, a reason string or a kind that is not on a closed list could carry file text into a log or a reply (rule 14)."""
    rec = importlib.import_module("app.recognition")
    assert rec.RECOGNITION_TYPES == ("invoice_transactions", "customer_revenue", "monthly_revenue_report", "revenue_summary",
                                     "financial_statements", "arr_mrr_snapshot")
    assert rec.NON_REVENUE_KINDS == ("forecast", "budget", "bookings", "pipeline", "contract_value", "run_rate")
    assert set(rec.REASONS) == {"no_candidate", "tie", "row_missing", "label_mismatch", "marked_not_revenue",
                                "not_numeric_enough", "no_revenue_word", "rules_and_model_disagree", "after_as_of",
                                "not_a_period", "turnover_question", "type_unsure"}
    assert all(re.fullmatch(r"[a-z0-9_]+", w) for w in rec.REASONS + rec.RECOGNITION_TYPES + rec.NON_REVENUE_KINDS)
    for name in MASKED_FIXTURES + METRIC_FIXTURES + ("arr_snapshot.csv", "invoice_transactions.csv", "pnl_single_period.csv"):
        got = rec.recognise(_recognition_sheets(name), as_of_month="2025-12-31")
        assert set(got["missing"]) <= {"revenue_measure", "reference", "numbers"}
        for part in (got["revenue_row"], got["last_actual"]):
            assert part is None or part["reason"] is None or part["reason"] in rec.REASONS
        assert all(r["kind"] in rec.NON_REVENUE_KINDS for r in got["non_revenue_rows"])
        assert got["file_type"] is None or got["file_type"] in rec.RECOGNITION_TYPES
        assert all(b["analysis"] in ("customer_level", "trend", "monthly_series") for b in got["blocked"])
        assert set(got["supports"]) <= {"customer_level", "trend", "claim_matching"}


# --- Cells under a customer-type header with no customer list (PR 147, review H1, George 2026-10-21) -------------------------------
def _customer_table():
    return [{"row": 1, "col": 1, "text": "Top customers"}, {"row": 1, "col": 2, "text": "ARR"},
            {"row": 2, "col": 1, "text": "Acme Retail"}, {"row": 2, "col": 2, "text": "€1.2M"},
            {"row": 3, "col": 1, "text": "Northwind Logistics"}, {"row": 3, "col": 2, "text": "€0.8M"},
            {"row": 4, "col": 1, "text": "Account 2024"}, {"row": 4, "col": 2, "text": "12"}]


@pytest.mark.parametrize("header", ["Customer", "Top customers", "Client", "Accounts", "Logo", "Key logos"])
def test_cells_under_a_customer_header_are_blanked_when_the_app_has_no_customer_list(header):
    cells = [{**c, "text": header} if (c["row"], c["col"]) == (1, 1) else c for c in _customer_table()]
    out, _ = structure_redact.redact_structure(cells, "Target Co", {}, deck=True)
    sent = structure_redact.structure_text(out)
    for needle in ("Acme Retail", "Northwind Logistics"):
        assert needle not in sent, f"{needle!r} reached the provider"
    texts = {(c["row"], c["col"]): c["text"] for c in out}
    assert texts[(2, 1)] == texts[(3, 1)] == structure_redact.REDACTED
    assert texts[(1, 1)] == header and texts[(2, 2)] == "€1.2M" and texts[(4, 2)] == "12", "headers and figures stay"


def test_a_customer_list_swaps_the_names_and_the_header_rule_stays_off():
    out, _ = structure_redact.redact_structure(_customer_table(), "Target Co", {"Acme Retail": "Customer_01"}, deck=True)
    texts = {(c["row"], c["col"]): c["text"] for c in out}
    assert texts[(2, 1)] == "Customer_01" and texts[(3, 1)] == "Northwind Logistics", "the known limit: names not in the list"


def test_a_deck_table_reaches_the_provider_with_no_customer_name_when_there_is_no_customer_list():
    text = structure_redact.structure_text(_customer_table())
    result, adapter = _send(_structure_db(), text, reply=NO_LABELS)
    assert result.status == "refused" and adapter.calls == 0, "the gateway re-runs the redaction and refuses unblanked names"
    blanked, _ = structure_redact.redact_structure(_customer_table(), "Target Co", {}, deck=True)
    result, adapter = _send(_structure_db(), structure_redact.structure_text(blanked), reply=NO_LABELS)
    assert result.status == "read" and adapter.calls == 1, result.reason
    sent = json.loads(adapter.payloads[0])["text"]
    assert "Acme Retail" not in sent and "Northwind Logistics" not in sent


def test_a_sheet_step_keeps_its_header_text_the_customer_header_rule_is_for_decks_only():
    cells = [{"row": 1, "col": 1, "text": "Accounts through 2028"}, {"row": 2, "col": 1, "text": "Account ID"}]
    out, counts = structure_redact.redact_structure(cells, "Target Co", {})
    assert out == cells and "customer_blanked" not in counts


# ---------------------------------------------------------------------------
# The page_claims step (docs/design/ai-reads-rules-check.md, docs/design/page-image.md; CLAUDE.md rules 16 and 31): the text lines of
# deck pages that carry a figure as digits, as printed with no redaction, may reach the provider, prose lines included, as
# `p<page>r<row>c<col>: <text>` lines through one builder (structures.pages.batch_text) with one PNG of each whole page in the same
# request, at most 5 pages per call, only while DECK_PIPELINE is "ai" and only to a Sonnet or Opus model. A page with no figure as
# digits never goes (a number-word is no figure; a line that is only the page's own number is none); a line over 200 characters is
# left out by the sender and refused by the gateway; a figure outside a cell line is refused; a page that cannot be drawn goes as text alone (the deck's `images_unavailable`, a flag). The
# reply holds closed words, numbers and cell ids only; nothing sent, the label cell's text included, is stored or logged.
# ---------------------------------------------------------------------------
from unittest import mock  # noqa: E402

from app.decks import pages as page_grids  # noqa: E402
from app.llm.schemas import PAGE_CLAIMS, PageClaim, page_claims_output_schema  # noqa: E402
from app.structures import page_check, pages as page_pipeline  # noqa: E402

PROSE_LINE = ("Our platform served 2,300 companies in 2025 while the founding team kept its focus on product quality and on the "
              "customers who had asked for a European data centre; the roadmap for next year adds two more regions and a partner programme.")
SHORT_PROSE = ("Our platform served 2,300 companies in 2025 while the team kept its focus on product quality and on the customers "
               "who had asked for a European data centre.")
GOOD_PAGE = ("p5r1c1 title: Traction\np5r2c1: 800 Paying Users\np5r3c1: $150,000 annual revenue run rate\np5r4c1: 97% margins\n"
             "p5r5c1: " + SHORT_PROSE)
PAGE_REPLY = json.dumps({"type": PAGE_CLAIMS, "claims": [
    {"metric": "customers", "value": 800, "value_high": None, "unit": "count", "unit_other": None, "period": None,
     "basis": "actual", "scenario": "none", "value_cell": "p5r2c1", "label_cell": "p5r1c1", "period_cell": None},
    {"metric": "revenue", "value": 150000, "value_high": None, "unit": "USD", "unit_other": None, "period": None,
     "basis": "actual", "scenario": "none", "value_cell": "p5r3c1", "label_cell": "p5r3c1", "period_cell": None}]})
NO_CLAIMS = json.dumps({"type": PAGE_CLAIMS, "claims": []})


def _page(n, *lines):
    return "\n".join(f"p{n}r{i}c1: {text}" for i, text in enumerate(lines, 1))


def _send_page(db, text, reply=NO_CLAIMS, pipeline="ai"):
    """read_structure on a page text with the switch DECK_PIPELINE at `pipeline` (None: unset, the default "ai")."""
    env = {"DECK_PIPELINE": pipeline} if pipeline is not None else {}
    with mock.patch.dict(os.environ, env, clear=False):
        if pipeline is None:
            os.environ.pop("DECK_PIPELINE", None)
        return _send(db, text, PAGE_CLAIMS, reply)


def test_a_page_text_with_a_figure_as_digits_reaches_the_provider_as_printed_with_its_image():
    assert len(SHORT_PROSE) <= 200 < len(PROSE_LINE)
    result, adapter = _send_page(_structure_db(), GOOD_PAGE, PAGE_REPLY)
    assert result.status == "read" and adapter.calls == 1, result.reason
    sent = json.loads(adapter.payloads[0])
    assert sent == {"type": PAGE_CLAIMS, "text": GOOD_PAGE}
    assert all(re.match(r"^p\d+r\d+c\d+( title)?: ", line) for line in sent["text"].splitlines())
    assert adapter.requests[-1]["cache_system"] is True, "the fixed prompt is marked for the provider's prompt cache"
    assert adapter.requests[-1]["images"] == [page_images.png(1)], "the page image goes in the same request as the text"
    assert [c["value_cell"] for c in result.claims] == ["p5r2c1", "p5r3c1"]


def test_the_page_step_is_refused_unless_the_switch_is_ai():
    """L1 of the review of PR #159: whatever the caller, the gateway reads a page only while DECK_PIPELINE is "ai". Since
    2026-10-22 "ai" is the default: unset, or a typo, reads; "rules" (the former path) refuses."""
    result, adapter = _send_page(_structure_db(), GOOD_PAGE, PAGE_REPLY, pipeline="rules")
    assert (result.status, result.reason, adapter.calls) == ("refused", "refused: pipeline_rules", 0)
    for pipeline in (None, "ai", "pages"):
        result, adapter = _send_page(_structure_db(), GOOD_PAGE, PAGE_REPLY, pipeline=pipeline)
        assert (result.status, adapter.calls) == ("read", 1), pipeline
    result, adapter = _send(_structure_db(), GOOD_ITEMS)
    assert result.status == "read", "a structure read does not depend on the switch"


def test_a_page_line_over_200_characters_is_left_out_by_the_sender_and_refused_by_the_gateway():
    """Decision 1b on the review of PR #159 (George, 2026-10-22): as for a structure cell, a line over 200 characters never goes."""
    long_line = ("Our CEO Michael Smith and our CTO Desmond Okafor founded the company with Northwind Trading as the first "
                 "customer; Jane Doe (jane.doe@northwind.com, +44 20 7946 0958) leads sales across three regions and the team of 12 "
                 "keeps growing with every product line launched since the spring")
    assert len(long_line) > 200
    cells = [{"row": 1, "col": 1, "text": "Team", "box": 1}, {"row": 2, "col": 1, "text": long_line, "box": 1},
             {"row": 3, "col": 1, "text": "12 engineers", "box": 2}]
    grid = {"type": "page", "page": 3, "header_rows": 0, "cells": cells, "tables": [], "figure": True}
    sent, = page_pipeline.sendable_pages([grid])
    assert [c["text"] for c in sent["cells"]] == ["Team", "12 engineers"], "the sender leaves the long line out"
    text = page_pipeline.batch_text([sent])
    result, adapter = _send_page(_structure_db(), text)
    assert result.status == "read" and adapter.calls == 1, result.reason
    with_long = structure_redact.page_text([(3, cells)])
    result, adapter = _send_page(_structure_db(), with_long)
    assert (result.status, result.reason, adapter.calls) == ("refused", "refused: cell_too_long", 0)
    grid_only_long = {**grid, "cells": cells[:2]}
    assert page_pipeline.sendable_pages([grid_only_long]) == [], "a page whose only figure sat in the long line has none"


def test_a_page_goes_as_printed_names_customers_emails_phones_and_file_names_included():
    """page-image.md section 2: no redaction on a page. The customer rules (the pseudonym swap, the customer-type header blanking), the
    people-name rule and the e-mail and phone rules do not run on it: sender and gateway both leave every printed line as it is."""
    cells = [{"row": 1, "col": 1, "text": "Traction", "box": 1, "title": True},
             {"row": 2, "col": 1, "text": "Customers", "box": 2}, {"row": 3, "col": 1, "text": "Northwind Trading (12 seats)", "box": 2},
             {"row": 4, "col": 1, "text": "Globex Corp", "box": 2},
             {"row": 5, "col": 1, "text": "800 Paying Users", "box": 3},
             {"row": 6, "col": 1, "text": "Rohan Mehta, Co-Founder", "box": 4}, {"row": 7, "col": 1, "text": "Jane Doe", "box": 4},
             {"row": 8, "col": 1, "text": "jane.doe@northwind.com +44 20 7946 0958", "box": 4},
             {"row": 9, "col": 1, "text": "See Acme_Board_Q3.pptx", "box": 4},
             {"row": 10, "col": 1, "text": "Clients", "table": 1, "trow": 1, "tcol": 1},
             {"row": 10, "col": 2, "text": "ARR", "table": 1, "trow": 1, "tcol": 2},
             {"row": 11, "col": 1, "text": "Initech Ltd", "table": 1, "trow": 2, "tcol": 1},
             {"row": 11, "col": 2, "text": "€1.5M", "table": 1, "trow": 2, "tcol": 2}]
    grid = {"type": "page", "page": 5, "header_rows": 0, "cells": cells, "tables": [{"table": 1, "header_rows": 1}], "figure": True}
    sent, = page_pipeline.sendable_pages([grid])
    assert [c["text"] for c in sent["cells"]] == [c["text"] for c in cells], "the sender changes no line"
    text = page_pipeline.batch_text([sent])
    mapping = {"Northwind Trading": "Customer_01"}              # a customer list the page no longer reads
    result, adapter = _send_page(_structure_db(mapping), text)
    assert result.status == "read" and adapter.calls == 1, result.reason
    seen = json.loads(adapter.payloads[0])["text"]
    for printed in ("Northwind Trading (12 seats)", "Globex Corp", "Rohan Mehta, Co-Founder", "Jane Doe", "jane.doe@northwind.com",
                    "+44 20 7946 0958", "Acme_Board_Q3.pptx", "Initech Ltd"):
        assert printed in seen, f"{printed!r} is printed on the page and goes as printed"
    for placeholder in ("[redacted]", "[person]", "[email]", "[phone]", "Customer_01"):
        assert placeholder not in seen
    # The page rules are gone from the code, not only unused; a deck structure keeps its own.
    assert not hasattr(structure_redact, "redact_page") and not hasattr(structure_redact, "redact_role_names")
    out, _ = structure_redact.redact_structure(cells[:5], "Target Co", mapping, deck=True)
    assert "Northwind" not in " ".join(c["text"] for c in out)


def test_a_page_read_without_its_image_goes_as_text_alone():
    adapter = t.FakeAdapter(replies=[PAGE_REPLY])
    result = asyncio.run(gateway.read_structure(_structure_db(), STRUCTURE_AUDIT["id"], GOOD_PAGE, PAGE_CLAIMS, adapter=adapter,
                                                sleep=t._noop_sleep))
    assert result.status == "read" and adapter.calls == 1 and not adapter.requests[-1].get("images")
    assert json.loads(adapter.payloads[0]) == {"type": PAGE_CLAIMS, "text": GOOD_PAGE}


def test_the_image_goes_only_with_a_page_read():
    """A structure, a sheet or a column mapping is text alone, whatever the caller passes."""
    adapter = t.FakeAdapter(replies=[STRUCTURE_REPLY])
    result = asyncio.run(gateway.read_structure(_structure_db(), STRUCTURE_AUDIT["id"], GOOD_ITEMS, "table", adapter=adapter,
                                                sleep=t._noop_sleep, images=[page_images.png(1)]))
    assert result.status == "read" and adapter.calls == 1 and adapter.requests[-1]["images"] is None


@pytest.mark.parametrize("name, text, reason", [
    ("a page with no figure as digits", _page(2, "Our mission", "We help teams ship faster"), "no_figure"),
    ("a page whose only figure is a number-word", _page(2, "Revenue grew to three million euros", "Twenty customers"), "no_figure"),
    ("a page whose only digits are a year", _page(2, "Founded in 2019", "Launched in January 2021"), "no_figure"),
    ("a page whose only digits are its page number", _page(3, "Our mission", "We help teams ship faster", "3"), "no_figure"),
    ("a page number as Slide n", _page(3, "Our mission", "Slide 3"), "no_figure"),
    ("a page number as n/m", _page(3, "Our mission", "3/25"), "no_figure"),
    ("a page number as Page n of m", _page(3, "Our mission", "Page 3 of 25"), "no_figure"),
    ("a page number between dashes", _page(3, "Our mission", "- 3 -"), "no_figure"),
    ("a page line over 200 characters", GOOD_PAGE + "\np5r6c1: " + PROSE_LINE, "cell_too_long"),
    ("one page without a figure among pages with one", GOOD_PAGE + "\n" + _page(6, "Our mission"), "no_figure"),
    ("six pages", "\n".join(_page(n, "Revenue €1.5M") for n in range(1, 7)), "too_many_pages"),
    ("a figure outside a cell line", GOOD_PAGE + "\n1,200", "not_cells"),
    ("a prose line with no cell id", GOOD_PAGE + "\nRevenue grew from £1.2m to £1.5m between FY2025 and FY2026.", "not_cells"),
    ("structure lines with no page", "r1c1: Revenue\nr1c2: £1,200,000", "not_cells"),
    ("pages out of order", _page(6, "Revenue €1.5M") + "\n" + _page(5, "Costs €1M"), "not_cells"),
    ("raw bytes", "p5r1c1: x\x00\x01", "raw_bytes"),
])
def test_the_page_step_refuses_what_rule_16_does_not_allow(name, text, reason):
    assert reason in STRUCTURE_REFUSALS
    result, adapter = _send_page(_structure_db({"Northwind Trading": "Customer_01"}), text)
    assert adapter.calls == 0 and getattr(adapter, "counted", 0) == 0, f"{name} reached the provider"
    assert (result.status, result.reason) == ("refused", f"refused: {reason}"), name


def test_the_page_step_reads_no_deck_text_and_the_page_cells_come_through_the_one_builder(monkeypatch):
    """read_structure is handed the page text; it reads consent and names from the audit, the pseudonym map, its own cache and the
    call log, never the parsed-text or candidate collections. The text itself is built by structures.pages.batch_text from the
    page grids and the page redaction, nowhere else: no gateway module names the parsed-text fields (the static test above)."""
    monkeypatch.setenv("DECK_PIPELINE", "ai")
    db, reads = _recording_db()
    db["audits"].docs[0].update({"company_name": "Target Co", "structure_reading_consent": True})
    adapter = t.FakeAdapter(replies=[PAGE_REPLY])
    asyncio.run(gateway.read_structure(db, RUN_ID, GOOD_PAGE, PAGE_CLAIMS, adapter=adapter, sleep=t._noop_sleep,
                                       images=[page_images.png(1)]))
    assert adapter.calls == 1
    touched = {name for name, _, _ in reads}
    assert not touched & {decks.TEXT_COLLECTION, decks.CANDIDATES_COLLECTION}, touched
    assert DECK_SENTINEL not in adapter.payloads[0] and DECK_FILE not in adapter.payloads[0]
    source = inspect.getsource(page_pipeline.sendable_pages) + inspect.getsource(page_pipeline.batch_text)
    assert "redact_page" not in source and "page_text" in source
    assert "page_text" not in inspect.getsource(gateway.read_structure), "the gateway builds no page text"


def test_no_sent_page_text_and_no_label_text_in_a_log_llm_calls_or_llm_structures(caplog, monkeypatch):
    """Rule 17 on the page step: after a deck is read page by page, the stored reading holds the model's claims (closed words,
    numbers, cell ids), the checked items (ids, values, statuses, checks) and metadata; no line of the page, no label text, no
    period text is stored or logged."""
    pytest.importorskip("pdfplumber")
    from pdf_support import pdf_page
    from app.decks.parser import parse_deck
    monkeypatch.setenv("DECK_PIPELINE", "ai")
    texts = [(50, 560, 20, "Traction"), (50, 500, 12, "800 Paying Users"), (50, 480, 12, "$150,000 annual revenue run rate"),
             (50, 460, 12, "FY2025"), (50, 320, 10, SHORT_PROSE), (50, 300, 10, PROSE_LINE)]
    parsed = parse_deck(pdf_page(texts, []), "deck.pdf")
    db = _structure_db()
    db["datasets"].docs.append({"audit_id": STRUCTURE_AUDIT["id"], "dtype": "revenue", "mapped_at": "2026-10-05",
                                "mapping": {"customer_id": "Customer"}})
    db[decks.TEXT_COLLECTION].docs.append({"audit_id": STRUCTURE_AUDIT["id"], "deck_id": "d1", "file": "deck.pdf",
                                           "page_unit": "page", "blocks": parsed["blocks"], "structures": []})
    reply = json.dumps({"type": PAGE_CLAIMS, "claims": [
        {"metric": "revenue", "value": 150000, "value_high": None, "unit": "USD", "unit_other": None, "period": "FY2025",
         "basis": "actual", "scenario": "none", "value_cell": "p1r3c1", "label_cell": "p1r1c1", "period_cell": "p1r4c1"}]})
    adapter = t.FakeAdapter(replies=[reply])
    with caplog.at_level(logging.DEBUG):
        status = asyncio.run(structures.process_deck(db, STRUCTURE_AUDIT["id"], "d1", adapter=adapter, sleep=t._noop_sleep))
    assert status == "read" and adapter.calls == 1
    sent = json.loads(adapter.payloads[0])["text"]
    assert "800 Paying Users" in sent and SHORT_PROSE in sent, "fixture: the page went, prose included"
    assert PROSE_LINE not in sent, "the line over 200 characters is left out (decision 1b, review of PR #159)"
    stored = json.dumps(db[gateway.STRUCTURES_COLLECTION].docs + db["llm_calls"].docs, ensure_ascii=False, default=str)
    logs = caplog.text + "".join(r.getMessage() for r in caplog.records)
    # The model's period "FY2025" is a closed form, kept in its reply as today's labels keep theirs; the header's text is not.
    for needle in ("Traction", "Paying Users", "150,000", "annual revenue", "2,300 companies", PROSE_LINE[:40], sent):
        assert needle not in stored, f"{needle!r} was stored"
        assert needle not in logs, f"{needle!r} was logged"
    record, = db[gateway.STRUCTURES_COLLECTION].docs
    assert record["items"]["claims"][0]["label_cell"] is None and record["items"]["claims"][0]["label_is_title"] is True \
        and record["items"]["claims"][0]["period_cells"] == ["r4c1"], "the page title is no label (George, 2026-10-27): the cell id of a title is not stored as one"
    assert record["statuses"] == ["matched"]
    assert "pages checked: run_id=audit-boundary deck_id=d1 pages=1 matched=1 unmatched=0 dropped=0 flagged=1" in caplog.text
    row, = db[decks.CANDIDATES_COLLECTION].docs
    assert row["label_from"] is None and row["ai_label"] == "Matched to page", "the row, for the analyst, reads the page"


def test_every_page_reply_field_is_a_closed_word_a_number_or_a_cell_id():
    """The new rule of the design note (section 5): a reply can carry no deck text into a log. Every property of the schema is an
    enum, a number or a string the reply model checks as a cell id or an ISO currency code."""
    schema = page_claims_output_schema()
    claim = schema["properties"]["claims"]["items"]
    assert schema["properties"]["type"] == {"type": "string", "enum": [PAGE_CLAIMS]}
    free = set()
    for name, prop in claim["properties"].items():
        options = prop.get("anyOf", [prop])
        kinds = {o.get("type") for o in options if "enum" not in o}
        if "string" in kinds:
            free.add(name)
        assert kinds <= {"number", "null", "string", "boolean"}, (name, prop)
    # label_text, segment and share_of (docs/design/page-image.md section 4) are short printed text: model JSON output, stored under rule 17,
    # at most PAGE_TEXT_MAX characters, never written to a log (the next test); is_quote is a boolean.
    assert free == {"value_cell", "label_cell", "period_cell", "unit_other", "period", "label_text", "segment", "share_of"}
    base = {"metric": "revenue", "value": 1.0, "value_high": None, "unit": "EUR", "unit_other": None, "period": None,
            "basis": "actual", "scenario": "none", "value_cell": "p5r2c1", "label_cell": None, "period_cell": None}
    assert PageClaim.model_validate(base).value_cell == "p5r2c1"
    for field in ("value_cell", "label_cell", "period_cell"):
        with pytest.raises(Exception):
            PageClaim.model_validate({**base, field: "Revenue £1,200,000 at Target Co"})
        with pytest.raises(Exception):
            PageClaim.model_validate({**base, field: "r2c1"})               # a cell id without its page
    with pytest.raises(Exception):
        PageClaim.model_validate({**base, "unit": "other", "unit_other": "Revenue at Target Co"})
    with pytest.raises(Exception):
        PageClaim.model_validate({**base, "period": "in 2025 at Target Co"})
    with pytest.raises(Exception):
        PageClaim.model_validate({**base, "metric": "Revenue at Target Co"})
    with pytest.raises(Exception):
        PageClaim.model_validate({**base, "note": "free text"})


def test_a_page_reply_not_json_is_rejected_a_claim_citing_a_cell_not_sent_or_carrying_text_is_dropped_alone(caplog):
    first = json.loads(PAGE_REPLY)["claims"][0]
    with caplog.at_level(logging.DEBUG):
        result, adapter = _send_page(_structure_db(), GOOD_PAGE, "Revenue £1,200,000 at Target Co")
        assert (result.status, adapter.calls) == ("not_read", 2)
        # The reason the page step logs: the check word and the failed schema fields, never a value.
        assert result.rejected == "schema" and "Target Co" not in result.rejected
        # George, 2026-10-23: a schema failure or a cell not sent inside one claim drops that claim only, the reply is read
        for sent in (json.dumps({"type": PAGE_CLAIMS, "claims": [{**first, "value_cell": "p9r9c9"}]}),
                     json.dumps({"type": PAGE_CLAIMS, "claims": [{**first, "label_cell": "p5r1c1 Traction"}]}),
                     json.dumps({"type": PAGE_CLAIMS, "claims": [{**first, "metric": "Revenue at Target Co"}]}),
                     json.dumps({"type": PAGE_CLAIMS, "claims": [{**first, "segment": "Target Co " * 20}]}),
                     json.dumps({"type": PAGE_CLAIMS, "claims": [{**first, "is_quote": "Target Co"}]})):
            result, adapter = _send_page(_structure_db(), GOOD_PAGE, sent)
            assert (result.status, adapter.calls) == ("read", 1)
    lines = [r.getMessage() for r in caplog.records if r.getMessage().startswith("structure not read:")]
    assert [line.rsplit(" check=", 1)[1] for line in lines] == ["schema"]
    # The schema fields that failed, as closed names from the schema (PR B: the run of 2026-10-22 lost 10 replies to check=schema
    # with no way to tell which field); a cell not sent or a reply that is not JSON names none.
    assert [line.split(" fields=")[1].split(" ")[0] for line in lines] == ["-"]
    rejected = [r.getMessage() for r in caplog.records if "claim rejected" in r.getMessage()]
    assert [re.sub(r"page \d+", "page n", line) for line in rejected] == ["page n: claim rejected (cell not in batch)",
                                                                            "page n: claim rejected (label_cell)",
                                                                            "page n: claim rejected (metric)",
                                                                            "page n: claim rejected (segment)",
                                                                            "page n: claim rejected (is_quote)"]
    assert set(gateway.PAGE_FIELDS) == {"type", "claims", "metric", "value", "value_high", "unit", "unit_other", "period",
                                        "basis", "scenario", "value_cell", "label_cell", "period_cell", "label_text", "segment",
                                        "share_of", "is_quote"}
    logs = caplog.text + "".join(r.getMessage() for r in caplog.records)
    for needle in ("Traction", "Paying Users", "Target Co", "p9r9c9", "Input should"):
        assert needle not in logs, f"{needle!r} was logged"


def test_label_text_segment_and_share_of_reach_the_approval_row_only_never_a_log_or_the_stored_read(caplog):
    """docs/design/page-image.md section 4 with rule 17: the three short printed texts of a page claim label the row the analyst sees; no log line,
    llm_calls or llm_structures document holds them (the stored read keeps cell ids, values and is_quote)."""
    import test_checker as tc
    import test_claim_corrections as cc
    parsed = tc.deck([tc.TITLE, (50, 480, 12, "800")])
    db = tc._page_db(parsed)
    claim_ = {**tc.claim(cc.at(parsed, "800"), 800, metric="customers", unit="count"), "label_text": "Zebra subscribers",
              "segment": "Zebra EMEA", "share_of": "Zebra total", "is_quote": False}
    adapter = t.FakeAdapter(replies=[json.dumps({"type": PAGE_CLAIMS, "claims": [claim_]})])
    with caplog.at_level(logging.DEBUG):
        assert asyncio.run(structures.process_deck(db, cc.AUDIT, "d1", adapter=adapter, sleep=t._noop_sleep)) == "read"
    assert "Zebra" not in caplog.text + "".join(r.getMessage() for r in caplog.records)
    assert "Zebra" not in json.dumps(db["llm_calls"].docs, default=str)
    doc, = db["llm_structures"].docs
    assert "Zebra" in json.dumps({k: v for k, v in doc.items() if k != "items"}, default=str), "the model's reply is kept as it came (rule 17: model JSON output)"
    assert "Zebra" not in json.dumps(doc["items"], default=str) and doc["items"]["claims"][0]["is_quote"] is False
    assert [r["label_from"] for r in db[decks.CANDIDATES_COLLECTION].docs] == ["Zebra subscribers"]


def test_a_page_reply_cut_at_the_output_limit_is_logged_as_truncated_with_its_tokens_and_no_text(caplog):
    """PR B (live run of 2026-10-22): a reply cut at max_tokens is reported by the adapter as "truncated" with the tokens it billed. The
    gateway logs the closed reason, bills the call in llm_calls and answers "truncated" (the page step splits the batch); the cut
    reply, which may hold a claim half-written, is never logged or stored."""
    cut = gateway.GatewayError("truncated", 'cut: {"metric": "revenue", "value_cell": "p5r3c1", "label": "Paying Users at Target Co')
    cut.tokens = (3000, 4000)
    adapter = t.FakeAdapter(raise_with=cut)
    with caplog.at_level(logging.DEBUG), mock.patch.dict(os.environ, {"DECK_PIPELINE": "ai"}):
        result = asyncio.run(gateway.read_structure(_structure_db(), STRUCTURE_AUDIT["id"], GOOD_PAGE, PAGE_CLAIMS, adapter=adapter,
                                                    sleep=t._noop_sleep, images=[page_images.png(1)]))
    assert (result.status, result.reason, adapter.calls) == ("truncated", gateway.NOT_READ, 1), "no reask: the same batch would be cut again"
    assert (result.input_tokens, result.output_tokens) == (3000, 4000)
    lines = [r.getMessage() for r in caplog.records if r.getMessage().startswith("structure not read:")]
    assert len(lines) == 1 and " tokens=3000/4000 " in lines[0] and lines[0].endswith("reason=truncated status=- type=- fields=- check=-")
    logs = caplog.text + "".join(r.getMessage() for r in caplog.records)
    for needle in ("Paying Users", "Target Co", "p5r3c1", "cut:"):
        assert needle not in logs, f"{needle!r} was logged"


def test_the_page_grid_layout_fields_are_kept_with_the_grid_and_never_sent():
    """next_to, bar_year, tick, box and table are Python's layout, stored with the grid for page_check and never in the text."""
    pytest.importorskip("pdfplumber")
    from pdf_support import pdf_page
    from app.decks.parser import parse_deck
    texts = [(50, 560, 20, "Revenue forecast")] + [(x, 380 + 20 * i, 12, bar) for i, (bar, x) in
                                                   enumerate(zip(("€1.48MM", "€3.39MM", "€5.3MM"), (100, 255, 410)))]
    texts += [(x, 340, 12, year) for year, x in zip(("2022", "2023", "2024"), (150, 300, 450))]
    grid, = page_grids.page_grids(parse_deck(pdf_page(texts, []), "deck.pdf")["blocks"])
    assert sorted(c["bar_year"] for c in grid["cells"] if "bar_year" in c) == ["2022", "2023", "2024"]
    text = page_pipeline.batch_text(page_pipeline.batches([grid])[0])
    for word in ("next_to", "bar_year", "tick", "box", "table"):
        assert word not in text
    assert set(structure_redact.parse_page_text(text)) == {1}
    cells = structure_redact.parse_page_text(text)[1]
    assert [c["text"] for c in cells] == [c["text"] for c in sorted(grid["cells"], key=lambda c: (c["row"], c["col"]))]
    assert set(page_check.check(grid, [], 12)) == {"items", "matched", "unmatched", "dropped", "periods_corrected"}


def test_the_layout_evidence_is_python_only_stored_without_text_and_never_sent_or_logged(caplog):
    """Source association (docs/specs/source-association.md, 2026-10-26): rotation, places, rectangles, colours, scopes, gaps and the association
    record are Python's. The page text is one `p<n>r<row>c<col>: <text>` line per cell (a turned line is an ordinary line); the deck's layout is stored
    as numbers, colour hex and closed words; a dropped claim leaves a Diagnostics line of closed words and numbers; no log line holds a line of the page."""
    pytest.importorskip("pdfplumber")
    import logging
    import test_checker as tc
    from pdf_support import pdf_page
    from app.decks.parser import GAPS, parse_deck
    from app.structures import association, checker as page_checker
    secret = "Northwind quarterly plan"
    parsed = parse_deck(pdf_page([(50, 560, 20, secret), (120, 460, 12, "42"), (600, 70, 12, "Enterprise")],
                                 fills=[(100, 100, 60, 340, (0.12, 0.47, 0.71)), (560, 68, 22, 18, (0.12, 0.47, 0.71))],
                                 rotated=[(60, 150, 12, "Revenue (EUR M)", "up")], pictures=[(10, 10, 400, 300)]), "deck.pdf")
    grid, = page_grids.page_grids(parsed["blocks"], parsed["layout"])
    text = page_pipeline.batch_text(page_pipeline.batches([grid])[0])
    assert all(re.fullmatch(r"p1r\d+c\d+(?: title)?: .*", line) for line in text.split("\n")), "one id and one text per line, nothing else"
    assert [line.split(": ", 1)[1] for line in text.split("\n")] == [c["text"] for c in sorted(grid["cells"], key=lambda c: (c["row"], c["col"]))], \
        "after the id, each line is the cell's text and nothing the layout adds"
    assert "Revenue (EUR M)" in text and "1F78B5" not in text and "scope" not in text and "gaps" not in text
    assert not {"at", "rot", "lbbox", "layout", "rects", "fill", "scope", "gaps", "association"} & decks.GATEWAY_READABLE_FIELDS
    layout, = parsed["layout"]
    assert set(layout) == {"page", "rects", "gaps"} and set(layout["gaps"]) <= set(GAPS)
    for r in layout["rects"]:
        assert set(r) == {"at", "fill", "scope"} and re.fullmatch(r"[0-9A-F]{6}", r["fill"]) and all(isinstance(x, float) for x in r["at"])
    assert secret not in json.dumps(layout)
    item, = page_checker.check(grid, [tc.claim("r2c1", 999_000)], 12)["items"]
    assert item["drop"] is None and item["flags"][-1] == page_checker.EXTRACTION_INCOMPLETE, "the picture is a gap: kept for extraction review"
    assert set(item["association"]) == {"label", "period", "unit", "flags", "options"}
    assert {item["association"][k] for k in ("label", "period", "unit")} <= set(association.SOURCES) | {None}
    assert set(item["association"]["flags"]) <= set(association.FLAGS)
    stored = json.dumps(page_pipeline._without_text(item))
    for needle in (secret, "Enterprise", "Revenue (EUR M)"):
        assert needle not in stored, f"{needle!r} in a stored item"
    complete = parse_deck(pdf_page([(50, 560, 20, secret), (50, 480, 12, "$150,000 annual revenue run rate")]), "deck.pdf")
    done, = page_grids.page_grids(complete["blocks"], complete["layout"])
    dropped, = page_checker.check(done, [tc.claim("r2c1", 999_000)], 12)["items"]
    assert dropped["drop"] == page_checker.NO_PRINTED_SOURCE == "no printed source"
    db = tc._page_db(complete)
    db[decks.TEXT_COLLECTION].docs[0]["layout"] = complete["layout"]
    adapter = t.FakeAdapter(replies=[json.dumps({"type": "page_claims", "claims": [tc.claim("p1r2c1", 999_000, period="2025")]})])
    with caplog.at_level(logging.DEBUG):
        assert asyncio.run(structures.process_deck(db, tc.AUDIT, "d1", adapter=adapter, sleep=t._noop_sleep)) == "read"
    doc, = db[decks.TEXT_COLLECTION].docs
    assert doc["diagnostics"] == [{"page": 1, "kind": "no printed source", "metric": "revenue", "value": 999_000, "period": "2025"}]
    assert secret not in caplog.text and "annual revenue run rate" not in caplog.text and secret not in json.dumps(doc["diagnostics"])


def test_the_compare_script_prints_counts_values_cells_and_statuses_and_no_line_of_a_deck(capsys):
    """Rule 17 for scripts/compare_pipelines.py: its table and its rows-only lists carry deck file names, counts, rates, values, cell
    ids, types and statuses, never the text of a line. --diagnostic prints the text and runs on the 10 public test decks only."""
    pytest.importorskip("pdfplumber")
    pytest.importorskip("pptx")
    import importlib.util
    from app.decks import parser
    spec = importlib.util.spec_from_file_location("compare_pipelines", BACKEND.parent / "scripts" / "compare_pipelines.py")
    script = importlib.util.module_from_spec(spec)
    spec.loader.exec_module(script)
    names = ["03-buffer.pptx", "05-zero2hero.pdf"]
    report = script.main(["--deck", names[0], "--deck", names[1]])
    assert set(report["decks"]) == set(names) and report["totals"]["rows_b"] > 0 and report["totals"]["b_matched"] > 0
    assert report["decks"][names[1]]["b"]["recall"] >= 0.85, "the recorded replies of zero2hero p17 and p19 are found"
    printed = capsys.readouterr().out
    words = set()
    for name in names:
        for block in parser.parse_deck((script.DECKS / name).read_bytes(), name)["blocks"]:
            text = block["text"]
            if len(text) >= 4 and any(c.isalpha() for c in text) and not any(text in n for n in names):
                words.add(text)
    assert len(words) > 50 and "Traction" in words, "the check sees the decks' lines"
    assert not sorted(w for w in words if w in printed), "a line of a deck was printed"
    with pytest.raises(SystemExit, match="public test decks only"):
        script.main(["--diagnostic", "--deck", "client-deck.pdf"])


def test_the_pages_run_prints_and_reports_no_line_of_a_deck_and_records_only_the_models_claims(monkeypatch, tmp_path, capsys, caplog):
    """Rule 17 for scripts/consistency_run.py --pages (docs/design/ai-reads-rules-check.md section 7, PR B): its progress lines, its report
    (kept in docs/test-runs), its JSON and the gateway's log lines for a stopped run carry deck file names, counts, tokens, cost and
    cell ids, never a line of a deck. The recorded replies (kept under the fixtures) hold the model's claims as the schema allows
    them (closed words, numbers, cell ids and a period) and, to find the page again, its first line as sent: public decks only."""
    pytest.importorskip("pdfplumber")
    pytest.importorskip("pptx")
    import importlib.util
    import logging
    import tempfile
    from app.decks import pages as page_grids
    from app.decks import parser
    from app.llm.schemas import PageClaimsReply
    spec = importlib.util.spec_from_file_location("consistency_run", BACKEND.parent / "scripts" / "consistency_run.py")
    script = importlib.util.module_from_spec(spec)
    spec.loader.exec_module(script)
    monkeypatch.setattr(tempfile, "tempdir", str(tmp_path))
    caplog.set_level(logging.WARNING)
    names = ["03-buffer.pptx", "05-zero2hero.pdf"]
    report = script.main(["--pages", "--fake", "--max-spend", "0.15", "--deck", names[0], "--deck", names[1],
                          "--out", str(tmp_path / "pages_report.json")])
    assert report["stopped"] and report["decks_complete"] == 1, "the run stopped at --max-spend, so the stop's log lines are checked too"
    assert set(report) == {"mode", "model", "prompt", "prompt_version", "passes", "decks_asked", "decks_complete", "max_spend_usd",
                           "spent_usd", "stopped", "agreement_pct_all", "target_agreement_pct", "totals", "per_deck"}
    assert set(report["totals"]) == {"calls", "batches", "read", "claims", "matched", "unmatched", "dropped", "input_tokens",
                                     "output_tokens", "truncated", "cache_hits", "recorded"}
    for deck in report["per_deck"].values():
        assert set(deck) <= {"pages", "sent", "passes", "agreement_pct", "complete", "agreed", "claims_union"}
        assert all(set(row) == {"claims", "matched", "unmatched", "dropped", "input_tokens", "output_tokens", "batches", "read",
                                "calls", "cost_usd", "truncated", "cache_hits", "recorded"} for row in deck["passes"])
    written = capsys.readouterr().out + "".join(p.read_text(encoding="utf-8") for p in tmp_path.glob("pages_*.md")) + \
        (tmp_path / "pages_report.json").read_text(encoding="utf-8") + "\n".join(r.getMessage() for r in caplog.records)
    assert "reason=spend_stop" in written
    words, lines = set(), set()
    for name in names:
        deck = parser.parse_deck((script.DECKS / name).read_bytes(), name)
        for block in deck["blocks"]:
            text = block["text"]
            if len(text) >= 8 and any(c.isalpha() for c in text) and not any(text in n for n in names):
                words.add(text)
        for grid in page_grids.page_grids(deck["blocks"]):
            lines |= {c["text"] for c in grid["cells"] if len(c["text"]) >= 12}
    assert len(words) > 50 and "Traction" in words, "the check sees the decks' lines"
    assert not sorted(w for w in words | lines if w in written), "a line of a deck was written or printed"
    recorded = sorted(tmp_path.glob("pages_*/pass-*/*.json"))
    assert recorded, "the stopped run recorded the pass it paid for"
    for path in recorded:
        body = json.loads(path.read_text(encoding="utf-8"))
        assert set(body) == {"file", "page", "match", "pass", "prompt_version", "model", "claims"} and body["file"] in script.PUBLIC_DECKS
        assert re.fullmatch(rf"p{body['page']}r\d+c\d+( title)?: .*", body["match"]) and "\n" not in body["match"]
        PageClaimsReply.model_validate({"type": "page_claims", "claims": body["claims"]})
        assert all(re.fullmatch(r"p\d+r\d+c\d+", str(c[k])) for c in body["claims"] for k in ("value_cell",)
                   ) and all(set(c) == set(PageClaimsReply.model_fields["claims"].annotation.__args__[0].model_fields) for c in body["claims"])


def test_the_checker_flags_basis_and_scenario_are_closed_words_and_never_gateway_readable():
    """The flags of app/structures/checker.py (built 2026-10-22) are closed strings shown in the review tab; a claim's basis and
    scenario are closed words of the schema. None of the three fields is readable by the gateway, and no flag names deck text."""
    from app.llm import schemas
    from app.structures import checker
    assert checker.FLAGS == ("value not on page", "currency assumed", "unit assumed", "scenario label not on page",
                             "period not on page", "page out of range", "use-of-funds share", "acquisition channel",
                             "turnover, not revenue", "year under bar differs", "pipeline column", "axis tick", "not a plan line",
                             "matches analyst claim", "basis and scenario disagree", "basis conflicts with period",
                             "cell corrected", "sign not on page", "period from page heading", "period from another cell",
                             "period is a range", "growth dated to start year", "label mismatch", "sources disagree", "source ambiguous",
                             "source not associated", "extraction incomplete", "source ambiguous, noted", "no label read")
    assert (checker.DERIVED_FIGURE, checker.LAUNCH_NO_PRODUCT, checker.NO_PRINTED_SOURCE, checker.QUOTED_FIGURE) == \
        ("derived figure, not a claim", "launch date without product name", "no printed source", "quoted figure, not a company claim"), \
        "the Diagnostics kinds of a dropped claim are closed words (George, 2026-10-26; no printed source, source association 2026-10-26)"
    assert checker.INFORMATIONAL_FLAGS == ("cell corrected", "currency assumed", "unit assumed", "period from page heading", "period from another cell",
                                       "source ambiguous, noted"), \
        "the flags that leave a claim in the main list (George, 2026-10-23: the unit and the period only flag)"
    assert set(checker.REVIEW_FLAGS) == set(checker.FLAGS) - set(checker.INFORMATIONAL_FLAGS)
    assert checker.RAISE_CUES == ("use of funds", "use of proceeds", "funding ask", "investment opportunity", "equity investment", "round size", "our ask")
    assert checker.NOT_PLAN_EXEMPT_METRICS == ("market", "pricing", "use_of_funds"), "closed lists of words: no flag carries deck text"
    assert checker.is_informational("year label below: 2022") and not checker.is_informational("year label below: Q3 2022")
    assert checker.BRACKET_NEGATIVE_METRICS == ("ebitda", "net_profit", "gross_profit", "costs", "cash", "burn")
    assert set(schemas.PAGE_BASES) == {"actual", "projection", "scenario"} and set(schemas.PAGE_SCENARIOS) == {"best", "realistic", "worst", "none"}
    assert not {"flags", "basis", "scenario", "figure_type"} & decks.GATEWAY_READABLE_FIELDS
    from schemas import metrics
    assert not set(metrics.DECK_CLAIM_FIELDS) & decks.GATEWAY_READABLE_FIELDS, "contract 6's deck claim fields are server-only"
    from app.llm import gateway as _gateway
    assert _gateway.PROMPT_MISSING_TEXT == "AI reading not run – deck_reading prompt missing in prompt store", "a closed sentence, no deck text"
    grid = {"type": "page", "page": 1, "header_rows": 0, "tables": [], "cells": [
        {"row": 1, "col": 1, "text": "Traction at Target Co", "box": 1, "title": True},
        {"row": 2, "col": 1, "text": "Northbridge revenue $150,000", "box": 2}]}
    claim = {"metric": "revenue", "value": 999, "value_high": None, "unit": "EUR", "unit_other": None, "period": "2027",
             "basis": "scenario", "scenario": "best", "value_cell": "r2c1", "label_cell": "r1c1", "period_cell": None}
    item, = checker.check(grid, [claim], 12)["items"]
    assert set(item["flags"]) <= set(checker.FLAGS) and item["basis"] in schemas.PAGE_BASES and item["scenario"] in schemas.PAGE_SCENARIOS
    assert checker.YEAR_LABEL_FLAG.fullmatch(checker.year_label_flag("2022")) and not checker.YEAR_LABEL_FLAG.fullmatch("year label below: Q3 2022"), \
        "the one flag with a part of the page: a four-digit year label, never a line"
    text = json.dumps(item)
    for needle in ("Target Co", "Northbridge", "Traction", "150,000"):
        assert needle not in text, f"{needle!r} in a checked item"


def test_the_page_cache_hit_log_line_holds_ids_a_hash_and_a_type_never_page_text(monkeypatch, caplog):
    """The "structure cache hit" line (George, 2026-10-23) carries the audit, the step, the content hash and the type only."""
    import logging
    import test_page_usage_attribution as pu
    db = pu.two_audits(monkeypatch)
    adapter = t.FakeAdapter(replies=[json.dumps(pu.REPLY)])
    pu.read(db, pu.AUDIT, "d1", adapter)
    with caplog.at_level(logging.DEBUG):
        pu.read(db, pu.AUDIT, "d1", adapter)
    line, = [r.getMessage() for r in caplog.records if "structure cache hit" in r.getMessage()]
    assert re.fullmatch(r"structure cache hit: run_id=\S+ step=structures hash=[0-9a-f]{64} type=page_claims", line), line
    for needle in ("Paying Users", "annual revenue", "margins", "Traction", "deck.pdf"):
        assert needle not in caplog.text, f"{needle!r} in the log"


def test_the_claim_rejected_line_holds_a_page_number_and_schema_field_names_never_a_value(caplog):
    """"page n: claim rejected (fields)" (George, 2026-10-23): a digit string and closed schema names. A claim whose period, unit_other
    and metric quote a name and a figure never reaches the log."""
    import logging
    bad = {"metric": "Jane Doe (CEO)", "value": 123456789.0, "value_high": None, "unit": "USD", "unit_other": "Northbridge",
           "period": "Northbridge 2022 plan", "basis": "projection", "scenario": "realistic", "value_cell": "p4r1c1",
           "label_cell": None, "period_cell": None}
    reply = json.dumps({"type": "page_claims", "claims": [bad]})
    with caplog.at_level(logging.DEBUG):
        gateway.parse_page_reply(reply, "p4r1c1: x")
    line, = [r.getMessage() for r in caplog.records if "claim rejected" in r.getMessage()]
    assert re.fullmatch(r"page (\d+|\?): claim rejected \((cell not in batch|no value|claim|[a-z_]+)(, [a-z_]+)*\)", line), line
    assert set(re.findall(r"[a-z_]+", line.split("(", 1)[1])) <= gateway.PAGE_FIELDS | {"claim"}
    for needle in ("Jane", "Northbridge", "123456789", "2022"):
        assert needle not in caplog.text, f"{needle!r} in the log"


def test_deck_page_reads_log_and_store_page_numbers_and_counts_only(monkeypatch, caplog):
    """Deck page reads (2026-10-23): the "page n not read" line, the deck's progress, totals, skipped and unread pages are page numbers,
    counts and closed words (a status, "timeout"); none is readable by the gateway and none holds a line of the deck."""
    import logging
    from test_page_reads_speed import Gauge, _page_db, _run_deck, many_pages
    from test_structure_reading import AUDIT
    monkeypatch.setenv("DECK_PIPELINE", "ai")
    caplog.set_level(logging.INFO)
    parsed = many_pages(4, plain={2})
    db = _page_db(parsed)
    _run_deck(db, Gauge(hold=0.01, failing={3}))
    doc, = db[decks.TEXT_COLLECTION].docs
    assert {"read_progress", "reading_since", "skipped_pages", "unread_pages", "retry_pages"}.isdisjoint(decks.GATEWAY_READABLE_FIELDS)
    assert doc["skipped_pages"] == [2] and doc["unread_pages"] == [3]
    lines = [r.getMessage() for r in caplog.records if r.getMessage().startswith("page ")]
    assert lines == [f"page 3 not read: run_id={AUDIT} deck_id=d1 status=not_read reason=-"]
    for needle in ("Traction", "Paying Users", "150,000", "Zero2Hero"):
        assert needle not in json.dumps({k: doc[k] for k in ("skipped_pages", "unread_pages", "retry_pages")}) + " ".join(lines)


@pytest.mark.real_render
def test_the_page_image_fallback_and_the_blocked_read_log_codes_only_and_store_flags_and_closed_words(monkeypatch, caplog):
    """Page images (2026-10-27): `images_unavailable` is a flag the gateway never reads; a deck read blocked by the model allow-list is
    status "blocked" with the panel line PAGE_READER_BLOCKED; the log lines of the drawing are run id, deck id, page, a format
    word and counts, never a file name or a line of the deck."""
    import logging
    from app.decks import render
    from test_page_reads_speed import Gauge, _page_db, _run_deck, many_pages
    from test_structure_reading import AUDIT
    monkeypatch.setenv("DECK_PIPELINE", "ai")
    caplog.set_level(logging.INFO)
    assert "images_unavailable" not in decks.GATEWAY_READABLE_FIELDS and "images_unavailable" not in gateway.OUTBOUND_FIELDS
    assert "images_need_file" not in decks.GATEWAY_READABLE_FIELDS and "images_need_file" not in gateway.OUTBOUND_FIELDS
    # 1. a deck that cannot be drawn, then a page that cannot: the two lines of pages.py
    monkeypatch.setattr(render, "prepare", lambda *a: (_ for _ in ()).throw(render.PageImageError(render.NOT_RENDERED)))
    db = _page_db(many_pages(2))
    _run_deck(db, Gauge(hold=0.0))
    doc, = db[decks.TEXT_COLLECTION].docs
    assert doc["images_unavailable"] is True and [r["status"] for r in doc["page_reads"]] == ["read", "read"]

    async def prepared(*args):
        return b"pdf"
    monkeypatch.setattr(render, "prepare", prepared)
    monkeypatch.setattr(render, "render_page", lambda pdf, page: (_ for _ in ()).throw(render.PageImageError(render.NOT_RENDERED)))
    _run_deck(_page_db(many_pages(1)), Gauge(hold=0.0))
    # 2. the three lines of render.py: no converter, a conversion that fails, a page count that differs
    monkeypatch.setattr(render.shutil, "which", lambda name: None)
    with pytest.raises(render.PageImageError):
        render.to_pdf(b"x", "pptx")
    monkeypatch.undo()
    monkeypatch.setenv("DECK_PIPELINE", "ai")
    monkeypatch.setattr(render, "to_pdf", lambda content, fmt: b"pdf")
    monkeypatch.setattr(render, "page_count", lambda pdf: 7)
    stored = _page_db(many_pages(1))
    asyncio.run(render.save(stored, AUDIT, "d1", b"x", "pdf"))
    with pytest.raises(render.PageImageError):
        asyncio.run(render.prepare(stored, AUDIT, "d1", "pdf", 3))
    lines = [r.getMessage() for r in caplog.records if r.name in ("growth.structures", "growth.decks") and r.getMessage().startswith("page image")]
    shapes = (rf"page images unavailable: run_id={AUDIT} deck_id=d1", rf"page image unavailable: run_id={AUDIT} deck_id=d1 page=\d+",
              r"page images: no converter for format=pptx", r"page images: format=pdf rendered 7 pages, parsed 3")
    for shape in shapes:
        assert [line for line in lines if re.fullmatch(shape, line)], f"no log line like {shape!r}: {lines}"
    assert all(any(re.fullmatch(shape, line) for shape in shapes) for line in lines), lines
    for needle in ("Traction", "Paying Users", "150,000", "deck.pdf", "Zero2Hero"):
        assert needle not in caplog.text, f"{needle!r} in the log"
    # 3. the model allow-list: status "blocked", the panel line, nothing sent, a code-only log line
    caplog.clear()
    monkeypatch.setenv("PAGE_READER_MODEL", "claude-fable-5-1")
    blocked, adapter = _page_db(many_pages(2)), Gauge(hold=0.0)
    _run_deck(blocked, adapter)
    doc, = blocked[decks.TEXT_COLLECTION].docs
    assert adapter.calls == 0 and doc["ai_status"] == "blocked" and doc["ai_message"] == gateway.PAGE_READER_BLOCKED
    assert [r["status"] for r in doc["page_reads"]] == ["blocked", "blocked"] and doc["images_unavailable"] is False
    assert all(r["reason"] == gateway.PAGE_READER_BLOCKED for r in doc["page_reads"]), "a closed text, no deck line"
    page_lines = [r.getMessage() for r in caplog.records if re.match(r"page \d+ ", r.getMessage())]
    assert page_lines and all(re.fullmatch(rf"page \d+ not read: run_id={AUDIT} deck_id=d1 status=blocked reason=-", line) for line in page_lines)
    assert "page reader model claude-fable-5-1 not permitted" in caplog.text and "Traction" not in caplog.text


# --- 2026-10-23 (engine contract 8): the monthly gross revenue churn series is the register's, not the narrative's --------------

MONTHLY_CHURN_SENTINEL = 0.4321


def test_the_monthly_churn_series_never_reaches_a_narrative_payload():
    assert "monthly_series" not in gateway.OUTBOUND_FIELDS
    db = _db()
    block = db["audits"].docs[0]["results"]["gross_churn"] = cf_gross_churn()
    assert block["monthly_series"], "fixture: the series is in the stored results"
    _, sent = _generate(db)
    churn = sent["metrics"]["gross_churn"]
    assert "monthly_series" not in churn and churn["series"], "the trailing-12-month series still goes"
    assert str(MONTHLY_CHURN_SENTINEL) not in json.dumps(sent) and "monthly_series" not in set(_keys(sent))


def cf_gross_churn():
    import contract_fixtures as cf
    return cf.gross_churn(0.061, series=[{"month": "2026-06", "churn_pct": 0.061}],
                          monthly=[{"month": "2026-06", "churn_pct": MONTHLY_CHURN_SENTINEL}])


def test_the_forecast_reason_and_the_ask_reasons_are_closed_sentences():
    """George, 2026-10-25, amended 2026-10-26 (forecast-and-payback.md section 3): the reason of a forecast claim is "Forecast: " and a closed
    phrase, a calculated figure or one of the closed reasons of forecast_calc, no deck text; the reasons that ask the analyst to set something are
    the closed texts of claim_matching."""
    assert claim_matching.FORECAST_REASON.format(text=claim_matching.FORECAST_NO_CALC) == "Forecast: no comparison is calculated for this metric"
    assert all(r == r.strip() for r in forecast_calc.REASONS)
    assert all(r == r.strip() and "{" not in r for r in claim_matching.ASK_REASONS)


def test_corrections_and_dropped_claims_log_values_and_cell_ids_never_deck_text():
    """George, 2026-10-26: what a check corrected or dropped is logged for Diagnostics (the deck document) and stored with the verified
    claims (llm_structures): period labels, cell ids, model values and the closed kind words, never the text of a label cell or a line."""
    import asyncio
    import test_claim_corrections as cc
    from app.llm import gateway as _gateway
    tc = cc.tc
    parsed = tc.deck([(50, 560, 12, "Overview"), (50, 500, 12, "Paying customers"), (50, 480, 12, "800"), (400, 200, 12, "Footnote words"),
                      (50, 440, 12, "Launch Q3 2025"), (50, 420, 12, "$150,000 annual revenue run rate")])
    db = tc._page_db(parsed)
    launch = tc.claim(cc.at(parsed, "Launch Q3 2025"), None, metric="product", unit=None, period="2025-Q3", period_cell=cc.at(parsed, "Launch Q3 2025"))
    label = tc.claim(cc.at(parsed, "800"), 800, metric="customers", unit="count", label_cell=cc.at(parsed, "Footnote words"))
    # A period the model set that the cited cell reads otherwise is kept and logged as "sources disagree" (review of PR #193): the line holds
    # the model's period, the period token the page reads and its cell id, never the cell's text.
    kept = tc.claim(cc.at(parsed, "$150,000 annual revenue run rate"), 150_000, period="2023", period_cell=cc.at(parsed, "Launch Q3 2025"))
    adapter = cc.t.FakeAdapter(replies=[json.dumps({"type": "page_claims", "claims": [launch, label, kept]})])
    assert asyncio.run(cc.structures.process_deck(db, cc.AUDIT, "d1", adapter=adapter, sleep=cc.t._noop_sleep)) == "read"
    deck_doc, = db[cc.TEXT_COLLECTION].docs
    assert {d["kind"] for d in deck_doc["diagnostics"]} == {"corrected", "launch date without product name", "sources disagree"}
    disagree, = [d for d in deck_doc["diagnostics"] if d["kind"] == "sources disagree"]
    assert set(disagree) == {"page", "kind", "field", "original", "read"} and disagree["original"] == "2023"
    assert [(r["period"], r["cells"]) for r in disagree["read"]] == [("2025-Q3", [tc.cell_of(parsed, "Launch Q3 2025")])]
    logged = json.dumps(deck_doc["diagnostics"]) + json.dumps(db[_gateway.STRUCTURES_COLLECTION].docs[0]["items"])
    for needle in ("Paying customers", "Footnote", "Launch Q3", "annual revenue"):
        assert needle not in logged, f"{needle!r} in a Diagnostics line or a stored item"


# ---------------------------------------------------------------------------
# docs/specs/forecast-and-payback.md section 6: the forecast calculation, the payback definition and the annual payback are pure Python over the
# register's rows and the stored results. Nothing reaches the gateway, nothing is logged, and the new fields and reasons hold figures and closed
# words only (CLAUDE.md rules 14 and 17).
# ---------------------------------------------------------------------------
def test_the_forecast_calculation_and_the_payback_definition_import_nothing_from_the_gateway_and_log_nothing():
    for name in ("forecast_calc.py", "payback_definition.py"):
        tree = ast.parse((BACKEND / "app" / name).read_text(encoding="utf-8"))
        offenders = [f"{m} -> {n}" for m, n in _imports(tree)
                     if (set(re.split(r"[.\s\"'()]+", m)) | {n}) & {"llm", "gateway", "anthropic", "structures", "<dynamic>", "server", "logging", "claim_matching"}]
        assert not offenders, f"{name} reaches for the gateway or logs:\n  " + "\n  ".join(offenders)
        calls = [ast.unparse(n.func) for n in ast.walk(tree) if isinstance(n, ast.Call) and isinstance(n.func, ast.Attribute)
                 and n.func.attr in ("debug", "info", "warning", "error", "exception", "critical", "log")]
        assert not calls, f"{name} writes a log line: {calls}"


def test_the_new_register_fields_and_the_yearly_scope_are_names_the_gateway_never_mentions():
    from app import decks as decks_mod
    names = {"calc_lines", "calc_details", "sm_scope", "net_new_arr", "arr_year_end", "arr_prior_year_end"}   # sm_expense is the monthly P&L field the gateway already names
    assert {"calc_lines", "calc_details"} <= set(claim_matching.FIELDS)
    offenders = []
    for path in sorted((BACKEND / "app" / "llm").glob("*.py")):
        tree = ast.parse(path.read_text(encoding="utf-8"))
        offenders += [f"{path.name}:{n.lineno}: {n.value!r}" for n in ast.walk(tree) if isinstance(n, ast.Constant) and n.value in names]
    assert not offenders, "the gateway names the forecast and payback fields:\n  " + "\n  ".join(offenders)
    assert not names & gateway.OUTBOUND_FIELDS and not names & decks_mod.GATEWAY_READABLE_FIELDS
    assert not ({"calc_lines", "calc_details"} & set(gateway.OUTBOUND_TOP_LEVEL))


def test_the_calculation_lines_hold_figures_and_closed_words_never_the_claims_words():
    """A claim whose words hold a sentinel: the register row carries its figures and the closed texts, not its snippet."""
    import forecast_support as fs
    secret = "Falcon Orchid hiring plan"
    row = fs.row_of(fs.claim(f"Revenue €2.1M in FY2026, {secret}", 2_073_600, "FY2026", "2026", basis="projection"))
    assert row["calc_lines"] and all(secret not in str(v) for v in (*row["calc_lines"], *row["calc_details"], row["reason"], row["relationship"]))
    annual = fs.row_of(fs.payback_claim(f"CAC payback 16 months in FY2024, {secret}, sales and marketing over net new ARR, gross margin", 16),
                       results=fs.payback_results())
    assert annual["calc_lines"] and all(secret not in str(v) for v in (*annual["calc_lines"], *annual["calc_details"], annual["reason"]))


def test_the_not_calculated_reasons_hold_closed_words_and_a_period_token_never_the_claims_words():
    """The annual payback reason (ANNUAL_NOT_CALCULATED) names the fiscal year of the claim's period and a closed reason; the forecast reason of a
    claim whose amounts are not comparable (forecast_calc.unavailable) holds the closed mismatch text and the figures. A sentinel in the claim's
    words reaches neither (CLAUDE.md rules 14 and 17)."""
    import forecast_support as fs
    secret = "Falcon Orchid hiring plan"
    assert claim_matching.ANNUAL_NOT_CALCULATED.format(year="FY2021", reason=claim_matching.ANNUAL_NO_YEAR) == \
        "annual payback not calculated for FY2021: no actual year in the yearly P&L"
    text = f"CAC payback 16 months in FY2021, {secret}, sales and marketing over net new ARR, gross margin"
    row = fs.row_of(fs.payback_claim(text, 16, year=2021), results=fs.payback_results())
    assert row["reason"] == "annual payback not calculated for FY2021: no actual year in the yearly P&L", row["reason"]
    unavailable = fs.row_of(fs.claim(f"180 registered users in FY2026, {secret}", 180, "FY2026", "2026", claim_type="customers", currency=None,
                                     basis="projection", population="registered"))
    assert unavailable["calc_lines"][0].startswith("Required growth: not calculated – population differs")
    for r in (row, unavailable):
        assert all(secret not in str(v) for v in (*r["calc_lines"], *r["calc_details"], r["reason"], r["relationship"], r["gloss"]))


def test_the_deck_actuals_reasons_and_details_hold_figures_and_closed_words_and_the_one_label_of_the_inputs_and_reach_no_model():
    """docs/design/deck-actuals.md: "deck differs from file" and "Deck only" are closed sentences with figures; the details block is built from stored
    data with no model call, and the label as printed is held in its inputs and in `not_checked` only (both shown in the register and the memo; and in the segment name of a part, as the segment rule already allows)."""
    from app import deck_actuals
    sentinel = "zephyr holdings"      # lower case: a capitalised leftover word would name a segment, and the row would stop at "segment not in the data"
    base = {"id": "x1", "status": "approved", "origin": "ai", "ai_status": "matched", "claim_type": "revenue", "currency": "EUR", "target_date": "2026-12",
            "period_text": "Dec 2026", "snippet": "EUR 100", "label_from": f"{sentinel} Revenue", "file": "deck.pptx", "order": 0, "basis": "actual",
            "sources": [{"file": "deck.pptx", "slide": 3}]}
    from app.decks import claims as deck_claims
    results = {**_results_with_the_series(), "as_of_month": "2026-12", "mrr_series": {"months": ["2026-12"], "segments": [], "source": SOURCE,
                                                                                    "data": [{"month": "2026-12", "total": 500.0}]}}
    seen = {}
    # 100.0 reconciles, 4800.0 is far beyond 50% off (the reconcile check runs before the Definition check), 4800.0 marked a projection carries the mark
    # note, and a period the files hold nothing for is used as it stands ("Deck only"): every closed reason string and the details text is reached.
    cases = ((15_000_000.0, "2026-12", {}), (SERIES_SENTINEL, "2026-12", {}), (15_000_000.0, "2026-12", {"basis": "projection"}), (4800.0, "2025-12", {}))
    for value, period, extra in cases:
        out = claim_matching.build_register([deck_claims.resolve_period({**base, "value": value, "target_date": period, **extra}, 12)], results,
                                            {"fiscal_year_end": 12, "reporting_currency": "EUR", "fx": {"EUR": 1.0}})[0]
        assert out["deck_status"] in (deck_actuals.DIFFERS_STATUS, deck_actuals.TAG, deck_actuals.CONSISTENT), out["reason"]
        seen[out["deck_status"]] = out
        elsewhere = {k: v for k, v in out.items() if k not in ("details", "segment", "relationship", "gate_sentence")}
        assert sentinel not in json.dumps(elsewhere, default=str)
        assert out["details"], "every deck actual has its block"
        assert [i["label"] for i in out["details"]["inputs"]] == [f"{sentinel} Revenue"]
        everything_else = {k: v for k, v in out["details"].items() if k not in ("inputs", "not_checked")}
        assert sentinel not in json.dumps(everything_else), "formula, checks and note say the metric and closed words, never the label"
        assert out["details"]["not_checked"] in ([], [f"{sentinel} Revenue, Dec 2026 (slide 3)"]) or out["details"]["not_checked"][0].startswith(f"{sentinel} Revenue")
    assert set(seen) == {deck_actuals.DIFFERS_STATUS, deck_actuals.TAG, deck_actuals.CONSISTENT}, "both new reasons and the consistent row were reached"
    assert seen[deck_actuals.DIFFERS_STATUS]["reason"].startswith(deck_actuals.DIFFERS) and seen[deck_actuals.TAG]["reason"] == deck_actuals.DECK_ONLY_REASON
    assert " by " in seen[deck_actuals.DIFFERS_STATUS]["details"]["checked"][0] and "matches" not in seen[deck_actuals.DIFFERS_STATUS]["details"]["checked"][0]
    noted = claim_matching.build_register([deck_claims.resolve_period({**base, "value": 15_000_000.0, "basis": "projection"}, 12)], results,
                                          {"fiscal_year_end": 12, "reporting_currency": "EUR", "fx": {"EUR": 1.0}})[0]["details"]["note"]
    assert noted == deck_actuals.mark_note({"basis": "projection"}) and sentinel not in noted
    assert deck_actuals.DIFFERS_REASON.format(deck="€1", file="€2") == "deck differs from file: deck €1, file €2"
    assert "app.llm" not in inspect.getsource(deck_actuals) and "gateway" not in inspect.getsource(deck_actuals)
