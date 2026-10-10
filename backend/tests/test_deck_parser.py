"""Deck parser and candidate claims (docs/specs/deck-parser.md sections 1, 2 and 5).

Small decks are built in the test with the same libraries the parser reads them with; the
public decks in tests/fixtures/decks/decks/ check references on real layouts. Endpoints run
against the in-memory Mongo stub.
"""
import io
import os
import sys
from pathlib import Path

import pytest

pptx = pytest.importorskip("pptx")
docx = pytest.importorskip("docx")
pytest.importorskip("pdfplumber")

BACKEND = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(BACKEND))
sys.path.insert(0, str(BACKEND / "tests"))

from app import decks  # noqa: E402
from app.decks import claims, parser  # noqa: E402
from app.structures import items, verify  # noqa: E402

DECKS = BACKEND.parent / "tests" / "fixtures" / "decks" / "decks"


def _pptx(slides, notes=None, groups=None, tables=None) -> bytes:
    """slides: list of lists of text-box texts. notes/groups/tables: {slide index: ...}."""
    from pptx.util import Inches
    prs = pptx.Presentation()
    for i, texts in enumerate(slides):
        slide = prs.slides.add_slide(prs.slide_layouts[6])
        for j, text in enumerate(texts):
            slide.shapes.add_textbox(Inches(1), Inches(1 + j), Inches(6), Inches(1)).text_frame.text = text
        for text in (groups or {}).get(i, []):
            slide.shapes.add_group_shape().shapes.add_textbox(Inches(1), Inches(5), Inches(4), Inches(1)).text_frame.text = text
        if i in (tables or {}):
            rows = tables[i]
            table = slide.shapes.add_table(len(rows), len(rows[0]), Inches(1), Inches(3), Inches(6), Inches(2)).table
            for r, row in enumerate(rows):
                for c, text in enumerate(row):
                    table.cell(r, c).text = text
        if i in (notes or {}):
            slide.notes_slide.notes_text_frame.text = notes[i]
    buf = io.BytesIO()
    prs.save(buf)
    return buf.getvalue()


def _pdf(pages: int) -> bytes:
    """A PDF of blank pages: no text, as a scanned file would have."""
    import pypdfium2 as pdfium
    pdf = pdfium.PdfDocument.new()
    for _ in range(pages):
        pdf.new_page(612, 792)
    buf = io.BytesIO()
    pdf.save(buf)
    return buf.getvalue()


def _has(blocks, **fields):
    """A block with these fields (it may also carry layout fields: box, bbox, title)."""
    return any(all(b.get(k) == v for k, v in fields.items()) for b in blocks)


def _by_value(candidates, value):
    """The candidates stating the value, a table row candidate among its values by period."""
    return [c for c in candidates if any(v["value"] == value for v in claims.claim_values(c))]


# ---------------------------------------------------------------------------
# Parsing and references
# ---------------------------------------------------------------------------
def test_pptx_text_boxes_groups_tables_and_notes_keep_their_slide():
    content = _pptx(
        [["Traction", "ARR grew to $3.6M by Q4 2024"], ["Plan"], ["Recap: ARR grew to $3.6M by Q4 2024"]],
        notes={1: "Team grows to 40 hires in 2025"},
        groups={0: ["Churn is 5%"]},
        tables={1: [["", "2023", "2024"], ["Revenue", "$1.2M", "$2.5M"]]},
    )
    deck = parser.parse_deck(content, "board.pptx")
    assert (deck["format"], deck["page_unit"], deck["pages"]) == ("pptx", "slide", 3)
    assert _has(deck["blocks"], slide=1, kind="text", text="Churn is 5%"), "text inside a group shape"
    assert _has(deck["blocks"], slide=2, kind="notes", text="Team grows to 40 hires in 2025")
    assert _has(deck["blocks"], slide=2, kind="table", table=1, row=2, col=3, text="$2.5M")

    found = claims.detect_candidates(deck["blocks"], "board.pptx")
    arr, = _by_value(found, 3600000)
    assert (arr["claim_type"], arr["currency"], arr["target_date"]) == ("revenue", "USD", "2024-Q4")
    assert arr["sources"] == [{"file": "board.pptx", "slide": 1, "kind": "text"},
                              {"file": "board.pptx", "slide": 3, "kind": "text"}], "merged, every source kept"
    churn, = _by_value(found, 5)
    assert (churn["claim_type"], churn["unit"]) == ("retention", "%")
    revenue, = _by_value(found, 2500000)
    cell = {"file": "board.pptx", "slide": 2, "kind": "table", "table": 1, "row": 2}
    assert revenue["sources"] == [{**cell, "col": 2}, {**cell, "col": 3}], "one candidate for the row"
    assert (revenue["snippet"], revenue["value"], revenue["target_date"]) == ("Revenue | $1.2M | $2.5M", None, None)
    assert revenue["by_period"] == [
        {"value": 1200000, "value_high": None, "target_date": "2023", "period": "2023", "period_text": "2023",
         "period_start": "2023-01-01", "period_end": "2023-12-31", "source": {**cell, "col": 2}},
        {"value": 2500000, "value_high": None, "target_date": "2024", "period": "2024", "period_text": "2024",
         "period_start": "2024-01-01", "period_end": "2024-12-31", "source": {**cell, "col": 3}}]
    hires, = _by_value(found, 40)
    assert hires["claim_type"] == "people" and hires["target_date"] == "2025"
    assert hires["sources"] == [{"file": "board.pptx", "slide": 2, "kind": "notes"}]


def test_docx_pages_follow_page_breaks_and_new_page_sections():
    from docx.enum.section import WD_SECTION
    doc = docx.Document()
    doc.add_paragraph("Board update")
    doc.add_page_break()
    doc.add_paragraph("Revenue was €2M in 2023")                     # page 2
    doc.add_section(WD_SECTION.NEW_PAGE)
    table = doc.add_table(rows=2, cols=2)                             # page 3
    table.cell(0, 0).text, table.cell(0, 1).text = "Metric", "Value"
    table.cell(1, 0).text, table.cell(1, 1).text = "Churn", "4%"
    doc.add_section(WD_SECTION.CONTINUOUS)
    doc.add_paragraph("Headcount 25 by March 2026")                   # still page 3
    buf = io.BytesIO()
    doc.save(buf)

    deck = parser.parse_deck(buf.getvalue(), "plan.docx")
    assert deck["pages"] == 3
    assert _has(deck["blocks"], page=2, kind="text", text="Revenue was €2M in 2023")
    assert _has(deck["blocks"], page=3, kind="table", table=1, row=2, col=2, text="4%")
    assert _has(deck["blocks"], page=3, kind="text", text="Headcount 25 by March 2026")
    found = claims.detect_candidates(deck["blocks"], "plan.docx")
    churn, = _by_value(found, 4)
    assert churn["sources"] == [{"file": "plan.docx", "page": 3, "kind": "table", "table": 1, "row": 2, "col": 2}]
    people, = _by_value(found, 25)
    assert (people["claim_type"], people["target_date"]) == ("people", "2026-03")


@pytest.mark.parametrize("file, page, value, kind", [
    ("01-front-b.pptx", 13, 150, "text"),             # 150% net retention rate at 1 year
    ("03-buffer.pptx", 5, 150000, "text"),            # $150,000 annual revenue run rate
    ("02-moz.pdf", 21, 25, "text"),                   # Churn Rate in 1st 2 Paid Months ~25%
    ("09-genesisai-2024.pdf", 5, 8000, "text"),       # $8,000 revenue in 2022
    ("05-zero2hero.pdf", 19, 130550, "table"),        # Revenue row of the projections table
    ("04-clevergig.docx", 6, 260, "text"),            # €260 MRR per client
    ("04-clevergig.docx", 9, 200000000, "text"),      # TAM of €200M
    ("04-clevergig.docx", 11, 20, "text"),            # teams to a total of 20 end of 2020
    ("07-equals-seed.docx", 10, 24, "text"),          # To fund an initial team for 24 months
])
def test_public_decks_cite_the_right_slide_or_page(file, page, value, kind):
    content = (DECKS / file).read_bytes()
    deck = parser.parse_deck(content, file)
    hits = [s for c in _by_value(claims.detect_candidates(deck["blocks"], file), value) for s in c["sources"]]
    assert any(s.get("slide", s.get("page")) == page and s["kind"] == kind for s in hits), hits


def test_zero2hero_table_figures_cite_row_and_column():
    file = "05-zero2hero.pdf"
    deck = parser.parse_deck((DECKS / file).read_bytes(), file)
    found = claims.detect_candidates(deck["blocks"], file)
    row, = [c for c in found if c["sources"][0] == {**c["sources"][0], "page": 19, "kind": "table", "row": 4}]
    assert (row["claim_type"], row["currency"]) == ("revenue", "GBP")
    assert [(i["source"]["row"], i["source"]["col"], i["value"], i["period"], i["target_date"]) for i in row["by_period"]] == [
        (4, 2, 130550, "Y/E 22", "2022"), (4, 3, 150000, "Y/E 23", "2023"), (4, 4, 250000, "Y/E 24", "2024"),
        (4, 5, 1000000, "Y/E 25", "2025"), (4, 6, 2500000, "Y/E 26", "2026")]


# ---------------------------------------------------------------------------
# Refusals
# ---------------------------------------------------------------------------
def test_a_file_with_no_readable_text_is_refused_with_the_spec_message():
    expected = ("No readable text found in this file. It may be scanned or made of images. "
                "Please upload a text-based version.")
    for content, name in ((_pdf(2), "scan.pdf"), (_pptx([[]]), "pictures.pptx")):
        with pytest.raises(parser.DeckError) as err:
            parser.parse_deck(content, name)
        assert err.value.message == expected


@pytest.mark.parametrize("name, message", [
    ("deck.key", "We cannot read Keynote files. Please export it as PowerPoint or PDF first."),
    ("deck.ppt", parser.UNSUPPORTED),
    ("https://docs.google.com/presentation/d/abc", parser.UNSUPPORTED),
])
def test_out_of_scope_formats_are_refused(name, message):
    with pytest.raises(parser.DeckError) as err:
        parser.parse_deck(b"anything", name)
    assert err.value.message == message


def test_limits_50_mb_and_200_pages():
    with pytest.raises(parser.DeckError) as err:
        parser.parse_deck(b"\0" * (parser.MAX_BYTES + 1), "big.pdf")
    assert err.value.message == "This file is larger than 50 MB. Please upload a file of 50 MB or less."
    with pytest.raises(parser.DeckError) as err:
        parser.parse_deck(_pdf(201), "long.pdf")
    assert err.value.message.startswith("This file has 201 pages. We read up to 200 slides or pages per file.")
    with pytest.raises(parser.DeckError) as err:
        parser.parse_deck(_pptx([["ARR $1M"]] * 201), "long.pptx")
    assert err.value.message.startswith("This file has 201 slides.")
    assert parser.parse_deck(_pptx([["ARR $1M"]] * 200), "ok.pptx")["pages"] == 200


def test_a_damaged_file_is_refused_without_quoting_it():
    with pytest.raises(parser.DeckError) as err:
        parser.parse_deck(b"Jane Doe (CEO) not a zip", "board.pptx")
    assert err.value.message == parser.UNREADABLE.format(kind="PowerPoint")
    assert "Jane Doe" not in err.value.message


# ---------------------------------------------------------------------------
# Candidate detection
# ---------------------------------------------------------------------------
def _line(text):
    return claims.line_candidates(text, [(0, len(text), {"file": "f.pdf", "page": 1, "kind": "text"})])


def test_a_line_needs_a_number_and_a_keyword():
    assert _line("We grew fast and churn is low") == []          # keyword, no number
    assert _line("Section 111, Row 15") == []                    # numbers, no keyword, nothing counted
    assert [c["value"] for c in _line("800 paying users, ARR $1.2M")] == [800, 1200000]


@pytest.mark.parametrize("text, value, unit, currency, date", [
    ("Revenues stand at €15K MRR", 15000, None, "EUR", None),
    ("2011 Estimated Revenue $12 - $13 million", 12000000, None, "USD", "2011"),
    ("Current Revenue Run Rate (June) ~$10.8 million", 10800000, None, "USD", None),
    ("Revenue £ 150,000", 150000, None, "GBP", None),
    ("4.9x more growth per user", 4.9, "x", None, None),
    ("150% net retention rate", 150, "%", None, None),
    ("Our BHAG is €5M ARR by end of 2024", 5000000, None, "EUR", "2024"),
    ("To fund an initial team for 24 months.", 24, "months", None, None),
    ("A TAM of USD 200M", 200000000, None, "USD", None),
    ("Market size ($52 B)", 52000000000, None, "USD", None),
])
def test_figures_are_read_with_unit_currency_and_date(text, value, unit, currency, date):
    first = _line(text)[0]
    assert (first["value"], first["unit"], first["currency"], first["target_date"]) == (value, unit, currency, date)


@pytest.mark.parametrize("text, dates", [
    ("Launched web app January 2011", ["2011-01"]),
    ("1981 2001 Feb. 2007 Oct. 2008 July 2011", ["2007-02", "2008-10", "2011-07"]),
    ("Q1 17\tQ2 17\tQ3 17", ["2017-Q1", "2017-Q2", "2017-Q3"]),
    ("Release planned 2021 Q3", ["2021-Q3"]),
    ("Iphone dev license applied for Nov28,08", ["2008-11"]),
    ("Aiming for $xx ARR by end of 2018", ["2018"]),
])
def test_a_line_whose_figures_are_dates_gives_one_candidate_per_date(text, dates):
    found = _line(text)
    assert [c["target_date"] for c in found] == dates and all(c["value"] is None for c in found)


def test_ordinals_and_lone_years_next_to_a_month_are_not_figures():
    assert _line("We are raising a seed round in the 2nd half of 2019") == []
    assert _line("Launch an app-store, host 3rd party apps") == []
    assert _line("Mayor of 2021") == []                           # "May" is a month only as a word


def test_snippet_is_at_most_300_characters_and_shows_the_figure():
    text = "Background " * 40 + "ARR reached $9.9M in 2024 " + "and more " * 40
    snippet = _line(text)[0]["snippet"]
    assert len(snippet) <= 300 and "$9.9M" in snippet


# ---------------------------------------------------------------------------
# Endpoints and storage
# ---------------------------------------------------------------------------
@pytest.fixture()
def api(monkeypatch):
    pytest.importorskip("fastapi")
    pytest.importorskip("motor")
    os.environ.setdefault("MONGO_URL", "mongodb://localhost:27017")
    os.environ.setdefault("DB_NAME", "deck_parser_test")
    from fastapi.testclient import TestClient
    import server
    import test_llm_gateway as t
    db = t.FakeDB()
    db["audits"].docs += [{"id": "audit-1", "company_name": "Testco", "results": None}, {"id": "audit-2", "company_name": "Other Co", "results": None}]
    monkeypatch.setattr(server, "db", db)
    return TestClient(server.app, raise_server_exceptions=False), db


def _upload(client, audit, name, content):
    return client.post(f"/api/audits/{audit}/decks/upload", files={"file": (name, content)})


def test_upload_stores_parsed_text_and_candidates_linked_to_the_audit(api):
    client, db = api
    r = _upload(client, "audit-1", "03-buffer.pptx", (DECKS / "03-buffer.pptx").read_bytes())
    assert r.status_code == 200, r.text
    assert r.json()["pages"] == 13 and r.json()["candidates"] > 0
    stored, = db[decks.TEXT_COLLECTION].docs
    assert stored["audit_id"] == "audit-1"
    assert _has(stored["blocks"], slide=5, kind="text", text="$150,000 annual revenue run rate")
    assert all(c["audit_id"] == "audit-1" and c["status"] == "pending" for c in db[decks.CANDIDATES_COLLECTION].docs)

    listed = client.get("/api/audits/audit-1/decks").json()
    assert [d["file"] for d in listed["decks"]] == ["03-buffer.pptx"] and "blocks" not in listed["decks"][0]
    run_rate = next(c for c in listed["candidates"] if c["value"] == 150000)
    assert run_rate["sources"][0] == {"file": "03-buffer.pptx", "slide": 5, "kind": "text"}

    # The same file again replaces its earlier parse.
    _upload(client, "audit-1", "03-buffer.pptx", (DECKS / "03-buffer.pptx").read_bytes())
    assert len(db[decks.TEXT_COLLECTION].docs) == 1
    assert len(db[decks.CANDIDATES_COLLECTION].docs) == r.json()["candidates"]


def test_refusals_reach_the_upload_screen_as_400_with_the_message(api):
    client, _ = api
    r = _upload(client, "audit-1", "scan.pdf", _pdf(1))
    assert r.status_code == 400 and r.json()["detail"] == parser.NO_TEXT
    r = _upload(client, "audit-1", "deck.key", b"keynote")
    assert r.status_code == 400 and r.json()["detail"] == parser.KEYNOTE
    assert _upload(client, "no-such-audit", "deck.pdf", _pdf(1)).status_code == 404


def test_approve_reject_and_edit_a_candidate(api):
    client, db = api
    _upload(client, "audit-1", "03-buffer.pptx", (DECKS / "03-buffer.pptx").read_bytes())
    first, second = db[decks.CANDIDATES_COLLECTION].docs[:2]
    url = "/api/audits/audit-1/decks/candidates/{}"
    assert client.put(url.format(first["id"]), json={"status": "approved"}).json()["status"] == "approved"
    assert client.put(url.format(first["id"]), json={"status": "rejected"}).json()["status"] == "rejected"

    before = {k: second[k] for k in ("claim_type", "value", "value_high", "unit", "currency", "target_date")}
    edited = client.put(url.format(second["id"]), json={"claim_type": "sales", "value": 42, "unit": None}).json()
    assert edited["status"] == "edited" and (edited["claim_type"], edited["value"], edited["unit"]) == ("sales", 42, None)
    assert edited["parsed"] == before, "what the parser found is kept next to the edit"
    assert edited["snippet"] == second["snippet"] and edited["sources"] == second["sources"]
    ranged = client.put(url.format(second["id"]), json={"value": 12000000, "value_high": 13000000}).json()
    assert (ranged["value"], ranged["value_high"], ranged["parsed"]) == (12000000, 13000000, before)

    assert client.put(url.format(second["id"]), json={"snippet": "typed"}).status_code == 400, "evidence is not editable"
    assert client.put(url.format(second["id"]), json={"status": "approved", "value": 1}).status_code == 400
    assert client.put(url.format(second["id"]), json={"target_date": "June"}).status_code == 422
    assert client.put(url.format(second["id"]), json={"unit": "x" * 41}).status_code == 422
    assert client.put("/api/audits/audit-2/decks/candidates/" + second["id"], json={"status": "approved"}).status_code == 404


def test_delete_audit_removes_its_parsed_text_and_candidates_only(api):
    client, db = api
    for audit in ("audit-1", "audit-2"):
        assert _upload(client, audit, "03-buffer.pptx", (DECKS / "03-buffer.pptx").read_bytes()).status_code == 200
    r = client.request("DELETE", "/api/audits/audit-1", json={"confirm": "Testco"})
    assert r.status_code == 200, r.text
    assert r.json()["decks_purged"]["deck_text"] == 1 and r.json()["decks_purged"]["deck_candidates"] > 0
    for name in (decks.TEXT_COLLECTION, decks.CANDIDATES_COLLECTION):
        assert {d["audit_id"] for d in db[name].docs} == {"audit-2"}


def test_reseeding_the_demo_audits_removes_their_parsed_text_and_candidates(api):
    """The startup re-seed deletes the old demo audits; their deck text goes with them."""
    import asyncio
    import server
    client, db = api
    db["audits"].docs.append({"id": "old-demo", "demo": True, "seed_version": 2, "results": None})
    for audit in ("old-demo", "audit-1"):
        assert _upload(client, audit, "03-buffer.pptx", (DECKS / "03-buffer.pptx").read_bytes()).status_code == 200
    asyncio.run(server.seed_demo())
    for name in (decks.TEXT_COLLECTION, decks.CANDIDATES_COLLECTION):
        assert {d["audit_id"] for d in db[name].docs} == {"audit-1"}


# ---------------------------------------------------------------------------
# Borrowing a keyword and a date from nearby text (spec section 2)
# ---------------------------------------------------------------------------
def _slide(boxes, title=None, table=None) -> bytes:
    """One slide. boxes: [(text, left_in, top_in)]; table: (rows, left_in, top_in)."""
    from pptx.util import Inches
    prs = pptx.Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[5 if title else 6])
    if title:
        slide.shapes.title.text = title
    for text, left, top in boxes:
        slide.shapes.add_textbox(Inches(left), Inches(top), Inches(2), Inches(0.5)).text_frame.text = text
    if table:
        rows, left, top = table
        grid = slide.shapes.add_table(len(rows), len(rows[0]), Inches(left), Inches(top), Inches(4), Inches(1)).table
        for r, row in enumerate(rows):
            for c, text in enumerate(row):
                grid.cell(r, c).text = text
    buf = io.BytesIO()
    prs.save(buf)
    return buf.getvalue()


def _found(content, name="deck.pptx"):
    deck = parser.parse_deck(content, name)
    return claims.detect_candidates(deck["blocks"], name)


def test_a_figure_borrows_its_label_from_the_same_text_box():
    found = _found(_slide([("Gross churn\n4.5%", 1, 2)]))
    churn, = _by_value(found, 4.5)
    assert (churn["claim_type"], churn["snippet"], churn["label_from"]) == ("retention", "4.5%", "Gross churn")
    assert churn["sources"] == [{"file": "deck.pptx", "slide": 1, "kind": "text"}], "the figure's own line is cited"


def test_a_figure_borrows_its_label_and_date_from_the_table_column_header():
    found = _found(_slide([], table=([["Year", "ARR"], ["2024", "$3M"], ["2025", "$5M"]], 1, 2)))
    arr, = _by_value(found, 5000000)
    assert (arr["claim_type"], arr["target_date"]) == ("revenue", "2025")
    found = _found(_slide([], table=([["Metric", "2024", "2025"], ["Pipeline", "$1M", "$2M"]], 1, 2)))
    row, = found
    assert [(i["value"], row["claim_type"], i["target_date"]) for i in row["by_period"]] == \
        [(1000000, "sales", "2024"), (2000000, "sales", "2025")]


def test_a_figure_borrows_by_position_from_the_same_row_or_above_never_below():
    beside = _found(_slide([("Gross churn", 1, 3), ("6%", 3.5, 3)]))
    assert [(c["value"], c["claim_type"]) for c in beside] == [(6, "retention")]
    above = _found(_slide([("Gross churn", 1, 3), ("6%", 1, 3.6)]))
    assert [(c["value"], c["claim_type"]) for c in above] == [(6, "retention")]
    below = _found(_slide([("6%", 1, 3), ("Gross churn", 1, 3.6)]))
    assert below == []
    too_far = _found(_slide([("Gross churn", 0.2, 0.2), ("6%", 8, 6.5)]))
    assert too_far == [], f"beyond REACH ({claims.REACH} of the slide)"


def test_a_figure_borrows_the_slide_title_last():
    found = _found(_slide([("$4.2M", 8, 6.5)], title="Pipeline coverage"))
    pipeline, = _by_value(found, 4200000)
    assert (pipeline["claim_type"], pipeline["snippet"], pipeline["label_from"]) == ("sales", "$4.2M", "Pipeline coverage")
    nearer = _found(_slide([("Win rate", 7, 6.5), ("31%", 8.5, 6.5)], title="Pipeline coverage"))
    assert [c["claim_type"] for c in _by_value(nearer, 31)] == ["sales"]
    assert _by_value(nearer, 31)[0]["label_from"] == "Win rate", "the nearest label wins over the title"


def test_a_product_line_without_a_figure_takes_a_nearby_date():
    found = _found(_slide([("Launch the API\nOctober 2026", 1, 2)]))
    api, = found
    assert (api["claim_type"], api["value"], api["target_date"]) == ("product", None, "2026-10")
    assert (api["snippet"], api["label_from"], api["date_from"]) == ("Launch the API", None, "October 2026")
    hire = _found(_slide([("Hire a CFO\nOctober 2026", 1, 2)]))
    assert [c["snippet"] for c in hire] == ["October 2026"], "a line without a figure counts only when it is a product line"


def test_a_roadmap_bullet_takes_the_quarter_beside_it():
    file = "10-tea.pdf"
    found = claims.detect_candidates(parser.parse_deck((DECKS / file).read_bytes(), file)["blocks"], file)
    gluon = [c for c in found if c["snippet"] == "Gluon wallet"]
    assert [(c["claim_type"], c["value"], c["target_date"]) for c in gluon] == [("product", None, "2021-Q2")]
    assert [s["page"] for s in gluon[0]["sources"]] == [11] and gluon[0]["date_from"] == "2021 Q2"


def test_a_pdf_value_borrows_the_label_on_its_row():
    file = "02-moz.pdf"
    found = claims.detect_candidates(parser.parse_deck((DECKS / file).read_bytes(), file)["blocks"], file)
    cac, = [c for c in _by_value(found, 100) if c["currency"] == "USD"]
    assert (cac["claim_type"], cac["snippet"], cac["label_from"]) == ("cac", "~$100", "Avg. Cost of Paid Acquisition")


@pytest.mark.parametrize("text, family", [
    ("Turnover £49,284", "revenue"), ("Avg. Customer Lifetime Value ~$900", "ltv"), ("LTV $240", "ltv"),
    ("Implied Customer Life ~9 Months", "customer_lifetime"), ("% of Free Trials Converting to Paid ~57%", "sales"),
    ("30% of our leads come via referrals", "sales"),
    ("Recruit 3 engineers", "people"), ("Low attrition: 0", "people"), ("Ship v2 in 40 days", "product"),
    ("Milestone 3 reached", "product"),
])
def test_added_keywords(text, family):
    assert _line(text)[0]["claim_type"] == family


@pytest.mark.parametrize("text", ["Marketing spend is 20% of budget", "30% of marketplace transactions"])
def test_market_is_a_whole_word(text):
    assert _line(text) == []


@pytest.mark.parametrize("text, low, high, unit, currency", [
    ("2011 Estimated Revenue $12 - $13 million", 12000000, 13000000, None, "USD"),
    ("Churn of 5-10%", 5, 10, "%", None),
    ("Team grows from 40 to 100", 40, 100, None, None),
])
def test_a_range_is_one_candidate_with_low_and_high(text, low, high, unit, currency):
    c, = _line(text)
    assert (c["value"], c["value_high"], c["unit"], c["currency"]) == (low, high, unit, currency)


# ---------------------------------------------------------------------------
# Unpack cap for zip-based formats (spec section 1)
# ---------------------------------------------------------------------------
def _zip(parts=1, unpacked=0) -> bytes:
    import zipfile
    buf = io.BytesIO()
    with zipfile.ZipFile(buf, "w", zipfile.ZIP_DEFLATED, compresslevel=1) as z:
        for i in range(parts):
            z.writestr(f"ppt/part{i}.xml", b"")
        if unpacked:
            with z.open("ppt/big.xml", "w", force_zip64=True) as f:
                chunk = b"\0" * (1 << 20)
                for _ in range(unpacked // len(chunk)):
                    f.write(chunk)
    return buf.getvalue()


@pytest.mark.parametrize("name", ["deck.pptx", "plan.docx"])
def test_a_zip_that_unpacks_too_far_is_refused_before_it_is_opened(name):
    assert parser.MAX_UNPACKED_BYTES == 250 * 1024 * 1024 and parser.MAX_PARTS == 5000
    big = _zip(unpacked=251 * 1024 * 1024)
    assert len(big) < 2 * 1024 * 1024, "small on disk, large unpacked"
    with pytest.raises(parser.DeckError) as err:
        parser.parse_deck(big, name)
    assert err.value.message == ("This file unpacks to more than 250 MB. We read files that unpack to 250 MB or less. "
                                 "Please save a copy with fewer or smaller pictures and upload it again.")
    with pytest.raises(parser.DeckError) as err:
        parser.parse_deck(_zip(parts=5001), name)
    assert err.value.message == ("This file holds more than 5,000 parts. We read files with up to 5,000 parts. "
                                 "Please save a simpler copy and upload it again.")


def test_a_zip_at_the_limits_is_opened():
    """5,000 parts and 250 MB pass the cap; this zip then fails as not a real deck."""
    for content in (_zip(parts=5000), _zip(unpacked=250 * 1024 * 1024)):
        with pytest.raises(parser.DeckError) as err:
            parser.parse_deck(content, "deck.pptx")
        assert err.value.message == parser.UNREADABLE.format(kind="PowerPoint")


# ---------------------------------------------------------------------------
# Claim types (customers, users, gross margin, usage, growth by noun)
# ---------------------------------------------------------------------------
def test_buffer_slide_5_types_and_units():
    """The browser test that found the wrong types: each line keeps its own keyword."""
    file = "03-buffer.pptx"
    found = claims.detect_candidates(parser.parse_deck((DECKS / file).read_bytes(), file)["blocks"], file)
    slide5 = [(c["claim_type"], c["value"], c["unit"], c["currency"], c["snippet"], c["label_from"])
              for c in found if c["sources"][0]["slide"] == 5]
    assert slide5 == [
        ("customers", 800, "paying users", None, "800 Paying Users", None),
        ("revenue", 150000, None, "USD", "$150,000 annual revenue run rate", None),
        ("gross_margin", 97, "%", None, "97% margins", None),
        ("users", 55000, "users", None, "55,000 users, growing 40% per month", None),
        ("user_growth", 40, "%", None, "55,000 users, growing 40% per month", None),
        # No plan keyword of its own: it borrows the nearest label in its text box, shown to the analyst.
        ("users", 1500000, "updates", None, "1.5 million updates Buffered", "55,000 users, growing 40% per month"),
    ]


@pytest.mark.parametrize("text, family", [
    ("Revenues stand at €15K MRR and are growing 15% MoM", "revenue_growth"),
    ("55,000 users, growing 40% per month", "user_growth"),
    ("Proven record of 2X+ growth for 4 years", "growth"),
    ("ARR grew 3x", "revenue_growth"),
    ("A CAGR of 17.2% from 2021 to 2026", "growth"),
])
def test_growth_takes_its_type_from_the_noun_on_its_line(text, family):
    assert [c["claim_type"] for c in _line(text) if c["unit"] in ("%", "x")] == [family]


def test_an_amount_beside_a_growth_word_takes_the_noun_type():
    assert [c["claim_type"] for c in _line("ARR grew to $3.6M")] == ["revenue"]
    assert [c["claim_type"] for c in _line("Users grew to 55,000")] == ["users"]


@pytest.mark.parametrize("text, family, unit", [
    ("800 Paying Users", "customers", "paying users"),
    ("or some 1,500 clients", "customers", "clients"),
    ("12 enterprise accounts", "customers", None),
    ("2,300 companies use the product", "customers", "companies"),
    ("55 agencies now book shifts", "customers", "agencies"),
    ("10K+ subscribers", "customers", "subscribers"),
    ("2,600+ users", "users", "users"),
    ("Gross Margins ~82%", "gross_margin", "%"),
    ("Avg. Customer Lifetime Value ~$900", "ltv", None),
])
def test_customers_users_and_margin(text, family, unit):
    c = _line(text)[0]
    assert c["claim_type"] == family and (unit is None or c["unit"] == unit)


@pytest.mark.parametrize("text, unit", [
    (">50 Dutch temporary work agencies", "agencies"),
    ("12 enterprise accounts", "accounts"),
    ("40 new mid-market paying customers", "customers"),
    ("300 big US enterprise brand customers", "big"),
])
def test_the_unit_is_a_keyword_noun_within_four_words_else_the_next_word(text, unit):
    assert _line(text)[0]["unit"] == unit


def test_the_unit_search_stops_at_the_next_figure():
    units = {c["value"]: c["unit"] for c in _line("5 advisors & 15 clients now signed")}
    assert units == {5: "advisors", 15: "clients"}


def test_there_is_no_usage_type_a_count_without_a_label_is_not_a_candidate():
    assert "usage" not in claims.CLAIM_TYPES
    assert _line("1.5 million updates Buffered") == []


def test_a_line_with_its_own_keyword_never_borrows_a_label():
    found = _found(_slide([("Revenue\n800 Paying Users\n97% margins", 1, 2)], title="ARR"))
    assert [(c["claim_type"], c["label_from"]) for c in found] == [("customers", None), ("gross_margin", None)]
    count = _found(_slide([("Revenue\n1.5 million updates", 1, 2)]))
    assert [(c["claim_type"], c["label_from"]) for c in count] == [("revenue", "Revenue")], "a count without a keyword borrows"


# ---------------------------------------------------------------------------
# What approve, edit and reject do (the instruction text above the approval list)
# ---------------------------------------------------------------------------
def test_approved_and_edited_claims_make_the_register_rejected_stay_on_record(api):
    client, db = api
    _upload(client, "audit-1", "03-buffer.pptx", (DECKS / "03-buffer.pptx").read_bytes())
    one, two, three = (dict(c) for c in db[decks.CANDIDATES_COLLECTION].docs[:3])
    url = "/api/audits/audit-1/decks/candidates/{}"
    client.put(url.format(one["id"]), json={"status": "approved"})
    edited = client.put(url.format(two["id"]), json={"value": 900}).json()
    assert edited["status"] == "edited" and edited["parsed"]["value"] == two["value"]
    again = client.put(url.format(two["id"]), json={"status": "approved"}).json()
    assert again["status"] == "edited" and again["parsed"]["value"] == two["value"], "approving an edit keeps it edited"
    client.put(url.format(three["id"]), json={"status": "rejected"})

    register = client.get("/api/audits/audit-1/claims").json()["claims"]
    assert [(c["id"], c["status"]) for c in register] == [(one["id"], "approved"), (two["id"], "edited")]
    assert register[1]["value"] == 900 and register[1]["parsed"]["value"] == two["value"], "original kept beside the edit"
    listed = client.get("/api/audits/audit-1/decks").json()["candidates"]
    assert next(c for c in listed if c["id"] == three["id"])["status"] == "rejected", "rejected stays on record"
    assert client.get("/api/audits/no-such-audit/claims").status_code == 404


def test_a_re_upload_keeps_approved_edited_and_rejected_claims(api):
    client, db = api
    content = (DECKS / "03-buffer.pptx").read_bytes()
    first = _upload(client, "audit-1", "03-buffer.pptx", content).json()
    one, two, three = (dict(c) for c in db[decks.CANDIDATES_COLLECTION].docs[:3])
    url = "/api/audits/audit-1/decks/candidates/{}"
    client.put(url.format(one["id"]), json={"status": "approved"})
    client.put(url.format(two["id"]), json={"value": 900})
    client.put(url.format(three["id"]), json={"status": "rejected"})

    again = _upload(client, "audit-1", "03-buffer.pptx", content).json()
    assert again["kept_reviewed"] == 3 and again["candidates"] == first["candidates"] - 3, "no duplicate of a reviewed claim"
    stored = {c["id"]: c for c in db[decks.CANDIDATES_COLLECTION].docs}
    assert (stored[one["id"]]["status"], stored[two["id"]]["status"], stored[three["id"]]["status"]) == \
        ("approved", "edited", "rejected")
    assert stored[two["id"]]["value"] == 900 and stored[two["id"]]["parsed"]["value"] == two["value"]
    assert len(db[decks.CANDIDATES_COLLECTION].docs) == first["candidates"]
    assert {c["deck_id"] for c in db[decks.CANDIDATES_COLLECTION].docs} == {again["deck_id"]}


@pytest.mark.parametrize("text, families", [
    ("800 Paying Users", ["customers"]),
    ("Avg. Customer Lifetime Value", ["ltv"]),
    ("Implied Customer Life", ["customer_lifetime"]),
    ("Avg. Cost of Paid Acquisition", ["cac"]),
    ("LTV / CAC", ["ltv_cac"]),
    ("LTV/CAC", ["ltv_cac"]),
    ("LTV:CAC", ["ltv_cac"]),
    ("LTV to CAC ratio", ["ltv_cac"]),
])
def test_of_two_overlapping_keywords_the_longer_one_counts(text, families):
    assert [k["family"] for k in claims._keywords(text)] == families


# ---------------------------------------------------------------------------
# One heading, an unknown type and the confidence of a claim (docs/specs/deck-parser.md sections 2 and 6)
# ---------------------------------------------------------------------------
def test_label_from_is_the_one_heading_that_names_the_type_not_every_heading_of_its_box():
    long_box = "Regions and territories we serve today\nCustomers\nStages of the sales process we follow"
    found = _found(_slide([(long_box, 1, 3), ("5,000", 3.5, 3)]))
    five, = _by_value(found, 5000)
    assert (five["claim_type"], five["label_from"]) == ("customers", "Customers"), "a box of several headings gives one"
    titled = _found(_slide([("5,000", 8, 6.5)], title="Our plan for the coming years and what follows\nPaying users\nHow we get there"))
    assert [(c["claim_type"], c["label_from"]) for c in _by_value(titled, 5000)] == [("customers", "Paying users")]
    wrapped = _found(_slide([("Our plan for the coming years and what follows\nMARKET\nSIZE\nHow we get there", 1, 3), ("$2.5B", 3.5, 3)]))
    assert _by_value(wrapped, 2500000000)[0]["label_from"] == "MARKET SIZE", "a keyword that wraps keeps both lines"


def test_a_box_short_enough_to_be_one_label_stays_whole_even_when_it_wraps():
    found = _found(_slide([("Spend as\n% of revenue", 1, 3), ("12%", 3.5, 3)]))
    assert [(c["claim_type"], c["label_from"]) for c in _by_value(found, 12)] == [("revenue", "Spend as % of revenue")]


def test_a_figure_no_heading_names_a_type_for_is_unknown_not_guessed_as_product():
    found = _found(_slide([("120 in Q3 2025", 1, 3)]))
    one, = _by_value(found, 120)
    assert (one["claim_type"], one["type_from"], one["label_from"]) == ("unknown", None, None)
    dated = _found(_slide([("Q3 2025", 1, 3)]))
    assert [(c["claim_type"], c["target_date"]) for c in dated] == [("unknown", "2025-Q3")]
    named = _found(_slide([("Launch the API Q3 2025", 1, 3)]))
    assert [(c["claim_type"], c["type_from"]) for c in named] == [("product", "line")]
    under = _found(_slide([("Roadmap", 1, 2), ("Q3 2025", 1, 3)]))
    assert [(c["claim_type"], c["type_from"]) for c in under] == [("product", "heading")]


def test_a_type_comes_from_the_nearest_heading_never_from_one_further_away():
    """2026-10-08: zero2hero p11's "AR/VR ($52 B)" and "E-Learning Market ($374B)" sit under "Market" (bare "market" names
    no type); the keyword "Acquisition" of a box further up the page used to type them Sales. They are Unknown."""
    file = "05-zero2hero.pdf"
    page = [c for c in claims.detect_candidates(parser.parse_deck((DECKS / file).read_bytes(), file)["blocks"], file)
            if _page(c) == 11]
    big = [c for c in page if c["value"] in (52e9, 374e9)]
    assert sorted(c["value"] for c in big) == [52e9, 374e9]
    assert [(c["claim_type"], c["type_from"], c["label_from"]) for c in big] == [("unknown", None, None)] * 2


def test_the_nearest_heading_names_the_type_or_the_figure_has_none():
    far = _found(_slide([("Content Acquisition", 1, 1.0), ("Market", 1, 2.6), ("($52 B)", 1, 3.2)]))
    assert [(c["claim_type"], c["type_from"], c["label_from"]) for c in _by_value(far, 52e9)] == [("unknown", None, None)], \
        "'Market' is nearer than the keyword and names no type"
    near = _found(_slide([("Content Acquisition", 1, 2.6), ("($52 B)", 1, 3.2)]))
    assert [(c["claim_type"], c["label_from"]) for c in _by_value(near, 52e9)] == [("sales", "Content Acquisition")]
    # A line of figures is not a heading: it is passed over to the label beyond it.
    passed = _found(_slide([("Spend as % of revenue", 1, 1.5), ("$5M $6M", 1, 2.2), ("12%", 1, 2.9)]))
    assert [c["claim_type"] for c in _by_value(passed, 12)] == ["revenue"]
    # Nothing names a type anywhere and there is no date word: not a candidate, as before.
    assert _found(_slide([("Market", 1, 2.6), ("($52 B)", 1, 3.2)])) == []


def test_the_ten_test_decks_keep_225_candidates_and_the_nearest_heading_leaves_31_unknown_more():
    counts, unknown = 0, 0
    for file in sorted(p.name for p in DECKS.iterdir() if not p.name.startswith(".")):
        found = claims.detect_candidates(parser.parse_deck((DECKS / file).read_bytes(), file)["blocks"], file)
        counts += len(found)
        unknown += sum(c["claim_type"] == "unknown" for c in found)
    assert (counts, unknown) == (225, 37), "8 before the rule, 31 more after it, 2 of them (zero2hero p22) Use of funds since; no row appears or disappears"


def test_the_same_figure_on_one_page_is_one_row_and_never_a_deck_inconsistency():
    """2026-10-08: the same type, value, currency and period is one claim whatever noun it is counted in; a different
    value for it is the only deck inconsistency."""
    found = _found(_slide([("5 customers in 2024", 1, 1), ("5 customers in 2024", 1, 3)]))
    one, = found
    assert (one["claim_type"], one["value"], one["inconsistent_dates"]) == ("customers", 5, [])
    same = claims.detect_candidates([{"slide": 1, "kind": "text", "text": "Revenue 2024: £150K", "box": 1},
                                     {"slide": 4, "kind": "text", "text": "Revenue 2024: £ 150,000", "box": 2}], "d.pptx")
    one, = same
    assert (one["value"], len(one["sources"]), one["inconsistent_dates"]) == (150000, 2, [])
    differ = _found(_slide([("5 customers in 2024", 1, 1), ("6 customers in 2024", 1, 3)]))
    assert sorted((c["value"], c["inconsistent_dates"]) for c in differ) == [(5, ["2024"]), (6, ["2024"])]
    rate = _found(_slide([("Churn 5% in 2024", 1, 1), ("Churn 5x in 2024", 1, 3)]))
    assert len(rate) == 2, "a rate and a multiple are different claims"


def test_a_direction_with_no_figure_is_stored_as_a_direction_and_is_at_most_medium():
    found = _found(_slide([("Positive EBITDA", 1, 1), ("Q2 2024", 1, 1.5)]))
    one, = found
    assert (one["claim_type"], one["value"], one["claim_direction"], one["target_date"]) == ("ebitda", None, "positive", "2024-Q2")
    conf = claims.confidence(one, [one])
    assert conf["level"] == "Medium" and conf["failed"] == ["no figure"], "every other check passes, and it is still not High"
    negative, = _found(_slide([("EBITDA negative until Q4 2024", 1, 1)]))
    assert negative["claim_direction"] == "negative"
    plain, = _found(_slide([("Break-even EBITDA", 1, 1), ("Q2 2024", 1, 1.5)]))
    assert plain["claim_direction"] is None and claims.confidence(plain, [plain])["level"] == "High"
    figure, = _found(_slide([("EBITDA of $2M in 2025", 1, 1)]))
    assert figure["claim_direction"] is None and figure["value"] == 2000000
    zero2hero = "05-zero2hero.pdf"
    page = [c for c in claims.detect_candidates(parser.parse_deck((DECKS / zero2hero).read_bytes(), zero2hero)["blocks"], zero2hero)
            if _page(c) == 19 and c["claim_type"] == "ebitda"]
    assert [(c["claim_direction"], c["target_date"]) for c in page] == [("positive", "2024-Q2")]


def test_an_edit_that_types_a_figure_ends_the_direction_and_the_list_is_grouped_then_by_page(api):
    client, db = api
    _upload(client, "audit-1", "05-zero2hero.pdf", (DECKS / "05-zero2hero.pdf").read_bytes())
    listed = client.get("/api/audits/audit-1/decks").json()["candidates"]
    groups = [c["group"] for c in listed]
    assert groups == sorted(groups) and set(groups) >= {1, 2, 3, 4, 6}, "revenue, P&L, customers and sales, hiring and roadmap, Unknown"
    for group in set(groups):
        pages = [_page(c) for c in listed if c["group"] == group]
        assert pages == sorted(pages), f"ascending by page within group {group}"
    assert {c["claim_type"] for c in listed if c["group"] == 1} <= {"revenue", "revenue_growth"}
    assert {c["claim_type"] for c in listed if c["group"] == 6} == {"unknown"}
    direction = next(c for c in listed if c.get("claim_direction"))
    assert direction["group"] == 2 and direction["claim_type"] == "ebitda"
    edited = client.put(f"/api/audits/audit-1/decks/candidates/{direction['id']}", json={"value": 3.0, "unit": "%"})
    assert edited.status_code == 200, edited.text
    assert edited.json()["claim_direction"] is None and edited.json()["value"] == 3.0


def test_every_claim_type_has_its_group():
    groups = {t: claims.type_group(t) for t in (*claims.CLAIM_TYPES, "usage", "use_of_funds", "other", "unknown")}
    assert {g: sorted(t for t, x in groups.items() if x == g) for g in range(1, 8)} == {
        1: ["revenue", "revenue_growth"],
        2: ["burn", "cash", "costs", "ebitda", "gross_margin", "gross_profit", "months_to_profitability", "net_profit", "runway"],
        3: ["cac", "customer_lifetime", "customers", "growth", "ltv", "ltv_cac", "retention", "sales", "trials_per_day", "usage",
            "user_growth", "users"],
        4: ["people", "product"], 5: ["market"], 6: ["unknown"], 7: ["other", "use_of_funds"]}
    assert claims.COLLAPSED_GROUPS == (6, 7)


def test_the_parser_and_the_model_reading_the_same_value_are_one_claim_only_when_type_value_currency_and_period_agree():
    base = {"claim_type": "gross_profit", "value": 150000, "value_high": None, "unit": None, "currency": "GBP", "target_date": "2023",
            "period_start": "2023-01-01", "period_end": "2023-12-31", "sources": [{"slide": 19}]}
    assert claims.same_claim(base, {**base, "sources": [{"slide": 19, "kind": "structure"}]})
    for change in ({"value": 50000}, {"currency": "EUR"}, {"claim_type": "revenue"}, {"target_date": "2024", "period_end": "2024-12-31"},
                   {"value_high": 160000}, {"unit": "%"}, {"value": None}):
        assert not claims.same_claim(base, {**base, **change}), change
    milestone = {**base, "claim_type": "product", "value": None, "currency": None}
    assert not claims.same_claim(milestone, dict(milestone)), "two lines with the same date are not one claim"


def test_claims_in_another_currency_carry_the_saved_rate_or_none(api):
    client, db = api
    db["audits"].docs[0].update(reporting_currency="EUR", as_of_month="2026-06")
    _upload(client, "audit-1", "05-zero2hero.pdf", (DECKS / "05-zero2hero.pdf").read_bytes())
    gbp = [c for c in client.get("/api/audits/audit-1/decks").json()["candidates"] if c["currency"] == "GBP"]
    assert gbp and all(c["fx"] == {"rate": None, "date": "2026-06-30", "currency": "EUR"} for c in gbp), "no rate saved"
    db["datasets"].docs.append({"audit_id": "audit-1", "dtype": "revenue", "fx": {"gbp": 1.14}})
    gbp = [c for c in client.get("/api/audits/audit-1/decks").json()["candidates"] if c["currency"] == "GBP"]
    assert all(c["fx"] == {"rate": 1.14, "date": "2026-06-30", "currency": "EUR"} for c in gbp)
    assert all(c["fx"] is None for c in client.get("/api/audits/audit-1/decks").json()["candidates"]
               if c["currency"] in (None, "EUR"))


def test_an_unknown_date_is_the_product_claim_with_the_same_date_when_there_is_one():
    blocks = [{"slide": 1, "kind": "text", "text": "Launch the API Q3 2025"}, {"slide": 2, "kind": "text", "text": "Q3 2025"}]
    found = claims.detect_candidates(blocks, "d.pptx")
    assert [(c["claim_type"], c["target_date"], c["type_from"], [s["slide"] for s in c["sources"]]) for c in found] == \
        [("product", "2025-Q3", "line", [1, 2])], "no row appears or disappears: both were product and merged"
    alone = claims.detect_candidates(blocks[1:], "d.pptx")
    assert [(c["claim_type"], c["target_date"]) for c in alone] == [("unknown", "2025-Q3")]


def _claim(**fields):
    return {"claim_type": "revenue", "value": 100, "value_high": None, "unit": None, "currency": "EUR",
            "target_date": "2025", "type_from": "line", "sources": [{"file": "d.pptx", "slide": 2, "kind": "text"}], **fields}


@pytest.mark.parametrize("fields, level, failed", [
    ({"sources": [{"slide": 2, "kind": "text"}, {"slide": 9, "kind": "text"}]}, "High", []),
    ({}, "Medium", ["not corroborated"]),
    ({"target_date": None, "sources": [{"slide": 2}, {"slide": 5}]}, "Medium", ["no date"]),
    ({"currency": None, "sources": [{"slide": 2}, {"slide": 5}]}, "Medium", ["no unit"]),
    ({"claim_type": "unknown", "type_from": None, "sources": [{"slide": 2}, {"slide": 5}]}, "Medium", ["no heading"]),
    ({"target_date": None, "claim_type": "unknown", "type_from": None, "sources": [{"slide": 2}, {"slide": 5}]},
     "Low", ["no date", "no heading"]),
    ({"target_date": None, "currency": None}, "Low", ["no date", "no unit", "not corroborated"]),
    ({"value": None, "target_date": "2025-Q3", "claim_type": "product"}, "High", []),
    ({"value": None, "target_date": None, "claim_type": "unknown", "type_from": None}, "Low", ["no date", "no heading"]),
])
def test_confidence_counts_the_failed_checks_never_a_model_score(fields, level, failed):
    got = claims.confidence(_claim(**fields), [])
    assert (got["level"], got["failed"]) == (level, failed)
    assert got["text"] == level + (" – " + ", ".join(failed) if failed else "")


def test_a_value_is_corroborated_by_another_candidate_of_the_deck_at_another_place_and_a_table_row_is_one_place():
    one, same, other_value = _claim(), _claim(sources=[{"slide": 7, "kind": "text"}]), _claim(value=5, sources=[{"slide": 7}])
    assert claims.confidence(one, [one, same])["failed"] == []
    assert claims.confidence(one, [one, other_value])["failed"] == ["not corroborated"]
    assert claims.confidence(one, [one, _claim(claim_type="costs", sources=[{"slide": 7}])])["failed"] == \
        ["not corroborated"], "the same number under another type is a coincidence"
    row = _claim(value=None, target_date=None, by_period=[
        {"value": 100, "value_high": None, "target_date": "2024", "source": {"slide": 3, "kind": "table", "table": 1, "row": 2, "col": c}}
        for c in (2, 3)], sources=[{"slide": 3, "kind": "table", "table": 1, "row": 2, "col": c} for c in (2, 3)])
    assert claims.confidence(row, [row])["failed"] == ["not corroborated"], "two cells of one row are not two places"


def test_a_claim_stored_before_the_confidence_has_no_type_from_and_counts_as_named_unless_its_type_is_unknown():
    old = _claim()
    del old["type_from"]
    assert "no heading" not in claims.confidence(old, [])["failed"]
    assert "no heading" in claims.confidence({**old, "claim_type": "unknown"}, [])["failed"]


def test_an_edited_claim_names_its_type_whatever_the_parser_found(api):
    edited = _claim(claim_type="revenue", type_from=None, parsed={"claim_type": "unknown"})
    assert "no heading" not in claims.confidence(edited, [])["failed"]


def test_an_unknown_claim_is_approved_only_once_its_type_is_chosen(api):
    client, _ = api
    _upload(client, "audit-1", "02-moz.pdf", (DECKS / "02-moz.pdf").read_bytes())
    unknown = next(c for c in client.get("/api/audits/audit-1/decks").json()["candidates"] if c["claim_type"] == "unknown")
    url = f"/api/audits/audit-1/decks/candidates/{unknown['id']}"
    assert client.put(url, json={"status": "approved"}).status_code == 400
    assert client.put(url, json={"target_date": "2011-07"}).status_code == 400, "an edit approves, so it needs the type too"
    done = client.put(url, json={"claim_type": "product"})
    assert done.status_code == 200 and done.json()["status"] == "edited"
    assert client.put(url, json={"claim_type": "market"}).status_code == 200


# ---------------------------------------------------------------------------
# Several decks per audit: removal and order
# ---------------------------------------------------------------------------
def _page(c):
    return min(s.get("slide", s.get("page")) for s in c["sources"])


def test_decks_list_newest_first_and_claims_ascend_by_page_across_decks_then_reading_order(api):
    client, db = api
    old = _upload(client, "audit-1", "03-buffer.pptx", (DECKS / "03-buffer.pptx").read_bytes()).json()
    new = _upload(client, "audit-1", "02-moz.pdf", (DECKS / "02-moz.pdf").read_bytes()).json()
    old_claims = [c for c in db[decks.CANDIDATES_COLLECTION].docs if c["deck_id"] == old["deck_id"]]
    late = max(old_claims, key=_page)
    client.put(f"/api/audits/audit-1/decks/candidates/{late['id']}", json={"status": "approved"})
    # A kept or merged claim can carry an early "order"; the slide still decides its place.
    last = max((c for c in old_claims if c["id"] != late["id"]), key=_page)
    last["order"] = -1

    listed = client.get("/api/audits/audit-1/decks").json()
    assert [d["deck_id"] for d in listed["decks"]] == [new["deck_id"], old["deck_id"]], "most recent deck first"
    rows = listed["candidates"]
    assert {c["deck_id"] for c in rows} == {new["deck_id"], old["deck_id"]}
    assert [c["group"] for c in rows] == sorted(c["group"] for c in rows), "by type group first"
    for group in {c["group"] for c in rows}:
        members = [c for c in rows if c["group"] == group]
        pages = [_page(c) for c in members]
        assert pages == sorted(pages), "ascending by slide or page across all decks, not grouped by deck"
        for deck in (old, new):
            mine = [_page(c) for c in members if c["deck_id"] == deck["deck_id"]]
            assert mine == sorted(mine), "and within each deck"
    assert len({c["deck_id"] for c in rows[:len(rows) // 2]}) == 2, "the decks interleave"
    assert [c["id"] for c in rows if _page(c) == _page(late)].count(late["id"]) == 1
    same_group = [c for c in rows if c["group"] == next(x["group"] for x in rows if x["id"] == late["id"])]
    assert [_page(c) for c in same_group].index(_page(late)) <= [c["id"] for c in same_group].index(late["id"]), \
        "a reviewed claim keeps the place its page gives it: it does not move below the others"
    page_two = [c for c in rows if c["deck_id"] == new["deck_id"] and _page(c) == 2 and c.get("reading")
                and c["group"] == max((x["group"] for x in rows if x["deck_id"] == new["deck_id"] and _page(x) == 2), key=lambda g: sum(
                    1 for x in rows if x["deck_id"] == new["deck_id"] and _page(x) == 2 and x["group"] == g))]
    assert len(page_two) > 3
    keys = [(round(c["reading"][0] / 0.02), c["reading"][1]) for c in page_two]
    assert keys == sorted(keys), "within a page: top to bottom, then left to right"


def test_a_claim_of_a_deck_shows_its_confidence_and_the_checks_it_failed(api):
    client, _ = api
    _upload(client, "audit-1", "02-moz.pdf", (DECKS / "02-moz.pdf").read_bytes())
    rows = client.get("/api/audits/audit-1/decks").json()["candidates"]
    seen = {c["confidence"]["level"] for c in rows}
    assert seen == {"High", "Medium", "Low"}, "the 46 moz claims use all three levels"
    for c in rows:
        conf = c["confidence"]
        assert conf["text"] == conf["level"] + (" – " + ", ".join(conf["failed"]) if conf["failed"] else "")
        assert len(conf["failed"]) == {"High": 0, "Medium": 1, "Low": len(conf["failed"])}[conf["level"]]
        assert conf["level"] != "Low" or len(conf["failed"]) >= 2


def test_remove_deck_deletes_its_text_and_all_its_claims_including_reviewed(api):
    client, db = api
    keep = _upload(client, "audit-1", "09-genesisai-2024.pdf", (DECKS / "09-genesisai-2024.pdf").read_bytes()).json()
    gone = _upload(client, "audit-1", "03-buffer.pptx", (DECKS / "03-buffer.pptx").read_bytes()).json()
    other = _upload(client, "audit-2", "03-buffer.pptx", (DECKS / "03-buffer.pptx").read_bytes()).json()
    mine = [c for c in db[decks.CANDIDATES_COLLECTION].docs if c["deck_id"] == gone["deck_id"]]
    url = "/api/audits/audit-1/decks/candidates/{}"
    client.put(url.format(mine[0]["id"]), json={"status": "approved"})
    client.put(url.format(mine[1]["id"]), json={"value": 1})
    client.put(url.format(mine[2]["id"]), json={"status": "rejected"})

    r = client.delete(f"/api/audits/audit-1/decks/{gone['deck_id']}")
    assert r.status_code == 200 and r.json() == {"removed": gone["deck_id"], "deck_text": 1,
                                                 "deck_candidates": gone["candidates"]}
    assert {d["deck_id"] for d in db[decks.TEXT_COLLECTION].docs} == {keep["deck_id"], other["deck_id"]}
    assert {c["deck_id"] for c in db[decks.CANDIDATES_COLLECTION].docs} == {keep["deck_id"], other["deck_id"]}
    assert all(c["deck_id"] != gone["deck_id"] for c in client.get("/api/audits/audit-1/claims").json()["claims"])
    assert client.delete(f"/api/audits/audit-1/decks/{gone['deck_id']}").status_code == 404
    assert client.delete(f"/api/audits/audit-2/decks/{keep['deck_id']}").status_code == 404, "a deck of another audit"


# ---------------------------------------------------------------------------
# Plan claims only (spec section 2): whole numbers, axis ticks, labelled market, drops
# ---------------------------------------------------------------------------
def test_pdf_axis_labels_are_read_as_whole_numbers():
    """zero2hero page 4: the axis labels 100 ... 800 came out as single digits when rows were
    sorted by position alone."""
    file = "05-zero2hero.pdf"
    texts = [b["text"] for b in parser.parse_deck((DECKS / file).read_bytes(), file)["blocks"] if b["page"] == 4]
    assert {"100", "200", "300", "400", "500", "600", "700", "800"} <= set(texts)
    assert not {"5", "6", "7", "8"} & set(texts)


def test_evenly_spaced_numbers_in_a_row_or_column_are_axis_ticks():
    row = _found(_slide([("Gross churn", 0.5, 1)] + [(f"{v}%", 1 + i, 2) for i, v in enumerate((0, 10, 20, 30))]))
    assert row == []
    column = _found(_slide([("ARR", 0.5, 1), ("$100K\n$200K\n$300K", 1, 1.5)]))
    assert column == []
    uneven = _found(_slide([("Gross churn", 0.5, 2)] + [(f"{v}%", 1 + i, 2) for i, v in enumerate((12, 18, 31))]))
    assert sorted(c["value"] for c in uneven) == [12, 18, 31], "uneven figures are data, not ticks"
    assert claims._evenly_spaced([0, 25, 50, 75, 100]) and not claims._evenly_spaced([49284, 181193, 278085])


def test_table_row_numbers_are_ticks():
    rows = [["#", "Metric", "Value"], ["1", "Revenue", "$1.2M"], ["2", "Revenue next year", "$2.0M"], ["3", "ARR", "$3.1M"]]
    found = _found(_slide([], table=(rows, 1, 2)))
    assert sorted(c["value"] for c in found) == [1200000, 2000000, 3100000]


@pytest.mark.parametrize("text, kept", [
    ("A TAM of €200M", True), ("Serviceable addressable market of $1B", True), ("Market size $2B", True),
    ("SAM $400M", True), ("This covers 50% of entire US market", False), ("Overall Market: $4.2B annually", False),
])
def test_market_counts_only_when_labelled(text, kept):
    found = [c for c in _line(text) if c["claim_type"] == "market"]
    assert bool(found) == kept


@pytest.mark.parametrize("title", ["The Problem", "Why Now?", "Macroeconomic Trends", "Social Media Landscape",
                                   "TEA Token Allocation"])
def test_figures_on_background_slides_are_dropped(title):
    assert _found(_slide([("ARR $1.2M, 40% growth", 1, 2)], title=title)) == []


def test_figures_on_a_page_that_cites_outside_research_are_dropped():
    assert _found(_slide([("Market size $230B", 1, 2), ("Source: Transparency Market Research", 1, 6)])) == []
    footnotes = _found(_slide([("Market size $374.3 billion", 1, 2), ("1. https://example.org/report", 1, 6),
                               ("2. Research at MIT", 1, 6.5)]))
    assert footnotes == []
    not_footnotes = _found(_slide([("3.9x more ARR per user", 1, 2), ("4.9x more MRR per user", 1, 3)]))
    assert sorted(c["value"] for c in not_footnotes) == [3.9, 4.9], "a decimal is not a footnote number"


@pytest.mark.parametrize("text", [
    "We raised $2M in our seed round", "Raising: $20-$25 Million", "Seed to Series A $3.1m ARR",
    "Team and Community: 5% vesting per month", "100 million total supply",
    "As a member of the exec team, scaled Intercom from $1M ARR to $180M+ ARR",
    "Co-Founder, took the idea to revenue in 7 weeks", "Growing to a $3.5B industry by 2010",
])
def test_lines_about_funds_tokens_careers_or_the_industry_are_dropped(text):
    assert _found(_slide([(text, 1, 2)])) == []


def test_plan_lines_that_mention_funds_or_a_team_are_kept():
    found = _found(_slide([("Funds will be used to grow the team to 20 by end of 2020", 1, 2),
                           ("To fund an initial team for 24 months.", 1, 3)]))
    assert sorted(c["value"] for c in found) == [20, 24]


# ---------------------------------------------------------------------------
# Tables by period, periods, finance types and deck inconsistencies (browser test, zero2hero p19)
# ---------------------------------------------------------------------------
def test_zero2hero_page_19_reads_as_one_candidate_per_row_and_flags_the_panel():
    """The deck's key financial slide: a projections table under "Y/E 22" ... "Y/E 26" and a panel
    headed "23 Y/E" that gives Gross Profit £150K where the table gives £ 50,000."""
    file = "05-zero2hero.pdf"
    deck = parser.parse_deck((DECKS / file).read_bytes(), file)
    page = [c for c in claims.detect_candidates(deck["blocks"], file) if _page(c) == 19]
    rows = {c["snippet"].split(" | ")[0]: c for c in page if c.get("by_period")}
    assert {k: (c["claim_type"], [i["value"] for i in c["by_period"]]) for k, c in rows.items()} == {
        "Registered Users": ("users", [200, 5000, 20000, 30000, 50000]),
        "Registered Institutions": ("customers", [1, 2, 25, 50]),
        "Revenue": ("revenue", [130550, 150000, 250000, 1000000, 2500000]),
        "Direct Costs": ("costs", [83403, 100000, 150000, 500000, 1000000]),
        "Gross Profit": ("gross_profit", [42638, 50000, 100000, 500000, 1500000])}
    assert [(i["period"], i["target_date"]) for i in rows["Revenue"]["by_period"]] == \
        [("Y/E 22", "2022"), ("Y/E 23", "2023"), ("Y/E 24", "2024"), ("Y/E 25", "2025"), ("Y/E 26", "2026")]
    panel = sorted((c["claim_type"], c["value"], c["target_date"], c["inconsistent_dates"]) for c in page
                   if not c.get("by_period"))
    assert panel == [("ebitda", None, "2024-Q2", []), ("gross_profit", 150000, "2023", ["2023"]),
                     ("users", 5000, "2023", [])], "23 Y/E is a period, never a value"
    assert rows["Gross Profit"]["inconsistent_dates"] == ["2023"]
    # The page-17 bars are labelled Turnover (see the page-17 test): a turnover figure is not compared with Revenue.
    assert rows["Revenue"]["inconsistent_dates"] == [], "page 17 is Turnover: never compared with this Revenue row"
    assert all(not c["inconsistent_dates"] for k, c in rows.items() if k not in ("Gross Profit", "Revenue"))


@pytest.mark.parametrize("text, date", [
    ("Y/E 22", "2022"), ("23 Y/E", "2023"), ("YE 2021", "2021"), ("FY23", "2023"), ("FY 2024", "2024"),
    ("2023E", "2023"), ("2025F", "2025"), ("H1 24", "2024-H1"), ("1H2025", "2025-H1"), ("H2'24", "2024-H2"),
    ("Q3 25", "2025-Q3"), ("3Q25", "2025-Q3"),
])
def test_periods_are_read_as_dates_never_as_values(text, date):
    assert [d["date"] for d in claims.find_dates(text)] == [date]
    assert claims.find_numbers(text, claims.find_dates(text)) == []


@pytest.mark.parametrize("text", ["1/1/18 5/1/18", "9/1/18 12/1/18", "3/1/15", "12/31/2018", "Paid on 5/1/2018"])
def test_a_date_written_with_slashes_is_a_date_with_no_period(text):
    """Issue #53, option a (decision of 2026-10-07): the order of day and month is not stated ("5/1/18" is 1 May or
    5 January), so none of its parts is a figure, it is no period cell and nothing is dated from it, the year of
    "12/31/2018" included."""
    assert claims.figures(text) == [] and claims.find_dates(text) == [], text
    assert verify.figures(text) == [] and claims.period_cell(text) is None, text


@pytest.mark.parametrize("text, date", [("Q1-Q2 2023", "2023-H1"), ("Q3-Q4 23", "2023-H2"), ("2023 Q1-Q2", "2023-H1"),
                                        ("2023 Q3-Q4", "2023-H2")])
def test_q1_q2_and_q3_q4_are_the_halves_with_the_year_before_or_after(text, date):
    """Issue #55 (decision of 2026-10-06): tea p11's "Mainnet starts" borrows "2023 Q1-Q2" from its box."""
    assert [d["date"] for d in claims.find_dates(text)] == [date] and claims.figures(text) == []


@pytest.mark.parametrize("text", ["Q2-Q3 2023", "Q1-Q3 23", "2023 Q2-Q3", "Q1-Q4 2023", "Launch Q2-Q3 2023"])
def test_a_quarter_range_that_is_no_half_has_no_period(text):
    """Decision of 2026-10-07 on issue #55: "Q2-Q3 2023" read 2023-Q2. A quarter range that is no half dates nothing,
    and its year is no figure."""
    assert claims.find_dates(text) == [] and claims.figures(text) == [], text
    assert verify.figures(text) == [] and claims.period_cell(text) is None, text


def test_a_pair_with_no_year_or_a_part_over_31_is_no_slash_date():
    assert [n["value"] for n in claims.figures("Building a predictable sales organization (1/2)")] == [1, 2]
    assert [n["value"] for n in claims.figures("45/3/20")] == [45, 3, 20]


def test_front_b_p14_and_p16_slash_dates_give_no_candidate_and_p16_no_panel():
    """Issue #53: front-b p14's axis dates ("3/1/15" to "11/1/17") and p16's ("1/1/18 5/1/18", "9/1/18 12/1/18") gave
    11 candidates from their parts; p16's panel held no other figure, so it is no panel and its 12 items go."""
    file = "01-front-b.pptx"
    deck = parser.parse_deck((DECKS / file).read_bytes(), file)
    found = claims.detect_candidates(deck["blocks"], file)
    assert not [c for c in found if "/1/1" in c["snippet"]], "no candidate from a slash date's parts"
    assert not [s for s in deck["structures"] if s.get("slide") == 16], "front-b p16 holds no structure"


def test_a_column_header_period_beats_a_date_nearby():
    rows = [["", "Y/E 22", "Y/E 23"], ["Revenue", "£ 130,550", "£ 150,000"]]
    row, = [c for c in _found(_slide([("Q2 2024", 1, 1.6)], table=(rows, 1, 2))) if c.get("by_period")]
    assert [i["target_date"] for i in row["by_period"]] == ["2022", "2023"]


def test_a_period_header_row_is_never_ticks_or_a_claim_and_a_row_is_never_ticks():
    rows = [["", "Q1 24", "Q2 24", "Q3 24"], ["Revenue", "$1M", "$2M", "$3M"]]
    found = _found(_slide([], table=(rows, 1, 2)))
    assert [(c["claim_type"], [(i["value"], i["target_date"]) for i in c["by_period"]]) for c in found] == \
        [("revenue", [(1000000, "2024-Q1"), (2000000, "2024-Q2"), (3000000, "2024-Q3")])]


def test_a_lone_figure_or_figures_of_different_types_in_a_row_stay_separate():
    lone = _found(_slide([], table=([["", "2024"], ["Revenue", "$3M"]], 1, 2)))
    assert [(c["value"], c["target_date"], c.get("by_period")) for c in lone] == [(3000000, "2024", None)]
    mixed = _found(_slide([], table=([["", "ARR", "Customers"], ["2024", "$3M", "40"]], 1, 2)))
    assert sorted((c["claim_type"], c["value"]) for c in mixed) == [("customers", 40), ("revenue", 3000000)]


def test_a_period_at_the_top_of_a_text_box_dates_every_figure_in_it():
    found = _found(_slide([("23 Y/E\nGross Profit £150K\n5K Users", 1, 2)]))
    assert sorted((c["claim_type"], c["value"], c["target_date"]) for c in found) == \
        [("gross_profit", 150000, "2023"), ("users", 5000, "2023")]
    found = _found(_slide([("FY24\nARR $3M\nBeta launch in Q1 2023", 1, 2)]))
    arr, = _by_value(found, 3000000)
    assert (arr["target_date"], arr["date_from"]) == ("2024", "FY24"), "the box period beats a quarter lower down"


@pytest.mark.parametrize("text, family", [
    ("Gross Profit £150K", "gross_profit"), ("Gross margin £1.2M", "gross_profit"), ("Gross margin 82%", "gross_margin"),
    ("Direct costs £83,403", "costs"), ("Opex of $2M", "costs"), ("Costs $500K", "costs"),
    ("EBITDA of $1.5M", "ebitda"), ("Profitability of 12% by 2026", "ebitda"), ("40 institutions", "customers"),
    ("Net income $1.5M", "net_profit"), ("Net profit 12%", "net_profit"), ("Net loss of $2M", "net_profit"),
])
def test_gross_profit_costs_ebitda_and_institutions(text, family):
    assert [c["claim_type"] for c in _line(text)] == [family]


def test_a_net_loss_is_a_negative_net_profit():
    assert [(c["value"], c["value_high"]) for c in _line("Net loss of $2M in 2023")] == [(-2000000, None)]
    assert [(c["value"], c["value_high"]) for c in _line("Net losses of $1-2M")] == [(-2000000, -1000000)]
    assert [(c["value"], c["value_high"]) for c in _line("Net income $1.5M")] == [(1500000, None)]


def test_a_break_even_milestone_takes_its_date():
    assert [(c["claim_type"], c["target_date"]) for c in _line("Break-even by Q3 2025")] == [("ebitda", "2025-Q3")]
    assert [(c["claim_type"], c["target_date"]) for c in _line("Profitability in H2 2026")] == [("ebitda", "2026-H2")]
    found = _found(_slide([("Positive EBITDA\nQ2 2024", 1, 2)]))
    assert [(c["claim_type"], c["value"], c["target_date"]) for c in found] == [("ebitda", None, "2024-Q2")]
    titled = _found(_slide([("Break-even", 1, 2)], title="Targets Q3 2025"))
    assert ("ebitda", "2025-Q3") in [(c["claim_type"], c["target_date"]) for c in titled]


# ---------------------------------------------------------------------------
# Claim types of issue #45 (spec section 2, decisions of 2026-10-06)
# ---------------------------------------------------------------------------
NEW_TYPES = ("cash", "burn", "runway", "ltv", "cac", "customer_lifetime", "ltv_cac", "trials_per_day",
             "months_to_profitability")


def test_the_new_claim_types_follow_the_old_ones():
    assert claims.CLAIM_TYPES == ("revenue", "revenue_growth", "growth", "retention", "sales", "customers", "users",
                                  "user_growth", "gross_margin", "gross_profit", "costs", "ebitda", "net_profit",
                                  "people", "product", "market") + NEW_TYPES


@pytest.mark.parametrize("text, family, value, unit, currency", [
    ("Cash on hand $7m", "cash", 7000000, None, "USD"),
    ("Net burn of $200K per month", "burn", 200000, None, "USD"),
    ("Burn rate €150K", "burn", 150000, None, "EUR"),
    ("Runway: 18 months", "runway", 18, "months", None),
    ("LTV $240", "ltv", 240, None, "USD"),
    ("Customer lifetime value of $900", "ltv", 900, None, "USD"),
    ("CAC $100", "cac", 100, None, "USD"),
    ("Cost per customer acquisition $120", "cac", 120, None, "USD"),
    ("Acquisition cost of $80", "cac", 80, None, "USD"),
    ("LTV / CAC 2.5x", "ltv_cac", 2.5, "x", None),
    ("Customer lifetime of 24 months", "customer_lifetime", 24, "months", None),
    ("Implied Customer Life ~9 Months", "customer_lifetime", 9, "months", None),
    ("# of New Free Trials / Day ~100", "trials_per_day", 100, None, None),
    ("200 trials per day", "trials_per_day", 200, None, None),
    ("Profitable in 10 months", "months_to_profitability", 10, "months", None),
])
def test_the_new_claim_types_and_their_keywords(text, family, value, unit, currency):
    found, = _line(text)
    assert (found["claim_type"], found["value"], found["currency"]) == (family, value, currency)
    assert unit is None or found["unit"] == unit


@pytest.mark.parametrize("text, family", [
    ("Payback in 12 months", "sales"), ("Customer acquisition up 30%", "sales"), ("Win rate 25%", "sales"),
    ("Churn of 5%", "retention"), ("Profitability of 12% by 2026", "ebitda"), ("Break-even in 18 months", "ebitda"),
])
def test_the_keywords_left_behind_keep_their_type(text, family):
    assert _line(text)[0]["claim_type"] == family


def test_cash_flow_is_no_cash_keyword():
    assert claims._keywords("Cash flow of $2M") == [] and claims._keywords("Cash-flow positive") == []
    assert [k["family"] for k in claims._keywords("Cash on hand")] == ["cash"]


def test_profitable_is_months_to_profitability_only_with_a_figure_in_months():
    """Spec section 2 (decision of 2026-10-06): "Profitable in 10 months" is months to profitability; a dated
    "profitable" line stays an EBITDA milestone, and a bare date under it takes the EBITDA type."""
    assert [(c["claim_type"], c["value"], c["unit"]) for c in _line("Profitable in 10 months")] == \
        [("months_to_profitability", 10, "months")]
    assert [(c["claim_type"], c["target_date"]) for c in _line("Profitable by Q3 2025")] == [("ebitda", "2025-Q3")]
    found = _found(_slide([("Profitable\nQ2 2024", 1, 2)]))
    assert [(c["claim_type"], c["value"], c["target_date"]) for c in found] == [("ebitda", None, "2024-Q2")]
    found = _found(_slide([("Profitable", 1, 1), ("in 18 months", 1, 1.6)]))
    assert [(c["claim_type"], c["value"]) for c in found] == [("months_to_profitability", 18)], "a borrowed label too"


def test_the_test_decks_retype_eleven_candidates_and_list_no_other_change():
    """Spec section 2: on the 10 test decks the keyword moves and the new keywords retype exactly these candidates."""
    retyped = {
        ("01-front-b.pptx", 12, 2.5): "ltv_cac", ("01-front-b.pptx", 12, 2.6): "ltv_cac",
        ("01-front-b.pptx", 12, 4.4): "ltv_cac", ("01-front-b.pptx", 15, 7000000): "cash",
        ("01-front-b.pptx", 15, 18): "runway", ("01-front-b.pptx", 15, 10): "months_to_profitability",
        ("02-moz.pdf", 20, 900): "ltv",
        ("02-moz.pdf", 20, 100): "cac", ("02-moz.pdf", 20, 9): "customer_lifetime",
        ("03-buffer.pptx", 7, 240): "ltv",
    }
    seen, profitable = {}, []
    for file in ("01-front-b.pptx", "02-moz.pdf", "03-buffer.pptx"):
        for c in claims.detect_candidates(parser.parse_deck((DECKS / file).read_bytes(), file)["blocks"], file):
            if c["snippet"] == "mozis profitable.":
                profitable.append((c["claim_type"], c["value"], c["target_date"]))
            if c["claim_type"] in NEW_TYPES:
                seen.setdefault((file, _page(c), c["value"]), []).append((c["claim_type"], c["snippet"]))
    moz_100 = seen.pop(("02-moz.pdf", 20, 100))
    assert sorted(moz_100) == [("cac", "~$100"), ("trials_per_day", "~100")], "the CAC and the trials, both 100"
    assert {k: [t for t, _ in v] for k, v in seen.items()} == {k: [t] for k, t in retyped.items() if k != ("02-moz.pdf", 20, 100)}
    assert profitable == [("ebitda", None, "2008-10")], "moz p2's dated profitable line stays an EBITDA milestone"


def test_only_stated_periods_with_different_values_are_a_deck_inconsistency():
    rows = [["", "Y/E 22", "Y/E 23"], ["Gross Profit", "£ 42,638", "£ 50,000"], ["Users", "200", "5,000"]]
    found = _found(_slide([("23 Y/E\nGross Profit £150K\n5K Users", 1, 1)], table=(rows, 4, 1)))
    flagged = sorted((c["claim_type"], bool(c.get("by_period"))) for c in found if c["inconsistent_dates"] == ["2023"])
    assert flagged == [("gross_profit", False), ("gross_profit", True)], "the same users figure is consistent"
    # Two figures that borrow the same year from the title state no period of their own.
    borrowed = _found(_slide([("Revenue $10M", 0.5, 2), ("Revenue run rate $8M", 6, 6)], title="Plan 2011"))
    assert sorted((c["value"], c["target_date"], c["inconsistent_dates"]) for c in borrowed) == \
        [(8000000, "2011", []), (10000000, "2011", [])]


def test_one_value_of_a_row_can_be_edited_and_the_row_approved_once(api):
    client, db = api
    _upload(client, "audit-1", "05-zero2hero.pdf", (DECKS / "05-zero2hero.pdf").read_bytes())
    row = next(c for c in db[decks.CANDIDATES_COLLECTION].docs if c.get("by_period") and c["claim_type"] == "revenue")
    url = f"/api/audits/audit-1/decks/candidates/{row['id']}"
    values = [{"value": i["value"], "value_high": None, "target_date": i["target_date"]} for i in row["by_period"]]
    values[1] = {**values[1], "value": 160000}
    edited = client.put(url, json={"by_period": values}).json()
    assert edited["status"] == "edited" and [i["value"] for i in edited["by_period"]][:2] == [130550, 160000]
    assert [i["period"] for i in edited["by_period"]] == [i["period"] for i in row["by_period"]], "periods stay"
    assert edited["by_period"][1]["source"] == row["by_period"][1]["source"], "cells stay"
    assert edited["parsed"]["by_period"][1]["value"] == 150000, "the parser's row kept beside the edit"
    assert client.put(url, json={"by_period": values[:2]}).status_code == 400, "one value per period"
    assert client.put(url, json={"value": 1}).status_code == 400, "a row is edited by period"
    other = next(c for c in db[decks.CANDIDATES_COLLECTION].docs if c.get("by_period") and c["claim_type"] == "users")
    assert client.put(f"/api/audits/audit-1/decks/candidates/{other['id']}", json={"status": "approved"}).json()["status"] \
        == "approved", "the analyst approves the row once"
    register = client.get("/api/audits/audit-1/claims").json()["claims"]
    assert sorted(c["claim_type"] for c in register) == ["revenue", "users"]
    again = _upload(client, "audit-1", "05-zero2hero.pdf", (DECKS / "05-zero2hero.pdf").read_bytes()).json()
    assert again["kept_reviewed"] == 2
    assert len([c for c in db[decks.CANDIDATES_COLLECTION].docs if c["claim_type"] == "revenue" and c.get("by_period")]) == 1
    half = client.put(f"/api/audits/audit-1/decks/candidates/{other['id']}", json={"by_period": [
        {"value": i["value"], "target_date": "2024-H1"} for i in other["by_period"]]})
    assert half.status_code == 200, "a half year is a target date"


def test_a_merged_claim_is_compared_when_any_of_its_sources_states_the_period():
    blocks = [{"slide": 1, "kind": "text", "text": "Plan 2023", "box": 1, "title": True},
              {"slide": 1, "kind": "text", "text": "Revenue £1M", "box": 2},
              {"slide": 2, "kind": "text", "text": "Revenue 2023: £1M", "box": 3},
              {"slide": 3, "kind": "text", "text": "Revenue 2023: £2M", "box": 4}]
    found = claims.detect_candidates(blocks, "deck.pptx")
    assert sorted((c["value"], len(c["sources"]), c["inconsistent_dates"]) for c in found) == \
        [(1000000, 2, ["2023"]), (2000000, 1, ["2023"])]


# ---------------------------------------------------------------------------
# Fiscal year-end (spec section 2, period rules): fiscal years are named by the calendar year in
# which they end, and every period resolves to a start and an end date. Display keeps the text.
# With a year-end other than December every year, quarter and half label is fiscal ("2025E",
# "Q1 25", "H1 25"); months stay calendar months, but a month under a year header falls inside that
# year: after the year-end month it is in the previous calendar year. With December nothing moves.
# ---------------------------------------------------------------------------
_FISCAL_BLOCKS = [
    {"slide": 1, "kind": "text", "text": "FY25 ARR $3M", "box": 1},
    {"slide": 2, "kind": "text", "text": "Y/E 22 revenue £1M", "box": 2},
    {"slide": 3, "kind": "text", "text": "FY2025/26 revenue £4M", "box": 3},
    {"slide": 4, "kind": "text", "text": "Revenue Q3 25 $2M", "box": 4},
    {"slide": 5, "kind": "text", "text": "Revenue H1 24 $5M", "box": 5},
    {"slide": 6, "kind": "text", "text": "Launch in March 2025", "box": 6},
    {"slide": 7, "kind": "text", "text": "Revenue in 2024 $6M", "box": 7},
    {"slide": 8, "kind": "table", "table": 1, "row": 1, "col": 1, "text": "Metric"},
    {"slide": 8, "kind": "table", "table": 1, "row": 1, "col": 2, "text": "FY23"},
    {"slide": 8, "kind": "table", "table": 1, "row": 1, "col": 3, "text": "FY24"},
    {"slide": 8, "kind": "table", "table": 1, "row": 2, "col": 1, "text": "Users"},
    {"slide": 8, "kind": "table", "table": 1, "row": 2, "col": 2, "text": "200"},
    {"slide": 8, "kind": "table", "table": 1, "row": 2, "col": 3, "text": "5,000"},
    {"slide": 9, "kind": "text", "text": "Revenue 2025E $7M", "box": 9},
    {"slide": 10, "kind": "text", "text": "Revenue Q1 FY25 $8M", "box": 10},
    {"slide": 11, "kind": "text", "text": "FY26 H2 revenue $9M", "box": 11},
    {"slide": 12, "kind": "table", "table": 2, "row": 1, "col": 1, "text": "Metric"},
    {"slide": 12, "kind": "table", "table": 2, "row": 1, "col": 2, "text": "FY2025", "col_span": 2},
    {"slide": 12, "kind": "table", "table": 2, "row": 2, "col": 2, "text": "Q3"},
    {"slide": 12, "kind": "table", "table": 2, "row": 2, "col": 3, "text": "Q4"},
    {"slide": 12, "kind": "table", "table": 2, "row": 3, "col": 1, "text": "Customers"},
    {"slide": 12, "kind": "table", "table": 2, "row": 3, "col": 2, "text": "30"},
    {"slide": 12, "kind": "table", "table": 2, "row": 3, "col": 3, "text": "40"},
    {"slide": 13, "kind": "table", "table": 3, "row": 1, "col": 1, "text": "Metric"},
    {"slide": 13, "kind": "table", "table": 3, "row": 1, "col": 2, "text": "FY2025", "col_span": 3},
    {"slide": 13, "kind": "table", "table": 3, "row": 1, "col": 5, "text": "2025"},
    {"slide": 13, "kind": "table", "table": 3, "row": 2, "col": 2, "text": "Mar"},
    {"slide": 13, "kind": "table", "table": 3, "row": 2, "col": 3, "text": "Apr"},
    {"slide": 13, "kind": "table", "table": 3, "row": 2, "col": 4, "text": "Dec"},
    {"slide": 13, "kind": "table", "table": 3, "row": 2, "col": 5, "text": "Jul"},
    {"slide": 13, "kind": "table", "table": 3, "row": 3, "col": 1, "text": "Users"},
    {"slide": 13, "kind": "table", "table": 3, "row": 3, "col": 2, "text": "10"},
    {"slide": 13, "kind": "table", "table": 3, "row": 3, "col": 3, "text": "20"},
    {"slide": 13, "kind": "table", "table": 3, "row": 3, "col": 4, "text": "30"},
    {"slide": 13, "kind": "table", "table": 3, "row": 3, "col": 5, "text": "50"},
    {"slide": 14, "kind": "text", "text": "Revenue Apr FY25 $10M", "box": 14},
]


def _periods(found):
    """{stated text: (target_date, start, end)} over every value of every candidate."""
    out = {}
    for c in found:
        for v in c.get("by_period") or [c]:
            out[v["period_text"]] = (v["target_date"], v["period_start"], v["period_end"])
    return out


@pytest.mark.parametrize("year_end, expected", [
    (12, {"FY25": ("2025", "2025-01-01", "2025-12-31"), "Y/E 22": ("2022", "2022-01-01", "2022-12-31"),
          "FY2025/26": ("2026", "2026-01-01", "2026-12-31"), "FY23": ("2023", "2023-01-01", "2023-12-31"),
          "FY24": ("2024", "2024-01-01", "2024-12-31"), "2024": ("2024", "2024-01-01", "2024-12-31"),
          "2025E": ("2025", "2025-01-01", "2025-12-31"), "Q3 25": ("2025-Q3", "2025-07-01", "2025-09-30"),
          "H1 24": ("2024-H1", "2024-01-01", "2024-06-30"), "Q1 FY25": ("2025-Q1", "2025-01-01", "2025-03-31"),
          "FY26 H2": ("2026-H2", "2026-07-01", "2026-12-31"), "Q3 FY2025": ("2025-Q3", "2025-07-01", "2025-09-30"),
          "Q4 FY2025": ("2025-Q4", "2025-10-01", "2025-12-31"),
          "March 2025": ("2025-03", "2025-03-01", "2025-03-31"),
          "Mar FY2025": ("FY2025-03", "2025-03-01", "2025-03-31"), "Apr FY2025": ("FY2025-04", "2025-04-01", "2025-04-30"),
          "Dec FY2025": ("FY2025-12", "2025-12-01", "2025-12-31"), "Jul 2025": ("FY2025-07", "2025-07-01", "2025-07-31"),
          "Apr FY25": ("FY2025-04", "2025-04-01", "2025-04-30")}),
    (3, {"FY25": ("2025", "2024-04-01", "2025-03-31"), "Y/E 22": ("2022", "2021-04-01", "2022-03-31"),
         "FY2025/26": ("2026", "2025-04-01", "2026-03-31"), "FY23": ("2023", "2022-04-01", "2023-03-31"),
         "FY24": ("2024", "2023-04-01", "2024-03-31"), "2024": ("2024", "2023-04-01", "2024-03-31"),
         "2025E": ("2025", "2024-04-01", "2025-03-31"), "Q3 25": ("2025-Q3", "2024-10-01", "2024-12-31"),
         "H1 24": ("2024-H1", "2023-04-01", "2023-09-30"), "Q1 FY25": ("2025-Q1", "2024-04-01", "2024-06-30"),
         "FY26 H2": ("2026-H2", "2025-10-01", "2026-03-31"), "Q3 FY2025": ("2025-Q3", "2024-10-01", "2024-12-31"),
         "Q4 FY2025": ("2025-Q4", "2025-01-01", "2025-03-31"),
         "March 2025": ("2025-03", "2025-03-01", "2025-03-31"),
         # A month under a year header: up to the year-end month in the named year, after it the year before.
         "Mar FY2025": ("FY2025-03", "2025-03-01", "2025-03-31"), "Apr FY2025": ("FY2025-04", "2024-04-01", "2024-04-30"),
         "Dec FY2025": ("FY2025-12", "2024-12-01", "2024-12-31"), "Jul 2025": ("FY2025-07", "2024-07-01", "2024-07-31"),
         "Apr FY25": ("FY2025-04", "2024-04-01", "2024-04-30")}),
    (6, {"FY25": ("2025", "2024-07-01", "2025-06-30"), "Q3 25": ("2025-Q3", "2025-01-01", "2025-03-31"),
         "H1 24": ("2024-H1", "2023-07-01", "2023-12-31"), "March 2025": ("2025-03", "2025-03-01", "2025-03-31"),
         "Mar FY2025": ("FY2025-03", "2025-03-01", "2025-03-31"), "Apr FY2025": ("FY2025-04", "2025-04-01", "2025-04-30"),
         "Dec FY2025": ("FY2025-12", "2024-12-01", "2024-12-31"), "Jul 2025": ("FY2025-07", "2024-07-01", "2024-07-31"),
         "Apr FY25": ("FY2025-04", "2025-04-01", "2025-04-30")}),
])
def test_every_year_quarter_and_half_follows_the_year_end_and_months_stay_calendar(year_end, expected):
    periods = _periods(claims.detect_candidates(_FISCAL_BLOCKS, "deck.pptx", fiscal_year_end=year_end))
    assert {k: periods.get(k) for k in expected} == expected


def test_every_period_resolves_to_a_start_and_an_end_and_the_default_year_end_is_december():
    found = claims.detect_candidates(_FISCAL_BLOCKS, "deck.pptx")
    assert _periods(found)["FY25"] == ("2025", "2025-01-01", "2025-12-31")
    for year_end in (12, 3):
        values = [v for c in claims.detect_candidates(_FISCAL_BLOCKS, "deck.pptx", fiscal_year_end=year_end)
                  for v in (c.get("by_period") or [c])]
        assert values and all(v["period_start"] and v["period_end"] and v["period_start"] <= v["period_end"]
                              for v in values if v["target_date"]), year_end


def test_changing_the_year_end_re_maps_the_stored_periods():
    found = claims.detect_candidates(_FISCAL_BLOCKS, "deck.pptx", fiscal_year_end=12)
    claims.remap_periods(found, 3)
    assert _periods(found)["FY25"] == ("2025", "2024-04-01", "2025-03-31")
    assert _periods(found)["Q3 25"] == ("2025-Q3", "2024-10-01", "2024-12-31")
    assert _periods(found)["March 2025"] == ("2025-03", "2025-03-01", "2025-03-31")
    claims.remap_periods(found, 12)
    assert _periods(found)["Q3 25"] == ("2025-Q3", "2025-07-01", "2025-09-30")


def test_the_audit_year_end_is_used_at_upload_and_a_change_re_maps_the_stored_claims(api):
    client, db = api
    assert client.put("/api/audits/audit-1", json={"fiscal_year_end": 3}).status_code == 200
    assert _upload(client, "audit-1", "fy.pptx", _slide([("FY25 ARR $3M", 1, 2), ("Revenue Q3 25 $2M", 1, 4)])).status_code == 200
    stored = {c["period_text"]: c for c in db[decks.CANDIDATES_COLLECTION].docs}
    assert (stored["FY25"]["target_date"], stored["FY25"]["period_start"], stored["FY25"]["period_end"]) == \
        ("2025", "2024-04-01", "2025-03-31")
    assert (stored["Q3 25"]["period_start"], stored["Q3 25"]["period_end"]) == ("2024-10-01", "2024-12-31")
    assert client.put("/api/audits/audit-1", json={"fiscal_year_end": 12}).status_code == 200
    stored = {c["period_text"]: c for c in db[decks.CANDIDATES_COLLECTION].docs}
    assert (stored["FY25"]["period_start"], stored["FY25"]["period_end"]) == ("2025-01-01", "2025-12-31")
    assert (stored["Q3 25"]["period_start"], stored["Q3 25"]["period_end"]) == ("2025-07-01", "2025-09-30")
    for bad in (0, 13):
        assert client.put("/api/audits/audit-1", json={"fiscal_year_end": bad}).status_code == 422


def test_an_edited_date_follows_the_year_end_and_drops_the_stated_text(api):
    client, db = api
    client.put("/api/audits/audit-1", json={"fiscal_year_end": 3})
    _upload(client, "audit-1", "fy.pptx", _slide([("FY25 ARR $3M", 1, 2)]))
    arr, = db[decks.CANDIDATES_COLLECTION].docs
    url = f"/api/audits/audit-1/decks/candidates/{arr['id']}"
    kept = client.put(url, json={"value": 3100000, "target_date": "2025"}).json()
    assert (kept["period_text"], kept["period_start"]) == ("FY25", "2024-04-01"), "same date: the stated period stays"
    moved = client.put(url, json={"target_date": "2026-Q1"}).json()
    assert (moved["period_text"], moved["period_start"], moved["period_end"]) == (None, "2025-04-01", "2025-06-30")
    month = client.put(url, json={"target_date": "2026-02"}).json()
    assert (month["period_start"], month["period_end"]) == ("2026-02-01", "2026-02-28"), "a month stays calendar"
    in_year = client.put(url, json={"target_date": "FY2026-04"}).json()
    assert (in_year["period_start"], in_year["period_end"]) == ("2025-04-01", "2025-04-30"), "April of FY2026"


def test_a_table_row_with_months_under_a_year_header_can_be_edited(api):
    client, db = api
    client.put("/api/audits/audit-1", json={"fiscal_year_end": 3})
    blocks = [b for b in _FISCAL_BLOCKS if b.get("slide") == 13]
    db[decks.CANDIDATES_COLLECTION].docs.extend(
        {**c, "id": f"c{i}", "audit_id": "audit-1", "deck_id": "d1", "status": "pending"}
        for i, c in enumerate(claims.detect_candidates(blocks, "deck.pptx", fiscal_year_end=3)))
    row = next(c for c in db[decks.CANDIDATES_COLLECTION].docs if c.get("by_period"))
    edit = [{"value": i["value"] + 1, "value_high": None, "target_date": i["target_date"]} for i in row["by_period"]]
    r = client.put(f"/api/audits/audit-1/decks/candidates/{row['id']}", json={"by_period": edit})
    assert r.status_code == 200, r.text
    assert [(i["target_date"], i["period_start"]) for i in r.json()["by_period"]][:2] == \
        [("FY2025-03", "2025-03-01"), ("FY2025-04", "2024-04-01")]


# ---------------------------------------------------------------------------
# Structure detection (spec section 7): tables, charts, KPI panels, roadmaps and timelines, and the
# hiring, unit-economics and use-of-funds tables. The thresholds were fixed on the 10 test decks.
# ---------------------------------------------------------------------------
from app.structures import redact as structure_redact  # noqa: E402

# front-b p16's panel held only dates written with slashes, so it is none (issue #53): 21 KPI panels.
DECK_STRUCTURES = {
    "01-front-b.pptx": [(11, "kpi_panel"), (12, "kpi_panel"), (14, "kpi_panel"), (15, "kpi_panel"), (18, "kpi_panel")],
    "02-moz.pdf": [(2, "roadmap"), (13, "kpi_panel"), (20, "kpi_panel"), (21, "kpi_panel"), (23, "kpi_panel"),
                   (32, "kpi_panel")],
    "03-buffer.pptx": [(6, "roadmap")],
    "04-clevergig.docx": [(7, "kpi_panel")],
    "05-zero2hero.pdf": [(11, "kpi_panel"), (17, "kpi_panel"), (19, "table"), (19, "kpi_panel"), (22, "hiring_table")],
    "06-uber.pdf": [],
    "07-equals-seed.docx": [],
    "08-genesisai-2021.pdf": [(13, "kpi_panel"), (14, "kpi_panel")],
    "09-genesisai-2024.pdf": [(5, "kpi_panel"), (13, "kpi_panel"), (14, "kpi_panel")],
    "10-tea.pdf": [(6, "kpi_panel"), (9, "kpi_panel"), (11, "roadmap")],
}


@pytest.mark.parametrize("file", sorted(DECK_STRUCTURES))
def test_the_structures_found_on_the_test_decks(file):
    deck = parser.parse_deck((DECKS / file).read_bytes(), file)
    assert [(s.get("slide") or s.get("page"), s["type"]) for s in deck["structures"]] == DECK_STRUCTURES[file]
    for s in deck["structures"]:
        assert s["type"] in parser.STRUCTURE_TYPES and s["cells"]
        assert all(len(c["text"]) <= parser.CELL_MAX for c in s["cells"])
        assert len({(c["row"], c["col"]) for c in s["cells"]}) == len(s["cells"]), "one cell per position"


def test_chart_axes_drawn_as_text_are_not_a_timeline_and_a_timeline_needs_a_product_word():
    # front-b slides 10-13: quarter labels under charts drawn as text boxes, evenly spaced.
    deck = parser.parse_deck((DECKS / "01-front-b.pptx").read_bytes(), "01-front-b.pptx")
    assert not [s for s in deck["structures"] if s["type"] == "roadmap"]
    dated = _slide([("Q1 2024", 1, 1), ("Q3 2024", 3, 1), ("Q2 2025", 5, 1), ("Launch the API", 1, 2)])
    assert [s["type"] for s in parser.parse_deck(dated, "d.pptx")["structures"]] == ["roadmap"]
    plain = _slide([("Q1 2024", 1, 1), ("Q3 2024", 3, 1), ("Q2 2025", 5, 1), ("Hired a VP", 1, 2)])
    assert not [s for s in parser.parse_deck(plain, "d.pptx")["structures"] if s["type"] == "roadmap"]
    axis = _slide([("Q1 2024", 1, 1), ("Q2 2024", 3, 1), ("Q3 2024", 5, 1), ("Launch the API", 1, 2)])
    assert not [s for s in parser.parse_deck(axis, "d.pptx")["structures"] if s["type"] == "roadmap"], \
        "evenly spaced, distinct dates in a row are an axis"


def test_the_timeline_grid_keeps_each_box_as_a_column_of_its_band():
    deck = parser.parse_deck((DECKS / "10-tea.pdf").read_bytes(), "10-tea.pdf")
    roadmap, = [s for s in deck["structures"] if s["type"] == "roadmap"]
    cells = {(c["row"], c["col"]): c for c in roadmap["cells"]}
    assert (cells[(1, 1)]["text"], cells[(2, 1)]["text"], cells[(2, 2)]["text"]) == ("2021", "Q2", "Gluon wallet")
    assert cells[(1, 1)]["box"] == cells[(2, 1)]["box"] != cells[(2, 2)]["box"]


def test_a_kpi_box_is_short_lines_with_a_figure_and_a_label_never_a_wrapped_sentence():
    found = parser.parse_deck(_slide([("2.5 hours\nper user per day", 1, 1), ("64%\nDAU / MAU ratio", 4, 1),
                                      ("We raised revenue of\n$1.1M from customers", 1, 3), ("40%", 4, 3)]), "d.pptx")
    panel, = found["structures"]
    assert panel["type"] == "kpi_panel"
    assert [c["text"] for c in sorted(panel["cells"], key=lambda c: (c["row"], c["col"]))] == \
        ["2.5 hours", "64%", "per user per day", "DAU / MAU ratio"], "the wrapped sentence and the bare 40% stay out"


# ---------------------------------------------------------------------------
# KPI panels keep their label boxes, the value boxes beside them and a title as label (spec section 7, decisions of
# 2026-10-06, issue #48). Built from the public test decks: moz p20 and p21, front-b p12 and p15.
# ---------------------------------------------------------------------------
MOZ_P20 = [["2011 Estimated Revenue", "$12 -$13 million"], ["Current Revenue Run Rate (June)", "~$10.8 million"],
           ["Number of PRO Subscribers", "~13,500"], ["# of New Free Trials / Day", "~100"],
           ["Avg. Customer Lifetime Value", "~$900"], ["Implied Customer Life", "~9 Months"],
           ["Avg. Cost of Paid Acquisition", "~$100"], ["Avg. Monthly Revenue / Subscriber", "~$93"]]
MOZ_P21 = [["% of Free Trials Converting to Paid", "~57%"], ["Churn Rate in 1st2 Paid Months", "~25%"],
           ["Monthly Visits to Moz+ OSE", "~1.25 million"], ["Email Subscribers", "~300K"], ["Gross Margins", "~82%"],
           ["Estimated Net Profit in 2011", "~$1 million"], ["Staffing Costs", "~$650K / Month"],
           ["Crawling, Serving, Hosting + Processing", "~$180K / Month"]]
# The value boxes that hold only a figure and never reached the model before (issue #48).
MOZ_VALUE_BOXES = ["~13,500", "~100", "~$900", "~$100", "~$93", "~300K", "~82%", "~57%", "~25%"]


def _panel(file, page):
    deck = parser.parse_deck((DECKS / file).read_bytes(), file)
    panel, = [s for s in deck["structures"] if (s.get("slide") or s.get("page")) == page and s["type"] == "kpi_panel"]
    return panel


def _grid_rows(panel):
    """Each grid row's cell texts, left to right."""
    rows = {}
    for c in sorted(panel["cells"], key=lambda c: (c["row"], c["col"])):
        rows.setdefault(c["row"], []).append(c["text"])
    return list(rows.values())


def _panel_texts(boxes, title=None):
    """The cell texts of the KPI panels of a built slide."""
    found = parser.parse_deck(_slide(boxes, title=title), "d.pptx")["structures"]
    return [c["text"] for s in found if s["type"] == "kpi_panel" for c in s["cells"]]


def test_moz_p20_and_p21_every_value_joins_the_panel_beside_its_label():
    """Every label of these pages is its own text box, and nine value boxes hold only a figure: before, the labels
    were left out and the nine never reached the model."""
    p20, p21 = _grid_rows(_panel("02-moz.pdf", 20)), _grid_rows(_panel("02-moz.pdf", 21))
    assert p20 == MOZ_P20 and p21 == MOZ_P21
    assert [row[1] for row in p20 + p21 if row[1] in MOZ_VALUE_BOXES] == \
        ["~13,500", "~100", "~$900", "~$100", "~$93", "~57%", "~25%", "~300K", "~82%"]


def test_the_slide_title_labels_a_value_box_with_no_label_of_its_own_and_is_marked():
    """Decision 2 of 2026-10-06: moz p20's "2011 Estimated Revenue" is the page title, beside "$12 -$13 million"."""
    for page, title in ((20, "2011 Estimated Revenue"), (21, "% of Free Trials Converting to Paid")):
        panel = _panel("02-moz.pdf", page)
        assert [(c["row"], c["col"], c["text"]) for c in panel["cells"] if c.get("title")] == [(1, 1, title)]
    text = structure_redact.structure_text(_panel("02-moz.pdf", 20)["cells"])
    assert text.splitlines()[:3] == ["r1c1 title: 2011 Estimated Revenue", "r1c2: $12 -$13 million",
                                     "r2c1: Current Revenue Run Rate (June)"]
    kpi = ("Profitable in 10 months", 6, 5)               # a KPI box, so the slide has a panel
    assert _panel_texts([("$7m", 1, 2), kpi], title="Runway") == ["Runway", "$7m", "Profitable in 10 months"], \
        "the title right above a value box with no label is its label"
    assert "Runway" not in _panel_texts([("Cash on hand", 3.2, 2), ("$7m", 1, 2), kpi], title="Runway"), \
        "a value box with a label box beside it keeps that label; the title stays out"
    assert "Runway 2" not in _panel_texts([("$7m", 1, 2), kpi], title="Runway 2"), "a title with a figure is no label"
    from pptx.util import Inches
    prs = pptx.Presentation(io.BytesIO(_slide([("$7m", 1, 2.2), kpi], title="Runway")))
    prs.slides[0].shapes.title.top = Inches(3)               # the title 0.3 inch below the value box
    buf = io.BytesIO()
    prs.save(buf)
    found = parser.parse_deck(buf.getvalue(), "d.pptx")["structures"]
    assert [c["text"] for s in found for c in s["cells"]] == ["Profitable in 10 months"], "never a title below it"


def test_front_b_p15_keeps_cash_on_hand_next_to_7m_left_and_runway():
    rows = _grid_rows(_panel("01-front-b.pptx", 15))
    row, = [r for r in rows if "$7m left" in r]
    assert row[row.index("$7m left") - 1] == "Cash on hand" and {"18 months", "Runway *"} <= set(row)


def test_front_b_p12_keeps_its_label_wrapped_over_two_lines_and_ltv_cac():
    """"Spend as / % of revenue" ends its first line on "as", yet read as one line it is a 21-character label."""
    texts = [c["text"] for c in _panel("01-front-b.pptx", 12)["cells"]]
    assert {"Spend as", "% of revenue", "LTV / CAC", "18% 18% 19%", "2.5 2.6 4.4"} <= set(texts)


def test_the_label_length_is_the_longest_test_deck_label_plus_half_and_one_over_stays_out():
    """Decision 1 of 2026-10-06: the longest label on the 10 test decks (moz p21) plus 50%, rounded up."""
    longest = "Crawling, Serving, Hosting + Processing"
    assert longest in [row[0] for row in MOZ_P21] and parser.LABEL_MAX == -(-len(longest) * 3 // 2) == 59
    assert parser.KPI_LINE_MAX == 30, "a KPI box's lines keep their 30 characters"
    for length, kept in ((parser.LABEL_MAX, True), (parser.LABEL_MAX + 1, False)):
        label = ("Average monthly revenue per paying subscriber " + "x" * 60)[:length]
        assert (label in _panel_texts([(label, 1, 1), ("$7m left", 1, 1.6)])) is kept, length


def test_a_wrapped_sentence_still_stays_out():
    sentence = "We took one round of\nfinancing and grew our\nsubscriber base steadily"     # 68 characters as one line
    assert _panel_texts([("$7m left", 1, 1), (sentence, 3.2, 1), ("~40%", 5.4, 1),
                         ("We raised revenue of\n$1.1M from customers", 1, 3)]) == ["$7m left"], \
        "neither a label box nor a KPI box; the value box beside it has no label"
    assert _panel_texts([("$7m left", 1, 1), ("Spend as\n% of revenue", 3.2, 1), ("~40%", 5.4, 1)]) == \
        ["$7m left", "Spend as", "% of revenue", "~40%"], "a label wrapped over two lines is one label"


def test_directly_next_to_is_beside_or_within_a_tenth_of_the_page_with_no_box_between():
    kpi = ("Profitable in 10 months", 6, 5)               # a KPI box apart from the rest, so the slide has a panel
    assert "Cash on hand" in _panel_texts([kpi, ("Cash on hand", 1, 1), ("$7m", 1, 1.6)]), "0.1 inch above"
    assert "Cash on hand" not in _panel_texts([kpi, ("Cash on hand", 1, 1), ("$7m", 1, 2.4)]), \
        "0.9 inch above: more than a tenth of a 7.5-inch slide"
    assert "Cash on hand" not in _panel_texts([kpi, ("Cash on hand", 1, 1), ("In the bank", 1, 1.6), ("$7m", 1, 2.2)]), \
        "another box between them"
    texts = [c["text"] for c in _panel("05-zero2hero.pdf", 17)["cells"]]
    assert not {"Turnover(£/year)", "278,085", "181,193"} & set(texts), \
        "zero2hero p17: the chart title sits a quarter of the page above its data labels"


def test_an_axis_tick_or_a_date_written_with_slashes_is_no_value_box():
    texts = [c["text"] for c in _panel("05-zero2hero.pdf", 17)["cells"]]
    assert not {"Traction", "800,000"} & set(texts), "zero2hero p17: the axis' top tick sits right under the title"
    rows = _grid_rows(_panel("01-front-b.pptx", 14))
    row, = [r for r in rows if "Recommend to a friend" in r]
    assert row[:4] == ["100%", "Recommend to a friend", "100%", "Approve of CEO"]
    assert not [t for r in rows for t in r if "/1/1" in t], "front-b p14: \"3/1/15\" is a date, no figure (#53)"


def test_a_page_with_no_kpi_box_has_no_panel():
    assert _panel_texts([("Cash on hand", 1, 1), ("$7m", 1, 1.6)], title="Our runway") == []


# ---------------------------------------------------------------------------
# A tall box merges two visual rows into one grid band (spec section 7, issue #50), so each KPI cell keeps the cells of
# other boxes whose line is directly next to its own. Built from the public test deck front-b, slides 15, 12 and 16.
# ---------------------------------------------------------------------------
def _cell_of(panel, text):
    cell, = [c for c in panel["cells"] if c["text"] == text]
    return cell


def _cell_id(cell):
    return f"r{cell['row']}c{cell['col']}"


def test_front_b_p15_one_grid_row_holds_two_visual_rows_and_next_to_tells_them_apart():
    """"“Default alive” † / Profitable in 10 months" spans the label row and the value row, so all five boxes share a
    grid row; "Cash on hand" sits above "$7m left", "Runway *" above "18 months"."""
    panel = _panel("01-front-b.pptx", 15)
    cash, left, months, runway = (_cell_of(panel, t) for t in ("Cash on hand", "$7m left", "18 months", "Runway *"))
    assert len({cash["row"], left["row"], months["row"], runway["row"]}) == 1, "one grid row"
    assert _cell_id(cash) in left["next_to"] and _cell_id(left) in cash["next_to"], "directly above it"
    assert _cell_id(cash) not in months["next_to"] and _cell_id(runway) in months["next_to"]


def test_front_b_p12_a_tall_box_between_two_lines_on_one_visual_row_keeps_them_apart():
    panel = _panel("01-front-b.pptx", 12)
    weeks = _cell_of(panel, "2 weeks ago")
    assert not {_cell_id(_cell_of(panel, t)) for t in ("LTV / CAC", "Spend as", "% of revenue")} & set(weeks["next_to"])
    assert _cell_id(_cell_of(panel, "The team isn’t one year old")) in _cell_of(panel, "joined 6 months ago")["next_to"]


def test_next_to_is_measured_line_to_line_not_box_to_box():
    """Built after front-b p16 (whose panel went with issue #53): the legend entry "Cash" sits on the "Gross margin"
    line's row, beside the box "$4M $6M / Gross margin", not on its figures' row, though the grid puts it there."""
    panel, = [s for s in parser.parse_deck(_slide([("Cash", 1, 1.25), ("$4M $6M\nGross margin", 3.5, 1)]),
                                           "d.pptx")["structures"] if s["type"] == "kpi_panel"]
    cash, figures = _cell_of(panel, "Cash"), _cell_of(panel, "$4M $6M")
    assert cash["row"] == figures["row"], "one grid row"
    assert _cell_id(cash) not in figures["next_to"] and _cell_id(cash) in _cell_of(panel, "Gross margin")["next_to"]
    assert all(i["headers"] == [] for i in items.list_items(panel)["items"]), "the legend entry labels no figure"


def _date_boxes(file, page):
    """{text line: the lines of its date box} for the lines of a roadmap that keep a date_box."""
    deck = parser.parse_deck((DECKS / file).read_bytes(), file)
    roadmap, = [s for s in deck["structures"] if (s.get("slide") or s.get("page")) == page and s["type"] == "roadmap"]
    return {c["text"]: [d["text"] for d in roadmap["cells"] if d["box"] == c["date_box"]]
            for c in roadmap["cells"] if "date_box" in c}


def test_a_roadmap_text_box_keeps_the_one_date_box_directly_next_to_it():
    """Spec section 7 (decision of 2026-10-06 on issue #47, option a): a date box holds only dates or parts of one;
    each line of another text box keeps the box number of the one date box directly next to its box, box to box."""
    moz = _date_boxes("02-moz.pdf", 2)
    assert sorted(v for v, in moz.values()) == ["1981", "1997", "2001", "2004", "Feb. 2007", "July 2011",
                                                 "Nov. 2007", "Oct. 2008", "Sept. 2010"], "one per paragraph"
    assert moz["Gillian (Rand’s Mom) founds the company that will become SEOmoz"] == ["1981"]
    tea = _date_boxes("10-tea.pdf", 11)
    assert (tea["Gluon wallet"], tea["Seed round secured including investment from Hashkey"],
            tea["Preview 1 version launch"]) == (["2021", "Q2"], ["2021", "Q2"], ["2021", "Q3"]), \
        "every line of the box, on either side"
    assert (tea["Majority of business logic migrated from layer-1 to layer-2"], tea["TEA framework dev guide released"],
            tea["Layer-1 EVM smart contract compatibility"]) == (["2022", "Q2"], ["2022", "Q2"], ["2022", "Q3"]), \
        "box 11 stands between box 10 and the 2022 Q3 box"
    assert len(tea) == 19 and tea["Mainnet starts"] == ["2023", "Q1-Q2"], \
        "issue #55: \"Q1-Q2\" is the first half, so \"2023\" / \"Q1-Q2\" is a date box; 19 bullets, 3 of them wrapped"
    assert _date_boxes("03-buffer.pptx", 6) == {}, "one box"


def test_roadmap_cells_keep_next_to_measured_line_to_line():
    """Spec section 7 (issue #56): roadmap cells keep next_to, as KPI panel cells do. moz p2's paragraphs sit directly
    over or under their dates, a one-box timeline (buffer p6) has no other box, and tea p11's wrapped bullet "Majority
    of business logic migrated from layer-1 to layer-2" (one cell since issue #55, measured over both its lines) has
    its date box "2022" / "Q2" on its left, and the bullet "Layer-1 EVM …" and another bullet's Q3 on its right."""
    near = {}
    for file, page in (("02-moz.pdf", 2), ("03-buffer.pptx", 6), ("10-tea.pdf", 11)):
        deck = parser.parse_deck((DECKS / file).read_bytes(), file)
        roadmap, = [s for s in deck["structures"] if (s.get("slide") or s.get("page")) == page]
        assert roadmap["type"] == "roadmap", (file, page)
        near[file] = {f"r{c['row']}c{c['col']}": c["next_to"] for c in roadmap["cells"]}
    assert near["02-moz.pdf"]["r1c1"] == ["r1c2", "r2c1"] and near["02-moz.pdf"]["r4c1"] == ["r3c1", "r4c2"]
    assert set(map(tuple, near["03-buffer.pptx"].values())) == {()}
    assert near["10-tea.pdf"]["r8c2"] == ["r8c1", "r9c1", "r8c3", "r9c4"]


# ---------------------------------------------------------------------------
# Roadmaps: a paragraph wrapped over the lines of a text box is one cell (issue #49)
# ---------------------------------------------------------------------------
def _roadmap(file, page):
    deck = parser.parse_deck((DECKS / file).read_bytes(), file)
    roadmap, = [s for s in deck["structures"] if (s.get("slide") or s.get("page")) == page and s["type"] == "roadmap"]
    return deck, roadmap


# moz p2's grid: four paragraphs over their dates, five under theirs, each paragraph one cell over or under its date.
MOZ_P2_GRID = [
    ["Rand starts working w/ Gillian building websites for small, local businesses",
     "Deeply in debt, and failing to get traffic to clients’ sites, Rand starts the SEOmoz Blog as part of learning "
     "the SEO process.",
     "SEOmoz takes an investment of $1.1M from Ignition Partners & Curious Office",
     "Moz’scollection of tools becomes a singular, campaign-based web app. Prices rise to $99 / $499 / $1999 per "
     "month."],
    ["1997", "2004", "Nov. 2007", "Sept. 2010"],
    ["1981", "2001", "Feb. 2007", "Oct. 2008", "July 2011"],
    ["Gillian (Rand’s Mom) founds the company that will become SEOmoz",
     "Rand drops out of UW, 2 classes from graduation to work full time w/ Gillian",
     "SEOmoz launches its first subscription software product, “PRO” for $39/month",
     "Linkscape, SEOmoz’sweb index and link graph, launches. By December, mozis profitable.",
     "SEOmozis moving from just “SEO” to social media, content marketing, analytics, local and video. To this end, "
     "we’ve acquired “Moz.com.”"],
]


def test_moz_p2_each_dated_paragraph_is_one_cell():
    """Issue #49: the pdf wraps each milestone over 2-6 lines, and a row per line made 40 text lines of the page's 9
    dated paragraphs. Each line continues the one above (lower case, "&", or after punctuation, a wrap word or an open
    bracket: "Gillian (Rand’s / Mom) founds the"), so each paragraph is one cell, its lines joined in order."""
    _, roadmap = _roadmap("02-moz.pdf", 2)
    assert _grid_rows(roadmap) == MOZ_P2_GRID


def test_buffer_p6_keeps_a_row_per_line():
    """Issue #49: buffer p6 is one box that alternates lines and dates, so every line stays a row."""
    deck, roadmap = _roadmap("03-buffer.pptx", 6)
    lines = [b["text"] for b in deck["blocks"] if b.get("slide") == 6 and b["kind"] == "text"]
    assert len(roadmap["cells"]) == 12 and all(c["text"] in lines for c in roadmap["cells"])


TEA_P11_WRAPPED = {"Seed round secured including investment from Hashkey": ("Seed round secured including investment",
                                                                           "from Hashkey"),
                   "Begin Go2Market strategy starting with miners' economy": (
                       "Begin Go2Market strategy starting with miners'", "economy"),
                   "Majority of business logic migrated from layer-1 to layer-2": (
                       "Majority of business logic migrated from", "layer-1 to layer-2")}


def test_tea_p11_each_wrapped_bullet_is_one_cell():
    """Issue #55: tea p11's boxes are bullet lists. A bullet wrapped over two lines gave a row per line, and each line
    could become its own milestone. The box splits into items at each line that does not continue the one above, and
    each item's lines are joined: 38 cells become 35, 30 text lines 27. "Second milestone ongoing in 2021" ends on a
    figure, not on the wrap word "in", so "Gluon wallet" below it stays its own bullet."""
    deck, roadmap = _roadmap("10-tea.pdf", 11)
    texts = [c["text"] for c in roadmap["cells"]]
    lines = [b["text"] for b in deck["blocks"] if b.get("page") == 11 and b["kind"] == "text"]
    assert len(texts) == 35 and all(t in lines or t in TEA_P11_WRAPPED for t in texts)
    assert set(TEA_P11_WRAPPED) <= set(texts), "each wrapped bullet is one cell, its lines joined in order"
    assert not {line for pair in TEA_P11_WRAPPED.values() for line in pair} & set(texts), "\"from Hashkey\" is no cell"
    assert {"Second milestone ongoing in 2021", "Gluon wallet"} <= set(texts)
    assert len([t for t in texts if not verify.is_date_line(t)]) == 27, "text lines"


def _roadmap_rows(boxes):
    """The grid rows of the roadmap of a built slide: three date boxes, then the given text boxes below them."""
    dates = [("Q1 2025", 1, 1), ("Q3 2025", 3.5, 1), ("Q2 2026", 6, 1)]
    roadmap, = [s for s in parser.parse_deck(_slide(dates + boxes), "d.pptx")["structures"] if s["type"] == "roadmap"]
    return _grid_rows(roadmap)


def test_a_roadmap_box_is_one_cell_only_when_each_line_continues_the_one_above():
    rows = _roadmap_rows([("We launch the app in\nthree new markets", 1, 2), ("Launch EU\n100 new hires", 3.5, 2),
                          ("Launch the API.\nHire a CFO", 6, 2)])
    assert rows == [["Q1 2025", "Q3 2025", "Q2 2026"],
                    ["We launch the app in three new markets", "Launch EU", "Launch the API. Hire a CFO"],
                    ["100 new hires"]], \
        "a figure after a line with no end punctuation starts a new item; a capital after a full stop does not"
    assert _roadmap_rows([("We ship the app in\nQ4 2026", 1, 2)])[1:] == [["We ship the app in"], ["Q4 2026"]], \
        "a date line is never joined, even after a wrap word"
    assert _roadmap_rows([("Hire a CFO\nLaunch the app in\nthree new markets\nOpen Berlin", 1, 2)])[1:] == \
        [["Hire a CFO"], ["Launch the app in three new markets"], ["Open Berlin"]], \
        "issue #55: a bullet list splits at each line that does not continue the one above; a wrapped bullet is one cell"
    assert _roadmap_rows([("Ship the beta to 12\nOpen Berlin", 1, 2)])[1:] == [["Ship the beta to 12"], ["Open Berlin"]], \
        "a line ending on a figure ends on no wrap word (tea p11: \"Second milestone ongoing in 2021\")"
    long =["We launch offices in three new markets across the region and", "hire local teams to sell the platform to",
            "mid-sized firms, with a partner programme that brings in", "resellers and integrators before the year ends"]
    assert all(len(line) <= parser.TIMELINE_LINE_MAX for line in long) and len(" ".join(long)) > parser.CELL_MAX
    assert _roadmap_rows([("\n".join(long), 1, 2)])[1:] == [[line] for line in long], \
        "a paragraph over 200 characters would be prose: its lines stay rows"


def test_front_b_p15_seed_to_series_a_joins_its_panel_like_its_twin():
    """"Seed to Series A" ends on a capital "A": a name, not the article that wraps a sentence over its lines."""
    texts = [c["text"] for c in _panel("01-front-b.pptx", 15)["cells"]]
    assert {"Seed to Series A", "$3.1m raised", "Series A to date", "$10m raised"} <= set(texts)
    assert _panel_texts([("Seed to Series A\n$3.1m raised", 1, 1), ("We closed a\n$3.1m seed round", 4, 1)]) == \
        ["Seed to Series A", "$3.1m raised"], "a line ending on a lower-case \"a\" still wraps a sentence"


def test_a_table_without_a_figure_is_prose_and_row_numbers_are_not_figures():
    deck = parser.parse_deck((DECKS / "07-equals-seed.docx").read_bytes(), "07-equals-seed.docx")
    assert deck["structures"] == [], "numbered text tables are prose"
    swot = parser.parse_deck((DECKS / "05-zero2hero.pdf").read_bytes(), "05-zero2hero.pdf")
    assert not [s for s in swot["structures"] if s.get("page") == 14], "a SWOT table numbered 1-12 is prose"


@pytest.mark.parametrize("header, caption, kind", [
    (["Role", "Start", "Salary"], None, "hiring_table"),
    (["Metric", "Value"], None, "table"),
    (["", "2024"], "Recruitment plan", "hiring_table"),
    (["Use of funds", "%"], None, "use_of_funds"),
    (["Metric", "2025"], "Unit economics", "unit_economics"),
    (["", "Value"], None, "unit_economics"),         # CAC in the first column
])
def test_a_table_takes_its_type_from_its_header_keywords_and_caption(header, caption, kind):
    rows = [header, ["CAC" if kind == "unit_economics" and caption is None else "Engineer", "£50,000"]]
    boxes = [(caption, 1, 1.4)] if caption else []
    deck = parser.parse_deck(_slide(boxes, table=(rows, 1, 2)), "d.pptx")
    assert [s["type"] for s in deck["structures"] if s.get("table")] == [kind]


def test_a_pptx_chart_is_read_from_its_xml_with_title_axis_titles_series_and_values():
    from pptx.chart.data import CategoryChartData
    from pptx.enum.chart import XL_CHART_TYPE
    from pptx.util import Inches
    prs = pptx.Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    data = CategoryChartData()
    data.categories = ["FY23", "FY24", "FY25"]
    data.add_series("ARR", (1.2, 2.5, 4.0))
    data.add_series("Customers", (40, 75, 120))
    chart = slide.shapes.add_chart(XL_CHART_TYPE.COLUMN_CLUSTERED, Inches(1), Inches(1), Inches(6), Inches(4), data).chart
    chart.has_title = True
    chart.chart_title.text_frame.text = "Growth plan"
    chart.value_axis.has_title = True
    chart.value_axis.axis_title.text_frame.text = "£m"
    buf = io.BytesIO()
    prs.save(buf)
    deck = parser.parse_deck(buf.getvalue(), "chart.pptx")
    found, = deck["structures"]
    assert (found["type"], found["slide"], found["header_rows"]) == ("chart", 1, 3)
    text = structure_redact.structure_text(found["cells"])
    assert text.splitlines()[:5] == ["r1c1: Growth plan (r1c1:r1c3)", "r2c2: £m (r2c2:r2c3)", "r3c2: ARR",
                                     "r3c3: Customers", "r4c1: FY23"]
    assert "r6c2: 4" in text and "r6c3: 120" in text


def _merged_table_slide(rows, merges):
    """A slide with one table; merges: [(row, col, rows, cols)], zero-based."""
    from pptx.util import Inches
    prs = pptx.Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    grid = slide.shapes.add_table(len(rows), len(rows[0]), Inches(1), Inches(2), Inches(6), Inches(2)).table
    for r, row in enumerate(rows):
        for c, text in enumerate(row):
            if text:
                grid.cell(r, c).text = text
    for r, c, nr, nc in merges:
        grid.cell(r, c).merge(grid.cell(r + nr - 1, c + nc - 1))
    buf = io.BytesIO()
    prs.save(buf)
    return buf.getvalue()


def test_a_merged_header_keeps_its_span_in_the_header_stack_and_the_structure_text():
    rows = [["", "2025", "", "", ""], ["", "Q1", "Q2", "Q3", "Q4"], ["Revenue", "$1M", "$2M", "$3M", "$4M"]]
    deck = parser.parse_deck(_merged_table_slide(rows, [(0, 1, 1, 4)]), "m.pptx")
    assert _has(deck["blocks"], slide=1, kind="table", row=1, col=2, text="2025", col_span=4)
    table, = deck["structures"]
    assert table["header_rows"] == 2
    text = structure_redact.structure_text(table["cells"])
    assert text.splitlines()[:2] == ["r1c2: 2025 (r1c2:r1c5)", "r2c2: Q1"]


def test_a_quarter_under_its_year_is_a_two_cell_period_and_with_no_year_it_has_none():
    rows = [["", "2025", "", "", ""], ["", "Q1", "Q2", "Q3", "Q4"], ["Revenue", "$1M", "$2M", "$3M", "$4M"]]
    row, = claims.detect_candidates(parser.parse_deck(_merged_table_slide(rows, [(0, 1, 1, 4)]), "m.pptx")["blocks"], "m.pptx")
    assert [(i["value"], i["target_date"], i["period"]) for i in row["by_period"]] == [
        (1000000, "2025-Q1", "Q1 2025"), (2000000, "2025-Q2", "Q2 2025"), (3000000, "2025-Q3", "Q3 2025"),
        (4000000, "2025-Q4", "Q4 2025")]
    no_year = _slide([("Plan 2024", 1, 1)], table=([["", "Q1", "Q2"], ["Revenue", "$1M", "$2M"]], 1, 2))
    row, = claims.detect_candidates(parser.parse_deck(no_year, "n.pptx")["blocks"], "n.pptx")
    assert [i["target_date"] for i in row["by_period"]] == [None, None], "never inferred from the slide title"
    relative = _slide([("Plan 2024", 1, 1)], table=([["", "Year 1", "Year 2"], ["Revenue", "$1M", "$2M"]], 1, 2))
    row, = claims.detect_candidates(parser.parse_deck(relative, "r.pptx")["blocks"], "r.pptx")
    assert [i["target_date"] for i in row["by_period"]] == [None, None], "relative columns have no period"


@pytest.mark.parametrize("months, expected", [
    (["Mär 2025", "Okt 2025"], ["2025-03", "2025-10"]),
    (["март 2025", "ДЕКЕМВРИ 2025"], ["2025-03", "2025-12"]),
    (["january 2025", "SEPT 2025"], ["2025-01", "2025-09"]),
])
def test_month_names_in_a_table_header_are_english_german_or_bulgarian_in_any_case(months, expected):
    row, = _found(_slide([], table=([[""] + months, ["Revenue", "€1M", "€2M"]], 1, 2)))
    assert [i["target_date"] for i in row["by_period"]] == expected


def test_a_docx_merged_cell_starts_in_its_grid_column_and_keeps_its_span():
    document = docx.Document()
    table = document.add_table(rows=2, cols=3)
    merged = table.cell(0, 0).merge(table.cell(0, 1))
    merged.text = "Plan"
    table.cell(0, 2).text = "FY2025"
    for c, text in enumerate(["Revenue", "€1M", "€2M"]):
        table.cell(1, c).text = text
    buf = io.BytesIO()
    document.save(buf)
    deck = parser.parse_deck(buf.getvalue(), "t.docx")
    assert _has(deck["blocks"], kind="table", row=1, col=1, text="Plan", col_span=2)
    assert _has(deck["blocks"], kind="table", row=1, col=3, text="FY2025"), "the cell after a merge starts in grid column 3"
    assert _has(deck["blocks"], kind="table", row=2, col=3, text="€2M")


def test_structures_are_stored_with_the_parsed_text_and_never_listed(api):
    client, db = api
    _upload(client, "audit-1", "05-zero2hero.pdf", (DECKS / "05-zero2hero.pdf").read_bytes())
    stored, = db[decks.TEXT_COLLECTION].docs
    assert [s["type"] for s in stored["structures"]] == [t for _, t in DECK_STRUCTURES["05-zero2hero.pdf"]]
    listed = client.get("/api/audits/audit-1/decks").json()["decks"][0]
    assert "structures" not in listed and "blocks" not in listed, "cells are deck text: never in a listing"
    client.delete(f"/api/audits/audit-1/decks/{stored['deck_id']}")
    assert db[decks.TEXT_COLLECTION].docs == [], "removed with the deck"


def test_bars_with_a_year_label_under_each_take_that_label_as_their_date():
    # Three bars: the value above each, the year under it, the bars' heights set by where the value sits.
    bars = _slide([("Revenue (£/year)", 1, 0.5), ("Revenue £1.2M", 1, 3), ("Revenue £2.4M", 3.5, 2.2),
                   ("Revenue £4.1M", 6, 1.4), ("FY2023", 1, 5), ("FY2024", 3.5, 5), ("FY2025", 6, 5)])
    found = sorted(_found(bars), key=lambda c: c["value"])
    assert [(c["value"], c["target_date"], c["period_text"], c["date_from"]) for c in found] == [
        (1200000, "2023", "FY2023", "FY2023"), (2400000, "2024", "FY2024", "FY2024"), (4100000, "2025", "FY2025", "FY2025")]
    assert all(c["currency"] == "GBP" for c in found)
    assert [c["period_start"] for c in found] == ["2023-01-01", "2024-01-01", "2025-01-01"]
    # Confidence is scored with the date attached: no "no date" check fails.
    assert all("no date" not in claims.confidence(c, found)["failed"] for c in found)

    # No label under the bars: no date, as before.
    bare = sorted(_found(_slide([("Revenue (£/year)", 1, 0.5), ("Revenue £1.2M", 1, 3), ("Revenue £2.4M", 3.5, 2.2),
                         ("Revenue £4.1M", 6, 1.4)])),
                  key=lambda c: c["value"])
    assert [c["target_date"] for c in bare] == [None, None, None]
    assert all("no date" in claims.confidence(c, bare)["failed"] for c in bare)

    # A year that stands alone under one figure is not an axis: two or more labels on one row make one.
    lone = _found(_slide([("Revenue £1.2M", 1, 3), ("2023", 1, 5)]))
    assert [c["target_date"] for c in lone] == [None]


def test_every_refused_save_names_the_rejected_field_and_the_reason(api):
    client, db = api
    _upload(client, "audit-1", "fy.pptx", _slide([("FY25 ARR €3M", 1, 2)]))
    arr, = db[decks.CANDIDATES_COLLECTION].docs
    url = f"/api/audits/audit-1/decks/candidates/{arr['id']}"

    def reasons(r):
        detail = r.json()["detail"]
        return [f"{d['loc'][-1]}: {d['msg']}" for d in detail] if isinstance(detail, list) else detail

    # a claim changed to a count type with no unit and a currency still saves; the form sends unit "count" and no currency
    assert client.put(url, json={"claim_type": "users", "unit": "count", "currency": None}).status_code == 200
    # a refusal by the field validators names the field (the screen shows "unit: ...", not "Could not save")
    r = client.put(url, json={"unit": ""})
    assert r.status_code == 422 and reasons(r) == ["unit: String should have at least 1 character"]
    r = client.put(url, json={"currency": "EURO"})
    assert r.status_code == 422 and reasons(r)[0].startswith("currency: ")
    r = client.put(url, json={"target_date": "June"})
    assert r.status_code == 422 and reasons(r)[0].startswith("target_date: ")
    # a refusal by the route's own rules names the field in the same form
    r = client.put(url, json={"status": "approved", "value": 1})
    assert r.status_code == 400 and reasons(r).startswith("status: ")
    r = client.put(url, json={"by_period": [{"value": 1}]})
    assert r.status_code == 400 and reasons(r).startswith("by_period: ")
    other = {k: v for k, v in arr.items() if k != "parsed"} | {"id": "u1", "claim_type": "unknown", "status": "pending"}
    db[decks.CANDIDATES_COLLECTION].docs.append(other)
    r = client.put("/api/audits/audit-1/decks/candidates/u1", json={"status": "approved"})
    assert r.status_code == 400 and reasons(r).startswith("claim_type: ")


def test_zero2hero_page_17_bars_take_the_year_under_each_bar_and_are_never_compared_with_page_19_revenue():
    file = "05-zero2hero.pdf"
    deck = parser.parse_deck((DECKS / file).read_bytes(), file)
    found = claims.detect_candidates(deck["blocks"], file)
    bars = sorted((c["value"], c["target_date"], c["date_from"], c["period_basis"]) for c in found if _page(c) == 17 and c["claim_type"] == "revenue")
    assert bars == [(278085, "2021", "2021", "per year"), (415107, "2022", "2022", "per year"), (550508, "2023", "2023", "per year")]
    flagged = {c["value"]: c["inconsistent_dates"] for c in found if _page(c) == 17 and c["claim_type"] == "revenue"}
    assert flagged == {278085: [], 415107: [], 550508: []}, "Turnover (page 17) is never compared with Revenue (page 19)"
    assert {c["value"]: claims.deck_label(c) for c in found if _page(c) == 17 and c["claim_type"] == "revenue"} == \
        {278085: "Turnover", 415107: "Turnover", 550508: "Turnover"}
    assert all("no date" not in claims.confidence(c, found)["failed"] for c in found if _page(c) == 17 and c["claim_type"] == "revenue")


def test_an_explicit_null_clears_the_as_of_month_and_an_absent_field_leaves_it(api):
    client, db = api
    assert client.put("/api/audits/audit-1", json={"as_of_month": "2026-06-30"}).json()["as_of_month"] == "2026-06-30"
    assert client.put("/api/audits/audit-1", json={"fiscal_year_end": 3}).json()["as_of_month"] == "2026-06-30"
    assert client.put("/api/audits/audit-1", json={"as_of_month": None}).json()["as_of_month"] is None
    client.put("/api/audits/audit-1", json={"target_date": "2028-12-31"})
    assert client.put("/api/audits/audit-1", json={"target_date": None}).json()["target_date"] == "2028-12-31", "never cleared"


def _revenue_figure(snippet, value, page, label_from=None, date="2023"):
    return {"claim_type": "revenue", "value": value, "value_high": None, "unit": None, "currency": "GBP", "snippet": snippet,
            "label_from": label_from, "target_date": date, "period_start": f"{date}-01-01", "period_end": f"{date}-12-31",
            "sources": [{"file": "d.pdf", "page": page}], "_stated": True}


def test_turnover_is_compared_with_turnover_and_revenue_with_revenue_only():
    def flagged(*figures):
        found = [dict(f) for f in figures]
        claims._flag_inconsistencies(found)
        return [c["inconsistent_dates"] for c in found]
    turn1, turn2 = _revenue_figure("550,508", 550508, 17, "Turnover(£/year)"), _revenue_figure("GMV 600,000", 600000, 18)
    rev1, rev2 = _revenue_figure("Revenue 150,000", 150000, 19), _revenue_figure("Revenue 160,000", 160000, 20)
    assert flagged(turn1, rev1) == [[], []], "turnover against revenue is never an inconsistency"
    assert flagged(turn1, turn2) == [["2023"], ["2023"]]
    assert flagged(rev1, rev2) == [["2023"], ["2023"]]
    found = [dict(f) for f in (turn1, turn2)]
    claims._flag_inconsistencies(found)
    pair = found[0]["inconsistencies"][0]
    assert (pair["this"]["label"], pair["other"]["label"]) == ("Turnover", "Turnover")
    assert claims.deck_label(_revenue_figure("ARR 100", 100, 3, "Turnover")) == "Revenue", "ARR is never turnover"


# ---------------------------------------------------------------------------
# Use of funds (zero2hero p22 "Investor Proposition", £250K raise, pie chart "Use of Funds")
# ---------------------------------------------------------------------------
P22_TEXT = [("Metaverse Development & Enhancement", 40), ("Content Acquisition & Development", 20),
            ("Marketing & Brand Awareness", 25), ("Strategic Partnerships & Collaborations", 10),
            ("Operational Expenses & Talent Acquisition", 5)]
KEYWORDED = ["Customer acquisition", "Team hiring", "Product launch", "Cost control", "Win rate"]     # each line names a type
P22_CHART = [42, 21, 26, 11]


def _funds_slide(text=P22_TEXT, chart=P22_CHART, head=("£250K",)):
    """One slide: a raise figure, the use of funds as text lines ("Category (40%)") and the pie's data labels as boxes,
    under a heading that gives them a type to borrow (the real p22's legend does)."""
    return [*head, *[f"{name} ({pct}%)" for name, pct in text], *(["Customer acquisition"] if chart else []),
            *[f"{v}%" for v in chart]]


def _funds_candidates(slides):
    deck = parser.parse_deck(_pptx(slides), "deck.pptx")
    return claims.detect_candidates(deck["blocks"], "deck.pptx")


def test_zero2hero_page_22_percentages_are_use_of_funds_and_never_sales():
    """The text keyword "Acquisition" typed 40/20/25/10/5 and the pie's 42/26 Sales, and the pie's 21/11 Unknown."""
    file = "05-zero2hero.pdf"
    deck = parser.parse_deck((DECKS / file).read_bytes(), file)
    page = [c for c in claims.detect_candidates(deck["blocks"], file) if _page(c) == 22 and c["unit"] == "%"]
    assert sorted(c["value"] for c in page) == [5, 10, 11, 20, 21, 25, 26, 40, 42]
    assert {c["claim_type"] for c in page} == {"use_of_funds"}
    funds = {c["value"]: c for c in page}
    assert funds[40]["inconsistencies"][0]["other"]["value"] == 42, "the text 40 is paired with the chart's 42"
    assert funds[5]["inconsistencies"] == [], "the dropped category has no chart value"


def test_a_raise_amount_with_percentages_summing_to_about_100_is_use_of_funds_without_a_cue_word():
    for pcts in ([40, 20, 25, 10, 5], [50, 30, 20], [48, 30, 20], [52, 30, 22]):       # 100, 100, 98, 104
        found = _funds_candidates([_funds_slide([(f"Sales and acquisition {n}", p) for n, p in enumerate(pcts)], [])])
        assert found and {c["claim_type"] for c in found if c["unit"] == "%"} == {"use_of_funds"}, pcts
    assert not any(c["claim_type"] in ("sales", "revenue") for c in found if c["unit"] == "%")


@pytest.mark.parametrize("pcts, head", [([50, 30, 10], "£250K"), ([50, 30, 26], "£250K"), ([50, 30, 20], "ARR is growing")])
def test_percentages_outside_95_to_105_or_without_a_money_figure_are_not_use_of_funds(pcts, head):
    found = _funds_candidates([[head, *[f"Customer acquisition {n} ({p}%)" for n, p in enumerate(pcts)]]])
    assert found and "use_of_funds" not in {c["claim_type"] for c in found}


@pytest.mark.parametrize("title", ["Use of funds", "Use of proceeds", "The Ask", "Investor Proposition", "Raise"])
def test_a_funds_cue_in_the_title_makes_the_percentages_use_of_funds_whatever_they_sum_to(title):
    slide = [title, "Customer acquisition (30%)", "Product (15%)"]
    deck = parser.parse_deck(_pptx([slide]), "deck.pptx")
    for b in deck["blocks"]:
        b["title"] = b["text"] == title                    # the builder has no title placeholder: mark the first box
    found = claims.detect_candidates(deck["blocks"], "deck.pptx")
    assert {(c["value"], c["claim_type"]) for c in found if c["unit"] == "%"} == {(30, "use_of_funds"), (15, "use_of_funds")}


def test_a_page_without_a_raise_amount_or_cue_keeps_its_types():
    found = _funds_candidates([["Win rate (60%)", "Customers (40%)"]])
    assert {c["claim_type"] for c in found} == {"sales", "customers"}


def test_text_and_chart_percentages_that_differ_are_tagged_on_each_figure_with_the_dropped_category():
    """The p22 shape (the real page is tested above): the chart drops the smallest category and rescales the rest to 100."""
    text = list(zip(KEYWORDED, (50, 30, 20)))
    found = [c for c in _funds_candidates([_funds_slide(text, [63, 38])]) if c["unit"] == "%"]
    by_value = {c["value"]: c for c in found}
    assert sorted(by_value) == [20, 30, 38, 50, 63]
    note = "chart excludes Product launch, rescaled"
    for t_value, c_value in ((50, 63), (30, 38)):
        mine, theirs = by_value[t_value]["inconsistencies"][0], by_value[c_value]["inconsistencies"][0]
        assert (mine["this"]["value"], mine["this"]["label"], mine["other"]["value"], mine["other"]["label"]) == \
            (t_value, "text", c_value, "chart")
        assert (theirs["this"]["value"], theirs["this"]["label"], theirs["other"]["value"], theirs["other"]["label"]) == \
            (c_value, "chart", t_value, "text")
        assert mine["note"] == theirs["note"] == note
    assert by_value[20]["inconsistencies"] == []
    assert all(c["claim_type"] == "use_of_funds" for c in found), "a deck inconsistency is never Contradicted"


def test_chart_values_that_are_not_the_text_rescaled_are_tagged_without_a_rescale_note_and_equal_values_are_not_tagged():
    text = list(zip(KEYWORDED, (60, 40)))
    other = [c for c in _funds_candidates([_funds_slide(text, [65, 35])]) if c["unit"] == "%"]
    assert len(other) == 4 and all(c["inconsistencies"] and c["claim_type"] == "use_of_funds" for c in other)
    assert all("note" not in p for c in other for p in c["inconsistencies"])
    same = [c for c in _funds_candidates([_funds_slide(list(zip(KEYWORDED, (60, 40))), [40, 60])]) if c["unit"] == "%"]
    assert same and all(c["inconsistencies"] == [] for c in same)
    unpaired = [c for c in _funds_candidates([_funds_slide(text, [60, 30, 10])]) if c["unit"] == "%"]
    assert unpaired and all(c["inconsistencies"] == [] for c in unpaired), "no pairing is guessed"


def test_rates_on_a_slide_with_a_money_figure_are_not_a_use_of_funds_even_when_they_sum_to_about_100():
    """04-clevergig.docx page 6: 15% MoM growth, 50% and 30% of leads sum to 95 beside "€15K in MRR"."""
    found = _funds_candidates([["€15K in MRR, 15% MoM growth over last 15 months", "€260 MRR per client, MRR expands by 50%",
                                "30% of our leads come via worker referrals"]])
    assert [c for c in found if c["unit"] == "%"] and "use_of_funds" not in {c["claim_type"] for c in found}
