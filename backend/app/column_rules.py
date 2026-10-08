"""Column rules: reading a sheet, finding its header, detecting its type and mapping its columns by rule
(docs/specs/chat-upload.md sections 3 and 4.1). Pure Python: no database, no model, no network.

The model never decides here. A column is mapped when its header and its values agree (confidence at or
above MAPPING_THRESHOLD, no tie); anything else is shown to the analyst, with or without the model's
proposal (CLAUDE.md rule 18: model output is never Verified on its own).
"""
import csv
import io
import re
import warnings
from dataclasses import dataclass
from datetime import date, datetime
from typing import Dict, List, Optional, Tuple

import pandas as pd

# ---------------------------------------------------------------------------
# Dataset field definitions (normalized names + fuzzy-match aliases)
# ---------------------------------------------------------------------------
FIELD_DEFS = {
    "revenue": {
        "required": {
            "customer_id": ["customer", "customer id", "account", "client", "cust"],
            "invoice_date": ["invoice date", "date", "billing date", "posted"],
            "amount": ["amount", "value", "revenue", "total", "arr", "mrr"],
            "currency": ["currency", "ccy", "curr"],
        },
        "optional": {
            "service_start": ["service start", "start date", "period start", "term start"],
            "service_end": ["service end", "end date", "period end", "term end"],
            "segment": ["segment", "tier", "size", "band"],
            "revenue_type": ["revenue type", "type", "recurring", "rec/one-off"],
        },
        "dates": ["invoice_date", "service_start", "service_end"],
        "numeric": ["amount"],
    },
    "crm": {
        "required": {
            "deal_id": ["deal id", "opportunity id", "deal", "id"],
            "created_date": ["created", "create date", "created date", "open date"],
            "close_date": ["close date", "closed", "won date", "close"],
            "stage": ["stage", "status", "outcome"],
            "amount": ["amount", "value", "deal value", "acv"],
        },
        "optional": {
            "segment": ["segment", "tier", "size"],
            "founder_involved": ["founder", "founder involved", "founder-led", "exec involved"],
        },
        "dates": ["created_date", "close_date"],
        "numeric": ["amount"],
    },
    "pnl": {
        "required": {
            "month": ["month", "period", "date", "fiscal month"],
            "sm_expense": ["sales & marketing", "s&m", "sales and marketing", "marketing expense", "sm expense"],
            "revenue": ["revenue", "total revenue", "sales", "turnover"],
            "cost_of_revenue": ["cost of revenue", "cogs", "cost of sales", "cost"],
        },
        "optional": {},
        "dates": ["month"],
        "numeric": ["sm_expense", "revenue", "cost_of_revenue"],
    },
}

MAPPING_THRESHOLD = 80          # confidence at or above this, with no tie, is accepted without a click (Q3)
HEADER_SCAN_ROWS = 10           # the header row is looked for in the first 10 sheet rows
DETECTION_MIN_SHARE = 0.5       # a type is detected when the rules map at least half of its required fields

_UNNAMED = re.compile(r"^Unnamed: \d+$")
_DEDUPED = re.compile(r"^(.*)\.\d+$")
_YEAR = re.compile(r"^(?:19|20)\d{2}(?:\.0)?$")
_PURE_NUMBER = re.compile(r"^[-+]?\d+(?:\.\d+)?$")
# 03/04/2024, 3.4.24, 03-04-2024 (+ optional time): day and month in either order (server.parse_date_column).
_DM_DATE = re.compile(r"^\s*(\d{1,2})[/.\-](\d{1,2})[/.\-](\d{4}|\d{2})(?:[ T].*)?$")


def _score(field, alias, lc):
    if lc == alias:
        return 100
    if lc.startswith(alias) or alias.startswith(lc):
        return 80 if min(len(lc), len(alias)) >= 4 else 30
    if alias in lc:  # alias is a substring of the column header
        return 60
    if lc in alias and len(lc) >= 4:  # column header is a substring of the alias
        return 40
    return 0


def _all_fields(dtype: str) -> Dict[str, List[str]]:
    defs = FIELD_DEFS[dtype]
    return {**defs["required"], **defs["optional"]}


def field_score(dtype: str, field: str, header: str) -> int:
    """The best score of a header against one field's name and aliases (100, 80, 60, 40, 30 or 0)."""
    lc = (header or "").lower().strip()
    if not lc:
        return 0
    aliases = [field.replace("_", " ")] + _all_fields(dtype)[field]
    return max(_score(field, alias, lc) for alias in aliases)


def suggest_mapping(dtype: str, columns: list, headers: Optional[list] = None) -> dict:
    """{field: column} by the greedy order of the aliases: best score first. `headers` are the header cell
    texts when they differ from the column names (a pandas-style "Customer ID.1" is scored as "Customer ID")."""
    all_fields = _all_fields(dtype)
    texts = headers if headers is not None else deduped_texts(columns)
    candidates = []
    for field in all_fields:
        for col, text in zip(columns, texts):
            s = field_score(dtype, field, text)
            if s:
                candidates.append((s, field, col))
    candidates.sort(reverse=True, key=lambda x: x[0])
    mapping = {f: None for f in all_fields}
    used = set()
    for s, field, col in candidates:
        if mapping[field] is None and col not in used:
            mapping[field] = col
            used.add(col)
    return mapping


def deduped_texts(columns: list) -> list:
    """The header text of each column when only the names are known: a "Name.1" whose "Name" is also a column
    is a de-duplicated header and is read as "Name"."""
    names = set(columns)
    out = []
    for c in columns:
        m = _DEDUPED.match(str(c))
        out.append(m.group(1) if m and m.group(1) in names else str(c))
    return out


# ---------------------------------------------------------------------------
# Reading a sheet (section 3)
# ---------------------------------------------------------------------------
@dataclass
class Sheet:
    sheet: str
    header_row: int                 # the sheet row number of the header (1 = the first row)
    columns: List[str]              # unique names; a blank header is "Unnamed: n", a repeat "Name.1"
    headers: List[str]              # the header cell text of each column, as written ("" when blank)
    frame: pd.DataFrame             # the data rows below the header, blank lines dropped
    row_numbers: List[int]          # the sheet row number of each frame row


class SheetError(Exception):
    """The file could not be read as a sheet."""


def _blank(v) -> bool:
    return v is None or (isinstance(v, float) and v != v) or (isinstance(v, str) and not v.strip())


def _grid(content: bytes, filename: str) -> Tuple[List[List], str, bool]:
    """(rows as cell lists, sheet name, is_csv). Index i of the grid is sheet row i + 1."""
    name = (filename or "").lower()
    try:
        if name.endswith(".csv"):
            try:
                text = content.decode("utf-8-sig")
            except UnicodeDecodeError:
                text = content.decode("latin-1")
            grid = [[None if _blank(c) else c for c in row] for row in csv.reader(io.StringIO(text))]
            return grid, "CSV", True
        if name.endswith(".xlsx"):
            from openpyxl import load_workbook
            wb = load_workbook(io.BytesIO(content), data_only=True)
            ws = wb[wb.sheetnames[0]]
            grid = [[None if _blank(c) else c for c in row] for row in ws.iter_rows(values_only=True)]
            return grid, ws.title, False
        xls = pd.ExcelFile(io.BytesIO(content))
        sheet = xls.sheet_names[0]
        frame = xls.parse(sheet, header=None, dtype=object)
        return [[None if _blank(c) else c for c in row] for row in frame.itertuples(index=False)], sheet, False
    except SheetError:
        raise
    except Exception as exc:
        raise SheetError(type(exc).__name__)


def _filled(row: List) -> int:
    return sum(1 for c in row if not _blank(c))


def _has_number(row: List) -> bool:
    """A number other than a year: a header row holds none."""
    for c in row:
        if _blank(c) or isinstance(c, (bool, datetime, date)):
            continue
        if isinstance(c, (int, float)) and not _YEAR.match(str(c)):
            return True
        if isinstance(c, str) and _PURE_NUMBER.match(c.strip()) and not _YEAR.match(c.strip()):
            return True
    return False


def find_header_row(grid: List[List]) -> int:
    """Index in the grid of the header row: among the first 10 rows, the first that fills at least half of the
    widest row's cells and holds no number other than a year. The first row when none qualifies."""
    top = grid[:HEADER_SCAN_ROWS]
    widest = max((_filled(r) for r in top), default=0)
    for i, row in enumerate(top):
        if _filled(row) and 2 * _filled(row) >= widest and not _has_number(row):
            return i
    return 0


def _header_text(v) -> str:
    if _blank(v):
        return ""
    if isinstance(v, float) and v.is_integer():
        return str(int(v))
    return str(v).strip()


def read_sheet(content: bytes, filename: str) -> Sheet:
    """The first sheet of a file as a header, its columns and its data rows with their sheet row numbers.

    Read without a header, short rows padded, blank lines counted: a one-cell title row cannot break a CSV.
    Rows above the header (titles, blanks) are dropped and never stored."""
    grid, sheet, is_csv = _grid(content, filename)
    while grid and _filled(grid[-1]) == 0:
        grid.pop()
    if not grid:
        raise SheetError("empty")
    h = find_header_row(grid)
    body = [(i + 1, row) for i, row in enumerate(grid) if i > h and _filled(row)]
    width = max([len(grid[h])] + [len(r) for _, r in body])
    pad = lambda r: list(r) + [None] * (width - len(r))  # noqa: E731
    headers = [_header_text(c) for c in pad(grid[h])]
    while width and not headers[width - 1] and all(_blank(pad(r)[width - 1]) for _, r in body):
        width -= 1
        headers.pop()
    columns, seen = [], {}
    for i, text in enumerate(headers):
        base = text or f"Unnamed: {i}"
        n = seen.get(base, 0)
        seen[base] = n + 1
        columns.append(base if n == 0 else f"{base}.{n}")
    frame = pd.DataFrame([pad(r)[:width] for _, r in body], columns=columns, dtype=object)
    if is_csv:                                   # CSV cells are text: a column that is all numbers is numeric
        for c in columns:
            col = frame[c]
            present = col.notna()
            if present.any():
                conv = pd.to_numeric(col, errors="coerce")
                if conv.notna().sum() == present.sum():
                    frame[c] = conv
    return Sheet(sheet, h + 1, columns, headers, frame, [n for n, _ in body])


# ---------------------------------------------------------------------------
# Value fit (section 4.1)
# ---------------------------------------------------------------------------
KIND_WORDS = {"date": "dates", "numeric": "numbers", "currency": "currency codes"}


def field_kind(dtype: str, field: str) -> str:
    defs = FIELD_DEFS[dtype]
    if field in defs["dates"]:
        return "date"
    if field in defs["numeric"]:
        return "numeric"
    return "currency" if field == "currency" else "text"


def _non_blank(values) -> list:
    return [v for v in values if not _blank(v)]


def _reads_as_date(values: list) -> int:
    ok, text = 0, []
    for v in values:
        if isinstance(v, (datetime, date, pd.Timestamp)):
            ok += 1
        elif isinstance(v, str):
            s = v.strip()
            m = _DM_DATE.match(s)
            if m:                                     # a day/month order it cannot settle still counts as a date
                a, b = int(m.group(1)), int(m.group(2))
                ok += 1 if 1 <= a <= 31 and 1 <= b <= 31 and (a <= 12 or b <= 12) else 0
            elif not _PURE_NUMBER.match(s):
                text.append(s)
    if text:
        with warnings.catch_warnings():
            warnings.simplefilter("ignore")
            ok += int(pd.to_datetime(pd.Series(text), errors="coerce").notna().sum())
    return ok


def _reads_as_number(values: list) -> int:
    ok, text = 0, []
    for v in values:
        if isinstance(v, bool):
            continue
        if isinstance(v, (int, float)):
            ok += 1 if v == v else 0
        elif isinstance(v, str):
            text.append(v)
    if text:
        ok += int(pd.to_numeric(pd.Series(text), errors="coerce").notna().sum())
    return ok


def value_fit(kind: str, values) -> Tuple[float, int, int]:
    """(fit, readable, total): the share of the column's non-blank values the engine can read as the field's
    kind. A text field accepts any non-blank value; a column with no non-blank value has a fit of 0."""
    present = _non_blank(values)
    if not present:
        return 0.0, 0, 0
    if kind == "date":
        ok = _reads_as_date(present)
    elif kind == "numeric":
        ok = _reads_as_number(present)
    elif kind == "currency":
        ok = sum(1 for v in present if isinstance(v, str) and re.fullmatch(r"[A-Za-z]{3}", v.strip()))
    else:
        ok = len(present)
    return ok / len(present), ok, len(present)


def fit_note(kind: str, ok: int, total: int) -> Optional[str]:
    """S9: "0 of 20 values are numbers", shown when the fit lowers the confidence."""
    if kind == "text" or ok == total:
        return None
    if total == 0:
        return "no values"
    return f"{ok} of {total} values are {KIND_WORDS[kind]}"


# ---------------------------------------------------------------------------
# Mapping rule, step 1 (section 4.1)
# ---------------------------------------------------------------------------
@dataclass
class Proposal:
    column: str
    field: Optional[str]
    score: int
    confidence: Optional[int]
    note: Optional[str]
    state: str                      # "auto" | "unsure" | "undecided" | "unused"
    tie: bool = False
    tie_fields: tuple = ()          # the fields this column is tied for


def analyse(dtype: str, sheet: Sheet) -> Tuple[List[Proposal], List[str]]:
    """(one Proposal per column, the fields still open). The rules propose one column per field, in the greedy
    order of suggest_mapping. Auto: confidence at or above the threshold and no tie. Unsure: below it, or tied
    (two columns share a field's best score, or one column shares its best score between two fields). Undecided:
    no proposal while a field is still open. Unused: no proposal and no open field."""
    fields = _all_fields(dtype)
    mapping = suggest_mapping(dtype, sheet.columns, sheet.headers)
    by_col = {c: f for f, c in mapping.items() if c}
    scores = {(f, c): field_score(dtype, f, t) for f in fields for c, t in zip(sheet.columns, sheet.headers)}
    tied, tie_fields = set(), {c: [] for c in sheet.columns}
    for f in fields:                                 # two columns share a field's best score
        best = max((scores[(f, c)] for c in sheet.columns), default=0)
        same = [c for c in sheet.columns if best and scores[(f, c)] == best]
        if len(same) > 1:
            tied |= set(same)
            for c in same:
                tie_fields[c].append(f)
    for c in sheet.columns:                          # one column shares its best score between two fields
        best = max((scores[(f, c)] for f in fields), default=0)
        same = [f for f in fields if best and scores[(f, c)] == best]
        if len(same) > 1:
            tied.add(c)
            tie_fields[c] += [f for f in same if f not in tie_fields[c]]
    out = []
    for c in sheet.columns:
        field = by_col.get(c)
        if field is None:
            out.append(Proposal(c, None, 0, None, None, "undecided", c in tied, tuple(tie_fields[c])))
            continue
        fit, ok, total = value_fit(field_kind(dtype, field), sheet.frame[c].tolist())
        conf = round(scores[(field, c)] * fit)
        note = fit_note(field_kind(dtype, field), ok, total)
        accepted = conf >= MAPPING_THRESHOLD and c not in tied
        out.append(Proposal(c, field, scores[(field, c)], conf, note, "auto" if accepted else "unsure", c in tied,
                            tuple(tie_fields[c])))
    held = {p.field for p in out if p.state == "auto"}
    open_fields = [f for f in fields if f not in held]
    for p in out:
        if p.state == "undecided" and not open_fields:
            p.state = "unused"
    return out, open_fields


def customer_columns(sheet: Sheet, proposals: List[Proposal]) -> List[str]:
    """Every column whose header matches a customer alias at any score, and every column tied for
    customer_id: these send a profile only, even when their values are numbers (section 4.2)."""
    out = [c for c, t in zip(sheet.columns, sheet.headers) if field_score("revenue", "customer_id", t) > 0]
    out += [p.column for p in proposals if "customer_id" in p.tie_fields]
    return list(dict.fromkeys(out))


# ---------------------------------------------------------------------------
# Type detection (section 3)
# ---------------------------------------------------------------------------
def detect_type(sheet: Sheet) -> Tuple[Optional[str], Dict[str, float]]:
    """(type, share per type). For each type, the share of its required fields that the rules map (score above
    0). The type with the unique highest share wins when that share is at least 50%; otherwise None."""
    shares = {}
    for dtype, defs in FIELD_DEFS.items():
        mapping = suggest_mapping(dtype, sheet.columns, sheet.headers)
        shares[dtype] = sum(1 for f in defs["required"] if mapping.get(f)) / len(defs["required"])
    best = max(shares.values())
    winners = [t for t, s in shares.items() if s == best]
    return (winners[0] if len(winners) == 1 and best >= DETECTION_MIN_SHARE else None), shares


def months_found(dtype: str, sheet: Sheet, mapping: Dict[str, str]) -> Optional[Dict]:
    """The distinct months of the date column (invoice date for revenue, close date for CRM, month for the P&L):
    {"count", "first", "last"} as YYYY-MM, or None before that column is mapped or when none reads."""
    field = {"revenue": "invoice_date", "crm": "close_date", "pnl": "month"}[dtype]
    col = mapping.get(field)
    if not col or col not in sheet.frame.columns:
        return None
    return months_of(sheet.frame[col].tolist())


def months_of(values: list) -> Optional[Dict]:
    """The distinct months of a date column, as the engine reads it: {"count", "first", "last"} (YYYY-MM), or None
    when no value reads. A day/month order the data cannot settle is not guessed: those values are left out."""
    present = _non_blank(values)
    parsed, text = [], []
    for v in present:
        if isinstance(v, (datetime, date, pd.Timestamp)):
            parsed.append(pd.Timestamp(v))
        elif isinstance(v, str):
            s = v.strip()
            m = _DM_DATE.match(s)
            if m:
                a, b, y = int(m.group(1)), int(m.group(2)), m.group(3)
                year = int(y) + (2000 if len(y) == 2 else 0)
                if a > 12 >= b:
                    parsed.append(pd.Timestamp(year, b, a))
                elif b > 12 >= a:
                    parsed.append(pd.Timestamp(year, a, b))
            elif not _PURE_NUMBER.match(s):
                text.append(s)
    if text:
        with warnings.catch_warnings():
            warnings.simplefilter("ignore")
            parsed += [t for t in pd.to_datetime(pd.Series(text), errors="coerce") if not pd.isna(t)]
    months = sorted({t.strftime("%Y-%m") for t in parsed})
    return {"count": len(months), "first": months[0], "last": months[-1]} if months else None


# ---------------------------------------------------------------------------
# Column states: what the mapping table shows and what the analyst still owes (sections 4.1-4.3, 5)
# ---------------------------------------------------------------------------
# The fixed list of reasons for a correction (S11). Anything else is refused.
REASONS = ("header_misleading", "other_column_right", "values_do_not_fit", "wrong_kind_of_date", "not_needed", "other")
PENDING = ("unsure", "ai", "needs")           # states that wait for a click
MAPPED = ("auto", "unsure", "ai", "confirmed", "corrected")


class DecisionError(Exception):
    """A decision that cannot be applied; `code` is a closed word, `column` the column it names (if any)."""

    def __init__(self, code: str, column: Optional[str] = None):
        super().__init__(code)
        self.code, self.column = code, column


def new_state(column: str, field: Optional[str], source: Optional[str], state: str, confidence: Optional[int] = None,
              note: Optional[str] = None) -> Dict:
    return {"column": column, "field": field, "source": source, "state": state, "confidence": confidence,
            "note": note, "decision": None, "reason": None}


def states_from_rules(proposals: List[Proposal]) -> List[Dict]:
    out = []
    for p in proposals:
        if p.state in ("auto", "unsure"):
            out.append(new_state(p.column, p.field, "rules", p.state, p.confidence, p.note))
        elif p.state == "undecided":
            out.append(new_state(p.column, None, "needs", "needs"))
        else:
            out.append(new_state(p.column, None, None, "unused"))
    return out


def to_send(states: List[Dict]) -> List[str]:
    """The columns the model is asked about: those the rules left unsure or undecided."""
    return [s["column"] for s in states if s["state"] in ("unsure", "needs")]


def apply_model(states: List[Dict], sent: List[str], proposed: Dict[str, str]) -> List[Dict]:
    """The model's proposals ({field: column}) for the columns that were sent. A proposal becomes an AI suggestion
    (confirm or correct); a sent column the reply leaves out is not used. Decided columns are never touched."""
    by_col = {c: f for f, c in proposed.items()}
    out = []
    for s in states:
        if s["column"] in sent:
            f = by_col.get(s["column"])
            s = new_state(s["column"], f, "ai", "ai") if f else new_state(s["column"], None, None, "unused")
        out.append(s)
    return out


def apply_model_failure(states: List[Dict], sent: List[str]) -> List[Dict]:
    """The model could not be read: every column that was sent needs the analyst's decision."""
    return [new_state(s["column"], None, "needs", "needs") if s["column"] in sent else s for s in states]


def states_from_saved(dtype: str, sheet: Sheet, saved: List[Dict], scale: bool) -> List[Dict]:
    """A saved mapping applied to a sheet. The same bytes: every mapped column is accepted. The same headers on
    other data (`scale`): each mapped column counts as confidence 100 scaled by its value fit on the new data."""
    by_col = {c["column"]: c for c in saved}
    out = []
    for col in sheet.columns:
        entry = by_col.get(col) or {}
        field = entry.get("field")
        if not field or field not in _all_fields(dtype):
            out.append(new_state(col, None, None, "unused"))
            continue
        conf, note = 100, None
        if scale:
            kind = field_kind(dtype, field)
            fit, ok, total = value_fit(kind, sheet.frame[col].tolist())
            conf, note = round(100 * fit), fit_note(kind, ok, total)
        out.append(new_state(col, field, "saved", "auto" if conf >= MAPPING_THRESHOLD else "unsure", conf, note))
    return out


def mapping_of(dtype: str, states: List[Dict]) -> Dict[str, Optional[str]]:
    """{field: column} of every mapped column, pending proposals included (they are shown, not yet confirmed)."""
    mapping = {f: None for f in _all_fields(dtype)}
    for s in states:
        if s["state"] in MAPPED and s["field"]:
            mapping[s["field"]] = s["column"]
    return mapping


def pending_count(states: List[Dict]) -> int:
    return sum(1 for s in states if s["state"] in PENDING)


def missing_required(dtype: str, states: List[Dict]) -> List[str]:
    mapping = mapping_of(dtype, states)
    return [f for f in FIELD_DEFS[dtype]["required"] if not mapping.get(f)]


def apply_decisions(dtype: str, states: List[Dict], decisions: List[Dict]) -> List[Dict]:
    """The analyst's decisions, one by one. Confirm accepts the proposal (a row that needs a decision takes the
    field chosen, or "Not used"). Correct takes a field or "Not used" and a reason from the fixed list. A field
    another column holds is refused and names that column. Raises DecisionError."""
    fields = _all_fields(dtype)
    states = [dict(s) for s in states]
    for d in decisions:
        state = next((s for s in states if s["column"] == d.get("column")), None)
        if state is None:
            raise DecisionError("unknown_column")
        action, field = d.get("action"), (d.get("field") or None)
        if action not in ("confirm", "correct"):
            raise DecisionError("bad_action", state["column"])
        if field is not None and field not in fields:
            raise DecisionError("unknown_field", state["column"])
        if action == "correct" and d.get("reason") not in REASONS:
            raise DecisionError("bad_reason", state["column"])
        if action == "confirm":
            if state["state"] in ("auto", "confirmed", "corrected", "unused") and field is None:
                continue                                       # nothing to accept
            if state["state"] != "needs":
                field = state["field"]
        if field:
            holder = next((s for s in states if s["field"] == field and s["column"] != state["column"]
                           and s["state"] in MAPPED), None)
            if holder:
                raise DecisionError("field_held", holder["column"])
        was_needs = state["state"] == "needs"
        state["field"] = field
        state["state"] = ("confirmed" if action == "confirm" else "corrected") if field else "unused"
        state["decision"] = action
        state["reason"] = d.get("reason") if action == "correct" else None
        if was_needs:
            state["source"] = "decision"
    return states
