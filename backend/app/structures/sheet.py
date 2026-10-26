"""Sheet structure reading (docs/specs/chat-upload.md section 26). Pure Python: no database, no model, no network.

A P&L the rules cannot map by columns is read as a statement: line labels down the rows, periods across the columns. Python reads
it from the grid (the header row, the period columns and their tags, the label column, each line's indent and numbers) and labels
every line by its own rules. The skeleton the model reads holds every text cell, redacted, and never a figure: the numbers are
"num" in each row's line (redact.sheet_text). The model's reply labels lines only (role, parent, meaning); Python keeps what it
can rebuild, applies the rules that win over the reply, checks every total against its sub-lines and builds what the map feeds: a
monthly P&L for the engine, company-reported KPI series, the actual years of the yearly-basis metrics and the forecast claims.
Every figure is read from its cell (rule 18).
"""
import math
import re
from datetime import date, datetime
from typing import Dict, List, Optional, Tuple

from openpyxl.utils import get_column_letter

from .. import column_rules as cr
from .. import forecast_claims
from ..llm.schemas import SHEET_BASES, SHEET_MEANINGS, SHEET_ROLES
from . import redact

MAX_ROWS = redact.SHEET_MAX_ROWS
TOLERANCE = 0.005               # a total equals the sum of its sub-lines within 0.5% per period
LONG_CELL = 200                 # a longer cell is prose: never sent
NOTE_LENGTH = 60                # a label over 60 characters with no number is a note
INHERITED = 60                  # the confidence of a meaning a sub-line takes from its parent
PNL_FIELDS = ("revenue", "cost_of_revenue", "sm_expense")       # the monthly P&L the engine reads
KPI_SERIES = ("arr", "mrr", "new_arr", "churn_rate")            # stored as company-reported series, never engine input
SM_CLAIM_LABEL = "S&M (sum of Sales and Marketing)"
NOT_PLAN_LINES = ("arpa", "contracted_revenue", "acv", "tcv", "new_arr")    # their label may hold "revenue", but they are no plan line
KPI_MEANINGS = ("arr", "mrr", "new_arr", "churned_mrr", "churn_rate", "customers", "arpa", "subscriptions")   # what makes a sheet a KPI sheet (26.14)
COSTS = ("cost_of_revenue", "sm_expense", "opex")               # a cost line written as negatives is read as positive
MONEY = ("revenue", "cost_of_revenue", "sm_expense", "gross_margin", "opex", "ebitda", "arr", "mrr", "new_arr", "churned_mrr",
         "bookings", "acv", "tcv", "cash_balance", "arpa", "contracted_revenue")
YEARLY_FIELDS = ("revenue", "cost_of_revenue", "sm_expense", "gross_margin")     # what the yearly-basis metrics read
TRACK_FIELDS = ("revenue", "cost_of_revenue", "sm_expense", "gross_margin", "opex", "ebitda")
FAILED_CHECKS = ("total_mismatch", "no_breakdown", "ytd_mismatch", "ltm_mismatch")              # these lines come first in the table
NOT_LINES = ("heading", "note")
# How a line is measured and which lines feed anything (26.15).
BALANCE_MEANINGS = ("arr", "mrr", "customers", "cash_balance")      # their basis is balance unless the label says otherwise
STOCK_MEANINGS = BALANCE_MEANINGS + ("headcount", "subscriptions")  # a balance basis is read for these
FEEDS = PNL_FIELDS + KPI_SERIES + ("gross_margin", "opex", "ebitda")    # the meanings a line can feed (a monthly P&L, a series, a yearly input)
ROW_TAGS = ("forecast", "budget")

# Meaning synonyms (26.10), first match in this order. "words" stand in the label as whole words; "whole" is the whole label.
MEANING_DEFS = (
    ("cost_of_revenue", {"words": ("cost of revenue", "cost of revenues", "cost of sales", "cost of goods sold", "cogs",
                                   "direct costs", "cost of services")}),
    ("sm_expense", {"words": ("sales and marketing", "sales & marketing", "s&m", "marketing", "marketing expenses", "advertising",
                              "sales expenses", "sales costs")}),
    ("gross_margin", {"words": ("gross margin", "gross profit")}),
    ("ebitda", {"words": ("ebitda", "adjusted ebitda")}),
    ("opex", {"words": ("operating expenses", "opex", "overheads", "general and administrative", "general & administrative", "g&a",
                        "research and development", "r&d", "personnel costs", "payroll", "salaries")}),
    ("new_arr", {"words": ("new arr", "net new arr", "new business arr")}),
    ("arpa", {"words": ("arpa", "average revenue per account", "average revenue per customer", "average revenue per user",
                        "average revenue per")}),
    ("contracted_revenue", {"words": ("contracted revenue",)}),
    ("churn_rate", {"words": ("churn rate",)}),
    ("churned_mrr", {"words": ("churned mrr", "lost mrr", "churned arr", "mrr churn")}),
    ("growth_rate", {"words": ("growth", "growth rate", "yoy", "mom")}),
    ("arr", {"words": ("arr", "annual recurring revenue")}),
    ("mrr", {"words": ("mrr", "monthly recurring revenue")}),
    ("bookings", {"words": ("bookings", "booking")}),
    ("acv", {"words": ("acv", "average contract value", "annual contract value")}),
    ("tcv", {"words": ("tcv", "total contract value")}),
    ("revenue", {"words": ("revenue", "revenues", "net sales", "total revenue"), "whole": ("sales",)}),
    ("customers", {"words": ("customers", "customer count", "clients", "subscribers", "logos")}),
    ("subscriptions", {"words": ("subscriptions", "subscription count")}),
    ("headcount", {"words": ("headcount", "head count", "fte", "ftes", "employees")}),
    ("cash_balance", {"words": ("cash balance", "closing cash", "ending cash"), "whole": ("cash",)}),
)
# Labels that name a meaning whatever else stands in them (26.15): "NEW ... (ARR)", "New sales (ACV)", a "% share" line.
_NEW_ARR = re.compile(r"^new\b.*\barr$")
_NEW_ACV = re.compile(r"^new\b.*\bacv$")
_SHARE = re.compile(r"%.*\bshare\b|\bshare\b.*%")
# Basis cues in a label (26.15), first match in this order; a cue wins over the reply.
_BASIS_CUES = ((re.compile(r"\bytd\b|\bcurrent year\b|\byear to date\b"), "ytd"),
               (re.compile(r"\bltm\b|\bttm\b|\blast 12 months\b|\blast twelve months\b|\btrailing (?:12|twelve) months\b"), "ltm"),
               (re.compile(r"\blifetime\b|\btotal\b.*\bachieved\b"), "lifetime"),
               (re.compile(r"\bquarterly\b"), "quarter"))
_TAG_CUES = ((re.compile(r"\bforecast(?:s|ed)?\b"), "forecast"), (re.compile(r"\btargets?\b"), "forecast"),
             (re.compile(r"\bbudget(?:s|ed)?\b"), "budget"))
# "plan" tags a row only as the plan figure (26.15), never a pricing or subscription plan: it starts the label, or is "(plan)" or "plan:".
_PLAN_LABEL = re.compile(r"(?i)^\W*(?:plan(?:ned)?|business plan)\b|\(\s*plan\s*\)|\bplan\s*:")
_PLAN_WORD = re.compile(r"\bplan(?:ned)?\b")
_RATE = re.compile(r"%|\bas % of\b|\b(?:rate|ratio|growth|yoy|mom)\b")
_MARGIN = re.compile(r"\bmargin\b")
_COST_SECTION = re.compile(r"(?i)\b(?:costs?|expenses?|opex|overheads|operating)\b")
_TOTAL = re.compile(r"(?i)^\s*(?:total|sum)\b")
# A result line ends a block of lines a "Total" adds up (no indents): a profit or margin is no cost line above a total.
_RESULT = re.compile(r"(?i)\b(?:gross (?:profit|margin)|ebitda|net (?:profit|income|loss)|operating (?:profit|income|result))\b")
# Tags (26.2): a cell above a period header, or a tag in the header itself.
_TAG_WORDS = ((re.compile(r"(?i)\b(?:actuals?|act)\b"), "actual"),
              (re.compile(r"(?i)\b(?:forecasts?|fcst|estimates?|est|projected|projections?)\b"), "forecast"),
              (re.compile(r"(?i)\b(?:budgets?|bud|plan)\b"), "budget"))
_TAG_LETTERS = {"A": "actual", "F": "forecast", "E": "forecast", "B": "budget"}
_HEADER_TAG = re.compile(r"^(?P<body>.*\d)\s?(?P<tag>[AFBE])$")
_PTD = re.compile(r"(?i)^(?:period to date|ptd|month to date|mtd|current (?:period|month))$")
_YTD = re.compile(r"(?i)^(?:year to date|ytd)$")
_FY_SHORT = re.compile(r"(?i)^FY\s?'?(\d{2})$")
_FY_SPLIT = re.compile(r"(?i)^(?:FY\s?)?((?:19|20)\d{2})\s?/\s?(\d{2}|\d{4})$")


# ---------------------------------------------------------------------------
# Cells
# ---------------------------------------------------------------------------
def _blank(v) -> bool:
    return cr._blank(v)


def _year_number(v) -> bool:
    return isinstance(v, (int, float)) and not isinstance(v, bool) and float(v).is_integer() and 1900 <= v <= 2100


def cell_kind(v, row: int, header_row: int) -> str:
    """"blank", "num" or "text". A number is num; so is text that reads as one. At or above the header row a year is a period
    header, so text; below it a year is a number like any other."""
    if _blank(v):
        return "blank"
    if isinstance(v, (datetime, date, bool)):
        return "text"
    if isinstance(v, (int, float)):
        return "text" if row <= header_row and _year_number(v) else "num"
    text = str(v).strip()
    if row > header_row and redact.is_figure(redact.blank_currency(text).strip() or text):
        return "num"
    return "num" if redact.sheet_figure(text) else "text"


def cell_text(v) -> str:
    """A text cell as the skeleton writes it: a date as its ISO date, a year as its number, text on one line."""
    if isinstance(v, datetime):
        return v.date().isoformat()
    if isinstance(v, date):
        return v.isoformat()
    if _year_number(v):
        return str(int(v))
    return re.sub(r"\s+", " ", str(v)).strip()


def number(v) -> Optional[float]:
    """A number cell's value, read as written: brackets make a negative, a currency sign and thousands commas are dropped, a percent
    is its fraction. None when it is no number."""
    if v is None or isinstance(v, bool):
        return None
    if isinstance(v, (int, float)):
        return float(v) if math.isfinite(v) else None
    if isinstance(v, (datetime, date)):
        return None
    text = str(v).strip()
    negative = text.startswith("(") and text.endswith(")")
    text = re.sub(r"[£$€¥\s]", "", text.strip("()")).replace("−", "-").replace(",", "")
    percent = text.endswith("%")
    try:
        value = float(text.rstrip("%"))
    except ValueError:
        return None
    if not math.isfinite(value):
        return None
    return (-value if negative else value) / (100 if percent else 1)


def _indent(v, row: int, col: int, fmt: Dict) -> int:
    """A text cell's indent level: the Excel indent, or the leading spaces of its text, two per level."""
    found = (fmt.get("indents") or {}).get((row, col), 0)
    if isinstance(v, str):
        spaces = len(v.expandtabs(2)) - len(v.expandtabs(2).lstrip(" "))
        found = max(found, math.ceil(spaces / 2))
    return int(found)


# ---------------------------------------------------------------------------
# Periods
# ---------------------------------------------------------------------------
def _add_months(year: int, month: int, n: int) -> Tuple[int, int]:
    k = year * 12 + (month - 1) + n
    return k // 12, k % 12 + 1


def _ym(ym: Tuple[int, int]) -> str:
    return f"{ym[0]:04d}-{ym[1]:02d}"


def _fiscal(year: int, fye: int, first: int, months: int) -> Tuple[str, str]:
    """(first month, last month) of a part of fiscal year `year` (named by the year it ends in): `first` months after the year's
    first month, `months` long."""
    start = _add_months(year, fye, -11 + first)
    return _ym(start), _ym(_add_months(*start, months - 1))


def period_of(value, fye: int = 12) -> Optional[Dict]:
    """A period header as {"grain", "year", "start", "end", "tag"}, or None. `year` is the fiscal year (the year it ends in) of a
    year, half, quarter or month; `start` and `end` its first and last month (YYYY-MM); `tag` the tag the header itself carries
    ("2025F", "FY24 B", "2023 Actual"). Year to date and period to date have no months."""
    if _blank(value):
        return None
    if isinstance(value, (datetime, date)):
        d = value if isinstance(value, date) else value.date()
        return _month(d.year, d.month, fye, None)
    if isinstance(value, bool):
        return None
    if isinstance(value, (int, float)):
        return _year(int(value), fye, None) if _year_number(value) else None
    text, tag = re.sub(r"\s+", " ", str(value)).strip(), None
    for rx, name in _TAG_WORDS:
        if rx.search(text):
            tag, text = name, re.sub(r"\s+", " ", rx.sub(" ", text)).strip(" -–:()")
    m = _HEADER_TAG.match(text)
    if m and (cr.parse_period(m.group("body").strip()) or _FY_SHORT.match(m.group("body").strip())):
        tag, text = tag or _TAG_LETTERS[m.group("tag")], m.group("body").strip()
    if _PTD.match(text):
        return {"grain": "period_to_date", "year": None, "start": None, "end": None, "tag": tag}
    if _YTD.match(text):
        return {"grain": "ytd", "year": None, "start": None, "end": None, "tag": tag}
    m = _FY_SHORT.match(text)
    if m:
        return _year(2000 + int(m.group(1)), fye, tag)
    m = _FY_SPLIT.match(text)
    if m:
        end = m.group(2)
        return _year(int(m.group(1)[:2] + end) if len(end) == 2 else int(end), fye, tag)
    found = cr.parse_period(text)
    if not found:
        return None
    grain, day = found
    if grain == "ytd":
        return {"grain": "ytd", "year": None, "start": None, "end": None, "tag": tag}
    if grain == "year":
        return _year(int(cr._P_YEAR.fullmatch(text).group(1)), fye, tag)
    if grain in ("quarter", "half"):
        upper = text.upper()
        q = re.search(r"[QH]([1-4])", upper)
        digits = re.search(r"\d{2,4}", upper[:q.start()] + " " + upper[q.end():]).group()
        year = int(digits) + (2000 if len(digits) == 2 else 0)
        size = 3 if grain == "quarter" else 6
        start, end = _fiscal(year, fye, size * (int(q.group(1)) - 1), size)
        return {"grain": grain, "year": year, "start": start, "end": end, "tag": tag}
    return _month(day.year, day.month, fye, tag)


def _year(year: int, fye: int, tag: Optional[str]) -> Dict:
    start, end = _fiscal(year, fye, 0, 12)
    return {"grain": "year", "year": year, "start": start, "end": end, "tag": tag}


def _month(year: int, month: int, fye: int, tag: Optional[str]) -> Dict:
    ym = _ym((year, month))
    return {"grain": "month", "year": year + (1 if month > fye else 0), "start": ym, "end": ym, "tag": tag}


def _cell_tag(text: str) -> Optional[str]:
    text = (text or "").strip()
    if text.upper() in _TAG_LETTERS and len(text) == 1:
        return _TAG_LETTERS[text.upper()]
    for rx, name in _TAG_WORDS:
        if rx.search(text):
            return name
    return None


def _tags(grid: List[List], header: int, cols: List[int], fmt: Dict, own: Dict[int, Optional[str]]) -> Dict[int, Tuple[str, str]]:
    """{column: (tag, how)} for the period columns (1-based): the header's own tag (`own`), else the nearest tag above it (a merged
    cell over its whole span, and no further); an untagged column after a tagged one that is not merged carries that tag; untagged
    columns before the first forecast column are actual; else unknown."""
    spans = fmt.get("merged") or []
    above: Dict[int, Optional[str]] = {}
    in_span = set()                         # columns a merged tag covers: its tag stops at the end of its span
    for c in cols:
        if own.get(c):
            above[c] = own[c]
            continue
        found = None
        for r in range(header - 1, 0, -1):          # sheet rows above the header, nearest first
            row = grid[r - 1] if r - 1 < len(grid) else []
            v = row[c - 1] if c - 1 < len(row) else None
            if _blank(v):
                for r1, c1, r2, c2 in spans:
                    if r1 <= r <= r2 and c1 <= c <= c2:
                        top = grid[r1 - 1] if r1 - 1 < len(grid) else []
                        v = top[c1 - 1] if c1 - 1 < len(top) else None
                        break
            if not _blank(v) and isinstance(v, str):
                found = _cell_tag(v)
                if found:
                    if any(r1 <= r <= r2 and c1 <= c <= c2 and c2 > c1 for r1, c1, r2, c2 in spans):
                        in_span.add(c)
                    break
        above[c] = found
    out, last = {}, None
    for c in cols:
        if above.get(c):
            last = None if c in in_span else above[c]
            out[c] = (above[c], "header" if own.get(c) else "sheet")
        elif last:
            out[c] = (last, "carried")
    forecast = [c for c in cols if out.get(c, (None,))[0] == "forecast"]
    for c in cols:
        if c not in out:
            out[c] = ("actual", "before_forecast") if forecast and c < forecast[0] else ("unknown", "none")
    return out


# ---------------------------------------------------------------------------
# Python's reading (26.2)
# ---------------------------------------------------------------------------
def _norm(text: str) -> str:
    return cr._norm(text)


def meaning_of(label: str) -> Tuple[str, int]:
    """(meaning, confidence) of a label by the synonyms: 100 when the label, a leading "total" left out, is a synonym; 80 when a
    synonym stands in it as whole words; ("other", 0) when none does."""
    text = _norm(label)
    bare = re.sub(r"^(?:total|sum)\s+", "", text)
    padded = f" {text} "
    if _SHARE.search((label or "").lower()):
        return "other", 100
    if _NEW_ARR.match(text):
        return "new_arr", 100
    if _NEW_ACV.match(text):
        return "acv", 100
    for meaning, defs in MEANING_DEFS:
        whole = defs.get("whole", ())
        if bare in defs["words"] or bare in whole or text in whole:
            return meaning, 100
        if any(f" {_norm(w)} " in padded for w in defs["words"]):
            return meaning, 80
    if " churn " in padded:
        return "churned_mrr", 80
    return "other", 0


def is_rate(label: str, meaning: str, values: List[Optional[float]]) -> bool:
    """A rate line (rule 4): its label holds "%", "as % of", rate, ratio or growth, or margin other than gross margin; or it is a
    gross margin, churn or growth line whose every value is between -1 and 1."""
    low = (label or "").lower()
    if _RATE.search(low) or (_MARGIN.search(low) and "gross margin" not in low):
        return True
    present = [v for v in values if v is not None]
    return meaning in ("gross_margin", "churn_rate", "churned_mrr", "growth_rate") and bool(present) \
        and all(-1 <= v <= 1 for v in present)


def basis_cue(label: str) -> Optional[str]:
    """The basis a label says outright (26.15): current year / YTD = ytd; last 12 months / LTM / TTM = ltm; lifetime / total ...
    achieved = lifetime; quarterly = quarter. None when it says none."""
    text = _norm(label)
    return next((basis for rx, basis in _BASIS_CUES if rx.search(text)), None)


def _every_third_month(line: Dict, periods: List[Dict]) -> bool:
    """True when the line holds a value only in every 3rd month column (at least two values, six month columns)."""
    months = []
    for k, p in enumerate(periods or []):
        found = period_of(p["header"]) if p.get("grain") == "month" else None
        if found:
            months.append((int(found["start"][:4]) * 12 + int(found["start"][5:]), line["values"][k]))
    present = sorted(m for m, v in months if v is not None)
    return len(months) >= 6 and len(present) >= 2 and all(b - a == 3 for a, b in zip(present, present[1:]))


def basis_of(line: Dict, meaning: str, periods: Optional[List[Dict]] = None) -> Optional[str]:
    """Python's reading of how a line is measured (26.15), None for a line with no figure: the label's cue; else balance for ARR, MRR,
    customers and cash; else quarter for values only in every 3rd month; else month."""
    if not _has_values(line):
        return None
    cue = basis_cue(line["label"])
    if cue:
        return cue
    if meaning in BALANCE_MEANINGS:
        return "balance"
    return "quarter" if _every_third_month(line, periods) else "month"


def row_tag(label: str) -> str:
    """A label that holds forecast, target (forecast) or budget as a word, or starts with plan, planned or business plan, or holds
    "(plan)" or "plan:" (budget), tags the whole row, whatever its columns say (26.15). "actual" when it holds none."""
    text = _norm(label)
    found = sorted((m.start(), tag) for rx, tag in _TAG_CUES for m in [rx.search(text)] if m)
    plan = _PLAN_WORD.search(text) if _PLAN_LABEL.search(label or "") else None
    if plan:
        found = sorted(found + [(plan.start(), "budget")])
    return found[0][1] if found else "actual"


def line_tag(line: Dict) -> str:
    """The row tag of a stored line: forecast, budget or actual (no row tag; the columns' tags apply)."""
    return line.get("tag") if line.get("tag") in ROW_TAGS else "actual"


def cell_tag(line: Dict, period: Dict) -> str:
    """The tag of one figure: its row's when the row is tagged forecast or budget, else its column's."""
    return line_tag(line) if line_tag(line) != "actual" else period["tag"]


def not_used(line: Dict) -> Optional[str]:
    """Why a line the engine could read is not read (26.15): its row tag (forecast, budget) or its basis (quarter, ytd, ltm, lifetime,
    balance for a meaning that is no balance). None when it is read. Only month, or balance for a balance, feeds anything."""
    return line_tag(line) if line_tag(line) != "actual" else basis_unused(line)


def basis_unused(line: Dict) -> Optional[str]:
    """The line's basis when it is none the engine reads (26.15), else None."""
    basis = line.get("basis")
    if basis is None or basis == "month" or (basis == "balance" and line["meaning"] in STOCK_MEANINGS):
        return None
    return basis


def _tag_lines(lines: List[Dict]) -> None:
    """Each line's row tag, in place: from its label, else its nearest parent's."""
    by_row = {line["row"]: line for line in lines}
    for line in lines:
        line["tag"] = row_tag(line["label"])
    for line in lines:
        if line["tag"] == "actual":
            for a in _ancestors(line, by_row):
                if row_tag(a["label"]) != "actual":
                    line["tag"] = row_tag(a["label"])
                    break


def label_column(grid: List[List], header: int, period_cols: List[int], width: int) -> Optional[int]:
    """The label column: of the columns before the first period column, the one with the most text cells below the header; a tie
    goes to the one with more finance words, then to the left one."""
    first = min(period_cols) if period_cols else width + 1
    candidates = [c for c in range(1, width + 1) if c < first] or [c for c in range(1, width + 1) if c not in period_cols]
    best = None
    for c in candidates:
        texts = [cell_text(row[c - 1]) for r, row in enumerate(grid[header:], header + 1)
                 if c - 1 < len(row) and cell_kind(row[c - 1], r, header) == "text"]
        key = (len(texts), sum(1 for t in texts if redact.has_finance_word(t)), -c)
        if texts and (best is None or key > best[0]):
            best = (key, c)
    return best[1] if best else None


def read(grid: List[List], fmt: Optional[Dict] = None, fiscal_year_end: int = 12, header: Optional[int] = None,
         label_col: Optional[int] = None) -> Optional[Dict]:
    """Python's reading of a statement, or None when the sheet is none (no period column or no label column). `header` and
    `label_col` (sheet positions) override Python's own when the reply gives better ones (26.4).

    {"header_row", "label_col", "width", "periods": [{"col", "header", "grain", "tag", "tag_how"}], "lines": [{"row", "label",
    "indent", "values"}], "rows": the skeleton's row facts, "cells": the text cells, "too_long", "titles"}; values are aligned with
    periods. Rows are read up to the 400th row with text, title and header rows included."""
    fmt = fmt or {}
    grid = [list(r) for r in grid or []]
    while grid and cr._filled(grid[-1]) == 0:
        grid.pop()
    if not grid:
        return None
    h = header if header is not None else cr.find_header_row(grid) + 1        # the header's sheet row
    if h > len(grid):
        return None
    width = max(len(r) for r in grid)
    head = grid[h - 1] + [None] * (width - len(grid[h - 1]))
    periods = []
    for c in range(1, width + 1):
        found = period_of(head[c - 1])
        if found:
            periods.append({"col": c, "header": cell_text(head[c - 1]), "grain": found["grain"], "own_tag": found["tag"]})
    if not periods:
        return None
    label = label_col if label_col is not None else label_column(grid, h, [p["col"] for p in periods], width)
    if label is None:
        return None
    tags = _tags(grid, h, [p["col"] for p in periods], fmt, {p["col"]: p.pop("own_tag") for p in periods})
    for p in periods:
        p["tag"], p["tag_how"] = tags[p["col"]]
    rows, cells, lines, texted, too_long = [], [], [], 0, False
    for r, row in enumerate(grid, 1):
        row = row + [None] * (width - len(row))
        kinds = [cell_kind(v, r, h) for v in row]
        if "text" in kinds:
            texted += 1
            if texted > MAX_ROWS:
                too_long = True
                break
        if all(k == "blank" for k in kinds):
            continue
        text_cols = [c for c, k in enumerate(kinds, 1) if k == "text"]
        first_text = label if kinds[label - 1] == "text" else (text_cols[0] if text_cols else None)
        rows.append({"row": r, "indent": _indent(row[first_text - 1], r, first_text, fmt) if first_text else 0,
                     "num": [c for c, k in enumerate(kinds, 1) if k == "num"],
                     "blank": [c for c, k in enumerate(kinds, 1) if k == "blank"]})
        cells += [{"row": r, "col": c, "text": cell_text(row[c - 1])} for c in text_cols]
        if r > h and kinds[label - 1] == "text":
            lines.append({"row": r, "label": cell_text(row[label - 1]), "indent": rows[-1]["indent"],
                          "values": [number(row[p["col"] - 1]) if kinds[p["col"] - 1] == "num" else None for p in periods]})
    for r1, c1, r2, c2 in fmt.get("merged") or []:
        for cell in cells:
            if (cell["row"], cell["col"]) == (r1, c1):
                if r2 > r1:
                    cell["row_span"] = r2 - r1 + 1
                if c2 > c1:
                    cell["col_span"] = c2 - c1 + 1
    titles = [(c["row"], c["col"], c["text"]) for c in cells if c["row"] < h]
    return {"header_row": h, "label_col": label, "width": width, "periods": periods, "lines": lines, "rows": rows,
            "cells": cells, "too_long": too_long, "titles": titles}


def period_count(grid: List[List], row: int) -> int:
    """How many period headers a sheet row holds (the header of the two, Python's or the model's, with more wins: 26.4)."""
    if row < 1 or row > len(grid or []):
        return 0
    return sum(1 for v in grid[row - 1] if period_of(v))


# ---------------------------------------------------------------------------
# Labels: Python's rules, the reply, the rules that win (26.2, 26.4)
# ---------------------------------------------------------------------------
def _parents_by_indent(lines: List[Dict]) -> Dict[int, Optional[int]]:
    out, stack = {}, []
    for line in lines:
        while stack and stack[-1]["indent"] >= line["indent"]:
            stack.pop()
        out[line["row"]] = stack[-1]["row"] if stack else None
        stack.append(line)
    return out


def _parents_by_blocks(lines: List[Dict]) -> Dict[int, Optional[int]]:
    """Without indents: a "Total" line is the parent of the lines directly above it, up to the previous heading, total or empty
    row."""
    out = {line["row"]: None for line in lines}
    for i, line in enumerate(lines):
        if not _TOTAL.match(line["label"]):
            continue
        j = i - 1
        while j >= 0:
            above, below = lines[j], lines[j + 1]
            if below["row"] - above["row"] > 1 or _TOTAL.match(above["label"]) or _RESULT.search(above["label"]) \
                    or not _has_values(above):
                break
            out[above["row"]] = line["row"]
            j -= 1
    return out


def _has_values(line: Dict) -> bool:
    return any(v is not None for v in line["values"])


def _section(line: Dict, by_row: Dict[int, Dict], lines: List[Dict], parent_of=lambda x: x.get("parent_row")) -> str:
    """The labels of the sections a line sits in: its parents' labels, and the nearest heading above it."""
    out, seen, parent = [], set(), parent_of(line)
    while parent and parent not in seen and parent in by_row:
        seen.add(parent)
        out.append(by_row[parent]["label"])
        parent = parent_of(by_row[parent])
    for other in reversed([x for x in lines if x["row"] < line["row"]]):
        if not _has_values(other):
            out.append(other["label"])
            break
    return " · ".join(out)


def python_labels(lines: List[Dict], periods: Optional[List[Dict]] = None) -> None:
    """Python's role, parent, meaning and basis for every line, in place, under "python" (26.2, 26.15)."""
    indented = any(line["indent"] for line in lines)
    parents = _parents_by_indent(lines) if indented else _parents_by_blocks(lines)
    children = {}
    for row, parent in parents.items():
        if parent:
            children.setdefault(parent, []).append(row)
    by_row = {line["row"]: line for line in lines}
    for line in lines:
        meaning, confidence = meaning_of(line["label"])
        if is_rate(line["label"], meaning, line["values"]):
            role = "rate"
            meaning = "churn_rate" if meaning == "churned_mrr" else meaning
        elif not _has_values(line):
            role = "note" if len(line["label"]) > NOTE_LENGTH else "heading"
        elif _TOTAL.match(line["label"]) or line["row"] in children:
            role = "total"
        elif parents[line["row"]]:
            role = "subline"
        else:
            role = "other"
        line["python"] = {"role": role, "parent_row": parents[line["row"]], "meaning": meaning, "confidence": confidence}

    for line in lines:                       # "Sales" alone in a cost or OPEX section is S&M expense (rule 4)
        if _norm(line["label"]) == "sales" and _COST_SECTION.search(
                _section(line, by_row, lines, lambda x: x["python"]["parent_row"])):
            line["python"].update(meaning="sm_expense", confidence=100)
    for line in lines:                       # a sub-line with no meaning of its own takes its parent's
        own, parent = line["python"], by_row.get(line["python"]["parent_row"])
        if parent and own["role"] != "rate" and parent["python"]["meaning"] != "other" and (
                own["meaning"] == "other" or (own["meaning"] == "opex" and parent["python"]["meaning"] in ("sm_expense",
                                                                                                         "cost_of_revenue"))):
            own.update(meaning=parent["python"]["meaning"], confidence=INHERITED)
    for line in lines:                       # the basis reads the final meaning, a sub-line's inherited one included
        line["python"]["basis"] = basis_of(line, line["python"]["meaning"], periods)


def label_lines(reading: Dict, reply: Optional[Dict] = None) -> List[Dict]:
    """Every line with its role, parent and meaning: the reply's for the rows it names (`source` "ai" where its meaning is not
    Python's), Python's for the others, then the rules that win (26.2). Returns the lines (a copy)."""
    lines = [dict(line) for line in reading["lines"]]
    python_labels(lines, reading.get("periods"))
    rows = {line["row"] for line in lines}
    replied = {r["row"]: r for r in (reply or {}).get("rows") or [] if r.get("row") in rows}
    for line in lines:
        mine, theirs = line["python"], replied.get(line["row"])
        if theirs:
            parent = theirs.get("parent_row")
            line.update(role=theirs["role"], meaning=theirs["meaning"],
                        parent_row=parent if parent in rows and parent != line["row"] else None)
            agrees = theirs["meaning"] == mine["meaning"]
            line.update(source="rules" if agrees else "ai", confidence=mine["confidence"] if agrees else None)
            line["basis"] = _reply_basis(line, theirs.get("basis"), mine["basis"])
        else:
            line.update(role=mine["role"], meaning=mine["meaning"], parent_row=mine["parent_row"], source="rules",
                        confidence=mine["confidence"], basis=mine["basis"])
    _break_cycles(lines)
    _rules_win(lines)
    _tag_lines(lines)
    return lines


def _reply_basis(line: Dict, basis: Optional[str], python: Optional[str]) -> Optional[str]:
    """The basis of a line the reply names (26.15): a cue in the label wins, then the reply's, then Python's. None for a line with
    no figure."""
    if not _has_values(line):
        return None
    return basis_cue(line["label"]) or (basis if basis in SHEET_BASES else None) or python


def _break_cycles(lines: List[Dict]) -> None:
    by_row = {line["row"]: line for line in lines}
    for line in lines:
        seen, parent = {line["row"]}, line.get("parent_row")
        while parent:
            if parent in seen:
                line["parent_row"] = None
                break
            seen.add(parent)
            parent = by_row[parent].get("parent_row") if parent in by_row else None


def _rules_win(lines: List[Dict]) -> None:
    """Rule 4, over any reply: "Sales" alone inside a cost or OPEX section is S&M expense; a rate line is role rate."""
    by_row = {line["row"]: line for line in lines}
    for line in lines:
        if line.get("decision"):
            continue                      # the analyst's own decision stands
        if _norm(line["label"]) == "sales" and _COST_SECTION.search(_section(line, by_row, lines)):
            if line["meaning"] != "sm_expense":
                line.update(meaning="sm_expense", source="rules", confidence=100)
        if is_rate(line["label"], line["meaning"], line["values"]) and line["role"] != "rate":
            line["role"] = "rate"
        if _SHARE.search(line["label"].lower()) and line["meaning"] != "other":
            line.update(meaning="other", source="rules", confidence=100)      # a "% share" line is a rate, never a figure


def apply_decisions(lines: List[Dict], decided: List[Dict]) -> List[Dict]:
    """The analyst's decisions (and a saved map's) over the lines: {row, role, meaning, parent_row, source, decision}."""
    by_row = {d["row"]: d for d in decided or []}
    for line in lines:
        d = by_row.get(line["row"])
        if not d:
            continue
        for key in ("role", "meaning"):
            if d.get(key) in (SHEET_ROLES if key == "role" else SHEET_MEANINGS):
                line[key] = d[key]
        if "parent_row" in d:
            line["parent_row"] = d["parent_row"]
        if d.get("basis") in SHEET_BASES and _has_values(line):
            line["basis"] = d["basis"]
        line.update(source=d.get("source") or line.get("source"), decision=d.get("decision"))
    _break_cycles(lines)
    _tag_lines(lines)
    return lines


# ---------------------------------------------------------------------------
# Checks (26.5)
# ---------------------------------------------------------------------------
def sublines(total: Dict, lines: List[Dict]) -> List[Dict]:
    """The lines a total adds up: those that hang from it; for a total with none, those that hang from the heading above it."""
    def counted(line):
        return line["role"] not in ("rate", "heading", "note") and _has_values(line) and line_tag(line) == line_tag(total)
    own = [x for x in lines if x.get("parent_row") == total["row"] and counted(x)]
    if own:
        return own
    heading = next((x for x in reversed(lines) if x["row"] < total["row"] and x["role"] == "heading"
                    and x.get("parent_row") == total.get("parent_row")), None)
    if heading is None:
        return []
    return [x for x in lines if x.get("parent_row") == heading["row"] and x["row"] != total["row"] and counted(x)]


def check(lines: List[Dict], periods: List[Dict]) -> None:
    """Python's checks, in place: each line's `checks` {code: [period headers]} (total_mismatch, no_breakdown, sign_differs,
    sign_convention)."""
    for line in lines:
        line["checks"] = {}
    for total in lines:
        if total["role"] != "total" or not _has_values(total):
            continue
        parts = sublines(total, lines)
        if not parts:
            continue
        for k, p in enumerate(periods):
            value = total["values"][k]
            if value is None:
                continue
            found = [x["values"][k] for x in parts if x["values"][k] is not None]
            if not found:
                total["checks"].setdefault("no_breakdown", []).append(p["header"])
                continue
            tolerance = TOLERANCE * abs(value) + 1e-9
            if abs(sum(found) - value) <= tolerance:
                continue
            if abs(-sum(found) - value) <= tolerance or abs(sum(abs(v) for v in found) - abs(value)) <= tolerance:
                total["checks"].setdefault("sign_differs", []).append(p["header"])
            else:
                total["checks"].setdefault("total_mismatch", []).append(p["header"])
    for line in lines:
        if line["meaning"] in COSTS and _negative(line):
            line["checks"]["sign_convention"] = []
    for found in basis_checks(lines, periods):
        if not found["ok"]:
            by_row = next(x for x in lines if x["row"] == found["row"])
            by_row["checks"].setdefault(f"{found['basis']}_mismatch", []).append(found["period"])


def basis_checks(lines: List[Dict], periods: List[Dict]) -> List[Dict]:
    """The yearly and rolling lines against the monthly line of the same meaning (26.15): the last month of a fiscal year of a ytd line
    equals the sum of the monthly line over that year, an ltm line equals the sum over its last 12 months, each within 0.5%.
    [{"row", "basis", "period", "line", "reference", "ok"}]; a period whose months are not all there is not checked. Costs are
    compared as positive amounts."""
    months = {periods[k]["start"]: k for k in monthly_columns(periods) if periods[k].get("start")}
    out = []
    for line in lines:
        if line.get("basis") not in ("ytd", "ltm") or line_tag(line) != "actual" or line["role"] in NOT_LINES + ("rate",):
            continue
        reference = pick(lines, line["meaning"])
        if not reference or not months:
            continue
        starts = sorted(months)

        def monthly(start_list):
            """The monthly line(s) summed over these months, None when one month is blank."""
            total = 0.0
            for start in start_list:
                for ref in reference:
                    v = ref["values"][months[start]]
                    if v is None:
                        return None
                    total += v
            return total

        def compare(k, window):
            value, ref = line["values"][k], monthly(window)
            if value is None or ref is None:
                return
            a, b = (abs(value), abs(ref)) if line["meaning"] in COSTS else (value, ref)
            out.append({"row": line["row"], "basis": line["basis"], "period": periods[k]["header"], "line": value, "reference": ref,
                        "ok": abs(a - b) <= TOLERANCE * abs(b) + 1e-9})
        if line["basis"] == "ytd":
            years = {}
            for start in starts:
                years.setdefault(periods[months[start]]["year"], []).append(start)
            for year_months in years.values():
                if len(year_months) == 12:
                    compare(months[year_months[-1]], year_months)
        else:
            for i, start in enumerate(starts):
                window = starts[i - 11:i + 1]
                if i >= 11 and _ym(_add_months(int(window[0][:4]), int(window[0][5:]), 11)) == start:
                    compare(months[start], window)
    return out


def _negative(line: Dict) -> bool:
    present = [v for v in line["values"] if v]
    return bool(present) and all(v < 0 for v in present)


# ---------------------------------------------------------------------------
# Use of the map (26.6)
# ---------------------------------------------------------------------------
def _ancestors(line: Dict, by_row: Dict[int, Dict]) -> List[Dict]:
    out, seen, parent = [], set(), line.get("parent_row")
    while parent and parent in by_row and parent not in seen:
        seen.add(parent)
        out.append(by_row[parent])
        parent = by_row[parent].get("parent_row")
    return out


def pick(lines: List[Dict], meaning: str, tag: str = "actual") -> List[Dict]:
    """The lines of a meaning that do not hang from a line of the same meaning, rates, headings and notes left out: summed, they
    are the field (S&M is the sum of the Sales and Marketing totals when no single S&M line holds both). Only a line measured by
    the month (a balance for ARR, MRR, customers and the like) and with no forecast or budget row tag is read (26.15): a ytd, ltm,
    quarter or lifetime line, or a forecast row, is never summed with them. `tag` "forecast" or "budget" picks the rows with that
    row tag instead (their basis still month or balance)."""
    by_row = {line["row"]: line for line in lines}
    return [x for x in lines if x["meaning"] == meaning and x["role"] not in ("rate",) + NOT_LINES and _has_values(x)
            and line_tag(x) == tag and basis_unused(x) is None
            and not any(a["meaning"] == meaning and line_tag(a) == tag for a in _ancestors(x, by_row))]


def _field_values(chosen: List[Dict], k: int, scale: float, meaning: str) -> Optional[float]:
    """The sum of the chosen lines in period k, blank when one of them is blank; a cost line written as negatives read as
    positive; money ×scale."""
    total = 0.0
    for line in chosen:
        v = line["values"][k]
        if v is None:
            return None
        total += abs(v) if meaning in COSTS and _negative(line) else v
    return total * (scale if meaning in MONEY else 1)


def with_periods(periods: List[Dict], fiscal_year_end: int) -> List[Dict]:
    """The periods with their fiscal year, first and last month by the audit's year-end, read again from each header."""
    out = []
    for p in periods:
        found = period_of(p["header"], fiscal_year_end) or {}
        out.append({**p, "year": found.get("year"), "start": found.get("start"), "end": found.get("end")})
    return out


def monthly_columns(periods: List[Dict]) -> List[int]:
    """Indexes of the month columns that may feed the engine: tagged actual or unknown, never forecast or budget."""
    return [k for k, p in enumerate(periods) if p["grain"] == "month" and p["tag"] in ("actual", "unknown")]


def monthly_pnl(lines: List[Dict], periods: List[Dict], scale: float) -> Optional[Dict]:
    """{"rows": [{"month", "revenue", "cost_of_revenue", "sm_expense"}], "lines": {field: [rows]}} when the map holds a monthly
    P&L (month columns and a line for each field), else None."""
    months = monthly_columns(periods)
    chosen = {f: pick(lines, f) for f in PNL_FIELDS}
    if not months or not all(chosen.values()):
        return None
    rows = [{"month": periods[k]["start"], **{f: _field_values(chosen[f], k, scale, f) for f in PNL_FIELDS}} for k in months]
    return {"rows": rows, "lines": {f: [x["row"] for x in chosen[f]] for f in PNL_FIELDS}}


def kpi_series(lines: List[Dict], periods: List[Dict], scale: float) -> List[Dict]:
    """The company-reported KPI series (arr, mrr, new_arr, churn_rate) by month: one per line that does not hang from a line of the
    same meaning, each value with its cell. Labelled Unverified; the engine never reads them."""
    months = monthly_columns(periods)
    if not months:
        return []
    by_row = {line["row"]: line for line in lines}
    out = []
    for line in lines:
        if line["meaning"] not in KPI_SERIES or line["role"] in NOT_LINES or not _has_values(line) or not_used(line):
            continue
        if any(a["meaning"] == line["meaning"] for a in _ancestors(line, by_row)):
            continue
        money = line["meaning"] in MONEY and line["role"] != "rate"
        out.append({"meaning": line["meaning"], "row": line["row"], "label": "Unverified",
                    "values": [{"month": periods[k]["start"], "value": (line["values"][k] * scale if money and line["values"][k]
                                                                         is not None else line["values"][k]),
                                "cell": f"{get_column_letter(periods[k]['col'])}{line['row']}"} for k in months]})
    return out


def used_rows(lines: List[Dict], periods: List[Dict]) -> set:
    """The lines the map uses: the monthly P&L fields (when it holds one), the KPI series and the yearly inputs."""
    used = set()
    if monthly_pnl(lines, periods, 1):
        used |= {x["row"] for f in PNL_FIELDS for x in pick(lines, f)}
    used |= {s["row"] for s in kpi_series(lines, periods, 1)}
    if any(p["grain"] == "year" and p["tag"] == "actual" for p in periods):
        used |= {x["row"] for f in YEARLY_FIELDS for x in pick(lines, f)}
    return used


def is_kpi_sheet(lines: List[Dict]) -> bool:
    """True when most mapped lines (a meaning other than `other`, a role that is no heading or note, a figure in some period) read as
    ARR, MRR, new ARR, churned MRR, churn rate or customers and none reads as cost of revenue or S&M expense (26.14)."""
    mapped = [x for x in lines if x["meaning"] != "other" and x["role"] not in NOT_LINES and _has_values(x)]
    kpi = [x for x in mapped if x["meaning"] in KPI_MEANINGS]
    return bool(kpi) and len(kpi) * 2 > len(mapped) and not any(x["meaning"] in ("cost_of_revenue", "sm_expense") for x in mapped)


def pending(line: Dict, used: set) -> bool:
    """A line waits for the analyst when the map uses it and its meaning is the model's (not Python's) or Python's at 80 or
    below, and it is not decided."""
    if line["row"] not in used or line.get("decision") or line.get("source") in ("saved", "you"):
        return False
    return line.get("source") == "ai" or (line.get("confidence") or 0) <= cr.MAPPING_THRESHOLD


_SM_TOTAL = re.compile(r"sales\s*(?:and|&|/)\s*marketing|\bs\s*&\s*m\b")
_SM_MARKETING = re.compile(r"\bmarketing\b|\badvertising\b")
_SM_SALES = re.compile(r"\bsales\b")


def sm_scope(labels: List[str]) -> str:
    """What the S&M line covers, from the labels of the lines that were summed (docs/specs/forecast-and-payback.md section 4.3): `sales and marketing`
    (one line that names both, a Sales line with a Marketing line, or "Sales" alone in a cost section, rule 4 of 26.2), `marketing only`
    (marketing or advertising), `sales only` (sales expenses or costs), else `unknown`. A marketing-only line is never total S&M."""
    texts = [str(label).lower().strip() for label in labels]
    if any(_SM_TOTAL.search(t) or t == "sales" for t in texts) or (any(_SM_SALES.search(t) for t in texts) and any(_SM_MARKETING.search(t) for t in texts)):
        return "sales and marketing"
    if texts and all(_SM_MARKETING.search(t) for t in texts):
        return "marketing only"
    if texts and all(_SM_SALES.search(t) for t in texts):
        return "sales only"
    return "unknown"


def yearly_input(lines: List[Dict], periods: List[Dict], scale: float, meta: Dict) -> Optional[Dict]:
    """The engine's yearly input (26.8): every actual year with its fields and cells, the forecast and budget years of the same
    lines (the track record) and the actual years, halves and quarters with a revenue line (against the revenue file). None
    without an actual year."""
    years = [k for k, p in enumerate(periods) if p["grain"] == "year" and p["tag"] == "actual"]
    if not years:
        return None
    chosen = {f: pick(lines, f) for f in TRACK_FIELDS}

    def fields(k):
        out = {}
        for f in TRACK_FIELDS:
            v = _field_values(chosen[f], k, scale, f) if chosen[f] else None
            if v is not None:
                out[f] = {"value": v, "cells": [f"{get_column_letter(periods[k]['col'])}{x['row']}" for x in chosen[f]]}
                if f == "sm_expense":
                    out[f]["scope"] = sm_scope([x["label"] for x in chosen[f]])
        return out
    actual = [{"period": periods[k]["header"], "year": periods[k]["year"], "start": periods[k]["start"], "end": periods[k]["end"],
               "fields": fields(k)} for k in years]
    planned = [{"period": periods[k]["header"], "year": periods[k]["year"], "tag": periods[k]["tag"], "fields": fields(k)}
               for k, p in enumerate(periods) if p["grain"] == "year" and p["tag"] in ("forecast", "budget")]
    reported = [{"period": periods[k]["header"], "start": periods[k]["start"], "end": periods[k]["end"],
                 "revenue": fields(k).get("revenue")} for k, p in enumerate(periods)
                if p["grain"] in ("year", "half", "quarter") and p["tag"] == "actual"]
    rows = sorted({x["row"] for f in TRACK_FIELDS for x in chosen[f]})
    return {"file": meta.get("file"), "sheet": meta.get("sheet"), "rows": rows, "actual": actual, "planned": planned,
            "reported": [r for r in reported if r["revenue"]]}


def claims_found(lines: List[Dict], periods: List[Dict], ds: Dict, fiscal_year_end: int, rows_only: bool = False) -> Optional[Dict]:
    """The forecast-table claims of a statement (section 20, amended by 26.6): its plan lines, found by section 20's own rules (no
    model), over the columns tagged forecast, budget or unknown; an actual-tagged column is no claim, and a line the map reads as a
    cost ("Sales" in an OPEX section) is none either. The shape of forecast_claims.read, each value with its column's tag. A row tagged
    forecast or budget (26.15) is a claim in every column it has a figure in, whatever its column's tag; with `rows_only` (a monthly
    P&L feeds the engine) those rows are the only claims. A line not measured by the month is none."""
    found = {}
    for line in lines:
        metric = forecast_claims.metric_of(line["label"]) if line["meaning"] not in COSTS + NOT_PLAN_LINES and line["role"] != "rate" \
            and not basis_unused(line) and (line_tag(line) != "actual" or not rows_only) else None
        if not metric:
            continue
        parts = []
        for k, p in enumerate(periods):
            value = line["values"][k]
            tag = cell_tag(line, p)
            if value is None or tag == "actual" or p["grain"] not in ("year", "quarter", "half", "month"):
                continue
            day = date.fromisoformat(p["start"] + "-01") if p["grain"] == "month" else None
            part = forecast_claims._part(p["header"], day, fiscal_year_end)
            if part:
                parts.append({"fy": part[0], "grain": part[1], "index": part[2], "value": value, "j": p["col"] - 1,
                              "header": p["header"], "tag": tag})
        if parts:
            found[line["row"]] = {"row": line["row"], "label": line["label"], "metric": metric, "parts": parts}
    plan = None if rows_only else _sm_claim(lines, periods, fiscal_year_end)
    if plan:
        found[plan["row"]] = plan
    return forecast_claims.assemble(found, {**ds, "first_col": 1}, fiscal_year_end)


def _sm_claim(lines: List[Dict], periods: List[Dict], fiscal_year_end: int) -> Optional[Dict]:
    """The S&M plan line (26.6): per forecast, budget or unknown period, the sum of the S&M total lines (`pick`), cited to their
    cells. A period where one of them is blank has no claim. The shape of a plan line of forecast_claims.assemble."""
    chosen = pick(lines, "sm_expense")
    if not chosen:
        return None
    parts = []
    for k, p in enumerate(periods):
        if p["tag"] == "actual" or p["grain"] not in ("year", "quarter", "half", "month"):
            continue
        value = _field_values(chosen, k, 1, "sm_expense")
        part = forecast_claims._part(p["header"], date.fromisoformat(p["start"] + "-01") if p["grain"] == "month" else None,
                                      fiscal_year_end)
        if value is not None and part:
            parts.append({"fy": part[0], "grain": part[1], "index": part[2], "value": value, "j": p["col"] - 1, "header": p["header"],
                          "tag": p["tag"], "cells": [f"{get_column_letter(p['col'])}{x['row']}" for x in chosen]})
    return {"row": chosen[0]["row"], "label": SM_CLAIM_LABEL, "metric": "sm_expense", "parts": parts} if parts else None


# ---------------------------------------------------------------------------
# The skeleton (26.3)
# ---------------------------------------------------------------------------
def skeleton(reading: Dict, company_name: Optional[str], mapping: Dict[str, str], withheld: Tuple[str, ...] = ()
             ) -> Tuple[str, Dict]:
    """(the text the model reads, the counts the analyst sees). Every text cell, redacted (the company name withheld, emails,
    phone numbers, names and customer names); a cell over 200 characters left out; a column the withhold rule takes keeps its title
    and header cells and is one line; one line per row with content. Built in memory, never stored."""
    long_cells = [c for c in reading["cells"] if len(c["text"]) > LONG_CELL]
    raw = [c for c in reading["cells"] if len(c["text"]) <= LONG_CELL]
    cells, _ = redact.redact_structure(raw, company_name, mapping, withheld)
    held = redact.withheld_columns(cells)       # read on the redacted cells, as the gateway reads them
    keep = [i for i, c in enumerate(cells) if c["col"] not in held or c["row"] <= reading["header_row"]]
    dropped = {col: len({cells[i]["text"] for i in range(len(cells)) if cells[i]["col"] == col and i not in keep})
               for col in held}
    text = redact.sheet_text([cells[i] for i in keep], reading["rows"], dropped)
    _, counts = redact.redact_structure([raw[i] for i in keep], company_name, mapping, withheld)   # what was masked in the text
    counts = {"lines": len(keep), "email": counts["email"], "phone": counts["phone"], "person": counts["person"],
              "customer": counts["customer"], "company": counts["withheld"], "columns": len(held), "long": len(long_cells)}
    return text, counts


# ---------------------------------------------------------------------------
# The reply (26.4)
# ---------------------------------------------------------------------------
def merged(grid: List[List], fmt: Optional[Dict], fiscal_year_end: int, reading: Dict, reply: Optional[Dict]) -> Dict:
    """Python's reading after the reply: the header row with more period headers (Python's on a tie), the model's label column
    when it holds a text cell below the header, Python's periods (read from their header cells) with the model's tag only where
    Python found none, and the lines labelled by label_lines."""
    if not reply:
        return reading
    header = reading["header_row"]
    if reply.get("header_row") and period_count(grid, reply["header_row"]) > period_count(grid, header):
        header = reply["header_row"]
    again = read(grid, fmt, fiscal_year_end, header, None) or reading
    candidate = reply.get("label_col")
    if candidate and candidate != again["label_col"]:
        other = read(grid, fmt, fiscal_year_end, again["header_row"], candidate)
        if other and other["lines"]:
            again = other
    tags = {p["col"]: p["tag"] for p in reply.get("periods") or []}
    for p in again["periods"]:
        if p["tag_how"] == "none" and tags.get(p["col"]) in ("actual", "forecast", "budget"):
            p["tag"], p["tag_how"] = tags[p["col"]], "ai"
    return again


def build(reading: Dict, reply: Optional[Dict] = None, decided: Optional[List[Dict]] = None) -> Dict:
    """The stored map of a statement file: its reading with every line labelled (the reply, then the analyst's decisions)."""
    lines = apply_decisions(label_lines(reading, reply), decided or [])
    return {"header_row": reading["header_row"], "label_col": reading["label_col"], "ai_read": reply is not None,
            "periods": [{k: p[k] for k in ("col", "header", "grain", "tag", "tag_how")} for p in reading["periods"]],
            "lines": [{k: line.get(k) for k in ("row", "label", "indent", "values", "role", "meaning", "basis", "tag", "parent_row",
                                                "source", "confidence", "decision", "python")} for line in lines],
            "too_long": reading["too_long"]}


def view(sheet_map: Dict, fiscal_year_end: int, scale: float) -> Dict:
    """What the table and the status line read, computed from the stored map: every line with its checks and whether it waits, in
    the order of 26.7, and the map's use (monthly P&L, series, yearly input, the missing fields)."""
    periods = with_periods(sheet_map["periods"], fiscal_year_end)
    lines = [dict(line) for line in sheet_map["lines"]]
    check(lines, periods)
    used = used_rows(lines, periods)
    series = {s["row"] for s in kpi_series(lines, periods, scale)}
    for line in lines:
        line["pending"] = pending(line, used)
        line["series"] = line["row"] in series
        line["used"] = line["row"] in used
        feeds = line["meaning"] in FEEDS and line["role"] not in ("rate",) + NOT_LINES and _has_values(line)
        line["not_used"] = not_used(line) if feeds and not line["used"] else None
    order = sorted(lines, key=lambda x: (0 if any(c in x["checks"] for c in FAILED_CHECKS) else 1 if x["pending"] else 2,
                                         x["row"]))
    months = [p for p in periods if p["grain"] == "month"]
    pnl = monthly_pnl(lines, periods, scale)
    dated = [p for p in periods if p["grain"] in ("month", "quarter", "half", "year")]
    return {"header_row": sheet_map["header_row"], "label_col": get_column_letter(sheet_map["label_col"]),
            "periods": [{"col": get_column_letter(p["col"]), "header": p["header"], "grain": p["grain"], "tag": p["tag"]}
                        for p in periods],
            "lines": [{k: x.get(k) for k in ("row", "label", "indent", "role", "meaning", "basis", "tag", "parent_row", "source",
                                             "confidence", "decision", "checks", "pending", "series", "used", "not_used")}
                      for x in order],
            "ai_read": bool(sheet_map.get("ai_read")),
            "pending": sum(1 for x in lines if x["pending"]), "too_long": bool(sheet_map.get("too_long")),
            "grain": "monthly" if months else "yearly", "monthly_pnl": bool(pnl),
            "missing_fields": [] if pnl or not months else [f for f in PNL_FIELDS if not pick(lines, f)],
            "one_period": not dated, "yearly": any(p["grain"] == "year" and p["tag"] == "actual" for p in periods),
            "series": len(series)}


def feed(sheet_map: Dict, fiscal_year_end: int, scale: float, meta: Dict) -> Dict:
    """What compute reads from the stored map: {"pnl": monthly P&L rows or None, "rows": the P&L lines' rows, "yearly": the yearly
    input or None}. Never a split of a year or a quarter into months."""
    periods = with_periods(sheet_map["periods"], fiscal_year_end)
    lines = sheet_map["lines"]
    pnl = monthly_pnl(lines, periods, scale)
    return {"pnl": pnl["rows"] if pnl else None,
            "rows": sorted({r for rows in (pnl or {}).get("lines", {}).values() for r in rows}),
            "revenue_row": ((pnl or {}).get("lines", {}).get("revenue") or [None])[0],
            "yearly": None if pnl else yearly_input(lines, periods, scale, meta)}


def claims(sheet_map: Dict, ds: Dict, fiscal_year_end: int) -> Optional[Dict]:
    """The forecast claims of a stored statement file. When it feeds a monthly P&L (its months are engine input) only its rows tagged
    forecast or budget are claims (26.15)."""
    periods = with_periods(sheet_map["periods"], fiscal_year_end)
    pnl = monthly_pnl(sheet_map["lines"], periods, 1)
    return claims_found(sheet_map["lines"], periods, ds, fiscal_year_end, rows_only=bool(pnl))


def saved_lines(sheet_map: Dict) -> List[Dict]:
    """A map as column_mappings keeps it: per line its row, role, meaning, parent, source and decision. No text."""
    return [{k: line.get(k) for k in ("row", "role", "meaning", "basis", "parent_row", "source", "decision")}
            for line in sheet_map["lines"]]


def pending_count(sheet_map: Optional[Dict]) -> int:
    """How many lines of a stored map wait for the analyst (the fiscal year-end moves no period's grain or tag)."""
    return view(sheet_map, 12, 1)["pending"] if sheet_map else 0


def months_found(sheet_map: Dict) -> Optional[Dict]:
    """The month columns of a map as the bubble's head shows them, {"count", "first", "last"} (YYYY-MM), or None."""
    months = sorted({p["start"] for p in with_periods(sheet_map["periods"], 12) if p["grain"] == "month" and p["start"]})
    return {"count": len(months), "first": months[0], "last": months[-1]} if months else None


def rules_cover(proposals) -> bool:
    """True when the column rules map every required P&L field at 80 or more: the file is read by columns, no sheet call (26.1)."""
    found = {p.field: p.confidence or 0 for p in proposals if p.field}
    return all(found.get(f, 0) >= cr.MAPPING_THRESHOLD for f in cr.FIELD_DEFS["pnl"]["required"])
