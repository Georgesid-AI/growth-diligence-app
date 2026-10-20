"""Revenue file recognition (docs/specs/revenue-file-recognition.md). Pure Python: no database, no model, no network.

A file is a revenue file when it holds a revenue measure, a reporting reference and a number (section 2); it is then classified into one of
six types (section 3), per sheet, by reading the labels of its rows as metric names or as entities (section 4). The model only proposes a
revenue row, a last actual column, a currency and the rows that are not revenue; Python checks every proposal before anything is used
(section 6, rule 18) and the analyst has the last word (section 6.6). `model_text` builds the text the model reads (section 5).

Everything a result says in words is one of the closed lists below; the cells of the file appear only as the labels of the rows an analyst
chooses between, never in a reason, a type or a log line (rule 14).
"""
import hashlib
import re
from datetime import date, datetime, time, timedelta
from decimal import Decimal
from types import SimpleNamespace
from typing import Dict, List, Optional, Set, Tuple

import pandas as pd
from openpyxl.utils import get_column_letter

from . import column_rules as cr
from .llm.schemas import NON_REVENUE_KINDS, RECOGNITION_TYPES  # noqa: F401  (the closed lists live beside the reply schema)
from .structures import redact
from .structures import sheet as sheet_mod

REASONS = ("no_candidate", "tie", "confidence_80", "row_missing", "label_mismatch", "marked_not_revenue", "not_numeric_enough",
           "no_revenue_word", "rules_and_model_disagree", "no_as_of_month", "after_as_of", "not_a_period", "turnover_question",
           "type_unsure")
MISSING = ("revenue_measure", "reference", "numbers")
METRIC_LABEL_SHARE = 0.5            # above this share of label rows holding a metric word, and at least MIN_METRIC_HITS of them
ENTITY_LABEL_SHARE = 0.2            # at most this share: the labels are entities; between the two the sheet is "unsure"
MIN_METRIC_HITS = 3
TOTAL_TOLERANCE = 0.005             # a total equals the sum of its rows within 0.5% (chat-upload.md 26.5)
MAX_TEXT_ROWS = redact.RECOGNITION_MAX_ROWS
ISO_CURRENCIES = ("EUR", "USD", "GBP", "CHF", "SEK", "NOK", "DKK", "PLN", "BGN", "CZK", "HUF", "RON", "CAD", "AUD", "NZD", "JPY",
                  "CNY", "HKD", "SGD", "INR", "AED", "ZAR", "BRL", "MXN", "TRY", "ILS")

NEEDS_BY_CUSTOMER = "revenue by customer and month (one row per customer, or one row per invoice with a customer column)"
NEEDS_CUSTOMER_COLUMN = "a customer column (name or ID) on each invoice row"
NEEDS_PERIODS = "the same figures for at least one more period (a prior year, or one column per month)"
NEEDS_SNAPSHOT_SERIES = "ARR or MRR for at least two periods (one column per month)"
NEEDS_MONTHLY = "one column per month"
ANALYSIS_NAMES = {"customer_level": "Customer-level analysis (NRR, gross churn, concentration, cohorts)",
                  "trend": "Trend analysis (MRR/ARR, growth, run-rate)", "monthly_series": "Monthly series missing"}
MISSING_NAMES = {"revenue_measure": "a revenue measure (recognised revenue, invoiced revenue, sales, turnover, recurring revenue, MRR or ARR)",
                 "reference": "a date, period or transaction ID", "numbers": "a number for the revenue measure"}
GROWTH_LABELS = ("Revenue growth", "MRR growth", "ARR growth", "Recurring revenue growth")

# ---------------------------------------------------------------------------
# Words
# ---------------------------------------------------------------------------
_KIND_WORDS = (
    ("forecast", ("forecast", "forecasts", "fcst", "projected", "projection", "estimate", "estimates", "outlook")),
    ("budget", ("budget", "budgeted", "target")),
    ("bookings", ("booking", "bookings", "booked")),
    ("pipeline", ("pipeline", "backlog", "opportunity", "opportunities")),
    ("contract_value", ("contract value", "contracted revenue", "total contract value", "annual contract value", "tcv", "acv",
                        "contracted arr", "committed arr", "committed mrr", "carr")),
    ("run_rate", ("run rate", "annualised", "annualized")))
_RECURRING = ("mrr", "arr", "recurring revenue")
_INVOICED = ("invoiced", "invoicing", "billing", "billed")
_RECOGNISED = ("revenue", "revenues", "sales", "turnover", "net sales")
_BALANCES = ("deferred", "unbilled", "accrued")      # balances, not the month's revenue: never a revenue row, never added (re-check Q1)
_EXCLUDED = _BALANCES + ("new", "net new", "lost", "churn", "churned", "added", "expansion", "contraction", "growth", "per", "average", "avg", "share",
             "ratio", "rate", "margin", "cost", "costs", "cogs", "expense", "expenses", "marketing", "payroll", "commission")
_SYNONYMS = frozenset(("revenue", "revenues", "net revenue", "recognised revenue", "recognized revenue", "net sales", "sales", "turnover",
                       "invoiced revenue", "invoiced", "recurring revenue", "mrr", "arr", "monthly recurring revenue",
                       "annual recurring revenue", "subscription revenue"))
_RANK = {"recognised": 0, "invoiced": 1, "recurring": 2}
# A cumulative or rolling line (issue #136): a running total for the year, the last 12 months, a lifetime or a cumulative figure is
# never the revenue row, whatever its revenue word (the whole label is read, a bracket's content included: "(current year)"). Summing
# such a line month by month multiplies the year.
_CUMULATIVE = re.compile(r"\b(?:ytd|year to date|current year|to date|cumulative|lifetime|ltm|ttm|last (?:12|twelve) months"
                         r"|trailing(?: (?:12|twelve))? months?|rolling|running total)\b")
_VOCABULARY = tuple("""revenue revenues sales turnover mrr arr arpa acv tcv customer customers client clients subscriber subscribers
subscription subscriptions logo logos churn churned retention nrr grr booking bookings booked forecast fcst budget pipeline backlog
contract contracted run-rate annualised annualized cost costs cogs expense expenses opex margin profit ebitda cash headcount fte payroll
salaries growth income invoiced billing billings rent""".split()) + ("run rate",)
_CUSTOMER_WORDS = ("customer", "customers", "client", "clients", "account", "accounts", "company", "companies", "subscriber",
                   "subscribers", "logo", "logos", "buyer", "buyers")
_NAME_HEADERS = ("name", "names")      # a column headed Name is entity-labelled, whatever its values (review L1)
_CATEGORY_WORDS = ("product", "products", "service", "services", "region", "regions", "country", "countries", "segment", "segments",
                   "channel", "channels", "category", "categories", "business unit", "division", "divisions", "geography", "tier",
                   "tiers", "line", "lines")
_AMOUNT_WORDS = ("amount", "net", "gross", "total", "value", "price", "revenue", "sales", "invoiced", "billing")
_CRM_WORDS = ("deal", "opportunity", "pipeline", "quote", "forecast", "stage")
_REFERENCE_WORDS = ("invoice", "transaction", "order", "receipt", "reference", "ref")
_PERIOD_HEADERS = ("month", "period", "date", "fiscal month")
_TOTAL_WORD = re.compile(r"(?i)\b(?:sub[\s-]?total|grand[\s-]?total|total|totals|sum)\b")
EXCLUDED_TAGS = ("forecast", "budget")          # the period columns tagged so are never read (section 6.2, review H5)
_CURRENCY_CODE = re.compile(r"(?<![A-Za-z])(EUR|USD|GBP|CHF|SEK|NOK|DKK|PLN)(?![A-Za-z])")
_SYMBOLS = {"€": "EUR", "£": "GBP", "$": "USD"}
_YEAR_IN_TEXT = re.compile(r"(?<!\d)(?:19|20)\d{2}(?!\d)")


def _clean(text) -> str:
    """A label as the vocabulary reads it: a bracket's content left out, every run of punctuation one space, lower case."""
    return cr._norm(re.sub(r"\([^)]*\)|\[[^\]]*\]", " ", str(text or "")))


def _has(clean: str, words) -> bool:
    padded = f" {clean} "
    return any(f" {w} " in padded for w in words)


def non_revenue_kind(label) -> Optional[str]:
    """The kind of a label that is recognised but never used as revenue (section 2.1), first match in the order of NON_REVENUE_KINDS.
    The whole label is read, a bracket's content included: "Revenue (forecast)" is a forecast row (review H1)."""
    clean = cr._norm(str(label or ""))
    return next((kind for kind, words in _KIND_WORDS if _has(clean, words)), None)


def revenue_kind(label) -> Optional[str]:
    """recurring, invoiced or recognised when the label holds a revenue word; None when it holds none (section 2.1)."""
    clean = _clean(label)
    if _has(clean, _RECURRING):
        return "recurring"
    if _has(clean, _INVOICED):
        return "invoiced"
    return "recognised" if _has(clean, _RECOGNISED) else None


def is_candidate_label(label) -> bool:
    """A label that may be the revenue row: a revenue word, none of the exclusion words, no "%" and no cumulative or rolling cue."""
    return revenue_kind(label) is not None and not _has(_clean(label), _EXCLUDED) and "%" not in str(label or "") \
        and non_revenue_kind(label) is None and not _CUMULATIVE.search(cr._norm(str(label or "")))


def label_score(label) -> int:
    clean = re.sub(r"^(?:total|sum)\s+", "", _clean(label))
    return 100 if clean in _SYNONYMS else 80


def _is_turnover(label) -> bool:
    return _has(_clean(label), ("turnover",))


def _month_key(v) -> Optional[str]:
    if cr._blank(v):
        return None
    if isinstance(v, (datetime, date, pd.Timestamp)):
        return f"{v.year:04d}-{v.month:02d}"
    text = str(v).strip()
    found = cr.parse_period(text)
    if found and found[0] == "month":
        return f"{found[1].year:04d}-{found[1].month:02d}"
    parsed = pd.to_datetime(text, errors="coerce", dayfirst=False)
    return None if pd.isna(parsed) else f"{parsed.year:04d}-{parsed.month:02d}"


def _num(v) -> Optional[float]:
    return sheet_mod.number(v)


def _cell(grid, row: int, col: int):
    r = grid[row - 1] if 0 < row <= len(grid) else []
    return r[col - 1] if 0 < col <= len(r) else None


def _codes(texts) -> Set[str]:
    out: Set[str] = set()
    for text in texts:
        text = str(text or "")
        out |= {m.upper() for m in _CURRENCY_CODE.findall(text)}
        out |= {code for sign, code in _SYMBOLS.items() if sign in text}
    return out


_AMOUNT_WITH_CURRENCY = re.compile(
    r"[€£$]\s*[-(]?\d|\d\s*[€£$]|(?<![A-Za-z])(?:EUR|USD|GBP|CHF|SEK|NOK|DKK|PLN)\s*[-(]?\d|\d\s*(?:EUR|USD|GBP|CHF|SEK|NOK|DKK|PLN)(?![A-Za-z])")


def currency_scan(sheet) -> Set[str]:
    """Every currency a sheet shows, in any layout (re-check, owner decision on the currency): its title rows and headers, the sign or
    code written in an amount cell ("€1,200", "EUR 1200", "1,200.00 EUR") and the title currency the sheet carries. A text cell that
    is no amount (a customer called "USD Logistics") shows none."""
    out = _codes(_title_texts(sheet) + [str(h or "") for h in sheet.headers])
    if getattr(sheet, "currency", None):
        out.add(sheet.currency)
    cells = (c for row in sheet.grid for c in row) if sheet.grid else (c for row in sheet.frame.itertuples(index=False) for c in row)
    for cell in cells:
        if isinstance(cell, str) and _AMOUNT_WITH_CURRENCY.search(cell):
            out |= _codes([cell])
    return out


def _close(a: Optional[float], b: float) -> bool:
    if a is None:
        return True
    return abs(a - b) <= TOTAL_TOLERANCE * abs(b) if b else a == 0


# ---------------------------------------------------------------------------
# Reading a sheet: the layout and the facts the rules need (independent of any proposal)
# ---------------------------------------------------------------------------
def _title_texts(sheet) -> List[str]:
    return [t for _, _, t in sheet.titles]


def _cost_groups(texts) -> int:
    labels = [cr._norm(t) for t in texts]
    return len({i for i, rx in enumerate(cr._COST_LINES) if any(rx.search(t) for t in labels)})


_METRIC_WORDS = tuple(w for w in _VOCABULARY if w not in _CUSTOMER_WORDS)


def _metric_labels(lines: List[Dict]) -> bool:
    """Decision 4 of 2026-10-19 (issue #136): the labels under a customer-style header read as metric names when their share of labels
    holding a vocabulary word other than a customer word passes the metric threshold and one of them is a revenue candidate ("Customer 1"
    is a customer; "Recognised revenue", "ARR", "Forecast revenue" are metrics)."""
    hits = sum(1 for ln in lines if _has(_clean(ln["label"]), _METRIC_WORDS))
    return bool(lines) and hits / len(lines) > METRIC_LABEL_SHARE and hits >= MIN_METRIC_HITS \
        and any(is_candidate_label(ln["label"]) for ln in lines)


def _label_mode(lines: List[Dict], customer_header: bool) -> Tuple[str, float, int, int]:
    hits = sum(1 for ln in lines if _has(_clean(ln["label"]), _VOCABULARY))
    share = hits / len(lines) if lines else 0.0
    if share > METRIC_LABEL_SHARE and hits >= MIN_METRIC_HITS and not customer_header:
        return "metric", share, hits, round(100 * share)
    if share <= ENTITY_LABEL_SHARE:
        return "entity", share, hits, round(100 * (1 - share))
    return "unsure", share, hits, 0


def _totals(labels: List[str], values: List[List[Optional[float]]]) -> Dict[int, bool]:
    """{index of a total row: whether it equals the sum of the rows it adds up}. A total row is a row whose label holds the word total,
    subtotal, grand total or sum anywhere, or an unlabelled row whose figures equal the sum of two or more rows since the previous total,
    or of two or more totals above it (section 6.1, review H3); one row above is never enough."""
    out: Dict[int, bool] = {}
    block: List[int] = []                    # the rows since the previous total
    rows: List[int] = []                     # every row used so far
    for i, label in enumerate(labels):
        figures = [k for k, v in enumerate(values[i]) if v is not None]

        def adds_up(group: List[int]) -> bool:
            return bool(group and figures) and all(_close(values[i][k], sum(values[j][k] or 0 for j in group)) for k in figures)

        sums = [g for g in (block, rows, list(out)) if g]
        labelled = bool(_TOTAL_WORD.search(label))
        by_value = (len(block) >= 2 and adds_up(block)) or (len(out) >= 2 and adds_up(list(out)))
        if labelled or by_value:
            out[i] = any(adds_up(g) for g in sums)
            block = []
        else:
            block.append(i)
            rows.append(i)
    return out


def _header_cell(grid, header_row: int, col: int) -> str:
    return sheet_mod.cell_text(_cell(grid, header_row, col)) if not cr._blank(_cell(grid, header_row, col)) else ""


LIST_MIN_ROWS = 3                   # a dated list holds at least three dated rows
SOURCE_TOLERANCE = 0.01             # the dated list sums to the revenue row month by month within 1% (section 6.7)
SERIES_SOURCES = ("metrics_block", "dated_list")


def _dated_lists(grid) -> List[Dict]:
    """The dated lists of a sheet (section 6.7): a header reading Date, Month or Period beside a text header, and below it at least three
    rows holding a date and a number. A list is read by its rows: its dates and values are summed by month."""
    found = []
    for r, row in enumerate(grid, 1):
        for c, cell in enumerate(row, 1):
            if not isinstance(cell, str) or _clean(cell) not in _PERIOD_HEADERS:
                continue
            beside = _cell(grid, r, c + 1)
            if not isinstance(beside, str) or not beside.strip() or cr.parse_period(beside.strip()):
                continue
            rows, rr = [], r + 1
            while rr <= len(grid) and _month_key(_cell(grid, rr, c)) and _num(_cell(grid, rr, c + 1)) is not None:
                rows.append((rr, _month_key(_cell(grid, rr, c)), _num(_cell(grid, rr, c + 1))))
                rr += 1
            if len(rows) >= LIST_MIN_ROWS:
                months: Dict[str, float] = {}
                month_rows: Dict[str, List[int]] = {}
                for row_no, month, value in rows:
                    months[month] = months.get(month, 0.0) + value
                    month_rows.setdefault(month, []).append(row_no)
                found.append({"header_row": r, "date_col": c, "value_col": c + 1, "first_row": rows[0][0], "last_row": rows[-1][0],
                              "rows": [x[0] for x in rows], "months": months, "month_rows": month_rows})
    return found


def _labels_info(sheet, rd: Dict, fye: int) -> Dict:
    grid = sheet.grid
    periods = []
    for p in rd["periods"]:
        found = sheet_mod.period_of(p["header"], fye)
        periods.append({"col": p["col"], "header": p["header"], "grain": p["grain"], "tag": p["tag"],
                        "end": found["end"] if found else None})
    pcols = [p["col"] for p in periods]
    lists = _dated_lists(grid)
    listed = {r for x in lists if x["date_col"] == rd["label_col"] for r in [x["header_row"]] + x["rows"]}
    lines = []
    for ln in rd["lines"]:
        if ln["row"] in listed:
            continue                                          # the rows of a dated list under the labels are another table
        cells = [_cell(grid, ln["row"], c) for c in pcols]
        if all(cr._blank(c) for c in cells):
            continue                                          # a note or a heading: not a label row (section 4.2)
        lines.append({"row": ln["row"], "label": ln["label"], "cells": cells, "values": [_num(c) for c in cells]})
    head = _header_cell(grid, rd["header_row"], rd["label_col"])
    customer_header = _has(_clean(head), _CUSTOMER_WORDS) or _clean(head) in _NAME_HEADERS
    mode, share, hits, confidence = _label_mode(lines, customer_header)
    # Decision 4 of 2026-10-19 (issue #136): the header says customer but the labels say metrics; the analyst chooses the type, and
    # until then the rows are never summed as customers.
    metric_labels = customer_header and _metric_labels(lines)
    marked = [{"row": ln["row"], "label": ln["label"], "kind": non_revenue_kind(ln["label"])} for ln in lines
              if (mode == "metric" or metric_labels) and non_revenue_kind(ln["label"])]
    texts = [t for _, _, t in rd["titles"]] + [h for h in [head] if h] + [p["header"] for p in periods]
    return {"layout": "labels", "sheet": sheet, "rd": rd, "periods": periods, "lines": lines, "lists": lists, "mode": mode, "share": share,
            "hits": hits, "confidence": confidence, "head": head, "marked": marked, "header_row": rd["header_row"],
            "metric_labels": metric_labels,
            "label_col": rd["label_col"], "codes": _codes(texts + [ln["label"] for ln in lines]),
            "titles": [t for _, _, t in rd["titles"]]}


def _frame_rows(sheet) -> List[Tuple[int, List]]:
    cols = list(sheet.columns)
    return [(n, [row[c] for c in cols]) for n, (_, row) in zip(sheet.row_numbers, sheet.frame.iterrows())]


def _date_header(header: str) -> bool:
    return _clean(header) in _PERIOD_HEADERS or cr.field_score("revenue", "invoice_date", header) >= cr.MAPPING_THRESHOLD \
        or bool(cr.parse_period(str(header or "").strip()))


def _numeric_column(rows: List[Tuple[int, List]], idx: int) -> int:
    return sum(1 for _, r in rows if _num(r[idx]) is not None)


def _snapshot_info(sheet) -> Optional[Dict]:
    """A label column and a column headed with a revenue measure, no period header, no date column (section 2)."""
    heads = list(sheet.headers)
    if sheet.wide or len(heads) < 2 or any(_date_header(h) for h in heads):
        return None
    rows = _frame_rows(sheet)
    measure = next((i for i, h in enumerate(heads) if i > 0 and revenue_kind(h) and not _has(_clean(h), _EXCLUDED)
                    and non_revenue_kind(h) is None), None)
    if measure is None or not rows:
        return None
    labels = [str(r[0]).strip() if not cr._blank(r[0]) else "" for _, r in rows]
    texts = [t for t in labels if t]
    if len(texts) * 2 < len(rows) or any(_num(r[0]) is not None for _, r in rows if not cr._blank(r[0])):
        return None
    kind = revenue_kind(heads[measure])
    titles = _title_texts(sheet)
    return {"layout": "snapshot", "sheet": sheet, "measure": measure, "kind": kind, "rows": rows, "labels": labels,
            "header_row": sheet.header_row, "periods": [], "titles": titles,
            "codes": _codes(titles + heads + texts), "measure_col": sheet.first_col + measure,
            "reference": any(_YEAR_IN_TEXT.search(t) or _month_key(t) for t in titles)}


def _columns_info(sheet) -> Optional[Dict]:
    """Dates down the rows, measures across (the accountant P&L): a period column and no invoice, transaction or customer column."""
    heads = list(sheet.headers)
    if sheet.wide:
        return None
    period = next((i for i, h in enumerate(heads) if _clean(h) in _PERIOD_HEADERS), None)
    if period is None:
        return None
    for h in heads:
        if _has(_clean(h), _REFERENCE_WORDS) or cr.field_score("revenue", "customer_id", h) >= cr.MAPPING_THRESHOLD:
            return None
    rows = _frame_rows(sheet)
    dated = [(n, r, _month_key(r[period])) for n, r in rows]
    dated = [d for d in dated if d[2]]
    titles = _title_texts(sheet)
    measures = []
    for i, h in enumerate(heads):
        if i == period or not is_candidate_label(h):
            continue
        fit = sum(1 for _, r, _ in dated if _num(r[i]) is not None) / len(dated) if dated else 0.0
        measures.append({"idx": i, "column": sheet.first_col + i, "label": h, "kind": revenue_kind(h),
                         "confidence": round(label_score(h) * fit), "numbers": sum(1 for _, r, _ in dated if _num(r[i]) is not None)})
    measures.sort(key=lambda m: (_RANK[m["kind"]], -m["confidence"], m["idx"]))
    return {"layout": "columns", "sheet": sheet, "period": period, "dated": dated, "measures": measures, "header_row": sheet.header_row,
            "periods": [], "titles": titles, "codes": _codes(titles + heads),
            "groups": _cost_groups(heads), "months": sorted({m for _, _, m in dated})}


def _records_info(sheet) -> Dict:
    heads = list(sheet.headers)
    rows = _frame_rows(sheet)
    crm = cr.type_shares(sheet).get("crm", 0) >= cr.DETECTION_MIN_SHARE
    amount = next((i for i, h in enumerate(heads) if _has(_clean(h), _AMOUNT_WORDS) and not _has(_clean(h), _CRM_WORDS)
                   and non_revenue_kind(h) is None and not _has(_clean(h), _EXCLUDED) and _numeric_column(rows, i)), None)
    if crm:
        amount = None
    date_col = next((i for i, h in enumerate(heads) if _date_header(h) and cr.value_fit("date", [r[i] for _, r in rows])[0] > 0), None)
    reference = next((i for i, h in enumerate(heads) if _has(_clean(h), _REFERENCE_WORDS)), None)
    customer = next((i for i, h in enumerate(heads) if cr.field_score("revenue", "customer_id", h) >= cr.MAPPING_THRESHOLD), None)
    currency_col = next((i for i, h in enumerate(heads) if cr.field_score("revenue", "currency", h) >= cr.MAPPING_THRESHOLD), None)
    titles = _title_texts(sheet)
    seen = set()
    if currency_col is not None:
        seen = {str(r[currency_col]).strip().upper() for _, r in rows if isinstance(r[currency_col], str)
                and re.fullmatch(r"[A-Za-z]{3}", r[currency_col].strip())}
    months = cr.months_of([r[date_col] for _, r in rows]) if date_col is not None else None
    return {"layout": "records", "sheet": sheet, "amount": amount, "date": date_col, "reference": reference, "customer": customer,
            "header_row": sheet.header_row, "periods": [], "titles": titles, "currency_values": seen,
            "codes": _codes(titles + heads) | (seen if len(seen) == 1 else set()), "month_count": (months or {}).get("count", 0),
            "multi_currency": len(seen) > 1, "rows": rows}


def analyse(sheet, fye: int = 12) -> Dict:
    """The layout of one sheet and the facts its rules need, in the fixed reading order: labels, snapshot, columns, records."""
    rd = sheet_mod.read(sheet.grid, sheet.fmt, fye) if sheet.grid else None
    info = _labels_info(sheet, rd, fye) if rd is not None else (_snapshot_info(sheet) or _columns_info(sheet) or _records_info(sheet))
    shown = currency_scan(sheet)      # the same in every layout: two currencies anywhere in the file are mixed currencies
    info["codes"] = set(info["codes"]) | shown
    info["currency_values"] = set(info.get("currency_values") or ()) | info["codes"]
    info["multi_currency"] = len(info["currency_values"]) > 1
    return info


# ---------------------------------------------------------------------------
# Candidates for the revenue row
# ---------------------------------------------------------------------------
def _candidates(info: Dict, out: Set[int], metric: Optional[bool] = None) -> List[Dict]:
    """The unmarked candidate rows of a metric-labelled sheet (section 2.1), ordered by rank, confidence and sheet order. `metric`
    reads the rows as metric labels whatever the sheet's mode (a customer-style header over metric labels, decision 4 of 2026-10-19)."""
    if not (info["mode"] == "metric" if metric is None else metric):
        return []
    ncols = len(info["periods"]) or 1
    found = []
    for ln in info["lines"]:
        if ln["row"] in out or not is_candidate_label(ln["label"]):
            continue
        numbers = sum(1 for v in ln["values"] if v is not None)
        kind = revenue_kind(ln["label"])
        found.append({"row": ln["row"], "label": ln["label"], "kind": kind, "confidence": round(label_score(ln["label"]) * numbers / ncols),
                      "numbers": numbers, "turnover": _is_turnover(ln["label"])})
    return sorted(found, key=lambda c: (_RANK[c["kind"]], -c["confidence"], c["row"]))


def _public(c: Dict) -> Dict:
    return {"row": c["row"], "label": c["label"], "kind": c["kind"], "confidence": c["confidence"]}


def _same_key(a: Dict, b: Dict) -> bool:
    return (_RANK[a["kind"]], a["confidence"]) == (_RANK[b["kind"]], b["confidence"])


# ---------------------------------------------------------------------------
# The result
# ---------------------------------------------------------------------------
def _asof(as_of_month) -> Optional[str]:
    text = str(as_of_month or "").strip()
    return text[:7] if re.match(r"^\d{4}-\d{2}", text) else None


def _type_of_labels(info: Dict) -> Optional[str]:
    if info["mode"] == "metric":
        return "financial_statements" if _cost_groups([ln["label"] for ln in info["lines"]]) >= 2 else "monthly_revenue_report"
    head = _clean(info["head"])
    if _has(head, _CUSTOMER_WORDS):
        return "customer_revenue"
    return "revenue_summary" if _has(head, _CATEGORY_WORDS) else None


def _supports(layout: str, ftype: Optional[str], periods: int, customer: bool, labels_without_month: bool) -> Tuple[List[str], List[Dict]]:
    if ftype is None:
        return [], []
    if ftype == "arr_mrr_snapshot":
        return ["claim_matching"], [{"analysis": "customer_level", "needs": NEEDS_BY_CUSTOMER},
                                    {"analysis": "trend", "needs": NEEDS_SNAPSHOT_SERIES}]
    supports, blocked = [], []
    if ftype == "customer_revenue" or (ftype == "invoice_transactions" and customer):
        supports.append("customer_level")
    else:
        blocked.append({"analysis": "customer_level", "needs": NEEDS_CUSTOMER_COLUMN if ftype == "invoice_transactions" else NEEDS_BY_CUSTOMER})
    if ftype == "invoice_transactions" or periods >= 2:
        supports.append("trend")
    else:
        blocked.append({"analysis": "trend", "needs": NEEDS_PERIODS})
    if labels_without_month:
        blocked.append({"analysis": "monthly_series", "needs": NEEDS_MONTHLY})
    return supports, blocked


def missing_lines(result: Dict) -> List[str]:
    """One line per blocked analysis for the data inventory (section 3.2): never Missing without the data."""
    return [f"{ANALYSIS_NAMES[b['analysis']]} – Missing: {b['needs']}" for b in result.get("blocked") or []]


def growth_label(result: Dict) -> Optional[str]:
    """The name of a growth figure read from this result (section 6.1): never "revenue growth" for an MRR or ARR row."""
    if result.get("series_basis") in (None, "snapshot"):
        return None
    row = result.get("revenue_row") or result.get("revenue_column")
    if row is None:
        return None if result.get("series_basis") in ("row", "column") else "Revenue growth"
    if row.get("kind") != "recurring":
        return "Revenue growth"
    clean = _clean(row.get("label"))
    if _has(clean, ("mrr", "monthly recurring revenue")):
        return "MRR growth"
    return "ARR growth" if _has(clean, ("arr", "annual recurring revenue")) else "Recurring revenue growth"


def choices_waiting(result: Dict) -> int:
    """The parts in click: type, revenue row, last actual month, currency, money kind. A turnover row's row and money kind are one choice."""
    n = 1 if result.get("type_state") == "click" else 0
    row = result.get("revenue_row") or result.get("revenue_column")
    asked = (result.get("money_kind") or {}).get("asked") and "answer" not in result["money_kind"]
    if row and row["state"] == "click" and row.get("reason") != "turnover_question":
        n += 1
    n += 1 if asked else 0
    n += 1 if (result.get("last_actual") or {}).get("state") == "click" else 0
    n += 1 if (result.get("currency") or {}).get("state") == "confirm" else 0
    n += 1 if (result.get("series_source") or {}).get("state") == "click" else 0
    return n


def _blank_result(info: Dict, missing: List[str], marked: List[Dict]) -> Dict:
    return {"recognised": False, "missing": missing, "message": message_for(missing), "sheet": info["sheet"].sheet,
            "header_row": info.get("header_row"), "file_type": None, "type_state": None, "type_source": None, "type_candidates": [],
            "type_reason": None, "label_mode": info.get("mode"), "label_share": info.get("share"), "confidence": info.get("confidence"),
            "series_basis": None, "periods": 0, "revenue_row": None, "revenue_column": None, "last_actual": None,
            "excluded_columns": [], "currency": {"code": None, "state": "confirm"}, "money_kind": {"asked": False},
            "non_revenue_rows": marked, "supports": [], "blocked": [], "other_sheets": [], "needs_click": False, "series": [],
            "series_rows": [], "total_rows": [], "tables": [], "series_source": None}


def message_for(missing: List[str]) -> Optional[str]:
    if not missing:
        return None
    return "Not a revenue file. Missing: " + "; ".join(MISSING_NAMES[m] for m in missing) + "."


def _currency(info: Dict, ftype: Optional[str], reporting_currency: Optional[str], decided: Optional[str]) -> Dict:
    if info.get("multi_currency"):            # several currencies in one file: no single currency to confirm (review L3)
        return {"code": None, "state": "confirm", "mixed": True, "codes": sorted(info.get("currency_values") or [])}
    if decided in ISO_CURRENCIES:
        return {"code": decided, "state": "you"}
    codes = info["codes"]
    if len(codes) == 1 and not info.get("multi_currency"):
        return {"code": next(iter(codes)), "state": "seen"}
    if ftype == "financial_statements" and not codes and reporting_currency:
        return {"code": reporting_currency, "state": "audit"}
    return {"code": None, "state": "confirm"}


def _boundary(info: Dict, row_values: Optional[List[Optional[float]]], as_of: Optional[str], proposed: Optional[int],
              decided: Optional[int]) -> Dict:
    """The last actual column and the columns after it (section 6.2); `row_values` are the revenue row's values by period."""
    periods = info["periods"]
    dated = [p for p in periods if p["end"]]
    usable = [p for p in dated if p["tag"] not in EXCLUDED_TAGS]      # a forecast or budget column is never the last actual column
    by_col = {p["col"]: p for p in dated}
    if not usable:
        return {"column": None, "month": None, "state": "click", "reason": "not_a_period", "candidates": []}
    readable = {p["col"] for p in usable}
    actual = [p for p in usable if p["tag"] == "actual"]
    if actual:
        rules = actual[-1]["col"]
    else:
        have = [p["col"] for k, p in enumerate(periods) if p in usable and row_values is not None and row_values[k] is not None]
        rules = have[-1] if have else usable[-1]["col"]
    candidates = [{"column": p["col"], "month": p["end"]} for p in usable if as_of is None or p["end"] <= as_of]
    column, state, reason = rules, "auto", None
    if decided in readable:
        column, state = decided, "you"
        if as_of is not None and by_col[decided]["end"] > as_of:
            state, reason = "click", "after_as_of"
    else:
        if proposed is not None:
            if proposed in by_col:
                column = proposed
                if proposed not in readable and (as_of is None or by_col[proposed]["end"] <= as_of):
                    column, reason = rules, "not_a_period"
            else:
                reason = "not_a_period"
        if reason is None and as_of is None:
            reason = "no_as_of_month"
        if reason is None and by_col[column]["end"] > as_of:
            reason = "after_as_of"
        if reason is None and proposed is not None and proposed != rules:
            reason = "rules_and_model_disagree"
        state = "click" if reason else "auto"
    return {"column": column, "month": by_col[column]["end"], "state": state, "reason": reason, "candidates": candidates}


def _row_decision(cands: List[Dict], info: Dict, prop_row: Optional[Dict], marked: Dict[int, str], periods: int, volume: Set[int],
                  decided: Optional[int], money: Optional[str]) -> Dict:
    """revenue_row (section 6.1): the rules' first candidate, the model's proposal checked in order, the analyst's decision."""
    if not cands:
        return {"row": None, "label": None, "kind": None, "confidence": None, "state": "click", "reason": "no_candidate",
                "candidates": []}
    first = cands[0]
    by_row = {c["row"]: c for c in cands}
    reason = None
    if prop_row is not None:
        lines = {ln["row"]: ln for ln in info["lines"]}
        ln = lines.get(prop_row["row"])
        need = min(3, periods)
        if ln is None:
            reason = "row_missing"
        elif prop_row["row"] in marked:
            reason = "marked_not_revenue"
        elif not _label_matches(ln["label"], prop_row):
            reason = "label_mismatch"
        elif sum(1 for v in ln["values"] if v is not None) < need:
            reason = "not_numeric_enough"
        elif revenue_kind(ln["label"]) is None or not is_candidate_label(ln["label"]):
            reason = "no_revenue_word"
        elif prop_row["row"] != first["row"]:
            reason = "rules_and_model_disagree"
    chosen = by_row.get(decided, first)
    if decided in by_row:
        state, reason = "you", None
    else:
        if reason is None:
            if first["turnover"]:
                reason = "turnover_question"
            elif len(cands) > 1 and _same_key(cands[0], cands[1]):
                reason = "tie"
            elif first["confidence"] <= cr.MAPPING_THRESHOLD:
                reason = "confidence_80"
        state = "click" if reason else "auto"
    if chosen["turnover"]:
        if money == "revenue":
            state, reason = "you", None
        elif money is None and state != "click":
            state, reason = "click", "turnover_question"
    return {**_public(chosen), "state": state, "reason": reason, "candidates": [_public(c) for c in cands]}


def _label_matches(label: str, prop_row: Dict) -> bool:
    if "label_matches" in prop_row:
        return bool(prop_row["label_matches"])
    norm = lambda t: re.sub(r"\s+", " ", str(t or "")).strip().casefold()  # noqa: E731
    return norm(label) == norm(prop_row.get("label"))


def _model_marks(info: Dict, prop: Optional[Dict], rules_marked: List[Dict]) -> List[Dict]:
    """The rows the rules mark, then the rows only the model marks (section 2.1: the union; a row either side marks is out)."""
    out = list(rules_marked)
    if not prop or info["mode"] != "metric":
        return out
    seen = {m["row"] for m in out}
    lines = {ln["row"]: ln for ln in info["lines"]}
    for item in prop.get("non_revenue_rows") or []:
        if item["row"] in lines and item["row"] not in seen:
            seen.add(item["row"])
            out.append({"row": item["row"], "label": lines[item["row"]]["label"], "kind": item["kind"]})
    return sorted(out, key=lambda m: m["row"])


def read_periods(info: Dict, last_col: Optional[int]) -> List[Tuple[int, Dict]]:
    """(index, period) of the columns a series reads: a dated period up to the boundary, not tagged forecast or budget, and one column
    per month, the first (section 6.2, review H5)."""
    out, seen = [], set()
    for k, p in enumerate(info["periods"]):
        if not p["end"] or p["tag"] in EXCLUDED_TAGS or (last_col is not None and p["col"] > last_col):
            continue
        if p["grain"] == "month" and p["end"] in seen:
            continue
        seen.add(p["end"])
        out.append((k, p))
    return out


def _sum_series(info: Dict, last_col: Optional[int], marked: Set[int]) -> Tuple[List[Dict], List[Dict], List[int]]:
    """(series, total_rows, rows used) of a sheet whose revenue is the sum of its rows (products, regions, customers): total rows, rows
    marked forecast, budget, bookings, pipeline, contract value or run-rate and balance rows (deferred, unbilled, accrued) are never
    added (section 6.1, review H2 and H3). A metrics sheet never sums: it reads its revenue row (re-check, owner decision)."""
    lines = [ln for ln in info["lines"] if ln["row"] not in marked and not non_revenue_kind(ln["label"])
             and not _has(_clean(ln["label"]), _BALANCES)]
    values = [ln["values"] for ln in lines]
    totals = _totals([ln["label"] for ln in lines], values)
    used = [i for i in range(len(lines)) if i not in totals]
    series = []
    for k, p in read_periods(info, last_col):
        figures = [values[i][k] for i in used if values[i][k] is not None]
        if figures:
            series.append({"column": p["col"], "month": p["end"], "grain": p["grain"], "value": _round(sum(figures))})
    return (series, [{"row": lines[i]["row"], "label": lines[i]["label"], "matches": ok} for i, ok in sorted(totals.items())],
            [lines[i]["row"] for i in used])


def _round(v: float):
    return int(v) if float(v).is_integer() else round(v, 6)


def _finalize_labels(info: Dict, ctx: Dict) -> Dict:
    prop, dec = ctx.get("proposal"), ctx.get("decisions") or {}
    marked_list = _model_marks(info, prop, info["marked"])
    marked = {m["row"]: m["kind"] for m in marked_list}
    money = dec.get("money_kind")
    out = set(marked)
    volume: Set[int] = set()
    mode = info["mode"]
    chosen_type = dec.get("file_type") if dec.get("file_type") in RECOGNITION_TYPES else None
    # Decision 4 of 2026-10-19 (issue #136): metric labels under a customer-style header are read as metrics until the analyst says
    # the rows are customers; "Monthly revenue report" makes the sheet a metrics sheet.
    unsure_metrics = bool(info.get("metric_labels")) and mode != "metric"
    if unsure_metrics and chosen_type == "monthly_revenue_report":
        mode = "metric"
    as_metrics = mode == "metric" or (unsure_metrics and chosen_type is None)
    cands = _candidates(info, out, as_metrics)
    turnover_first = bool(cands and cands[0]["turnover"])
    if money == "volume" and turnover_first:
        volume = {c["row"] for c in cands if c["turnover"]}
        cands = _candidates(info, out | volume, as_metrics)
    labels_text = _title_texts(info["sheet"]) + ([info["head"]] if info["head"] else [])
    a_ok = bool(cands) if as_metrics else (chosen_type is not None or any(revenue_kind(t) for t in labels_text))
    anynum = any(v is not None for ln in info["lines"] for v in ln["values"])
    c_ok = any(c["numbers"] for c in cands) if (as_metrics and cands) else anynum
    missing = [name for name, ok in (("revenue_measure", a_ok), ("reference", True), ("numbers", c_ok)) if not ok]
    if missing:
        return _blank_result(info, missing, marked_list)
    periods = [p for p in info["periods"] if p["end"]]
    rules_type = None if (unsure_metrics and chosen_type is None) else _type_of_labels({**info, "mode": mode})
    model_type = prop["file_type"] if prop and prop.get("file_type") in RECOGNITION_TYPES else None
    ftype, tstate, treason, tsource, tcands = rules_type, "auto", None, "rules", []
    if rules_type:
        if model_type and model_type != rules_type:
            tstate, treason, tcands = "click", "rules_and_model_disagree", [rules_type, model_type]
    else:
        tstate, tcands = "click", (["customer_revenue", "monthly_revenue_report"] if unsure_metrics else ["customer_revenue", "revenue_summary"])
        if model_type in tcands:
            ftype, treason, tsource = model_type, "type_unsure", "ai"
    if dec.get("file_type") in RECOGNITION_TYPES:
        ftype, tstate, treason, tsource, tcands = dec["file_type"], "auto", None, "you", []
        tstate = "auto"
    has_row = ftype in ("monthly_revenue_report", "financial_statements") or mode == "metric"      # a metrics sheet never sums its rows
    revenue_row, row_values = None, None
    if has_row:
        revenue_row = _row_decision(cands, info, (prop or {}).get("revenue_row") if mode == "metric" else None, marked, len(periods),
                                    volume, dec.get("revenue_row"), money)
        if revenue_row["row"] is not None:
            row_values = next(ln["values"] for ln in info["lines"] if ln["row"] == revenue_row["row"])
    last = _boundary(info, row_values, ctx.get("as_of"), (prop or {}).get("last_actual_column"), dec.get("last_actual"))
    last_col = last["column"]
    read = {p["col"] for _, p in read_periods(info, last_col)}
    excluded = [p["col"] for p in info["periods"] if p["end"] and p["col"] not in read]
    if has_row:
        series = []
        if revenue_row["row"] is not None:
            for k, p in read_periods(info, last_col):
                if row_values[k] is not None:
                    series.append({"column": p["col"], "month": p["end"], "grain": p["grain"], "value": _round(row_values[k])})
        basis, total_rows = "row", []
        series_rows = [revenue_row["row"]] if revenue_row["row"] is not None else []
    else:
        series, total_rows, series_rows = _sum_series(info, last_col, set(marked))
        basis = "sum_of_rows"
    source = None
    if has_row and revenue_row["row"] is not None and info["lists"]:
        source = _series_source(info, revenue_row, row_values, last, dec.get("series_source"))
        if source["supplier"] == "dated_list":
            listed = info["lists"][0]
            limit = last["month"] if last["month"] else None
            series = [{"column": None, "month": m, "grain": "month", "value": _round(v)} for m, v in sorted(listed["months"].items())
                      if limit is None or m <= limit]
            basis, series_rows = "column", sorted(r for x in series for r in listed["month_rows"][x["month"]])      # rows after the boundary are not cited
    month_cols = [p for p in periods if p["grain"] == "month"]
    supports, blocked = _supports("labels", ftype, len(periods), False, bool(periods) and not month_cols)
    currency = _currency(info, ftype, ctx.get("reporting_currency"), dec.get("currency"))
    asked = turnover_first or bool(revenue_row and revenue_row["row"] is not None and _is_turnover(revenue_row["label"]))
    money_kind = {"asked": asked}
    if asked and money in ("revenue", "volume"):
        money_kind["answer"] = money
    result = {
        "recognised": True, "missing": [], "message": None, "sheet": info["sheet"].sheet, "header_row": info["header_row"],
        "file_type": ftype, "type_state": tstate, "type_source": tsource, "type_candidates": tcands, "type_reason": treason,
        "label_mode": mode, "label_share": info["share"], "confidence": info["confidence"], "series_basis": basis,
        "periods": len(periods), "revenue_row": revenue_row, "revenue_column": None, "last_actual": last,
        "excluded_columns": excluded, "currency": currency, "money_kind": money_kind, "non_revenue_rows": marked_list,
        "supports": supports, "blocked": blocked, "other_sheets": [], "series": series, "total_rows": total_rows,
        "series_rows": series_rows, "tables": _tables(info, ftype), "series_source": source}
    return _finish(result)


def _tables(info: Dict, ftype: Optional[str]) -> List[Dict]:
    """The tables a sheet holds (section 6.7): the metrics block, and each dated list, each classified. Positions only, no text."""
    if not info["lists"]:
        return []
    block = {"kind": "metrics_block", "file_type": ftype, "header_row": info["header_row"], "label_col": info["label_col"]}
    return [block] + [{"kind": "dated_list", "file_type": "monthly_revenue_report", "header_row": x["header_row"], "date_col": x["date_col"],
                       "value_col": x["value_col"], "first_row": x["first_row"], "last_row": x["last_row"], "months": len(x["months"])}
                      for x in info["lists"]]


def _series_source(info: Dict, row: Dict, row_values: Optional[List[Optional[float]]], last: Dict, decided: Optional[str]) -> Dict:
    """Which table supplies the revenue series (section 6.7). A metrics row of recognised revenue wins; when the dated list sums to it
    month by month within 1% the list is a cross-check, not a second source. When they differ, or the row is not recognised revenue,
    the analyst chooses."""
    listed = info["lists"][0]
    limit = last.get("month")
    pairs = [(p, row_values[k]) for k, p in enumerate(info["periods"])
             if p["grain"] == "month" and p["end"] and p["tag"] not in EXCLUDED_TAGS and row_values[k] is not None
             and (limit is None or p["end"] <= limit)]
    gaps, off = [], []
    for p, value in pairs:
        have = listed["months"].get(p["end"])
        gap = None if have is None else abs(have - value) / abs(value) if value else (0.0 if have == 0 else 1.0)
        off.append(gap)
        if gap is None or gap > SOURCE_TOLERANCE:         # both figures and both sources, for each month that differs (review M4)
            found = listed["month_rows"].get(p["end"]) or []
            col = get_column_letter(listed["value_col"])
            gaps.append({"month": p["end"], "row_value": _round(value), "list_value": None if have is None else _round(have),
                         "row_cell": f"{get_column_letter(p['col'])}{row['row']}",
                         "list_cells": (f"{col}{min(found)}" if len(set(found)) < 2 else f"{col}{min(found)}:{col}{max(found)}")
                         if found else None})
    matched = bool(pairs) and all(g is not None and g <= SOURCE_TOLERANCE for g in off)
    cross = {"status": "match" if matched else "mismatch", "months": len(pairs),
             "months_off": sum(1 for g in off if g is None or g > SOURCE_TOLERANCE),
             "max_gap_pct": round(max((g for g in off if g is not None), default=0.0), 4)}
    wins = row.get("kind") == "recognised" and matched
    if decided in SERIES_SOURCES:
        return {"supplier": decided, "state": "you", "cross_check": cross, "gaps": gaps}
    return {"supplier": "metrics_block", "state": "auto" if wins else "click", "cross_check": cross, "gaps": gaps}


def _finish(result: Dict) -> Dict:
    result.setdefault("tables", [])
    result.setdefault("series_source", None)
    row = result.get("revenue_row") or result.get("revenue_column")
    last = result.get("last_actual")
    source = result.get("series_source")
    result["needs_click"] = bool(
        result["type_state"] == "click" or (row and row["state"] == "click") or (last and last["state"] == "click")
        or (source and source["state"] == "click")
        or result["currency"]["state"] == "confirm" or (result["money_kind"].get("asked") and "answer" not in result["money_kind"]))
    return result


def _finalize_snapshot(info: Dict, ctx: Dict) -> Dict:
    dec = ctx.get("decisions") or {}
    rows = info["rows"]
    m = info["measure"]
    figures = [_num(r[m]) for _, r in rows]
    missing = [name for name, ok in (("revenue_measure", True), ("reference", info["reference"]),
                                     ("numbers", any(v is not None for v in figures))) if not ok]
    if missing:
        return _blank_result(info, missing, [])
    rules_type = "arr_mrr_snapshot" if info["kind"] == "recurring" else "revenue_summary"
    ftype = dec["file_type"] if dec.get("file_type") in RECOGNITION_TYPES else rules_type
    totals = _totals(info["labels"], [[v] for v in figures])
    value = sum(v for i, v in enumerate(figures) if v is not None and i not in totals)
    supports, blocked = _supports("snapshot", ftype, 1, False, False)
    result = {
        "recognised": True, "missing": [], "message": None, "sheet": info["sheet"].sheet, "header_row": info["header_row"],
        "file_type": ftype, "type_state": "auto", "type_source": "you" if dec.get("file_type") else "rules", "type_candidates": [],
        "type_reason": None, "label_mode": "entity", "label_share": None, "confidence": None, "series_basis": "snapshot", "periods": 1,
        "revenue_row": None, "revenue_column": None, "last_actual": None, "excluded_columns": [],
        "currency": _currency(info, ftype, ctx.get("reporting_currency"), dec.get("currency")), "money_kind": {"asked": False},
        "non_revenue_rows": [], "supports": supports, "blocked": blocked, "other_sheets": [],
        "series": [{"column": info["measure_col"], "month": None, "grain": None, "value": _round(value)}],
        "series_rows": [rows[i][0] for i in range(len(rows)) if i not in totals and figures[i] is not None],
        "total_rows": [{"row": rows[i][0], "label": info["labels"][i], "matches": ok} for i, ok in sorted(totals.items())]}
    return _finish(result)


def _finalize_columns(info: Dict, ctx: Dict) -> Dict:
    dec, as_of = ctx.get("decisions") or {}, ctx.get("as_of")
    measures, dated = info["measures"], info["dated"]
    money = dec.get("money_kind")
    volume = set()
    if money == "volume" and measures and _is_turnover(measures[0]["label"]):
        volume = {m["idx"] for m in measures if _is_turnover(m["label"])}
        measures = [m for m in measures if m["idx"] not in volume]
    missing = [name for name, ok in (("revenue_measure", bool(measures)), ("reference", True),
                                     ("numbers", any(m["numbers"] for m in measures) if measures else bool(dated))) if not ok]
    if missing:
        return _blank_result(info, missing, [])
    first = measures[0]
    by_col = {m["column"]: m for m in measures}
    chosen = by_col.get(dec.get("revenue_row"), first)
    reason = None
    if chosen["column"] != dec.get("revenue_row"):
        if _is_turnover(first["label"]):
            reason = "turnover_question"
        elif len(measures) > 1 and (_RANK[measures[0]["kind"]], measures[0]["confidence"]) == \
                (_RANK[measures[1]["kind"]], measures[1]["confidence"]):
            reason = "tie"
        elif first["confidence"] <= cr.MAPPING_THRESHOLD:
            reason = "confidence_80"
    state = "you" if chosen["column"] == dec.get("revenue_row") else ("click" if reason else "auto")
    if _is_turnover(chosen["label"]):
        if money == "revenue":
            state, reason = "you", None
        elif money is None:
            state, reason = "click", "turnover_question"
    revenue_column = {"column": chosen["column"], "label": chosen["label"], "kind": chosen["kind"], "confidence": chosen["confidence"],
                      "state": state, "reason": reason, "candidates": [{"column": m["column"], "label": m["label"], "kind": m["kind"],
                                                                         "confidence": m["confidence"]} for m in measures]}
    idx = chosen["idx"]
    usable = [(n, r, mth) for n, r, mth in dated if _num(r[idx]) is not None]
    upto = [d for d in usable if as_of is None or d[2] <= as_of]
    candidates = [{"row": n, "month": mth} for n, _, mth in upto]
    last_row = upto[-1] if upto else (usable[-1] if usable else None)
    if as_of is None:
        last_state, last_reason = "click", "no_as_of_month"
    elif not upto:
        last_state, last_reason = "click", "after_as_of"
    else:
        last_state, last_reason = "auto", None
    if dec.get("last_actual") in [n for n, _, _ in usable] and not (as_of and next(m for n, _, m in usable if n == dec["last_actual"]) > as_of):
        last_row = next(d for d in usable if d[0] == dec["last_actual"])
        last_state, last_reason = "you", None
    last = {"row": last_row[0] if last_row else None, "month": last_row[2] if last_row else None, "state": last_state,
            "reason": last_reason, "candidates": candidates}
    limit = last["month"]                    # the analyst's last actual month ends the series too (review H4)
    series = [{"row": n, "month": mth, "grain": "month", "value": _round(_num(r[idx]))}
              for n, r, mth in (upto if as_of is not None else usable) if limit is None or mth <= limit]
    ftype = dec["file_type"] if dec.get("file_type") in RECOGNITION_TYPES else \
        ("financial_statements" if info["groups"] >= 2 else "monthly_revenue_report")
    supports, blocked = _supports("columns", ftype, len(info["months"]), False, False)
    result = {
        "recognised": True, "missing": [], "message": None, "sheet": info["sheet"].sheet, "header_row": info["header_row"],
        "file_type": ftype, "type_state": "auto", "type_source": "you" if dec.get("file_type") else "rules", "type_candidates": [],
        "type_reason": None, "label_mode": None, "label_share": None, "confidence": None, "series_basis": "column",
        "periods": len(info["months"]), "revenue_row": None, "revenue_column": revenue_column, "last_actual": last,
        "excluded_columns": [], "currency": _currency(info, ftype, ctx.get("reporting_currency"), dec.get("currency")),
        "money_kind": {"asked": _is_turnover(chosen["label"]), **({"answer": money} if _is_turnover(chosen["label"]) and money else {})},
        "non_revenue_rows": [], "supports": supports, "blocked": blocked, "other_sheets": [], "series": series, "total_rows": [],
        "series_rows": [x["row"] for x in series]}
    return _finish(result)


def _finalize_records(info: Dict, ctx: Dict) -> Dict:
    dec = ctx.get("decisions") or {}
    missing = [name for name, ok in (("revenue_measure", info["amount"] is not None),
                                     ("reference", info["date"] is not None or info["reference"] is not None),
                                     ("numbers", info["amount"] is not None)) if not ok]
    if missing:
        return _blank_result(info, missing, [])
    ftype = dec["file_type"] if dec.get("file_type") in RECOGNITION_TYPES else "invoice_transactions"
    supports, blocked = _supports("records", ftype, info["month_count"], info["customer"] is not None, False)
    result = {
        "recognised": True, "missing": [], "message": None, "sheet": info["sheet"].sheet, "header_row": info["header_row"],
        "file_type": ftype, "type_state": "auto", "type_source": "you" if dec.get("file_type") else "rules", "type_candidates": [],
        "type_reason": None, "label_mode": None, "label_share": None, "confidence": None, "series_basis": "records",
        "periods": info["month_count"], "revenue_row": None, "revenue_column": None, "last_actual": None, "excluded_columns": [],
        "currency": _currency(info, ftype, ctx.get("reporting_currency"), dec.get("currency")), "money_kind": {"asked": False},
        "non_revenue_rows": [], "supports": supports, "blocked": blocked, "other_sheets": [], "series": [], "total_rows": [],
        "series_rows": []}
    return _finish(result)


_FINALIZE = {"labels": _finalize_labels, "snapshot": _finalize_snapshot, "columns": _finalize_columns, "records": _finalize_records}


def _rank_of(result: Dict) -> int:
    row = result.get("revenue_row") or result.get("revenue_column")
    return _RANK.get((row or {}).get("kind"), 0)


def recognise(sheets, as_of_month: Optional[str] = None, proposal: Optional[Dict] = None, reporting_currency: Optional[str] = None,
              decisions: Optional[Dict] = None, fiscal_year_end: int = 12) -> Dict:
    """The recognition of a workbook (sections 2, 3 and 6): every sheet is read, the best recognised sheet is the result and the
    others are listed. `proposal` is the validated reply of the model (section 5.7, rows and columns as sheet positions, the revenue
    row with its label or `label_matches`); it is used only for the sheet Python chose. `decisions` are the analyst's choices."""
    as_of = _asof(as_of_month)
    infos = [analyse(s, fiscal_year_end) for s in sheets]
    base = {"as_of": as_of, "reporting_currency": reporting_currency}
    rules = [_FINALIZE[i["layout"]](i, base) for i in infos]
    recognised = [k for k, r in enumerate(rules) if r["recognised"]]
    dec = dict(decisions or {})
    chosen_name = dec.pop("sheet", None)
    if chosen_name is not None and any(r["sheet"] == chosen_name and r["recognised"] for r in rules):
        best = next(k for k, r in enumerate(rules) if r["sheet"] == chosen_name)
    elif recognised:
        best = min(recognised, key=lambda k: (-len(rules[k]["supports"]), _rank_of(rules[k]), -rules[k]["periods"], k))
    else:
        best = min(range(len(rules)), key=lambda k: (len(rules[k]["missing"]), k))
    use = proposal if proposal and proposal.get("sheet") == best + 1 else None
    result = _FINALIZE[infos[best]["layout"]](infos[best], {**base, "proposal": use, "decisions": dec})
    result["other_sheets"] = [{"sheet": rules[k]["sheet"], "file_type": rules[k]["file_type"], "recognised": True}
                              for k in recognised if k != best]
    return result


# ---------------------------------------------------------------------------
# The text the model reads (section 5)
# ---------------------------------------------------------------------------
def _sample_rows(columns: Dict[int, List[Tuple[int, str]]], file_hash: str) -> Dict[int, List[str]]:
    """Up to min(3, n ÷ 2) example values per column, from rows chosen by the file hash, none in common with the column before
    (section 5.1): no value is tied to a row and no row's series can be rebuilt from the text."""
    out: Dict[int, List[str]] = {}
    previous: Set[int] = set()
    for j, col in enumerate(sorted(columns), 1):
        found = columns[col]
        k = min(3, len(found) // 2)
        order = sorted(found, key=lambda rv: hashlib.sha256(f"{file_hash}:{j}:{rv[0]}".encode()).hexdigest())
        taken = [rv for rv in order if rv[0] not in previous][:k]
        previous = {r for r, _ in taken}
        out[col] = [v for _, v in sorted(taken)]
    return {c: v for c, v in out.items() if v}


def model_text(sheets, company_name: Optional[str], mapping: Dict[str, str], withheld: Tuple[str, ...], file_hash: str,
               fiscal_year_end: int = 12) -> str:
    """The redacted text the model reads for a file (section 5.2), built in memory and never stored or logged."""
    from . import structures as st
    out = []
    for sheet in sheets[:redact.FILE_MAX_SHEETS]:
        info = analyse(sheet, fiscal_year_end)
        grid = sheet.grid or []
        header = info["header_row"]
        limit, texted = len(grid), 0
        for r, row in enumerate(grid, 1):
            if any(sheet_mod.cell_kind(v, r, header) == "text" for v in row):
                texted += 1
                if texted == MAX_TEXT_ROWS:
                    limit = r
                    break
        keep = sorted({r for r, _, _ in sheet.titles if r <= limit})[-redact.FILE_MAX_TITLE_ROWS:]
        titles = [{"row": r, "col": c, "text": t[:st.FILE_TYPE_CELL_MAX]} for r, c, t in sheet.titles if r in keep]
        shown = sheet.raw_headers if sheet.raw_headers is not None else sheet.headers
        heads = [{"row": header, "col": c, "text": h[:st.FILE_TYPE_CELL_MAX]} for c, h in enumerate(shown[:redact.FILE_MAX_COLUMNS], 1) if h]
        labels, label_rows = None, []
        if info["layout"] == "labels":
            labels = "metric" if info["mode"] == "metric" else "masked"
            label_rows = [(ln["row"], ln["label"]) for ln in info["lines"] if ln["row"] <= limit]
        if labels == "masked":
            label_rows = [(r, f"Row {i}") for i, (r, _) in enumerate(label_rows, 1)]
        cells = redact.file_type_cells({"name": sheet.sheet[:st.FILE_TYPE_CELL_MAX], "titles": titles, "headers": heads}) + \
            [{"row": r, "col": 0, "text": t[:st.FILE_TYPE_CELL_MAX]} for r, t in label_rows]
        cells, _ = redact.redact_structure(cells, company_name, mapping, withheld)
        n_titles, n_heads = len(titles), len(heads)
        data = [r for r in range(header + 1, limit + 1)]
        customer_cols = {c for c, h in enumerate(shown, 1) if cr.field_score("revenue", "customer_id", h) > 0} \
            if info["layout"] == "records" else set()
        numeric: Dict[int, List[Tuple[int, str]]] = {}
        profiles: Dict[int, Dict] = {}
        for c in range(1, min(len(shown), redact.FILE_MAX_COLUMNS) + 1):
            values = [(r, _cell(grid, r, c)) for r in data if not cr._blank(_cell(grid, r, c))]
            if not values:
                continue
            as_samples = [(r, st._sample(v)) for r, v in values]
            hits = [(r, s) for r, s in as_samples if s]
            if len(hits) >= st.SAMPLE_SHARE * len(values) and c not in customer_cols:
                numeric[c] = hits
            else:
                profiles[c] = st._profile([v for _, v in values])
        out.append({"name": cells[0]["text"], "titles": cells[1:1 + n_titles], "headers": cells[1 + n_titles:1 + n_titles + n_heads],
                    "labels": labels, "label_rows": [(r["row"], r["text"]) for r in cells[1 + n_titles + n_heads:]],
                    "samples": _sample_rows(numeric, file_hash), "profiles": profiles})
    return redact.recognition_text(out)


# ---------------------------------------------------------------------------
# What the server stores and hands on (sections 5.5, 8 and 9): no I/O, closed words and numbers
# ---------------------------------------------------------------------------
NO_MAPPING_TYPES = ("monthly_revenue_report", "revenue_summary", "financial_statements", "arr_mrr_snapshot")
SAME_SOURCE = "Not applicable – same source"
CONFIRM_ONLY = ("confidence_80", "no_as_of_month")          # the reasons that wait for a confirmation only (section 7)
SERIES_RULES = {
    "row": "Revenue = the confirmed revenue row, one value per month up to the last actual month",
    "recurring": "Revenue series read from an MRR or ARR line, not computed from customers, one value per month up to the last actual month",
    "sum_of_rows": "Revenue = the sum of the customer or category rows, total rows left out, up to the last actual month",
    "column": "Revenue = the confirmed revenue column, one value per month up to the last actual month"}


def storable(value):
    """A cell as MongoDB can keep it: a number, text, a boolean, a datetime or nothing. A time of day or a duration (not BSON) becomes
    text or seconds, a date a datetime, a decimal or a numpy number a plain number, anything else text (review H8)."""
    if value is None or isinstance(value, (bool, str)):
        return value
    if isinstance(value, datetime):
        return value.to_pydatetime() if isinstance(value, pd.Timestamp) else value
    if isinstance(value, date):
        return datetime(value.year, value.month, value.day)
    if isinstance(value, time):
        return value.isoformat()
    if isinstance(value, timedelta):
        return value.total_seconds()
    if hasattr(value, "item") and not isinstance(value, (int, float)):
        try:
            return storable(value.item())
        except (ValueError, TypeError):
            return str(value)
    if isinstance(value, int):
        return value if -2 ** 63 <= value < 2 ** 63 else float(value)
    if isinstance(value, (float, Decimal)):
        return float(value)
    return str(value)


def pack_sheets(sheets) -> List[Dict]:
    """The sheets of a file as the dataset keeps them to read the file again: name, whether it was a CSV, every row (the caller refuses a dataset over the size limit; nothing is cut),
    each cell's indent and the merged spans. Every value is one MongoDB can store (review H8)."""
    out = []
    for s in sheets:
        fmt = s.fmt or {}
        out.append({"name": s.sheet, "csv": s.raw_headers is None and not fmt, "grid": [[storable(c) for c in r] for r in (s.grid or [])],
                    "indents": [[r, c, n] for (r, c), n in (fmt.get("indents") or {}).items()],
                    "merged": [list(m) for m in fmt.get("merged") or []]})
    return out


def kept_names(result: Dict, keep_rows: bool) -> List[str]:
    """The sheets a dataset keeps (owner decision 3 on PR 135): the recognised sheets only, and for a file that already keeps its rows
    (invoices, customers by period) the one sheet recognition read."""
    return [result["sheet"]] + ([] if keep_rows else [o["sheet"] for o in result.get("other_sheets") or []])


def renumber_proposal(proposal: Optional[Dict], sheets, names: List[str]) -> Optional[Dict]:
    """The model's kept proposal with its `sheet` counted among the kept sheets; one that named a sheet not kept names none (0)."""
    if not proposal:
        return proposal
    position = proposal.get("sheet")
    named = sheets[position - 1].sheet if isinstance(position, int) and 1 <= position <= len(sheets) else None
    return {**proposal, "sheet": names.index(named) + 1 if named in names else 0}


def records_facts(info: Dict) -> Dict:
    """The few facts of a records layout (one row per invoice) that its recognition needs: the dataset already keeps the rows, so the
    sheet is not kept a second time (owner decision 3 on PR 135)."""
    return {"sheet": info["sheet"].sheet, "header_row": info["header_row"], "amount": info["amount"], "date": info["date"],
            "reference": info["reference"], "customer": info["customer"], "month_count": info["month_count"],
            "multi_currency": info["multi_currency"], "currency_values": sorted(info["currency_values"]), "codes": sorted(info["codes"])}


def recognise_facts(facts: Dict, as_of_month: Optional[str] = None, reporting_currency: Optional[str] = None,
                    decisions: Optional[Dict] = None) -> Dict:
    """The recognition of a records layout from its stored facts: the same result as recognising the sheet."""
    info = {**facts, "layout": "records", "sheet": SimpleNamespace(sheet=facts["sheet"]), "periods": [], "titles": [],
            "currency_values": set(facts["currency_values"]), "codes": set(facts["codes"]), "rows": []}
    dec = {k: v for k, v in (decisions or {}).items() if k != "sheet"}
    result = _finalize_records(info, {"as_of": _asof(as_of_month), "reporting_currency": reporting_currency, "decisions": dec})
    result["other_sheets"] = []
    return result


def unpack_sheets(stored: List[Dict]):
    """The sheets of pack_sheets, read again by the same rules that read the file."""
    out = []
    for item in stored or []:
        fmt = {"indents": {(r, c): n for r, c, n in item.get("indents") or []}, "merged": [tuple(m) for m in item.get("merged") or []]} \
            if (item.get("indents") or item.get("merged")) else {}
        sheet = cr._read_grid([list(r) for r in item["grid"]], item["name"], bool(item.get("csv")), fmt)
        if sheet is not None:
            out.append(sheet)
    return out


def proposal_from(reply: Dict, sheets, company_name: Optional[str], mapping: Dict[str, str], withheld: Tuple[str, ...],
                  fiscal_year_end: int = 12) -> Optional[Dict]:
    """The proposal Python checks (section 6) from a validated reply (section 5.7): the label echo of `revenue_row` is compared with
    the label of that row (as written, or as the model was sent it) and dropped; only `label_matches` stays (section 5.5)."""
    if not reply or not 1 <= reply.get("sheet", 0) <= len(sheets):
        return None
    out = {k: reply.get(k) for k in ("sheet", "file_type", "last_actual_column", "currency")}
    out["non_revenue_rows"] = [{"row": x["row"], "kind": x["kind"]} for x in reply.get("non_revenue_rows") or []]
    row = reply.get("revenue_row")
    out["revenue_row"] = None
    if row:
        info = analyse(sheets[reply["sheet"] - 1], fiscal_year_end)
        label = next((ln["label"] for ln in info.get("lines") or [] if ln["row"] == row["row"]), None)
        sent = None
        if label is not None:
            cell, _ = redact.redact_structure([{"row": 0, "col": 0, "text": label}], company_name, mapping, withheld)
            sent = cell[0]["text"]
        matches = label is not None and (_label_matches(label, {"label": row["label"]}) or
                                         (sent is not None and _label_matches(sent, {"label": row["label"]})))
        out["revenue_row"] = {"row": row["row"], "label_matches": bool(matches)}
    return out


def has_monthly_series(result: Optional[Dict]) -> bool:
    """Whether a recognition gives a monthly revenue series. A snapshot, a yearly or quarterly report and a sheet with no revenue figure
    by month give none, and such a file is not the revenue file for the banner or Run audit (owner decision 1 on PR 135)."""
    return bool(result and result.get("recognised") and any(x.get("month") and x.get("grain") == "month" for x in result.get("series") or []))


def pnl_confirmed(result: Optional[Dict]) -> bool:
    """A P&L satisfies the required revenue file when its recognition is confirmed with two or more periods and gives a monthly series
    (section 8): no part in click and the revenue row or column decided by the analyst."""
    if not result or not result.get("recognised") or result.get("periods", 0) < 2 or result.get("needs_click") \
            or not has_monthly_series(result):
        return False
    row = result.get("revenue_row") or result.get("revenue_column")
    return bool(row and row["state"] == "you")


def engine_feed(result: Dict, file: str, sheet_name: str, unit_scale: int = 1, from_pnl: bool = False) -> Optional[Dict]:
    """What Run audit reads from a recognised file with no customer mapping (section 8): the revenue series by month up to the
    boundary, its rows for the citation and the closed words the contract carries (section 9). None for a records layout, which the
    column mapping reads."""
    if not result or not result.get("recognised") or result.get("series_basis") in (None, "records"):
        return None
    row = result.get("revenue_row") or result.get("revenue_column")
    waiting = (row and row["state"] == "click") or (result.get("last_actual") or {}).get("state") == "click" or \
        (result.get("money_kind") or {}).get("asked") and "answer" not in result["money_kind"]
    series = [] if waiting else [{"month": x["month"], "value": x["value"] * unit_scale} for x in result["series"]
                                  if x.get("month") and x.get("grain") == "month"]
    recurring = bool(row and row.get("kind") == "recurring")
    return {"file": file, "sheet": sheet_name, "basis": result["series_basis"], "file_type": result["file_type"],
            "recurring": recurring, "series": series, "rows": result.get("series_rows") or [],
            "rule": SERIES_RULES["recurring"] if recurring else SERIES_RULES.get(result["series_basis"]),
            "growth_label": growth_label(result), "periods": result["periods"],
            "excluded_columns": len(result.get("excluded_columns") or []),
            "currency_confirmed": result["currency"]["state"] in ("seen", "you", "audit"),
            "monthly": bool(series), "from_pnl": from_pnl,
            "series_source": (result.get("series_source") or {}).get("supplier"),
            "cross_check": ((result.get("series_source") or {}).get("cross_check") or {}).get("status"),
            "cross_check_months": ((result.get("series_source") or {}).get("cross_check") or {}).get("months"),
            "cross_check_gaps": [{**g, "row_value": g["row_value"] * unit_scale,
                                  "list_value": None if g["list_value"] is None else g["list_value"] * unit_scale}
                                 for g in (result.get("series_source") or {}).get("gaps") or []],
            "missing": [{"metric": ANALYSIS_NAMES[b["analysis"]], "reason": line, "unlocked_by": b["needs"]}
                        for b, line in zip(result.get("blocked") or [], missing_lines(result))]}
