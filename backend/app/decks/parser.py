"""A .pptx, .pdf or .docx file -> text blocks, each with its source reference.

A block is one text line, one speaker-notes line or one table cell:

    {"slide": 5, "kind": "text", "text": "...", "box": 12, "bbox": [0.1, 0.2, 0.4, 0.25], "title": True}
    {"page": 3, "kind": "table", "table": 1, "row": 2, "col": 4, "text": "...", "bbox": [...]}

The layout fields are what the claim detector borrows context from, where the format has
them: "box" groups the lines of one text box (pptx, docx) or of one stack of lines (pdf);
"bbox" is the line's position as fractions of the slide or page (pptx, pdf; a pptx line's
height is its share of its text box); "title" marks the slide title (pptx) or the topmost
text of a page (pdf).

A merged table cell keeps its span: "col_span" and "row_span" are set when they are above 1, and
"col" is the grid column the cell starts in.

Text only. Pictures, charts saved as pictures and scanned pages hold pixels, not text, so
they are not read; a file with no readable text is refused, never guessed.

parse_deck also lists the deck's structures (spec section 7, see detect_structures): tables,
charts read from the chart XML (pptx), KPI panels and roadmaps or timelines, each with its type,
its source reference and every cell with its row and column, so a value read from it can be cited.
"""
import io
import itertools
import re
import zipfile
from typing import Dict, List

MAX_BYTES = 50 * 1024 * 1024
MAX_PAGES = 200
MAX_UNPACKED_BYTES = 250 * 1024 * 1024      # pptx and docx are zip files
MAX_PARTS = 5000

NO_TEXT = ("No readable text found in this file. It may be scanned or made of images. "
           "Please upload a text-based version.")
TOO_LARGE = "This file is larger than 50 MB. Please upload a file of 50 MB or less."
TOO_MANY_PAGES = ("This file has {n} {unit}s. We read up to 200 slides or pages per file. "
                  "Please split it into smaller files.")
TOO_LARGE_UNPACKED = ("This file unpacks to more than 250 MB. We read files that unpack to 250 MB or less. "
                      "Please save a copy with fewer or smaller pictures and upload it again.")
TOO_MANY_PARTS = ("This file holds more than 5,000 parts. We read files with up to 5,000 parts. "
                  "Please save a simpler copy and upload it again.")
KEYNOTE = "We cannot read Keynote files. Please export it as PowerPoint or PDF first."
UNSUPPORTED = ("We read text from PowerPoint (.pptx), Word (.docx) and text-based PDF files. "
               "Please export this file as PowerPoint or PDF first.")
UNREADABLE = ("This file could not be opened as a {kind} file. It may be damaged or saved in another "
              "format. Please export it again as PowerPoint, Word or PDF.")

FORMATS = {".pptx": "pptx", ".pdf": "pdf", ".docx": "docx"}
_KIND_NAME = {"pptx": "PowerPoint", "pdf": "PDF", "docx": "Word"}


class DeckError(ValueError):
    """The file is refused; `message` is shown to the user as is."""

    def __init__(self, message: str):
        super().__init__(message)
        self.message = message


def deck_format(filename: str) -> str:
    """'pptx' | 'pdf' | 'docx' from the file name, or DeckError with the reason."""
    name = (filename or "").lower().strip()
    for ext, fmt in FORMATS.items():
        if name.endswith(ext):
            return fmt
    raise DeckError(KEYNOTE if name.endswith(".key") else UNSUPPORTED)


def parse_deck(content: bytes, filename: str) -> Dict:
    """{"file", "format", "page_unit", "pages", "blocks", "structures"} or DeckError."""
    fmt = deck_format(filename)
    if len(content) > MAX_BYTES:
        raise DeckError(TOO_LARGE)
    if fmt in ("pptx", "docx"):
        _check_unpacked(content, fmt)
    reader = {"pptx": _pptx_blocks, "pdf": _pdf_blocks, "docx": _docx_blocks}[fmt]
    charts = []
    try:
        if fmt == "pptx":
            pages, blocks, charts = reader(content)
        else:
            pages, blocks = reader(content)
    except DeckError:
        raise
    except Exception:
        # The library's message can quote file content, so it is neither shown nor logged.
        raise DeckError(UNREADABLE.format(kind=_KIND_NAME[fmt]))
    unit = "slide" if fmt == "pptx" else "page"
    if pages > MAX_PAGES:
        raise DeckError(TOO_MANY_PAGES.format(n=pages, unit=unit))
    if not any(re.search(r"\w", b["text"]) for b in blocks) and not charts:     # chart XML is text too
        raise DeckError(NO_TEXT)
    return {"file": filename, "format": fmt, "page_unit": unit, "pages": pages, "blocks": blocks,
            "structures": detect_structures(blocks, charts)}


def _check_unpacked(content: bytes, fmt: str) -> None:
    """Refuse a zip that would unpack too far before any library opens it. The sizes are the ones
    the zip declares; Python's zipfile never unpacks a part past its declared size."""
    try:
        parts = zipfile.ZipFile(io.BytesIO(content)).infolist()
    except (zipfile.BadZipFile, ValueError, OSError):
        raise DeckError(UNREADABLE.format(kind=_KIND_NAME[fmt]))
    if len(parts) > MAX_PARTS:
        raise DeckError(TOO_MANY_PARTS)
    if sum(part.file_size for part in parts) > MAX_UNPACKED_BYTES:
        raise DeckError(TOO_LARGE_UNPACKED)


def _clean(text: str) -> str:
    return re.sub(r"\s+", " ", text or "").strip()


def _lines(text: str) -> List[str]:
    """Split on line and paragraph breaks (pptx writes a line break as \\v); drop blank lines."""
    return [line for line in (_clean(part) for part in re.split(r"[\n\r\v]+", text or "")) if line]


# ---------------------------------------------------------------------------
# pptx: text boxes (including grouped shapes and placeholders), tables, speaker notes
# ---------------------------------------------------------------------------
def _pptx_blocks(content: bytes):
    from pptx import Presentation
    from pptx.enum.shapes import MSO_SHAPE_TYPE, PP_PLACEHOLDER

    prs = Presentation(io.BytesIO(content))
    slides = list(prs.slides)
    if len(slides) > MAX_PAGES:
        return len(slides), [], []
    width, height = prs.slide_width or 1, prs.slide_height or 1
    titles = (PP_PLACEHOLDER.TITLE, PP_PLACEHOLDER.CENTER_TITLE)
    boxes = itertools.count(1)

    def shapes(container, grouped=False):
        for shape in container:
            if shape.shape_type == MSO_SHAPE_TYPE.GROUP:
                yield from shapes(shape.shapes, True)
            else:
                yield shape, grouped

    def frac(x0, y0, x1, y1):
        return [round(x0 / width, 4), round(y0 / height, 4), round(x1 / width, 4), round(y1 / height, 4)]

    blocks, charts = [], []
    for n, slide in enumerate(slides, 1):
        tables = 0
        for shape, grouped in shapes(slide.shapes):
            if getattr(shape, "has_chart", False) and shape.has_chart:
                chart = _pptx_chart(shape.chart)
                if chart:
                    charts.append({"slide": n, **chart})
            # A grouped shape's position is in its group's own coordinates, so it has none here.
            geo = None if grouped else (shape.left, shape.top, shape.width, shape.height)
            if geo and None in geo:
                geo = None
            if shape.has_text_frame:
                box = next(boxes)
                title = shape.is_placeholder and shape.placeholder_format.type in titles
                paras = shape.text_frame.paragraphs
                for i, para in enumerate(paras):
                    for line in _lines(para.text):
                        block = {"slide": n, "kind": "text", "text": line, "box": box}
                        if geo:
                            left, top, w, h = geo
                            block["bbox"] = frac(left, top + h * i / len(paras), left + w, top + h * (i + 1) / len(paras))
                        if title:
                            block["title"] = True
                        blocks.append(block)
            if shape.has_table:
                tables += 1
                table = shape.table
                widths = [col.width for col in table.columns]
                heights = [row.height for row in table.rows]
                for r, row in enumerate(table.rows, 1):
                    for c, cell in enumerate(row.cells, 1):
                        if not _clean(cell.text) or cell.is_spanned:
                            continue
                        block = {"slide": n, "kind": "table", "table": tables, "row": r, "col": c, "text": _clean(cell.text)}
                        if cell.is_merge_origin:
                            block.update(_spans(cell.span_height, cell.span_width))
                        if geo and sum(widths) and sum(heights):
                            left, top, w, h = geo
                            sx, sy = w / sum(widths), h / sum(heights)
                            block["bbox"] = frac(left + sum(widths[:c - 1]) * sx, top + sum(heights[:r - 1]) * sy,
                                                 left + sum(widths[:c]) * sx, top + sum(heights[:r]) * sy)
                        blocks.append(block)
        if slide.has_notes_slide and slide.notes_slide.notes_text_frame is not None:
            box = next(boxes)
            for para in slide.notes_slide.notes_text_frame.paragraphs:
                blocks += [{"slide": n, "kind": "notes", "text": line, "box": box} for line in _lines(para.text)]
    return len(slides), blocks, charts


def _spans(rows: int, cols: int) -> Dict:
    """The span keys of a merged cell; none for a single cell."""
    out = {}
    if rows and rows > 1:
        out["row_span"] = rows
    if cols and cols > 1:
        out["col_span"] = cols
    return out


def _chart_number(value) -> str:
    """A chart value as text: 1200.0 -> "1200", 0.25 -> "0.25"."""
    if value is None:
        return ""
    return str(int(value)) if float(value).is_integer() else repr(float(value))


def _pptx_chart(chart):
    """A pptx chart from its XML: title, axis titles, series names, categories and the values the data
    labels show. None for a chart that cannot be read (no category plot)."""
    def title(owner):
        try:
            if owner.has_title and owner.chart_title.has_text_frame:
                return _clean(owner.chart_title.text_frame.text)
        except Exception:
            return ""
        return ""

    def axis_title(axis_name):
        try:
            axis = getattr(chart, axis_name)
            return _clean(axis.axis_title.text_frame.text) if axis.has_title else ""
        except Exception:
            return ""

    try:
        plot = chart.plots[0]
        categories = [_clean(str(c)) for c in plot.categories]
        series = [(_clean(s.name or ""), [_chart_number(v) for v in s.values]) for s in plot.series]
    except Exception:
        return None
    if not series or not categories:
        return None
    return {"title": title(chart), "category_title": axis_title("category_axis"),
            "value_title": axis_title("value_axis"), "categories": categories, "series": series}


# ---------------------------------------------------------------------------
# pdf: text lines, plus ruled tables read cell by cell
# ---------------------------------------------------------------------------
def _is_grid(rows) -> bool:
    """A real table has at least two rows with two filled cells. Slide frames and chart
    gridlines that pdfplumber also reports as tables do not."""
    return sum(1 for row in rows if sum(1 for cell in row if _clean(cell)) >= 2) >= 2


def _stack(blocks: List[Dict], boxes) -> None:
    """Give each pdf line a box: lines stacked under one another, close and overlapping
    sideways, are one box - a heading over its value, a year over its quarter, a bullet list."""
    open_boxes = []          # [box id, bbox of its lowest line]
    for block in sorted(blocks, key=lambda b: (b["bbox"][1], b["bbox"][0])):
        x0, y0, x1, y1 = block["bbox"]
        best = None
        for entry in open_boxes:
            lx0, ly0, lx1, ly1 = entry[1]
            gap = y0 - ly1
            if -0.25 * (y1 - y0) <= gap <= 0.6 * min(y1 - y0, ly1 - ly0) and min(x1, lx1) > max(x0, lx0):
                if best is None or gap < best[0]:
                    best = (gap, entry)
        if best:
            block["box"], best[1][1] = best[1][0], block["bbox"]
        else:
            block["box"] = next(boxes)
            open_boxes.append([block["box"], block["bbox"]])


def _pdf_blocks(content: bytes):
    """Text is read in segments: words on one line with no wide gap between them. A wide gap
    starts a new segment, so side-by-side columns stay apart (a label and its value, two
    roadmap columns) and keep their own position. Characters are taken in the order the file
    writes them, so a number stays whole even when text at other heights sits between its rows
    (axis labels 500, 600 ... were read as single digits when sorted by position alone)."""
    import pdfplumber

    blocks = []
    boxes = itertools.count(1)
    with pdfplumber.open(io.BytesIO(content)) as pdf:
        pages = len(pdf.pages)
        if pages > MAX_PAGES:
            return pages, []
        for n, page in enumerate(pdf.pages, 1):
            x0, top, x1, bottom = page.bbox
            width, height = (x1 - x0) or 1, (bottom - top) or 1

            def frac(a, b, c, d):
                return [round((a - x0) / width, 4), round((b - top) / height, 4),
                        round((c - x0) / width, 4), round((d - top) / height, 4)]

            grids = []
            for table in page.find_tables():
                rows = table.extract()
                if _is_grid(rows):
                    tx0, ttop, tx1, tbottom = table.bbox
                    grids.append(((max(tx0, x0), max(ttop, top), min(tx1, x1), min(tbottom, bottom)), rows, table))
            rest = page
            for bbox, _, _ in grids:
                rest = rest.outside_bbox(bbox)
            words = rest.extract_words(keep_blank_chars=True, x_tolerance_ratio=0.5, y_tolerance=3, use_text_flow=True)
            lines = [{"page": n, "kind": "text", "text": _clean(w["text"]),
                      "bbox": frac(w["x0"], w["top"], w["x1"], w["bottom"])}
                     for w in words if _clean(w["text"])]
            _stack(lines, boxes)
            if lines:
                first = min(lines, key=lambda b: b["bbox"][1])
                if first["bbox"][1] < 0.25:
                    for line in lines:
                        if line["box"] == first["box"]:
                            line["title"] = True
            blocks += lines
            for t, (_, rows, table) in enumerate(grids, 1):
                for r, row in enumerate(rows, 1):
                    cells = table.rows[r - 1].cells if r - 1 < len(table.rows) else []
                    for c, cell in enumerate(row, 1):
                        if _clean(cell):
                            block = {"page": n, "kind": "table", "table": t, "row": r, "col": c, "text": _clean(cell)}
                            # pdfplumber writes None for the grid cells a merged cell covers
                            covered = next((k for k, v in enumerate(row[c:]) if v is not None), len(row) - c)
                            block.update(_spans(1, 1 + covered))
                            if c - 1 < len(cells) and cells[c - 1]:
                                block["bbox"] = frac(*cells[c - 1])
                            blocks.append(block)
    return pages, blocks


# ---------------------------------------------------------------------------
# docx: paragraphs, text boxes and tables in document order
# ---------------------------------------------------------------------------
_W = "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}"
_MC_FALLBACK = "{http://schemas.openxmlformats.org/markup-compatibility/2006}Fallback"


def _docx_cell_span(tc) -> int:
    """How many grid columns a w:tc covers (w:gridSpan)."""
    pr = tc.find(_W + "tcPr")
    span = pr.find(_W + "gridSpan") if pr is not None else None
    try:
        return max(1, int(span.get(_W + "val"))) if span is not None else 1
    except (TypeError, ValueError):
        return 1


def _docx_grid_col(tr, tc) -> int:
    """The grid column a w:tc starts in: the spans of the cells before it, plus one."""
    col = 1
    for cell in tr:
        if cell is tc:
            return col
        if cell.tag == _W + "tc":
            col += _docx_cell_span(cell)
    return col


def _docx_vmerge(tc):
    """None, "restart" (a vertical merge starts here) or "continue" (the cell above covers this one)."""
    pr = tc.find(_W + "tcPr")
    merge = pr.find(_W + "vMerge") if pr is not None else None
    if merge is None:
        return None
    return "restart" if merge.get(_W + "val") == "restart" else "continue"


def _docx_row_span(rows, index: int, col: int, tc) -> int:
    """How many rows a w:tc covers: itself and the "continue" cells below it in its grid column."""
    if _docx_vmerge(tc) != "restart":
        return 1
    span = 1
    for tr in rows[index + 1:]:
        below = next((c for c in tr if c.tag == _W + "tc" and _docx_grid_col(tr, c) == col), None)
        if below is None or _docx_vmerge(below) != "continue":
            break
        span += 1
    return span


def _docx_blocks(content: bytes):
    """A .docx has no fixed pages, so the page number is counted from what Word wrote into the
    file: a hard page break, a rendered page break (w:lastRenderedPageBreak) and a section break
    that starts a new page. Signals with no text between them are one break, since Word writes a
    rendered break right after a hard one. A blank page between two pages of text is not counted."""
    import docx

    body = docx.Document(io.BytesIO(content)).element.body
    for fallback in list(body.iter(_MC_FALLBACK)):       # the same text box again, for old readers
        fallback.getparent().remove(fallback)
    section_types = []
    for sect in body.iter(_W + "sectPr"):
        start = sect.find(_W + "type")
        section_types.append(start.get(_W + "val") if start is not None else "nextPage")

    state = {"page": 1, "text_since_break": False, "section": 0}
    text_boxes = {}     # w:txbxContent element -> box id
    segments = {}       # (owner element, page) -> {"page", "kind", "parts", table ref...}, in order of appearance
    tables_on_page = {}

    def page_break():
        if state["text_since_break"]:
            state["page"] += 1
            state["text_since_break"] = False

    def owner_of(el):
        cell = para = None
        for anc in el.iterancestors():
            if para is None and anc.tag == _W + "p":
                para = anc
            if anc.tag == _W + "tc":
                cell = anc
                break
        return cell if cell is not None else para

    def segment(owner):
        key = (owner, state["page"])
        if key not in segments:
            seg = {"page": state["page"], "parts": []}
            if owner.tag == _W + "tc":
                row = owner.getparent()
                table = row.getparent()
                numbers = tables_on_page.setdefault(state["page"], {})
                rows = [c for c in table if c.tag == _W + "tr"]
                col = _docx_grid_col(row, owner)
                seg.update(kind="table", table=numbers.setdefault(table, len(numbers) + 1),
                           row=rows.index(row) + 1, col=col,
                           spans=_spans(_docx_row_span(rows, rows.index(row), col, owner), _docx_cell_span(owner)))
            else:
                seg["kind"] = "text"
                frame = next((a for a in owner.iterancestors() if a.tag == _W + "txbxContent"), None)
                if frame is not None:
                    seg["box"] = text_boxes.setdefault(frame, len(text_boxes) + 1)
            segments[key] = seg
        return segments[key]

    for top in body:
        for el in top.iter():
            tag = el.tag
            if tag == _W + "t" and el.text:
                owner = owner_of(el)
                if owner is not None:
                    segment(owner)["parts"].append(el.text)
                    state["text_since_break"] = state["text_since_break"] or bool(el.text.strip())
            elif tag == _W + "tab":
                owner = owner_of(el)
                if owner is not None:
                    segment(owner)["parts"].append("\t")
            elif tag in (_W + "br", _W + "cr"):
                if el.get(_W + "type") == "page":
                    page_break()
                else:
                    owner = owner_of(el)
                    if owner is not None:
                        segment(owner)["parts"].append("\n")
            elif tag == _W + "lastRenderedPageBreak":
                page_break()
            elif tag == _W + "p":
                owner = owner_of(el) if el is not top else None
                # A paragraph inside a table cell or a text box ends a line of its own.
                if owner is not None and owner.tag == _W + "tc":
                    segment(owner)["parts"].append("\n")
        # A section break sits in the last paragraph of its section; the next section's start
        # type says whether the following text begins on a new page.
        for sect in top.iter(_W + "sectPr"):
            if sect.getparent() is not None and sect.getparent().tag == _W + "pPr":
                state["section"] += 1
                following = section_types[state["section"]] if state["section"] < len(section_types) else "nextPage"
                if following != "continuous":
                    page_break()

    blocks = []
    for (owner, _), seg in segments.items():
        text = "".join(seg["parts"])
        if seg["kind"] == "table":
            if _clean(text):
                blocks.append({"page": seg["page"], "kind": "table", "table": seg["table"], "row": seg["row"],
                               "col": seg["col"], "text": _clean(text), **seg["spans"]})
        else:
            box = {"box": seg["box"]} if "box" in seg else {}
            blocks += [{"page": seg["page"], "kind": "text", "text": line, **box} for line in _lines(text)]
    return state["page"], blocks


# ---------------------------------------------------------------------------
# Structures (spec section 7): what CLAUDE.md rule 16 lets the gateway read. Python finds them and
# assigns their type; nothing here calls or imports the gateway.
# ---------------------------------------------------------------------------
STRUCTURE_TYPES = ("table", "chart", "kpi_panel", "roadmap", "hiring_table", "unit_economics", "use_of_funds")
# Fixed on the 10 test decks and written into deck-parser.md section 7. A table's type comes from
# the keywords of its header rows, its first column and its caption (the line right above it);
# when several match, the first type in this order wins.
TABLE_KEYWORDS = (
    ("use_of_funds", re.compile(r"(?i)\b(?:use of (?:funds|proceeds)|funds|proceeds)\b")),
    ("unit_economics", re.compile(r"\b(?:CAC|LTV|ARPU|ARPA|ACV)s?\b|(?i:\bpayback\b|\bunit economics\b|\bcontribution margin\b)")),
    ("hiring_table", re.compile(r"(?i)\b(?:hir(?:e|es|ing)|headcount|recruit(?:s|ed|ing|ment)?|roles?|positions?|FTEs?)\b")),
)
# How far above a table its caption may sit, as a share of the page; also how far a KPI panel's label box may sit
# directly above or below its box.
CAPTION_REACH = 0.1
KPI_LINE_MAX = 30           # a KPI box: every line at most this many characters...
KPI_MAX_LINES = 4           # ...at most this many lines, a figure and a word, and not a wrapped sentence
# The label length (decision of 2026-10-06): the longest label on the 10 test decks, "Crawling, Serving, Hosting +
# Processing" (39 characters), plus 50%, rounded up. A label box or a title used as a label, read as one line.
LABEL_MAX = 59
TIMELINE_MIN_DATES = 3      # date labels on a page (chart axes left out) that make it a roadmap, with a
TIMELINE_LINE_MAX = 60      # product keyword on the page; every line of a timeline box at most this long
CELL_MAX = 200              # a longer cell is prose and is left out of its structure
_WORD = re.compile(r"[^\W\d_]{2,}")
# A sentence wrapped over the lines of a box: a line that ends on one of these words or a comma, or
# a line after the first that starts with one ("We took one round of / financing in 2007").
_WRAP_WORDS = frozenset("""a an and are as at be by for from has have in into is it its of on or our than that the
their this to was we were which with""".split())
_LIST_NUMBER = re.compile(r"^\s*\d+\.\s+(?=\D)")     # "1. Data": a list number, not a figure


def detect_structures(blocks: List[Dict], charts: List[Dict] = ()) -> List[Dict]:
    """The deck's structures, in page order. Each is

        {"type": "table", "slide": 4 (or "page"), "header_rows": 1,
         "cells": [{"row": 1, "col": 2, "text": "FY2025", "col_span": 12}, ...]}

    with "table" (a table's number on its page), "chart" (a chart's) or, for a KPI panel or a
    roadmap, a "box" on every cell: the text box the line comes from; their cells also keep
    "next_to", and a roadmap's lines "date_box" (_band_cells). header_rows counts the rows
    above the first row that holds a figure. A structure holds at least one figure (a number or a
    date); text that is none of these structures is prose and is never one. Pages whose figures
    section 2 drops as background or cited research hold no structure.
    """
    from . import claims        # the period and figure rules are section 2's, shared with the claims

    excluded = claims.excluded_pages(blocks)
    found = _table_structures(blocks, excluded, claims) + _chart_structures(charts) + \
        _box_structures(blocks, excluded, claims)
    found.sort(key=lambda s: (s.get("slide") or s.get("page") or 0, STRUCTURE_ORDER.get(s["type"], 9)))
    return found


STRUCTURE_ORDER = {"table": 0, "hiring_table": 0, "unit_economics": 0, "use_of_funds": 0, "chart": 1,
                   "kpi_panel": 2, "roadmap": 2}


def _where(block: Dict) -> Dict:
    return {k: block[k] for k in ("slide", "page") if k in block}


def _page_key(block: Dict):
    return (block.get("slide"), block.get("page"))


def _cell(block: Dict) -> Dict:
    return {"row": block["row"], "col": block["col"], "text": block["text"],
            **{k: block[k] for k in ("row_span", "col_span") if k in block}}


def _figures(text: str, claims) -> List[Dict]:
    return claims.figures(_LIST_NUMBER.sub("", text))


def _has_figure(text: str, claims) -> bool:
    return bool(_figures(text, claims) or claims.find_dates(text))


def _header_rows(cells: List[Dict], claims) -> int:
    """Rows from the top until the first that holds a figure other than a date."""
    count = 0
    for r in sorted({c["row"] for c in cells}):
        if any(_figures(c["text"], claims) for c in cells if c["row"] == r):
            break
        count += 1
    return count


def _wrapped(box: List[Dict]) -> bool:
    """True when the lines of a box are one sentence wrapped, not separate labels. The article "a" counts only in
    lower case: "Seed to Series A" ends on a name."""
    words = [[w if w == "A" else w.lower() for w in re.findall(r"[^\W\d_]+|,", line["text"])] for line in box]
    return any(w and (w[-1] in _WRAP_WORDS or w[-1] == ",") for w in words[:-1]) or \
        any(w and w[0] in _WRAP_WORDS for w in words[1:])


_END_PUNCTUATION = re.compile(r"[.,;:!?/][\"'”’)\]]*$")      # "app.", "media,", "$499 /", "“Moz.com.”"


def _continues(above: str, line: str) -> bool:
    """Whether a roadmap line continues the line above it (spec section 7, issues #49 and #55): it starts with a
    lower-case letter or "&", or the line above ends with punctuation, a wrap word ("a" in lower case only; a line
    ending on a figure, "ongoing in 2021", ends on no word) or inside a bracket it opened ("Gillian (Rand’s / Mom)
    founds the")."""
    start = line.lstrip()[:1]
    tokens = [w if w == "A" else w.lower() for w in re.findall(r"[^\W_]+", above)]
    return start == "&" or (start.isalpha() and start.islower()) or bool(_END_PUNCTUATION.search(above.rstrip())) \
        or bool(tokens and tokens[-1] in _WRAP_WORDS) or above.count("(") > above.count(")")


def _paragraph(box: List[Dict], claims) -> List[Dict]:
    """A roadmap text box as its grid rows (spec section 7, issues #49 and #55): its lines split into items at each
    line that does not continue the one above, each item one row holding its lines joined with a space. A date label
    is an item of its own, never joined. An item whose joined text would be over CELL_MAX characters keeps a row per
    line. A paragraph is the one-item box; a bullet list gives one row per bullet, a wrapped bullet included."""
    items = []
    for line in box:
        dated = bool(_date_labels(line["text"], claims))
        if items and not dated and not items[-1][0] and _continues(items[-1][1][-1]["text"], line["text"]):
            items[-1][1].append(line)
        else:
            items.append((dated, [line]))
    out = []
    for _, lines in items:
        joined = " ".join(line["text"] for line in lines)
        if len(lines) < 2 or len(joined) > CELL_MAX:
            out += lines
            continue
        extent = _extent(lines)
        out.append({**lines[0], "text": joined, **({"bbox": list(extent)} if extent else {})})
    return out


def _caption(blocks: List[Dict], page, cells: List[Dict]) -> str:
    """The text line right above a table, when the layout says where the table is."""
    boxes = [c["bbox"] for c in cells if c.get("bbox")]
    if not boxes:
        return ""
    x0, top, x1 = min(b[0] for b in boxes), min(b[1] for b in boxes), max(b[2] for b in boxes)
    above = [b for b in blocks if b["kind"] == "text" and _page_key(b) == page and b.get("bbox")
             and b["bbox"][3] <= top + 1e-6 and top - b["bbox"][3] <= CAPTION_REACH
             and min(x1, b["bbox"][2]) > max(x0, b["bbox"][0])]
    return max(above, key=lambda b: b["bbox"][3])["text"] if above else ""


def table_type(text: str) -> str:
    for kind, rx in TABLE_KEYWORDS:
        if rx.search(text):
            return kind
    return "table"


def _table_structures(blocks: List[Dict], excluded: set, claims) -> List[Dict]:
    tables = {}
    for b in blocks:
        if b["kind"] == "table":
            tables.setdefault((_page_key(b), b["table"]), []).append(b)
    ticks = claims.tick_cells(blocks)
    out = []
    for (page, number), members in tables.items():
        if page in excluded:
            continue
        cells = [b for b in members if len(b["text"]) <= CELL_MAX]
        where = page[0] or page[1]
        if not any(_has_figure(c["text"], claims) for c in cells if (where, number, c["row"], c["col"]) not in ticks):
            continue
        header_rows = _header_rows(cells, claims)
        first_row, first_col = min(c["row"] for c in cells), min(c["col"] for c in cells)
        heads = [c["text"] for c in cells if c["row"] < first_row + max(header_rows, 1) or c["col"] == first_col]
        out.append({"type": table_type(" ".join(heads + [_caption(blocks, page, members)])), **_where(members[0]),
                    "table": number, "header_rows": header_rows, "cells": [_cell(c) for c in cells]})
    return out


def _chart_structures(charts: List[Dict]) -> List[Dict]:
    """A chart as a grid: its title, the value axis title over the series, the series names, then one
    row per category with the value of each series."""
    out = []
    per_slide = {}
    for chart in charts:
        n = len(chart["series"])
        rows, cells = 0, []
        if chart["title"]:
            rows += 1
            cells.append({"row": rows, "col": 1, "text": chart["title"], **_spans(1, n + 1)})
        if chart["value_title"]:
            rows += 1
            cells.append({"row": rows, "col": 2, "text": chart["value_title"], **_spans(1, n)})
        rows += 1
        if chart["category_title"]:
            cells.append({"row": rows, "col": 1, "text": chart["category_title"]})
        cells += [{"row": rows, "col": k + 2, "text": name} for k, (name, _) in enumerate(chart["series"]) if name]
        header_rows = rows
        for i, category in enumerate(chart["categories"]):
            rows += 1
            if category:
                cells.append({"row": rows, "col": 1, "text": category})
            for k, (_, values) in enumerate(chart["series"]):
                if i < len(values) and values[i]:
                    cells.append({"row": rows, "col": k + 2, "text": values[i]})
        cells = [c for c in cells if c["text"] and len(c["text"]) <= CELL_MAX]
        per_slide[chart["slide"]] = per_slide.get(chart["slide"], 0) + 1
        out.append({"type": "chart", "slide": chart["slide"], "chart": per_slide[chart["slide"]],
                    "header_rows": header_rows, "cells": cells})
    return out


def dates_only(texts: List[str], claims) -> bool:
    """A date box's lines (spec section 7, decision of 2026-10-06 on issue #47): each is a date label or a part of
    one ("2022" / "Q2", "Nov. 2007")."""
    return bool(texts) and all(_date_labels(text, claims) or claims.period_cell(text) for text in texts)


def _date_labels(text: str, claims) -> List[Dict]:
    """The dates of a line that is a date label: dates and at most two other words ("Nov. 2007",
    "Q1 17 Q2 17", "Launch Q3 2024"); [] for a sentence that holds a date."""
    dates = claims.find_dates(text)
    rest = text
    for d in reversed(dates):
        rest = rest[:d["start"]] + " " + rest[d["end"]:]
    return dates if dates and len(_WORD.findall(rest)) <= 2 else []


def _month_index(date_label: str, claims) -> int:
    start = claims.period_range(date_label)[0]
    return int(start[:4]) * 12 + int(start[5:7])


def _axis(labels: List[Dict], claims) -> bool:
    """Three or more distinct dates, evenly spaced, none repeated: the axis of a chart drawn as text."""
    months = [_month_index(d["date"], claims) for d in labels]
    if len(months) < 3 or len(set(months)) != len(months):
        return False
    months.sort()
    steps = {b - a for a, b in zip(months, months[1:])}
    return len(steps) == 1


def _axis_lines(lines: List[Dict], claims) -> set:
    """id() of the date-label lines that are chart axis ticks: a line of three or more evenly spaced
    dates, or such dates one per line in a row or a column."""
    labelled = [(line, _date_labels(line["text"], claims)) for line in lines]
    labelled = [(line, dates) for line, dates in labelled if dates]
    out = {id(line) for line, dates in labelled if _axis(dates, claims)}
    placed = [(line, dates) for line, dates in labelled if line.get("bbox")]
    for lo, hi in ((0, 2), (1, 3)):            # a column (overlapping x), then a row (overlapping y)
        for line, _ in placed:
            group = [(other, dates) for other, dates in placed
                     if min(line["bbox"][hi], other["bbox"][hi]) > max(line["bbox"][lo], other["bbox"][lo])]
            if len(group) > 1 and _axis([d for _, dates in group for d in dates], claims):
                out.update(id(other) for other, _ in group)
    return out


def _band_cells(boxes: List[List[Dict]], others=None, date_box=None) -> List[Dict]:
    """Text boxes as a grid: boxes that overlap in height form a band of rows, each box a column of
    its band in left-to-right order, each line a row. Every cell keeps its box; a title line is marked.
    With `others` (a KPI panel or a roadmap: the extents of every text box of its page and its title), a cell whose
    line has a position also keeps "next_to": the ids of the cells of other boxes whose line is directly next to
    its own line (_next_to, line to line). A tall box can put two visual rows in one band (front-b p15), so
    a cell's row neighbour need not sit next to it on the page. With `date_box` (a roadmap: whether a box is a date
    box), each line of a box that is none also keeps "date_box": the number of the one date box directly next to its
    box, box to box (_next_to), when exactly one is."""
    def extent(lines):
        placed = [l["bbox"] for l in lines if l.get("bbox")]
        return (min(b[1] for b in placed), max(b[3] for b in placed), min(b[0] for b in placed)) if placed else None

    placed = sorted((b for b in boxes if extent(b)), key=lambda b: (extent(b)[0], extent(b)[2]))
    bands = []
    for box in placed:
        top, bottom, _ = extent(box)
        if bands and top < bands[-1]["bottom"] and bottom > bands[-1]["top"]:
            bands[-1]["boxes"].append(box)
            bands[-1]["bottom"] = max(bands[-1]["bottom"], bottom)
        else:
            bands.append({"top": top, "bottom": bottom, "boxes": [box]})
    bands += [{"boxes": [box]} for box in boxes if not extent(box)]       # no layout: one box per band
    cells, lines, row, number, numbered = [], [], 0, 0, []
    for band in bands:
        members = sorted(band["boxes"], key=lambda b: extent(b)[2] if extent(b) else 0)
        for col, box in enumerate(members, 1):
            number += 1
            numbered.append((number, box))
            cells += [{"row": row + i + 1, "col": col, "text": line["text"], "box": number,
                       **({"title": True} if line.get("title") else {})} for i, line in enumerate(box)]
            lines += [line.get("bbox") for line in box]
        row += max(len(box) for box in members)
    if others is not None:
        for cell, bbox in zip(cells, lines):
            if bbox:
                cell["next_to"] = [f"r{o['row']}c{o['col']}" for o, near in zip(cells, lines) if near and
                                   o["box"] != cell["box"] and _next_to(tuple(bbox), tuple(near), others)]
    if others is not None and date_box is not None:
        dated = [(n, _extent(box)) for n, box in numbered if _extent(box) and date_box(box)]
        for n, box in numbered:
            near = [d for d, e in dated if _extent(box) and not date_box(box) and _next_to(_extent(box), e, others)]
            if len(near) == 1:
                for cell in cells:
                    if cell["box"] == n:
                        cell["date_box"] = near[0]
    return cells


def _extent(box: List[Dict]):
    """(left, top, right, bottom) of a text box, or None when its lines have no position."""
    placed = [l["bbox"] for l in box if l.get("bbox")]
    return (min(b[0] for b in placed), min(b[1] for b in placed), max(b[2] for b in placed),
            max(b[3] for b in placed)) if placed else None


def _overlap(a, b, axis: int) -> bool:
    """Whether two extents overlap along x (axis 0) or y (axis 1)."""
    return min(a[axis + 2], b[axis + 2]) > max(a[axis], b[axis])


def _next_to(a, b, others) -> bool:
    """Two box extents directly next to each other: in the same band (they overlap in height) or directly above or
    below (they overlap in width, at most CAPTION_REACH apart), with no other box between them."""
    for along, across in ((0, 1), (1, 0)):            # side by side, then one above the other
        if not _overlap(a, b, across) or _overlap(a, b, along):
            continue
        first, second = (a, b) if a[along + 2] <= b[along] else (b, a)
        if along == 1 and second[1] - first[3] > CAPTION_REACH:
            continue
        if not any(c[along] >= first[along + 2] - 1e-9 and c[along + 2] <= second[along] + 1e-9
                   and _overlap(c, a, across) and _overlap(c, b, across) for c in others if c is not a and c is not b):
            return True
    return False


def _kpi_box(box: List[Dict], claims) -> bool:
    return len(box) <= KPI_MAX_LINES and all(len(l["text"]) <= KPI_LINE_MAX for l in box) \
        and any(_figures(l["text"], claims) for l in box) and any(_WORD.search(l["text"]) for l in box) \
        and not _wrapped(box)


def _value_box(box: List[Dict], ticks: set, claims) -> bool:
    """One figure that is not a date and no word outside it ("~13,500", "$12 -$13 million"); never an axis tick."""
    if any(id(l) in ticks for l in box):
        return False
    texts = [_LIST_NUMBER.sub("", l["text"]) for l in box]
    found = [(i, f) for i, text in enumerate(texts) for f in claims.figures(text)]
    if len(found) != 1:
        return False
    (i, figure), = found
    texts[i] = texts[i][:figure["start"]] + " " + texts[i][figure["end"]:]
    return not any(_WORD.search(text) for text in texts)


def _label_box(box: List[Dict], claims) -> bool:
    """A word and no figure other than a date, at most LABEL_MAX characters read as one line (a label wrapped over
    two lines counts; a longer wrapped sentence is prose)."""
    return len(box) <= KPI_MAX_LINES and not any(_figures(l["text"], claims) for l in box) \
        and any(_WORD.search(l["text"]) for l in box) and len(" ".join(l["text"] for l in box)) <= LABEL_MAX


def _kpi_panel(boxes: List[List[Dict]], title: List[Dict], ticks: set, claims) -> List[List[Dict]]:
    """The boxes of a page's KPI panel (spec section 7): its KPI boxes, the value boxes a label box sits directly next
    to, the label boxes directly next to a KPI or value box, and the title line as the label of a value box with no
    label box of its own, beside it or above it. A page with no KPI box has no panel."""
    kpi = [box for box in boxes if _kpi_box(box, claims)]
    if not kpi:
        return []
    placed = {id(box): _extent(box) for box in boxes + ([title] if title else []) if _extent(box)}
    others = list(placed.values())

    def near(a, b):
        return id(a) in placed and id(b) in placed and _next_to(placed[id(a)], placed[id(b)], others)

    labels = [box for box in boxes if _label_box(box, claims)]
    values = [box for box in boxes if _value_box(box, ticks, claims)]
    labelled = [box for box in values if any(near(label, box) for label in labels)]
    labels = [label for label in labels if any(near(label, box) for box in kpi + labelled)]
    titled = [box for box in values if title and box not in labelled and _label_box(title, claims)
              and near(title, box) and placed[id(title)][1] < placed[id(box)][3]]
    kept = kpi + [box for box in labelled + titled if box not in kpi] + labels
    return kept + ([title] if titled else [])


def _box_structures(blocks: List[Dict], excluded: set, claims) -> List[Dict]:
    """Per page: a roadmap or timeline when TIMELINE_MIN_DATES date labels remain once chart axes are
    left out, holding every box that is not prose, a paragraph as one line (_paragraph); otherwise a KPI
    panel (_kpi_panel)."""
    pages, titles = {}, {}
    for b in blocks:
        if b["kind"] == "text" and _page_key(b) not in excluded:
            if b.get("title"):
                titles.setdefault(_page_key(b), []).append(b)
                continue
            key = b.get("box") if b.get("box") is not None else ("line", len(pages.get(_page_key(b), {})))
            pages.setdefault(_page_key(b), {}).setdefault(key, []).append(b)
    ticks = claims.tick_lines(blocks)
    out = []
    for page, boxes in pages.items():
        lines = [line for box in boxes.values() for line in box]
        axis = _axis_lines(lines, claims)
        labels = [line for line in lines if id(line) not in axis and _date_labels(line["text"], claims)]
        titled = [t["text"] for t in titles.get(page, ())]
        if len(labels) >= TIMELINE_MIN_DATES and any(claims.has_product_keyword(t) for t in titled + [l["text"] for l in lines]):
            kind = "roadmap"
            kept = [_paragraph(box, claims) for box in boxes.values()
                    if all(len(l["text"]) <= TIMELINE_LINE_MAX for l in box)
                    and not all(id(l) in axis for l in box)]
        else:
            kind = "kpi_panel"
            kept = _kpi_panel(list(boxes.values()), titles.get(page), ticks, claims)
        if not kept or not any(_has_figure(l["text"], claims) for box in kept for l in box):
            continue
        first = kept[0][0]
        page_boxes = list(boxes.values()) + ([titles[page]] if page in titles else [])
        others = [e for e in map(_extent, page_boxes) if e]
        dated = (lambda box: dates_only([line["text"] for line in box], claims)) if kind == "roadmap" else None
        out.append({"type": kind, **_where(first), "header_rows": 0, "cells": _band_cells(kept, others, dated)})
    return out
